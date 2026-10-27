# -*- coding: utf-8 -*-
"""Estudio de gestión de riesgos SM-107 (CUI 2673852) tal como se presentó, con tres anexos agregados al final:
Anexo 3 matriz de gestión de riesgos impresa, Anexo 4 ubicación georreferenciada de los puntos críticos y de los riesgos,
Anexo 5 panel fotográfico georreferenciado. El contenido presentado no se modifica.
Uso: python3 anexos_estudio_docx.py   (antes: fotos_sm107.py)"""
import os, sys
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import estudio_docx as ED
import levantamiento_obs_docx as LO
from estudio_docx import RS, par, vin, seccion

ORIG = os.path.join(os.path.dirname(RS.ENT), "insumos", "estudio_presentado", "ESTUDIO_GESTION_RIESGOS_EJECUCION_OBRA_SM107.docx")
SAL = os.path.join(RS.ENT, "ESTUDIO_GESTION_RIESGOS_EJECUCION_OBRA_SM107_CON_ANEXOS.docx")
T04, T08, T08A = "Tabla N.° 15", "Tabla N.° 16", "Tabla N.° 17"     # el estudio presentado llega a la Tabla N.° 14


def main():
    doc = Document(ORIG)
    n0 = len(doc.element.body)
    LO.T.update(t04=T04, t08=T08, t08a=T08A, panel="Anexo 5")

    # ------------------------------------------------------------------ anexo 3 (sigue en la sección horizontal de los anexos)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    ED.tit(doc, "13.3", "Anexo 3: Matriz de gestión de riesgos en la ejecución de la obra", 2)
    LO.anexo_matriz(doc, pref="A3", titulo=None)

    # ------------------------------------------------------------------ anexo 4
    s = seccion(doc, False); s.left_margin = s.right_margin = Cm(2.0)
    ED.tit(doc, "13.4", "Anexo 4: Ubicación georreferenciada de los puntos críticos y de los riesgos", 2)
    par(doc, "Complementa la Tabla N.° 04 (puntos y sectores críticos del inventario vial) y la Tabla N.° 08 (riesgos identificados) con las coordenadas UTM "
             "y la fotografía georreferenciada de cada ubicación (Anexo 5). Criterios:")
    for t in ["Fotografías: panel fotográfico del estudio de inventario vial, tomadas del 25 al 28 de setiembre de 2026 con cámara que imprime la fecha, la hora "
              "y las coordenadas UTM (WGS 84, zona 18 S) en cada imagen.",
              "Progresiva de cada fotografía: punto más cercano del eje georreferenciado del estudio (plano clave PC-01); %d de las %d fotografías están a menos de "
              "20 m del eje." % (sum(x["dist"] < 20 for x in LO.FOTOS), len(LO.FOTOS)),
              "Coordenadas de cada punto o sector: punto del eje en la progresiva de inicio y de fin.",
              "Puntos sin fotografía propia: se registran con la fotografía georreferenciada más cercana del mismo tramo, a no más de %d m, como vista del sector; "
              "se indica la distancia." % LO.RADIO,
              "Puntos sin fotografía a menos de %d m (%d): registro documental del estudio de inventario vial, ubicado con las coordenadas del eje; su registro "
              "fotográfico forma parte de la verificación de campo previa a la convocatoria." % (LO.RADIO, len(LO.VERIF)),
              "Los riesgos de gestión, contractuales, económicos, de permisos o de seguridad que abarcan todo el proyecto no tienen un punto físico y se indican "
              "como «No localizado»."]: vin(doc, t)
    LO.tabla_04(doc, T04 + ": Puntos y sectores críticos del inventario vial - coordenadas UTM (WGS 84 - zona 18 S) y fotografía")
    LO.fichas_04(doc, "Fichas de los puntos y sectores críticos de la Tabla N.° 04")
    LO.tabla_08(doc, T08 + ": Riesgos identificados - coordenadas y fotografía")
    LO.tabla_08a(doc, T08A + ": Ubicación georreferenciada de los riesgos (UTM WGS 84 - zona 18 S)")
    LO.fichas_08(doc, "Fichas de los riesgos identificados de la Tabla N.° 08")
    LO.accion_verificacion(doc)

    # ------------------------------------------------------------------ anexo 5
    s = seccion(doc, False); s.left_margin = s.right_margin = Cm(2.0)
    ED.tit(doc, "13.5", "Anexo 5: Panel fotográfico georreferenciado", 2)
    LO.panel(doc)

    doc.save(SAL)
    # el contenido presentado queda igual: se comparan los elementos del cuerpo anteriores a los anexos agregados
    o, d = Document(ORIG).element.body, Document(SAL).element.body
    iguales = all(a.xml == b.xml for a, b in zip(list(o)[:n0 - 1], list(d)[:n0 - 1]))
    print(SAL, "| original intacto:", iguales, "| fotos", len(LO.FOTOS))


if __name__ == "__main__":
    main()
