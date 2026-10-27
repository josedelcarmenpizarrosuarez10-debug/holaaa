# -*- coding: utf-8 -*-
"""Word del LEVANTAMIENTO DE OBSERVACIONES al estudio de gestión de riesgos en la ejecución de la obra - SM-107 (CUI 2673852).
Observación 1: matriz de gestión de riesgos impresa en el Anexo 2. Observaciones 2 y 3: tablas N.° 08 y N.° 04 con coordenadas UTM
y panel fotográfico georreferenciado (Anexo 3).
Uso: python3 levantamiento_obs_docx.py   (antes: riesgos_sm107.py, figuras_sm107.py y fotos_sm107.py)"""
import os, sys, collections
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import estudio_docx as ED
import fotos_sm107 as FS
from estudio_docx import RS, P, R, MC, LL, FIG, fuente, par, vin, mixto, tabla, sombra, seccion, pie, AZUL, C, L_, NIVCOL
from fotos_sm107 import km, punto

SAL = os.path.join(RS.ENT, "LEVANTAMIENTO_OBSERVACIONES_GESTION_RIESGOS_SM107.docx")
FOTOS = FS.fotos(); FID = {x["img"]: x["id"] for x in FOTOS}
pc = lambda v: "%.2f %%" % (v / P["PO"] * 100)

# ------------------------------------------------------------------ puntos y sectores críticos del inventario vial (Tabla N.° 04)
# progresiva (texto), inicio (m), fin (m o None), descripción, fotos (imagen del inventario)
SECTORES = [
    ("20+337 a 20+340", 20337, 20340, "Erosión h = 2.00 m (LI)", []),
    ("20+360 a 21+170", 20360, 21170, "Inundación por cultivos de arroz", []),
    ("22+270", 22270, None, "Empozamiento de agua (LI)", []),
    ("22+800 a 22+960", 22800, 22960, "Empozamiento de agua (LI)", []),
    ("26+924", 26924, None, "Empozamiento sin elemento de descarga", []),
    ("28+800 a 29+060", 28800, 29060, "Material en la vía (LI)", []),
    ("34+766 a 34+786", 34766, 34786, "Erosión h = 3 m (LD)", [47, 48]),
    ("34+954", 34954, None, "Empozamiento sin elemento de descarga", []),
    ("36+060 a 36+100", 36060, 36100, "Inestabilidad de talud h = 20 m (LI)", []),
    ("36+380 a 36+460", 36380, 36460, "Erosión h = 2.0 m (LD)", []),
    ("37+520 a 37+780", 37520, 37780, "Inestabilidad de talud h = 30 m (LI)", [51]),
    ("39+380 a 39+840", 39380, 39840, "Erosión h = 3 a 5 m (LD y LI)", [53]),
    ("40+380", 40380, None, "Empozamiento sin elemento de descarga", []),
    ("42+280", 42280, None, "Quebrada sin descarga, pendiente y contrapendiente, material de arrastre", [55]),
    ("43+620 a 43+632", 43620, 43632, "Erosión h = 1.0 m (LI)", []),
    ("43+660", 43660, None, "Empozamiento sin elemento de descarga", []),
    ("34+400 a 44+600 (varios)", 34400, 44600, "Plataforma angosta", [47, 50, 55]),
]

# ------------------------------------------------------------------ ubicación georreferenciada de los riesgos (Tabla N.° 08-A)
LOC = "Localidades del tramo (ver R-10)"
UBIC = [
    ("R-01", "Todo el tramo: empozamientos y pérdida de material por lluvia", 19343, 44693, [48, 52, 57]),
    ("R-02", "Quebrada del puente Tingo de Ponasa", 20300, None, [42]),
    ("R-02", "Sectores bajos e inundables", 19343, 23000, [12]),
    ("R-02", "Quebrada sin descarga", 42280, None, [55]),
    ("R-03", "Puente Tingo de Ponasa", 20300, None, [42]),
    ("R-03", "Puente Barco", 29300, None, []),
    ("R-03", "Puente Paucarillo", 37870, None, []),
    ("R-03", "Puente Apangurayco", 38800, None, []),
    ("R-03", "Puente Chanchamayco", 44040, None, []),
    ("R-04", "Talud inestable h = 20 m (LI)", 36060, 36100, []),
    ("R-04", "Talud inestable h = 30 m (LI)", 37520, 37780, [51]),
    ("R-04", "Laderas con cortes", 34000, 37000, [47, 50]),
    ("R-04", "Laderas con cortes", 43000, 44000, []),
    ("R-05", "Erosión h = 2.00 m (LI)", 20337, 20340, []),
    ("R-05", "Erosión h = 3 m (LD)", 34766, 34786, [47, 48]),
    ("R-05", "Erosión del borde de la plataforma (LD)", 37223, None, [49, 50]),
    ("R-05", "Erosión h = 2.0 m (LD)", 36380, 36460, []),
    ("R-05", "Erosión h = 3 a 5 m (LD y LI)", 39380, 39840, [53]),
    ("R-05", "Erosión h = 1.0 m (LI)", 43620, 43632, []),
    ("R-06", "Alcantarilla obstruida en Tingo de Ponasa", 19393, None, [12]),
    ("R-06", "Inundación por cultivos de arroz", 20360, 21170, []),
    ("R-06", "Empozamiento de agua (LI)", 22270, None, []),
    ("R-06", "Empozamiento de agua (LI)", 22800, 22960, []),
    ("R-06", "Empozamiento sin descarga", 26924, None, []),
    ("R-06", "Empozamiento sin descarga", 34954, None, []),
    ("R-06", "Empozamientos sobre la calzada", 34733, None, [48]),
    ("R-06", "Empozamientos sobre la calzada", 37083, 38013, [49, 51, 52]),
    ("R-06", "Empozamiento sin descarga", 40380, None, []),
    ("R-06", "Empozamientos sobre la calzada", 40763, 42993, [54, 57, 56]),
    ("R-06", "Empozamiento sin descarga", 43660, None, []),
    ("R-07", "Sectores en afirmado y terreno natural", 34400, 44600, [47, 49]),
    ("R-10", "Tingo de Ponasa", 19343, 19800, [6, 7, 41]),
    ("R-10", "Aypena", 21940, None, [15]),
    ("R-10", "Nuevo América", 26270, None, [17]),
    ("R-10", "Leoncio Prado", 28080, None, [18]),
    ("R-10", "San Antonio", 31420, None, [19, 40]),
    ("R-10", "Huañipo", 33780, None, [20]),
    ("R-10", "Paucar", 38020, None, [22]),
    ("R-10", "Alfonso Ugarte", 38730, None, [23]),
    ("R-10", "Shamboyacu", 44300, 44693, [43, 46]),
    ("R-11", "Pase de agua potable", 22425, None, []),
    ("R-11", "Tubería PVC 21\"", 25715, None, []),
    ("R-11", "Tubería PVC 6\"", 34688, None, [47]),
    ("R-11", "Postes en las localidades", 33780, None, [20, 46]),
    ("R-12", "Canal de riego paralelo a la vía", 31100, 31920, []),
    ("R-12", "Cruce de canal", 32578, None, []),
    ("R-12", "Cruce de canal", 35950, None, []),
    ("R-12", "Cruce de canal", 36120, None, []),
    ("R-12", "Cruce de canal", 37400, None, []),
    ("R-12", "Cruce de canal", 37794, None, []),
    ("R-12", "Cruce de canal", 38370, None, []),
    ("R-13", LOC, 19343, 44693, [6, 15, 17, 18, 19, 20, 22, 23, 46]),
    ("R-20", "Tránsito en localidades e intersecciones", 19343, 44693, [41, 14, 16, 20, 46]),
    ("R-21", "Puentes y cortes en ladera", 20300, 37223, [42, 50]),
    ("R-26", "Quebradas de los puentes", 20300, None, [42]),
    ("R-27", "Desvíos de los puentes y accesos a las localidades", 19343, 44693, [6, 41, 14, 42, 16, 57, 43]),
    ("R-32", "Sectores agrícolas (arrozales)", 20360, 21170, []),
    ("R-33", "Puente existente Tingo de Ponasa", 20300, None, [42]),
    ("R-33", "Puente existente Barco", 29329, None, []),
    ("R-33", "Puente existente Paucarillo", 37960, None, []),
    ("R-33", "Puente existente Apangurayco", 38919, None, []),
    ("R-33", "Puente existente Chanchamayco", 44149, None, []),
    ("R-34", "Puente Tingo de Ponasa", 20300, None, [42]),
    ("R-34", "Puente Barco", 29300, None, []),
]
NO_LOC = {"R-08": "Cantera de río y accesos: se georreferencian con la ubicación que fije el expediente técnico.",
          "R-09": "Fuente de agua y cantera de río: se georreferencian con la ubicación que fije el expediente técnico.",
          "R-14": "Cortes, canteras, DME y ensanches: ubicación según el plan de monitoreo arqueológico.",
          "R-25": "DME: se georreferencian con la ubicación aprobada en el estudio ambiental."}


RADIO = 500                                             # m: foto del mismo sector para un punto sin foto propia


def fotos_de(imgs):
    return ", ".join(FID[i] for i in imgs)


def foto_sector(a, b):
    """Foto más cercana (por progresiva) a un punto o sector sin foto propia, si está a menos de RADIO."""
    b = b or a
    d = lambda x: 0 if a <= x["prog"] <= b else min(abs(x["prog"] - a), abs(x["prog"] - b))
    x = min(FOTOS, key=d)
    return (x, d(x)) if d(x) <= RADIO else (None, d(x))


def celda_fotos(imgs, a, b):
    if imgs: return fotos_de(imgs)
    x, dd = foto_sector(a, b)
    return "%s (vista del sector, a %.0f m)" % (x["id"], dd) if x else "Registro documental"


def utm(prog):
    e, n = punto(prog); return "%.0f" % e, "%.0f" % n


def ctab(doc, texto):
    p = par(doc, texto, 9, True, C, antes=4, despues=3); p.paragraph_format.keep_with_next = True


def cita(doc, titulo, texto):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.width = Cm(16.0); sombra(c, "F2F2F2")
    p = c.paragraphs[0]; p.alignment = L_; fuente(p.add_run(titulo), 9.5, True, color=AZUL)
    q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; fuente(q.add_run(texto), 9.5, cur=True)
    par(doc, despues=4)


def h1(doc, texto):
    p = par(doc, texto, 12, True, L_, antes=12, despues=6, color=AZUL); p.paragraph_format.keep_with_next = True; return p


def h2(doc, texto):
    p = par(doc, texto, 10.5, True, L_, antes=8, despues=4, color=AZUL); p.paragraph_format.keep_with_next = True; return p


# ------------------------------------------------------------------ anexo 2: matriz impresa
def anexo_matriz(doc, pref="A2", titulo="ANEXO 2: MATRIZ DE GESTIÓN DE RIESGOS EN LA EJECUCIÓN DE LA OBRA"):
    if titulo: h1(doc, titulo)
    par(doc, "Impresión de la matriz MATRIZ_GESTION_RIESGOS_EJECUCION_OBRA_SM107.xlsx, que acompaña al estudio en formato Excel con sus fórmulas "
             "(celda amarilla = dato; celda verde = fórmula). Los efectos en costo se expresan en porcentaje del presupuesto de obra.", 9)
    sub = lambda n: "%s.%d" % (pref, n)

    h2(doc, pref + ".1 Datos del proyecto")
    datos = [("Código único de inversiones (CUI)", P["cui"], "Inventario vial / FTE"), ("Entidad", P["entidad"], "FTE"),
             ("Unidad ejecutora presupuestal", P["ue"], "FTE"), ("Ubicación", P["ubicacion"], "FTE"),
             ("Progresiva inicial", P["prog_ini"], "Inventario vial"), ("Progresiva final", P["prog_fin"], "Inventario vial"),
             ("Longitud del tramo (km)", "%.3f" % P["L_inv"], "Inventario vial"), ("Plazo de ejecución (días calendario)", "%d" % P["plazo"], "Resumen ejecutivo"),
             ("Modalidad de ejecución", P["modalidad"], "FTE 19"), ("Fuente de financiamiento", P["fuente_fin"], "FTE 20"),
             ("IMDA actual (veh/día)", "%d" % P["imda_actual"], "FTE 9.1"), ("Población beneficiaria directa (hab)", "{:,}".format(P["poblacion"]), "FTE 6"),
             ("Costo de un día de ampliación (gastos generales y supervisión), % del presupuesto de obra por día", "%.4f %%" % P["costo_dia_pct"],
              "Se reconoce cuando el retraso es atribuible a la Entidad")]
    tabla(doc, ["CONCEPTO", "VALOR", "FUENTE"], [list(d) for d in datos], [10.0, 9.0, 6.5], tam=8)

    h2(doc, pref + ".2 Escalas de probabilidad e impacto y matriz P x I")
    desc_p = ["Puede ocurrir solo en circunstancias excepcionales", "Podría ocurrir alguna vez durante la obra", "Puede ocurrir durante la obra",
              "Es probable que ocurra durante la obra", "Se espera que ocurra, casi con certeza"]
    tabla(doc, ["PROBABILIDAD", "VALOR", "DESCRIPCIÓN"], [[t, "%.2f" % v, d] for (t, v), d in zip(RS.PROB, desc_p)], [4.0, 2.0, 19.5], tam=8, alinear=[L_, C, L_])
    desc_i = ["Efecto mínimo; se absorbe con la holgura del frente", "Afecta un frente; se recupera dentro del mes",
              "Afecta varios frentes o un hito; requiere reprogramar", "Afecta la ruta crítica o el presupuesto de forma significativa",
              "Compromete el plazo total, el financiamiento o la viabilidad"]
    lim_c = ["%.1f %%" % (u * 100) for u in RS.IMPACTO_COSTO] + ["mayor"]; lim_t = ["%d" % u for u in RS.IMPACTO_PLAZO] + ["mayor"]
    tabla(doc, ["IMPACTO", "VALOR", "DESCRIPCIÓN", "COSTO HASTA (% PRESUPUESTO)", "PLAZO HASTA (días)"],
          [[t, "%.2f" % v, d, a, b] for (t, v), d, a, b in zip(RS.IMP, desc_i, lim_c, lim_t)], [3.0, 1.8, 13.2, 4.0, 3.5], tam=8, alinear=[L_, C, L_, C, C])
    tabla(doc, ["UMBRAL", "P x I", "INTERPRETACIÓN"],
          [["ALTO", ">= %.2f" % RS.UMBRAL_ALTO, "Respuesta obligatoria antes del inicio; seguimiento mensual del comité"],
           ["MODERADO", "%.2f a < %.2f" % (RS.UMBRAL_MOD, RS.UMBRAL_ALTO), "Respuesta planificada y seguimiento por el residente y el supervisor"],
           ["BAJO", "< %.2f" % RS.UMBRAL_MOD, "Se acepta y se vigila con el registro"]], [4.0, 4.0, 17.5], tam=8,
          colores=lambda i, j, v: NIVCOL.get(v) if j == 0 else None, alinear=[C, C, L_])
    filas = []
    for i in range(4, -1, -1):
        filas.append(["%s (%.2f)" % RS.PROB[i]] + ["%.3f" % (RS.PROB[i][1] * RS.IMP[j][1]) for j in range(5)])
    col = lambda v: (NIVCOL["ALTO"] if v >= RS.UMBRAL_ALTO - 1e-9 else NIVCOL["MODERADO"] if v >= RS.UMBRAL_MOD - 1e-9 else NIVCOL["BAJO"])
    ctab(doc, "Matriz P x I")
    tabla(doc, ["P \\ I"] + ["%s (%.2f)" % x for x in RS.IMP], filas, [4.5] + [4.2] * 5, tam=8,
          colores=lambda i, j, v: col(float(v)) if j > 0 else "D9D9D9", alinear=[C] * 6)

    h2(doc, pref + ".3 Identificación de riesgos")
    tabla(doc, ["COD.", "CATEGORÍA", "RIESGO (EVENTO)", "CAUSA", "EFECTO", "UBICACIÓN / PROGRESIVA", "EVIDENCIA EN LOS DOCUMENTOS"],
          [[r["cod"], r["cat"], r["riesgo"], r["causa"], r["efecto"], r["ubic"], r["evid"]] for r in R], [1.2, 2.2, 4.2, 5.0, 4.0, 3.6, 5.3], tam=6.5,
          alinear=[C, L_, L_, L_, L_, L_, L_])

    h2(doc, pref + ".4 Análisis cualitativo (P x I)")
    filas = [[r["cod"], r["riesgo"], r["P_txt"], "%.2f" % r["P"], "%.2f %%" % r["costo_pct"][1], "%d" % (r["ic"] + 1), "%d" % r["plazo"][1], "%d" % (r["ip"] + 1),
              r["I_txt"], "%.2f" % r["I"], "%.3f" % r["PxI"], r["nivel"], r["Pr_txt"], "%.2f" % r["Pr"], "%.3f" % r["PxIr"], r["nivel_r"]] for r in R]
    tabla(doc, ["COD.", "RIESGO", "PROBA-BILIDAD", "P", "COSTO MÁS PROB. (%)", "NIVEL COSTO", "PLAZO MÁS PROB. (días)", "NIVEL PLAZO", "IMPACTO", "I",
                "P x I", "NIVEL INHERENTE", "PROBAB. RESIDUAL", "P RES.", "P x I RES.", "NIVEL RESIDUAL"],
          filas, [1.1, 5.4, 1.6, 1.0, 1.6, 1.2, 1.5, 1.2, 1.5, 1.0, 1.2, 1.9, 1.6, 1.0, 1.2, 1.9], tam=6.5,
          colores=lambda i, j, v: NIVCOL.get(v) if j in (11, 15) else None, alinear=[C, L_] + [C] * 14)
    n_i = collections.Counter(r["nivel"] for r in R); n_r = collections.Counter(r["nivel_r"] for r in R)
    tabla(doc, ["RESUMEN", "INHERENTE", "RESIDUAL"], [["Riesgos de nivel " + k, "%d" % n_i[k], "%d" % n_r[k]] for k in ("ALTO", "MODERADO", "BAJO")],
          [8.0, 3.0, 3.0], tam=8, alinear=[L_, C, C])

    h2(doc, pref + ".5 Análisis cuantitativo: valor esperado")
    fe = {"Entidad": 1.0, "Compartido": 0.5, "Contratista": 0.0}; cd = P["costo_dia_pct"]
    filas = []; tot = [0.0, 0.0, 0.0, 0.0]
    for r in R:
        cm = sum(r["costo_pct"]) / 3.0; tm = r["t_media"]; f = fe[r["asig"]]
        ve = r["P"] * cm; vt = r["P"] * tm
        ei = r["P"] * cm * f * r["fr"] + (r["P"] * tm * f * cd if r["rc"] else 0)
        er = r["Pr"] * cm * f * r["fr"] + (r["Pr"] * tm * f * cd if r["rc"] else 0)
        for k, v in enumerate((ve, vt, ei, er)): tot[k] += v
        filas.append([r["cod"], r["riesgo"], r["asig"], "%.2f" % f, "%.2f" % r["fr"], "%.2f" % r["P"]] + ["%.2f" % v for v in r["costo_pct"]] +
                     ["%.2f" % cm, "%.3f" % ve] + ["%d" % v for v in r["plazo"]] + ["%.1f" % tm, "%.1f" % vt, "SI" if r["rc"] else "NO",
                                                                                   "%.3f" % ei, "%.2f" % r["Pr"], "%.3f" % er])
    filas.append(["", "TOTAL (suma simple, sin concurrencia de plazos)", "", "", "", "", "", "", "", "", "%.2f" % tot[0], "", "", "", "", "%.1f" % tot[1], "",
                  "%.2f" % tot[2], "", "%.2f" % tot[3]])
    tabla(doc, ["COD.", "RIESGO", "ASIGNADO A", "FACTOR ENTIDAD", "FACTOR PÓLIZA", "P", "COSTO MÍN. (%)", "COSTO MÁS PROB. (%)", "COSTO MÁX. (%)",
                "COSTO MEDIO (%)", "VALOR ESP. COSTO (%)", "PLAZO MÍN.", "PLAZO MÁS PROB.", "PLAZO MÁX.", "PLAZO MEDIO", "VALOR ESP. PLAZO (días)",
                "RUTA CRÍTICA", "EXPOSIC. ENTIDAD INHER. (%)", "P RES.", "EXPOSIC. ENTIDAD RESID. (%)"],
          filas, [1.0, 4.4, 1.6, 1.0, 1.0, 0.9, 1.1, 1.1, 1.1, 1.1, 1.2, 0.9, 1.0, 0.9, 1.0, 1.2, 1.0, 1.4, 0.9, 1.4], tam=6, alinear=[C, L_] + [C] * 18)
    par(doc, "Valor esperado = P x valor medio de la distribución triangular. Factor Entidad: 1 (Entidad), 0.5 (compartido), 0 (contratista); factor póliza: "
             "fracción que no cubre la póliza CAR. Los días de ruta crítica de la Entidad se valorizan con el costo diario de ampliación (" + sub(1) + "). La suma simple "
             "no considera que los retrasos se superponen; la simulación de Monte Carlo (" + sub(6) + ") sí lo hace.", 7.5, cur=True)

    h2(doc, pref + ".6 Simulación de Monte Carlo (20,000 iteraciones)")
    I_, Rr = MC["inherente"], MC["residual"]
    q = lambda d, fmt: [fmt(d[k]) for k in ("10", "50", "80", "90")]
    f2 = lambda v: pc(v); d0 = lambda v: "%.0f" % v
    filas = [["Efecto en costo para la Entidad - inherente (% del presupuesto de obra)"] + q(I_["entidad"], f2) + [f2(I_["entidad_media"])],
             ["Efecto en costo para la Entidad - residual (% del presupuesto de obra)"] + q(Rr["entidad"], f2) + [f2(Rr["entidad_media"])],
             ["Retraso total de la obra - inherente (días)"] + q(I_["dias"], d0) + [d0(I_["dias_media"])],
             ["Retraso total de la obra - residual (días)"] + q(Rr["dias"], d0) + [d0(Rr["dias_media"])],
             ["Retraso atribuible a la Entidad - inherente (días)"] + q(I_["diasE"], d0) + ["—"],
             ["Retraso atribuible a la Entidad - residual (días)"] + q(Rr["diasE"], d0) + ["—"],
             ["Efecto en costo de riesgos del contratista - inherente (%)"] + q(I_["contratista"], f2) + [f2(I_["contratista_media"])],
             ["Efecto en costo de riesgos del contratista - residual (%)"] + q(Rr["contratista"], f2) + [f2(Rr["contratista_media"])],
             ["Reserva de contingencia recomendada (P80 residual, % del presupuesto de obra)", "", "", f2(Rr["entidad"]["80"]), "", ""]]
    tabla(doc, ["RESULTADO", "P10", "P50", "P80", "P90", "MEDIA"], filas, [11.5, 2.8, 2.8, 2.8, 2.8, 2.8], tam=8, alinear=[L_] + [C] * 5,
          colores=lambda i, j, v: "FF7C80" if i == 8 and j in (0, 3) else None)
    for t in ["Cada riesgo ocurre en una iteración con su probabilidad (inherente o residual); si ocurre, su costo y su retraso se toman de una distribución "
              "triangular (mínimo, más probable, máximo).",
              "Efecto en costo para la Entidad: riesgos de la Entidad, 50 % de los compartidos (descontando la póliza CAR donde aplica) y los días de ampliación "
              "que le corresponden por el costo diario.",
              "Plazo: los riesgos de ruta crítica se agrupan por etapa (inicio, clima, terreno, gestión); en cada grupo se suma el mayor retraso y el 25 % de los "
              "demás, y los grupos se suman."]: vin(doc, t)
    p = doc.add_paragraph(); p.alignment = C; p.add_run().add_picture(os.path.join(FIG, "f5_montecarlo.png"), width=Cm(17.0))

    h2(doc, pref + ".7 Mapa de calor")
    for tit_, kp in (("Inherente (antes de la respuesta)", "p"), ("Residual (con la respuesta ejecutada)", "pr")):
        ctab(doc, "Mapa de calor - " + tit_)
        filas = []
        for i in range(4, -1, -1):
            filas.append([RS.PROB[i][0]] + [", ".join(r["cod"] for r in R if r[kp] == i and r["i"] == j) for j in range(5)])
        tabla(doc, ["PROBABILIDAD \\ IMPACTO"] + [x[0] for x in RS.IMP], filas, [4.5] + [4.2] * 5, tam=7.5, alinear=[C] * 6,
              colores=lambda i, j, v, f=filas: "D9D9D9" if j == 0 else col(RS.PROB[4 - i][1] * RS.IMP[j - 1][1]))

    h2(doc, pref + ".8 Plan de respuesta")
    tabla(doc, ["COD.", "RIESGO", "NIVEL INHERENTE", "ESTRATEGIA", "ACCIONES DE RESPUESTA", "RESPONSABLE", "PARTIDA O RECURSO", "MOMENTO",
                "ALERTA TEMPRANA (DISPARADOR)", "NIVEL RESIDUAL"],
          [[r["cod"], r["riesgo"], r["nivel"], r["resp"], r["acc"], r["quien"], r["partida"], r["momento"], r["alerta"], r["nivel_r"]] for r in R],
          [1.0, 3.4, 1.6, 1.5, 5.6, 2.3, 2.9, 1.8, 3.3, 1.6], tam=6.5,
          colores=lambda i, j, v: NIVCOL.get(v) if j in (2, 9) else None, alinear=[C, L_, C, C, L_, L_, L_, L_, L_, C])

    h2(doc, pref + ".9 Asignación de riesgos")
    crit = {"Entidad": "Depende de decisiones, permisos, predios, financiamiento o del contenido del expediente técnico, que controla la Entidad; "
                       "si ocurre, procede la modificación contractual o la ampliación de plazo.",
            "Contratista": "Depende de la planificación, los métodos, los recursos y la seguridad de la obra, que controla el contratista; "
                           "se cubre con sus gastos generales, sus pólizas y su programación.",
            "Compartido": "La causa tiene una parte de diseño o gestión de la Entidad y otra de ejecución del contratista; se reparte según la causa "
                          "que se compruebe en obra (en el análisis se toma 50 %)."}
    tabla(doc, ["COD.", "RIESGO", "ASIGNADO A", "CRITERIO"], [[r["cod"], r["riesgo"], r["asig"], crit[r["asig"]]] for r in R], [1.2, 8.8, 2.4, 13.1], tam=7,
          alinear=[C, L_, C, L_])
    n_as = collections.Counter(r["asig"] for r in R)
    tabla(doc, ["RIESGOS ASIGNADOS A", "CANTIDAD"], [[k, "%d" % n_as[k]] for k in ("Entidad", "Contratista", "Compartido")], [8.0, 3.0], tam=8, alinear=[L_, C])

    h2(doc, pref + ".10 Oportunidades")
    tabla(doc, ["COD.", "OPORTUNIDAD", "CAUSA", "PROBABILIDAD", "P", "AHORRO MÍN. (%)", "AHORRO MÁS PROB. (%)", "AHORRO MÁX. (%)", "VALOR ESPERADO",
                "ESTRATEGIA", "ACCIONES", "RESPONSABLE"],
          [[o["cod"], o["oport"], o["causa"], RS.PROB[o["p"]][0], "%.2f" % o["P"]] + ["%.2f" % v for v in o["beneficio_pct"]] +
           [("%.2f %%" % o["EMV_pct"] if o["EMV"] else "%.0f días" % o["EMV_plazo"]), o["resp"], o["acc"], o["quien"]] for o in RS.OPORT],
          [1.0, 3.4, 3.6, 1.6, 0.9, 1.4, 1.4, 1.4, 1.6, 1.6, 5.0, 2.6], tam=6.5, alinear=[C, L_, L_, C, C, C, C, C, C, C, L_, L_])

    h2(doc, pref + ".11 Formato de seguimiento mensual de riesgos (comité de gestión de riesgos)")
    par(doc, "Mes de obra: ______     Fecha de revisión: ____/____/______     Valorización N.°: ______", 8.5)
    tabla(doc, ["COD.", "RIESGO", "ESTADO", "PROBAB. ACTUAL", "IMPACTO ACTUAL", "NIVEL ACTUAL", "¿ALERTA ACTIVADA?", "ACCIÓN EJECUTADA / PENDIENTE",
                "RESPONSABLE", "COMENTARIOS"],
          [[r["cod"], r["riesgo"], "Activo", "", "", "", "", "", r["quien"], ""] for r in R] +
          [["R-%02d" % (35 + k), "", "Nuevo", "", "", "", "", "", "", ""] for k in range(5)],
          [1.0, 5.0, 1.4, 1.6, 1.6, 1.6, 1.6, 4.6, 3.3, 3.8], tam=6.5, alinear=[C, L_, C, C, C, C, C, L_, L_, L_])

    h2(doc, pref + ".12 Lluvia mensual de referencia y programación recomendada")
    MES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
    t_, j_ = LL["Tarapoto"], LL["Juanjuí"]; filas = []; s = [0, 0, 0]
    prog = {"LLUVIOSA": "Alcantarillas, cunetas, gaviones, fabricación de vigas, campamento; evitar terraplenes en zonas de arrozal y cimentación de puentes",
            "TRANSICIÓN": "Cortes con drenaje de coronación, subbase y base con protección; preparar desvíos de cauce",
            "ESTIAJE": "Cimentación y estribos de puentes, terraplenes en zonas bajas, base y carpeta asfáltica"}
    for m in range(1, 13):
        a = t_["meses"][str(m)]; b = j_["meses"][str(m)]; pr = (a[0] + b[0]) / 2
        tem = "LLUVIOSA" if pr >= 110 else ("TRANSICIÓN" if pr >= 75 else "ESTIAJE")
        s[0] += a[0]; s[1] += b[0]; s[2] += pr
        filas.append([MES[m - 1], "%.0f" % a[0], "%.0f" % b[0], "%.0f" % pr, "%.1f" % ((a[2] + b[2]) / 2), tem, prog[tem]])
    filas.append(["ANUAL", "%.0f" % s[0], "%.0f" % s[1], "%.0f" % s[2], "", "", ""])
    tc = {"LLUVIOSA": "FF7C80", "TRANSICIÓN": "FFD966", "ESTIAJE": "A9D08E"}
    tabla(doc, ["MES", "TARAPOTO (mm)", "JUANJUÍ (mm)", "PROMEDIO (mm)", "DÍAS >= 10 mm (prom.)", "TEMPORADA", "PROGRAMACIÓN RECOMENDADA"], filas,
          [2.3, 2.2, 2.2, 2.2, 2.4, 2.6, 11.6], tam=7.5, colores=lambda i, j, v: tc.get(v) if j == 5 else None, alinear=[L_, C, C, C, C, C, L_])
    par(doc, "Fuente: NOAA GHCN-Daily, estaciones Tarapoto (PE000084455) y Juanjuí (PEM00084474); meses con al menos 90 % de datos. Series cortas: el estudio "
             "hidrológico del expediente prevalece para el diseño.", 7.5, cur=True)


# ------------------------------------------------------------------ anexo 3: panel fotográfico
def celda_foto(c, x):
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    c.paragraphs[0].paragraph_format.space_after = Pt(0)
    p = c.paragraphs[0]; p.alignment = C
    w, h = Image.open(x["ruta"]).size
    if w >= h: p.add_run().add_picture(x["ruta"], width=Cm(7.6))
    else: p.add_run().add_picture(x["ruta"], height=Cm(7.8))
    q = c.add_paragraph(); q.alignment = C; q.paragraph_format.space_before = Pt(2)
    fuente(q.add_run("Foto N.° %s - km %s" % (x["id"], km(x["prog"]))), 8.5, True, color=AZUL)
    q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; q.paragraph_format.space_after = Pt(1); fuente(q.add_run(x["desc"]), 8)
    q = c.add_paragraph(); q.alignment = L_
    fuente(q.add_run("UTM WGS 84 - 18 S:  E %s   N %s\n" % (("%.0f" % x["e"]), ("%.0f" % x["n"]))), 7.5, True)
    fuente(q.add_run("Fecha y hora: %s\nRiesgos relacionados: %s" % (x["fecha"], ", ".join(x["riesgos"]))), 7.5)
    sec = [s[0] for s in SECTORES if x["img"] in s[4]]
    if sec: fuente(q.add_run("\nSector crítico (%s): %s" % (T["t04"], "; ".join(sec))), 7.5)
    vis = ["km %s (a %.0f m)" % (d["prog"], d["dist"]) for d in PUNTOS if d["foto"] is x]
    if vis: fuente(q.add_run("\nVista del sector para: %s" % "; ".join(vis)), 7.5, True)


def grilla(doc, items, fn):
    for k in range(0, len(items), 2):
        t = doc.add_table(rows=1, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        for j in range(2):
            c = t.rows[0].cells[j]; c.width = Cm(8.2)
            if k + j < len(items): fn(c, items[k + j])
        tr = t.rows[0]._tr; trPr = tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)


def sin_foto_propia():
    """Puntos de la Tabla N.° 04 y de la Tabla N.° 08-A sin fotografía propia, sin repetir progresivas, con la foto del sector si la hay."""
    vis = collections.OrderedDict()
    for prog, a, b, d, imgs in SECTORES:
        if not imgs: vis.setdefault(a, dict(prog=prog, a=a, b=b, desc=d, ref=["Tabla N.° 04"]))
    for cod, d, a, b, imgs in UBIC:
        if imgs: continue
        e = vis.setdefault(a, dict(prog=km(a) + (" a " + km(b) if b else ""), a=a, b=b, desc=d, ref=[]))
        if cod not in e["ref"]: e["ref"].append(cod)
    L = []
    for a, d in sorted(vis.items()):
        d["e"], d["n"] = utm(a); d["ref"] = ", ".join(d["ref"]); d["foto"], d["dist"] = foto_sector(d["a"], d["b"]); L.append(d)
    return L


PUNTOS = sin_foto_propia()
ASOC = [d for d in PUNTOS if d["foto"]]; VERIF = [d for d in PUNTOS if not d["foto"]]


# ------------------------------------------------------------------ bloques (se usan aquí y en los anexos del estudio)
def ref_08(r):
    """Columna de coordenadas y fotografía de la tabla de riesgos identificados."""
    cods = {u[0] for u in UBIC}
    if r["cod"] in cods:
        f = sorted({FID[i] for c_, d_, a, b, im in UBIC if c_ == r["cod"] for i in im})
        f += sorted({foto_sector(a, b)[0]["id"] for c_, d_, a, b, im in UBIC if c_ == r["cod"] and not im and foto_sector(a, b)[0]} - set(f))
        return "%s. Fotos: %s" % (T["t08a"], ", ".join(f) if f else "registro documental")
    if r["cod"] in NO_LOC: return "Ubicación por definir"
    return "No localizado"


T = {"t04": "Tabla N.° 04", "t08": "Tabla N.° 08", "t08a": "Tabla N.° 08-A", "panel": "Anexo 3"}


def tabla_08(doc, titulo):
    ctab(doc, titulo)
    tabla(doc, ["COD.", "CATEGORÍA", "RIESGO", "UBICACIÓN", "COORDENADAS Y FOTOGRAFÍA"],
          [[r["cod"], r["cat"], r["riesgo"], r["ubic"], ref_08(r)] for r in R], [1.2, 2.2, 5.3, 4.3, 3.0], tam=7, alinear=[C, L_, L_, L_, L_])
    par(doc, "Ubicación por definir: " + " ".join("%s: %s" % kv for kv in NO_LOC.items()), 7.5, cur=True)


def tabla_08a(doc, titulo):
    ctab(doc, titulo)
    filas = []
    for cod, d, a, b, imgs in UBIC:
        e1, n1 = utm(a)
        if b: e2, n2 = utm(b); pr = "%s\n%s" % (km(a), km(b)); e = "%s\n%s" % (e1, e2); n = "%s\n%s" % (n1, n2)
        else: pr, e, n = km(a), e1, n1
        filas.append([cod, d, pr, e, n, celda_fotos(imgs, a, b)])
    tabla(doc, ["COD.", "UBICACIÓN", "PROGRESIVA (inicio / fin)", "ESTE (m)", "NORTE (m)", "FOTO N.°"], filas, [1.2, 5.4, 2.4, 2.0, 2.2, 2.8], tam=7,
          alinear=[C, L_, C, C, C, C])


def tabla_04(doc, titulo):
    ctab(doc, titulo)
    filas = []
    for prog, a, b, d, imgs in SECTORES:
        e1, n1 = utm(a); e2, n2 = utm(b) if b else ("—", "—")
        filas.append([prog, d, e1, n1, e2, n2, celda_fotos(imgs, a, b)])
    tabla(doc, ["PROGRESIVA", "DESCRIPCIÓN", "ESTE INICIO", "NORTE INICIO", "ESTE FIN", "NORTE FIN", "FOTO N.°"], filas,
          [2.4, 4.2, 1.6, 1.9, 1.6, 1.9, 2.4], tam=7, alinear=[L_, L_, C, C, C, C, C])
    par(doc, "LI: lado izquierdo; LD: lado derecho. Cuando el punto o sector es puntual no se indican coordenadas de fin.", 7.5, cur=True)


def fijar_anchos(t, anchos):
    """Anchos fijos de columna (cm), como en estudio_docx.tabla."""
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    for e in tblPr.findall(qn("w:tblW")): tblPr.remove(e)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(anchos) * 567))); tw.set(qn("w:type"), "dxa"); tblPr.append(tw)
    for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), anchos): gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for c, w in zip(row.cells, anchos): c.width = Cm(w)


def ficha(doc, cabecera, lineas, fs, mas=""):
    """Ficha de dos columnas: datos y coordenadas a la izquierda; fotografías (máx. 2) a la derecha.
    lineas: (rótulo, texto); fs: [(foto, distancia o None)]; mas: otras fotos del mismo punto."""
    t = doc.add_table(rows=2, cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    fijar_anchos(t, [5.4, 11.6])
    cab = t.cell(0, 0).merge(t.cell(0, 1)); cab.text = ""; sombra(cab, "1F4E79"); cab.paragraphs[0].paragraph_format.keep_with_next = True
    fuente(cab.paragraphs[0].add_run(cabecera), 8.5, True, color=ED.RGBColor(255, 255, 255))
    izq, der = t.cell(1, 0), t.cell(1, 1); sombra(izq, "F2F2F2")
    izq.vertical_alignment = der.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = izq.paragraphs[0]
    for k, (rot, txt) in enumerate(lineas):
        if rot: fuente(p.add_run(("\n" if k else "") + rot), 8, True); fuente(p.add_run(txt), 8)
        else: fuente(p.add_run("\n" + txt), 7.5)
    q = der.paragraphs[0]; q.alignment = C
    if fs:
        ancho = 10.8 if len(fs) == 1 else 4.8
        for x, dd in fs:
            w, h = Image.open(x["ruta"]).size
            if w >= h: q.add_run().add_picture(x["ruta"], width=Cm(ancho))
            else: q.add_run().add_picture(x["ruta"], height=Cm(7.0 if len(fs) == 1 else 6.0))
            q.add_run(" ")
        r_ = der.add_paragraph(); r_.alignment = C
        txt = "; ".join("Foto N.° %s - km %s - E %.0f  N %.0f - %s%s" % (x["id"], km(x["prog"]), x["e"], x["n"], x["fecha"],
                                                                         "" if dd is None else " (vista del sector, a %.0f m)" % dd) for x, dd in fs)
        if mas: txt += ". Ver también: " + mas
        fuente(r_.add_run(txt), 7.5)
    else:
        q.paragraph_format.space_before = Pt(20); q.paragraph_format.space_after = Pt(20)
        fuente(q.add_run("Registro documental del estudio de inventario vial, ubicado con las coordenadas del eje. Su registro fotográfico forma parte "
                         "de la verificación de campo previa a la convocatoria (acción complementaria del plan de respuesta)."), 8, cur=True)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    par(doc, despues=4)


def fichas_04(doc, titulo):
    """Una ficha por punto o sector crítico de la tabla de puntos críticos."""
    par(doc, titulo, 10, True, antes=6, despues=4).paragraph_format.keep_with_next = True
    for k, (prog, a, b, d, imgs) in enumerate(SECTORES, 1):
        if imgs: fs = [(next(x for x in FOTOS if x["img"] == i), None) for i in imgs[:2]]
        else:
            x, dd = foto_sector(a, b); fs = [(x, dd)] if x else []
        e1, n1 = utm(a)
        lin = [("Progresiva: ", prog), ("Descripción: ", d + "\n"), ("Coordenadas UTM (WGS 84 - 18 S)", ""), (None, "Inicio:  E %s   N %s" % (e1, n1))]
        if b: lin.append((None, "Fin:      E %s   N %s" % utm(b)))
        rr = sorted({c_ for c_, d_, a_, b_, im in UBIC if a_ == a})
        if rr: lin.append(("\nRiesgos: ", ", ".join(rr)))
        ficha(doc, "FICHA %02d - km %s - %s" % (k, prog, d), lin, fs, fotos_de(imgs[2:]))


def fichas_08(doc, titulo):
    """Una ficha por riesgo con ubicación física de la tabla de riesgos identificados."""
    par(doc, titulo, 10, True, antes=6, despues=4).paragraph_format.keep_with_next = True
    cods = [u[0] for u in UBIC]
    for r in [r for r in R if r["cod"] in cods]:
        filas = [u for u in UBIC if u[0] == r["cod"]]
        propias = []
        for u in filas:
            for i in u[4]:
                if FID[i] not in [x["id"] for x in propias]: propias.append(next(x for x in FOTOS if x["img"] == i))
        sector = []
        for c_, d_, a, b, im in filas:
            if im: continue
            x, dd = foto_sector(a, b)
            if x and x not in propias and x not in [y for y, _ in sector]: sector.append((x, dd))
        fs = ([(x, None) for x in propias] + sector)[:2]
        resto = [x["id"] for x in propias][2:] + [x["id"] for x, _ in sector][max(0, 2 - len(propias)):]
        lin = [("Categoría: ", r["cat"]), ("Riesgo: ", r["riesgo"]), ("Ubicación: ", r["ubic"] + "\n"), ("Puntos (UTM WGS 84 - 18 S)", "")]
        for c_, d_, a, b, im in filas:
            e, n = utm(a)
            lin.append((None, "km %s%s: E %s  N %s - %s" % (km(a), " a " + km(b) if b else "", e, n, d_)))
        ficha(doc, "%s - %s" % (r["cod"], r["riesgo"]), lin, fs, ", ".join(resto))
    nl = [r["cod"] for r in R if r["cod"] not in cods]
    par(doc, "Riesgos sin ubicación física (gestión, contractuales, económicos, permisos o seguridad de todo el proyecto, o con ubicación por definir): %s."
        % ", ".join(nl), 8, cur=True)


def accion_verificacion(doc):
    par(doc, "Acción complementaria del plan de respuesta", 10, True, antes=6, despues=3)
    tabla(doc, ["RIESGOS", "ACCIÓN", "RESPONSABLE", "MOMENTO"],
          [[", ".join(sorted({c for d in VERIF for c in d["ref"].split(", ") if c.startswith("R-")})),
            "Verificación de campo de los %d puntos de registro documental (%s, C), con fotografía georreferenciada (fecha y coordenadas UTM "
            "impresas) y descripción del estado actual; actualización de la matriz de riesgos con el resultado." % (len(VERIF), T["panel"]),
            "Entidad (DRTC San Martín)", "Antes de la convocatoria"]], [2.6, 8.4, 2.6, 2.4], tam=7.5, alinear=[C, L_, L_, C])


def panel(doc):
    par(doc, "Fuente: estudio de inventario vial SM-107, panel fotográfico (setiembre 2026). Coordenadas UTM WGS 84, zona 18 S, impresas por la cámara en cada "
             "fotografía; progresiva referida al eje georreferenciado del estudio. Las fotografías se ordenan por progresiva.", 8.5, cur=True)
    h2(doc, "A. Fotografías de los riesgos y sectores críticos")
    ctab(doc, "Relación de fotografías")
    tabla(doc, ["FOTO N.°", "PROGRESIVA", "ESTE (m)", "NORTE (m)", "FECHA Y HORA", "RIESGOS", "SECTOR CRÍTICO (%s)" % T["t04"].upper()],
          [[x["id"], km(x["prog"]), "%.0f" % x["e"], "%.0f" % x["n"], x["fecha"], ", ".join(x["riesgos"]),
            "; ".join(s[0] for s in SECTORES if x["img"] in s[4]) or "—"] for x in FOTOS],
          [1.4, 1.8, 1.8, 2.0, 2.6, 4.0, 3.4], tam=7, alinear=[C, C, C, C, C, L_, L_])
    grilla(doc, FOTOS, celda_foto)
    h2(doc, "B. Puntos registrados con la fotografía de su sector")
    par(doc, "Puntos de la %s y de la %s sin fotografía propia, registrados con la fotografía georreferenciada más cercana del mismo "
             "tramo (a no más de %d m). Las coordenadas son las del punto en el eje." % (T["t04"], T["t08a"], RADIO), 8.5)
    tabla(doc, ["N.°", "PROGRESIVA", "DESCRIPCIÓN", "ESTE (m)", "NORTE (m)", "RELACIONADO CON", "FOTO DEL SECTOR"],
          [["%d" % i, d["prog"], d["desc"], d["e"], d["n"], d["ref"], "%s (km %s, a %.0f m)" % (d["foto"]["id"], km(d["foto"]["prog"]), d["dist"])]
           for i, d in enumerate(ASOC, 1)], [0.9, 2.4, 4.1, 1.5, 1.7, 2.8, 3.6], tam=7, alinear=[C, C, L_, C, C, L_, C])
    h2(doc, "C. Puntos de registro documental")
    par(doc, "Puntos registrados en el estudio de inventario vial sin fotografía georreferenciada a menos de %d m; se ubican con las coordenadas del eje. "
             "Su registro fotográfico forma parte de la verificación de campo previa a la convocatoria (acción complementaria del plan de respuesta)." % RADIO, 8.5)
    tabla(doc, ["N.°", "PROGRESIVA", "DESCRIPCIÓN", "ESTE (m)", "NORTE (m)", "RELACIONADO CON", "FOTO MÁS CERCANA"],
          [["%d" % i, d["prog"], d["desc"], d["e"], d["n"], d["ref"], "a %.0f m" % d["dist"]] for i, d in enumerate(VERIF, 1)],
          [0.9, 2.4, 4.1, 1.5, 1.7, 3.6, 2.8], tam=7, alinear=[C, C, L_, C, C, L_, C])


# ------------------------------------------------------------------ documento
def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = ED.FT; st.font.size = Pt(10)
    s0 = doc.sections[0]; s0.page_width, s0.page_height = Cm(21.0), Cm(29.7)
    s0.left_margin = s0.right_margin = Cm(2.5); s0.top_margin = Cm(2.5); s0.bottom_margin = Cm(2.0)
    con_foto = [s for s in SECTORES if s[4]]
    sec_asoc = [s for s in SECTORES if not s[4] and foto_sector(s[1], s[2])[0]]

    par(doc, "GOBIERNO REGIONAL SAN MARTÍN", 12, True, C, despues=2, color=AZUL)
    par(doc, "DIRECCIÓN REGIONAL DE TRANSPORTES Y COMUNICACIONES", 10.5, True, C, despues=16, color=AZUL)
    par(doc, "LEVANTAMIENTO DE OBSERVACIONES", 16, True, C, despues=6)
    par(doc, "ESTUDIO DE GESTIÓN DE RIESGOS EN LA PLANIFICACIÓN DE LA EJECUCIÓN DE LA OBRA", 11.5, True, C, despues=10)
    par(doc, "“" + P["nombre"] + "”", 10, True, C, despues=4)
    par(doc, "CUI N.° %s" % P["cui"], 10.5, True, C, despues=14)

    h1(doc, "CUADRO RESUMEN")
    tabla(doc, ["N.°", "OBSERVACIÓN", "ABSOLUCIÓN", "UBICACIÓN EN EL ESTUDIO"],
          [["1", "La MATRIZ DE GESTIÓN DE RIESGOS EJECUCIÓN OBRA SM107 mencionada en el Anexo 2 no se observa en el anexo.",
            "Se incorpora la matriz impresa completa en el Anexo 2 (12 apartados: datos, escalas, identificación, análisis cualitativo y cuantitativo, Monte Carlo, "
            "mapa de calor, plan de respuesta, asignación, oportunidades, seguimiento y calendario). El Excel con fórmulas se mantiene como archivo adjunto.",
            "Anexo 2"],
           ["2", "Tabla N.° 8 Riesgos identificados: adjuntar fotografías con coordenadas de acuerdo con la identificación realizada.",
            "La Tabla N.° 08 incorpora la referencia a las fotografías; se agrega la Tabla N.° 08-A con la ubicación de cada riesgo localizado (progresiva y "
            "coordenadas UTM) y el Anexo 3, panel fotográfico georreferenciado (%d fotografías con fecha y coordenadas)." % len(FOTOS),
            "Capítulo 5 y Anexo 3"],
           ["3", "Tabla N.° 04 Puntos y sectores críticos del inventario vial: adjuntar fotografías con coordenadas de acuerdo con la descripción del sector crítico.",
            "La Tabla N.° 04 incorpora las coordenadas UTM de inicio y fin de cada sector y el número de la fotografía del sector; las fotografías están en "
            "el Anexo 3. Los puntos sin fotografía a menos de %d m se presentan como registro documental con sus coordenadas." % RADIO,
            "Capítulo 3 y Anexo 3"]],
          [1.0, 4.6, 7.6, 2.8], tam=8, alinear=[C, L_, L_, C])

    # ------------------------------------------------------------------ observación 1
    h1(doc, "OBSERVACIÓN N.° 1")
    cita(doc, "Anexo 2. Documentos que acompañan el estudio",
         "Se menciona la MATRIZ DE GESTIÓN DE RIESGOS EJECUCIÓN OBRA SM107, la cual no se observa en el Anexo, por lo cual se recomienda adjuntar al entregable.")
    mixto(doc, [("Absolución: ", True), ("se adjunta la matriz al entregable de dos formas: (a) impresa dentro del estudio, en el Anexo 2, con todas sus hojas; "
                                         "y (b) como archivo Excel con fórmulas (MATRIZ_GESTION_RIESGOS_EJECUCION_OBRA_SM107.xlsx), que permite actualizarla durante la obra. "
                                         "El Anexo 2 queda con el contenido siguiente:", False)])
    tabla(doc, ["APARTADO", "HOJA DE LA MATRIZ", "CONTENIDO"],
          [["A2.1", "DATOS", "Datos del proyecto y costo diario de ampliación en porcentaje"],
           ["A2.2", "ESCALAS", "Escalas de probabilidad e impacto, umbrales y matriz P x I"],
           ["A2.3", "IDENTIFICACION", "Registro de riesgos: evento, causa, efecto, ubicación y evidencia"],
           ["A2.4", "ANALISIS CUALITATIVO", "Probabilidad, impacto en costo y plazo, P x I y nivel (inherente y residual)"],
           ["A2.5", "ANALISIS CUANTITATIVO", "Efecto en costo (%) y plazo triangulares, valor esperado y exposición de la Entidad"],
           ["A2.6", "MONTE CARLO", "Resultados de la simulación (20,000 iteraciones), inherente y residual"],
           ["A2.7", "MAPA DE CALOR", "Ubicación de cada riesgo en la matriz P x I, antes y después de la respuesta"],
           ["A2.8", "PLAN DE RESPUESTA", "Estrategia, acciones, responsable, partida, momento y alerta temprana"],
           ["A2.9", "ASIGNACION", "Parte que asume cada riesgo y criterio"],
           ["A2.10", "OPORTUNIDADES", "Oportunidades y su estrategia"],
           ["A2.11", "SEGUIMIENTO", "Formato de seguimiento mensual para el comité de gestión de riesgos"],
           ["A2.12", "CALENDARIO", "Lluvias mensuales de referencia y programación recomendada"]],
          [2.0, 4.4, 9.6], tam=8, alinear=[C, L_, L_])

    # ------------------------------------------------------------------ observación 2
    h1(doc, "OBSERVACIÓN N.° 2")
    cita(doc, "Tabla N.° 8 Riesgos identificados",
         "Se contemplan de acuerdo con la ubicación de las progresivas; sin embargo, es necesario adjuntar fotografías con coordenadas de acuerdo con la "
         "identificación realizada.")
    mixto(doc, [("Absolución: ", True), ("se georreferencian los riesgos que tienen una ubicación física en el tramo y se adjunta el panel fotográfico. "
                                         "Criterios:", False)])
    for t in ["Fotografías: panel fotográfico del estudio de inventario vial, tomadas del 25 al 28 de setiembre de 2026 con cámara que imprime la fecha, la hora "
              "y las coordenadas UTM (WGS 84, zona 18 S) en cada imagen.",
              "Progresiva de cada fotografía: punto más cercano del eje georreferenciado del estudio (plano clave PC-01); %d de las %d fotografías están a menos de "
              "20 m del eje." % (sum(x["dist"] < 20 for x in FOTOS), len(FOTOS)),
              "Coordenadas de cada riesgo (Tabla N.° 08-A): punto del eje en la progresiva de inicio y de fin del sector.",
              "Puntos sin fotografía propia: se asocian a la fotografía georreferenciada más cercana del mismo tramo, a no más de %d m, como vista del sector; "
              "en la tabla se indica la distancia. Las coordenadas del punto son las de la tabla, no las de la fotografía." % RADIO,
              "Puntos sin fotografía a menos de %d m (%d): son registros documentales del estudio de inventario vial, ubicados con las coordenadas del eje "
              "(Anexo 3, C); su registro fotográfico forma parte de la verificación de campo previa a la convocatoria, acción complementaria del plan de respuesta."
              % (RADIO, len(VERIF)),
              "Los riesgos de gestión, contractuales, económicos, de permisos o de seguridad que abarcan todo el proyecto no tienen un punto físico; en la Tabla N.° 08 "
              "se indican como «No localizado»."]: vin(doc, t)
    par(doc, "La Tabla N.° 08 queda así:", antes=4)
    tabla_08(doc, "Tabla N.° 08: Riesgos identificados")
    tabla_08a(doc, "Tabla N.° 08-A: Ubicación georreferenciada de los riesgos (UTM WGS 84 - zona 18 S)")
    fichas_08(doc, "Fichas de los riesgos identificados de la Tabla N.° 08")

    # ------------------------------------------------------------------ observación 3
    h1(doc, "OBSERVACIÓN N.° 3")
    cita(doc, "Tabla N.° 04. Puntos y sectores críticos del inventario vial",
         "Se ubican los sectores críticos por progresivas; sin embargo, es necesario adjuntar fotografías con coordenadas de acuerdo con la descripción del "
         "sector crítico.")
    mixto(doc, [("Absolución: ", True), ("se agregan a la Tabla N.° 04 las coordenadas UTM de inicio y fin de cada punto o sector crítico, tomadas del eje "
                                         "georreferenciado en la progresiva del inventario, y el número de la fotografía del Anexo 3. De los %d puntos y sectores, %d tienen "
                                         "fotografía propia y %d se registran con la fotografía georreferenciada de su sector (a no más de %d m, distancia indicada); los %d "
                                         "restantes son registros documentales ubicados con las coordenadas del eje (Anexo 3, C)."
                                         % (len(SECTORES), len(con_foto), len(sec_asoc), RADIO, len(SECTORES) - len(con_foto) - len(sec_asoc)), False)])
    tabla_04(doc, "Tabla N.° 04: Puntos y sectores críticos del inventario vial (UTM WGS 84 - zona 18 S)")
    fichas_04(doc, "Fichas de los puntos y sectores críticos de la Tabla N.° 04")
    accion_verificacion(doc)

    # ------------------------------------------------------------------ anexo 2 (horizontal)
    sa = seccion(doc, True)
    anexo_matriz(doc)

    # ------------------------------------------------------------------ anexo 3 (vertical)
    sb = seccion(doc, False); sb.left_margin = sb.right_margin = Cm(2.0)
    h1(doc, "ANEXO 3: PANEL FOTOGRÁFICO GEORREFERENCIADO")
    panel(doc)

    for s in doc.sections: pie(s, "Levantamiento de observaciones - Estudio de gestión de riesgos en la ejecución de obra - SM-107 - CUI N.° %s" % P["cui"])
    doc.save(SAL); print(SAL, "fotos", len(FOTOS), "con foto del sector", len(ASOC), "a verificar", len(VERIF))


if __name__ == "__main__":
    main()
