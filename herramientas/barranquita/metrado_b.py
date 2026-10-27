"""Metrado de drenaje pluvial (red de cunetas) de la I.E. N.° 0590 Barranquita (C.I. 468901), con formulas.

Formato y nombres de partidas del metrado de drenaje pluvial del CAR Mujeres (CUI 2717013); cantidades propias de este proyecto
(cantidades_b.py: perfiles de cada eje, tramos tapados, tapas de inspeccion y descargas).
Hojas: RESUMEN, PLANILLA GENERAL DE METRADOS, CONCRETO EN CUNETAS, ENCOFRADO Y DESENCOFRADO, METRADO ACERO,
METRADO JUNTA DE DILATACION, METRADO DE CURADO, METRADO DE REJILLAS, PLANILLA TAPAS INSPECCION, PARAMETROS y Listas.
Las hojas de detalle calculan cada eje; la PLANILLA GENERAL suma los ejes y el RESUMEN toma los totales de la planilla.
"""
import os, sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.pagebreak import Break
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import cantidades_b as Q

SAL = os.path.join(RAIZ, "entregables_barranquita", "METRADO_DRENAJE_PLUVIAL_IE0590_BARRANQUITA.xlsx")
CALC = os.path.join(RAIZ, "entregables_barranquita", "_calc")
ENTIDAD = "GOBIERNO REGIONAL SAN MARTIN - GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO"
PROYECTO = ("“MEJORAMIENTO Y AMPLIACIÓN DEL SERVICIO DE EDUCACIÓN PRIMARIA EN I.E. N° 0590 DISTRITO DE BARRANQUITA DE LA "
            "PROVINCIA DE LAMAS DEL DEPARTAMENTO DE SAN MARTÍN” - C.I. N° 468901")
UBIC = "SAN MARTIN - LAMAS - BARRANQUITA"
FECHA = "OCTUBRE 2026"
_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
LETRA = "Arial Narrow"
CAB, SEC, PAR_, AZC, VER, GRS = "305496", "D9EAF7", "FFF2CC", "DDEBF7", "F0F7E8", "F2F2F2"


def F(**k):
    return Font(name=LETRA, size=k.pop("size", 11), **k)


def fill(c): return PatternFill("solid", fgColor=c)


def cel(ws, r, c, v, fmt=None, bold=False, bg=None, color=None, al="center", wrap=False, borde=True, size=11):
    x = ws.cell(r, c, v); x.font = F(bold=bold, size=size, color=color)
    x.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if fmt: x.number_format = fmt
    if bg: x.fill = fill(bg)
    if borde: x.border = BOR
    return x


def cab_hoja(ws, titulo, ncol, col0=1, fila=1, lab=1):
    """Titulo y datos del proyecto tomados del RESUMEN (como en el metrado del CAR Mujeres); lab = columnas del rotulo."""
    cel(ws, fila, col0, titulo, bold=True, size=14, borde=False)
    ws.merge_cells(start_row=fila, start_column=col0, end_row=fila, end_column=col0 + ncol - 1)
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        r = fila + 1 + i
        cel(ws, r, col0, k, bold=True, al="left"); cel(ws, r, col0 + lab, ref, al="left", wrap=True)
        if lab > 1: ws.merge_cells(start_row=r, start_column=col0, end_row=r, end_column=col0 + lab - 1)
        ws.merge_cells(start_row=r, start_column=col0 + lab, end_row=r, end_column=col0 + ncol - 1)
    ws.row_dimensions[fila + 2].height = 33


def nota(ws, r, c0, c1, texto, alto=30):
    """Nota al pie en celdas combinadas (no ensancha el area de impresion)."""
    cel(ws, r, c0, texto, al="left", wrap=True, borde=False)
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1); ws.row_dimensions[r].height = alto


def encabezado(ws, r, textos, anchos=None, c0=1, bg=CAB, color="FFFFFF", alto=32):
    for j, t in enumerate(textos):
        cel(ws, r, c0 + j, t, bold=True, bg=bg, color=color, wrap=True)
        if anchos: ws.column_dimensions[CL(c0 + j)].width = anchos[j]
    ws.row_dimensions[r].height = alto


# ---------------------------------------------------------------------------- PARAMETROS (columna B)
PAR = [("Ancho exterior de la cuneta (m)", Q.BT, "be", "Planos DC-02: muros 0.10 + 0.40 libres + 0.10"),
       ("Luz libre interior (m)", Q.B, "b", "Planos DC-02"), ("Espesor de muros (m)", Q.E, "em", "Planos DC-02"),
       ("Espesor de losa de fondo (m)", Q.EF, "ef", "Planos DC-02"), ("Espesor de losa tapa en tramos tapados (m)", Q.ET, "et", "Planos DC-02"),
       ("Espesor de solado f'c = 100 kg/cm2 (m)", Q.SOLADO, "sol", "Solado e = 4\" bajo la losa de fondo"),
       ("Sobreancho de excavacion a cada lado (m)", Q.SOBREANCHO, "sa", "Espacio de trabajo y encofrado"),
       ("Ancho de excavacion (m)", "=B5+2*B11", "bexc", "Ancho exterior + 2 sobreanchos"),
       ("Profundidad adicional bajo la cota de fondo: losa + solado (m)", "=B8+B10", "hadd", "Losa de fondo + solado"),
       ("Factor de esponjamiento para eliminacion", Q.ESPONJ, "esp", "Eliminacion = (excavacion - relleno) x factor"),
       ("Junta de dilatacion en la seccion en U cada (m)", Q.JUNTA_CADA, "jc", "Juntas en toda la seccion en U"),
       ("Recubrimiento del acero (m)", Q.REC, "rec", "Planos DC-02"), ("Peso de la barra de 3/8\" (kg/m)", Q.KG38, "kg", "Hoja Listas"),
       ("Separacion de las U de 3/8\" (m)", Q.SEP_T, "st", "Planos DC-02"), ("Separacion de barras longitudinales de 3/8\" (m)", Q.SEP_L, "sl", "Planos DC-02"),
       ("Profundidad minima con acero (m): h >= 0.40 o tramo tapado", Q.H_ACERO, "hac", "Con h < 0.40 m la cuneta abierta es de concreto simple"),
       ("Tapa de inspeccion: lado (m)", Q.TAPA[0], "tl", "Planos DC-01: tapa 0.60 x 0.60"), ("Tapa de inspeccion: espesor (m)", Q.TAPA[2], "te", "Planos DC-01"),
       ("Tapa de inspeccion: separacion de la parrilla de 3/8\" (m)", 0.15, "tsep", "Parrilla simple en ambos sentidos")]
P = {k: "PARAMETROS!$B$%d" % (5 + i) for i, (_, _, k, _) in enumerate(PAR)}


def construir():
    q = Q.por_eje(); dsc = Q.descargas()
    tapas = json.load(open(os.path.join(CALC, "tapas.json"))) if os.path.exists(os.path.join(CALC, "tapas.json")) else []
    wb = Workbook()
    wr = wb.active; wr.title = "RESUMEN"
    wp = wb.create_sheet("PLANILLA GENERAL DE METRADOS"); wc = wb.create_sheet("CONCRETO EN CUNETAS")
    we = wb.create_sheet("ENCOFRADO Y DESENCOFRADO"); wa = wb.create_sheet("METRADO ACERO"); wj = wb.create_sheet("METRADO JUNTA DE DILATACION")
    wu = wb.create_sheet("METRADO DE CURADO"); wg = wb.create_sheet("METRADO DE REJILLAS"); wt = wb.create_sheet("PLANILLA TAPAS INSPECCION")
    wq = wb.create_sheet("PARAMETROS"); wl = wb.create_sheet("Listas")
    hw = {x["eje"]: round(x["area_muro"] / x["L"], 3) for x in q}          # altura libre media del muro de cada eje

    # ======================================================================== PARAMETROS y Listas
    cel(wq, 1, 1, "PARÁMETROS DEL METRADO - RED DE CUNETAS", bold=True, size=14, borde=False); wq.merge_cells("A1:D1")
    cel(wq, 2, 1, "Valores editables: las hojas de detalle los usan en sus formulas.", al="left", borde=False)
    encabezado(wq, 4, ["Descripción", "Valor", "Unidad", "Criterio"], [62, 12, 9, 58])
    for i, (t, v, k, crit) in enumerate(PAR):
        r = 5 + i; u = "kg/m" if k == "kg" else ("-" if k == "esp" else "m")
        cel(wq, r, 1, t, al="left"); cel(wq, r, 2, v, "0.000", bold=True, bg=PAR_); cel(wq, r, 3, u); cel(wq, r, 4, crit, al="left")
    nota(wq, 6 + len(PAR), 1, 4, "Fuente: planos DC-01 a DC-19 (secciones tipicas DC-02, perfiles DC-03 a DC-11, secciones DC-12 a DC-14, descargas DC-16 a DC-19).")
    cel(wl, 1, 1, "LISTAS, FACTORES Y PARÁMETROS EDITABLES", bold=True, size=14, borde=False); wl.merge_cells("A1:F1")
    encabezado(wl, 3, ["DIÁMETRO", "FACTOR kg/m", "", "TIPO DE CUNETA"], [12, 13, 3, 18])
    for i, (d, f) in enumerate((("6 mm", 0.222), ("8 mm", 0.395), ("3/8\"", 0.56), ("12 mm", 0.888), ("1/2\"", 0.994), ("5/8\"", 1.552), ("3/4\"", 2.235), ("1\"", 3.973))):
        cel(wl, 4 + i, 1, d); cel(wl, 4 + i, 2, f, "0.000")
    for i, t in enumerate(("ABIERTA", "TAPADA", "MIXTA")): cel(wl, 4 + i, 4, t)
    cel(wl, 13, 1, "NOTA", bold=True, al="left", borde=False)
    cel(wl, 14, 1, "Los factores kg/m son editables. MIXTA: eje con tramos con rejilla (abiertos) y tramos tapados.", al="left", borde=False)

    # ======================================================================== CONCRETO EN CUNETAS (bloques por eje, dos columnas)
    cab_hoja(wc, "CONCRETO F´C=175 KG/CM2 PARA CUNETAS - DRENAJE PLUVIAL", 14, col0=2, fila=2)
    cel(wc, 7, 2, "OBRAS DE CONCRETO ARMADO", bold=True, color="FFFFFF", bg="0070C0"); wc.merge_cells("B7:O7")
    for col, w in zip("ABCDEFGHIJKLMNO", [3, 14, 17, 12, 11, 11, 13, 4, 4, 14, 17, 12, 11, 11, 13]): wc.column_dimensions[col].width = w
    ref_conc = {}
    for i, x in enumerate(q):
        c0 = 2 if i % 2 == 0 else 10; r0 = 9 + (i // 2) * 16
        cc = [CL(c0 + k) for k in range(6)]          # B..G o J..O
        for k, t in ((0, "LONGITUDES Y AREAS"), (3, "m3"), (4, "Sub.TOTAL"), (5, "TOTAL (m3)")):
            cel(wc, r0, c0 + k, t, bold=True, color="FFFFFF", bg="0070C0")
        wc.merge_cells("%s%d:%s%d" % (cc[0], r0, cc[2], r0))
        filas = [("MUROS", "AREA", "=%.2f*%.3f" % (x["L"], hw[x["eje"]]), "0.00"), (None, "ESPESOR", "=%s" % P["em"], "0.00"), (None, "CANT.", 2, "0"),
                 ("LOSA", "ESPESOR", "=%s" % P["ef"], "0.00"), (None, "LONGITUD", x["L"], "0.00"), (None, "ANCHO", "=%s" % P["be"], "0.00"),
                 ("TAPA", "LONGITUD", x["L_tapado"], "0.00"), (None, "ESPESOR", "=%s" % P["et"], "0.00"), (None, "ANCHO", "=%s" % P["be"], "0.00"),
                 ("DESC. TAPAS", "CANT.", -x["tapas"], "0"), (None, "LADO", "=%s" % P["tl"], "0.00"), (None, "ESPESOR", "=%s" % P["te"], "0.00")]
        for k, (grupo, nom, v, fmt) in enumerate(filas):
            r = r0 + 1 + k
            cel(wc, r, c0 + 1, grupo, bold=True); cel(wc, r, c0 + 2, nom); cel(wc, r, c0 + 3, v, fmt)
            if k % 3 == 0:
                a, b_, c_ = "%s%d" % (cc[3], r), "%s%d" % (cc[3], r + 1), "%s%d" % (cc[3], r + 2)
                f = "=%s*%s*%s" % (a, b_, c_) if k < 9 else "=%s*%s*%s*%s" % (a, b_, b_, c_)
                cel(wc, r, c0 + 4, f, "0.000", bold=True)
            else:
                cel(wc, r, c0 + 4, None)
            cel(wc, r, c0 + 5, None)
        cel(wc, r0 + 1, c0, "CUNETA %s" % x["eje"], bold=True, wrap=True)
        wc.merge_cells("%s%d:%s%d" % (cc[0], r0 + 1, cc[0], r0 + 12))
        cel(wc, r0 + 1, c0 + 5, "=SUM(%s%d:%s%d)" % (cc[4], r0 + 1, cc[4], r0 + 12), "0.000", bold=True, bg=PAR_)
        wc.merge_cells("%s%d:%s%d" % (cc[5], r0 + 1, cc[5], r0 + 12))
        for k in range(4):
            wc.merge_cells("%s%d:%s%d" % (cc[1], r0 + 1 + 3 * k, cc[1], r0 + 3 + 3 * k))
        ref_conc[x["eje"]] = ("%s%d" % (cc[5], r0 + 1), "%s%d" % (cc[3], r0 + 1))   # total del eje y area de muro
        if i % 4 == 3 and i < len(q) - 1: wc.row_breaks.append(Break(id=r0 + 14))

    # ======================================================================== ENCOFRADO Y DESENCOFRADO
    cab_hoja(we, "ENCOFRADO Y DESENCOFRADO CUNETAS - DRENAJE PLUVIAL", 15, col0=2, fila=2)
    cel(we, 8, 2, "OBRAS DE CONCRETO ARMADO - ENCOFRADO Y DESENCOFRADO", bold=True, color="FFFFFF", bg="0070C0"); we.merge_cells("B8:P8")
    for col, w in zip("ABCDEFGHIJKLMNOP", [3, 12, 12, 17, 16, 11, 12, 13, 4, 12, 12, 17, 16, 11, 12, 13]): we.column_dimensions[col].width = w
    ref_enc = {}
    for i, x in enumerate(q):
        c0 = 2 if i % 2 == 0 else 10; r0 = 11 + (i // 2) * 11
        cc = [CL(c0 + k) for k in range(7)]          # B..H o J..P
        tipo = "TAPADA" if x["L_rejilla"] <= 0 else ("MIXTA" if x["L_tapado"] > 0 else "ABIERTA")
        for k, t in ((0, "LONGITUDES Y AREAS"), (5, "Sub.TOTAL"), (6, "TOTAL (m2)")):
            cel(we, r0, c0 + k, t, bold=True, color="FFFFFF", bg="0070C0")
        we.merge_cells("%s%d:%s%d" % (cc[0], r0, cc[4], r0))
        h = hw[x["eje"]]
        filas = [("MURO EXTERNO", "AREA", "=%.2f*(%.3f+%s)" % (x["L"], h, P["ef"]), "=%s{r}*%s{r1}"), (None, "CANT.", 2, None),
                 ("MURO INTERNO", "AREA", "=%.2f*%.3f" % (x["L"], h), "=%s{r}*%s{r1}"), (None, "CANTIDAD", 2, None),
                 ("LOSA TAPA", "LONGITUD", x["L_tapado"], "=%s{r}*%s{r1}"), (None, "FONDO + CANTOS", "=%s+2*%s" % (P["b"], P["et"]), None)]
        for k, (grupo, nom, v, f) in enumerate(filas):
            r = r0 + 1 + k
            cel(we, r, c0 + 2, grupo, bold=True); cel(we, r, c0 + 3, nom); cel(we, r, c0 + 4, v, "0.00")
            cel(we, r, c0 + 5, (f % (cc[4], cc[4])).format(r=r, r1=r + 1) if f else None, "0.00", bold=True); cel(we, r, c0 + 6, None)
            if k % 2 == 0: we.merge_cells("%s%d:%s%d" % (cc[2], r, cc[2], r + 1))
        cel(we, r0 + 1, c0, x["eje"], bold=True); cel(we, r0 + 1, c0 + 1, tipo, bold=True)
        we.merge_cells("%s%d:%s%d" % (cc[0], r0 + 1, cc[0], r0 + 6)); we.merge_cells("%s%d:%s%d" % (cc[1], r0 + 1, cc[1], r0 + 6))
        cel(we, r0 + 1, c0 + 6, "=SUM(%s%d:%s%d)" % (cc[5], r0 + 1, cc[5], r0 + 6), "0.00", bold=True, bg=PAR_)
        we.merge_cells("%s%d:%s%d" % (cc[6], r0 + 1, cc[6], r0 + 6))
        ref_enc[x["eje"]] = "%s%d" % (cc[6], r0 + 1)
        if i % 6 == 5 and i < len(q) - 1: we.row_breaks.append(Break(id=r0 + 9))
    rr = 11 + ((len(q) + 1) // 2) * 11
    nota(we, rr, 2, 16, "Areas de muro = longitud x altura libre media del eje (perfil cada 0.50 m); cara externa con la losa de fondo (+0.10 m). "
        "Losa tapa: fondo (0.40 m) y cantos (2 x 0.10 m) en tramos tapados.")

    # ======================================================================== METRADO ACERO
    cab_hoja(wa, "METRADO DE ACERO EN CUNETAS - DRENAJE PLUVIAL", 23, lab=2)
    cel(wa, 8, 1, "METRADO DETALLADO POR EJE (ACERO SOLO CON h >= 0.40 m O EN TRAMO TAPADO)", bold=True, al="left", borde=False)
    grupos = [(1, 6, "IDENTIFICACIÓN Y GEOMETRÍA", "D9E1F2"), (7, 12, "ACERO LONGITUDINAL DEL CUERPO U", AZC),
              (13, 17, "ACERO TRANSVERSAL DEL CUERPO U", "E2EFDA"), (18, 22, "ACERO DE LA LOSA TAPA (SOLO TRAMOS TAPADOS)", "FCE4D6"), (23, 23, "RESULTADO", PAR_)]
    for a, b_, t, bg in grupos:
        cel(wa, 9, a, t, bold=True, bg=bg, wrap=True)
        if b_ > a: wa.merge_cells(start_row=9, start_column=a, end_row=9, end_column=b_)
    hdr = ["N°", "EJE", "TIPO DE CUNETA", "LONGITUD EJE (m)", "LONG. CON ACERO (m)", "h PROM. CON ACERO (m)",
           "Ø CUERPO", "FACTOR (kg/m)", "SEP. LONG. (m)", "N° BARRAS LONG.", "LONG. ACERO LONG. (m)", "KG LONG.",
           "SEP. TRANSV. U (m)", "DESARROLLO U (m)", "CANT. U", "LONG. TRANSV. (m)", "KG U",
           "LONG. TAPADA (m)", "SEP. TRANSV. (m)", "DESARROLLO (m)", "BARRAS LONG.", "KG LOSA TAPA", "TOTAL ACERO (kg)"]
    encabezado(wa, 10, hdr, [5, 9, 11, 10, 10, 10, 8, 9, 8, 9, 11, 10, 9, 10, 9, 11, 10, 10, 9, 10, 9, 10, 12], bg="5B9BD5", alto=45)
    fila_acero = {}
    for i, x in enumerate(q):
        r = 11 + i; fila_acero[x["eje"]] = r
        tipo = "TAPADA" if x["L_rejilla"] <= 0 else ("MIXTA" if x["L_tapado"] > 0 else "ABIERTA")
        vals = [(i + 1, "0"), (x["eje"], None), (tipo, None), (x["L"], "0.00"), (x["L_acero"], "0.00"), (x["h_acero"], "0.000"),
                ("3/8\"", None), ('=IFERROR(VLOOKUP(G{r},Listas!$A$4:$B$11,2,FALSE),0)', "0.000"), ("=%s" % P["sl"], "0.00"),
                ("=IF(E{r}>0,2*(CEILING(ROUND(F{r}/I{r},6),1)+1)+(CEILING(ROUND(({be}-2*{rec})/I{r},6),1)+1),0)", "0"),
                ("=E{r}*J{r}", "0.00"), ("=K{r}*H{r}", "0.00"),
                ("=%s" % P["st"], "0.00"), ("=IF(E{r}>0,2*(F{r}+{ef}-2*{rec})+({be}-2*{rec})+0.2,0)", "0.000"), ("=IF(E{r}>0,E{r}/M{r},0)", "0.00"),
                ("=N{r}*O{r}", "0.00"), ("=P{r}*H{r}", "0.00"),
                (x["L_tapado"], "0.00"), ("=%s" % P["sl"], "0.00"), ("=IF(R{r}>0,{be}-2*{rec}+0.2,0)", "0.000"), ("=IF(R{r}>0,3,0)", "0"),
                ("=IF(AND(E{r}>0,R{r}>0),((R{r}/S{r})*T{r}+U{r}*R{r})*H{r},0)", "0.00"), ("=L{r}+Q{r}+V{r}", "0.00")]
        for j, (v, fmt) in enumerate(vals, 1):
            if isinstance(v, str) and v.startswith("="): v = v.format(r=r, **P)
            cel(wa, r, j, v, fmt, bold=(j == 23), bg=PAR_ if j == 23 else None)
    rt = 11 + len(q); fila_acero["_total"] = rt
    cel(wa, rt, 1, "TOTAL", bold=True); wa.merge_cells(start_row=rt, start_column=1, end_row=rt, end_column=22)
    cel(wa, rt, 23, "=SUM(W11:W%d)" % (rt - 1), "0.00", bold=True, bg=PAR_)
    nota(wa, rt + 2, 1, 23, "Acero solo donde h >= 0.40 m (de la rejilla al fondo) o en tramo tapado; con h < 0.40 m la cuneta abierta es de concreto simple. "
        "U de 3/8\" @0.25 m con ganchos de 0.10 m; longitudinales de 3/8\" @0.20 m en muros y losa de fondo; losa tapa: 3/8\" @0.20 m transversal y 3 barras longitudinales.")

    # ======================================================================== METRADO JUNTA DE DILATACION
    cel(wj, 1, 1, "PLANILLA DE METRADOS - JUNTA DE DILATACIÓN E=1\" CON SELLADOR ASFALTICO", bold=True, size=14, borde=False); wj.merge_cells("A1:J1")
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        cel(wj, 2 + i, 1, k, bold=True, al="left"); cel(wj, 2 + i, 2, ref, al="left", wrap=True); wj.merge_cells(start_row=2 + i, start_column=2, end_row=2 + i, end_column=10)
    wj.row_dimensions[3].height = 33
    cel(wj, 7, 1, "PARÁMETROS", bold=True, al="left", borde=False)
    cel(wj, 8, 1, "Espaciamiento entre juntas (m)", al="left"); cel(wj, 8, 2, "=%s" % P["jc"], "0.00", bg=PAR_)
    cel(wj, 9, 1, "Junta en toda la seccion en U: 2 muros + losa de fondo", al="left"); cel(wj, 9, 2, "2 (h + ef) + be")
    cel(wj, 10, 1, "Espesor de junta (referencial, no interviene en el calculo)", al="left"); cel(wj, 10, 2, "1\" = 0.0254 m")
    encabezado(wj, 12, ["ITEM", "EJE / DESCRIPCIÓN", "UND.", "LONGITUD EJE (m)", "h PROM. (m)", "ESPACIAMIENTO (m)", "N° DE JUNTAS",
                        "DESARROLLO DE LA U (m)", "", "PARCIAL (ml)"], [48, 34, 7, 12, 10, 14, 10, 14, 3, 12], alto=45)
    cel(wj, 13, 1, "01.04.02.03", bold=True, bg=PAR_); cel(wj, 13, 2, "JUNTA DE DILATACION E=1\" CON SELLADOR ASFALTO", bold=True, bg=PAR_, al="left", wrap=True)
    cel(wj, 13, 3, "ml", bg=PAR_)
    for j in range(4, 10): cel(wj, 13, j, None, bg=PAR_)
    rj0 = 14
    for i, x in enumerate(q):
        r = rj0 + i
        vals = [(i + 1, "0"), (x["eje"], None), ("ml", None), (x["L"], "0.00"), (x["h_prom"], "0.000"), ("=$B$8", "0.00"),
                ("=INT(ROUND(D%d/F%d,6))" % (r, r), "0"), ("=2*(E%d+%s)+%s" % (r, P["ef"], P["be"]), "0.000"), (None, None), ("=G%d*H%d" % (r, r), "0.00")]
        for j, (v, fmt) in enumerate(vals, 1): cel(wj, r, j, v, fmt)
    rj1 = rj0 + len(q) - 1
    cel(wj, 13, 10, "=SUM(J%d:J%d)" % (rj0, rj1), "0.00", bold=True, bg=PAR_)
    r = rj1 + 3
    cel(wj, r, 1, "01.04.02.04", bold=True, bg=PAR_); cel(wj, r, 2, "JUNTA E=1\" CON SELLADOR ASFALTICO ENTRE CUNETA Y PISO ADYACENTE", bold=True, bg=PAR_, al="left", wrap=True)
    cel(wj, r, 3, "ml", bg=PAR_)
    for j in range(4, 10): cel(wj, r, j, None, bg=PAR_)
    rp = r; ok_p = [x for x in q if x["junta_piso"] > 0]
    encabezado(wj, r + 1, ["ITEM", "EJE", "UND.", "LONGITUD EJE (m)", "BORDE JUNTO A PISO (m)", "", "", "", "", "PARCIAL (ml)"], alto=32)
    for i, x in enumerate(ok_p):
        rr = r + 2 + i
        for j, (v, fmt) in enumerate([(i + 1, "0"), (x["eje"], None), ("ml", None), (x["L"], "0.00"), (x["junta_piso"], "0.00"),
                                       (None, None), (None, None), (None, None), (None, None), ("=E%d" % rr, "0.00")], 1):
            cel(wj, rr, j, v, fmt)
    rp1 = r + 1 + len(ok_p)
    cel(wj, rp, 10, "=SUM(J%d:J%d)" % (r + 2, rp1), "0.00", bold=True, bg=PAR_)
    nota(wj, rp1 + 2, 1, 10, "Borde junto a piso: longitud del borde de la cuneta que colinda con vereda, patio o piso (medida a cada lado del eje en la planta DC-01).")

    # ======================================================================== METRADO DE CURADO
    cel(wu, 1, 1, "METRADO DE CURADO DE CONCRETO EN CUNETAS", bold=True, size=14, borde=False, color="FFFFFF", bg=CAB); wu.merge_cells("A1:N1")
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        cel(wu, 2 + i, 1, k, bold=True, al="left"); cel(wu, 2 + i, 3, ref, al="left", wrap=True); wu.merge_cells(start_row=2 + i, start_column=3, end_row=2 + i, end_column=14)
        wu.merge_cells(start_row=2 + i, start_column=1, end_row=2 + i, end_column=2)
    wu.row_dimensions[3].height = 33
    encabezado(wu, 6, ["N°", "EJE", "TIPO DE CUNETA", "LONGITUD (m)", "LONG. CON REJILLA (m)", "LONG. TAPADA (m)", "ANCHO INTERIOR (m)", "ALTURA LIBRE PROM. (m)",
                       "ÁREA 2 MUROS (m²)", "ÁREA FONDO (m²)", "CORONAS DE MURO (m²)", "ÁREA TAPA (m²)", "TOTAL CURADO (m²)", "UND."],
               [6, 10, 13, 12, 12, 12, 12, 12, 13, 12, 13, 12, 13, 7], alto=45)
    ru0 = 7
    for i, x in enumerate(q):
        r = ru0 + i; tipo = "TAPADA" if x["L_rejilla"] <= 0 else ("MIXTA" if x["L_tapado"] > 0 else "ABIERTA")
        vals = [(i + 1, "0"), (x["eje"], None), (tipo, None), (x["L"], "0.00"), (x["L_rejilla"], "0.00"), (x["L_tapado"], "0.00"),
                ("=%s" % P["b"], "0.00"), (hw[x["eje"]], "0.000"), ("=2*D{r}*H{r}", "0.00"), ("=D{r}*G{r}", "0.00"),
                ("=2*{em}*E{r}", "0.00"), ("=F{r}*{be}", "0.00"), ("=SUM(I{r}:L{r})", "0.00"), ("m²", None)]
        for j, (v, fmt) in enumerate(vals, 1):
            if isinstance(v, str) and v.startswith("="): v = v.format(r=r, **P)
            cel(wu, r, j, v, fmt, bold=(j == 13), bg=VER if j == 13 else None)
    ru1 = ru0 + len(q); fila_cur = ru1
    cel(wu, ru1, 1, "TOTAL CURADO DE CONCRETO EN CUNETA", bold=True); wu.merge_cells(start_row=ru1, start_column=1, end_row=ru1, end_column=12)
    cel(wu, ru1, 13, "=SUM(M%d:M%d)" % (ru0, ru1 - 1), "0.00", bold=True, bg=PAR_); cel(wu, ru1, 14, "m²", bold=True)
    nota(wu, ru1 + 2, 1, 14, "Superficies expuestas: caras interiores de los dos muros, fondo interior, coronas de muro en tramos con rejilla y cara superior de la losa tapa.")

    # ======================================================================== METRADO DE REJILLAS
    cel(wg, 1, 1, "METRADO DE REJILLAS PARA CUNETAS", bold=True, size=14, borde=False, color="FFFFFF", bg=CAB); wg.merge_cells("A1:H1")
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        cel(wg, 2 + i, 1, k, bold=True, al="left"); cel(wg, 2 + i, 3, ref, al="left", wrap=True); wg.merge_cells(start_row=2 + i, start_column=3, end_row=2 + i, end_column=8)
        wg.merge_cells(start_row=2 + i, start_column=1, end_row=2 + i, end_column=2)
    wg.row_dimensions[3].height = 45
    encabezado(wg, 6, ["N°", "EJE", "TIPO DE CUNETA", "LONGITUD DEL EJE (m)", "DESCUENTO TRAMOS TAPADOS (m)", "LONGITUD CON REJILLA (m)", "N° DE TAPAS DE INSPECCIÓN", "UND."],
               [6, 10, 14, 14, 16, 15, 14, 7], alto=45)
    rg0 = 7; con_rej = [x for x in q if x["L_rejilla"] > 0]
    for i, x in enumerate(q):
        r = rg0 + i; tipo = "TAPADA" if x["L_rejilla"] <= 0 else ("MIXTA" if x["L_tapado"] > 0 else "ABIERTA")
        for j, (v, fmt) in enumerate([(i + 1, "0"), (x["eje"], None), (tipo, None), (x["L"], "0.00"), (x["L_tapado"], "0.00"),
                                       ("=IF(D{r}-E{r}<0,0,D{r}-E{r})".format(r=r), "0.00"), (x["tapas"], "0"), ("m", None)], 1):
            cel(wg, r, j, v, fmt, bold=(j == 6), bg=VER if j == 6 else None)
    rg1 = rg0 + len(q); fila_rej = rg1
    cel(wg, rg1, 1, "TOTAL REJILLA METALICA", bold=True); wg.merge_cells(start_row=rg1, start_column=1, end_row=rg1, end_column=5)
    cel(wg, rg1, 6, "=SUM(F%d:F%d)" % (rg0, rg1 - 1), "0.00", bold=True, bg=PAR_); cel(wg, rg1, 7, "=SUM(G%d:G%d)" % (rg0, rg1 - 1), "0", bold=True)
    cel(wg, rg1, 8, "m", bold=True)
    nota(wg, rg1 + 2, 1, 8, "Rejilla en todo tramo abierto (platinas de 1\" x 3/16\" sobre angulo de 1\" x 1\" x 3/16\"); en los tramos tapados va losa con tapas de inspeccion.")

    # ======================================================================== PLANILLA TAPAS INSPECCION
    cel(wt, 1, 1, "METRADO DE TAPAS DE INSPECCIÓN EN CUNETAS TAPADAS", bold=True, size=14, borde=False); wt.merge_cells("A1:K1")
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        cel(wt, 2 + i, 1, k, bold=True, al="left"); cel(wt, 2 + i, 2, ref, al="left", wrap=True); wt.merge_cells(start_row=2 + i, start_column=2, end_row=2 + i, end_column=11)
    wt.row_dimensions[3].height = 33
    encabezado(wt, 7, ["ITEM", "DESCRIPCIÓN", "Nº DE VECES", "LARGO", "ANCHO", "ALTURA", "FACT.", "AREA", "PARCIAL", "TOTAL", "UND"],
               [14, 52, 11, 10, 10, 10, 8, 10, 11, 11, 7])
    cel(wt, 8, 1, "01.04.02.02", bold=True, bg=PAR_); cel(wt, 8, 2, "TAPA DE INSPECCION 0.60m x 0.60m", bold=True, al="left", bg=PAR_)
    for j in range(3, 10): cel(wt, 8, j, None, bg=PAR_)
    cel(wt, 8, 11, "und", bold=True, bg=PAR_)
    rt0 = 9
    for i, t in enumerate(tapas):
        r = rt0 + i
        cel(wt, r, 1, i + 1, "0"); cel(wt, r, 2, "%s - progresiva %.2f m (tramo tapado)" % (t["eje"], t["progresiva"]), al="left"); cel(wt, r, 3, 1, "0")
        for j in range(4, 9): cel(wt, r, j, None)
        cel(wt, r, 9, "=C%d" % r, "0"); cel(wt, r, 10, None); cel(wt, r, 11, "und")
    rt1 = rt0 + len(tapas) - 1; fila_tapa = 8
    cel(wt, 8, 10, "=SUM(I%d:I%d)" % (rt0, rt1), "0", bold=True, bg=PAR_)
    r = rt1 + 2
    cel(wt, r, 2, "COMPOSICIÓN DE CADA TAPA (INCLUIDA EN LA UNIDAD)", bold=True, al="left", borde=False)
    comp = [("Concreto f'c=175 kg/cm2 en tapa", "=$J$8", "={tl}", "={tl}", "={te}", "=C{r}*D{r}*E{r}*F{r}", "m3"),
            ("Encofrado lateral de tapa", "=$J$8", "=4*{tl}", None, "={te}", "=C{r}*D{r}*F{r}", "m2"),
            ("Acero 3/8\" parrilla simple en ambos sentidos", "=$J$8", "=2*({tl}-2*{rec})", None, None,
             "=C{r}*D{r}*(CEILING(ROUND(({tl}-2*{rec})/{tsep},6),1)+1)*{kg}", "kg")]
    for i, (d, n, l, a, h, f, u) in enumerate(comp):
        rr = r + 1 + i
        cel(wt, rr, 2, d, al="left")
        for j, v in ((3, n), (4, l), (5, a), (6, h)):
            cel(wt, rr, j, v.format(r=rr, **P) if isinstance(v, str) else v, "0.000" if j > 3 else "0")
        cel(wt, rr, 9, f.format(r=rr, **P), "0.000"); cel(wt, rr, 11, u)
    nota(wt, r + 5, 2, 11, "Tapas a cada extremo de tramo tapado, en cada empalme y a 10 m como maximo entre tapas (planta DC-01). Asas de fierro liso de 3/8\".")

    # ======================================================================== PLANILLA GENERAL
    cel(wp, 2, 1, "PLANILLA GENERAL DE METRADOS - DRENAJE PLUVIAL", bold=True, size=14, borde=False); wp.merge_cells("A2:L2")
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        cel(wp, 3 + i, 1, k, bold=True, al="left"); cel(wp, 3 + i, 2, ref, al="left", wrap=True); wp.merge_cells(start_row=3 + i, start_column=2, end_row=3 + i, end_column=12)
    wp.row_dimensions[4].height = 33
    encabezado(wp, 8, ["ITEM", "DESCRIPCION", "Nº DE VECES", "LARGO", "ANCHO", "ALTURA", "FACT.", "AREA", "VOLUMEN", "PARCIAL", "TOTAL", "UND"],
               [14, 52, 10, 10, 10, 10, 8, 10, 10, 12, 13, 7])
    row = [9]; resumen = []

    def titulo(item, desc):
        r = row[0]; cel(wp, r, 1, item, bold=True, bg=SEC, al="left"); cel(wp, r, 2, desc, bold=True, bg=SEC, al="left")
        for j in range(3, 13): cel(wp, r, j, None, bg=SEC)
        resumen.append((item, desc, None, None)); row[0] += 1

    def partida(item, desc, und, filas, total=None):
        r0 = row[0]; cel(wp, r0, 1, item, bold=True, bg=PAR_, al="left"); cel(wp, r0, 2, desc, bold=True, bg=PAR_, al="left", wrap=True)
        for j in range(3, 11): cel(wp, r0, j, None, bg=PAR_)
        cel(wp, r0, 12, und, bold=True, bg=PAR_)
        row[0] += 1; ri = row[0]
        for f in filas:
            r = row[0]; d, C_, D_, E_, F_, G_, form = f
            cel(wp, r, 1, None); cel(wp, r, 2, d, al="left")
            for j, v in ((3, C_), (4, D_), (5, E_), (6, F_), (7, G_)):
                if isinstance(v, str) and v.startswith("="): v = v.format(r=r, **P)
                cel(wp, r, j, v, "0.00")
            cel(wp, r, 8, None); cel(wp, r, 9, None)
            cel(wp, r, 10, form.format(r=r, **P), "0.00"); cel(wp, r, 11, None); cel(wp, r, 12, None)
            row[0] += 1
        cel(wp, r0, 11, total if total else "=SUM(J%d:J%d)" % (ri, row[0] - 1), "0.00", bold=True, bg=PAR_)
        resumen.append((item, desc, und, "=+'PLANILLA GENERAL DE METRADOS'!K%d" % r0)); return r0

    L = lambda x: x["L"]
    titulo("01", "SISTEMA DE DRENAJE PLUVIAL - RED DE CUNETAS")
    titulo("01.01", "TRABAJOS PRELIMINARES")
    partida("01.01.01", "LIMPIEZA MANUAL DE TERRENO", "m²", [(x["eje"], 1, L(x), "={bexc}", None, None, "=C{r}*D{r}*E{r}") for x in q])
    partida("01.01.02", "TRAZO, NIVELES Y REPLANTEO PRELIMINAR", "m²", [(x["eje"], 1, L(x), "={bexc}", None, None, "=C{r}*D{r}*E{r}") for x in q])
    titulo("01.02", "MOVIMIENTO DE TIERRAS")
    r_exc = partida("01.02.01", "EXCAVACIÓN MANUAL DE ZANJAS P/CUNETAS", "m³",
                    [(x["eje"] + " (h prom. %.3f m)" % x["h_prom"], 1, L(x), "={bexc}", x["h_prom"], None, "=C{r}*D{r}*E{r}*(F{r}+{ef}+{sol})") for x in q])
    partida("01.02.02", "REFINE, NIVELACIÓN Y COMPACTACIÓN EN TERRENO NORMAL", "m²", [(x["eje"], 1, L(x), "={be}", None, None, "=C{r}*D{r}*E{r}") for x in q])
    r_rel = partida("01.02.03", "RELLENO COMPACTADO C/MATERIAL PROPIO SELECCIONADO", "m³",
                    [(x["eje"] + " (a los lados de la cuneta)", 1, L(x), "={bexc}-{be}", x["h_prom"], None, "=C{r}*D{r}*E{r}*(F{r}+{ef}+{sol})") for x in q])
    partida("01.02.04", "ELIMINACIÓN DE MATERIAL EXCEDENTE D<1Km", "m³",
            [("(Excavacion - relleno) x esponjamiento", None, None, None, None, "={esp}", "=(K%d-K%d)*G{r}" % (r_exc, r_rel))])
    partida("01.02.05", "PERFILADO MANUAL DE TERRENO EN DESCARGAS", "m²",
            [("Descarga %s (C.F. %.3f)" % (d["D"], d["cf"]), 1, d["sale_a_m"], "={be}", None, None, "=C{r}*D{r}*E{r}") for d in dsc])
    titulo("01.03", "OBRAS DE CONCRETO SIMPLE")
    partida("01.03.01", "CONCRETO f'c=100 kg/cm2 PARA SOLADO E=4''", "m²", [(x["eje"], 1, L(x), "={be}", None, None, "=C{r}*D{r}*E{r}") for x in q])
    titulo("01.04", "OBRAS DE CONCRETO ARMADO")
    titulo("01.04.01", "CUNETAS")
    partida("01.04.01.01", "CONCRETO f'c=175 kg/cm2 PARA CUNETAS", "m³",
            [(x["eje"] + " (hoja CONCRETO EN CUNETAS)", None, None, None, None, None, "=+'CONCRETO EN CUNETAS'!%s" % ref_conc[x["eje"]][0]) for x in q])
    partida("01.04.01.02", "ENCOFRADO Y DESENCOFRADO DE CUNETAS", "m²",
            [(x["eje"] + " (hoja ENCOFRADO Y DESENCOFRADO)", None, None, None, None, None, "=+'ENCOFRADO Y DESENCOFRADO'!%s" % ref_enc[x["eje"]]) for x in q])
    partida("01.04.01.03", "ACERO f'y=4200 kg/cm2 PARA CUNETA", "kg",
            [(x["eje"] + " (hoja METRADO ACERO)", None, None, None, None, None, "=+'METRADO ACERO'!W%d" % fila_acero[x["eje"]]) for x in q])
    partida("01.04.01.04", "CURADO DE CONCRETO EN CUNETA", "m²",
            [(x["eje"] + " (hoja METRADO DE CURADO)", None, None, None, None, None, "=+'METRADO DE CURADO'!M%d" % (ru0 + i)) for i, x in enumerate(q)])
    titulo("01.04.02", "VARIOS")
    partida("01.04.02.01", "REJILLA METALICA P/CUNETA C/PL 1''x 3/16'' Y ANGULO 1'' x 1'' x 3/16''", "m",
            [(x["eje"] + " (hoja METRADO DE REJILLAS)", None, None, None, None, None, "=+'METRADO DE REJILLAS'!F%d" % (rg0 + i))
             for i, x in enumerate(q) if x["L_rejilla"] > 0])
    partida("01.04.02.02", "TAPA DE INSPECCION 0.60m x 0.60m", "und",
            [("Tapas de inspeccion (hoja PLANILLA TAPAS INSPECCION)", None, None, None, None, None, "=+'PLANILLA TAPAS INSPECCION'!J8")])
    partida("01.04.02.03", "JUNTA DE DILATACION E=1\" CON SELLADOR ASFALTO", "m",
            [(x["eje"] + " (hoja METRADO JUNTA DE DILATACION)", None, None, None, None, None, "=+'METRADO JUNTA DE DILATACION'!J%d" % (rj0 + i))
             for i, x in enumerate(q) if x["n_juntas"]])
    partida("01.04.02.04", "JUNTA E=1\" CON SELLADOR ASFALTICO ENTRE CUNETA Y PISO ADYACENTE", "m",
            [(x["eje"] + " (hoja METRADO JUNTA DE DILATACION)", None, None, None, None, None, "=+'METRADO JUNTA DE DILATACION'!J%d" % (rp + 2 + i))
             for i, x in enumerate(ok_p)])

    # ======================================================================== RESUMEN
    cel(wr, 2, 1, "RESUMEN METRADOS DE DRENAJE PLUVIAL", bold=True, size=14, borde=False); wr.merge_cells("A2:D2")
    for i, (k, v) in enumerate((("ENTIDAD", ENTIDAD), ("PROYECTO", PROYECTO), ("UBICACIÓN", UBIC), ("FECHA", FECHA))):
        cel(wr, 4 + i, 1, k, bold=True); cel(wr, 4 + i, 2, v, al="left", wrap=True); wr.merge_cells(start_row=4 + i, start_column=2, end_row=4 + i, end_column=4)
    wr.row_dimensions[5].height = 48
    encabezado(wr, 8, ["Item", "Descripción", "Unid.", "Total."], [16, 74, 9, 15], alto=20)
    for i, (it, d, u, f) in enumerate(resumen):
        r = 9 + i; t = f is None
        cel(wr, r, 1, it, "@", bold=t, bg=SEC if t else None, al="left"); cel(wr, r, 2, d, bold=t, bg=SEC if t else None, al="left", wrap=True)
        cel(wr, r, 3, u, bg=SEC if t else None); cel(wr, r, 4, f, "#,##0.00", bg=SEC if t else None)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    for ws in (wr, wj, wg, wt, wq, wl): ws.page_setup.orientation = "portrait"
    import caratula_b; caratula_b.excel(wb, "metrado")                # caratula comun (foto aerea del local)
    wb.save(SAL); print(SAL)
    return resumen


if __name__ == "__main__":
    construir()
