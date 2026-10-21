"""Anexo de analisis estadistico de la memoria hidrologica (I.E. N.° 0590 Barranquita) con el formato de los cuadros del estudio
hidrologico del CAR Mujeres (CUI 2717013): cada hoja lleva Proyecto / Entidad / Ubicacion / Fecha.
PARTE I: A. registro de precipitaciones maximas en 24 h; B. modelos de distribucion (B.1 Normal, B.2 Log Normal,
B.3 Log Pearson III, B.4 Gumbel); C. prueba de Kolmogorov-Smirnov y C.4 resumen.
PARTE II: A. parametros de intensidades maximas (Cuadros N.° 04 a 06); B. calculo de intensidades I = K T^m / t^n;
C. intensidades de diseno para T = 25 anos.
"""
import math
import numpy as np
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import hidrologia_b as H
import word_b as W

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
DN = {"normal": "Normal", "lognormal": "Log Normal", "lp3": "Log Pearson III", "gumbel": "Gumbel"}
AZ, CL, AM, GR = "1F3A68", "D9E2F3", "FFF2CC", "F2F2F2"


def _ancho(t, anchos):
    t.autofit = False
    tblW = t._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW"); t._tbl.tblPr.append(tblW)
    tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(int(sum(anchos) * 567)))
    for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(anchos[j] * 567)))
    for row in t.rows:
        for j, w in enumerate(anchos): row.cells[j].width = Cm(w)


class Hoja:
    """Tabla tipo hoja de calculo: celdas por (fila, columna), combinaciones y estilos; se vuelca a Word al final."""

    def __init__(self, ncol, anchos, tam=7.5):
        self.nc, self.anchos, self.tam = ncol, anchos, tam
        self.c = {}; self.m = []; self.r = 0

    def put(self, r, c, v, neg=False, bg=None, al="c", blanco=False, c1=None):
        self.c[(r, c)] = (v, neg, bg, al, blanco)
        if c1 is not None and c1 > c: self.m.append((r, c, r, c1))
        self.r = max(self.r, r + 1)

    def fila(self, valores, c0=0, **k):
        r = self.r
        for j, v in enumerate(valores): self.put(r, c0 + j, v, **k)
        return r

    def cabecera(self, titulo_parte=None, lab=2):
        for i, (k, v) in enumerate((("Proyecto", W.PROYECTO + " - " + W.CI), ("Entidad", W.ENTIDAD), ("Ubicación", "SAN MARTÍN - LAMAS - BARRANQUITA"),
                                    ("Fecha", W.FECHA))):
            self.put(i, 0, k, neg=True, al="l", c1=lab - 1); self.put(i, lab, v, al="l", c1=self.nc - 1)
        self.r = 5
        if titulo_parte:
            self.put(self.r, 0, titulo_parte, neg=True, bg=AZ, blanco=True, al="l", c1=self.nc - 1); self.r += 1

    def titulo(self, t, c0=0, c1=None, r=None):
        r = self.r if r is None else r
        self.put(r, c0, t, neg=True, bg=CL, al="l", c1=self.nc - 1 if c1 is None else c1)
        return r

    def volcar(self, D):
        t = D.d.add_table(rows=self.r, cols=self.nc); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        tapadas = {(r, c) for r0, c0, r1, c1 in self.m for r in range(r0, r1 + 1) for c in range(c0, c1 + 1) if (r, c) != (r0, c0)}
        for (r, c), (v, neg, bg, al, blanco) in self.c.items():
            if (r, c) in tapadas: continue                     # celdas cubiertas por una combinacion: sin texto
            cc = t.cell(r, c); cc.text = ""
            if v is None or v == "":
                if bg: W._sombra(cc, bg)
                continue
            run = cc.paragraphs[0].add_run(str(v)); run.font.size = Pt(self.tam); run.bold = neg
            if blanco: run.font.color.rgb = W.RGBColor(0xFF, 0xFF, 0xFF)
            cc.paragraphs[0].alignment = {"c": WD_ALIGN_PARAGRAPH.CENTER, "l": WD_ALIGN_PARAGRAPH.LEFT, "r": WD_ALIGN_PARAGRAPH.RIGHT}[al]
            cc.paragraphs[0].paragraph_format.space_after = Pt(0)
            if bg: W._sombra(cc, bg)
        for r0, c0, r1, c1 in self.m:
            t.cell(r0, c0).merge(t.cell(r1, c1))
        for row in t.rows:
            for cc in row.cells:
                for p in cc.paragraphs:
                    p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0); p.paragraph_format.line_spacing = 1.0
        _ancho(t, self.anchos)
        D.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t


def f(v, d=2): return ("%." + str(d) + "f") % (round(v, d) + 0.0)


def w_z(T):
    p = 1.0 / T; w = math.sqrt(math.log(1.0 / p ** 2))
    z = w - (2.515517 + 0.802853 * w + 0.010328 * w * w) / (1 + 1.432788 * w + 0.189269 * w * w + 0.001308 * w ** 3)
    return p, w, z


def anexo(D, HID):
    A = HID["analisis"]; e = H.ESTACION; T_ = H.PERIODOS; el = A["elegida"]
    # ---------------------------------------------------------------------------------------------- PARTE I - A. registro
    D.p("A continuación se presentan las hojas de cálculo del análisis estadístico de la precipitación máxima en 24 horas y de las "
        "intensidades de diseño, con el mismo ordenamiento del estudio hidrológico de los proyectos de Morales.")
    h = Hoja(14, [1.15] + [1.03] * 12 + [1.35], tam=6.5)
    h.cabecera("PARTE I : DETERMINACIÓN DE LA PRECIPITACIÓN MÁXIMA DE DISEÑO")
    h.titulo("A. REGISTRO DE PRECIPITACIONES MÁXIMAS EN 24 HORAS - ESTACIÓN %s (SENAMHI, CÓDIGO %s)" % (e["nombre"].upper(), e["codigo"])); h.r += 1
    for k1, v1, k2, v2 in (("Estación", e["nombre"], "Dpto.", e["departamento"].upper()), ("Tipo", e["tipo"], "Provincia", e["provincia"].upper()),
                           ("Altura", "%d msnm." % e["altitud"], "Distrito", e["distrito"].upper())):
        r = h.r; h.put(r, 0, k1, neg=True, al="l"); h.put(r, 1, v1, al="l", c1=4); h.put(r, 6, k2, neg=True, al="l", c1=7); h.put(r, 8, ": " + v2, al="l", c1=13)
    h.titulo("Precipitación Máxima en 24 Horas (mm) - periodo %s (%d años completos más recientes del registro diario)" % (e["periodo"], e["anios"]))
    pass
    h.fila(["Año"] + MESES + ["P Máxima"], neg=True, bg=CL)
    for a, ms in H.SERIE.items():
        vals = [p for p in ms if p is not None]
        h.fila([a] + [("%.1f" % p) if p is not None else "S/D" for p in ms] + [f(max(vals))])
        h.put(h.r - 1, 13, f(max(vals)), neg=True, bg=AM)
    h.put(h.r, 0, "Fuente: SENAMHI, datos diarios con control de calidad de la estación %s. S/D: mes sin dato completo (se conservan los "
          "máximos de los otros meses)." % e["nombre"], al="l", c1=13)
    h.volcar(D)
    # ---------------------------------------------------------------------------------------------- PARTE I - B. modelos
    D.salto()
    h = Hoja(11, [2.0, 1.25, 1.15, 0.25, 1.2, 1.3, 1.15, 1.15, 1.2, 1.3, 1.4], tam=6.5)
    h.cabecera("B. MODELOS DE DISTRIBUCIÓN DE DATOS ESTADÍSTICOS")
    r0 = h.r
    h.put(r0, 0, "P máx. anual (mm); T en años", neg=True, bg=CL, c1=2)
    h.put(r0 + 1, 0, "Año", neg=True, bg=CL); h.put(r0 + 1, 1, "P máx", neg=True, bg=CL); h.put(r0 + 1, 2, "Log P", neg=True, bg=CL)
    for i, (a, p) in enumerate(A["serie"]):
        h.put(r0 + 2 + i, 0, a); h.put(r0 + 2 + i, 1, f(p)); h.put(r0 + 2 + i, 2, f(math.log10(p), 3))
    x = [p for _, p in A["serie"]]; lx = [math.log10(p) for p in x]
    rs = r0 + 2 + len(x)
    for k, (t1, v1, v2) in enumerate((("m (media)", f(A["media"]), f(A["media_log"], 3)), ("var (varianza)", f(A["desv"] ** 2), f(A["desv_log"] ** 2, 4)),
                                       ("s (desv. est.)", f(A["desv"], 3), f(A["desv_log"], 3)), ("Coef. asimetría", f(A["cs"], 3), f(A["cs_log"], 3)),
                                       ("k (factor k)", "", f(A["cs_log"] / 6, 3)), ("n (n° de datos)", A["n"], A["n"]))):
        h.put(rs + k, 0, t1, neg=True, al="l", bg=GR); h.put(rs + k, 1, v1, bg=GR); h.put(rs + k, 2, v2, bg=GR)
    # columna derecha: B.1, B.2 y B.3
    r = r0
    h.put(r, 4, "B.1. Distribución Normal", neg=True, bg=CL, c1=10); r += 1
    for j, t in enumerate(["T", "P", "w", "z", "P (mm)"]): h.put(r, 4 + j, t, neg=True, bg=CL)
    r += 1
    for T in T_:
        p, w, z = w_z(T)
        for j, v in enumerate([T, f(p, 4), f(w, 3), f(z, 3), f(A["tabla"][T]["normal"], 1)]): h.put(r, 4 + j, v)
        r += 1
    h.put(r, 4, "B.2. Distribución Log Normal", neg=True, bg=CL, c1=10); r += 1
    for j, t in enumerate(["T", "P", "w", "z", "Log P", "P (mm)"]): h.put(r, 4 + j, t, neg=True, bg=CL)
    r += 1
    for T in T_:
        p, w, z = w_z(T); P = A["tabla"][T]["lognormal"]
        for j, v in enumerate([T, f(p, 4), f(w, 3), f(z, 3), f(math.log10(P), 3), f(P, 1)]): h.put(r, 4 + j, v)
        r += 1
    h.put(r, 4, "B.3. Distribución Log Pearson III", neg=True, bg=CL, c1=10); r += 1
    for j, t in enumerate(["T", "P", "w", "z", "KT", "Log P", "P (mm)"]): h.put(r, 4 + j, t, neg=True, bg=CL)
    r += 1
    for T in T_:
        p, w, z = w_z(T); P = A["tabla"][T]["lp3"]
        for j, v in enumerate([T, f(p, 4), f(w, 3), f(z, 2), f(A["tabla"][T]["kt_lp3"], 3), f(math.log10(P), 3), f(P, 1)]): h.put(r, 4 + j, v)
        r += 1
    # B.4 Gumbel a la izquierda, bajo los estadisticos de la serie
    rg = rs + 7
    h.put(rg, 0, "B.4. Gumbel (Extrema Tipo I)", neg=True, bg=CL, c1=2)
    for j, t in enumerate(["T", "KT", "P (mm)"]): h.put(rg + 1, j, t, neg=True, bg=CL)
    for i, T in enumerate(T_):
        for j, v in enumerate([T, f(A["tabla"][T]["kt_gumbel"], 4), f(A["tabla"][T]["gumbel"], 1)]): h.put(rg + 2 + i, j, v)
    rg2 = rg + 2 + len(T_)
    h.put(rg2, 0, "α = √6·s/π"); h.put(rg2, 1, f(A["gumbel_alfa"], 3)); h.put(rg2 + 1, 0, "u = x̄ − 0.5772·α"); h.put(rg2 + 1, 1, f(A["gumbel_u"], 3))
    h.r = max(h.r, r)
    h.volcar(D)
    # ---------------------------------------------------------------------------------------------- PARTE I - C. K-S y resumen
    D.salto()
    h = Hoja(11, [1.8, 1.7, 1.4, 0.25, 1.1, 1.25, 1.25, 1.5, 1.2, 1.2, 1.3], tam=6.5)
    h.cabecera("C. PRUEBA DE BONDAD DE AJUSTE Y RESUMEN DE RESULTADOS")
    h.titulo("C.1 a C.3 Prueba de Kolmogorov-Smirnov (P de excedencia observada m/(n+1); α = 0.05; Dα = [−ln(α/2)/(2n)]^0.5 − 1/(6n))")
    h.fila(["Distribución", "Δ máx", "", "", "Δ crítico", "", "Resultado", "", "", "", ""], neg=True, bg=CL)
    r = h.r - 1
    h.m += [(r, 0, r, 0), (r, 1, r, 3), (r, 4, r, 5), (r, 6, r, 10)]
    for k in ("normal", "lognormal", "lp3", "gumbel"):
        r = h.fila([DN[k], f(A["ks"][k], 4), "", "", f(A["ks_critico"], 4), "", ("SE ACEPTA" if A["ks"][k] < A["ks_critico"] else "NO SE ACEPTA")
                    + (" - MENOR Δ (ADOPTADA)" if k == el else ""), "", "", "", ""], neg=(k == el), bg=AM if k == el else None)
        h.m += [(r, 1, r, 3), (r, 4, r, 5), (r, 6, r, 10)]
    h.r += 1
    h.titulo("C.4 RESUMEN DE LOS RESULTADOS DE LAS DISTRIBUCIONES (P máx. 24 h en mm)")
    r0 = h.r
    h.put(r0, 0, "Normales", neg=True, bg=CL, c1=1); h.put(r0, 2, "Valor", neg=True, bg=CL)
    for j, t in enumerate(["T (años)", "Normal", "Log Normal", "Log Pearson III", "Gumbel", "Diseño", ""]): h.put(r0, 4 + j, t, neg=True, bg=CL)
    h.m.append((r0, 9, r0, 10))
    n = A["n"]; vcs = 6.0 / n                       # como la hoja D. Resumen de los CAR
    norm = [("Media aritmética", f(A["media"], 3)), ("Desviación estándar", f(A["desv"], 3)), ("Coeficiente de asimetría", f(A["cs"], 3)),
            ("Número de datos", n), ("Variación del coef. de asimetría", f(vcs, 3)), ("Desv. est. coef. de asimetría", f(math.sqrt(vcs), 3))]
    loga = [("Media aritmética", f(A["media_log"], 3)), ("Desviación estándar", f(A["desv_log"], 3)), ("Varianza", f(A["desv_log"] ** 2, 4)),
            ("Coeficiente de asimetría", f(A["cs_log"], 3)), ("Número de datos", n), ("K", f(A["cs_log"] / 6, 3))]
    for i, (t, v) in enumerate(norm): h.put(r0 + 1 + i, 0, t, al="l", c1=1); h.put(r0 + 1 + i, 2, v)
    rl = r0 + 2 + len(norm)
    h.put(rl, 0, "Logarítmica", neg=True, bg=CL, c1=1); h.put(rl, 2, "Valor", neg=True, bg=CL)
    for i, (t, v) in enumerate(loga): h.put(rl + 1 + i, 0, t, al="l", c1=1); h.put(rl + 1 + i, 2, v)
    for i, T in enumerate(T_):
        r = r0 + 1 + i; tb = A["tabla"][T]
        for j, v in enumerate([T, f(tb["normal"]), f(tb["lognormal"]), f(tb["lp3"]), f(tb["gumbel"])]): h.put(r, 4 + j, v)
        h.put(r, 9, f(tb[el]), neg=True, bg=AM if T == H.T_DISENO else None, c1=10)
    h.r = max(rl + 1 + len(loga), r0 + 1 + len(T_)) + 1
    h.put(h.r, 0, "Distribución adoptada: %s (menor Δ de Kolmogorov-Smirnov). Precipitación máxima en 24 h de diseño para T = %d años: %.2f mm."
          % (DN[el], H.T_DISENO, HID["P24"]), neg=True, al="l", c1=10)
    h.volcar(D)
    # ---------------------------------------------------------------------------------------------- PARTE II - A. parametros
    D.salto()
    DUR = [5, 10, 15, 20, 30, 60]; Td = sorted(T_, reverse=True)
    h = Hoja(8, [1.5, 1.9, 1.9, 1.9, 1.9, 1.9, 1.9, 1.9], tam=7)
    h.cabecera("PARTE II: DETERMINACIÓN DE LA TORMENTA DE DISEÑO")
    h.titulo("A. PARÁMETROS DE INTENSIDADES MÁXIMAS")
    for nc, tit, fn, fmt in (("04", "Lluvias máximas (mm)", lambda T, d: HID["bell_p"][T][d], "%.2f"),
                             ("05", "Intensidades máximas (mm/hora)", lambda T, d: HID["bell_i"][T][d], "%.1f")):
        h.r += 1; h.put(h.r, 0, "Cuadro N° %s: %s" % (nc, tit), neg=True, al="l", c1=7)
        r = h.r; h.put(r, 0, "T", neg=True, bg=CL); h.put(r, 1, "P. Máx", neg=True, bg=CL); h.put(r, 2, "Duración en minutos", neg=True, bg=CL, c1=7)
        pass
        h.fila(["años", "24 horas"] + DUR, neg=True, bg=CL)
        for T in Td:
            P = HID["p24"][T]; h.fila([T, f(P)] + [fmt % fn(T, d) for d in DUR])
        h.put(h.r, 0, "Fuente: Modelo de Bell, P60(10) = 0.4602 · P24(10)^0.876 = %.2f mm (distribución %s)." % (HID["P60_10"], DN[el]), al="l", c1=7)
    h.volcar(D)
    # Cuadro 06: regresion lineal multiple log I = log K + m log T - n log t
    RG = HID["regresion"]; M, V, Mi, sol = np.array(RG["M"]), np.array(RG["V"]), np.array(RG["inv"]), RG["sol"]     # intensidades del Cuadro N° 05
    h = Hoja(8, [3.2, 1.8, 1.8, 1.8, 1.8, 1.6, 1.6, 1.8], tam=7)
    h.put(0, 0, "Cuadro N° 06: Ecuaciones de la regresión lineal múltiple (X1 = log T, X2 = log t, Y = log I)", neg=True, bg=CL, c1=7); h.r = 1
    for t1, t2 in (("Ecuación de la regresión lineal múltiple", "m(X1) − n(X2) + logK = Y"),
                   ("Sistema de ecuaciones normalizadas (1)", "mΣ(X1) + nΣ(X2) + (logK)(N) = Σ(Y)"),
                   ("Sistema de ecuaciones normalizadas (2)", "mΣ(X1²) + nΣ(X1)(X2) + (logK)ΣX1 = Σ(X1)(Y)"),
                   ("Sistema de ecuaciones normalizadas (3)", "mΣ(X1)(X2) + nΣ(X2²) + (logK)ΣX2 = Σ(X2)(Y)")):
        r = h.r; h.put(r, 0, t1, al="l"); h.put(r, 1, t2, al="l", c1=7)
    h.r += 1
    h.fila(["Matriz del sistema", "m", "n", "logK", "Valor", "Inversa de la matriz", "", ""], neg=True, bg=CL)
    h.m.append((h.r - 1, 5, h.r - 1, 7))
    for i in range(3):
        h.fila(["Ecuación (%d)" % (i + 1), f(M[i, 0]), f(M[i, 1]), f(M[i, 2]), f(V[i]), f(Mi[i, 0], 4), f(Mi[i, 1], 4), f(Mi[i, 2], 4)])
    h.r += 1
    h.fila(["Resultado", "m = %.4f" % sol[0], "n = %.4f" % (-sol[1]), "logK = %.4f" % sol[2], "K = %.2f" % HID["K"], "", "", ""], neg=True, bg=AM)
    h.m.append((h.r - 1, 5, h.r - 1, 7))
    h.volcar(D)
    # ---------------------------------------------------------------------------------------------- PARTE II - B. intensidades
    D.salto()
    K, m_, n_ = HID["K"], HID["m"], HID["n"]
    TT = [2, 5, 10, 20, 25, 50, 100]
    h = Hoja(8, [2.4, 1.75, 1.75, 1.75, 1.75, 1.75, 1.75, 1.75], tam=7)
    h.cabecera("B. CÁLCULO DE INTENSIDADES MÁXIMAS (mm/h)")
    r = h.r; h.put(r, 0, "I = K · T^m / t^n", neg=True, al="l", c1=7)
    for t, v in (("K", f(K)), ("m", f(m_, 4)), ("n", f(n_, 4))): h.fila([t, v] + [""] * 6, neg=False); h.m.append((h.r - 1, 1, h.r - 1, 7))
    r = h.r; h.put(r, 0, "Duración (t)", neg=True, bg=CL); h.put(r, 1, "Período de retorno (T) en años", neg=True, bg=CL, c1=7)
    h.fila(["(minutos)"] + TT, neg=True, bg=CL)
    for d in range(10, 130, 10):
        h.fila([d] + [f(K * T ** m_ / d ** n_) for T in TT])
    h.put(h.r, 0, "Ecuación de intensidad válida para duraciones de 5 a 60 minutos (regresión del Cuadro N° 06); de 70 a 120 minutos es extrapolación.",
          al="l", c1=7)
    h.volcar(D)
    # ---------------------------------------------------------------------------------------------- PARTE II - C. diseno
    h = Hoja(4, [3.6, 3.0, 3.0, 3.0], tam=7.5)
    h.cabecera("C. INTENSIDADES DE DISEÑO PARA PRECIPITACIÓN MÁXIMA (T = %d AÑOS)" % H.T_DISENO, lab=1)
    r = h.r; h.put(r, 0, "Precipitación Pmax 24 horas, T = %d años (mm)" % H.T_DISENO, neg=True, al="l", c1=2); h.put(r, 3, f(HID["P24"]), neg=True, bg=AM)
    h.fila(["D (min)", "D (horas)", "P (mm)", "I (mm/h)"], neg=True, bg=CL)
    for d in (10, 20, 30, 40, 50, 60, 90, 120, 240, 360, 420, 480, 600, 660, 720, 1440):
        P = H.p_dur(HID["P24"], d); r = h.fila([d, f(d / 60.0), f(P), f(P * 60.0 / d)])
        if d == int(H.TC):
            for j in range(4): h.put(r, j, h.c[(r, j)][0], neg=True, bg=AM)
    h.put(h.r, 0, "Intensidad de diseño (tc = %d min): I = %.2f mm/h. Fuente: elaboración propia con base en SENAMHI - estación %s."
          % (int(H.TC), HID["I"], e["nombre"]), neg=True, al="l", c1=3)
    h.volcar(D)
    return {"K": K, "m": m_, "n": n_}
