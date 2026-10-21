"""Caratula comun de los documentos de drenaje pluvial de la I.E. N.° 0590 Barranquita (C.I. 468901), con la foto aerea del local.

Foto del usuario: insumos/barranquita/foto_aerea_ie0590.webp (recortada sin las esquinas oscuras del lente en
entregables_barranquita/figuras_memoria/foto_caratula.jpg).
Se usa en: memoria hidrologica y memoria descriptiva (word_b), especificaciones tecnicas (docx_espec.js), metrado y memoria de
calculo (openpyxl) y analisis estadistico (LibreOffice UNO). Mismo contenido y orden en todos:
entidad, expediente tecnico, proyecto, C.I., foto, franja con el titulo del documento, especialidad, ubicacion y fecha.
"""
import os
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))

FOTO_USUARIO = os.path.join(RAIZ, "insumos", "barranquita", "foto_aerea_ie0590.webp")
FOTO = os.path.join(RAIZ, "entregables_barranquita", "figuras_memoria", "foto_caratula.jpg")
RECORTE = (45, 30, 1555, 1150)                 # sin las esquinas oscuras del lente (foto de 1600 x 1198)

ENTIDAD = "GOBIERNO REGIONAL SAN MARTÍN"
GERENCIA = "GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO"
EXPEDIENTE = "EXPEDIENTE TÉCNICO"
PROYECTO = ("“MEJORAMIENTO Y AMPLIACIÓN DEL SERVICIO DE EDUCACIÓN PRIMARIA EN I.E. N° 0590 DISTRITO DE BARRANQUITA "
            "DE LA PROVINCIA DE LAMAS DEL DEPARTAMENTO DE SAN MARTÍN”")
CI = "C.I. N° 468901"
LEYENDA = "Vista aérea del local de la I.E. N° 0590 - Barranquita"
ESPECIALIDAD = "ESPECIALIDAD: INSTALACIONES SANITARIAS - DRENAJE PLUVIAL"
UBICACION = "Distrito de Barranquita  -  Provincia de Lamas  -  Departamento de San Martín"
FECHA = "OCTUBRE 2026"
AZUL, AZUL_CLARO, GRIS = "1F3A68", "D9E2F3", "595959"

DOCS = {
    "metrado": ("METRADO", "DRENAJE PLUVIAL - RED DE CUNETAS"),
    "hidrologica": ("ESTUDIO HIDROLÓGICO Y DE DRENAJE PLUVIAL", "MEMORIA HIDROLÓGICA - RED DE CUNETAS"),
    "descriptiva": ("MEMORIA DESCRIPTIVA", "DRENAJE PLUVIAL - RED DE CUNETAS"),
    "calculo": ("MEMORIA DE CÁLCULO", "HIDROLOGÍA E HIDRÁULICA DE LA RED DE CUNETAS"),
    "especificaciones": ("ESPECIFICACIONES TÉCNICAS", "DRENAJE PLUVIAL - RED DE CUNETAS"),
    "analisis": ("ANÁLISIS ESTADÍSTICO", "PRECIPITACIÓN MÁXIMA EN 24 HORAS Y TORMENTA DE DISEÑO"),
}


def foto():
    """Ruta de la foto recortada (la prepara si falta o si la foto del usuario es mas reciente)."""
    if not os.path.exists(FOTO) or (os.path.exists(FOTO_USUARIO) and os.path.getmtime(FOTO_USUARIO) > os.path.getmtime(FOTO)):
        from PIL import Image
        Image.open(FOTO_USUARIO).convert("RGB").crop(RECORTE).save(FOTO, quality=85, optimize=True)
    return FOTO


def proporcion():
    from PIL import Image
    w, h = Image.open(foto()).size
    return h / w


def datos(clave):
    """Textos de la caratula (para docx_espec.js y para el analisis estadistico por UNO)."""
    t, s = DOCS[clave]
    return dict(entidad=ENTIDAD, gerencia=GERENCIA, expediente=EXPEDIENTE, proyecto=PROYECTO, ci=CI, foto=foto(), leyenda=LEYENDA,
                proporcion=proporcion(), titulo=t, subtitulo=s, especialidad=ESPECIALIDAD, ubicacion=UBICACION, fecha=FECHA)


# ------------------------------------------------------------------------------------------------ Word (python-docx)
def word(d, clave, ancho_cm=15.5):
    """Agrega la caratula al inicio del documento python-docx `d` (pagina propia, sin encabezado ni pie)."""
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    t, s = DOCS[clave]
    rgb = lambda h: RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    def par(txt, tam, neg=False, color=None, antes=0, despues=4):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.0
        r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.font.name = "Arial"
        if color: r.font.color.rgb = rgb(color)
        return p

    def linea(p, color=AZUL, tam=12):
        pPr = p._p.get_or_add_pPr(); bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", str(tam)), ("space", "4"), ("color", color)): b.set(qn("w:" + k), v)
        bdr.append(b); pPr.append(bdr)

    sec = d.sections[0]; sec.different_first_page_header_footer = True
    par(ENTIDAD, 15, True, AZUL, despues=0)
    linea(par(GERENCIA, 11, True, GRIS, despues=6))
    par(EXPEDIENTE, 12, True, AZUL, antes=6, despues=4)
    par(PROYECTO, 11.5, True, None, despues=2)
    par(CI, 11, True, AZUL, despues=8)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(foto(), width=Cm(ancho_cm))
    par(LEYENDA, 8, False, GRIS, despues=8).runs[0].italic = True
    tb = d.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.rows[0].cells[0]; c.width = Cm(ancho_cm)
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), AZUL)
    c._tc.get_or_add_tcPr().append(sh)
    for i, (txt, tam) in enumerate(((t, 20 if len(t) < 30 else 17), (s, 11))):
        q = c.paragraphs[0] if i == 0 else c.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_before = Pt(6 if i == 0 else 0); q.paragraph_format.space_after = Pt(0 if i == 0 else 6)
        r = q.add_run(txt); r.bold = True; r.font.size = Pt(tam); r.font.name = "Arial"; r.font.color.rgb = rgb("FFFFFF")
    par(ESPECIALIDAD, 10.5, True, None, antes=10, despues=2)
    par(UBICACION, 10, False, None, despues=8)
    par(FECHA, 12, True, AZUL, despues=0)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------------------------------------ Excel (openpyxl)
def excel(wb, clave, nombre="CARÁTULA"):
    """Hoja de caratula al inicio del libro openpyxl `wb` (A4 vertical, una pagina)."""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.drawing.image import Image as XImage
    from openpyxl.utils import get_column_letter as CL
    t, s = DOCS[clave]
    ws = wb.create_sheet(nombre, 0); wb.active = 0
    for j in range(1, 9): ws.column_dimensions[CL(j)].width = 11.5
    def fila(r, txt, tam, neg=False, color="000000", bg=None, alto=None, cursiva=False):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
        c = ws.cell(r, 1, txt); c.font = Font(name="Arial", size=tam, bold=neg, color=color, italic=cursiva)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if bg:
            for j in range(1, 9): ws.cell(r, j).fill = PatternFill("solid", fgColor=bg)
        if alto: ws.row_dimensions[r].height = alto
    fila(1, ENTIDAD, 15, True, AZUL, alto=24)
    fila(2, GERENCIA, 11, True, GRIS, alto=18)
    for j in range(1, 9): ws.cell(2, j).border = Border(bottom=Side(style="medium", color=AZUL))
    fila(3, "", 6, alto=8)
    fila(4, EXPEDIENTE, 12, True, AZUL, alto=20)
    fila(5, PROYECTO, 11, True, alto=46)
    fila(6, CI, 11, True, AZUL, alto=20)
    fila(7, "", 6, alto=6)
    # foto estirada a la celda A8:H8 (ancla de dos celdas): ancho de las 8 columnas (11.5 caracteres = 85.5 px = 64.1 pt cada una)
    # y alto de la fila con la proporcion de la foto
    from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
    ancho_pt = 8 * (11.5 * 7 + 5) * 0.75
    ws.row_dimensions[8].height = round(ancho_pt * proporcion(), 1)
    ws.merge_cells(start_row=8, start_column=1, end_row=8, end_column=8)
    img = XImage(foto())
    img.anchor = TwoCellAnchor(editAs="oneCell", _from=AnchorMarker(col=0, row=7), to=AnchorMarker(col=8, row=8))
    ws.add_image(img)
    fila(9, LEYENDA, 8, False, GRIS, alto=14, cursiva=True)
    fila(10, "", 6, alto=6)
    fila(11, t, 20 if len(t) < 30 else 17, True, "FFFFFF", AZUL, alto=34)
    fila(12, s, 11, True, "FFFFFF", AZUL, alto=22)
    fila(13, "", 6, alto=10)
    fila(14, ESPECIALIDAD, 10.5, True, alto=18)
    fila(15, UBICACION, 10, False, alto=16)
    fila(16, "", 6, alto=8)
    fila(17, FECHA, 12, True, AZUL, alto=20)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1
    ws.print_options.horizontalCentered = True; ws.print_options.verticalCentered = True; ws.print_area = "A1:H17"
    ws.page_margins.top = 0.8; ws.page_margins.bottom = 0.6
    return ws


# ------------------------------------------------------------------------------------------------ Calc (LibreOffice UNO)
def uno(doc, ctx, clave, nombre="CARÁTULA"):
    """Hoja de caratula al inicio de un documento Calc abierto por UNO (analisis estadistico, que conserva formulas y graficos).
    ctx: contexto de componentes del LibreOffice remoto (para el GraphicProvider)."""
    import uno as U
    from com.sun.star.awt import Size, Point
    from com.sun.star.table.CellHoriJustify import CENTER
    from com.sun.star.table.CellVertJustify import CENTER as VCENTER
    from com.sun.star.table import BorderLine2, CellRangeAddress
    from com.sun.star.beans import PropertyValue
    t, s = DOCS[clave]
    sh = doc.Sheets
    if sh.hasByName(nombre): sh.removeByName(nombre)
    sh.insertNewByName(nombre, 0); ws = sh.getByName(nombre)
    ANCHO = 2250                                                     # 8 columnas de 22.5 mm = 18 cm
    for j in range(8): ws.Columns.getByIndex(j).Width = ANCHO
    col = lambda h: int(h, 16)
    altos = {}
    def fila(r, txt, tam, neg=False, color="000000", bg=None, alto=None, cursiva=False):
        rg = ws.getCellRangeByName("A%d:H%d" % (r, r)); rg.merge(True)
        ws.getCellRangeByName("A%d" % r).setString(txt)
        rg.CharFontName = "Arial"; rg.CharHeight = tam; rg.CharWeight = 150 if neg else 100; rg.CharColor = col(color)
        if cursiva: rg.CharPosture = U.Enum("com.sun.star.awt.FontSlant", "ITALIC")
        rg.HoriJustify = CENTER; rg.VertJustify = VCENTER; rg.IsTextWrapped = True
        if bg: rg.CellBackColor = col(bg)
        if alto:
            altos[r] = int(alto * 35.28); ws.Rows.getByIndex(r - 1).Height = altos[r]       # puntos -> 1/100 mm
    fila(1, ENTIDAD, 15, True, AZUL, alto=24)
    fila(2, GERENCIA, 11, True, GRIS, alto=18)
    b = BorderLine2(); b.Color = col(AZUL); b.OuterLineWidth = 50; b.LineWidth = 50
    ws.getCellRangeByName("A2:H2").BottomBorder = b
    fila(3, "", 6, alto=8); fila(4, EXPEDIENTE, 12, True, AZUL, alto=20); fila(5, PROYECTO, 11, True, alto=46)
    fila(6, CI, 11, True, AZUL, alto=20); fila(7, "", 6, alto=6)
    w = 8 * ANCHO; h = int(w * proporcion())
    ws.getCellRangeByName("A8:H8").merge(True); ws.Rows.getByIndex(7).Height = h + 100
    fila(9, LEYENDA, 8, False, GRIS, alto=14, cursiva=True); fila(10, "", 6, alto=6)
    fila(11, t, 20 if len(t) < 30 else 17, True, "FFFFFF", AZUL, alto=34); fila(12, s, 11, True, "FFFFFF", AZUL, alto=22)
    fila(13, "", 6, alto=10); fila(14, ESPECIALIDAD, 10.5, True, alto=18); fila(15, UBICACION, 10, False, alto=16)
    fila(16, "", 6, alto=8); fila(17, FECHA, 12, True, AZUL, alto=20)
    # foto anclada a la celda A8
    prov = ctx.ServiceManager.createInstanceWithContext("com.sun.star.graphic.GraphicProvider", ctx)
    pv = PropertyValue(); pv.Name = "URL"; pv.Value = U.systemPathToFileUrl(foto())
    shp = doc.createInstance("com.sun.star.drawing.GraphicObjectShape")
    shp.Graphic = prov.queryGraphic((pv,))
    ws.DrawPage.add(shp)
    shp.setSize(Size(w, h)); shp.setPosition(Point(0, sum(altos[r] for r in range(1, 8)) + 50))
    shp.Anchor = ws.getCellRangeByName("A8")
    # pagina A4 vertical, una hoja, sin encabezado
    ps = doc.StyleFamilies.getByName("PageStyles")
    if not ps.hasByName("Caratula"):
        ps.insertByName("Caratula", doc.createInstance("com.sun.star.style.PageStyle"))
    st = ps.getByName("Caratula")
    st.IsLandscape = False; st.Width = 21000; st.Height = 29700; st.ScaleToPagesX = 1; st.ScaleToPagesY = 1
    st.CenterHorizontally = True; st.CenterVertically = True; st.HeaderIsOn = False; st.FooterIsOn = False; st.TopMargin = 2000; st.BottomMargin = 1500
    ws.PageStyle = "Caratula"
    a = CellRangeAddress(); a.Sheet = 0; a.StartColumn = 0; a.StartRow = 0; a.EndColumn = 7; a.EndRow = 16
    ws.setPrintAreas((a,))
    return ws
