"""Utilidades para los Word de drenaje pluvial de la I.E. N.° 0590 Barranquita (python-docx): caratula, titulos numerados,
parrafos justificados, vinetas, tablas con encabezado, figuras con leyenda, encabezado y pie con numero de pagina."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AZUL = RGBColor(0x1F, 0x3A, 0x68)
PROYECTO = ("“MEJORAMIENTO Y AMPLIACIÓN DEL SERVICIO DE EDUCACIÓN PRIMARIA EN I.E. N° 0590 DISTRITO DE BARRANQUITA DE LA "
            "PROVINCIA DE LAMAS DEL DEPARTAMENTO DE SAN MARTÍN”")
CI = "C.I. N° 468901"
ENTIDAD = "GOBIERNO REGIONAL SAN MARTÍN - GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO"
UBIC = "Localidad y distrito de Barranquita, provincia de Lamas, departamento de San Martín"
FECHA = "OCTUBRE 2026"
ESPECIALIDAD = "INSTALACIONES SANITARIAS - DRENAJE PLUVIAL"


def _sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)


def _campo(run, instr):
    for t, v in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if t == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = v
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        run._r.append(e)


class Doc:
    def __init__(self, corto):
        self.d = Document(); self.n = [0, 0, 0]; self.nt = 0; self.nf = 0; self.indice = []
        st = self.d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        pf = st.paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.15
        for k, tam in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5)):
            h = self.d.styles[k]; h.font.name = "Arial"; h.font.size = Pt(tam); h.font.bold = True; h.font.color.rgb = AZUL
            h.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial"); h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        s = self.d.sections[0]; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = Cm(3.0); s.right_margin = Cm(2.5); s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.2)
        hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = hp.add_run("I.E. N° 0590 BARRANQUITA - %s - %s" % (CI, corto)); r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Página "); r.font.size = Pt(8); r2 = fp.add_run(); r2.font.size = Pt(8); _campo(r2, "PAGE")

    # ------------------------------------------------------------------ caratula e indice
    def caratula(self, titulo, subtitulo=None, clave=None):
        if clave:                                    # caratula comun con la foto aerea del local (caratula_b)
            import caratula_b
            caratula_b.word(self.d, clave); return
        d = self.d
        for _ in range(4): d.add_paragraph()
        for txt, tam, neg in ((ENTIDAD, 11, True), ("", 10, False), (PROYECTO, 13, True), (CI, 11, True), ("", 10, False),
                              (titulo, 18, True), (subtitulo or "", 12, False), ("", 10, False), ("ESPECIALIDAD: " + ESPECIALIDAD, 10.5, True),
                              ("UBICACIÓN: " + UBIC, 10, False), ("", 10, False), (FECHA, 11, True)):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg
            if tam >= 13: r.font.color.rgb = AZUL
        self.salto()

    def contenido(self, titulos):
        self.h_simple("CONTENIDO")
        for t in titulos:
            p = self.d.add_paragraph(t); p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.6 * t.split(" ")[0].count("."))
        self.salto()

    def salto(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ titulos y texto
    def h_simple(self, t):
        p = self.d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL

    def h(self, nivel, t):
        self.n[nivel - 1] += 1
        for k in range(nivel, 3): self.n[k] = 0
        num = ".".join(str(x) for x in self.n[:nivel])
        txt = "%s%s %s" % (num, "." if nivel == 1 else "", t.upper() if nivel == 1 else t)
        self.d.add_heading(txt, level=nivel); self.indice.append(txt)

    def p(self, t, neg=None, centro=False, tam=None, cursiva=False):
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
        if neg:
            r = par.add_run(neg); r.bold = True
            if tam: r.font.size = Pt(tam)
        r = par.add_run(t); r.italic = cursiva
        if tam: r.font.size = Pt(tam)
        return par

    def vinetas(self, items, estilo="List Bullet"):
        for it in items:
            par = self.d.add_paragraph(style=estilo); par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if isinstance(it, tuple):
                r = par.add_run(it[0]); r.bold = True; par.add_run(it[1])
            else:
                par.add_run(it)

    def formula(self, t):
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(t); r.italic = True; r.font.size = Pt(11)

    # ------------------------------------------------------------------ tablas y figuras
    def tabla(self, titulo, cab, filas, anchos=None, tam=8.5, fuente=None, total=False):
        self.nt += 1
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run("Tabla N° %02d. %s" % (self.nt, titulo)); r.bold = True; r.font.size = Pt(9)
        t = self.d.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(cab):
            cc = t.rows[0].cells[j]; cc.text = ""; rr = cc.paragraphs[0].add_run(str(c)); rr.bold = True; rr.font.size = Pt(tam)
            cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; _sombra(cc, "D9E2F3")
        for i, f in enumerate(filas):
            row = t.add_row().cells
            for j, v in enumerate(f):
                row[j].text = ""; rr = row[j].paragraphs[0].add_run("" if v is None else str(v)); rr.font.size = Pt(tam)
                row[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 or (isinstance(v, str) and len(v) > 18)) else WD_ALIGN_PARAGRAPH.CENTER
                if total and i == len(filas) - 1:
                    rr.bold = True; _sombra(row[j], "F2F2F2")
        if anchos:
            t.autofit = False
            tblW = t._tbl.tblPr.find(qn("w:tblW"))
            if tblW is None:
                tblW = OxmlElement("w:tblW"); t._tbl.tblPr.append(tblW)
            tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(int(sum(anchos) * 567)))
            for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
                if j < len(anchos): gc.set(qn("w:w"), str(int(anchos[j] * 567)))
            for row in t.rows:
                for j, w in enumerate(anchos): row.cells[j].width = Cm(w)
        if fuente:
            par = self.d.add_paragraph(); r = par.add_run("Fuente: " + fuente); r.italic = True; r.font.size = Pt(8)
        else:
            self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figura(self, ruta, titulo, ancho=15.5, fuente=None):
        self.nf += 1
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.add_run().add_picture(ruta, width=Cm(ancho))
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run("Figura N° %02d. %s" % (self.nf, titulo)); r.bold = True; r.font.size = Pt(9)
        if fuente:
            r = par.add_run("\nFuente: " + fuente); r.italic = True; r.font.size = Pt(8)

    def guardar(self, ruta):
        self.d.save(ruta); print(ruta)
