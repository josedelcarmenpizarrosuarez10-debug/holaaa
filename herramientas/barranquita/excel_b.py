"""Cuadro de tramos del planteamiento preliminar de cunetas - I.E. N.° 0590, Barranquita (Excel con formulas).
Lee entregables_barranquita/_calc/diseno.json. Uso: python3 herramientas/barranquita/excel_b.py
"""
import os, sys, json
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL

SAL = os.path.join(RAIZ, "entregables_barranquita", "CUADRO_TRAMOS_CUNETAS_IE0590_BARRANQUITA.xlsx")
AZUL, AMA, VER, GRIS = "1F4E79", "FFF2CC", "E2EFDA", "F2F2F2"
_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
F = lambda **k: Font(name="Arial", size=k.pop("size", 9), **k)
PROY = ("MEJORAMIENTO Y AMPLIACION DEL SERVICIO DE EDUCACION PRIMARIA EN LA I.E. N.° 0590, DISTRITO DE BARRANQUITA, PROVINCIA DE LAMAS, "
        "DEPARTAMENTO DE SAN MARTIN - C.I. N.° 468901")
sys.path.insert(0, AQUI)
import diseno_b as DS
SIS = {k: "SISTEMA %d - %s (descarga %s)" % (k, v[0], v[1]) for k, v in DS.SISTEMAS.items()}
DESC_SIS = {k: v[1] for k, v in DS.SISTEMAS.items()}


def laminas_txt():
    """Relacion de laminas segun el indice de planos_b.py (_calc/indice.json), agrupada por titulo."""
    f = os.path.join(RAIZ, "entregables_barranquita", "_calc", "indice.json")
    if not os.path.exists(f): return ""
    grupos = []
    for cod, tit, _, _ in json.load(open(f)):
        if grupos and grupos[-1][2] == tit: grupos[-1][1] = cod
        else: grupos.append([cod, cod, tit])
    return "; ".join(("%s a %s" % (a, b) if a != b else a) + " " + t.lower() for a, b, t in grupos)


def cab(ws, fila, textos, anchos=None):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t); c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOR
    ws.row_dimensions[fila].height = 30
    if anchos:
        for j, w in enumerate(anchos, 1): ws.column_dimensions[CL(j)].width = w


def titulo(ws, t, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol); ws["A1"] = t
    ws["A1"].font = F(bold=True, size=12, color=AZUL); ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol); ws["A2"] = PROY
    ws["A2"].font = F(size=8); ws["A2"].alignment = Alignment(horizontal="center", wrap_text=True); ws.row_dimensions[2].height = 24


DS_EM = 0.10      # espesor de muro (diseno_b.E_MURO)


def main():
    R = json.load(open(os.path.join(RAIZ, "entregables_barranquita", "_calc", "diseno.json")))
    c = R["criterios"]; nom = R["nombre"]
    wb = Workbook()
    # ---- criterios
    wc = wb.active; wc.title = "CRITERIOS"
    titulo(wc, "CRITERIOS DEL PLANTEAMIENTO PRELIMINAR DE CUNETAS", 4)
    cab(wc, 4, ["PARAMETRO", "VALOR", "UND", "OBSERVACION"], [52, 12, 8, 70])
    par = [("Ancho libre de la cuneta (b)", c["b_int"], "m", "Seccion rectangular de concreto"),
           ("Espesor de muros", c["muro"], "m", "Ancho total = b + 2 muros"),
           ("Ancho total", "=B5+2*B6", "m", "Criterio del usuario: 0.60 m"),
           ("Profundidad al inicio de cada ramal", c["h0"], "m", "Criterio del usuario"),
           ("Profundidad minima", c["hmin"], "m", "Donde el terreno cae mas que la cuneta"),
           ("Pendiente de la cota de fondo", c["S"], "m/m", "0.5 %; mayor donde el terreno cae mas"),
           ("Borde superior de la cuneta (rasante)", c["rasante"], "m", "Indicacion del usuario: NPT +%.2f (pisos y pasadizos), area verde +%.2f" % (c["npt"], c["nivel_area_verde"])),
           ("Coeficiente de Manning (concreto)", c["n_manning"], "", "Para la capacidad de cada tramo"),
           ("Borde libre", c["borde_libre"], "m", "Tirante maximo = h - borde libre")]
    for i, (t, v, u, o) in enumerate(par, 5):
        wc.cell(i, 1, t).font = F(); cc = wc.cell(i, 2, v); cc.font = F(bold=True); cc.fill = PatternFill("solid", fgColor=VER if isinstance(v, str) else AMA)
        wc.cell(i, 3, u).font = F(); wc.cell(i, 4, o).font = F()
        for j in range(1, 5): wc.cell(i, j).border = BOR
    wc.cell(15, 1, "Celda amarilla = dato; celda verde = formula. Los cambios de criterio se reflejan en las capacidades (hoja TRAMOS); "
                   "las cotas de fondo vienen del trazo (rehacer el calculo si cambian h inicial o S).").font = F(italic=True, size=8)
    # ---- tramos
    wt = wb.create_sheet("TRAMOS")
    titulo(wt, "CUADRO DE TRAMOS DE CUNETAS - COTAS DE FONDO, PROFUNDIDADES, PENDIENTES Y CAPACIDAD", 17)
    cab(wt, 4, ["SISTEMA", "EJE", "DE", "A", "L (m)", "RASANTE DE", "RASANTE A", "COTA DE FONDO DE", "COTA DE FONDO A", "h DE (m)", "h A (m)",
                "S (%)", "TAPADO (m)", "CON REJILLA (m)", "TIRANTE MAX. (m)", "Q CAPACIDAD (l/s)", "DESCRIPCION DEL EJE"],
        [9, 8, 7, 7, 9, 11, 11, 13, 13, 9, 9, 8, 10, 11, 11, 12, 60])
    f = 5
    for r in R["ramales"]:
        for t in r["tramos"]:
            vals = [r["sistema"], r["codigo"], nom[t["de"]], nom[t["a"]], t["L"], t["ras_de"], t["ras_a"], t["cf_de"], t["cf_a"],
                    "=F%d-H%d" % (f, f), "=G%d-I%d" % (f, f), "=100*(H%d-I%d)/E%d" % (f, f, f), t["L_tapado"], "=E%d-M%d" % (f, f),
                    "=MAX(MIN(J%d,K%d)-CRITERIOS!$B$13,0.01)" % (f, f),
                    "=1000*CRITERIOS!$B$5*O{0}*((CRITERIOS!$B$5*O{0})/(CRITERIOS!$B$5+2*O{0}))^(2/3)*SQRT(MAX(L{0}/100,CRITERIOS!$B$10))/CRITERIOS!$B$12".format(f),
                    r["descripcion"]]
            for j, v in enumerate(vals, 1):
                cc = wt.cell(f, j, v); cc.font = F(); cc.border = BOR
                cc.alignment = Alignment(horizontal="left" if j == 17 else "center", vertical="center")
                if j in (6, 7, 10, 11, 13, 14, 15): cc.number_format = "0.00"
                if j in (8, 9): cc.number_format = "0.000"
                if j == 5: cc.number_format = "0.00"
                if j in (12,): cc.number_format = "0.00"
                if j == 16: cc.number_format = "0.0"
                if isinstance(v, str) and v.startswith("="): cc.fill = PatternFill("solid", fgColor=VER)
            f += 1
    ft = f - 1
    wt.cell(f, 4, "TOTAL").font = F(bold=True)
    for j, col in ((5, "E"), (13, "M"), (14, "N")):
        cc = wt.cell(f, j, "=SUM(%s5:%s%d)" % (col, col, ft)); cc.font = F(bold=True); cc.number_format = "0.00"; cc.fill = PatternFill("solid", fgColor=VER); cc.border = BOR
    wt.freeze_panes = "E5"
    # ---- nudos
    wn = wb.create_sheet("NUDOS")
    titulo(wn, "NUDOS DE LA RED (COORDENADAS UTM WGS84 DEL PLANO DE ARQUITECTURA)", 5)
    cab(wn, 4, ["NUDO", "ESTE (m)", "NORTE (m)", "EJES QUE PASAN", "DESCARGA / OBSERVACION"], [9, 14, 15, 28, 70])
    pasa = {}
    for r in R["ramales"]:
        for n in r["nudos"]: pasa.setdefault(n, []).append(r["codigo"])
    orden = sorted(nom, key=lambda n: (nom[n] in R["descargas"], nom[n]))
    for i, n in enumerate(orden, 5):
        x, y = R["nudos"][n]
        for j, v in enumerate([nom[n], x, y, ", ".join(pasa[n]), R["descargas"].get(n, "")], 1):
            cc = wn.cell(i, j, v); cc.font = F(bold=nom[n] in R["descargas"]); cc.border = BOR
            if j in (2, 3): cc.number_format = "0.000"
    # ---- resumen
    wr = wb.create_sheet("RESUMEN", 0)
    titulo(wr, "RESUMEN DEL PLANTEAMIENTO PRELIMINAR DE CUNETAS POR SISTEMA", 6)
    cab(wr, 4, ["SISTEMA", "LONGITUD (m)", "CON REJILLA (m)", "TAPADO (m)", "h MAXIMA (m)", "DESCARGA"], [48, 14, 15, 13, 13, 70])
    sis = sorted({r["sistema"] for r in R["ramales"]})
    for i, k in enumerate(sis, 5):
        d = DESC_SIS[k]
        vals = [SIS[k], "=SUMIF(TRAMOS!$A$5:$A$%d,%d,TRAMOS!$E$5:$E$%d)" % (ft, k, ft), "=SUMIF(TRAMOS!$A$5:$A$%d,%d,TRAMOS!$N$5:$N$%d)" % (ft, k, ft),
                "=SUMIF(TRAMOS!$A$5:$A$%d,%d,TRAMOS!$M$5:$M$%d)" % (ft, k, ft),
                max(max(t["h_max"] for t in r["tramos"]) for r in R["ramales"] if r["sistema"] == k), d + ": " + R["descargas"][d]]
        for j, v in enumerate(vals, 1):
            cc = wr.cell(i, j, v); cc.font = F(); cc.border = BOR
            if j in (2, 3, 4, 5): cc.number_format = "0.00"; cc.alignment = Alignment(horizontal="center")
            if isinstance(v, str) and v.startswith("="): cc.fill = PatternFill("solid", fgColor=VER)
    ft_r = 5 + len(sis)
    wr.cell(ft_r, 1, "TOTAL").font = F(bold=True)
    for j, col in ((2, "B"), (3, "C"), (4, "D")):
        cc = wr.cell(ft_r, j, "=SUM(%s5:%s%d)" % (col, col, ft_r - 1)); cc.font = F(bold=True); cc.number_format = "0.00"; cc.fill = PatternFill("solid", fgColor=VER); cc.border = BOR
    notas = ["Planteamiento preliminar: ejes de cuneta en las lineas magenta del plano PLANTA_GENERAL_590 ULTIMA VERSION (lados de cada bloque, techos a dos aguas) y uniones hasta las descargas.",
             "Sentidos de flujo y descargas segun las curvas de nivel C-TOPO del mismo plano. Techos a dos aguas: cada alero descarga a su lado.",
             "Area de intervencion = interior del cerco perimetrico: todas las cunetas quedan dentro del cerco (las perimetrales a 0.60 m de su eje).",
             "Interior del cerco nivelado: pisos y pasadizos NPT +184.65; borde de cuneta +184.65 (al ras del piso); areas verdes +184.60; patio de maniobra +184.50 (cuneta al ras, con rejilla). Fuera del cerco, el terreno natural.",
             "D1, D2, D3 y D4 salen por un pase en el cimiento del cerco. Fuera del cerco no hay obra.",
             "Tramos de fuerte pendiente con V > 3.0 m/s: gradas de concreto (hoja HIDRAULICA). Verificacion completa en la hoja VERIFICACION.",
             "Ver planos PLANTEAMIENTO_CUNETAS_IE0590_BARRANQUITA.dxf: " + laminas_txt() + "."]
    wr.cell(11, 1, "NOTAS").font = F(bold=True, color=AZUL)
    for i, t in enumerate(notas, 12):
        wr.cell(i, 1, "%d. %s" % (i - 11, t)).font = F(size=8)
    # ---- verificacion (10 comprobaciones) e hidraulica con el caudal de diseno
    Vf = json.load(open(os.path.join(RAIZ, "entregables_barranquita", "_calc", "verificacion.json")))
    wv = wb.create_sheet("VERIFICACION", 1)
    titulo(wv, "VERIFICACION DEL PLANTEAMIENTO (%d COMPROBACIONES)" % len(Vf["comprobaciones"]), 4)
    cab(wv, 4, ["N.", "COMPROBACION", "RESULTADO", "DETALLE"], [5, 62, 22, 120])
    for i, c_ in enumerate(Vf["comprobaciones"], 5):
        for j, v in enumerate([c_["n"], c_["comprobacion"], c_["estado"], c_["detalle"]], 1):
            cc = wv.cell(i, j, v); cc.font = F(bold=(j == 3)); cc.border = BOR; cc.alignment = Alignment(wrap_text=True, vertical="top")
            if j == 3: cc.fill = PatternFill("solid", fgColor=VER if c_["estado"] == "CUMPLE" else AMA)
        wv.row_dimensions[i].height = 15 * max(1, len(str(c_["detalle"])) // 95 + 1)
    r0 = 5 + len(Vf["comprobaciones"]) + 1
    wv.cell(r0, 2, "DESCARGAS").font = F(bold=True, color=AZUL)
    cab(wv, r0 + 1, ["", "DESCARGA", "COTA DE FONDO", "SALIDA POR GRAVEDAD"])
    for i, d in enumerate(Vf["descargas"], r0 + 2):
        txt = (("poza de salida: rebalsa a %.2f; el fondo exterior no sube y baja 0.25 m a %.1f m" % (d["nivel_rebalse"], d["sale_a_m"])) if d["tipo"] == "poza"
               else ("pase en el cimiento del cerco: el terreno exterior queda bajo la C.F. a %.1f m (perfilado maximo %.2f m, sin obra)" % (d["sale_a_m"], d["corte_max"])))
        for j, v in enumerate(["", d["D"], d["cf"], txt], 1):
            cc = wv.cell(i, j, v); cc.font = F(); cc.border = BOR
    h = Vf["hidrologia"]
    wh = wb.create_sheet("HIDRAULICA", 3)
    titulo(wh, "CAUDAL DE DISENO, TIRANTE Y VELOCIDAD POR TRAMO (REFERENCIAL)", 13)
    par = [("C (escorrentia)", h["C"]), ("I (mm/h)", h["I"]), ("Factor de seguridad", h["FS"]), ("Velocidad maxima (m/s)", 3.0)]
    for i, (t, v) in enumerate(par, 4):
        wh.cell(i, 1, t).font = F(); cc = wh.cell(i, 2, v); cc.font = F(bold=True); cc.fill = PatternFill("solid", fgColor=AMA); cc.border = BOR
    wh.cell(8, 1, "I: tc = %.0f min, TR = 25 anos, P10 = %.2f mm (estacion usada en el Hogar de Refugio - Morales); reemplazar por la del proyecto. "
                  "Area tributaria: cada m2 dentro del cerco a la cuneta mas cercana." % (h["tc"], h["P10"])).font = F(italic=True, size=8)
    cab(wh, 10, ["EJE", "DE", "A", "AREA (m2)", "Q (l/s)", "S (%)", "TIRANTE y (m)", "h (m)", "BORDE LIBRE (m)", "V (m/s)", "CAPACIDAD", "VELOCIDAD", "OBSERVACION"],
        [9, 7, 7, 11, 10, 8, 12, 8, 13, 9, 12, 12, 46])
    for i, fh in enumerate(Vf["hidraulica"], 11):
        vals = [fh["ramal"], fh["de"], fh["a"], fh["A_m2"], "=1000*$B$4*$B$5*D%d/3600000*$B$6" % i, fh["S"], fh["y"], fh["h"], "=H%d-G%d" % (i, i),
                fh["V"], '=IF(I%d>=0.05,"CUMPLE","NO CUMPLE")' % i, '=IF(J%d<=$B$7,"CUMPLE","GRADAS")' % i,
                "gradas de concreto (escalones) en este tramo" if fh["vel"] != "CUMPLE" else ""]
        for j, v in enumerate(vals, 1):
            cc = wh.cell(i, j, v); cc.font = F(); cc.border = BOR; cc.alignment = Alignment(horizontal="center" if j < 13 else "left")
            if j in (7, 9): cc.number_format = "0.000"
            if j in (5, 10): cc.number_format = "0.00"
            if isinstance(v, str) and v.startswith("="): cc.fill = PatternFill("solid", fgColor=VER)
    wh.freeze_panes = "D11"
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    # ---- secciones de los perfiles (las mismas de las laminas de perfiles y de secciones transversales)
    fs = os.path.join(RAIZ, "entregables_barranquita", "_calc", "secciones.json")
    if os.path.exists(fs):
        ws = wb.create_sheet("SECCIONES")
        titulo(ws, "SECCIONES DE LOS PERFILES LONGITUDINALES POR EJE", 10)
        cab(ws, 4, ["SECCION", "EJE", "PROGRESIVA (m)", "NCT", "NCF", "H (m)", "EL (m)", "EM (m)", "TIPO", "OBSERVACION"],
            [10, 10, 15, 11, 11, 9, 9, 9, 14, 40])
        i = 5
        for eje, lst in json.load(open(fs)).items():
            for x in lst:
                for j, v in enumerate(["S-%02d" % x["n"], eje, x["s"], x["nct"], x["ncf"], "=D%d-E%d" % (i, i), 0.10, DS_EM,
                                       "TAPADA" if x["tap"] else "CON REJILLA", x["nota"]], 1):
                    cc = ws.cell(i, j, v); cc.font = F(); cc.border = BOR
                    if j in (3, 4, 5, 6, 7, 8): cc.number_format = "0.00"
                    if j == 6: cc.fill = PatternFill("solid", fgColor=VER)
                i += 1
    # ---- tapas de inspeccion de los tramos tapados (las mismas de la planta DC-01)
    ft = os.path.join(RAIZ, "entregables_barranquita", "_calc", "tapas.json")
    if os.path.exists(ft):
        ws = wb.create_sheet("TAPAS INSPECCION")
        titulo(ws, "TAPAS DE INSPECCION EN TRAMOS TAPADOS (0.60 x 0.60, CONCRETO REMOVIBLE)", 6)
        cab(ws, 4, ["N.", "EJE", "PROGRESIVA (m)", "ESTE (m)", "NORTE (m)", "UBICACION"], [8, 10, 15, 15, 15, 46])
        i = 5
        for k, t in enumerate(json.load(open(ft)), 1):
            for j, v in enumerate(["T-%02d" % k, t["eje"], t["progresiva"], t["x"], t["y"],
                                   "Inicio / fin de tramo tapado, empalme o a 10 m como maximo"], 1):
                cc = ws.cell(i, j, v); cc.font = F(); cc.border = BOR
                if j == 3: cc.number_format = "0.00"
                if j in (4, 5): cc.number_format = "0.000"
            i += 1
        c = ws.cell(i, 2, "TOTAL"); c.font = F(bold=True)
        c = ws.cell(i, 3, "=COUNTA(B5:B%d)" % (i - 1)); c.font = F(bold=True); c.border = BOR
    wb.save(SAL); print(SAL)


if __name__ == "__main__":
    main()
