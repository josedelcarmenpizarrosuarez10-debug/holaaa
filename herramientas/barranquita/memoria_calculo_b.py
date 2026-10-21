"""Memoria de calculo hidrologico e hidraulico (Excel con formulas) de la red de cunetas de la I.E. N.° 0590 Barranquita.

Hojas: RESUMEN, DATOS, SERIE (Pmax 24 h de la estacion Pongo de Caynarachi, 1989-2013), FRECUENCIAS (Normal, Log-Normal, Log-Pearson III,
Gumbel), KS (Kolmogorov-Smirnov), IDF (Dick y Peschke y regresion I = K T^m / t^n), SUPERFICIES (coeficiente de escorrentia),
CAUDALES (metodo racional y Manning por tramo) y DESCARGAS. Los valores coinciden con hidrologia_b.py y verifica_b.py.
"""
import os, sys, json, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import diseno_b as DS, hidrologia_b as H
import metrado_b as MB

CALC = os.path.join(RAIZ, "entregables_barranquita", "_calc")
SAL = os.path.join(RAIZ, "entregables_barranquita", "MEMORIA_CALCULO_DRENAJE_PLUVIAL_IE0590_BARRANQUITA.xlsx")
_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
F = lambda **k: Font(name="Arial", size=k.pop("size", 9), **k)
AZ, GR, AM, VE = "DDEBF7", "F2F2F2", "FFF2CC", "E2EFDA"
MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SET", "OCT", "NOV", "DIC"]


def cel(ws, r, c, v, fmt=None, bold=False, fill=None, borde=True, wrap=False, al=None):
    x = ws.cell(r, c, v); x.font = F(bold=bold)
    if fmt: x.number_format = fmt
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    if borde: x.border = BOR
    if wrap or al: x.alignment = Alignment(wrap_text=wrap, horizontal=al, vertical="center")
    return x


def cabecera(ws, titulo, ncol):
    """Titulo y datos del proyecto en las filas 1 a 3 (las tablas empiezan en la fila 4 o despues)."""
    ws.cell(1, 1, titulo).font = F(size=11, bold=True)
    ws.cell(2, 1, "PROYECTO: " + MB.PROYECTO).font = F(size=8)
    ws.cell(3, 1, "ENTIDAD: %s  -  UBICACION: %s  -  FECHA: %s" % (MB.ENTIDAD, MB.UBIC, MB.FECHA)).font = F(size=8)


def tabla_cab(ws, r, textos, anchos=None):
    for j, t in enumerate(textos, 1):
        cel(ws, r, j, t, bold=True, fill=AZ, wrap=True, al="center")
        if anchos: ws.column_dimensions[CL(j)].width = anchos[j - 1]
    ws.row_dimensions[r].height = 30


def yn(Q, S, n, b=DS.B_INT):
    lo, hi = 1e-5, 3.0
    for _ in range(100):
        y = (lo + hi) / 2; A = b * y; R = A / (b + 2 * y)
        if A * R ** (2 / 3) * math.sqrt(S) / n > Q: hi = y
        else: lo = y
    return (lo + hi) / 2


def construir():
    HID = H.resultados(); A = HID["analisis"]
    V = json.load(open(os.path.join(CALC, "verificacion.json"))); hv = V["hidrologia"]
    wb = Workbook(); wr = wb.active; wr.title = "RESUMEN"
    wd = wb.create_sheet("DATOS"); ws_ = wb.create_sheet("SERIE"); wf = wb.create_sheet("FRECUENCIAS"); wk = wb.create_sheet("KS")
    wi = wb.create_sheet("IDF"); wsu = wb.create_sheet("SUPERFICIES"); wc = wb.create_sheet("CAUDALES"); wdc = wb.create_sheet("DESCARGAS")
    # ------------------------------------------------------------------ DATOS
    cabecera(wd, "DATOS GENERALES Y PARAMETROS DE DISENO", 4)
    DN = {"normal": "Normal", "lognormal": "Log-Normal", "lp3": "Log-Pearson III", "gumbel": "Gumbel"}; EL = A["elegida"]
    pac = H.ESTACIONES_ZONA[[x[0] for x in H.ESTACIONES_ZONA].index("Pongo de Caynarachi")][4]
    datos = [("Precipitacion media anual - estacion Pongo de Caynarachi (mm)", pac, "0.0",
              "Referencial: SENAMHI 2013-2017, Cuadro N° 05 del estudio hidrologico de Barranquita"),
             ("Precipitacion media anual - llanura de Barranquita (mm)", H.P_ANUAL_BARRANQUITA, "0.0",
              "Referencial: promedio de Navarro (2,349.79), Pelejo (2,332.92) y Shanusi (2,444.00), SENAMHI 2013-2017"),
             ("Factor aplicado a la serie de maximas", H.FACTOR, "0.0000",
              "Serie medida de Pongo de Caynarachi (misma subcuenca del rio Caynarachi): sin transposicion"),
             ("Periodo de retorno de diseno T (anos)", H.T_DISENO, "0", "Mismo criterio de los proyectos de Morales (T = 25 anos)"),
             ("Tiempo de concentracion de cunetas tc (min)", H.TC, "0.0", "Duracion minima (Norma CE.040)"),
             ("Coeficiente de Manning para capacidad (concreto con rejilla)", 0.015, "0.000", "CE.040; mismo valor de los proyectos de Morales"),
             ("Coeficiente de Manning para velocidad (concreto liso)", 0.013, "0.000", "CE.040"),
             ("Luz libre de la cuneta b (m)", DS.B_INT, "0.00", "Planos DC-02: 0.40 m libres + muros 0.10 m = 0.60 m"),
             ("Borde libre minimo (m)", DS.BL, "0.00", "Criterio de diseno"),
             ("Velocidad maxima en concreto (m/s)", 3.0, "0.00", "CE.040"),
             ("Precipitacion maxima en 24 h de diseno, T = 25 anos (mm)", "=FRECUENCIAS!H%d" % 0, "0.00", "Hoja FRECUENCIAS (%s)" % DN[EL]),
             ("Precipitacion en tc - Dick y Peschke P = P24 (tc/1440)^0.25 (mm)", "=B15*(B9/1440)^0.25", "0.00", "Dick y Peschke"),
             ("Intensidad de diseno I = P x 60 / tc (mm/h)", "=B16*60/B9", "0.00", "Para el metodo racional"),
             ("Coeficiente de escorrentia - techos y coberturas", H.C_SUP["techo"], "0.00", "CE.040 Anexo I, Tabla 1.a: concreto / techo, T = 25 anos"),
             ("Coeficiente de escorrentia - areas libres de concreto (veredas y patios)", H.C_SUP["pavimento"], "0.00", "CE.040 Anexo I, Tabla 1.a: concreto, T = 25 anos"),
             ("Coeficiente de escorrentia - patio de maniobra y estacionamiento", H.C_SUP["estacionamiento"], "0.00", "CE.040 Anexo I, Tabla 1.a: concreto, T = 25 anos"),
             ("Coeficiente de escorrentia - areas verdes y terreno", H.C_SUP["verde"], "0.00", "CE.040 Anexo I, Tabla 1.a: zona verde, valor mayor de la tabla (rellenos compactados), T = 25 anos"),
             ("Factor de seguridad de los caudales FS", 1.15, "0.00", "Mismo criterio de los proyectos de Morales (Q = C I A FS)")]
    tabla_cab(wd, 4, ["PARAMETRO", "VALOR", "", "FUENTE / CRITERIO"], [70, 12, 2, 80])
    for i, (t, v, fmt, fte) in enumerate(datos):
        r = 5 + i; cel(wd, r, 1, t); cel(wd, r, 2, v, fmt, bold=True, fill=AM); cel(wd, r, 4, fte)
    # ------------------------------------------------------------------ SERIE
    cabecera(ws_, "PRECIPITACION MAXIMA EN 24 HORAS - ESTACION PONGO DE CAYNARACHI (SENAMHI, CODIGO 003317)", 16)
    e = H.ESTACION
    cel(ws_, 8, 1, "Estacion: %s (codigo %s) - %s - Altitud %d m s. n. m. - %s, %s, %s - periodo %s (%d anos completos mas recientes del registro diario)" % (
        e["nombre"], e["codigo"], e["tipo"], e["altitud"], e["distrito"], e["provincia"], e["departamento"], e["periodo"], e["anios"]), borde=False)
    tabla_cab(ws_, 10, ["AÑO"] + MESES + ["P MAX ANUAL (mm)", "P MAX DE DISENO (mm)", "LOG10 P"], [7] + [6] * 12 + [12, 13, 9])
    r0 = 11
    for i, (a, m) in enumerate(H.SERIE.items()):
        r = r0 + i; cel(ws_, r, 1, a)
        for j, v in enumerate(m, 2): cel(ws_, r, j, v if v is not None else "S/D", "0.0")
        cel(ws_, r, 14, "=MAX(B%d:M%d)" % (r, r), "0.00", bold=True)
        cel(ws_, r, 15, "=N%d*DATOS!$B$7" % r, "0.00", bold=True, fill=VE)
        cel(ws_, r, 16, "=LOG10(O%d)" % r, "0.0000")
    r1 = r0 + len(H.SERIE) - 1
    cel(ws_, r1 + 2, 1, "Maximas mensuales obtenidas de los registros diarios del SENAMHI. S/D: enero de 2005 y julio de 2008 sin datos (se conservan los maximos de "
        "los otros once meses). 2014 (solo enero a abril) no se considera.", borde=False)
    rngO, rngP = "SERIE!$O$%d:$O$%d" % (r0, r1), "SERIE!$P$%d:$P$%d" % (r0, r1)
    # ------------------------------------------------------------------ FRECUENCIAS
    cabecera(wf, "ANALISIS DE FRECUENCIAS DE LA SERIE DE PONGO DE CAYNARACHI 1989-2013 (mm)", 9)
    est = [("Numero de datos n", "=COUNT(%s)" % rngO, "0"), ("Media aritmetica", "=AVERAGE(%s)" % rngO, "0.00"),
           ("Desviacion estandar", "=STDEV(%s)" % rngO, "0.00"), ("Coeficiente de asimetria", "=SKEW(%s)" % rngO, "0.000"),
           ("Media de log P", "=AVERAGE(%s)" % rngP, "0.0000"), ("Desviacion estandar de log P", "=STDEV(%s)" % rngP, "0.0000"),
           ("Coeficiente de asimetria de log P", "=SKEW(%s)" % rngP, "0.0000"),
           ("Gumbel: alfa = raiz(6) s / pi", "=SQRT(6)*B7/PI()", "0.000"), ("Gumbel: u = media - 0.5772 alfa", "=B6-0.5772*B12", "0.000")]
    for i, (t, f_, fmt) in enumerate(est):
        cel(wf, 5 + i, 1, t); cel(wf, 5 + i, 2, f_, fmt, bold=True)
    wf.column_dimensions["A"].width = 38
    tabla_cab(wf, 16, ["T (anos)", "z", "NORMAL", "LOG-NORMAL", "KT LP III", "LOG-PEARSON III", "KT GUMBEL", "GUMBEL", "DISEÑO (%s)" % DN[EL]],
              [38, 9, 11, 12, 10, 13, 11, 11, 14])
    filaT = {}
    for i, T in enumerate(H.PERIODOS):
        r = 17 + i; filaT[T] = r
        cel(wf, r, 1, T, "0", bold=True)
        cel(wf, r, 2, "=NORMSINV(1-1/A%d)" % r, "0.000")
        cel(wf, r, 3, "=$B$6+B%d*$B$7" % r, "0.00")
        cel(wf, r, 4, "=10^($B$9+B%d*$B$10)" % r, "0.00")
        cel(wf, r, 5, "=(2/$B$11)*((1+$B$11/6*B%d-($B$11/6)^2)^3-1)" % r, "0.000")
        cel(wf, r, 6, "=10^($B$9+E%d*$B$10)" % r, "0.00")
        cel(wf, r, 7, "=-SQRT(6)/PI()*(0.5772+LN(LN(A%d/(A%d-1))))" % (r, r), "0.000")
        cel(wf, r, 8, "=$B$6+G%d*$B$7" % r, "0.00")
        cel(wf, r, 9, "=%s%d" % ({"normal": "C", "lognormal": "D", "lp3": "F", "gumbel": "H"}[EL], r), "0.00", bold=True, fill=VE)
    # P24 de diseno en DATOS (T = 25)
    wd.cell(15, 2).value = "=FRECUENCIAS!I%d" % filaT[H.T_DISENO]
    cel(wf, 27, 1, "Se adopta la distribucion %s: menor estadistico de Kolmogorov-Smirnov (hoja KS)." % DN[EL], borde=False)
    # ------------------------------------------------------------------ KS
    cabecera(wk, "PRUEBA DE BONDAD DE AJUSTE DE KOLMOGOROV-SMIRNOV (alfa = 0.05)", 11)
    tabla_cab(wk, 8, ["m", "P ordenada (mm)", "F empirica m/(n+1)", "F Normal", "F Log-Normal", "F Log-Pearson III", "F Gumbel",
                      "|D| Normal", "|D| Log-Normal", "|D| LP III", "|D| Gumbel"], [6, 12, 12, 10, 11, 12, 10, 10, 11, 10, 10])
    n = A["n"]; ord_ = A["orden"]
    # F Log-Pearson III: no tiene funcion directa en la hoja; se escribe el valor calculado (Wilson-Hilferty) y se verifica con la tabla
    import hidrologia_b as HH
    ml, sl, csl = A["media_log"], A["desv_log"], A["cs_log"]
    def F_lp3(v):
        lo, hi = -8.0, 8.0
        for _ in range(100):
            zz = (lo + hi) / 2; k = csl / 6; kt = (2 / csl) * ((1 + k * zz - k * k) ** 3 - 1)
            if 10 ** (ml + kt * sl) < v: lo = zz
            else: hi = zz
        return HH._phi((lo + hi) / 2)
    for i in range(n):
        r = 9 + i
        cel(wk, r, 1, i + 1, "0"); cel(wk, r, 2, "=SMALL(%s,A%d)" % (rngO, r), "0.00"); cel(wk, r, 3, "=A%d/(FRECUENCIAS!$B$5+1)" % r, "0.0000")
        cel(wk, r, 4, "=NORMDIST(B%d,FRECUENCIAS!$B$6,FRECUENCIAS!$B$7,TRUE())" % r, "0.0000")
        cel(wk, r, 5, "=NORMDIST(LOG10(B%d),FRECUENCIAS!$B$9,FRECUENCIAS!$B$10,TRUE())" % r, "0.0000")
        cel(wk, r, 6, round(F_lp3(ord_[i][0]), 6), "0.0000")
        cel(wk, r, 7, "=EXP(-EXP(-(B%d-FRECUENCIAS!$B$13)/FRECUENCIAS!$B$12))" % r, "0.0000")
        for j, col in enumerate("DEFG"):
            cel(wk, r, 8 + j, "=ABS(C%d-%s%d)" % (r, col, r), "0.0000")
    rm = 9 + n
    cel(wk, rm, 7, "D max", bold=True)
    for j in range(4): cel(wk, rm, 8 + j, "=MAX(%s9:%s%d)" % (CL(8 + j), CL(8 + j), rm - 1), "0.0000", bold=True, fill=VE)
    cel(wk, rm + 1, 7, "D critico", bold=True)
    for j in range(4): cel(wk, rm + 1, 8 + j, A["ks_critico"], "0.00")
    cel(wk, rm + 2, 7, "Ajuste", bold=True)
    for j in range(4): cel(wk, rm + 2, 8 + j, '=IF(%s%d<%s%d,"ACEPTA","RECHAZA")' % (CL(8 + j), rm, CL(8 + j), rm + 1))
    cel(wk, rm + 4, 1, "F Log-Pearson III por Wilson-Hilferty (valor calculado). D critico para n = %d y alfa = 0.05: %.2f (tabla de Smirnov)." % (A["n"], A["ks_critico"]), borde=False)
    # ------------------------------------------------------------------ IDF
    cabecera(wi, "LLUVIAS E INTENSIDADES MAXIMAS (DICK Y PESCHKE) Y CURVAS IDF", 11)
    cel(wi, 8, 1, "P (mm) para la duracion D: P = P24 (D / 1440)^0.25", borde=False, bold=True)
    tabla_cab(wi, 9, ["D (min)"] + ["T = %d" % T for T in H.PERIODOS], [10] + [10] * len(H.PERIODOS))
    for i, D in enumerate(H.DURACIONES):
        r = 10 + i; cel(wi, r, 1, D, "0", bold=True)
        for j, T in enumerate(H.PERIODOS, 2):
            cel(wi, r, j, "=FRECUENCIAS!$I$%d*($A%d/1440)^0.25" % (filaT[T], r), "0.00")
    r2 = 10 + len(H.DURACIONES) + 2
    cel(wi, r2, 1, "I (mm/h) = P x 60 / D", borde=False, bold=True)
    tabla_cab(wi, r2 + 1, ["D (min)"] + ["T = %d" % T for T in H.PERIODOS])
    for i, D in enumerate(H.DURACIONES):
        r = r2 + 2 + i; cel(wi, r, 1, D, "0", bold=True)
        for j in range(len(H.PERIODOS)):
            cel(wi, r, 2 + j, "=%s%d*60/$A%d" % (CL(2 + j), 10 + i, r), "0.00", fill=VE if (D == H.TC and H.PERIODOS[j] == H.T_DISENO) else None)
    r3 = r2 + 2 + len(H.DURACIONES) + 2
    cel(wi, r3, 1, "Regresion multiple I = K T^m / t^n (duraciones de 5 a 60 min):", borde=False, bold=True)
    for k, (t, v) in enumerate((("K", HID["K"]), ("m", HID["m"]), ("n", HID["n"]))):
        cel(wi, r3 + 1 + k, 1, t, bold=True); cel(wi, r3 + 1 + k, 2, round(v, 4), "0.0000", fill=AM)
    cel(wi, r3 + 5, 1, "Comprobacion: I (T = 25, t = 10 min) con la ecuacion IDF (mm/h)", borde=False)
    cel(wi, r3 + 5, 6, "=B%d*%d^B%d/%d^B%d" % (r3 + 1, H.T_DISENO, r3 + 2, int(H.TC), r3 + 3), "0.00", bold=True)
    # ------------------------------------------------------------------ SUPERFICIES
    cabecera(wsu, "COEFICIENTE DE ESCORRENTIA PONDERADO (AREA DENTRO DEL CERCO)", 5)
    tabla_cab(wsu, 8, ["TIPO DE SUPERFICIE", "AREA (m2)", "C", "C x A", "% DEL AREA"], [58, 12, 8, 12, 12])
    sup = hv.get("superficies", {})
    filas_s = [("techo", "DATOS!$B$18"), ("pavimento", "DATOS!$B$19"), ("estacionamiento", "DATOS!$B$20"), ("verde", "DATOS!$B$21")]
    rt = 9 + len(filas_s)
    for i, (k, cref) in enumerate(filas_s):
        r = 9 + i
        cel(wsu, r, 1, H.NOM_SUP[k]); cel(wsu, r, 2, sup.get(k, {}).get("A", 0), "#,##0"); cel(wsu, r, 3, "=" + cref, "0.00")
        cel(wsu, r, 4, "=B%d*C%d" % (r, r), "#,##0.0"); cel(wsu, r, 5, "=B%d/$B$%d" % (r, rt), "0.0%")
    cel(wsu, rt, 1, "TOTAL / C PONDERADO", bold=True); cel(wsu, rt, 2, "=SUM(B9:B%d)" % (rt - 1), "#,##0", bold=True)
    cel(wsu, rt, 3, "=D%d/B%d" % (rt, rt), "0.000", bold=True, fill=VE); cel(wsu, rt, 4, "=SUM(D9:D%d)" % (rt - 1), "#,##0.0", bold=True)
    cel(wsu, rt, 5, "=SUM(E9:E%d)" % (rt - 1), "0.0%", bold=True)
    # detalle por zona (cuadro de areas de la memoria de arquitectura)
    det = hv.get("detalle_superficies", []); r1 = rt + 3
    cel(wsu, r1 - 1, 1, "DETALLE POR ZONA (NOMBRES DEL CUADRO DE AREAS DE ARQUITECTURA)", borde=False, bold=True)
    tabla_cab(wsu, r1, ["ZONA", "AREA (m2)", "C", "C x A", "TIPO"], [58, 12, 8, 12, 12])
    TIPO = {"techo": "TECHO", "pavimento": "AREA LIBRE", "estacionamiento": "ESTACIONAM.", "verde": "AREA VERDE"}
    for i, d in enumerate(det):
        r = r1 + 1 + i; cref = dict(filas_s)[d["tipo"]]
        cel(wsu, r, 1, d["zona"]); cel(wsu, r, 2, d["A"], "#,##0.0"); cel(wsu, r, 3, "=" + cref, "0.00")
        cel(wsu, r, 4, "=B%d*C%d" % (r, r), "#,##0.0"); cel(wsu, r, 5, TIPO[d["tipo"]])
    if det:
        r = r1 + 1 + len(det)
        cel(wsu, r, 1, "TOTAL (AREA DENTRO DEL CERCO)", bold=True); cel(wsu, r, 2, "=SUM(B%d:B%d)" % (r1 + 1, r - 1), "#,##0.0", bold=True)
        cel(wsu, r, 3, "=D%d/B%d" % (r, r), "0.000", bold=True, fill=VE); cel(wsu, r, 4, "=SUM(D%d:D%d)" % (r1 + 1, r - 1), "#,##0.0", bold=True)
        rn = r + 2
    else:
        rn = rt + 2
    cel(wsu, rn, 1, "Areas medidas a escala real en la planta del proyecto dentro del cerco. Tabla superior: celdas de 1 m2 asignadas a cada tramo "
                    "(base del calculo de caudales). Detalle: poligonos de cada zona (techos = proyeccion de coberturas, incluye patio civico "
                    "y losa multiusos techados; areas libres = veredas, patios, atrio y circulaciones; patio de maniobra y estacionamiento; "
                    "el resto, areas verdes y terreno nivelado).", borde=False)
    # ------------------------------------------------------------------ CAUDALES
    cabecera(wc, "CAUDALES DE DISENO (METODO RACIONAL) Y VERIFICACION HIDRAULICA (MANNING) POR TRAMO", 20)
    hdr = ["EJE", "DE", "A", "AREA ACUM. (m2)", "C POND.", "I (mm/h)", "Q (m3/s)", "Q (l/s)", "S (%)", "h LIBRE (m)", "y (m) n=0.015",
           "A=b y (m2)", "P=b+2y (m)", "R (m)", "Q MANNING (m3/s)", "BORDE LIBRE (m)", "CAPACIDAD", "y (m) n=0.013", "V (m/s)", "VELOCIDAD"]
    tabla_cab(wc, 8, hdr, [8, 7, 7, 10, 7, 9, 9, 8, 7, 8, 9, 9, 8, 8, 10, 9, 10, 9, 8, 10])
    for i, f in enumerate(V["hidraulica"]):
        r = 9 + i; Qm = f["Q_ls"] / 1000.0
        y = yn(Qm, f["S"] / 100.0, 0.015); yv = yn(Qm, f["S"] / 100.0, 0.013)
        vals = [(f["ramal"], None), (f["de"], None), (f["a"], None), (f["A_m2"], "#,##0"), (f.get("C", hv["C"]), "0.000"),
                ("=DATOS!$B$17", "0.00"), ("=E{r}*F{r}*D{r}*DATOS!$B$22/3600000".format(r=r), "0.0000"), ("=G{r}*1000".format(r=r), "0.0"),
                (f["S"], "0.00"), (f["h"], "0.00"), (round(y, 4), "0.000"), ("=DATOS!$B$12*K{r}".format(r=r), "0.0000"),
                ("=DATOS!$B$12+2*K{r}".format(r=r), "0.000"), ("=L{r}/M{r}".format(r=r), "0.0000"),
                ("=L{r}*N{r}^(2/3)*SQRT(I{r}/100)/DATOS!$B$10".format(r=r), "0.0000"), ("=J{r}-K{r}".format(r=r), "0.000"),
                ('=IF(P{r}>=DATOS!$B$13-0.001,"CUMPLE","NO CUMPLE")'.format(r=r), None), (round(yv, 4), "0.000"),
                ("=G{r}/(DATOS!$B$12*R{r})".format(r=r), "0.00"), ('=IF(S{r}<=DATOS!$B$14,"CUMPLE","NO CUMPLE")'.format(r=r), None)]
        for j, (v, fmt) in enumerate(vals, 1): cel(wc, r, j, v, fmt)
    rr = 9 + len(V["hidraulica"])
    cel(wc, rr + 1, 1, "y: tirante normal por Manning resuelto por iteracion (Q Manning = Q de diseno). En tramos tapados, h libre = altura bajo la losa tapa. "
                       "Area acumulada: celdas de 1 m2 asignadas a la cuneta mas cercana mas las areas de los ejes que entregan aguas arriba.", borde=False)
    # ------------------------------------------------------------------ DESCARGAS
    cabecera(wdc, "DESCARGAS POR EL CERCO PERIMETRICO", 6)
    tabla_cab(wdc, 8, ["DESCARGA", "COTA DE FONDO (m)", "TERRENO AL PIE (m)", "SALE A (m del cerco)", "ESTADO", "RECEPTOR"], [11, 13, 13, 13, 22, 90])
    for i, d in enumerate(V["descargas"]):
        r = 9 + i
        for j, (v, fmt) in enumerate([(d["D"], None), (d["cf"], "0.000"), (d["terreno"], "0.00"), (d.get("sale_a_m"), "0.00"),
                                      (d["estado"], None), (d["receptor"], None)], 1):
            cel(wdc, r, j, v, fmt)
    # ------------------------------------------------------------------ RESUMEN
    cabecera(wr, "MEMORIA DE CALCULO - DRENAJE PLUVIAL (RED DE CUNETAS)", 3)
    res = [("Estacion de referencia", "Pongo de Caynarachi (SENAMHI, codigo 003317), %s, %d anos" % (H.ESTACION["periodo"], A["n"]), None),
           ("Factor aplicado a la serie", "=DATOS!B7", "0.0000"),
           ("Distribucion adoptada", "%s (menor D de Kolmogorov-Smirnov)" % DN[EL], None),
           ("Periodo de retorno (anos)", "=DATOS!B8", "0"), ("Precipitacion maxima en 24 h de diseno (mm)", "=DATOS!B15", "0.00"),
           ("Tiempo de concentracion (min)", "=DATOS!B9", "0.0"), ("Intensidad de diseno (mm/h)", "=DATOS!B17", "0.00"),
           ("Coeficiente de escorrentia ponderado", "=SUPERFICIES!C13", "0.000"), ("Area de aporte dentro del cerco (m2)", "=SUPERFICIES!B13", "#,##0"),
           ("Caudal maximo en un tramo (l/s)", "=MAX(CAUDALES!H9:H%d)" % (rr - 1), "0.0"),
           ("Menor borde libre (m)", "=MIN(CAUDALES!P9:P%d)" % (rr - 1), "0.000"),
           ("Velocidad maxima (m/s)", "=MAX(CAUDALES!S9:S%d)" % (rr - 1), "0.00"),
           ("Tramos que no cumplen capacidad", '=COUNTIF(CAUDALES!Q9:Q%d,"NO CUMPLE")' % (rr - 1), "0"),
           ("Seccion de cuneta", "0.60 m de ancho total: muros de 0.10 m y 0.40 m libres; profundidad variable 0.30 a 0.91 m", None)]
    tabla_cab(wr, 8, ["CONCEPTO", "VALOR", ""], [52, 60, 2])
    for i, (t, v, fmt) in enumerate(res):
        cel(wr, 9 + i, 1, t, bold=True); cel(wr, 9 + i, 2, v, fmt, fill=VE if fmt else None)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    wb.save(SAL); print(SAL)


if __name__ == "__main__":
    construir()
