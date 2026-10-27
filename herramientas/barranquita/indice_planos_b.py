"""Caratula e indice de planos de drenaje pluvial de la I.E. N.° 0590 Barranquita (C.I. 468901) en un solo Word.

Lee (sin modificar) el juego de planos impreso por el usuario: PLANTEAMIENTO_CUNETAS_IE0590_BARRANQUITA.dxf
(laminas de _calc/indice.json, el mismo cuadro de la lamina DC-00; subtitulo del rotulo de cada lamina; ejes y
progresivas de los perfiles; ejes y numero de secciones de las laminas de secciones).
Pagina 1: caratula comun de los documentos (caratula_b, clave "planos"). Pagina 2: indice de planos por grupos.
Salida: entregables_barranquita/CARATULA_E_INDICE_DE_PLANOS_IE0590_BARRANQUITA.docx
"""
import os, re, sys, json
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import caratula_b as C

ENT = os.path.join(RAIZ, "entregables_barranquita")
DXF = os.path.join(ENT, "PLANTEAMIENTO_CUNETAS_IE0590_BARRANQUITA.dxf")
SAL = os.path.join(ENT, "CARATULA_E_INDICE_DE_PLANOS_IE0590_BARRANQUITA.docx")
GRUPOS = (("PLANTA", "PLANTA GENERAL"), ("CUADRO", "CUADRO DE TRAMOS Y SECCIONES TÍPICAS"), ("PERFILES", "PERFILES LONGITUDINALES"),
          ("SECCIONES", "SECCIONES TRANSVERSALES"), ("CORTES", "CORTES GENERALES"), ("DESCARGA", "DETALLES DE DESCARGAS POR EL CERCO"))
TILDES = ((r"\bSECCION\b", "SECCIÓN"), (r"\bTIPICAS\b", "TÍPICAS"), (r"\bTIPICA\b", "TÍPICA"), (r"\bISOMETRICO\b", "ISOMÉTRICO"))


def tildes(t):
    for a, b in TILDES: t = re.sub(a, b, t)
    return t


def leer_planos():
    """Laminas del juego impreso con su subtitulo, ejes y progresivas (o secciones) leidos del DXF."""
    import ezdxf
    indice = json.load(open(os.path.join(ENT, "_calc", "indice.json"), encoding="utf-8"))
    cajas = json.load(open(os.path.join(ENT, "_calc", "laminas.json"), encoding="utf-8"))
    txt = []
    for e in ezdxf.readfile(DXF).modelspace().query("TEXT MTEXT"):
        p = e.dxf.insert; txt.append((p.x, p.y, e.plain_text() if e.dxftype() == "MTEXT" else e.dxf.text))
    prog = lambda s: float(s.replace("+", "")) if s else 0.0
    planos = []
    for cod, titulo, formato, escala in indice:
        x0, y0, x1, y1 = cajas[cod]
        dentro = [(x, y, t) for x, y, t in txt if x0 <= x <= x1 and y0 <= y <= y1]
        px, py = [(x, y) for x, y, t in dentro if t.strip() == "PLANO:"][0]
        rot = [t for _, _, t in sorted((-y, x, t) for x, y, t in dentro if py - 0.03 * (y1 - y0) < y < py + 0.01 * (y1 - y0) and x > px and t.strip() != "PLANO:")]
        sub = rot[1] if len(rot) > 1 else ""
        ejes = {}
        for _, _, t in dentro:
            m = re.match(r"PERFIL LONGITUDINAL (EJE-\d\d) - .* - PROG\. (\d\+\d+\.\d+) A (\d\+\d+\.\d+)$", t.strip())
            if m:
                a, b = ejes.get(m.group(1), (m.group(2), m.group(3)))
                ejes[m.group(1)] = (min(a, m.group(2), key=prog), max(b, m.group(3), key=prog))
        secc = {}
        for _, _, t in dentro:
            m = re.match(r"(EJE-\d\d) - PROG\. ", t.strip())
            if m and titulo.startswith("SECCIONES"): secc[m.group(1)] = secc.get(m.group(1), 0) + 1
        if titulo.startswith("PERFILES"):
            det = "; ".join("%s (%s A %s)" % (k, a, b) for k, (a, b) in sorted(ejes.items()))
        elif titulo.startswith("SECCIONES"):
            k = sorted(secc); det = "%d SECCIONES DEL %s AL %s" % (sum(secc.values()), k[0], k[-1])
        else:
            det = sub
            fin = sorted(set(re.findall(r"ultimos [\d.]+ m del (EJE-\d\d)", " ".join(t for _, _, t in dentro))))
            if fin: det = "SALIDA DEL %s - PLANTA, PERFIL E ISOMETRICO A ESC. 1/25" % fin[0]
        grupo = next(g for k, g in GRUPOS if k in titulo)
        planos.append(dict(codigo=cod, titulo=tildes(titulo), detalle=tildes(det), formato=formato, escala=escala, grupo=grupo))
    return planos


def generar(sal=SAL):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    planos = leer_planos()
    rgb = lambda h: RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))
    d = Document()
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10); st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    s = d.sections[0]; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = Cm(3.0); s.right_margin = Cm(2.5); s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.2)
    if d.paragraphs: d.element.body.remove(d.paragraphs[0]._p)
    C.word(d, "planos")                                                   # pagina 1: caratula

    def par(txt, tam, neg=False, color=None, antes=0, despues=2, centro=True):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.0
        r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg
        if color: r.font.color.rgb = rgb(color)
        return p

    def sombra(c, color):
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color)
        c._tc.get_or_add_tcPr().append(sh)

    def celda(c, lineas, centro=True):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for i, (txt, tam, neg, color) in enumerate(lineas):
            q = c.paragraphs[0] if i == 0 else c.add_paragraph()
            q.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
            q.paragraph_format.space_before = Pt(1 if i == 0 else 0); q.paragraph_format.space_after = Pt(1); q.paragraph_format.line_spacing = 1.0
            r = q.add_run(txt); r.font.size = Pt(tam); r.bold = neg
            if color: r.font.color.rgb = rgb(color)

    # ------------------------------------------------------------------------ pagina 2: indice
    par("ÍNDICE DE PLANOS", 15, True, C.AZUL, despues=2)
    par(C.EXPEDIENTE, 10.5, True, C.AZUL, despues=4)
    par(C.PROYECTO, 9, True, despues=2)
    par(C.CI, 9.5, True, C.AZUL, despues=6)
    anchos = (0.8, 1.6, 9.2, 1.9, 2.0)
    tb = d.add_table(rows=1, cols=5); tb.style = "Table Grid"; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for gc, w in zip(tb._tbl.tblGrid.findall(qn("w:gridCol")), anchos): gc.set(qn("w:w"), str(int(w * 567)))
    for c, w, t in zip(tb.rows[0].cells, anchos, ("N°", "LÁMINA", "DESCRIPCIÓN", "FORMATO", "ESCALA")):
        c.width = Cm(w); sombra(c, C.AZUL); celda(c, [(t, 8, True, "FFFFFF")])
    trPr = tb.rows[0]._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    grupo = None; ng = 0
    for i, p in enumerate(planos, 1):
        if p["grupo"] != grupo:
            grupo = p["grupo"]; ng += 1
            fila = tb.add_row().cells; m = fila[0].merge(fila[-1]); sombra(m, C.AZUL_CLARO)
            celda(m, [("%d. %s" % (ng, grupo), 8.5, True, C.AZUL)], centro=False)
        fila = tb.add_row()
        fila._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit")); fila = fila.cells
        for c, w in zip(fila, anchos): c.width = Cm(w)
        celda(fila[0], [(str(i), 8.5, False, None)])
        celda(fila[1], [(p["codigo"], 8.5, True, C.AZUL)])
        celda(fila[2], [(p["titulo"], 8.5, True, None), (p["detalle"], 7, False, C.GRIS)], centro=False)
        celda(fila[3], [(p["formato"], 8.5, False, None)])
        celda(fila[4], [(p["escala"], 8.5, False, None)])
    formatos = sorted(set(p["formato"] for p in planos))
    par("Total: %d láminas, formato %s (841 x 594 mm)." % (len(planos), " / ".join(formatos)), 9, True, antes=6, despues=2, centro=False)
    par("Archivo: PLANTEAMIENTO_CUNETAS_IE0590_BARRANQUITA.dxf (cada lámina en su presentación, imprimir a escala 1:1). "
        "Coordenadas UTM WGS84; cotas de fondo, niveles y dimensiones en metros.", 8, False, C.GRIS, despues=0, centro=False)
    d.core_properties.author = ""; d.core_properties.title = "Carátula e índice de planos - Drenaje pluvial I.E. N° 0590 Barranquita"
    d.save(sal)
    return sal, planos


if __name__ == "__main__":
    sal, planos = generar()
    for p in planos: print(p["codigo"], "|", p["titulo"], "|", p["detalle"], "|", p["formato"], p["escala"])
    print(sal)
