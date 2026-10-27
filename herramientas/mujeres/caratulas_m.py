# -*- coding: utf-8 -*-
"""Carátulas de los documentos de drenaje pluvial - CAR Mujeres (CUI 2717013), todas en un solo Word.

Diseño con imágenes del propio proyecto (se extraen, sin modificar los originales, de los documentos de Wilma):
  - vista 3D del Centro de Acogida Residencial: insumos/wilma/RESUMEN_EJECUTIVO_CAR_MUJERES_perfil.docx (word/media/image18.jpg),
    sin la sombra del borde inferior;
  - vista aérea del terreno con su lindero: carátula de insumos/wilma/MEMORIA_HIDROLOGICA_CAR_MUJERES.docx (word/media/image1.png),
    solo la foto;
  - logo del Gobierno Regional San Martín: insumos/wilma/MEMORIA_IISS_CAR_MUJERES.docx (word/media/image1.png).
Contenido: entidad, "SERVICIO DE DISEÑO DE DRENAJE PLUVIAL PARA ELABORACIÓN DEL EXPEDIENTE TÉCNICO DE PI:", proyecto con
CUI, vista 3D, franja granate con el título del documento, vista aérea con la ubicación y la fecha. Sin nombre ni CIP del
proyectista. Una página por documento, en el orden de la carpeta de entrega del usuario.

Uso: python3 herramientas/mujeres/caratulas_m.py
Salida: entregables_mujeres/CARATULAS_DRENAJE_PLUVIAL_CAR_MUJERES.docx (imágenes en entregables_mujeres/figuras_caratula/)
"""
import io
import os
import zipfile

import docx
from docx.text.paragraph import Paragraph
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
W = os.path.join(RAIZ, 'insumos/wilma')
FIG = os.path.join(RAIZ, 'entregables_mujeres/figuras_caratula')
SALIDA = os.path.join(RAIZ, 'entregables_mujeres/CARATULAS_DRENAJE_PLUVIAL_CAR_MUJERES.docx')
IMAGENES = {   # nombre: (documento de origen, archivo dentro del docx, recorte)
    'vista_3d_car.jpg': ('RESUMEN_EJECUTIVO_CAR_MUJERES_perfil.docx', 'word/media/image18.jpg', (0, 0, 1289, 790)),
    'vista_aerea_terreno.jpg': ('MEMORIA_HIDROLOGICA_CAR_MUJERES.docx', 'word/media/image1.png', (215, 860, 995, 1245)),
    'escudo_gore_san_martin.png': ('MEMORIA_IISS_CAR_MUJERES.docx', 'word/media/image1.png', (60, 12, 166, 124)),
}
GRANATE, GRANATE_CLARO, GRIS = '7B1E2B', 'F3E6E8', '595959'
ENTIDAD = 'GOBIERNO REGIONAL SAN MARTÍN'
GERENCIA = 'GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO'
SERVICIO = 'SERVICIO DE DISEÑO DE DRENAJE PLUVIAL PARA ELABORACIÓN DEL EXPEDIENTE TÉCNICO DE PI:'
PROYECTO = ('“CREACIÓN DEL SERVICIO DE PROTECCIÓN INTEGRAL A NIÑAS, NIÑOS Y ADOLESCENTES SIN CUIDADOS PARENTALES O EN RIESGO '
            'DE PERDERLOS EN CENTRO DE ACOGIDA RESIDENCIAL - MUJERES DISTRITO DE MORALES DE LA PROVINCIA DE SAN MARTÍN DEL '
            'DEPARTAMENTO DE SAN MARTÍN”')
CUI = 'CUI N° 2717013'
FECHA = 'SEPTIEMBRE - 2026'
ESPECIALIDAD = 'ESPECIALIDAD DE DRENAJE PLUVIAL'
DOCUMENTOS = [   # (título, subtítulo) en el orden de la carpeta de entrega
    ('METRADOS', ESPECIALIDAD),
    ('MEMORIA HIDROLÓGICA', ESPECIALIDAD),
    ('MEMORIA DESCRIPTIVA', ESPECIALIDAD),
    ('MEMORIA DE CÁLCULO', 'COLECTOR PLUVIAL FRONTAL - ' + ESPECIALIDAD),
    ('MEMORIA DE CÁLCULO', 'DRENAJE PLUVIAL INTERNO - ' + ESPECIALIDAD),
    ('ESPECIFICACIONES TÉCNICAS', ESPECIALIDAD),
]
ANCHO = 17.0     # cm útiles (A4 con márgenes de 2.0 cm)


def imagenes():
    from PIL import Image
    os.makedirs(FIG, exist_ok=True)
    rutas = {}
    for nombre, (doc, miembro, caja) in IMAGENES.items():
        ruta = os.path.join(FIG, nombre)
        if not os.path.exists(ruta):
            with zipfile.ZipFile(os.path.join(W, doc)) as z:
                im = Image.open(io.BytesIO(z.read(miembro)))
            im = im.crop(caja)
            if nombre.endswith('.jpg'):
                im.convert('RGB').save(ruta, quality=90, optimize=True)
            else:
                im.save(ruta)
        rutas[nombre] = ruta
    return rutas


ORDEN_TBLPR = ('tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize', 'tblStyleColBandSize', 'tblW', 'jc',
               'tblCellSpacing', 'tblInd', 'tblBorders', 'shd', 'tblLayout', 'tblCellMar', 'tblLook')
ORDEN_TCPR = ('cnfStyle', 'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap', 'tcMar', 'textDirection',
              'tcFitText', 'vAlign', 'hideMark')


def _poner_en_orden(padre, elem, orden):
    """Inserta (o reemplaza) elem en padre respetando el orden del esquema de Word."""
    nombre = elem.tag.split('}')[1]
    viejo = padre.find(qn('w:' + nombre))
    if viejo is not None:
        padre.remove(viejo)
    despues = orden[orden.index(nombre) + 1:]
    for i, ch in enumerate(padre):
        if ch.tag.split('}')[1] in despues:
            padre.insert(i, elem); return
    padre.append(elem)


def _sombra(c, color):
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    _poner_en_orden(c._tc.get_or_add_tcPr(), sh, ORDEN_TCPR)


def _bordes(c, **lados):
    """Bordes de la celda: lados = {'left': (tam, color), ...}; los no indicados quedan sin borde."""
    tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + lado)
        if lado in lados:
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(lados[lado][0])); e.set(qn('w:color'), lados[lado][1])
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)
    _poner_en_orden(tcPr, b, ORDEN_TCPR)


def _sin_bordes(tb):
    tblPr = tb._tbl.tblPr; b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + lado); e.set(qn('w:val'), 'nil'); b.append(e)
    _poner_en_orden(tblPr, b, ORDEN_TBLPR)


def _anchos(tb, anchos):
    tb.autofit = False
    tblPr = tb._tbl.tblPr
    for tag, val in (('w:tblW', sum(anchos)), ('w:tblInd', 0)):
        e = OxmlElement(tag); e.set(qn('w:w'), str(int(val * 567))); e.set(qn('w:type'), 'dxa')
        _poner_en_orden(tblPr, e, ORDEN_TBLPR)
    lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); _poner_en_orden(tblPr, lay, ORDEN_TBLPR)
    cm = OxmlElement('w:tblCellMar')                                  # sin relleno lateral: la tabla mide lo mismo que la foto
    for lado in ('left', 'right'):
        e = OxmlElement('w:' + lado); e.set(qn('w:w'), '0'); e.set(qn('w:type'), 'dxa'); cm.append(e)
    _poner_en_orden(tblPr, cm, ORDEN_TBLPR)
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(int(w * 567)))
    for fila in tb.rows:
        for c, w in zip(fila.cells, anchos):
            c.width = Cm(w)


def _par(cont, txt='', tam=10, neg=False, color=None, centro=True, antes=0, despues=0, esp=1.0, cursiva=False, sep=0):
    p = cont if isinstance(cont, Paragraph) else cont.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = esp
    if txt:
        r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.italic = cursiva; r.font.name = 'Arial'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        if sep:
            sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(sep)); r._element.rPr.append(sp)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
    return p


def _linea(p, color=GRANATE, tam=18):
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', str(tam)), ('space', '4'), ('color', color)):
        b.set(qn('w:' + k), v)
    bdr.append(b); pPr.append(bdr)


def caratula(d, img, titulo, subtitulo, primera):
    if not primera:                                                  # cada carátula en su página
        p = _par(d); p.paragraph_format.page_break_before = True
        rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rpr.append(sz)
        p._p.get_or_add_pPr().append(rpr)                           # párrafo mínimo: no empuja la carátula
    # ---- entidad con logo
    tb = d.add_table(rows=1, cols=2); _sin_bordes(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    _anchos(tb, (2.6, ANCHO - 2.6))
    c0, c1 = tb.rows[0].cells
    c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _par(c0.paragraphs[0]).add_run().add_picture(img['escudo_gore_san_martin.png'], width=Cm(2.2))
    _par(c1.paragraphs[0], ENTIDAD, 17, True, GRANATE, centro=False, sep=10)
    _linea(_par(c1, GERENCIA, 10.5, True, GRIS, centro=False, despues=2, sep=6))
    # ---- servicio y proyecto
    _par(d, SERVICIO, 11, True, GRANATE, antes=18, despues=6)
    _par(d, PROYECTO, 10, True, despues=3, esp=1.05)
    _par(d, CUI, 11, True, GRANATE, despues=14)
    # ---- vista 3D del CAR
    tb = d.add_table(rows=1, cols=1); _sin_bordes(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; _anchos(tb, (ANCHO,))
    _par(tb.rows[0].cells[0].paragraphs[0]).add_run().add_picture(img['vista_3d_car.jpg'], width=Cm(ANCHO))
    _par(d, 'Vista 3D del Centro de Acogida Residencial - Mujeres', 8, False, GRIS, cursiva=True, antes=2, despues=12)
    # ---- franja con el título del documento
    tb = d.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; _anchos(tb, (ANCHO,))
    c = tb.rows[0].cells[0]; _sombra(c, GRANATE); _bordes(c)
    _par(c.paragraphs[0], titulo, 24, True, 'FFFFFF', antes=10, despues=1, sep=8)
    _par(c, subtitulo, 11, True, 'F2D7DB', despues=10, sep=4)
    # ---- ubicación: vista aérea del terreno y datos
    tb = d.add_table(rows=1, cols=2); _sin_bordes(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    _anchos(tb, (8.6, ANCHO - 8.6))
    c0, c1 = tb.rows[0].cells
    c0.vertical_alignment = c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _par(c0.paragraphs[0], antes=22).add_run().add_picture(img['vista_aerea_terreno.jpg'], width=Cm(8.3))
    _par(c0, 'Vista aérea del terreno del proyecto', 7.5, False, GRIS, cursiva=True, antes=1)
    _bordes(c1, left=(24, GRANATE))
    _par(c1.paragraphs[0], '   UBICACIÓN DEL PROYECTO', 10, True, GRANATE, centro=False, antes=22, despues=4, sep=6)
    for etq, val in (('Distrito', 'Morales'), ('Provincia', 'San Martín'), ('Departamento', 'San Martín')):
        p = _par(c1, '   %s:  ' % etq, 10, False, GRIS, centro=False, despues=2)
        r = p.add_run(val); r.font.size = Pt(10); r.bold = True; r.font.name = 'Arial'
    _par(c1, '   ' + FECHA, 13, True, GRANATE, centro=False, antes=14, sep=6)


def generar():
    img = imagenes()
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); st.paragraph_format.space_after = Pt(0)
    s = d.sections[0]; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = Cm(2.3); s.bottom_margin = Cm(1.2)
    if d.paragraphs:
        d.element.body.remove(d.paragraphs[0]._p)
    for k, (t, sub) in enumerate(DOCUMENTOS):
        caratula(d, img, t, sub, k == 0)
    d.core_properties.author = ''
    d.core_properties.title = 'Carátulas - Drenaje pluvial CAR Mujeres'
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
