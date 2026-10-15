"""CAR Mujeres (CUI 2717013): analisis de costos unitarios de las partidas GLOBALES (glb) del metrado de instalaciones
sanitarias (insumos/wilma/METRADOS_IISS_CAR_MUJERES_DREA.xlsx, hoja PLANILLA).

Partidas: 4.2.3.1.1 a 4.2.3.1.7 (lineas de cisterna y tanque elevado, prueba hidraulica) y 4.3.6.3.3 (tecle de palanca
para caballete). Las cantidades de materiales se leen de la planilla (no se cambia el metrado); las filas con unidad o
parcial irregular se toman como se indica en OBS y se avisa en cada ACU.
Mano de obra: tabla oficial CAPECO - FTCCP 2026 (operario 28.95, oficial 22.68, peon 20.54 S/ hh).
Materiales: precios de venta publicados (Sodimac Peru, Promart, etc.) con IGV, pasados a precio sin IGV (/1.18).
Celda amarilla = dato editable; verde = formula. El libro se recalcula con LibreOffice para dejar los valores guardados.
Uso: python3 herramientas/mujeres/acu_globales_m.py
"""
import os, math, shutil, subprocess, tempfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__)); HERR = os.path.dirname(AQUI); RAIZ = os.path.dirname(HERR)
METRADO = os.path.join(RAIZ, "insumos", "wilma", "METRADOS_IISS_CAR_MUJERES_DREA.xlsx")
SALIDA = os.path.join(RAIZ, "entregables_mujeres", "ACU_PARTIDAS_GLOBALES_IISS_CAR_MUJERES.xlsx")
PROYECTO = ("CREACIÓN DEL SERVICIO DE PROTECCIÓN INTEGRAL A NIÑAS, NIÑOS Y ADOLESCENTES SIN CUIDADOS PARENTALES O EN RIESGO DE "
            "PERDERLOS EN CENTRO DE ACOGIDA RESIDENCIAL - MUJERES, DISTRITO DE MORALES, PROVINCIA Y DEPARTAMENTO DE SAN MARTÍN "
            "(CAR Mujeres, CUI 2717013)")
LUGAR = "San Martín - San Martín - Morales"
FECHA = "Octubre 2026"
IGV = 1.18

# ------------------------------------------------------------------ estilos
F_T = Font(name="Arial", size=12, bold=True, color="FFFFFF")
F_B = Font(name="Arial", size=9, bold=True)
F_N = Font(name="Arial", size=9)
F_S = Font(name="Arial", size=8, italic=True, color="404040")
F_H = Font(name="Arial", size=9, bold=True, color="FFFFFF")
AZUL = PatternFill("solid", fgColor="1F4E78")
CEL = PatternFill("solid", fgColor="DDEBF7")
GRIS = PatternFill("solid", fgColor="F2F2F2")
AMA = PatternFill("solid", fgColor="FFF2CC")
VER = PatternFill("solid", fgColor="E2EFDA")
NAR = PatternFill("solid", fgColor="FCE4D6")
LIN = Side(style="thin", color="808080")
BOR = Border(left=LIN, right=LIN, top=LIN, bottom=LIN)
IZQ = Alignment(horizontal="left", vertical="center", wrap_text=True)
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DER = Alignment(horizontal="right", vertical="center")


def put(ws, ref, v, font=F_N, fill=None, al=IZQ, fmt=None, borde=True):
    c = ws[ref]
    if not (v is None and type(c).__name__ == "MergedCell"):
        c.value = v
    c.font = font; c.alignment = al
    if fill: c.fill = fill
    if borde: c.border = BOR
    if fmt: c.number_format = fmt
    return c


def banda(ws, r, texto, c1="A", c2="H", fill=AZUL, font=F_H, alto=18):
    ws.merge_cells("%s%d:%s%d" % (c1, r, c2, r))
    for col in range(ord(c1), ord(c2) + 1):
        put(ws, "%s%d" % (chr(col), r), None, fill=fill)
    put(ws, "%s%d" % (c1, r), texto, font=font, fill=fill)
    ws.row_dimensions[r].height = alto


def texto_largo(ws, r, texto, c1="A", c2="H", alto=30, font=F_S, fill=None):
    ws.merge_cells("%s%d:%s%d" % (c1, r, c2, r))
    put(ws, "%s%d" % (c1, r), texto, font=font, fill=fill, borde=False)
    ws.row_dimensions[r].height = alto


def anchos(ws, d):
    for k, v in d.items():
        ws.column_dimensions[k].width = v


def pagina(ws, horizontal=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if horizontal else "portrait"
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.sheet_view.showGridLines = False


# ------------------------------------------------------------------ mano de obra CAPECO - FTCCP 2026
HOJA_MO = "MANO DE OBRA 2026"
MO_FUENTE = ("Tabla oficial CAPECO - FTCCP «Remuneración y condiciones de trabajo en construcción civil 2026» (Acta final de "
             "negociación colectiva, Exp. N.° 339-2025-MTPE/2.14-NC), costo hora-hombre en edificación del 01.01.2026 al 31.12.2026.")
MO_FILAS = [  # concepto, operario, oficial, peon, sustento
    ("Remuneración básica (jornal diario)", 89.30, 69.75, 62.80, "R.M. 197-2025-TR; tabla oficial"),
    ("Leyes y beneficios sociales sobre el jornal básico (112.20 %)", "=ROUND({c}9*112.2/100,2)", None, None, "tabla oficial"),
    ("Bonificación unificada de construcción BUC (32 % / 30 % / 30 %)", "=0", "=0", "=0", "tabla oficial"),
    ("Leyes y beneficios sociales sobre el BUC (12.25 %)", "=ROUND({c}11*12.25/100,2)", None, None, "tabla oficial"),
    ("Seguro de vida ESSALUD + Vida", 0.16, 0.16, 0.16, "tabla oficial"),
    ("Fondo de capacitación", 0.40, 0.31, 0.28, "tabla oficial"),
    ("Bonificación por movilidad (R.M. N.° 139-2024-TR)", 8.60, 8.60, 8.60, "S/ 8.60 por día laborado"),
    ("Overol (R.D. N.° 777-87-DR-LIM): (2 x S/ 132.80) / 300", "=ROUND(2*132.8/300,2)", None, None, "tabla oficial"),
]


def construir_mo(wb):
    ws = wb.create_sheet(HOJA_MO)
    anchos(ws, {"A": 5, "B": 54, "C": 13, "D": 13, "E": 13, "F": 34})
    banda(ws, 1, "COSTO DE LA HORA-HOMBRE 2026 - RÉGIMEN DE CONSTRUCCIÓN CIVIL", "A", "F", font=F_T, alto=24)
    texto_largo(ws, 2, PROYECTO, "A", "F", alto=36)
    texto_largo(ws, 3, MO_FUENTE, "A", "F", alto=30)
    banda(ws, 5, "DESGLOSE DEL COSTO DIARIO EN EDIFICACIÓN (S/ por día de 8 horas)", "A", "F", fill=CEL, font=F_B)
    for c, t in zip("ABCDEF", ("N.°", "CONCEPTO", "OPERARIO", "OFICIAL", "PEÓN", "SUSTENTO")):
        put(ws, "%s6" % c, t, F_B, GRIS, CEN)
    r0 = 7
    for i, (t, a, b, p, s) in enumerate(MO_FILAS):
        r = r0 + i
        put(ws, "A%d" % r, i + 1, al=CEN); put(ws, "B%d" % r, t); put(ws, "F%d" % r, s, F_S)
        vals = (a, b if b is not None else a, p if p is not None else a)
        for col, v in zip("CDE", vals):
            if isinstance(v, str):
                v = v.replace("{c}9", "%s%d" % (col, r0)).replace("{c}11", "%s%d" % (col, r0 + 2))
                put(ws, "%s%d" % (col, r), v, fill=VER, al=DER, fmt="0.00")
            else:
                put(ws, "%s%d" % (col, r), v, fill=AMA, al=DER, fmt="0.00")
    # la fila BUC usa columnas distintas por categoria
    rb = r0 + 2
    ws["C%d" % rb] = "=ROUND(C%d*0.32,2)" % r0; ws["D%d" % rb] = "=ROUND(D%d*0.30,2)" % r0; ws["E%d" % rb] = "=ROUND(E%d*0.30,2)" % r0
    rt = r0 + len(MO_FILAS)
    put(ws, "B%d" % rt, "Total por día de 8 horas (S/)", F_B, GRIS)
    put(ws, "B%d" % (rt + 1), "COSTO DE LA HORA-HOMBRE (S/ hh) = total / 8", F_B, GRIS)
    for col in "CDE":
        put(ws, "%s%d" % (col, rt), "=SUM(%s%d:%s%d)" % (col, r0, col, rt - 1), F_B, VER, DER, "0.00")
        put(ws, "%s%d" % (col, rt + 1), "=ROUND(%s%d/8,2)" % (col, rt), F_B, VER, DER, "0.00")
    for r in (rt, rt + 1):
        put(ws, "A%d" % r, None, fill=GRIS); put(ws, "F%d" % r, "tabla oficial: 231.62 / 181.46 / 164.34 y 28.95 / 22.68 / 20.54" if r == rt + 1 else None, F_S, GRIS)
    pagina(ws)
    return {"Operario": "'%s'!$C$%d" % (HOJA_MO, rt + 1), "Oficial": "'%s'!$D$%d" % (HOJA_MO, rt + 1),
            "Peón": "'%s'!$E$%d" % (HOJA_MO, rt + 1)}


# ------------------------------------------------------------------ precios de insumos
# clave: (insumo ACU, und ACU, producto publicado, presentacion, precio publicado con IGV, contenido en und ACU, tienda, url, estado)
PRECIOS = {}
from precios_globales_m import PRECIOS  # noqa: E402  (tabla de precios con fuentes)

HOJA_P = "PRECIOS DE INSUMOS"


def construir_precios(wb):
    ws = wb.create_sheet(HOJA_P)
    anchos(ws, {"A": 8, "B": 40, "C": 6, "D": 44, "E": 13, "F": 11, "G": 10, "H": 11, "I": 20, "J": 50, "K": 26})
    banda(ws, 1, "PRECIOS DE INSUMOS (materiales y equipos) - SUSTENTO", "A", "K", font=F_T, alto=24)
    texto_largo(ws, 2, PROYECTO, "A", "K", alto=30)
    texto_largo(ws, 3, "Precio publicado con IGV en la tienda indicada (venta nacional; Sodimac y Promart tienen tienda en Tarapoto) y "
                       "precio unitario sin IGV = precio publicado / contenido / 1.18. Las ACU leen la columna H. Celda amarilla = dato "
                       "editable (el ingeniero de costos puede reemplazarlo por su cotización y todo se recalcula); verde = fórmula.",
                "A", "K", alto=30)
    cab = ("CÓDIGO", "INSUMO (como figura en la ACU)", "UND", "PRODUCTO PUBLICADO", "PRESENTACIÓN", "PRECIO CON IGV (S/)",
           "CONTENIDO (und ACU)", "PRECIO UNIT. SIN IGV (S/)", "TIENDA / FUENTE", "ENLACE", "ESTADO")
    for c, t in zip("ABCDEFGHIJK", cab):
        put(ws, "%s5" % c, t, F_B, GRIS, CEN)
    ws.row_dimensions[5].height = 30
    ref = {}
    r = 6
    banda(ws, r, "1. INSUMOS COMERCIALES", "A", "K", fill=CEL, font=F_B); r += 1
    for k, (ins, und, prod, pres, pv, cont, tienda, url, est) in PRECIOS.items():
        if k.startswith("_"):
            continue
        put(ws, "A%d" % r, k, al=CEN); put(ws, "B%d" % r, ins); put(ws, "C%d" % r, und, al=CEN); put(ws, "D%d" % r, prod)
        put(ws, "E%d" % r, pres, al=CEN)
        put(ws, "F%d" % r, pv, fill=AMA, al=DER, fmt="#,##0.00")
        put(ws, "G%d" % r, cont, fill=AMA if not isinstance(cont, str) else VER, al=DER, fmt="0.000")
        fac = 1.0 if k == "AGUA" else IGV
        put(ws, "H%d" % r, "=ROUND(F%d/G%d/%s,2)" % (r, r, fac), F_B, VER, DER, "#,##0.00")
        put(ws, "I%d" % r, tienda, F_S); put(ws, "J%d" % r, url, F_S); put(ws, "K%d" % r, est, F_S)
        ref[k] = "'%s'!$H$%d" % (HOJA_P, r)
        ws.row_dimensions[r].height = 30
        r += 1
    ref["NFG15"] = ref["TFG15"]  # niple F°G° 1 1/2" L=0.60 se valoriza por metro de tubo
    # insumos compuestos: bridas rompe agua (plancha + niple F°G° + soldadura)
    r += 1
    banda(ws, r, "2. BRIDAS ROMPE AGUA (insumo compuesto, fabricado en obra): plancha LAC 1/8\" (24.91 kg/m²) + niple F°G° + soldadura",
          "A", "K", fill=CEL, font=F_B); r += 1
    for c, t in zip("ABCDEFGHIJK", ("CÓDIGO", "BRIDA", "UND", "PLANCHA (m x m)", "PLANCHA (kg)", "NIPLE F°G° (Ø)", "NIPLE L (m)",
                                     "PRECIO SIN IGV (S/)", "SOLDADURA (kg)", "CÁLCULO", "OBS.")):
        put(ws, "%s%d" % (c, r), t, F_B, GRIS, CEN)
    r += 1
    for k, (nom, lado, tubo, L, sold, obs) in BRIDAS.items():
        put(ws, "A%d" % r, k, al=CEN); put(ws, "B%d" % r, nom); put(ws, "C%d" % r, "und", al=CEN)
        put(ws, "D%d" % r, lado, fill=AMA, al=DER, fmt="0.00")
        put(ws, "E%d" % r, "=ROUND(D%d^2*24.91,2)" % r, fill=VER, al=DER, fmt="0.00")
        put(ws, "F%d" % r, PRECIOS[tubo][0].split("(")[0].strip(), al=CEN)
        put(ws, "G%d" % r, L, fill=AMA, al=DER, fmt="0.00")
        put(ws, "I%d" % r, sold, fill=AMA, al=DER, fmt="0.00")
        put(ws, "H%d" % r, "=ROUND(E%d*%s+G%d*%s+I%d*%s,2)" % (r, ref["PLA"], r, ref[tubo], r, ref["SOLD"]), F_B, VER, DER, "#,##0.00")
        put(ws, "J%d" % r, "plancha kg x precio kg + niple m x precio m del tubo F°G° + soldadura kg x precio kg", F_S)
        put(ws, "K%d" % r, obs, F_S)
        ref[k] = "'%s'!$H$%d" % (HOJA_P, r)
        ws.row_dimensions[r].height = 30
        r += 1
    ws.freeze_panes = "C6"
    pagina(ws, horizontal=True)
    return ref


# brida: (descripcion, lado plancha m, tubo F°G° para el niple, largo niple m, soldadura kg, obs)
BRIDAS = {
    "BR01": ('Brida rompe agua Ø1" PL 0.15x0.15 e=1/8" (niple F°G° 3/4" L=0.35)', 0.15, "TFG34", 0.35, 0.10, ""),
    "BR02": ('Brida rompe agua Ø1 1/2" PL 0.15x0.15 e=1/8" (niple F°G° 1 1/2" L=0.25)', 0.15, "TFG15", 0.25, 0.12, ""),
    "BR03": ('Brida rompe agua Ø2" PL 0.25x0.25 e=1/8" (niple F°G° 2" L=0.40)', 0.25, "TFG2", 0.40, 0.15, ""),
    "BR04": ('Brida rompe agua Ø2" PL 0.15x0.15 e=1/8" (niple F°G° 2" L=0.25)', 0.15, "TFG2", 0.25, 0.15, ""),
    "BR05": ('Brida rompe agua Ø2" PL 0.20x0.20 e=1/8" (niple F°G° 2" L=0.45)', 0.20, "TFG2", 0.45, 0.15, "rosca a un lado"),
}

# ------------------------------------------------------------------ rendimientos por tipo de componente (und/dia de la cuadrilla)
REND = {  # tipo: (descripcion, rendimiento und/dia, und)
    "tubo1": ("Tubería PVC SP Ø1\" (tendido, corte, unión y fijación)", 25, "m"),
    "tubo15": ("Tubería PVC SP Ø1 1/2\"", 20, "m"),
    "tubo2": ("Tubería PVC SP Ø2\"", 18, "m"),
    "acc15": ("Accesorio PVC hasta Ø1 1/2\" (codo, tee, unión, niple, reducción, tapón)", 20, "und"),
    "acc2": ("Accesorio PVC Ø2\" (codo, tee, unión, adaptador, niple)", 16, "und"),
    "fg": ("Accesorio F°G° roscado (unión, codo, niple)", 16, "und"),
    "valv": ("Válvula roscada (compuerta, esférica, check)", 8, "und"),
    "flot": ("Válvula flotadora con boya (instalación y regulación)", 4, "und"),
    "brida": ("Brida rompe agua (corte, soldadura, roscado y fijación al encofrado)", 3, "und"),
    "abr": ("Abrazadera F°G° (perforación, tarugo y fijación)", 16, "und"),
    "elec": ("Electronivel (montaje, regulación y conexión al tablero de bombas)", 3, "und"),
    "canast": ("Canastilla de succión", 8, "und"),
    "somb": ("Sombrero de ventilación", 10, "und"),
}

# ------------------------------------------------------------------ partidas y lectura de la planilla
# fila: (clave precio, tipo de rendimiento, cantidad corregida o None, factor desperdicio, observacion)
PARTIDAS = [
    ("4.2.3.1.1", 437, ("1op", "0.5pe"), {
        438: ("TB1", "tubo1", None, 1.05, ""),
        439: ("VF1", "flot", None, 1, ""),
        440: ("UFG1", "fg", 1, 1, "fila 440: la planilla no lleva el parcial (columna I vacía); se toma N.° veces 1 x cantidad 1 = 1 und"),
        441: ("BR01", "brida", None, 1, ""),
    }, "Ingreso a cisterna. El título dice D=1 1/4\" pero el desagregado de la planilla es de Ø1\"; se analiza lo desagregado."),
    ("4.2.3.1.2", 443, ("1op", "0.5pe"), {
        444: ("BR02", "brida", None, 1, ""),
        445: ("TB15", "tubo15", None, 1.05, "fila 445: «Canastilla de PVC-U Ø1 1/2\"» figura con 4.64 m; se analiza como 4.64 m de tubería "
                                         "Ø1 1/2\" del tramo de succión + 2 canastillas (una por bomba, igual que las 2 tees, 2 tapones y 2 electroniveles)"),
        "445b": ("CAN15", "canast", 2, 1, ""),
        446: ("SK15", "acc15", None, 1, ""),
        447: ("TE15", "acc15", None, 1, ""),
        448: ("TP15", "acc15", None, 1, ""),
        449: ("ELN", "elec", None, 1, ""),
    }, ""),
    ("4.2.3.1.3", 452, ("1op", "0.5pe"), {
        453: ("CD15", "acc15", None, 1, ""), 454: ("VCH15", "valv", None, 1, ""), 455: ("UU15", "acc15", None, 1, ""),
        456: ("VEP15", "valv", None, 1, ""), 457: ("TE15", "acc15", None, 1, ""), 458: ("TB15", "tubo15", None, 1.05, ""),
        459: ("NP15", "acc15", None, 1, ""), 460: ("RD15", "acc15", None, 1, ""), 461: ("UU15", "acc15", None, 1, ""),
        462: ("VEB15", "valv", None, 1, ""), 463: ("TP15", "acc15", None, 1, ""), 464: ("NP15", "acc15", None, 1, ""),
        465: ("AB15", "abr", None, 1, ""), 466: ("SK15", "acc15", None, 1, ""), 467: ("CFG15", "fg", None, 1, ""),
        468: ("BR02", "brida", None, 1, "fila 468: la brida es Ø1 1/2\" pero el texto dice niple F°G° 1\"; se analiza con niple Ø1 1/2\" (el del diámetro de la línea)"),
        469: ("NFG15", "fg", None, 1, ""), 470: ("UFG15", "fg", None, 1, ""), 471: ("VF15", "flot", None, 1, ""),
        472: ("ELN", "elec", None, 1, ""),
    }, "Impulsión. El título dice D=1\" pero el desagregado de la planilla es de Ø1 1/2\"; se analiza lo desagregado."),
    ("4.2.3.1.4", 474, ("1op", "0.5pe"), {
        475: ("BR03", "brida", None, 1, ""), 476: ("SK2", "acc2", None, 1, ""),
        477: ("CD2", "acc2", 1, 1, "filas 477-478: las unidades están cruzadas (codo 2\" con 17 «m» y tubería 2\" con 1 «und»); "
                                  "se analiza tubería Ø2\" = 17 m (la medida está en la columna LARGO) y codo Ø2\" x 90° = 1 und"),
        478: ("TB2", "tubo2", 17, 1.05, ""),
        479: ("UPR2", "acc2", None, 1, ""), 480: ("UU2", "acc2", None, 1, ""), 481: ("NP2", "acc2", None, 1, ""),
        482: ("VCB2", "valv", None, 1, ""), 483: ("AB2", "abr", None, 1, ""), 484: ("TE2", "acc2", None, 1, ""),
        485: ("CD245", "acc2", None, 1, ""),
    }, ""),
    ("4.2.3.1.5", 487, ("1op", "0.5pe"), {
        488: ("RD2", "acc2", None, 1, ""), 489: ("TB1", "tubo1", None, 1.05, ""), 490: ("AB1", "abr", None, 1, ""),
        491: ("CD1", "acc15", None, 1, ""), 492: ("BR04", "brida", None, 1, ""), 493: ("TE2", "acc2", None, 1, ""),
        494: ("NP2", "acc2", None, 1, ""), 495: ("SK2", "acc2", None, 1, ""), 496: ("AB2", "abr", None, 1, ""),
        497: ("TB2", "tubo2", None, 1.05, ""), 498: ("SOMB2", "somb", None, 1, ""), 499: ("TE2", "acc2", None, 1, ""),
        500: ("UPR2", "acc2", None, 1, ""), 501: ("UU2", "acc2", None, 1, ""), 502: ("NP2", "acc2", None, 1, ""),
        503: ("VCB2", "valv", None, 1, ""), 504: ("CD2", "acc2", None, 1, ""),
    }, ""),
    ("4.2.3.1.6", 506, ("1op", "0.5pe"), {
        507: ("BR05", "brida", 2, 1, "fila 507: la planilla no lleva el parcial; se toma N.° veces 1 x cantidad 2 = 2 und"),
        508: ("CD2", "acc2", None, 1, ""), 509: ("NP2", "acc2", None, 1, ""),
        510: ("CD2", "acc2", 2, 1, "fila 510: la planilla no lleva el parcial; se toma 2 und (codo Ø2\" + malla galvanizada soldada)"),
        "510b": ("MALLA", None, 0.10, 1, ""),
    }, ""),
]
# consumibles: (uniones pegadas, uniones roscadas) por tipo de accesorio
UNIONES = {"TB1": (0.2, 0), "TB15": (0.2, 0), "TB2": (0.2, 0), "CD1": (2, 0), "CD15": (2, 0), "CD2": (2, 0), "CD245": (2, 0),
           "TE15": (3, 0), "TE2": (3, 0), "SK15": (1, 1), "SK2": (1, 1), "TP15": (0, 1), "UU15": (0, 2), "UU2": (0, 2),
           "UPR2": (1, 1), "NP15": (0, 2), "NP2": (0, 2), "RD15": (2, 0), "RD2": (2, 0), "VCH15": (0, 2), "VEP15": (0, 2),
           "VEB15": (0, 2), "VCB2": (0, 2), "VF1": (0, 1), "VF15": (0, 1), "UFG1": (0, 2), "UFG15": (0, 2), "CFG15": (0, 2),
           "NFG15": (0, 2), "BR01": (0, 2), "BR02": (0, 2), "BR03": (0, 2), "BR04": (0, 2), "BR05": (0, 1), "CAN15": (0, 1),
           "SOMB2": (1, 0), "ELN": (0, 0), "AB1": (0, 0), "AB15": (0, 0), "AB2": (0, 0), "MALLA": (0, 0)}


def leer_planilla():
    wb = openpyxl.load_workbook(METRADO, data_only=True)
    ws = wb["PLANILLA"]
    return ws


def cantidad(ws, r):
    v = ws.cell(r, 9).value  # PARCIAL
    if v in (None, ""):
        return None
    return float(v)


def hoja_acu(wb, nombre, item, titulo, metrado, mo, mats, equipos, rend_tabla, cuadrilla, obs, P, HH, extra_rend=None):
    """mats: lista (codigo precio, descripcion, und, cantidad formula/numero, sustento)."""
    ws = wb.create_sheet(nombre)
    anchos(ws, {"A": 8, "B": 50, "C": 7, "D": 10, "E": 11, "F": 11, "G": 12, "H": 42})
    banda(ws, 1, "ANÁLISIS DE COSTOS UNITARIOS", font=F_T, alto=24)
    texto_largo(ws, 2, PROYECTO, alto=30, font=F_S)
    put(ws, "A3", "Partida", F_B, GRIS); ws.merge_cells("B3:H3"); put(ws, "B3", "%s  %s" % (item, titulo), F_B, CEL)
    for c in "CDEFGH":
        ws["%s3" % c].border = BOR
    ws.row_dimensions[3].height = 30
    put(ws, "A4", "Rend.", F_B, GRIS); put(ws, "B4", None, F_B, VER, DER, "0.0000 \"glb/día\"")
    put(ws, "C4", "MO", F_B, GRIS, CEN); put(ws, "D4", None, F_B, VER, DER, "0.0000")
    put(ws, "E4", "EQ", F_B, GRIS, CEN); put(ws, "F4", None, F_B, VER, DER, "0.0000")
    put(ws, "G4", "Costo unit. directo por glb (S/)", F_B, GRIS, CEN); put(ws, "H4", None, Font(name="Arial", size=11, bold=True), VER, DER, "#,##0.00")
    ws.row_dimensions[4].height = 30
    put(ws, "A5", "Metrado", F_B, GRIS); put(ws, "B5", "%s glb (planilla, hoja RESUMEN)" % metrado, F_N)
    put(ws, "C5", "Lugar", F_B, GRIS, CEN); ws.merge_cells("D5:F5"); put(ws, "D5", LUGAR)
    put(ws, "G5", "Fecha", F_B, GRIS, CEN); put(ws, "H5", FECHA)
    r = 7
    for c, t in zip("ABCDEFGH", ("CÓDIGO", "DESCRIPCIÓN DEL RECURSO", "UND", "CUADRILLA", "CANTIDAD", "PRECIO S/", "PARCIAL S/", "SUSTENTO")):
        put(ws, "%s%d" % (c, r), t, F_H, AZUL, CEN)
    r += 1
    # tabla de rendimiento al final; primero fijamos su posicion
    filas_mo, filas_mat, filas_eq = [], [], []
    # --- mano de obra
    banda(ws, r, "MANO DE OBRA", fill=CEL, font=F_B, alto=16); r += 1
    r_mo0 = r
    for cat, cu in cuadrilla:
        put(ws, "A%d" % r, {"Operario": "47-0001", "Oficial": "47-0002", "Peón": "47-0003"}[cat], al=CEN)
        put(ws, "B%d" % r, cat); put(ws, "C%d" % r, "hh", al=CEN)
        put(ws, "D%d" % r, cu, fill=AMA, al=DER, fmt="0.00")
        put(ws, "E%d" % r, "=ROUND(D%d*8/$B$4,4)" % r, fill=VER, al=DER, fmt="0.0000")
        put(ws, "F%d" % r, "=" + HH[cat], fill=VER, al=DER, fmt="0.00")
        put(ws, "G%d" % r, "=ROUND(E%d*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
        put(ws, "H%d" % r, "cantidad = cuadrilla x 8 h / rendimiento; costo HH CAPECO - FTCCP 2026", F_S)
        filas_mo.append(r); r += 1
    r_mo = r
    put(ws, "B%d" % r, "Subtotal mano de obra", F_B, GRIS); put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (r_mo0, r - 1), F_B, VER, DER, "#,##0.00")
    for c in "ACDEFH":
        put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r += 1
    # --- materiales
    banda(ws, r, "MATERIALES", fill=CEL, font=F_B, alto=16); r += 1
    r_m0 = r
    for cod, desc, und, cant, sus in mats:
        put(ws, "A%d" % r, cod, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, und, al=CEN)
        put(ws, "D%d" % r, None)
        put(ws, "E%d" % r, cant, fill=VER if isinstance(cant, str) else AMA, al=DER, fmt="0.0000")
        put(ws, "F%d" % r, "=" + P[cod], fill=VER, al=DER, fmt="#,##0.00")
        put(ws, "G%d" % r, "=ROUND(E%d*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
        put(ws, "H%d" % r, sus, F_S)
        ws.row_dimensions[r].height = 24 if len(desc) > 55 or len(sus) > 50 else 15
        filas_mat.append(r); r += 1
    r_mat = r
    put(ws, "B%d" % r, "Subtotal materiales", F_B, GRIS); put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (r_m0, r - 1), F_B, VER, DER, "#,##0.00")
    for c in "ACDEFH":
        put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r += 1
    # --- equipos
    banda(ws, r, "EQUIPOS Y HERRAMIENTAS", fill=CEL, font=F_B, alto=16); r += 1
    r_e0 = r
    for e in equipos:
        if e[0] == "%MO":
            pct = e[1]
            put(ws, "A%d" % r, "37-0001", al=CEN); put(ws, "B%d" % r, "Herramientas manuales"); put(ws, "C%d" % r, "%MO", al=CEN)
            put(ws, "D%d" % r, None); put(ws, "E%d" % r, pct, fill=AMA, al=DER, fmt="0.00")
            put(ws, "F%d" % r, "=G%d" % r_mo, fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "G%d" % r, "=ROUND(E%d/100*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "H%d" % r, e[2], F_S)
        else:
            cod, desc, und, cant, sus = e
            put(ws, "A%d" % r, cod, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, und, al=CEN)
            put(ws, "D%d" % r, None); put(ws, "E%d" % r, cant, fill=VER if isinstance(cant, str) else AMA, al=DER, fmt="0.0000")
            put(ws, "F%d" % r, "=" + P[cod], fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "G%d" % r, "=ROUND(E%d*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "H%d" % r, sus, F_S)
        ws.row_dimensions[r].height = 24
        filas_eq.append(r); r += 1
    r_eq = r
    put(ws, "B%d" % r, "Subtotal equipos y herramientas", F_B, GRIS); put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (r_e0, r - 1), F_B, VER, DER, "#,##0.00")
    for c in "ACDEFH":
        put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r += 1
    r_tot = r
    ws.merge_cells("A%d:F%d" % (r, r))
    put(ws, "A%d" % r, "COSTO UNITARIO DIRECTO POR glb (S/)", F_H, AZUL, DER)
    for c in "BCDEF":
        ws["%s%d" % (c, r)].border = BOR; ws["%s%d" % (c, r)].fill = AZUL
    put(ws, "G%d" % r, "=G%d+G%d+G%d" % (r_mo, r_mat, r_eq), Font(name="Arial", size=10, bold=True, color="FFFFFF"), AZUL, DER, "#,##0.00")
    put(ws, "H%d" % r, None, fill=AZUL)
    ws.row_dimensions[r].height = 20
    r += 2
    # --- calculo del rendimiento
    banda(ws, r, "CÁLCULO DEL RENDIMIENTO (días-cuadrilla por componente)", fill=CEL, font=F_B, alto=16); r += 1
    for c, t in zip("ABCDEFGH", ("TIPO", "COMPONENTE", "UND", "", "CANTIDAD", "REND. (und/día)", "DÍAS", "CRITERIO")):
        put(ws, "%s%d" % (c, r), t, F_B, GRIS, CEN)
    r += 1
    r_r0 = r
    for tipo, (desc, rnd, und), cant, crit in rend_tabla:
        put(ws, "A%d" % r, tipo, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, und, al=CEN); put(ws, "D%d" % r, None)
        put(ws, "E%d" % r, cant, fill=VER if isinstance(cant, str) else AMA, al=DER, fmt="0.00")
        put(ws, "F%d" % r, rnd, fill=AMA, al=DER, fmt="0.00")
        put(ws, "G%d" % r, "=E%d/F%d" % (r, r), fill=VER, al=DER, fmt="0.000")
        put(ws, "H%d" % r, crit, F_S)
        ws.row_dimensions[r].height = 24 if len(desc) > 55 else 15
        r += 1
    put(ws, "B%d" % r, "Total días-cuadrilla para 1 glb", F_B, GRIS)
    put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (r_r0, r - 1), F_B, VER, DER, "0.000")
    for c in "ACDEFH":
        put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r_dias = r; r += 1
    put(ws, "B%d" % r, "Rendimiento = 1 / total días (glb/día)", F_B, GRIS)
    put(ws, "G%d" % r, "=ROUND(1/G%d,4)" % r_dias, F_B, VER, DER, "0.0000")
    for c in "ACDEFH":
        put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r_rend = r; r += 2
    # cabecera
    ws["B4"] = "=G%d" % r_rend
    ws["D4"] = "=SUM(E%d:E%d)" % (r_mo0, r_mo - 1)
    ws["F4"] = '=SUMIF(C%d:C%d,"hm",E%d:E%d)' % (r_e0, r_eq - 1, r_e0, r_eq - 1)
    ws["H4"] = "=G%d" % r_tot
    # observaciones
    if obs:
        banda(ws, r, "OBSERVACIONES", fill=NAR, font=F_B, alto=16); r += 1
        for o in obs:
            texto_largo(ws, r, "• " + o, alto=max(15, 13 * math.ceil(len(o) / 135.0)), font=F_N); r += 1
    pagina(ws)
    ws.print_title_rows = "7:7"
    return "'%s'!$G$%d" % (nombre, r_tot), "'%s'!$G$%d" % (nombre, r_rend)


def main():
    ws_pl = leer_planilla()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    res = wb.create_sheet("RESUMEN")
    HH = construir_mo(wb)
    P = construir_precios(wb)
    filas_res = []
    CUA = {"1op": ("Operario", 1.0), "0.5pe": ("Peón", 0.5), "1pe": ("Peón", 1.0), "1of": ("Oficial", 1.0)}
    # ---------------- lineas 4.2.3.1.1 a 4.2.3.1.6
    for item, rp, cuad, filas, nota in PARTIDAS:
        titulo = ws_pl.cell(rp, 2).value.strip()
        met = ws_pl.cell(rp, 10).value
        mats, rend_cnt, obs = [], {}, []
        peg = {1: 0.0, 15: 0.0, 2: 0.0}; rosc = 0.0; agr = {}
        for fila, (cod, tipo, cfix, desp, ob) in filas.items():
            if isinstance(fila, int):
                desc = ws_pl.cell(fila, 2).value.strip()
                cpl = cantidad(ws_pl, fila)
                c = cfix if cfix is not None else cpl
                und_pl = (ws_pl.cell(fila, 3).value or "").strip()
                n2 = lambda x: ("%.2f" % x).rstrip("0").rstrip(".")
                sus = "planilla fila %d: %s" % (fila, ("%s %s" % (n2(cpl), und_pl)) if cpl is not None else "sin parcial")
                if cfix is not None and cfix != cpl:
                    sus += " (se toma %s, ver obs.)" % n2(cfix)
            else:
                c = cfix; desc = {"CAN15": 'Canastilla de succión PVC Ø1 1/2"', "MALLA": "Malla galvanizada para codos de ventilación (2 x 0.05 m²)"}[cod]
                sus = "ver observación"
            if cod.startswith("BR"):
                desc = BRIDAS[cod][0]
            if fila == 445:
                desc = 'Tubería PVC SP C-10 Ø1 1/2" del tramo de succión (planilla: «Canastilla de PVC-U Ø1 1/2"» 4.64 m)'
            g = agr.setdefault(cod, {"desc": desc, "c": [], "sus": [], "desp": desp})
            g["c"].append(c); g["sus"].append(sus)
            if ob:
                obs.append(ob)
            if tipo:
                rend_cnt[tipo] = rend_cnt.get(tipo, 0) + c
            pg, rs = UNIONES[cod]
            dia = 1 if cod in ("TB1", "CD1", "UFG1", "VF1") else (
                2 if cod in ("TB2", "CD2", "CD245", "TE2", "SK2", "UU2", "UPR2", "NP2", "RD2", "SOMB2") else 15)
            peg[dia] += pg * c; rosc += rs * c
        for cod, g in agr.items():  # un renglon por insumo (como en S10 / Delphin)
            und = "und" if cod.startswith("BR") else PRECIOS.get(cod, (None, "m"))[1]
            num = lambda x: ("%.2f" % x).rstrip("0").rstrip(".")
            suma = "+".join(num(x) for x in g["c"])
            sus = " + ".join(g["sus"]) if len(g["sus"]) < 3 else "planilla filas: " + ", ".join(x.split(":")[0].replace("planilla fila ", "") for x in g["sus"])
            if cod == "NFG15":
                und = "m"; c_txt = "=(%s)*0.60" % suma; sus += " (niple L = 0.60 m)"
            elif g["desp"] != 1:
                c_txt = "=(%s)*%s" % (suma, g["desp"]); sus += "; + %d %% de desperdicio" % round((g["desp"] - 1) * 100)
            elif len(g["c"]) > 1:
                c_txt = "=" + suma
            else:
                c_txt = g["c"][0]
            mats.append((cod, g["desc"], und, c_txt, sus))
        # consumibles
        gal = round(peg[1] * 0.003 + peg[15] * 0.005 + peg[2] * 0.007, 4)
        if gal > 0:
            mats.append(("PEG", "Pegamento para PVC", "gal", gal,
                         "%.0f uniones pegadas (≈0.003 / 0.005 / 0.007 gal por unión de Ø1\" / 1 1/2\" / 2\")" % sum(peg.values())))
        if rosc > 0:
            mats.append(("TEF", "Cinta teflón 3/4\" x 10 m", "und", round(rosc / 2 * 0.25, 2),
                         "%.0f extremos roscados ≈ %.0f roscas macho; 1 rollo cada 4 roscas" % (rosc, rosc / 2)))
        if rend_cnt.get("abr"):
            mats.append(("TAR", "Perno de anclaje 3/8\" para fijar abrazaderas", "und", 2 * rend_cnt["abr"], "2 por abrazadera (una por oreja)"))
        equipos = [("%MO", 3, "3 % de la mano de obra (tarraja, llaves, sierra, taladro)")]
        nbr = sum(c for (cod, tipo, cfix, desp, ob), c in [(v, (v[2] if v[2] is not None else (cantidad(ws_pl, f) or 0))) for f, v in filas.items()] if cod.startswith("BR"))
        if nbr:
            equipos.append(("SOLM", "Máquina de soldar eléctrica 250 A", "hm", round(nbr * 0.5, 2), "0.5 h por brida rompe agua (%.0f bridas)" % nbr))
        rend_tabla = [(t, REND[t], v, "rendimiento adoptado (referencial, editable)") for t, v in rend_cnt.items()]
        if nota:
            obs.insert(0, nota)
        cuadrilla = [CUA[k] for k in cuad]
        cu, rr = hoja_acu(wb, item, item, titulo, met, None, mats, equipos, rend_tabla, cuadrilla, obs, P, HH)
        filas_res.append((item, titulo, ws_pl.cell(rp, 3).value, met, cu))
    # ---------------- 4.2.3.1.7 prueba hidraulica
    item, rp = "4.2.3.1.7", 512
    titulo = ws_pl.cell(rp, 2).value.strip(); met = ws_pl.cell(rp, 10).value
    VC, VT = 32.0, 15.0
    mats = [("AGUA", "Agua para llenado de cisterna (32 m³) y tanque elevado (15 m³) + 10 % de enjuague", "m3", "=(32+15)*1.10",
             "volúmenes del programa arquitectónico del perfil (cisterna 32 m³, tanque elevado 15 m³)"),
            ("HIPO", "Hipoclorito de calcio al 70 %", "kg", "=ROUND((32+15)*50/1000/0.70,2)",
             "50 mg/l de cloro libre durante 24 h (RNE IS.010); 47 m³ x 50 g/m³ / 0.70"),
            ("KIT", "Comparador de cloro residual y pH (con reactivos)", "und", 1, "lectura de cloro residual al final de la desinfección")]
    rend_tabla = [("prep", ("Preparación, taponeo de salidas y llenado de cisterna y tanque", 2, "und"), 1, "2 estructuras por día"),
                  ("lect", ("Prueba de estanqueidad: marcado de nivel y lectura a las 24 h, revisión de filtraciones", 4, "und"), 1, "control del nivel; 1/4 de día"),
                  ("desi", ("Desinfección: preparación de la solución, aplicación, 24 h de contacto y medición", 2, "und"), 1, "1/2 día"),
                  ("vac", ("Vaciado, enjuague y limpieza final (2 estructuras)", 1.333, "und"), 1, "3/4 de día")]
    obs = ["La planilla no desagrega esta partida; el análisis se arma con los volúmenes del programa arquitectónico (cisterna 32 m³ y "
           "tanque elevado 15 m³). Si el diseño de IISS tiene otros volúmenes, cambiar las fórmulas de agua e hipoclorito.",
           "Se llena una sola vez: la desinfección se hace con el agua de la prueba de estanqueidad, agregando el cloro; el vaciado se "
           "hace con las electrobombas de la partida 4.2.3.2.1 (sin costo de equipo adicional).",
           "El agua se valoriza con la tarifa oficial de EMAPA San Martín (categoría comercial y otros, tramo de 30 m³ a más), que es sin IGV."]
    cu, rr = hoja_acu(wb, item, item, titulo, met, None, mats, [("%MO", 3, "3 % de la mano de obra")], rend_tabla,
                      [("Operario", 1.0), ("Peón", 1.0)], obs, P, HH)
    filas_res.append((item, titulo, ws_pl.cell(rp, 3).value, met, cu))
    # ---------------- 4.3.6.3.3 tecle de palanca para caballete
    item, rp = "4.3.6.3.3", 1021
    titulo = ws_pl.cell(rp, 2).value.strip(); met = ws_pl.cell(rp, 10).value
    L_ipn = cantidad(ws_pl, 1022)
    mats = [("IPN", "Perfil IPN 100 (8.34 kg/m)", "m", "=%s*1.05" % L_ipn, "planilla filas 1022-1026: %.2f m; + 5 %% de desperdicio" % L_ipn),
            ("TECLE", "Tecle de palanca", "und", cantidad(ws_pl, 1027), "planilla fila 1027"),
            ("PLA", "Tapa: plancha LAC 1/8\" 0.40 x 0.40 m (24.91 kg/m²)", "kg", "=ROUND(0.40*0.40*24.91,2)",
             "planilla fila 1028: «Tapa» 1 und; medida adoptada, confirmar con el plano"),
            ("RUEDA", "Rueda giratoria 360° con freno", "und", cantidad(ws_pl, 1029), "planilla fila 1029"),
            ("SOLD", "Soldadura cellocord E6011 1/8\"", "kg", 2.5, "≈2.5 % del peso del caballete (≈100 kg); incluye la soldadura de las placas de las ruedas"),
            ("ANTI", "Pintura anticorrosiva", "gal", 0.30, "≈4.7 m² (0.40 m²/m de IPN 100) a 2 manos, 32 m²/gal"),
            ("ESM", "Esmalte sintético", "gal", 0.30, "≈4.7 m² a 2 manos, 32 m²/gal"),
            ("THIN", "Thinner", "gal", 0.15, "dilución y limpieza"),
            ("LIJA", "Lija para fierro N.° 80", "und", 4, "preparación de superficie")]
    peso = "=ROUND(%s*1.05*8.34,1)" % L_ipn
    rend_tabla = [("fab", ("Habilitado, corte y soldadura del caballete de IPN 100", 60, "kg"), peso, "60 kg/día de estructura liviana"),
                  ("pin", ("Limpieza, lijado y pintura (anticorrosivo + esmalte)", 25, "m2"), 4.7, "25 m²/día"),
                  ("rue", ("Montaje de ruedas con freno", 8, "und"), cantidad(ws_pl, 1029), "8 und/día"),
                  ("tec", ("Montaje y prueba del tecle de palanca y de la tapa", 2, "und"), 1, "1/2 día")]
    obs = ["Caballete móvil de perfil IPN 100 con 4 ruedas giratorias con freno y tecle de palanca, para levantar las tapas de "
           "inspección del tanque séptico (1.00 x 1.00 x 0.20 m ≈ 480 kg): el tecle debe ser de 1.5 t o más.",
           "La planilla pone «Tapa» 1 und sin medida; se analiza como plancha LAC 1/8\" de 0.40 x 0.40 m (cabezal del caballete). Confirmar con el plano."]
    equipos = [("%MO", 5, "5 % de la mano de obra (amoladora, discos de corte, escuadras, prensas)"),
               ("SOLM", "Máquina de soldar eléctrica 250 A", "hm", "=ROUND(%s/60*8,2)" % peso.lstrip("="), "horas de fabricación del caballete")]
    cu, rr = hoja_acu(wb, item, item, titulo, met, None, mats, equipos, rend_tabla, [("Operario", 1.0), ("Peón", 1.0)], obs, P, HH)
    filas_res.append((item, titulo, ws_pl.cell(rp, 3).value, met, cu))
    # ---------------- resumen
    ws = res
    anchos(ws, {"A": 12, "B": 70, "C": 7, "D": 10, "E": 16, "F": 16})
    banda(ws, 1, "ANÁLISIS DE COSTOS UNITARIOS - PARTIDAS GLOBALES (glb) DE INSTALACIONES SANITARIAS", "A", "F", font=F_T, alto=26)
    texto_largo(ws, 2, PROYECTO, "A", "F", alto=30, font=F_B)
    texto_largo(ws, 3, "Lugar: %s.   Fecha: %s.   Metrado: hoja PLANILLA del metrado de IISS (sin cambios)." % (LUGAR, FECHA), "A", "F", alto=16, font=F_N)
    for c, t in zip("ABCDEF", ("ÍTEM", "DESCRIPCIÓN", "UND", "METRADO", "PRECIO UNIT. (S/)", "PARCIAL (S/)")):
        put(ws, "%s5" % c, t, F_H, AZUL, CEN)
    r = 6
    for item, tit, und, met, cu in filas_res:
        put(ws, "A%d" % r, item, al=CEN); put(ws, "B%d" % r, tit); put(ws, "C%d" % r, und, al=CEN)
        put(ws, "D%d" % r, met, al=DER, fmt="0.00")
        c = put(ws, "E%d" % r, "=" + cu, fill=VER, al=DER, fmt="#,##0.00")
        c.hyperlink = "#'%s'!A1" % item
        put(ws, "F%d" % r, "=ROUND(D%d*E%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
        ws.row_dimensions[r].height = 26
        r += 1
    ws.merge_cells("A%d:E%d" % (r, r))
    put(ws, "A%d" % r, "COSTO DIRECTO DE LAS PARTIDAS GLOBALES (S/, sin IGV)", F_H, AZUL, DER)
    for c in "BCDE":
        ws["%s%d" % (c, r)].border = BOR; ws["%s%d" % (c, r)].fill = AZUL
    put(ws, "F%d" % r, "=SUM(F6:F%d)" % (r - 1), Font(name="Arial", size=10, bold=True, color="FFFFFF"), AZUL, DER, "#,##0.00")
    r += 2
    notas = [
        "Mano de obra: costo hora-hombre 2026 de la tabla oficial CAPECO - FTCCP (operario S/ 28.95, oficial S/ 22.68, peón S/ 20.54); ver hoja MANO DE OBRA 2026.",
        "Materiales: precios de venta publicados (con IGV) en tiendas de alcance nacional, pasados a precio sin IGV; cada precio tiene su tienda y enlace en la hoja PRECIOS DE INSUMOS.",
        "Las cantidades de materiales son las del desagregado de la planilla (más 5 % de desperdicio en tuberías). Las filas con unidad o parcial irregular se explican en las observaciones de cada ACU; el metrado no se ha modificado.",
        "Rendimientos: se calculan por componente (días-cuadrilla) en cada ACU; son referenciales y editables (celdas amarillas).",
        "Los precios unitarios son costo directo sin IGV; gastos generales, utilidad e IGV se agregan en el presupuesto.",
    ]
    banda(ws, r, "NOTAS", "A", "F", fill=CEL, font=F_B); r += 1
    for n in notas:
        texto_largo(ws, r, "• " + n, "A", "F", alto=28, font=F_N); r += 1
    pagina(ws)
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    recalcular(SALIDA)
    return SALIDA


def recalcular(f):
    tmp = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", tmp, f],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.move(os.path.join(tmp, os.path.basename(f)), f)
    shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    print(main())
