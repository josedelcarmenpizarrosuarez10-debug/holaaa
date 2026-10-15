"""CAR Mujeres (CUI 2717013): ACU de las partidas de cisterna, tanque elevado y tecle SIN unidad global.

Las 8 partidas que en el metrado de IISS estan en glb (4.2.3.1.1 a 4.2.3.1.7 y 4.3.6.3.3) se desagregan en sub-partidas
valorizables (opcion B):
  - lineas 4.2.3.1.1 a 4.2.3.1.6: la tuberia por metro (m) y cada accesorio, valvula, brida, abrazadera, etc. por unidad (und),
    con la cantidad de su fila en la planilla (una sub-partida por fila; las filas repetidas del mismo insumo se suman);
  - 4.2.3.1.7 prueba hidraulica de cisterna y tanque elevado: por m3 (32 + 15 m3 del programa arquitectonico);
  - 4.3.6.3.3 tecle de palanca para caballete: por und.
El cambio de unidad es solo en este libro de ACU; la planilla de metrados no se toca.
Mano de obra CAPECO - FTCCP 2026 y precios de insumos con fuente: los mismos de acu_globales_m.py / precios_globales_m.py.
Uso: python3 herramientas/mujeres/acu_por_metro_m.py
"""
import os, sys
import openpyxl
from openpyxl.styles import Font

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import acu_globales_m as G  # noqa: E402
from acu_globales_m import (put, banda, texto_largo, anchos, pagina, F_B, F_N, F_S, F_H, F_T, AZUL, CEL, GRIS, AMA, VER, NAR,  # noqa: E402
                            IZQ, CEN, DER, BOR, PROYECTO, LUGAR, FECHA, REND, BRIDAS, PARTIDAS, UNIONES, PRECIOS, cantidad)

SALIDA = os.path.join(G.RAIZ, "entregables_mujeres", "ACU_IISS_CISTERNA_TANQUE_Y_TECLE_POR_METRO_Y_UNIDAD_CAR_MUJERES.xlsx")
GAL = {1: 0.003, 15: 0.005, 2: 0.007}                      # gal de pegamento por union pegada
DIAM2 = ("TB2", "CD2", "CD245", "TE2", "SK2", "UU2", "UPR2", "NP2", "RD2", "SOMB2", "AB2")
DIAM1 = ("TB1", "CD1", "UFG1", "VF1", "AB1")
UND_INS = {"BR01": "und", "BR02": "und", "BR03": "und", "BR04": "und", "BR05": "und"}


def diam(cod):
    return 1 if cod in DIAM1 else (2 if cod in DIAM2 else 15)


def nombre_ins(cod):
    return BRIDAS[cod][0] if cod.startswith("BR") else PRECIOS[cod][0]


def und_ins(cod):
    return UND_INS.get(cod) or PRECIOS[cod][1]


def bloque(ws, r, item, titulo, und, metrado, cuadrilla, rend, rend_txt, mats, equipos, P, HH, rend_tabla=None, obs=()):
    """Escribe un ACU a partir de la fila r; devuelve (fila siguiente, celda del costo unitario)."""
    banda(ws, r, "%s   %s" % (item, titulo), fill=CEL, font=F_B, alto=30); r += 1
    rr = r
    put(ws, "A%d" % r, "Rend.", F_B, GRIS)
    put(ws, "B%d" % r, rend if rend_tabla is None else None, F_B, AMA if rend_tabla is None else VER, DER, '0.0000 "%s/día"' % und)
    put(ws, "C%d" % r, "MO", F_B, GRIS, CEN); put(ws, "D%d" % r, None, F_B, VER, DER, "0.0000")
    put(ws, "E%d" % r, "EQ", F_B, GRIS, CEN); put(ws, "F%d" % r, None, F_B, VER, DER, "0.0000")
    put(ws, "G%d" % r, "Costo unit. por %s (S/)" % und, F_B, GRIS, CEN)
    put(ws, "H%d" % r, None, Font(name="Arial", size=11, bold=True), VER, DER, "#,##0.00")
    ws.row_dimensions[r].height = 26; r += 1
    put(ws, "A%d" % r, "Metrado", F_B, GRIS); put(ws, "B%d" % r, metrado, F_N, al=IZQ, fmt='0.00 "%s"' % und)
    ws.merge_cells("C%d:H%d" % (r, r)); put(ws, "C%d" % r, rend_txt, F_S)
    r += 1
    for c, t in zip("ABCDEFGH", ("CÓDIGO", "DESCRIPCIÓN DEL RECURSO", "UND", "CUADRILLA", "CANTIDAD", "PRECIO S/", "PARCIAL S/", "SUSTENTO")):
        put(ws, "%s%d" % (c, r), t, F_H, AZUL, CEN)
    r += 1
    banda(ws, r, "MANO DE OBRA", fill=GRIS, font=F_B, alto=15); r += 1
    m0 = r
    for cat, cu in cuadrilla:
        put(ws, "A%d" % r, {"Operario": "47-0001", "Oficial": "47-0002", "Peón": "47-0003"}[cat], al=CEN)
        put(ws, "B%d" % r, cat); put(ws, "C%d" % r, "hh", al=CEN); put(ws, "D%d" % r, cu, fill=AMA, al=DER, fmt="0.00")
        put(ws, "E%d" % r, "=ROUND(D%d*8/$B$%d,4)" % (r, rr), fill=VER, al=DER, fmt="0.0000")
        put(ws, "F%d" % r, "=" + HH[cat], fill=VER, al=DER, fmt="0.00")
        put(ws, "G%d" % r, "=ROUND(E%d*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
        put(ws, "H%d" % r, "cuadrilla x 8 h / rendimiento; HH CAPECO - FTCCP 2026", F_S); r += 1
    rmo = r
    put(ws, "B%d" % r, "Subtotal mano de obra", F_B, GRIS); put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (m0, r - 1), F_B, VER, DER, "#,##0.00")
    for c in "ACDEFH": put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r += 1
    banda(ws, r, "MATERIALES", fill=GRIS, font=F_B, alto=15); r += 1
    a0 = r
    for cod, desc, u, cant, sus in mats:
        put(ws, "A%d" % r, cod, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, u, al=CEN); put(ws, "D%d" % r, None)
        put(ws, "E%d" % r, cant, fill=VER if isinstance(cant, str) else AMA, al=DER, fmt="0.0000")
        put(ws, "F%d" % r, "=" + P[cod], fill=VER, al=DER, fmt="#,##0.00")
        put(ws, "G%d" % r, "=ROUND(E%d*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
        put(ws, "H%d" % r, sus, F_S)
        ws.row_dimensions[r].height = 24 if (len(desc) > 55 or len(sus) > 55) else 15
        r += 1
    rma = r
    put(ws, "B%d" % r, "Subtotal materiales", F_B, GRIS); put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (a0, r - 1), F_B, VER, DER, "#,##0.00")
    for c in "ACDEFH": put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r += 1
    banda(ws, r, "EQUIPOS Y HERRAMIENTAS", fill=GRIS, font=F_B, alto=15); r += 1
    e0 = r
    for e in equipos:
        if e[0] == "%MO":
            put(ws, "A%d" % r, "37-0001", al=CEN); put(ws, "B%d" % r, "Herramientas manuales"); put(ws, "C%d" % r, "%MO", al=CEN)
            put(ws, "D%d" % r, None); put(ws, "E%d" % r, e[1], fill=AMA, al=DER, fmt="0.00")
            put(ws, "F%d" % r, "=G%d" % rmo, fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "G%d" % r, "=ROUND(E%d/100*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "H%d" % r, e[2], F_S)
        else:
            cod, desc, u, cant, sus = e
            put(ws, "A%d" % r, cod, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, u, al=CEN); put(ws, "D%d" % r, None)
            put(ws, "E%d" % r, cant, fill=VER if isinstance(cant, str) else AMA, al=DER, fmt="0.0000")
            put(ws, "F%d" % r, "=" + P[cod], fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "G%d" % r, "=ROUND(E%d*F%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "H%d" % r, sus, F_S)
        r += 1
    req = r
    put(ws, "B%d" % r, "Subtotal equipos y herramientas", F_B, GRIS); put(ws, "G%d" % r, "=SUM(G%d:G%d)" % (e0, r - 1), F_B, VER, DER, "#,##0.00")
    for c in "ACDEFH": put(ws, "%s%d" % (c, r), None, fill=GRIS)
    r += 1
    rt = r
    ws.merge_cells("A%d:F%d" % (r, r))
    put(ws, "A%d" % r, "COSTO UNITARIO DIRECTO POR %s (S/)" % und, F_H, AZUL, DER)
    for c in "BCDEF":
        ws["%s%d" % (c, r)].border = BOR; ws["%s%d" % (c, r)].fill = AZUL
    put(ws, "G%d" % r, "=G%d+G%d+G%d" % (rmo, rma, req), Font(name="Arial", size=10, bold=True, color="FFFFFF"), AZUL, DER, "#,##0.00")
    put(ws, "H%d" % r, None, fill=AZUL)
    r += 1
    if rend_tabla:
        r += 1
        for c, t in zip("ABCDEFGH", ("TIPO", "COMPONENTE", "UND", "", "CANTIDAD", "REND. (und/día)", "DÍAS", "CRITERIO")):
            put(ws, "%s%d" % (c, r), t, F_B, GRIS, CEN)
        r += 1; t0 = r
        for tipo, (desc, rnd, u), cant, crit in rend_tabla:
            put(ws, "A%d" % r, tipo, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, u, al=CEN); put(ws, "D%d" % r, None)
            put(ws, "E%d" % r, cant, fill=VER if isinstance(cant, str) else AMA, al=DER, fmt="0.00")
            put(ws, "F%d" % r, rnd, fill=AMA, al=DER, fmt="0.00")
            put(ws, "G%d" % r, "=E%d/F%d" % (r, r), fill=VER, al=DER, fmt="0.000"); put(ws, "H%d" % r, crit, F_S); r += 1
        put(ws, "B%d" % r, "Rendimiento = 1 / total días (%s/día)" % und, F_B, GRIS)
        put(ws, "G%d" % r, "=ROUND(1/SUM(G%d:G%d),4)" % (t0, r - 1), F_B, VER, DER, "0.0000")
        for c in "ACDEFH": put(ws, "%s%d" % (c, r), None, fill=GRIS)
        ws["B%d" % rr] = "=G%d" % r
        r += 1
    ws["D%d" % rr] = "=SUM(E%d:E%d)" % (m0, rmo - 1)
    ws["F%d" % rr] = '=SUMIF(C%d:C%d,"hm",E%d:E%d)' % (e0, req - 1, e0, req - 1)
    ws["H%d" % rr] = "=G%d" % rt
    for o in obs:
        texto_largo(ws, r, "• " + o, alto=max(15, 13 * (len(o) // 130 + 1)), font=F_N, fill=NAR); r += 1
    return r + 2, "'%s'!$G$%d" % (ws.title, rt)


def hoja(wb, nombre):
    ws = wb.create_sheet(nombre)
    anchos(ws, {"A": 8, "B": 50, "C": 7, "D": 10, "E": 11, "F": 11, "G": 12, "H": 42})
    banda(ws, 1, "ANÁLISIS DE COSTOS UNITARIOS", font=F_T, alto=24)
    texto_largo(ws, 2, PROYECTO, alto=30, font=F_S)
    texto_largo(ws, 3, "Lugar: %s.   Fecha: %s.   Precios sin IGV." % (LUGAR, FECHA), alto=15, font=F_S)
    pagina(ws)
    return ws


def main():
    ws_pl = G.leer_planilla()
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    res = wb.create_sheet("RESUMEN")
    HH = G.construir_mo(wb)
    P = G.construir_precios(wb)
    P["NFG15"] = P["TFG15"]
    filas_res = []          # ("grupo", item, titulo, und_original) | (item, desc, und, metrado, cu, sustento)
    CUA = [("Operario", 1.0), ("Peón", 0.5)]
    # ---------------- lineas: tuberia por m y accesorios por und
    for item, rp, cuad, filas, nota in PARTIDAS:
        titulo = ws_pl.cell(rp, 2).value.strip()
        ws = hoja(wb, item)
        r = 5
        banda(ws, r, "%s %s  (en la planilla: 1 %s)" % (item, titulo, ws_pl.cell(rp, 3).value), fill=AZUL, font=F_H, alto=30)
        r += 1
        if nota:
            texto_largo(ws, r, nota, alto=26, font=F_N, fill=NAR); r += 1
        r += 1
        filas_res.append(("grupo", item, titulo, ws_pl.cell(rp, 3).value))
        # agrupa por (codigo, fila especial) para sumar filas repetidas del mismo insumo
        sub = {}
        for fila, (cod, tipo, cfix, desp, ob) in filas.items():
            if cod == "MALLA":
                continue
            clave = "CD2M" if fila == 510 else cod
            if isinstance(fila, int):
                cpl = cantidad(ws_pl, fila)
                c = cfix if cfix is not None else cpl
                desc = ws_pl.cell(fila, 2).value.strip()
                ref = "fila %d" % fila
            else:
                c = cfix; desc = 'Canastilla de succión PVC Ø1 1/2"'; ref = "fila 445 (ver obs.)"
            if fila == 445:
                desc = 'Tubería PVC SP C-10 de Ø=1 1/2" - tramo de succión (planilla: «Canastilla de PVC-U Ø1 1/2"» 4.64 m)'
            clave = (clave, " ".join(desc.lower().split()))   # solo se suman filas con el mismo insumo y la misma descripcion
            s = sub.setdefault(clave, {"cod": cod, "tipo": tipo, "desc": desc, "c": 0.0, "ref": [], "obs": []})
            s["c"] += c; s["ref"].append(ref)
            if ob:
                s["obs"].append(ob)
        # tuberias primero, luego el resto en el orden de la planilla
        orden = sorted(sub.items(), key=lambda kv: (0 if kv[1]["tipo"].startswith("tubo") else 1))
        for k, (clave, s) in enumerate(orden, 1):
            cod, tipo = s["cod"], s["tipo"]
            it = "%s.%02d" % (item, k)
            es_tubo = tipo.startswith("tubo")
            und = "m" if es_tubo else "und"
            pg, rs = UNIONES[cod]
            mats, equipos = [], [("%MO", 3, "3 % de la mano de obra")]
            if es_tubo:
                mats.append((cod, nombre_ins(cod), "m", 1.05, "1 m + 5 % de desperdicio"))
            elif cod == "NFG15":
                mats.append(("NFG15", nombre_ins("TFG15"), "m", 0.60, "niple de L = 0.60 m cortado y roscado del tubo"))
            elif clave[0] == "CD2M":
                mats.append(("CD2", nombre_ins("CD2"), "und", 1, "codo de ventilación"))
                mats.append(("MALLA", nombre_ins("MALLA"), "m2", 0.05, "malla galvanizada soldada en la boca del codo"))
            else:
                mats.append((cod, nombre_ins(cod), und_ins(cod), 1, "1 und"))
            if cod.startswith("AB"):
                mats.append(("TAR", "Perno de anclaje 3/8\" para fijar abrazaderas", "und", 2, "una por oreja"))
            gal = pg * GAL[diam(cod)]
            if gal:
                mats.append(("PEG", "Pegamento para PVC", "gal", round(gal, 4),
                             ("%.1f uniones pegadas por m" % pg if es_tubo else "%d uniones pegadas" % pg) + " x %.3f gal" % GAL[diam(cod)]))
            if rs:
                mats.append(("TEF", "Cinta teflón 3/4\" x 10 m", "und", round(rs / 2 * 0.25, 4), "%d extremos roscados; 1 rollo cada 4 roscas macho" % rs))
            if cod.startswith("BR"):
                equipos.append(("SOLM", "Máquina de soldar eléctrica 250 A", "hm", 0.5, "0.5 h por brida (soldadura de la plancha al niple)"))
            tipo_r = REND[tipo]
            rend_txt = "Rendimiento: %s (%s %s/día, cuadrilla 1 operario + 0.5 peón; referencial, editable)." % (tipo_r[0], tipo_r[1], und)
            titulo_s = s["desc"]
            r, cu = bloque(ws, r, it, titulo_s, und, s["c"], CUA, tipo_r[1], rend_txt, mats, equipos, P, HH, obs=s["obs"])
            filas_res.append((it, titulo_s, und, s["c"], cu, "planilla " + ", ".join(s["ref"])))
    # ---------------- 4.2.3.1.7 prueba hidraulica por m3
    item, rp = "4.2.3.1.7", 512
    titulo = ws_pl.cell(rp, 2).value.strip()
    ws = hoja(wb, item); r = 5
    banda(ws, r, "%s %s  (en la planilla: 1 glb)" % (item, titulo), fill=AZUL, font=F_H, alto=30); r += 2
    filas_res.append(("grupo", item, titulo, "glb"))
    mats = [("AGUA", "Agua (llenado + 10 % de enjuague)", "m3", 1.10, "1 m³ de volumen + 10 % para enjuague"),
            ("HIPO", "Hipoclorito de calcio al 70 %", "kg", "=ROUND(50/1000/0.70,4)", "50 mg/l de cloro libre durante 24 h (RNE IS.010): 50 g/m³ / 0.70"),
            ("KIT", "Comparador de cloro residual y pH", "und", "=ROUND(1/47,4)", "1 kit para los 47 m³")]
    obs = ["Metrado: cisterna 32 m³ + tanque elevado 15 m³ = 47 m³ (programa arquitectónico del perfil). La planilla tiene esta partida como 1 glb sin desagregado.",
           "Se llena una sola vez: la desinfección usa el agua de la prueba; el vaciado se hace con las electrobombas de la partida 4.2.3.2.1.",
           "Agua con la tarifa oficial de EMAPA San Martín (categoría comercial, tramo de 30 m³ a más), sin IGV."]
    r, cu = bloque(ws, r, item, "Prueba hidráulica de estanqueidad y desinfección de cisterna y tanque elevado", "m3", "=32+15",
                   [("Operario", 1.0), ("Peón", 1.0)], 23.5, "Rendimiento: 47 m³ en 2 días de cuadrilla (llenado, lectura a 24 h, desinfección, vaciado y limpieza) = 23.5 m³/día; editable.",
                   mats, [("%MO", 3, "3 % de la mano de obra")], P, HH, obs=obs)
    filas_res.append((item + ".01", "Prueba hidráulica de estanqueidad y desinfección de cisterna (32 m³) y tanque elevado (15 m³)", "m3", 47, cu,
                      "programa arquitectónico: 32 + 15 m³"))
    # ---------------- 4.3.6.3.3 tecle por und
    item, rp = "4.3.6.3.3", 1021
    titulo = ws_pl.cell(rp, 2).value.strip()
    ws = hoja(wb, item); r = 5
    banda(ws, r, "%s %s  (en la planilla: 1 glb)" % (item, titulo), fill=AZUL, font=F_H, alto=30); r += 2
    filas_res.append(("grupo", item, titulo, "glb"))
    L_ipn = cantidad(ws_pl, 1022)
    peso = "ROUND(%s*1.05*8.34,1)" % L_ipn
    mats = [("IPN", "Perfil IPN 100 (8.34 kg/m)", "m", "=%s*1.05" % L_ipn, "planilla filas 1022-1026: %.2f m; + 5 %% de desperdicio" % L_ipn),
            ("TECLE", "Tecle de palanca 1.5 t", "und", cantidad(ws_pl, 1027), "planilla fila 1027"),
            ("PLA", "Tapa: plancha LAC 1/8\" 0.40 x 0.40 m (24.91 kg/m²)", "kg", "=ROUND(0.40*0.40*24.91,2)", "planilla fila 1028: «Tapa» 1 und; medida adoptada, confirmar con el plano"),
            ("RUEDA", "Rueda giratoria 360° con freno", "und", cantidad(ws_pl, 1029), "planilla fila 1029"),
            ("SOLD", "Soldadura cellocord E6011 1/8\"", "kg", 2.5, "≈2.5 % del peso del caballete; incluye las placas de las ruedas"),
            ("ANTI", "Pintura anticorrosiva", "gal", 0.30, "≈4.7 m² a 2 manos, 32 m²/gal"),
            ("ESM", "Esmalte sintético", "gal", 0.30, "≈4.7 m² a 2 manos, 32 m²/gal"),
            ("THIN", "Thinner acrílico", "gal", 0.15, "dilución y limpieza"),
            ("LIJA", "Lija para fierro N.° 80", "und", 4, "preparación de superficie")]
    rend_tabla = [("fab", ("Habilitado, corte y soldadura del caballete de IPN 100", 60, "kg"), "=" + peso, "60 kg/día de estructura liviana"),
                  ("pin", ("Limpieza, lijado y pintura (anticorrosivo + esmalte)", 25, "m2"), 4.7, "25 m²/día"),
                  ("rue", ("Montaje de ruedas con freno", 8, "und"), cantidad(ws_pl, 1029), "8 und/día"),
                  ("tec", ("Montaje y prueba del tecle de palanca y de la tapa", 2, "und"), 1, "1/2 día")]
    obs = ["Caballete móvil de IPN 100 con 4 ruedas con freno y tecle de palanca para levantar las tapas del tanque séptico (≈480 kg): tecle de 1.5 t.",
           "La planilla pone «Tapa» 1 und sin medida; se analiza como plancha LAC 1/8\" de 0.40 x 0.40 m. Confirmar con el plano."]
    equipos = [("%MO", 5, "5 % de la mano de obra (amoladora, discos, escuadras, prensas)"),
               ("SOLM", "Máquina de soldar eléctrica 250 A", "hm", "=ROUND(%s/60*8,2)" % peso, "horas de fabricación del caballete")]
    r, cu = bloque(ws, r, item, "Caballete móvil de IPN 100 con tecle de palanca", "und", 1, [("Operario", 1.0), ("Peón", 1.0)], None,
                   "Rendimiento calculado por componentes (tabla al pie del análisis).", mats, equipos, P, HH, rend_tabla=rend_tabla, obs=obs)
    filas_res.append((item + ".01", "Caballete móvil de IPN 100 con 4 ruedas con freno y tecle de palanca", "und", 1, cu, "planilla filas 1022-1029"))
    # ---------------- resumen
    ws = res
    anchos(ws, {"A": 13, "B": 66, "C": 6, "D": 10, "E": 14, "F": 14, "G": 30})
    banda(ws, 1, "ANÁLISIS DE COSTOS UNITARIOS - CISTERNA, TANQUE ELEVADO Y TECLE (SIN PARTIDAS GLOBALES)", "A", "G", font=F_T, alto=26)
    texto_largo(ws, 2, PROYECTO, "A", "G", alto=30, font=F_B)
    texto_largo(ws, 3, "Las 8 partidas que la planilla de IISS tiene en glb se desagregan para valorizar: tubería por metro, accesorios por unidad, "
                       "prueba de cisterna y tanque por m³ y tecle por unidad. Los metrados son los del desagregado de la planilla (sin cambios). "
                       "Lugar: %s. Fecha: %s." % (LUGAR, FECHA), "A", "G", alto=30, font=F_N)
    for c, t in zip("ABCDEFG", ("ÍTEM", "DESCRIPCIÓN", "UND", "METRADO", "PRECIO UNIT. (S/)", "PARCIAL (S/)", "SUSTENTO DEL METRADO")):
        put(ws, "%s5" % c, t, F_H, AZUL, CEN)
    r = 6; grupos = []; g_ini = None
    for f in filas_res:
        if f[0] == "grupo":
            if g_ini:
                ws["F%d" % g_ini] = "=SUM(F%d:F%d)" % (g_ini + 1, r - 1)
            _, item, tit, u0 = f
            put(ws, "A%d" % r, item, F_B, CEL, CEN); put(ws, "B%d" % r, tit, F_B, CEL); put(ws, "C%d" % r, None, fill=CEL)
            put(ws, "D%d" % r, None, fill=CEL); put(ws, "E%d" % r, "subtotal", F_S, CEL, DER)
            put(ws, "F%d" % r, None, F_B, CEL, DER, "#,##0.00"); put(ws, "G%d" % r, "en la planilla: 1 %s" % u0, F_S, CEL)
            c = ws["A%d" % r]; c.hyperlink = "#'%s'!A1" % item
            g_ini = r; grupos.append(r); ws.row_dimensions[r].height = 26
        else:
            item, desc, u, met, cu, sus = f
            put(ws, "A%d" % r, item, al=CEN); put(ws, "B%d" % r, desc); put(ws, "C%d" % r, u, al=CEN)
            put(ws, "D%d" % r, met, fill=AMA, al=DER, fmt="0.00")
            put(ws, "E%d" % r, "=" + cu, fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "F%d" % r, "=ROUND(D%d*E%d,2)" % (r, r), fill=VER, al=DER, fmt="#,##0.00")
            put(ws, "G%d" % r, sus, F_S)
            ws.row_dimensions[r].height = 24 if len(desc) > 70 else 15
        r += 1
    ws["F%d" % g_ini] = "=SUM(F%d:F%d)" % (g_ini + 1, r - 1)
    ws.merge_cells("A%d:E%d" % (r, r))
    put(ws, "A%d" % r, "COSTO DIRECTO TOTAL (S/, sin IGV)", F_H, AZUL, DER)
    for c in "BCDE":
        ws["%s%d" % (c, r)].border = BOR; ws["%s%d" % (c, r)].fill = AZUL
    put(ws, "F%d" % r, "=" + "+".join("F%d" % g for g in grupos), Font(name="Arial", size=10, bold=True, color="FFFFFF"), AZUL, DER, "#,##0.00")
    put(ws, "G%d" % r, None, fill=AZUL)
    r += 2
    notas = [
        "Valorización: la tubería se paga por metro instalado y cada accesorio, válvula, brida, abrazadera, electronivel o canastilla por unidad instalada; "
        "la prueba de cisterna y tanque por m³ probado y desinfectado; el caballete con tecle por unidad.",
        "Cada sub-partida sale de una fila del desagregado de la planilla (columna SUSTENTO). Filas con unidad o parcial irregular: ver las observaciones de cada ACU.",
        "La partida 4.2.3.1.8 (prueba hidráulica de las líneas, 67 m) no se toca: corresponde a impulsión 15 + distribución 17 + rebose 17 + 18 = 67 m "
        "(no incluye ingreso 6 m ni succión 4.64 m).",
        "Mano de obra: tabla oficial CAPECO - FTCCP 2026 (operario S/ 28.95, oficial S/ 22.68, peón S/ 20.54 por hora). Materiales: hoja PRECIOS DE INSUMOS (tienda y enlace).",
        "Costo directo sin IGV; gastos generales, utilidad e IGV van en el presupuesto.",
    ]
    banda(ws, r, "NOTAS", "A", "G", fill=CEL, font=F_B); r += 1
    for n in notas:
        texto_largo(ws, r, "• " + n, "A", "G", alto=28, font=F_N); r += 1
    ws.freeze_panes = "A6"
    pagina(ws)
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    G.recalcular(SALIDA)
    return SALIDA


if __name__ == "__main__":
    print(main())
