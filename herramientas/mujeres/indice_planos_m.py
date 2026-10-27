# -*- coding: utf-8 -*-
"""Índice de planos de drenaje pluvial - CAR Mujeres (CUI 2717013), en Word.

Fuente: el último juego de planos en PDF del usuario (PLANOS_ANDFY_MUJERES, 46 láminas, setiembre 2026;
orden y códigos en insumos/wilma/LISTA_PLANOS_ANDFY_MUJERES.txt). De cada PDF se lee, sin modificarlo:
formato del papel, título y escala del rótulo, escalas de los dibujos y el contenido que distingue a las
láminas de un mismo título (secciones, tramos, progresivas, bancas, cortes).
El encabezado (logo y nombre del proyecto) es el de los documentos del proyecto
(insumos/wilma/MEMORIA_IISS_CAR_MUJERES.docx, solo lectura), con la especialidad DRENAJE PLUVIAL.

Uso:
  python3 herramientas/mujeres/indice_planos_m.py <carpeta_pdf_planos>   lee los PDF -> entregables_mujeres/_calc/indice_planos.json y Word
  python3 herramientas/mujeres/indice_planos_m.py                        Word desde el JSON
Salida: entregables_mujeres/INDICE_DE_PLANOS_DRENAJE_PLUVIAL_CAR_MUJERES.docx
"""
import copy
import json
import os
import re
import subprocess
import sys
import unicodedata

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PLANTILLA = os.path.join(RAIZ, 'insumos/wilma/MEMORIA_IISS_CAR_MUJERES.docx')
LISTA = os.path.join(RAIZ, 'insumos/wilma/LISTA_PLANOS_ANDFY_MUJERES.txt')
DATOS = os.path.join(RAIZ, 'entregables_mujeres/_calc/indice_planos.json')
SALIDA = os.path.join(RAIZ, 'entregables_mujeres/INDICE_DE_PLANOS_DRENAJE_PLUVIAL_CAR_MUJERES.docx')
GRANATE, GRIS_CLARO, GRIS = '7B1E2B', 'F2F2F2', '595959'

# Título de cada lámina tal como figura en su rótulo (con espacios y tildes). Se comprueba contra el texto del PDF.
TITULOS = {
    'RGDPLV-01': 'RED GENERAL DRENAJE PLUVIAL',
    'ST-DPLV': 'SECCIONES TRANSVERSALES',
    'PL-DPLV-01': 'PERFILES LONGITUDINALES',
    'IM-DPLV': 'DETALLE INSTALACIÓN MONTANTE PLUVIAL - EDIF %s',
    'IM-ING-DPLV': 'DETALLE INSTALACIÓN MONTANTE PLUVIAL - INGRESO %s',
    'PCPLV-01': 'PLANTA DEL COLECTOR PLUVIAL',
    'PLGC-01': 'PERFIL LONGITUDINAL GENERAL DEL COLECTOR',
    'PLDT-CG': 'PERFIL LONGITUDINAL DETALLADO DEL COLECTOR %s',
    'ST-CG': 'SECCIONES TRANSVERSALES DEL COLECTOR-%s',
    'ISM-ST': 'ISOMÉTRICOS DE LAS SECCIONES TRANSVERSALES-%s',
    'STP-CG': 'SECCIONES TÍPICAS DEL COLECTOR',
    'D-RL': 'DETALLE DE REGISTRO DE LIMPIEZA - TAPA CIEGA DE CONCRETO ARMADO 0.68x0.68x0.08 EN REBAJE',
    'IC-01': 'ISOMÉTRICOS CONSTRUCTIVOS',
    'DCTSA-01': 'DETALLES DE CANALETA DE TECHO, SOPORTE Y ACCESORIOS',
    'D-ECC': 'DETALLE DEL EMPALME DE CUNETA AL COLECTOR Y DE LA JUNTA DE DILATACIÓN',
    'ESDAE-01': 'ESTRUCTURA DE SALIDA CON DOS ALEROS Y EMBOQUILLADO',
    'ISM-MPLV': 'ISOMÉTRICO DEL MONTANTE PLUVIAL TÍPICO',
    'DS-01': 'ISOMÉTRICO DE LA CAJA SUMIDERO DE PISO - DETALLE DE CAJA SUMIDERO DE PISO: CORTE, PLANTA Y DESAGREGADO',
    'DFC-MPLV-01': 'DETALLE FALSA COLUMNA - MONTANTE PLUVIAL',
    'DT-RDTJ': 'DETALLE RED DRENAJE TECHO JARDINERAS',
    'IDP-BJ': 'ISOMÉTRICO DEL DRENAJE PLUVIAL DE LA BANCA DE JARDINERÍA',
    'DI-CBCB/1': 'DETALLE ISOMÉTRICO DE CANALETA, BOQUILLA, COLGADOR Y BAJADA DE LA BANCA',
    'DI-CBCB/2': 'CORTES A-A Y B-B DEL DRENAJE PLUVIAL DE LA BANCA DE JARDINERÍA',
}


def _llano(t):
    """Mayúsculas sin tildes ni espacios (para comparar con el texto del PDF)."""
    t = unicodedata.normalize('NFD', t.upper())
    return re.sub(r'[^A-Z0-9.,:]', '', ''.join(c for c in t if unicodedata.category(c) != 'Mn'))


def titulo(cod, n_dicbcb):
    if cod.startswith('ST-DPLV'):
        return TITULOS['ST-DPLV']
    if cod.startswith('IM-ING-DPLV') or cod.startswith('IM-DPLV'):
        return TITULOS[cod.split()[0]] % cod.split()[-1]
    if cod.startswith('PLDT-CG'):
        a, b = cod.split()[1][1:], cod.split()[2]
        return TITULOS['PLDT-CG'] % ('T%s-T%s' % (a, b))
    if cod.startswith('ST-CG') or cod.startswith('ISM-ST'):
        return TITULOS[cod[:-3]] % cod[-2:]
    if cod.startswith('DT-RDTJ'):
        return TITULOS['DT-RDTJ']
    if cod == 'DI-CBCB':
        return TITULOS['DI-CBCB/%d' % n_dicbcb]
    return TITULOS[cod]


def _rango(nums, fmt='%02d'):
    nums = sorted(set(nums))
    return (fmt % nums[0]) + ' A ' + (fmt % nums[-1]) if len(nums) > 1 else fmt % nums[0]


def contenido(cod, t):
    """Lo que distingue a la lámina, leído de sus dibujos."""
    if cod.startswith('ST-DPLV') or cod == 'PL-DPLV-01':
        s = [int(x) for x in re.findall(r'SECCION (\d\d)\b', t)]
        return ('SECCIONES %s' if cod.startswith('ST') else 'CUNETAS, CON LAS SECCIONES %s') % _rango(s)
    if cod.startswith('IM-DPLV') or cod.startswith('IM-ING'):
        c = sorted(set(re.findall(r'CORTE ([A-Z])-\1', t)))
        return ('CORTE %s-%s' % (c[0], c[0]) if len(c) == 1 else 'CORTES ' + ' Y '.join('%s-%s' % (x, x) for x in c)) if c else ''
    if cod == 'PLGC-01' or cod.startswith('PLDT'):
        m = re.findall(r'(PERFIL|TRAMO) (\d+): (\d\+\d{3}\.\d\d) A (\d\+\d{3}\.\d\d)', t)
        n = [int(x[1]) for x in m]
        return '%s %s: PROG. %s A %s' % ({'PERFIL': 'PERFILES', 'TRAMO': 'TRAMOS'}[m[0][0]], _rango(n, '%d'), m[0][2], m[-1][3])
    if cod.startswith('ST-CG') or cod.startswith('ISM-ST'):
        s = [int(x) for x in re.findall(r'(?:SECCION|ISOMETRICO) S-(\d\d)', t)]
        return ('SECCIONES S-%s' if cod.startswith('ST') else 'ISOMÉTRICOS S-%s') % _rango(s).replace(' A ', ' A S-')
    if cod == 'IC-01':
        return '; '.join(sorted(set(re.findall(r'ISOMETRICO \d: [A-Z ]+?(?= ISOMETRICO|\n|$)', t)))).replace('ISOMETRICO', 'ISOMÉTRICO')
    if cod == 'ESDAE-01':
        return 'ELEVACIÓN FRONTAL, CORTES A-A Y B-B'
    if cod.startswith('DT-RDTJ'):
        bc = sorted(set(int(x) for x in re.findall(r'\bBC-(\d\d)\b', t)))
        return ('BC-%02d' % bc[0] if len(bc) == 1 else 'BC-%02d A BC-%02d' % (bc[0], bc[-1])) if bc else ''
    if cod == 'IDP-BJ':
        return 'BANCA DE JARDINERÍA TIPO 01' if 'TIPO 01' in t else ''
    if cod == 'DI-CBCB':
        if 'CORTE A-A (LATERAL)' in t:
            return 'CORTE A-A (LATERAL) Y CORTE B-B (FRONTAL)'
        return 'CORTE DE LA CANALETA E ISOMÉTRICO DEL EXTREMO DE LA CANALETA'
    return ''


def leer_pdf(carpeta):
    planos = []; n_dicbcb = 0
    for ln in open(LISTA, encoding='utf-8').read().splitlines()[1:]:
        if ' | ' not in ln:
            continue
        arch, cod = [x.strip() for x in ln.split(' | ')]
        f = os.path.join(carpeta, arch)
        t = subprocess.run(['pdftotext', '-raw', f, '-'], capture_output=True, text=True, check=True).stdout
        info = subprocess.run(['pdfinfo', f], capture_output=True, text=True, check=True).stdout
        formato = re.search(r'Page size:.*\((A\d)\)', info).group(1)
        lin = t.splitlines()
        i = max(k for k, l in enumerate(lin) if '2717013' in l)
        j = next(k for k, l in enumerate(lin) if k > i and l.startswith('COORDINADOR'))
        rotulo = ' '.join(lin[i + 1:j])
        k = next(k for k, l in enumerate(lin) if k > j and re.search(r'20\d\d$', l.strip()))
        escala = lin[k + 1].strip()
        if cod == 'DI-CBCB':
            n_dicbcb += 1
        tit = titulo(cod, n_dicbcb)
        if not _llano(rotulo).endswith(_llano(tit)):
            raise SystemExit('Título distinto en %s: rótulo "%s" / índice "%s"' % (cod, rotulo[-120:], tit))
        esc_dib = sorted(set(re.sub(r'\s', '', e) for e in re.findall(r'ESC(?:ALA)?\.?\s*:?\s*(1\s*/\s*\d+)', t)),
                         key=lambda e: int(e.split('/')[1]))
        planos.append(dict(n=int(arch.split()[0]), archivo=arch, codigo=cod, titulo=tit, contenido=contenido(cod, t),
                           formato=formato, escala=escala, escalas_dibujo=esc_dib))
    os.makedirs(os.path.dirname(DATOS), exist_ok=True)
    json.dump(planos, open(DATOS, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return planos


# ----------------------------------------------------------------------------------------------- Word
def _sombra(c, color):
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    c._tc.get_or_add_tcPr().append(sh)


def _celda(c, lineas, centro=True):
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i, (txt, tam, neg, color) in enumerate(lineas):
        q = c.paragraphs[0] if i == 0 else c.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
        pf = q.paragraph_format; pf.space_before = Pt(1 if i == 0 else 0); pf.space_after = Pt(1); pf.line_spacing = 1.0
        r = q.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.font.name = 'Arial'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def generar(planos):
    d = docx.Document(PLANTILLA)
    body = d.element.body
    # encabezado de la seccion del cuerpo de la memoria (logo y proyecto), en todas las paginas
    sect = copy.deepcopy(d.sections[1]._sectPr)
    for el in list(body):
        body.remove(el)
    for tp in sect.findall(qn('w:titlePg')):
        sect.remove(tp)
    body.append(sect)
    for s in d.sections:
        for h in (s.header, s.first_page_header):
            for t in h._element.iter(qn('w:t')):
                if t.text == 'INSTALACIONES SANITARAS SANITARIAS':
                    t.text = 'DRENAJE PLUVIAL'
    s = d.sections[0]
    ancho = (s.page_width - s.left_margin - s.right_margin) / 360000.0

    def par(txt, tam, neg=False, color=None, antes=0, despues=2, centro=True, subr=False):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.0
        r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.underline = subr; r.font.name = 'Arial'
        r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        return p

    par('ÍNDICE DE PLANOS', 13, True, despues=2, subr=True)
    par('ESPECIALIDAD: DRENAJE PLUVIAL', 10, True, despues=2)
    par('CUI N° 2717013', 9.5, True, despues=8)
    anchos = (0.7, 2.6, ancho - 0.7 - 2.6 - 1.75 - 1.9, 1.75, 1.9)
    tb = d.add_table(rows=1, cols=5); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(int(w * 567)))
    for c, w, t in zip(tb.rows[0].cells, anchos, ('N°', 'LÁMINA', 'DESCRIPCIÓN', 'FORMATO', 'ESCALA')):
        c.width = Cm(w); _sombra(c, GRANATE); _celda(c, [(t, 7, True, 'FFFFFF')])
    trPr = tb.rows[0]._tr.get_or_add_trPr(); h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); trPr.append(h)
    for p in planos:
        fila = tb.add_row()
        fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        cel = fila.cells
        for c, w in zip(cel, anchos):
            c.width = Cm(w)
        _celda(cel[0], [('%02d' % p['n'], 8, False, None)])
        _celda(cel[1], [(p['codigo'], 7.5, True, GRANATE)])
        lineas = [(p['titulo'], 8, True, None)]
        if p['contenido']:
            lineas.append((p['contenido'], 6.5, False, GRIS))
        _celda(cel[2], lineas, centro=False)
        _celda(cel[3], [(p['formato'], 8, False, None)])
        esc = [(p['escala'], 8, False, None)]
        if p['escalas_dibujo']:
            esc.append((', '.join(p['escalas_dibujo']), 6.5, False, GRIS))
        _celda(cel[4], esc)
    cuenta = {}
    for p in planos:
        cuenta[p['formato']] = cuenta.get(p['formato'], 0) + 1
    resumen = ', '.join('%d en %s' % (cuenta[f], f) for f in sorted(cuenta))
    par('Total: %d láminas (%s).' % (len(planos), resumen), 9, True, antes=8, despues=2, centro=False)
    par('Escala del rótulo: INDICADA; debajo, en gris, las escalas de los dibujos de cada lámina. Fecha de las láminas: setiembre 2026.',
        8, False, GRIS, despues=0, centro=False)
    d.core_properties.author = ''
    d.core_properties.title = 'Índice de planos - Drenaje pluvial CAR Mujeres'
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    if len(sys.argv) > 1:
        planos = leer_pdf(sys.argv[1])
    else:
        planos = json.load(open(DATOS, encoding='utf-8'))
    for p in planos:
        print('%02d %-15s %-3s %-9s %-60s | %s | %s' % (p['n'], p['codigo'], p['formato'], p['escala'], p['titulo'][:60], p['contenido'], ','.join(p['escalas_dibujo'])))
    print(generar(planos))
