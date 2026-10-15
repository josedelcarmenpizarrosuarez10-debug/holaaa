"""Hogar de Refugio Temporal (Mujeres Violentadas, CUI 2675514): laminas DJ-01 a DJ-03 de las cunetas de piso.
DJ-01: detalle de la partida 01.04.03.04.03.03
JUNTA DE DILATACION E=1" CON SELLADOR ASFALTO en cunetas, en toda la seccion en U (2 muros + losa de fondo), segun la
EG-2013 MTC, Seccion 510 (cunetas revestidas de concreto).

Escala real (1 unidad = 1 m) en una lamina A1 a 1/10; las vistas ampliadas o reducidas se dibujan con su factor e indican
su escala. Seccion de cuneta de la hoja CONCRETO EN CUNETAS (interior 0.40, muros y losa de 0.10, losa de 0.60) y metrado
de la hoja METRADO JUNTA DE DILATACION (juntas cada 3.00 m; longitud por junta = 2 x altura del muro + 0.40).
DJ-02 (proceso constructivo de la junta en isometrico) y DJ-03 (rejilla) en planos_junta_cunetas_iso.py.
Uso: python3 herramientas/planos_junta_cunetas.py
"""
import os, sys, warnings
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI); sys.path.insert(0, AQUI)
import openpyxl
import dxf_base as B
import dxf_layouts as LY
from ezdxf.enums import TextEntityAlignment as TA
from dxf_detalles import hatch_conc, tecnopor, barra, punto_barra, D38

PLANILLA = os.path.join(RAIZ, "insumos", "metrados_vigentes", "METRADO_DRENAJE_MUJERES_VIOLENTADAS.xlsx")
SALIDA = os.path.join(RAIZ, "entregables", "PLANOS_DETALLE_CUNETAS_JUNTA_Y_REJILLA_HOGAR_REFUGIO.dxf")
BI, EM, EL, ES = 0.40, 0.10, 0.10, 0.10       # ancho interior, muros, losa de fondo, solado (m)
BE = BI + 2 * EM                               # 0.60
HD = 0.50                                      # altura interior dibujada (variable segun el eje)
EJ, PR = 0.0254, 0.0508                        # junta 1" y profundidad del sellador 2"
SEP = 3.00


def metrado():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = openpyxl.load_workbook(PLANILLA, data_only=True)["METRADO JUNTA DE DILATACION"]
    assert abs(ws["B10"].value - BI) < 1e-9 and abs(ws["B8"].value - SEP) < 1e-9, "la hoja no tiene la junta en toda la seccion"
    filas = []
    for r in range(13, 25):
        filas.append(dict(eje=ws["B%d" % r].value, L=ws["D%d" % r].value, h=ws["F%d" % r].value, n=int(ws["H%d" % r].value),
                          lj=ws["I%d" % r].value, p=ws["J%d" % r].value))
    return filas, ws["J12"].value


class V:
    """Vista con origen en papel (mm) y factor k (k = 1: escala de la lamina)."""
    def __init__(self, lam, xmm, ymm, k):
        self.l, self.O, self.k = lam, lam.P(xmm, ymm), k
    def __call__(self, u, v): return (self.O[0] + u * self.k, self.O[1] + v * self.k)
    def pts(self, ps): return [self(*p) for p in ps]


def dj01(doc, F, total):
    lam = B.Lamina(doc, 0.0, 0.0, 10, "DJ-01", "DETALLE DE JUNTA DE DILATACION E=1\" CON SELLADOR ASFALTICO EN CUNETAS (TODA LA SECCION)",
                   "PARTIDA 01.04.03.04.03.03 - JUNTA EN 2 MUROS + LOSA DE FONDO CADA 3.00 m - EG-2013 MTC, SECCION 510 - ESC. INDICADAS")
    nj = sum(f["n"] for f in F)
    # ------------------------------------------------ A. seccion transversal en la junta (1/5)
    v = V(lam, 95, 345, 2.0)
    zt = EL + HD
    U = [(0, 0), (BE, 0), (BE, zt), (BE - EM, zt), (BE - EM, EL), (EM, EL), (EM, zt), (0, zt)]
    tecnopor(lam, v.pts(U))
    # sellador 1" x 2" en la cara interior (2 muros + fondo): franja de 2" desde la cara interior
    for q in ([(EM - PR, EL - PR), (EM, EL - PR), (EM, zt), (EM - PR, zt)], [(BE - EM, EL - PR), (BE - EM + PR, EL - PR), (BE - EM + PR, zt), (BE - EM, zt)],
              [(EM, EL - PR), (BE - EM, EL - PR), (BE - EM, EL), (EM, EL)]):
        lam.relleno(v.pts(q), "CORTES")
    lam.poli(v.pts(U), "CONCRETO", cerrada=True, ancho=0.004)
    sol = [(-0.05, -ES), (BE + 0.05, -ES), (BE + 0.05, 0), (-0.05, 0)]; lam.poli(v.pts(sol), "SOLADO", cerrada=True); lam.achurado(v.pts(sol), "TERRENO-ACHURADO", escala_mm=0.5)
    for s in (-1, 1):
        x0 = 0 if s < 0 else BE
        lam.poli([v(x0, zt), v(x0 + s * 0.30, zt)], "TERRENO", ancho=0.003)
    lam.cota(v(0, -ES), v(EM, -ES), -6, texto="0.10"); lam.cota(v(EM, -ES), v(BE - EM, -ES), -6, texto="0.40"); lam.cota(v(BE - EM, -ES), v(BE, -ES), -6, texto="0.10")
    lam.cota(v(0, -ES), v(BE, -ES), -13, texto="0.60")
    lam.cota(v(BE, 0), v(BE, EL), 34, horizontal=False, texto="0.10"); lam.cota(v(BE, EL), v(BE, zt), 34, horizontal=False, texto="H (variable)")
    lam.cota(v(BE, -ES), v(BE, 0), 34, horizontal=False, texto="4\"")
    lam.juntar_llamadas()
    for p, t in (((BE - EM / 2, zt - 0.08), ["Plancha de tecnopor 1\" en TODA la seccion: los 2 muros", "y la losa de fondo (la junta separa la cuneta en tramos)"]),
                 ((BE - EM + PR / 2, EL + HD * 0.55), ["Sellador asfaltico 1\" x 2\" (RC-250 + arena) en la", "cara interior de los muros: 2 x H"]),
                 ((BI / 2 + EM, EL - PR / 2), ["Sellador asfaltico 1\" x 2\" en la cara interior del fondo: 0.40"]),
                 ((BE + 0.03, -ES / 2), ["Solado f'c = 100 kg/cm2, e = 4\" (no lleva junta)"])):
        lam.llamada(v(*p), (lam.P(260, 0)[0], v(*p)[1]), t, 2.0)
    lam.volcar_llamadas()
    lam.texto(v(BE / 2, zt + 0.03), "CARA DE LA JUNTA (plancha vista)", 1.9, "TEXTOS", TA.MIDDLE_CENTER)
    lam.titulo_vista(140, 300, "A. SECCION TRANSVERSAL EN LA JUNTA", "ESC. 1/5 - longitud por junta = 2 x H + 0.40 (cara interior sellada)", 190)
    # ------------------------------------------------ B. detalle de la junta en el muro (1/2), corte en planta
    w = V(lam, 560, 470, 5.0)
    izq = [(-0.25, 0), (-EJ / 2, 0), (-EJ / 2, EM), (-0.25, EM)]; der = [(EJ / 2, 0), (0.25, 0), (0.25, EM), (EJ / 2, EM)]
    hatch_conc(lam, w.pts(izq), 0.8); hatch_conc(lam, w.pts(der), 0.8)
    tecnopor(lam, w.pts([(-EJ / 2, 0), (EJ / 2, 0), (EJ / 2, EM - PR), (-EJ / 2, EM - PR)]))
    lam.relleno(w.pts([(-EJ / 2, EM - PR), (EJ / 2, EM - PR), (EJ / 2, EM), (-EJ / 2, EM)]), "CORTES")
    ym = EM / 2
    barra(lam, w(-0.25, ym), w(-EJ / 2 - 0.05, ym), D38, 5.0, "ACERO-LONG"); barra(lam, w(EJ / 2 + 0.05, ym), w(0.25, ym), D38, 5.0, "ACERO-LONG")
    for x in (-0.20, 0.20): punto_barra(lam, w(x, ym), D38, 5.0)
    lam.cota(w(-EJ / 2, 0), w(EJ / 2, 0), -6, texto="1\""); lam.cota(w(EJ / 2, 0), w(EJ / 2 + 0.05, 0), -14, texto="0.05")
    lam.cota(w(0.25, EM - PR), w(0.25, EM), 6, horizontal=False, texto="2\" (50 mm)"); lam.cota(w(0.25, 0), w(0.25, EM - PR), 6, horizontal=False, texto="0.049")
    lam.texto(w(0, EM + 0.03), "CARA INTERIOR (agua)", 1.8, "TEXTOS", TA.MIDDLE_CENTER); lam.texto(w(-0.16, -0.035), "CARA EXTERIOR (terreno)", 1.8, "TEXTOS", TA.MIDDLE_CENTER)
    lam.juntar_llamadas()
    for p, t in (((0, EM - PR / 2), ["Sellador asfaltico 1\" x 2\": asfalto liquido RC-250", "+ arena gruesa, compactado con tacos hasta el ras"]),
                 ((0, (EM - PR) / 2), ["Tecnopor 1\" en todo el espesor, retirado 2\"", "en la cara interior antes de sellar"]),
                 ((-0.10, ym), ["Acero longitudinal 3/8\" interrumpido a 0.05", "de la junta (no la atraviesa)"]),
                 ((0.20, ym), ["Acero transversal en U 3/8\" @0.25 a cada lado"])):
        lam.llamada(w(*p), (lam.P(720, 0)[0], w(*p)[1]), t, 2.0)
    lam.volcar_llamadas()
    lam.titulo_vista(560, 420, "B. DETALLE DE LA JUNTA", "ESC. 1/2 - corte en planta por el muro (igual en la losa de fondo)", 170)
    # ------------------------------------------------ C. planta y D. corte longitudinal de un tramo (1/25)
    k = 0.4; LT = 7.0; xs = [0.5, 0.5 + SEP, 0.5 + 2 * SEP]
    p = V(lam, 40, 235, k)
    for a, b in zip([0.0] + [x + EJ / 2 for x in xs], [x - EJ / 2 for x in xs] + [LT]):
        lam.poli(p.pts([(a, 0), (b, 0), (b, BE), (a, BE)]), "CONCRETO", cerrada=True, ancho=0.003)
        for y0 in (0, BE - EM): lam.relleno(p.pts([(a, y0), (b, y0), (b, y0 + EM), (a, y0 + EM)]), "CONCRETO-ACHURADO")
        lam.linea(p(a, EM), p(b, EM), "CONCRETO"); lam.linea(p(a, BE - EM), p(b, BE - EM), "CONCRETO")
    for x in xs:
        lam.relleno(p.pts([(x - EJ / 2, 0), (x + EJ / 2, 0), (x + EJ / 2, BE), (x - EJ / 2, BE)]), "CORTES")
        lam.texto(p(x, -0.12), "J", 2.0, "TEXTOS", TA.MIDDLE_CENTER)
    lam.cota(p(xs[0], BE), p(xs[1], BE), 6, texto="3.00"); lam.cota(p(xs[1], BE), p(xs[2], BE), 6, texto="3.00")
    lam.cota(p(LT, 0), p(LT, BE), 6, horizontal=False, texto="0.60")
    lam.texto(p(LT / 2, BE / 2), "FONDO DE CUNETA (losa e = 0.10)", 1.8, "TEXTOS", TA.MIDDLE_CENTER)
    lam.llamada(p(xs[2], BE * 0.5), (lam.P(345, 0)[0], p(0, BE + 0.30)[1]), ["Junta J continua:", "muro - fondo - muro"], 2.0)
    lam.titulo_vista(180, 213, "C. PLANTA DE UN TRAMO DE CUNETA", "ESC. 1/25 - juntas cada 3.00 m en toda la seccion", 170)
    c = V(lam, 40, 140, k)
    zt = EL + HD
    for a, b in zip([0.0] + [x + EJ / 2 for x in xs], [x - EJ / 2 for x in xs] + [LT]):
        lam.relleno(c.pts([(a, 0), (b, 0), (b, EL), (a, EL)]), "CONCRETO-ACHURADO"); lam.poli(c.pts([(a, 0), (b, 0), (b, EL), (a, EL)]), "CONCRETO", cerrada=True)
        lam.poli(c.pts([(a, EL), (b, EL), (b, zt), (a, zt)]), "CONCRETO", cerrada=True, ancho=0.002)
        barra(lam, c(a + (0.05 if a > 0 else 0), EL / 2), c(b - (0.05 if b < LT else 0), EL / 2), D38, k, "ACERO-LONG")
    for x in xs:
        lam.relleno(c.pts([(x - EJ / 2, 0), (x + EJ / 2, 0), (x + EJ / 2, zt), (x - EJ / 2, zt)]), "CORTES")
        lam.texto(c(x, zt + 0.12), "J", 2.0, "TEXTOS", TA.MIDDLE_CENTER)
    lam.rect(*c(0, -ES), *c(LT, 0), "SOLADO")
    lam.texto(c(LT / 2, EL + HD / 2), "MURO DE CUNETA (vista interior)", 1.8, "TEXTOS", TA.MIDDLE_CENTER)
    lam.cota(c(LT, 0), c(LT, EL), 6, horizontal=False, texto="0.10"); lam.cota(c(LT, EL), c(LT, zt), 6, horizontal=False, texto="H")
    lam.llamada(c(xs[2], zt * 0.6), (lam.P(345, 0)[0], c(0, zt + 0.30)[1]), ["La junta J corta el muro y la", "losa de fondo; el solado es corrido"], 2.0)
    lam.titulo_vista(180, 118, "D. CORTE LONGITUDINAL POR EL EJE", "ESC. 1/25 - acero longitudinal interrumpido en cada junta", 170)
    # ------------------------------------------------ E. cuadro de metrado (hoja METRADO JUNTA DE DILATACION)
    filas = [[f["eje"], "%.2f" % f["L"], "%.3f" % f["h"], "%d" % f["n"], "%.3f" % f["lj"], "%.2f" % f["p"]] for f in F]
    filas.append(["TOTAL", "%.2f" % sum(f["L"] for f in F), "", "%d" % nj, "", "%.2f ml" % total])
    yb = lam.tabla(440, 345, ["EJE", "LONGITUD (m)", "ALTURA MURO H (m)", "N° JUNTAS (L/3.00)", "LONG. POR JUNTA 2H+0.40", "PARCIAL (ml)"],
                   filas, [28, 30, 34, 34, 42, 30], hmm=1.8, alto_mm=4.6, titulo="E. METRADO DE LA PARTIDA 01.04.03.04.03.03 (hoja METRADO JUNTA DE DILATACION)")
    lam.notas(440, yb - 8, "NOTAS", [
        "1. EG-2013 MTC, Seccion 510 (cunetas revestidas de concreto): la junta corta TODA la seccion en U, los dos muros y la losa de fondo, para que cada tramo de 3.00 m entre juntas trabaje independiente; una junta solo en los muros deja la losa amarrada y se fisura.",
        "2. Juntas cada 3.00 m y en los cambios de direccion; N° de juntas = REDONDEAR.MAS(L / 3.00) por eje, %d juntas en total." % nj,
        "3. Longitud por junta = 2 x H (altura del muro = area del muro / longitud, hoja CONCRETO EN CUNETAS) + 0.40 (ancho interior del fondo). Se mide en metros lineales de junta sellada.",
        "4. Procedimiento: cortar la plancha de tecnopor 1\" con la forma de la U y colocarla al vaciar; interrumpir el acero longitudinal a 0.05; retirar 2\" de tecnopor en la cara interior; calentar el RC-250, mezclarlo con arena gruesa y compactarlo con tacos hasta el ras.",
        "5. Control: junta continua y alineada en los dos muros y en el fondo, sellador compacto, al ras y sin filtraciones.",
    ], hmm=1.8, ancho_mm=200)
    lam.leyenda2(40, 70, [("concreto", "CONCRETO", "concreto f'c = 175 kg/cm2 de la cuneta"), ("rect", "JUNTAS", "tecnopor 1\" (poliestireno expandido)"),
                          ("relleno", "CORTES", "sellador asfaltico 1\" x 2\" (RC-250 + arena)"), ("linea2", "ACERO-LONG", "acero longitudinal 3/8\""),
                          ("rect", "SOLADO", "solado f'c = 100 kg/cm2, e = 4\"")], 1.8, ancho_col=110, filas_col=3)
    return lam


def construir():
    F, total = metrado()
    doc = B.nuevo_documento()
    from planos_junta_cunetas_iso import dj02, dj03
    lams = {"DJ-01": dj01(doc, F, total), "DJ-02": dj02(doc, 10.0, 0.0), "DJ-03": dj03(doc, 20.0, 0.0)}
    msp = doc.modelspace()
    msp.set_redraw_order({e.dxf.handle: -1 for e in msp.query("HATCH") if e.dxf.layer == "CONCRETO-ACHURADO"})
    LY.presentaciones(doc, lams)
    err = len(doc.audit().errors)
    doc.saveas(SALIDA)
    print("DJ-01, DJ-02, DJ-03 (%d errores de auditoria): %d juntas, %.2f ml -> %s" % (err, sum(f["n"] for f in F), total, os.path.relpath(SALIDA, RAIZ)))
    return SALIDA


if __name__ == "__main__":
    construir()
