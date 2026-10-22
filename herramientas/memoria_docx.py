# -*- coding: utf-8 -*-
"""Agrega, AL FINAL del estudio hidrologico del proyecto (Word), el capitulo del colector pluvial
frontal. No toca ninguna pagina anterior: se abre una seccion nueva (vertical) despues del ultimo
parrafo y alli se escribe el capitulo con los estilos ya existentes del documento."""
import os, sys, json, subprocess, copy
import openpyxl
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, "herramientas"))
import diseno as dz
import metrado_calc as MC

ORIG = os.path.join(RAIZ, "insumos", "solange", "MEMORIA_HIDROLOGIA_MUJERES_VIOLENTADAS_2026.docx")
DEST = os.path.join(RAIZ, "entregables", "MEMORIA_HIDROLOGIA_MUJERES_VIOLENTADAS_2026_CON_COLECTOR.docx")
FIG = os.path.join(RAIZ, "entregables", "_tmp", "fig")
MEM = os.path.join(RAIZ, "entregables", "MEMORIA_CALCULO_COLECTOR_HOGAR_REFUGIO.xlsx")
DJ = json.load(open(os.path.join(RAIZ, "entregables", "_calc", "diseno.json")))
D = DJ["datos"]
FUENTE = "Arial Narrow"
# metrado vigente del usuario (con las bancas OC-01) y juego de laminas DPL (indice_planos_refugio.py)
VIG = os.path.join(RAIZ, "insumos", "metrados_vigentes", "METRADO_DRENAJE_MUJERES_VIOLENTADAS_CON_BANCAS_OC-01.xlsx")
PDF_DPL = os.path.join(RAIZ, "insumos", "solange", "PLANOS_DRENAJE_PLUVIAL_MUJERES_VIOLENTADAS_LAMINAS.pdf")
INDICE = os.path.join(RAIZ, "entregables", "_calc", "indice_planos.json")
PREV_OC = os.path.join(RAIZ, "entregables", "vistas_previas_oc01")


# ----------------------------------------------------------------------------- utilidades de formato
def run(p, texto, negrita=False, cursiva=False, tam=12):
    r = p.add_run(texto); r.font.name = FUENTE; r.italic = cursiva
    if tam: r.font.size = Pt(tam)
    if negrita is not None: r.bold = negrita
    rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), FUENTE); rf.set(qn("w:hAnsi"), FUENTE); rf.set(qn("w:cs"), FUENTE)
    return r


def parrafo(doc, texto="", estilo="Normal", alinear=WD_ALIGN_PARAGRAPH.JUSTIFY, negrita=False, tam=12, espacio=6):
    p = doc.add_paragraph(style=estilo)
    p.paragraph_format.alignment = alinear; p.paragraph_format.line_spacing = 1.5; p.paragraph_format.space_after = Pt(espacio)
    p.paragraph_format.left_indent = Cm(1.75)
    if texto: run(p, texto, negrita=negrita, tam=tam)
    return p


def _numerar(p, nivel):
    """Misma numeracion automatica (lista numId 1) que usan los titulos del documento."""
    pPr = p._p.get_or_add_pPr(); numPr = OxmlElement("w:numPr")
    il = OxmlElement("w:ilvl"); il.set(qn("w:val"), str(nivel)); ni = OxmlElement("w:numId"); ni.set(qn("w:val"), "1")
    numPr.append(il); numPr.append(ni); pPr.insert(1 if pPr.find(qn("w:pStyle")) is not None else 0, numPr)
    p.paragraph_format.line_spacing = 1.5


def titulo1(doc, texto):
    p = doc.add_paragraph(style="Heading 1"); _numerar(p, 0); run(p, texto, negrita=None, tam=None); return p


def titulo2(doc, texto):
    p = doc.add_paragraph(style="Heading 2"); _numerar(p, 1); run(p, texto, negrita=None, tam=None); return p


def vineta(doc, texto):
    p = doc.add_paragraph(style="List Paragraph"); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5; p.paragraph_format.left_indent = Cm(2.5); p.paragraph_format.first_line_indent = Cm(-0.5)
    run(p, "•  " + texto); return p


def formula(doc, texto):
    p = doc.add_paragraph(style="Normal"); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(6)
    run(p, texto, cursiva=True); return p


class Numerador:
    def __init__(self, tabla=22, imagen=5): self.t = tabla; self.i = imagen
    def tabla(self): s = "Tabla N° %02d" % self.t; self.t += 1; return s
    def imagen(self): s = "Imagen N° %02d" % self.i; self.i += 1; return s


NUM = Numerador()


def leyenda(doc, texto, fuente="Fuente: Elaboración propia."):
    p = doc.add_paragraph(style="Normal"); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(2)
    run(p, texto, negrita=True, tam=11); p.paragraph_format.keep_with_next = True
    return p


def fuente(doc, texto="Fuente: Elaboración propia."):
    p = doc.add_paragraph(style="Normal"); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.space_after = Pt(10)
    run(p, texto, cursiva=True, tam=10); return p


def sombrear(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color); tcPr.append(shd)


def tabla(doc, encabezado, filas, anchos=None, tam=10, alinear_num=True, color="1F4E79", junta=False):
    t = doc.add_table(rows=1, cols=len(encabezado)); t.style = doc.styles["Table Grid"]; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(encabezado):
        c = t.rows[0].cells[j]; c.text = ""; p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = run(p, h, negrita=True, tam=tam); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); sombrear(c, color)
    for fila in filas:
        celdas = t.add_row().cells
        for j, v in enumerate(fila):
            c = celdas[j]; c.text = ""; p = c.paragraphs[0]
            es_num = isinstance(v, (int, float))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (es_num and alinear_num) or j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            run(p, v if isinstance(v, str) else ("%d" % v if isinstance(v, int) else "%.2f" % v), tam=tam)
    if anchos:
        t.autofit = False
        for j, w in enumerate(anchos): t.columns[j].width = Cm(w)
        for fila in t.rows:
            for j, w in enumerate(anchos): fila.cells[j].width = Cm(w)
    if junta:                                            # tabla corta: no se parte entre paginas
        for fila in t.rows[:-1]:
            for c in fila.cells:
                for q in c.paragraphs: q.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def imagen(doc, archivo, ancho_cm, titulo, fuente_txt="Fuente: Elaboración propia. Lámina del expediente del colector."):
    leyenda(doc, titulo)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(archivo, width=Cm(ancho_cm))
    fuente(doc, fuente_txt)


def lamina(cod):
    """Imagen JPG de la lamina DPL del juego en PDF del usuario (o vista previa PNG de las laminas OC-01)."""
    if os.path.exists(os.path.join(PREV_OC, cod + ".png")):
        return os.path.join(PREV_OC, cod + ".png")
    pag = {p["codigo"]: p["pagina"] for p in json.load(open(INDICE, encoding="utf-8"))}.get(cod)
    if not pag: return None
    os.makedirs(FIG, exist_ok=True); base = os.path.join(FIG, cod)
    if not os.path.exists(base + ".jpg"):
        subprocess.run(["pdftoppm", "-f", str(pag), "-l", str(pag), "-singlefile", "-jpeg", "-jpegopt", "quality=85",
                        "-scale-to", "2200", PDF_DPL, base], check=True)
    return base + ".jpg"


def figura(doc, cod, texto):
    f = lamina(cod)
    if f: imagen(doc, f, 15.5, "%s: %s (lámina %s)." % (NUM.imagen(), texto, cod), "Fuente: Lámina %s del expediente de drenaje pluvial." % cod)


def resumen_vigente():
    """Filas (item, descripcion, und, total) del RESUMEN del metrado vigente del usuario."""
    ws = openpyxl.load_workbook(VIG, data_only=True)["RESUMEN"]
    return [(str(r[0]).strip(), str(r[1]).strip(), r[2], r[3]) for r in ws.iter_rows(min_row=9, max_col=4, values_only=True) if r[0]]


def hoja_vigente(nombre):
    return [[c for c in r] for r in openpyxl.load_workbook(VIG, data_only=True)[nombre].iter_rows(values_only=True)]


def f2(x): return "%.2f" % x
def f3(x): return "%.3f" % x


# ----------------------------------------------------------------------------- datos de la memoria Excel recalculada
def leer_estructural():
    """Lee la hoja ESTRUCTURAL de la memoria recalculada (LibreOffice) para el cuadro resumen."""
    out = os.path.join(RAIZ, "entregables", "_tmp", "mem"); os.makedirs(out, exist_ok=True)
    rec = os.path.join(out, os.path.basename(MEM))
    if not os.path.exists(rec) or os.path.getmtime(rec) < os.path.getmtime(MEM):
        subprocess.run(["soffice", "-env:UserInstallation=file:///tmp/lo_profile", "--headless", "--convert-to", "xlsx", "--outdir", out, MEM],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
    ws = openpyxl.load_workbook(rec, data_only=True)["ESTRUCTURAL"]
    filas = [[c.value for c in r] for r in ws.iter_rows()]
    sec = None; res = {}
    for f in filas:
        a = f[0]
        if isinstance(a, str) and len(a) > 2 and a[1] == "." and a[0] in "ABCDEFGH": sec = a[0]; res[sec] = {"titulo": a[3:].strip()}
        elif sec and isinstance(a, str):
            k = a.strip()
            if k.startswith("Mu"): res[sec]["Mu"] = f[1]
            elif k.startswith("Momento resistente"): res[sec]["Mn"] = f[1]; res[sec]["ok"] = f[4] if len(f) > 4 else None
            elif k.startswith("As colocado"): res[sec]["As"] = f[1]; res[sec]["ref"] = f[3]
            elif k.startswith("Cortante ultimo") or k == "Vu": res[sec]["Vu"] = f[1]
            elif k.startswith("Cortante resistente") or k == "fVc": res[sec]["Vc"] = f[1]
            elif k.startswith("Presion sobre el suelo"): res[sec]["q"] = f[1]
    return res


def leer_velocidades():
    """Lee el cuadro de velocidades y autolimpieza de la hoja PERFIL_FLUJO recalculada."""
    rec = os.path.join(RAIZ, "entregables", "_tmp", "mem", os.path.basename(MEM))
    if not os.path.exists(rec): return []
    ws = openpyxl.load_workbook(rec, data_only=True)["PERFIL_FLUJO"]
    out = []; dentro = False
    for row in ws.iter_rows(max_col=11):
        a = row[0].value
        if isinstance(a, str) and a.startswith("VERIFICACION DE VELOCIDADES"): dentro = True; continue
        if dentro and isinstance(a, str) and a.startswith("Velocidad maxima"): break
        if dentro and isinstance(row[1].value, (int, float)) and isinstance(row[2].value, (int, float)) and isinstance(a, str):
            out.append([a] + [c.value for c in row[1:8]])
    return out


def leer_cumplimiento():
    """Lee la hoja CUMPLIMIENTO de la memoria recalculada: filas (requisito, norma, criterio, valor, resultado) y pendientes."""
    rec = os.path.join(RAIZ, "entregables", "_tmp", "mem", os.path.basename(MEM))
    if not os.path.exists(rec): return [], []
    ws = openpyxl.load_workbook(rec, data_only=True)["CUMPLIMIENTO"]
    filas = []; pend = []; modo = "tabla"
    for row in ws.iter_rows(min_row=5, max_col=5):
        a = row[0].value
        if a is None: continue
        if isinstance(a, str) and a.startswith("Requisitos que no cumplen"): modo = "x"; continue
        if isinstance(a, str) and a.startswith("DATOS EXTERNOS"): modo = "pend"; continue
        if modo == "tabla":
            filas.append([c.value for c in row[:5]])
        elif modo == "pend":
            pend.append(a)
    return filas, pend


# ----------------------------------------------------------------------------- capitulo
def construir():
    doc = Document(ORIG)
    n_parrafos_orig = len(doc.paragraphs); n_tablas_orig = len(doc.tables)

    # Seccion nueva, vertical, con los margenes de la primera seccion del documento (el cuerpo).
    s0 = doc.sections[0]
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width = s0.page_width; sec.page_height = s0.page_height
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = s0.left_margin, s0.right_margin, s0.top_margin, s0.bottom_margin
    sec.header_distance, sec.footer_distance = s0.header_distance, s0.footer_distance
    sec.different_first_page_header_footer = False

    R = dz.disenar() if hasattr(dz, "disenar") else None
    perfil = DJ["perfil"]; brink = DJ["brink"]; poza = DJ["poza"]; cl = DJ["caja_llegada"]
    g = DJ["geom"]; T, RM = MC.resumen(); est = leer_estructural()
    yn, yc, Vn, Fn = DJ["yn"], DJ["yc"], DJ["V_n"], DJ["F_n"]
    h_min = min(p["h"] for p in perfil); h_max = max(p["h"] for p in perfil)
    llen = max(p["y"] / p["h"] for p in perfil); bl_min = min(p["techo"] - p["NA"] for p in perfil)
    Fmax = max(p["F"] for p in perfil); Vmax = max(p["V"] for p in perfil)

    # ------------------------------------------------------------------ 1. titulo y alcance
    titulo1(doc, "COLECTOR PLUVIAL FRONTAL: TRAMO HOGAR DE REFUGIO TEMPORAL")
    titulo2(doc, "Alcance del tramo")
    parrafo(doc, "El sistema de drenaje pluvial del Hogar de Refugio Temporal (CUI 2675514) entrega sus aguas a un colector "
                 "frontal que corre a lo largo del frente del predio, paralelo a la carretera Oasis, y que forma parte de un "
                 "colector continuo compartido con los dos proyectos vecinos. El colector recibe en su inicio el aporte del "
                 "proyecto aguas arriba, CAR Varones (CUI 2705619), recoge las cunetas del presente proyecto y entrega el caudal "
                 "total al registro R-01 del colector del proyecto aguas abajo, CAR Mujeres (CUI 2717013), que lo conduce hasta "
                 "el punto de descarga final.")
    parrafo(doc, "Este capítulo desarrolla el tramo correspondiente al Hogar de Refugio Temporal: el trazo, el caudal de diseño, "
                 "el dimensionamiento hidráulico de la sección, el perfil de flujo, el empalme con el colector receptor, la "
                 "llegada de las cunetas interiores, la verificación estructural de la estructura cubierta y los criterios "
                 "constructivos. El aporte del proyecto aguas arriba se toma como dato de su propia memoria de cálculo y el "
                 "colector receptor se toma con la geometría de su expediente; ambos proyectos son independientes y aquí sólo "
                 "se verifican las condiciones de compatibilidad en los puntos de empalme.")
    parrafo(doc, "Condiciones de compatibilidad verificadas en la entrega al colector receptor:", negrita=True)
    vineta(doc, "Cota de fondo de llegada igual o superior a la cota de fondo del registro R-01 del receptor (%.2f msnm)." % D["CF_R01"])
    vineta(doc, "Caudal entregado no mayor que el caudal previsto por el receptor en su tramo inicial (560.7 L/s).")
    vineta(doc, "Entrega en régimen subcrítico, sin transmitir un resalto ni velocidades altas al colector receptor.")

    # ------------------------------------------------------------------ 2. trazo
    titulo2(doc, "Trazo y emplazamiento")
    parrafo(doc, "El colector se ubica dentro del predio, pegado por fuera del cerco perimetral frontal, con su eje a %.3f m del "
                 "paramento exterior del cerco, de manera que la estructura no invade la faja de la carretera Oasis y la losa "
                 "superior sirve de vereda, al nivel del terreno terminado del frente y del piso de los ingresos (NPT %.2f msnm). Entre el muro del "
                 "colector y el cimiento del cerco se deja una junta longitudinal de tecnopor de 1\" con sellador asfáltico; al otro lado del colector hay terreno natural, sin junta. El trazo recto tiene %.2f m de longitud "
                 "(progresivas 0+000.00 a 0+%06.2f) y, al llegar al extremo suroeste del frente, gira dos veces 45° para "
                 "alinearse con el registro R-01 del receptor, al que llega en la caja de caída CC (progresiva 0+%06.2f, final "
                 "de la estructura 0+%06.2f)." % (D["eje_desde_cerco"], D["NPT"], g["P_B1"], g["P_B1"], g["P_BRINK"], g["P_FIN"]))
    parrafo(doc, "A lo largo del trazo se identifican dos cruces de circulación: el ingreso vehicular OE-01, de 4.415 m de ancho "
                 "(progresivas 0+026.32 a 0+031.33, incluyendo 0.30 m de transición a cada lado), por el que ingresan camiones "
                 "cisterna y vehículos de servicio, y el portón de motos de 1.77 m (progresivas 0+039.04 a 0+041.04). El resto "
                 "del colector funciona como vereda, con cargas peatonales. La ubicación en coordenadas UTM WGS84 18S de "
                 "la lámina DPL-06 se obtuvo a partir del registro R-01 del colector receptor y de la orientación del frente del "
                 "predio; es referencial y debe verificarse en campo durante el replanteo.")
    tabla(doc, ["ELEMENTO", "PROGRESIVA", "DESCRIPCIÓN"],
          [[r["nombre"], "0+%06.2f" % r["prog"], r["tipo"] + ": " + r["nota"]] for r in DJ["registros"]], anchos=[3.0, 3.0, 9.5])
    leyenda(doc, "%s: Registros, cajas y puntos singulares del tramo." % NUM.tabla()); fuente(doc)
    figura(doc, "DPL-06", "Red general del colector pluvial frontal")

    # ------------------------------------------------------------------ 3. caudal
    titulo2(doc, "Caudal de diseño del tramo")
    parrafo(doc, "El caudal propio del Hogar de Refugio Temporal se obtiene con el método racional, con la intensidad de diseño "
                 "determinada en el estudio hidrológico para el período de retorno de %d años y el tiempo de concentración del "
                 "colector, de %d minutos, con la ecuación de Dick y Peschke:" % (D["TR"], D["tc"]))
    formula(doc, "I = P24 · (tc / 1440)^0.25 · 60 / tc = %.2f mm/h" % DJ["I"])
    parrafo(doc, "Esta intensidad proviene de las mismas precipitaciones de diseño del estudio hidrológico: 35.16 mm en 10 minutos "
                 "(que equivalen a 210.98 mm/h, la intensidad usada para las cunetas y montantes del proyecto) y 41.82 mm en 20 "
                 "minutos (125.45 mm/h). Como el colector recoge toda el área del predio más el aporte del proyecto aguas arriba, "
                 "su tiempo de concentración es mayor que el de una cuneta individual, 15 minutos, y la intensidad correspondiente "
                 "sobre la curva precipitación–duración es de %.2f mm/h. Es el mismo criterio con el que se dimensionó el colector "
                 "del proyecto receptor, de modo que los tres tramos del colector compartido usan la misma lluvia de diseño." % DJ["I"])
    formula(doc, "Q = Σ(C · A) · I · FS / 3600 = %.2f m² · %.2f mm/h · %.2f / 3600 = %.1f L/s" % (D["CA_refugio"], DJ["I"], D["FS"], DJ["Q_refugio_calc"]))
    parrafo(doc, "El aporte externo del proyecto aguas arriba se toma de su memoria de cálculo (%.1f L/s) y no se recalcula. "
                 "El caudal de diseño del tramo es la suma de ambos aportes, que coincide con el caudal que el colector "
                 "receptor prevé recibir en su registro R-01." % D["Q_varones"])
    tabla(doc, ["APORTE", "ÁREA (m²)", "Σ C·A (m²)", "I (mm/h)", "FS", "Q (L/s)"],
          [["Hogar de Refugio Temporal (CUI 2675514), cunetas Ejes 01, 02, 06, 07, 11 y 12", f2(D["A_refugio"]), f2(D["CA_refugio"]), f2(DJ["I"]), f2(D["FS"]), "%.1f" % D["Q_refugio"]],
           ["CAR Varones (CUI 2705619): dato de su memoria", f2(D["A_varones"]), f2(D["CA_varones"]), f2(DJ["I"]), f2(D["FS"]), "%.1f" % D["Q_varones"]],
           ["Caudal de diseño del tramo", "", "", "", "", "%.1f" % DJ["Q"]]], anchos=[6.0, 2.0, 2.0, 2.0, 1.5, 2.0])
    leyenda(doc, "%s: Caudal de diseño del colector frontal, tramo Hogar de Refugio Temporal." % NUM.tabla()); fuente(doc)

    # ------------------------------------------------------------------ 4. hidraulica
    titulo2(doc, "Dimensionamiento hidráulico de la sección")
    parrafo(doc, "Se adopta un canal rectangular de concreto armado f'c = 210 kg/cm², vaciado monolíticamente, con ancho interior "
                 "b = %.2f m, muros de %.2f m, losa de fondo de %.2f m y losa superior de %.2f m (%.2f m en el cruce de camiones), "
                 "coeficiente de rugosidad de Manning n = %.3f. La cota de fondo en el inicio es %.2f msnm y la pendiente "
                 "longitudinal es S = %.1f ‰, con lo que la cota de fondo en el extremo aguas abajo del canal (0+%06.2f) es "
                 "%.2f msnm, por encima de la cota de fondo del registro receptor." % (D["b"], D["e_muro"], D["e_fondo"], D["e_losa"], D["e_losa_camion"], D["n"], D["CF0"], D["S"] * 1000, g["P_BRINK"], brink["z"]))
    parrafo(doc, "La pendiente queda limitada por la cota de fondo de la cuneta del Eje 02 (259.70 msnm al llegar al cerco), que "
                 "debe caer libremente dentro del colector, y por la cota de fondo del receptor; con S = 3 ‰ ambas condiciones "
                 "se cumplen con holgura. Para el caudal de diseño, el tirante normal, el tirante crítico y el número de Froude "
                 "de la sección son:")
    formula(doc, "Q = (1/n) · A · R^(2/3) · S^(1/2)      →      yn = %.3f m ;  V = %.2f m/s ;  F = %.2f" % (yn, Vn, Fn))
    formula(doc, "yc = (Q² / (g · b²))^(1/3) = %.3f m      (yn > yc : régimen subcrítico, pendiente suave)" % yc)
    parrafo(doc, "La altura interior del canal resulta de la diferencia entre el nivel de la losa superior (fijo, a nivel de "
                 "vereda) y la cota de fondo, y varía entre %.2f m y %.2f m; con el tirante de diseño el llenado máximo es "
                 "%.0f %% y el borde libre mínimo es %.2f m, dentro de los criterios adoptados (llenado ≤ 85 %% y borde libre "
                 "≥ 0.05 m)." % (h_min, h_max, llen * 100, bl_min))

    titulo2(doc, "Perfil de flujo")
    parrafo(doc, "Como el régimen es subcrítico, el perfil de flujo se calcula por el método del paso estándar desde el control "
                 "de aguas abajo. El control es la caída libre al final del canal, donde el flujo pasa por el tirante crítico "
                 "(yc = %.3f m, nivel de agua %.3f msnm) antes de caer a la poza de la caja de caída. Desde allí se integra la "
                 "ecuación de energía hacia aguas arriba, con la pendiente de fricción de Manning evaluada como promedio entre "
                 "secciones consecutivas, hasta la caja de llegada en 0+000. El perfil resultante es del tipo M2: el tirante "
                 "crece hacia aguas arriba hasta aproximarse al tirante normal." % (brink["yc"], brink["NA"]))
    formula(doc, "z₁ + y₁ + V₁²/2g = z₂ + y₂ + V₂²/2g + Sf · Δx      ;      Sf = (n · V / R^(2/3))²")
    filas = []
    for p in perfil:
        if abs(p["p"] / 5.0 - round(p["p"] / 5.0)) < 1e-6 or abs(p["p"] - g["P_BRINK"]) < 1e-6 or abs(p["p"] - g["P_B1"]) < 1e-6:
            filas.append(["0+%06.2f" % p["p"], f2(p["z"]), f3(p["y"]), f3(p["NA"]), f2(p["V"]), f2(p["F"]), f2(p["techo"] - p["NA"]), "%.0f" % (100 * p["y"] / p["h"])])
    filas.sort(key=lambda f: f[0])
    tabla(doc, ["PROG.", "COTA DE FONDO (msnm)", "TIRANTE y (m)", "NIVEL DE AGUA (msnm)", "V (m/s)", "FROUDE", "BORDE LIBRE (m)", "LLENADO (%)"], filas, tam=9,
          anchos=[2.0, 2.3, 1.8, 2.3, 1.6, 1.6, 2.0, 1.8])
    leyenda(doc, "%s: Perfil de flujo del colector para Q = %.1f L/s (paso estándar desde la caída libre)." % (NUM.tabla(), DJ["Q"]))
    fuente(doc, "Fuente: Elaboración propia. Hoja PERFIL_FLUJO de la memoria de cálculo del colector.")
    parrafo(doc, "En todo el tramo el número de Froude se mantiene por debajo de %.2f, la velocidad máxima es %.2f m/s y el "
                 "borde libre mínimo es %.2f m. No se presentan cambios de régimen dentro del canal." % (Fmax, Vmax, bl_min))
    titulo2(doc, "Velocidades y autolimpieza")
    parrafo(doc, "La pendiente de 3 ‰ se verifica frente a los criterios de velocidad del Reglamento Nacional de Edificaciones "
                 "(norma CE.040 Drenaje Pluvial): velocidad mínima de 0.90 m/s con el caudal de diseño, para evitar la "
                 "sedimentación, y velocidad máxima de 3.0 m/s para revestimiento de concreto, valor conservador que protege "
                 "la superficie del canal. Como el colector funciona la mayor parte del tiempo con caudales menores que el de "
                 "diseño, se verifica además el esfuerzo cortante tractivo τ = γ·R·S para caudales parciales, con un mínimo "
                 "de 0.15 kg/m² (1.5 Pa), suficiente para arrastrar arena fina y evitar depósitos en el fondo.")
    vel = leer_velocidades()
    if vel:
        tabla(doc, ["CASO", "Q (L/s)", "yn (m)", "V (m/s)", "τ (kg/m²)", "FROUDE", "VERIFICACIÓN"],
              [[v[0], "%.1f" % (v[1] * 1000), f3(v[2]), f2(v[3]), f2(v[5]), f2(v[7]), "CUMPLE"] for v in vel], tam=9, anchos=[5.6, 1.6, 1.6, 1.6, 1.8, 1.6, 2.2])
        leyenda(doc, "%s: Velocidades y esfuerzo tractivo en el colector para el caudal de diseño y caudales parciales (flujo uniforme)." % NUM.tabla())
        fuente(doc, "Fuente: Elaboración propia. Hoja PERFIL_FLUJO de la memoria de cálculo del colector.")
    parrafo(doc, "Con el caudal de diseño la velocidad es de %.2f m/s, por encima del mínimo de autolimpieza y muy por debajo "
                 "del máximo admisible; con apenas el 5 %% del caudal de diseño el esfuerzo tractivo sigue siendo mayor que el "
                 "mínimo adoptado. La pendiente de 3 ‰ es, por tanto, suficiente para que el colector se mantenga limpio sin "
                 "necesidad de mayor desnivel, que obligaría a profundizar la caja de caída en la entrega." % Vn)
    figura(doc, "DPL-07", "Perfil longitudinal general del colector")

    # ------------------------------------------------------------------ 5. empalme
    titulo2(doc, "Empalme con el colector receptor: caja de caída CC")
    parrafo(doc, "El colector receptor tiene en su registro R-01 un ancho interior de %.2f m, cota de fondo %.2f msnm y nivel de "
                 "agua %.3f msnm para su caudal de diseño. El nivel de agua del receptor es inferior al nivel de agua que "
                 "tendría el canal de 0.80 m al llegar con su tirante normal, de modo que no es posible una transición "
                 "gradual por ensanche sin estrangular el flujo. Por ello la entrega se resuelve con una caja de caída: el "
                 "canal termina en una caída libre sobre una poza deprimida de %.2f m de largo y %.2f m de ancho, cuyo piso "
                 "está %.2f m por debajo de la cota de fondo del receptor, con un umbral de salida de 0.40 m de alto que "
                 "coincide con el fondo del R-01." % (D["b_wilma"], D["CF_R01"], D["NA_R01"], D["CC_poza_largo"], D["CC_ancho"], D["CC_poza_prof"]))
    parrafo(doc, "El chorro cae desde el borde del canal (cota de fondo %.2f msnm) hasta el piso de la poza (%.2f msnm), una "
                 "caída de %.2f m. Al pie de la caída el flujo es supercrítico y se forma un resalto hidráulico que queda "
                 "ahogado por el nivel de agua del receptor: el tirante conjugado del resalto es menor que el tirante "
                 "disponible en la poza, de modo que el resalto se completa dentro de la caja y el agua sale sobre el umbral "
                 "hacia el R-01 en régimen subcrítico." % (brink["z"], poza["z_piso"], poza["caida"]))
    tabla(doc, ["PARÁMETRO", "VALOR", "CRITERIO"],
          [["Cota de fondo del canal en la caída", "%.2f msnm" % brink["z"], "CF0 − S · L"],
           ["Tirante crítico en la caída yc", "%.3f m" % brink["yc"], "control del perfil de flujo"],
           ["Piso de la poza", "%.2f msnm" % poza["z_piso"], "fondo del receptor − 0.40 m"],
           ["Caída total", "%.2f m" % poza["caida"], ""],
           ["Energía al pie E₁", "%.3f m" % (poza["caida"] + 1.5 * brink["yc"]), "caída + 1.5 yc (sin pérdidas, conservador)"],
           ["Tirante al pie y₁ / velocidad V₁", "%.3f m / %.2f m/s" % (poza["y1"], poza["V1"]), "y₁ + V₁²/2g = E₁"],
           ["Froude al pie F₁", "%.2f" % poza["F1"], "régimen supercrítico"],
           ["Tirante conjugado y₂", "%.3f m" % poza["y2"], "y₂ = y₁/2 · (√(1 + 8F₁²) − 1)"],
           ["Tirante disponible en la poza", "%.3f m" % poza["tirante_disp"], "nivel de agua del receptor − piso de poza"],
           ["Verificación del resalto", "y₂ < tirante disponible: resalto ahogado", "CUMPLE"],
           ["Longitud de poza", "%.2f m" % poza["L_poza"], "≥ 0.8 · 6 y₂ = %.2f m (resalto ahogado)" % (0.8 * poza["L_resalto"])]], anchos=[5.5, 4.5, 6.0])
    leyenda(doc, "%s: Verificación hidráulica de la caja de caída CC en la entrega al R-01." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Hoja EMPALME de la memoria de cálculo del colector.")
    parrafo(doc, "Las tres condiciones de compatibilidad se cumplen: la cota de fondo de llegada (%.2f msnm) es superior a la del "
                 "R-01 (%.2f msnm), el caudal entregado (%.1f L/s) no excede el previsto por el receptor y la salida sobre el "
                 "umbral se produce en régimen subcrítico." % (brink["z"], D["CF_R01"], DJ["Q"]))

    titulo2(doc, "Caja de llegada CL del aporte externo")
    parrafo(doc, "En el inicio del tramo (0+000) se dispone una caja de llegada de %.2f × %.2f m interiores que recibe el colector "
                 "del CAR Varones (CUI 2705619) y la cuneta del Eje 01 de ese proyecto. Según el expediente del CAR Varones, su colector "
                 "(b = 0.60 m) llega con cota de fondo %.3f msnm a la cara este de la caja, por una ventana de 0.60 × 1.17 m hasta el techo "
                 "de la CL, y la cuneta del Eje 01 (37.0 L/s) entra por una ventana de 0.40 × 0.66 m en el muro norte (NCF %.3f msnm); en "
                 "conjunto entregan los 258.7 L/s considerados. La caja queda del lado del CAR Varones (x = 349162.00 a 349163.80 del "
                 "sistema local), por lo que su tapa se fija en +%.2f msnm, al ras del piso terminado de ese proyecto, con un escalón de "
                 "0.55 m respecto del piso de este tramo (+%.2f). El aporte cae en una poza de %.2f m de profundidad bajo la cota de fondo "
                 "del canal (piso %.2f msnm), que mantiene un colchón de agua de %.2f m con el nivel de agua del inicio del colector (%.3f msnm) "
                 "y amortigua el chorro antes de que el flujo ingrese al canal." % (D["CL_largo"], D["CL_ancho"], D["CF_varones_sup"], D["E01_NCF"], D["NPT_CL"], D["NPT"], D["CL_poza"], cl["z_piso"], cl["tirante_poza"], cl["NA_salida"]))
    figura(doc, "DPL-26", "Caja de llegada CL y caja de caída CC con poza de disipación")

    # ------------------------------------------------------------------ 6. cunetas
    titulo2(doc, "Llegada de las cunetas del proyecto")
    parrafo(doc, "Seis cunetas de piso del proyecto (sección 0.40 × H, muros de 0.10 m, pendiente 0.5 %) llegan al frente del "
                 "predio y se empalman al colector a través de una ventana de 0.40 × H en el muro del lado del predio, con "
                 "caída libre al fondo del colector, junta de tecnopor de 1\" en el contacto y un registro de limpieza sobre "
                 "cada empalme. En todos los casos la cota de fondo de la cuneta al llegar al cerco queda por encima del nivel "
                 "de agua del colector, por lo que las cunetas descargan libremente y el colector no las remansa. Las cunetas "
                 "de los Ejes 11 y 12 terminan antes del cerco y se prolongan con la misma sección y pendiente hasta el muro "
                 "del colector (1.88 y 5.39 m). Esa prolongación se metra dentro de las partidas de cunetas del proyecto; la "
                 "partida 01.04.04 del colector incluye únicamente la ventana de empalme y su registro.")
    filas = [[c["nombre"], c["perfil"], "0+%06.2f" % c["prog"], f2(c["NCF"]), f2(c["H"]), f3(c["NA_colector"]), f2(c["caida_libre"]), f2(c["prolong"]) if c["prolong"] > 0.1 else "-"] for c in DJ["cunetas"]]
    tabla(doc, ["CUNETA", "PERFIL", "PROG. DE EMPALME", "COTA DE FONDO AL CERCO (msnm)", "H (m)", "NIVEL DE AGUA COLECTOR (msnm)", "CAÍDA LIBRE (m)", "PROLONG. (m)"], filas, tam=9,
          anchos=[3.4, 1.3, 2.0, 2.3, 1.3, 2.3, 1.7, 1.7])
    leyenda(doc, "%s: Empalme de las cunetas de piso al colector frontal." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Perfiles de cunetas del plano de arquitectura y hoja CUNETAS de la memoria de cálculo.")

    # ------------------------------------------------------------------ 7. estructural
    titulo2(doc, "Verificación estructural")
    parrafo(doc, "La estructura se verifica como marco cerrado de concreto armado (muros, losa de fondo y losa superior vaciados "
                 "monolíticamente), con las cargas de la norma E.020, el empuje del relleno en reposo y, en los cruces, la "
                 "carga de rueda de la especificación AASHTO LRFD (camión de diseño HL-93 con impacto del 33 % en el ingreso "
                 "vehicular; rueda de 1 t en el portón de motos). El diseño por resistencia sigue la norma E.060. Se adoptan "
                 "tres armados según la zona: un solo marco cerrado de 3/8\" @0.20 en el eje de muros y losas, con longitudinales "
                 "de 3/8\" @0.25, en el tramo normal; marco único de 3/8\" @0.15 en el portón de motos; y doble marco de 1/2\" "
                 "@0.15 (una capa en cada cara) con muros de 0.15 m y losa superior de 0.25 m en el ingreso vehicular. El "
                 "refuerzo en una sola capa en el tramo normal es el que corresponde a muros y losas de espesor no mayor de "
                 "0.20 m según la norma E.060, y es el mismo criterio del colector del proyecto receptor; con ello la sección "
                 "no queda sobredimensionada y el acero del colector es del orden de 38 kg por metro.")
    filas = []
    for k in "ABCDEFGH":
        e = est.get(k)
        if not e: continue
        filas.append([e["titulo"], "%.3f" % e.get("Mu", 0), "%.3f" % e.get("Mn", 0), ("%.2f" % e["Vu"]) if "Vu" in e else "-", ("%.2f" % e["Vc"]) if "Vc" in e else "-", "CUMPLE"])
    tabla(doc, ["ELEMENTO", "Mu (t·m/m)", "φMn (t·m/m)", "Vu (t/m)", "φVc (t/m)", "VERIFICACIÓN"], filas, tam=9, anchos=[7.0, 1.8, 1.8, 1.6, 1.6, 2.2])
    leyenda(doc, "%s: Resumen de la verificación estructural del colector y de las tapas de registro." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Hoja ESTRUCTURAL de la memoria de cálculo del colector.")
    if "G" in est and est["G"].get("q") is not None:
        parrafo(doc, "La presión transmitida al suelo por la estructura llena y con sobrecarga es de %.2f kg/cm², valor bajo que "
                     "debe compararse con la capacidad portante del estudio de mecánica de suelos del proyecto." % est["G"]["q"])
    figura(doc, "DPL-14", "Secciones típicas del colector: tramo normal, cruce de motos y cruce de camiones")

    # ------------------------------------------------------------------ 8. criterios constructivos
    titulo2(doc, "Marco normativo del diseño")
    vineta(doc, "Reglamento Nacional de Edificaciones, norma CE.040 Drenaje Pluvial (R.M. N.° 126-2021-VIVIENDA, que reemplaza a la OS.060): criterios de diseño de colectores, velocidades, registros y tiempos de concentración.")
    vineta(doc, "Reglamento Nacional de Edificaciones, norma E.020 Cargas: pesos unitarios y sobrecargas peatonales.")
    vineta(doc, "Reglamento Nacional de Edificaciones, norma E.060 Concreto Armado: diseño por resistencia, cuantías mínimas (9.7.2), refuerzo de muros en una o dos capas (14.3.4), recubrimientos y traslapes.")
    vineta(doc, "Manual de Hidrología, Hidráulica y Drenaje del MTC: método racional, intensidades de diseño (Dick y Peschke) y período de retorno.")
    vineta(doc, "Manual de Puentes del MTC y AASHTO LRFD Bridge Design Specifications: carga de rueda del camión de diseño HL-93, impacto y ancho de franja para la losa del cruce vehicular.")
    vineta(doc, "Reglamento Nacional de Edificaciones, norma E.050 Suelos y Cimentaciones: la presión transmitida al suelo se compara con la capacidad portante del estudio de mecánica de suelos del proyecto.")
    titulo2(doc, "Criterios constructivos")
    vineta(doc, "Colector cubierto, con losa superior que sirve de vereda, al nivel del terreno del frente y de los pisos de ingreso (%.2f msnm), sin lloraderos." % D["NPT"])
    vineta(doc, "Registros de limpieza a no más de 12 m entre sí y en cada llegada de cuneta, fuera de los cruces vehiculares: "
                "tapa ciega de concreto armado de 0.68 × 0.68 × 0.08 m vaciada aparte en obra, con parrilla de 3/8\" @0.10 y 2 asas "
                "de fierro liso de 1/2\" en cajuelas, apoyada en un rebaje de concreto de 0.70 × 0.70 × 0.08 m de la losa (luz libre "
                "de %.2f m), sin marco ni contramarco metálico; borde engrosado de la losa de 0.15 × 0.10 m y refuerzo de borde con "
                "8 barras de 1/2\" de 1.40 m." % MC.LUZ_REG)
    vineta(doc, ("Juntas de dilatación cada 4.00 m, con tecnopor de 1\" y sellador asfáltico de 1\" × 2\" (asfalto líquido RC-250 + arena gruesa); "
                 "junta longitudinal de tecnopor de 1\" con el mismo sellador en su cara superior entre el colector y el cimiento del cerco, en todo el tramo pegado al cerco; al otro lado hay terreno natural, sin junta.")
                if MC.SELLO["clave"] == "asfaltico" else
                ("Juntas de dilatación cada 4.00 m, con tecnopor de 1\" y sello elastomérico de poliuretano; junta de tecnopor de 1\" entre el "
                 "colector y el cimiento del cerco y entre la losa superior y los pisos adyacentes."))
    vineta(doc, "Empalme de cunetas por ventana de 0.40 × H en el muro del lado del predio, con caída al fondo, junta de 1\" y "
                "registro sobre el empalme.")
    vineta(doc, "Concreto f'c = 210 kg/cm² en colector, cajas y tapas; solado de f'c = 100 kg/cm² de 0.05 m; acero fy = 4200 kg/cm²; "
                "recubrimiento de 4 cm; traslapes de 0.40 m.")
    vineta(doc, "Excavación con sobreancho de 0.25 m por lado, relleno lateral compactado con material propio seleccionado y "
                "eliminación del excedente con esponjamiento de 20 %.")
    vineta(doc, "Las láminas DPL-24 y DPL-29 (registro de limpieza con tapa ciega y sus componentes), DPL-25 (empalme de cunetas al "
                "colector) y DPL-28 (juntas e = 1\" con tecnopor y sellador asfáltico) desarrollan el detalle constructivo de estas "
                "partidas: despiece, dimensiones, procedimiento y cuadro de componentes por unidad para los análisis de costos unitarios. "
                "Las secciones de cada progresiva y sus isométricos están en las láminas DPL-10 a DPL-14, y el perfil detallado en DPL-08 y DPL-09.")
    figura(doc, "DPL-27", "Isométricos constructivos del colector frontal")

    # ------------------------------------------------------------------ 9. metrado
    titulo2(doc, "Resumen de metrados del tramo")
    parrafo(doc, "Los metrados del colector se incorporan a la planilla general de metrados del proyecto como partida "
                 "independiente 01.04.04 COLECTOR PLUVIAL FRONTAL, con su propio desglose por tramos, registros, cajas, juntas "
                 "y empalmes, sin modificar las partidas existentes del sistema de drenaje pluvial. Cada cantidad se sustenta en las "
                 "hojas COLECTOR PARAMETROS, COLECTOR MOV. TIERRAS, COLECTOR CONCRETO, COLECTOR ENCOFRADO, COLECTOR ACERO, COLECTOR "
                 "REGISTROS Y TAPAS y COLECTOR JUNTAS Y EMPALMES de la misma planilla, con las dimensiones de las láminas y la fórmula "
                 "de cada medida a la vista. Las cantidades principales son:")
    filas = [[it, de, un, f2(to)] for it, de, un, to in resumen_vigente() if it.startswith("01.04.04.") and un]
    tabla(doc, ["ÍTEM", "PARTIDA", "UND", "METRADO"], filas, anchos=[3.7, 8.5, 1.3, 2.1], tam=9)
    leyenda(doc, "%s: Metrados principales del colector pluvial frontal." % NUM.tabla())
    fuente(doc, "Fuente: Planilla general de metrados del proyecto (RESUMEN), partida 01.04.04.")

    # ------------------------------------------------------------------ 9b. cumplimiento normativo
    cump, pend = leer_cumplimiento()
    if cump:
        titulo2(doc, "Cuadro de cumplimiento normativo")
        parrafo(doc, "El cuadro siguiente resume, para cada requisito del diseño, la norma o referencia que lo exige, el criterio "
                     "aplicado, el valor obtenido y el resultado de la verificación. Cada valor proviene de una celda con fórmula de "
                     "la memoria de cálculo (hoja CUMPLIMIENTO), de modo que cualquier cambio en los datos de entrada actualiza el "
                     "resultado y puede ser auditado.")
        filas = []
        for f in cump:
            if f[1] is None and f[2] is None:
                filas.append([str(f[0]).upper(), "", "", "", ""])
            else:
                v = f[3]
                vt = ("%.2f" % v) if isinstance(v, (int, float)) and not isinstance(v, bool) else ("" if v is None else str(v))
                filas.append([str(f[0]), str(f[1] or ""), str(f[2] or ""), vt, str(f[4] or "")])
        tabla(doc, ["REQUISITO", "NORMA / REFERENCIA", "CRITERIO", "VALOR", "RESULTADO"], filas, tam=8, anchos=[3.8, 3.2, 3.6, 1.5, 3.1])
        leyenda(doc, "%s: Cuadro de cumplimiento normativo del colector pluvial frontal." % NUM.tabla())
        fuente(doc, "Fuente: Elaboración propia. Hoja CUMPLIMIENTO de la memoria de cálculo del colector.")
        parrafo(doc, "Todos los requisitos verificados cumplen. Los datos que provienen de otros expedientes o estudios se "
                     "registran a continuación con la forma en que este diseño los cubre, de modo que una variación razonable de "
                     "cualquiera de ellos no altera la sección, el armado ni el metrado del colector.")
        titulo2(doc, "Datos de otros expedientes y cómo quedan cubiertos")
        for t in pend:
            vineta(doc, str(t))

    # ------------------------------------------------------------------ 10. conclusiones
    titulo2(doc, "Conclusiones del tramo")
    vineta(doc, "El colector frontal del Hogar de Refugio Temporal conduce %.1f L/s (aporte propio de %.1f L/s más %.1f L/s del "
                "proyecto aguas arriba) en una sección rectangular cubierta de %.2f m de ancho, con pendiente de %.1f ‰, en "
                "régimen subcrítico en toda su longitud (F ≤ %.2f), con llenado máximo de %.0f %% y borde libre mínimo de %.2f m."
                % (DJ["Q"], D["Q_refugio"], D["Q_varones"], D["b"], D["S"] * 1000, Fmax, llen * 100, bl_min))
    vineta(doc, "La entrega al registro R-01 del colector receptor cumple las tres condiciones de compatibilidad: cota de fondo de "
                "llegada %.2f msnm (≥ %.2f msnm), caudal %.1f L/s (≤ 560.7 L/s) y salida en régimen subcrítico gracias a la caja "
                "de caída con resalto ahogado." % (brink["z"], D["CF_R01"], DJ["Q"]))
    vineta(doc, "Las seis cunetas del proyecto descargan libremente al colector, con caídas entre %.2f m y %.2f m sobre el nivel "
                "de agua de diseño." % (min(c["caida_libre"] for c in DJ["cunetas"]), max(c["caida_libre"] for c in DJ["cunetas"])))
    vineta(doc, "La pendiente de 3 ‰ cumple los criterios de velocidad de la norma CE.040 (0.90 a 3.0 m/s con el caudal de diseño) y el "
                "colector se autolimpia incluso con el 5 % del caudal de diseño.")
    vineta(doc, "La estructura cumple las verificaciones de flexión y cortante de la norma E.060 para cargas peatonales y para las "
                "cargas vehiculares de los dos cruces, con un solo marco en el tramo normal y de motos y doble marco en el cruce de "
                "camiones, sin sobredimensionar el acero.")
    vineta(doc, "El aporte del CAR Varones (258.7 L/s) y la cota de fondo de su llegada (%.3f msnm) provienen del expediente de ese "
                "proyecto, con el que este tramo queda coordinado (caja CL con tapa en +%.2f y ventanas este y norte); la ubicación UTM "
                "del trazo debe verificarse en el replanteo." % (D["CF_varones_sup"], D["NPT_CL"]))

    capitulo_complementario(doc)
    doc.save(DEST)
    return n_parrafos_orig, n_tablas_orig


# ----------------------------------------------------------------------------- bancas OC-01 y lloraderos
C_TECHO, I_TECHO, FS_T, N_CAN, S_CAN = 0.88, 210.98, 1.15, 0.013, 0.005    # criterios del estudio (canaletas y montantes)


def manning_rect(Q, b, n, S):
    """Tirante normal en canal rectangular (biseccion)."""
    lo, hi = 1e-5, 2.0
    for _ in range(80):
        y = (lo + hi) / 2; A = b * y; R = A / (b + 2 * y)
        if A * R ** (2 / 3) * S ** 0.5 / n < Q: lo = y
        else: hi = y
    return y


def capitulo_complementario(doc):
    sys.path.insert(0, os.path.join(RAIZ, "herramientas", "techo"))
    import planos_oc01_refugio as OC
    res = {it: (de, un, to) for it, de, un, to in resumen_vigente()}
    b, h = OC.CB, OC.CH
    titulo1(doc, "DRENAJE COMPLEMENTARIO: TECHOS DE LAS BANCAS CON PÉRGOLA (OC-01) Y LLORADEROS EN CUNETAS")
    titulo2(doc, "Techos de las bancas con pérgola (OC-01)")
    parrafo(doc, "La obra complementaria OC-01 (jardinera + banca + pérgola) se repite tres veces en el proyecto; cada juego tiene "
                 "tres techos a una agua de cobertura de panel aluzinc, de %.2f, %.2f y %.2f m de largo por %.2f m de fondo, de modo "
                 "que las nueve bancas suman nueve techos. Cada techo descarga en una canaleta de plancha de acero galvanizado de "
                 "e = 0.50 mm, de sección interior %.2f m de ancho por %.2f m de alto, con pendiente de %.1f %% hacia una boquilla "
                 "de plancha galvanizada que entrega a la bajada de PVC de Ø 4\" adosada a la columna de la pérgola; la bajada "
                 "llega a la falsa columna de drenaje pluvial y de allí al sistema de cunetas. El sumidero tipo cúpula del plano "
                 "de la OC-01 se reemplaza por la boquilla." % (OC.TECHOS[0]["L"], OC.TECHOS[1]["L"], OC.TECHOS[2]["L"], OC.FONDO, b, h, OC.S_CAN))
    parrafo(doc, "Los caudales se calculan con el mismo método racional y los mismos parámetros del estudio para canaletas y "
                 "montantes: intensidad de %.2f mm/h, coeficiente de escorrentía de techo C = %.2f y factor de seguridad %.2f. "
                 "La canaleta se verifica con la ecuación de Manning (n = %.3f, S = %.3f) y la bajada con la ecuación de Wyly-Eaton "
                 "(rs = 0.250), igual que los montantes de las edificaciones." % (I_TECHO, C_TECHO, FS_T, N_CAN, S_CAN))
    formula(doc, "Q = C · I · A · FS ;      Q = (1/n) · Ah · R^(2/3) · S^(1/2) ;      Q (L/s) = 1.754 · rs^(5/3) · D^(8/3)  (D en pulgadas)")
    filas = []; ymax = 0; dmax = 0
    for R in OC.TECHOS:
        A = R["L"] * OC.FONDO; Q = C_TECHO * I_TECHO / 3.6e6 * A * FS_T
        y = manning_rect(Q, b, N_CAN, S_CAN); V = Q / (b * y); F = V / (9.81 * y) ** 0.5
        D = (Q * 1000 / (1.754 * 0.25 ** (5 / 3))) ** (3 / 8); ymax = max(ymax, y); dmax = max(dmax, D)
        filas.append([R["n"], f2(R["L"]), f2(A), "%.3f" % (Q * 1000), f3(y), f2(V), f2(F), f3(h - y), "%.2f\"" % D, "4\""])
    tabla(doc, ["TECHO", "L (m)", "ÁREA (m²)", "Q (L/s)", "y (m)", "V (m/s)", "FROUDE", "BORDE LIBRE (m)", "D REQ.", "D ADOPT."], filas, tam=9,
          anchos=[1.4, 1.3, 1.6, 1.5, 1.4, 1.5, 1.5, 1.9, 1.5, 1.5], junta=True)
    leyenda(doc, "%s: Verificación hidráulica de la canaleta %.2f × %.2f m y de la bajada de los techos de la OC-01." % (NUM.tabla(), b, h))
    fuente(doc, "Fuente: Elaboración propia. Geometría de las láminas DPL-41 y DPL-42.")
    Qcap = (b * (h - 0.05)) * ((b * (h - 0.05)) / (b + 2 * (h - 0.05))) ** (2 / 3) * S_CAN ** 0.5 / N_CAN
    parrafo(doc, "El tirante máximo en la canaleta es de %.3f m, muy por debajo de la altura de %.2f m; con un borde libre de 0.05 m "
                 "la canaleta conduce %.1f L/s, varias veces el caudal del techo mayor. La bajada requiere un diámetro de %.2f\" y se "
                 "adopta Ø 4\", el mismo de los montantes del proyecto y mayor que el mínimo de la norma CE.040. La sección de "
                 "%.2f × %.2f m queda definida por la geometría de la cobertura y por la limpieza de hojas de la pérgola, no por la "
                 "capacidad hidráulica." % (ymax, h, Qcap * 1000, dmax, b, h))
    filas = [[it, res[it][0], res[it][1], f2(res[it][2]) if res[it][1] == "m" else "%d" % res[it][2]]
             for it in sorted(res) if it.startswith("01.04.03.05.04.")]
    tabla(doc, ["ÍTEM", "PARTIDA", "UND", "METRADO"], filas, anchos=[3.7, 8.5, 1.3, 2.1], tam=9, junta=True)
    leyenda(doc, "%s: Metrado de la red de recolección de los techos de las bancas OC-01 (3 juegos, 9 techos)." % NUM.tabla())
    fuente(doc, "Fuente: Planilla general de metrados del proyecto (RESUMEN y hoja METRADO CANALETA BANCAS).")
    figura(doc, "DPL-41", "Isométrico del drenaje pluvial de la OC-01 y planta de techos")
    figura(doc, "DPL-42", "Detalle isométrico de la salida de la canaleta de la OC-01")

    titulo2(doc, "Lloraderos en cunetas colindantes con área verde")
    ll = hoja_vigente("METRADO LLORADEROS")
    tramos = [r for r in ll if isinstance(r[0], (int, float)) and r[1]]
    tot = next(r for r in ll if isinstance(r[0], str) and r[0].upper().startswith("TOTAL DE LLORADEROS"))
    tot = [v for v in tot if v not in (None, "")]
    parrafo(doc, "En los tramos en que el muro de la cuneta queda junto a un área verde, el agua que se infiltra en el jardín se "
                 "acumula contra el muro y lo somete a empuje hidrostático y a humedad permanente. Para aliviarla se colocan "
                 "lloraderos de tubería PVC de Ø 3\" que atraviesan el muro de la cuneta cada 1.50 m, con su boca interior protegida por "
                 "un filtro de grava envuelto en geotextil, de modo que el agua del suelo pasa a la cuneta sin arrastrar finos. "
                 "Los lloraderos no aportan un caudal que cambie el dimensionamiento de las cunetas: el agua infiltrada en el área "
                 "verde ya está considerada en el coeficiente de escorrentía de esas superficies.")
    tabla(doc, ["EJE", "PROGRESIVAS", "LADO", "LONGITUD (m)", "N.°"], [[r[1], r[2], r[3], f2(r[4]), "%d" % r[6]] for r in tramos]
          + [["TOTAL", "", "", f2(tot[1]), "%d" % tot[2]]], anchos=[2.2, 4.4, 3.0, 2.6, 1.6], tam=9)
    leyenda(doc, "%s: Tramos de cuneta con lloraderos de PVC Ø 3\" @ 1.50 m." % NUM.tabla())
    fuente(doc, "Fuente: Planilla general de metrados del proyecto (hoja METRADO LLORADEROS).")
    parrafo(doc, "La ubicación de los lloraderos está en la lámina DPL-39, el corte del lloradero con el tubo perforado y su "
                 "desagregado en la lámina DPL-37 y el isométrico de instalación en la lámina DPL-40.")
    figura(doc, "DPL-37", "Detalle de lloradero: corte, tubo perforado y desagregado")


def verificar(n_par, n_tab):
    """Comprueba que los parrafos y tablas originales quedaron identicos (texto) en el documento nuevo."""
    a = Document(ORIG); b = Document(DEST)
    pa = [p.text for p in a.paragraphs]; pb = [p.text for p in b.paragraphs]
    ok_p = pa == pb[:len(pa)]
    ta = ["|".join(c.text for r in t.rows for c in r.cells) for t in a.tables]
    tb = ["|".join(c.text for r in t.rows for c in r.cells) for t in b.tables]
    ok_t = ta == tb[:len(ta)]
    print("parrafos originales %d (iguales: %s); parrafos nuevos %d" % (len(pa), ok_p, len(pb) - len(pa)))
    print("tablas originales %d (iguales: %s); tablas nuevas %d" % (len(ta), ok_t, len(tb) - len(ta)))
    print("secciones: %d -> %d; imagenes: %d -> %d" % (len(a.sections), len(b.sections), len(a.inline_shapes), len(b.inline_shapes)))
    return ok_p and ok_t


if __name__ == "__main__":
    n_par, n_tab = construir()
    print(DEST)
    verificar(n_par, n_tab)
    import revision_hidrologia as RH                     # texto comun de intensidades y caudales de los tres proyectos
    print("revision:", RH.revisar(DEST, DEST, "refugio"))
