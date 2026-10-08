"""ACU de la partida 01.04.03.04.03.04 LLORADERO DE TUBERIA PVC 3" (CAR Varones y Hogar de Refugio), en un libro aparte.

Mismo formato que la hoja ACU CANALETAS de la planilla del Hogar de Refugio: celda amarilla = dato (cuadrilla, rendimiento,
aporte o precio), verde = formula. Precios sin IGV, referenciales (base Mercado Tarapoto 2026 de la planilla; PVC 3" y
geotextil son precios de mercado a verificar). El metrado sale de la hoja METRADO LLORADEROS de cada planilla.
Uso: python3 herramientas/acu_lloraderos.py
"""
import os, shutil, subprocess, tempfile, warnings
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MET = os.path.join(RAIZ, "insumos", "metrados_vigentes")
PROY = {
    "varones": dict(planilla="METRADO_DRENAJE_PLUVIAL_CAR_VARONES.xlsx", salida=os.path.join(RAIZ, "entregables_varones", "ACU_LLORADEROS_CAR_VARONES.xlsx")),
    "refugio": dict(planilla="METRADO_DRENAJE_MUJERES_VIOLENTADAS.xlsx", salida=os.path.join(RAIZ, "entregables", "ACU_LLORADEROS_HOGAR_REFUGIO.xlsx")),
}
COD = "01.04.03.04.03.04"
FUENTE_MO = "CAPECO: costo hora-hombre 2024-2025 llevado al jornal basico 2025-2026"
FUENTE_MAT = "Mercado Tarapoto 2026, sin IGV (referencial)"

AZUL, CELESTE, AMARILLO, VERDE = "1F4E79", "D9EAF7", "FFF2CC", "F0F7E8"
fino = Side(style="thin", color="7F7F7F"); BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)


def leer_planilla(fn):
    """Datos del proyecto (RESUMEN B4:B7) y metrado de la partida, con valores calculados."""
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(fn, data_only=True)
    rs = wb["RESUMEN"]
    cab = [rs["B%d" % r].value for r in (4, 5, 6, 7)]
    desc = metrado = None
    for r in rs.iter_rows(values_only=True):
        if r[0] == COD: desc, metrado = r[1], r[3]
    if metrado is None:                       # planilla sin valores en cache: recalcular una copia
        tmp = tempfile.mkdtemp()
        try:
            shutil.copy(fn, os.path.join(tmp, "m.xlsx"))
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "o"), os.path.join(tmp, "m.xlsx")],
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
            return leer_planilla(os.path.join(tmp, "o", "m.xlsx"))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)
    return cab, desc, metrado


def construir(clave):
    P = PROY[clave]
    (entidad, proyecto, ubic, fecha), desc, metrado = leer_planilla(os.path.join(MET, P["planilla"]))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "ACU LLORADEROS"; MERGES = []
    for col, w in zip("ABCDEFGH", (8, 52, 9, 13, 13, 13, 14, 58)): ws.column_dimensions[col].width = w

    def celda(ref, valor, fill=None, b=False, sz=11, color=None, h="center", fmt=None, wrap=False):
        c = ws[ref]; c.value = valor; c.font = Font(name="Arial Narrow", size=sz, bold=b, color=color)
        c.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap); c.border = BORDE
        if fill: c.fill = PatternFill("solid", fgColor=fill)
        if fmt: c.number_format = fmt
        return c

    def fila_merge(r, texto, fill=CELESTE, b=True, sz=11, color=None, h="left", alto=None, c1="A", c2="H"):
        MERGES.append("%s%d:%s%d" % (c1, r, c2, r))
        for col in "ABCDEFGH"[ "ABCDEFGH".index(c1): "ABCDEFGH".index(c2) + 1]: celda("%s%d" % (col, r), None, fill)
        celda("%s%d" % (c1, r), texto, fill, b, sz, color, h, wrap=True)
        if alto: ws.row_dimensions[r].height = alto

    fila_merge(1, "ANALISIS DE COSTOS UNITARIOS - LLORADEROS DE PVC Ø3\" EN CUNETAS COLINDANTES CON AREA VERDE", AZUL, True, 14, "FFFFFF", "center", 26)
    for i, (et, v) in enumerate((("ENTIDAD", entidad), ("PROYECTO", proyecto), ("UBICACIÓN", ubic), ("FECHA", fecha))):
        r = 2 + i; MERGES.append("A%d:B%d" % (r, r)); MERGES.append("C%d:H%d" % (r, r))
        celda("A%d" % r, et, b=True); celda("B%d" % r, None)
        celda("C%d" % r, v, h="left", wrap=True, fmt="mmm-yy" if et == "FECHA" else None)
        for col in "DEFGH": celda("%s%d" % (col, r), None)
    ws.row_dimensions[3].height = 48
    fila_merge(7, "Analisis de la partida %s %s. Las cantidades salen del detalle de instalacion de lloraderos en cunetas del plano de drenaje "
                  "(PVC Ø3\" @1.50 m, filtro localizado de 0.30 x 0.30 x 0.30 m de grava envuelta en geotextil no tejido) y de la hoja METRADO LLORADEROS. "
                  "Celda amarilla = dato (cuadrilla, rendimiento, aporte o precio); verde = formula. Precios sin IGV, referenciales: al pasar a Delphin rigen los de la base de precios del presupuesto." % (COD, desc),
               b=True, alto=66)

    # ---------------------------------------------------------------- datos de calculo
    fila_merge(9, "DATOS DE CALCULO (del detalle del plano y de la especificacion tecnica)")
    for col, t in zip("ABCDH", ("N°", "DATO", "UND", "VALOR", "SUSTENTO")): celda("%s10" % col, t, AZUL, True, color="FFFFFF")
    for col in "EFG": celda("%s10" % col, None, AZUL)
    datos = [  # (dato, und, valor o formula, es_formula, sustento)
        ("Espaciamiento de lloraderos", "m", 1.50, False, "Plano: especificacion tecnica de lloraderos (L = 1.50 m)"),
        ("Longitud de tubo PVC por lloradero", "m", 0.30, False, "Muro de 0.10 m + 0.20 m dentro del filtro"),
        ("Lado del filtro localizado (cubo)", "m", 0.30, False, "Plano: filtro localizado 0.30 x 0.30 x 0.30"),
        ("Volumen de grava por lloradero", "m3", "=D13^3", True, "lado^3"),
        ("Area de geotextil por lloradero", "m2", "=6*D13^2*1.1", True, "6 caras del cubo + 10 % de traslapes"),
        ("Volumen de excavacion local por lloradero", "m3", "=D13^3", True, "hueco del filtro (incluido en la mano de obra)"),
        ("Mortero 1:3 de sellado alrededor del tubo", "m3", 0.0005, False, "anillo de 0.02 m alrededor del tubo en las dos caras del muro"),
        ("Cemento por m3 de mortero 1:3", "bol/m3", 10.5, False, "dosificacion de mortero 1:3"),
        ("Arena por m3 de mortero 1:3", "m3/m3", 1.05, False, "dosificacion de mortero 1:3"),
        ("Desperdicio de grava (compactacion y perdidas)", "-", 1.10, False, "criterio de obra"),
    ]
    for i, (t, u, v, f, s) in enumerate(datos):
        r = 11 + i
        celda("A%d" % r, i + 1); celda("B%d" % r, t, h="left"); celda("C%d" % r, u)
        celda("D%d" % r, v, VERDE if f else AMARILLO, fmt="0.0000"); celda("H%d" % r, s, h="left", wrap=True)
        for col in "EFG": celda("%s%d" % (col, r), None)
    # D11 esp, D12 L tubo, D13 lado, D14 grava, D15 geotextil, D16 excav, D17 mortero, D18 cem, D19 arena, D20 desperdicio

    # ---------------------------------------------------------------- ACU
    r0 = 23
    fila_merge(r0, "ACU: %s %s" % (COD, desc), alto=32)
    for k, (et, v, nota) in enumerate((("Unidad", "und", ""), ("Metrado", metrado, "hoja METRADO LLORADEROS de la planilla (RESUMEN %s)" % COD),
                                       ("Rendimiento (und/dia)", 24, "colocacion del tubo en el encofrado, destapado, excavacion del hueco, geotextil, grava y cierre; cuadrilla 0.5 op + 1 peon"))):
        r = r0 + 1 + k; MERGES.append("A%d:B%d" % (r, r)); celda("A%d" % r, et, CELESTE, True, h="left"); celda("B%d" % r, None, CELESTE)
        celda("C%d" % r, v, AMARILLO if et == "Rendimiento (und/dia)" else None, fmt="0.00" if et != "Unidad" else None)
        MERGES.append("D%d:H%d" % (r, r)); celda("D%d" % r, nota, h="left", wrap=True)
        for col in "EFGH": celda("%s%d" % (col, r), None)
    RREND = "$C$%d" % (r0 + 3); RMET = "C%d" % (r0 + 2)
    r = r0 + 4
    for col, t in zip("ABCDEFGH", ("N°", "RECURSO", "UND", "CUADRILLA", "CANTIDAD", "PRECIO S/", "PARCIAL S/", "SUSTENTO / FUENTE")):
        celda("%s%d" % (col, r), t, AZUL, True, color="FFFFFF", wrap=True)
    ws.row_dimensions[r].height = 30
    n = [0]; sub = {}

    def grupo(titulo, items, clave_):
        nonlocal r
        r += 1; MERGES.append("B%d:H%d" % (r, r)); celda("A%d" % r, None, CELESTE); celda("B%d" % r, titulo, CELESTE, True, h="left")
        for col in "CDEFGH": celda("%s%d" % (col, r), None, CELESTE)
        ini = r + 1
        for nombre, und, cuad, cant, precio, fuente in items:
            r += 1; n[0] += 1
            celda("A%d" % r, n[0]); celda("B%d" % r, nombre, h="left", wrap=True); celda("C%d" % r, und)
            celda("D%d" % r, cuad, AMARILLO if cuad is not None else None, fmt="0.00")
            cant_f = ("=D%d*8/%s" % (r, RREND)) if cuad is not None else cant
            celda("E%d" % r, cant_f, VERDE if isinstance(cant_f, str) else AMARILLO, fmt="0.0000")
            celda("F%d" % r, precio, VERDE if isinstance(precio, str) else AMARILLO, fmt="0.00")
            celda("G%d" % r, "=E%d*F%d" % (r, r), VERDE, fmt="0.00"); celda("H%d" % r, fuente, h="left", wrap=True)
        r += 1; celda("A%d" % r, None, CELESTE); celda("B%d" % r, "Subtotal " + titulo.lower(), CELESTE, True, h="left")
        for col in "CDEF": celda("%s%d" % (col, r), None, CELESTE)
        celda("G%d" % r, "=SUM(G%d:G%d)" % (ini, r - 1), CELESTE, True, fmt="0.00"); celda("H%d" % r, None, CELESTE)
        sub[clave_] = "G%d" % r

    grupo("MANO DE OBRA", [("Operario", "hh", 0.5, None, 30.42, FUENTE_MO), ("Peon", "hh", 1.0, None, 20.79, FUENTE_MO)], "mo")
    grupo("MATERIALES", [
        ("Tuberia PVC-U para desague 3\" (NTP 399.003)", "m", None, "=$D$12*1.05", 6.10, "L de tubo x 1.05 desperdicio; precio de mercado del tubo de 3 m (referencial, verificar)"),
        ("Grava filtrante (ripio) 20-40 mm", "m3", None, "=$D$14*$D$20", 101.69, "volumen del filtro x desperdicio; precio de la piedra chancada de la base (%s)" % FUENTE_MAT),
        ("Geotextil no tejido clase 2 (filtro)", "m2", None, "=$D$15*1.05", 4.80, "area con traslapes x 1.05 desperdicio; precio de mercado (referencial, verificar)"),
        ("Cemento Portland tipo I", "bol", None, "=$D$17*$D$18", 30.51, "mortero de sellado; " + FUENTE_MAT),
        ("Arena gruesa", "m3", None, "=$D$17*$D$19", 50.85, "mortero de sellado; " + FUENTE_MAT),
        ("Agua", "m3", None, "=$D$17*0.25", 5.00, "mortero de sellado; " + FUENTE_MAT),
    ], "mat")
    r += 1; MERGES.append("B%d:H%d" % (r, r)); celda("A%d" % r, None, CELESTE); celda("B%d" % r, "EQUIPOS", CELESTE, True, h="left")
    for col in "CDEFGH": celda("%s%d" % (col, r), None, CELESTE)
    r += 1; n[0] += 1
    celda("A%d" % r, n[0]); celda("B%d" % r, "Herramientas manuales", h="left"); celda("C%d" % r, "%mo"); celda("D%d" % r, None)
    celda("E%d" % r, 0.03, AMARILLO, fmt="0.00"); celda("F%d" % r, "=+%s" % sub["mo"], VERDE, fmt="0.00"); celda("G%d" % r, "=E%d*F%d" % (r, r), VERDE, fmt="0.00")
    celda("H%d" % r, "3 % del costo de mano de obra", h="left")
    r += 1; celda("A%d" % r, None, CELESTE); celda("B%d" % r, "Subtotal equipos", CELESTE, True, h="left")
    for col in "CDEF": celda("%s%d" % (col, r), None, CELESTE)
    celda("G%d" % r, "=G%d" % (r - 1), CELESTE, True, fmt="0.00"); celda("H%d" % r, None, CELESTE); sub["eq"] = "G%d" % r
    r += 1; MERGES.append("A%d:F%d" % (r, r)); celda("A%d" % r, "COSTO UNITARIO DIRECTO (S/ por und)", CELESTE, True, 12, h="left")
    for col in "BCDEF": celda("%s%d" % (col, r), None, CELESTE)
    celda("G%d" % r, "=+%s+%s+%s" % (sub["mo"], sub["mat"], sub["eq"]), CELESTE, True, 12, fmt="0.00"); celda("H%d" % r, None, CELESTE); rcu = r
    r += 1; MERGES.append("A%d:F%d" % (r, r)); celda("A%d" % r, "COSTO DIRECTO DE LA PARTIDA (metrado x costo unitario)", CELESTE, True, 12, h="left")
    for col in "BCDEF": celda("%s%d" % (col, r), None, CELESTE)
    celda("G%d" % r, "=G%d*%s" % (rcu, RMET), CELESTE, True, 12, fmt="#,##0.00"); celda("H%d" % r, None, CELESTE)
    r += 2
    fila_merge(r, "Notas: 1) La excavacion del hueco del filtro y el relleno a su alrededor estan dentro de la mano de obra de esta partida; no se suman a las partidas de "
                  "movimiento de tierras. 2) Los precios de tuberia PVC 3\" y de geotextil no estan en la base de la planilla: son referenciales de mercado y deben "
                  "reemplazarse por los de la base de precios del presupuesto.", None, False, alto=48)
    for m in MERGES: ws.merge_cells(m)
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_area = "A1:H%d" % r; ws.freeze_panes = None
    # guardar y recalcular con LibreOffice para que el archivo lleve los valores calculados
    tmp = tempfile.mkdtemp()
    try:
        crudo = os.path.join(tmp, os.path.basename(P["salida"])); wb.save(crudo)
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", os.path.join(tmp, "o"), crudo],
                       check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240)
        shutil.copy(os.path.join(tmp, "o", os.path.basename(P["salida"])), P["salida"])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        v = openpyxl.load_workbook(P["salida"], data_only=True)["ACU LLORADEROS"]
    print("%s: metrado %s und, costo unitario S/ %.2f, costo directo S/ %.2f -> %s" % (clave, metrado, v["G%d" % rcu].value, v["G%d" % (rcu + 1)].value, os.path.relpath(P["salida"], RAIZ)))


if __name__ == "__main__":
    for k in PROY: construir(k)
