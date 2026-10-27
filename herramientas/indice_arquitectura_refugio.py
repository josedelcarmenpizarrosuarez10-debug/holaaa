# -*- coding: utf-8 -*-
"""Índice de planos de ARQUITECTURA - Hogar de Refugio Temporal (CUI 2675514), en Word, con el mismo estilo de los índices de
drenaje pluvial (indice_planos_refugio.py): encabezado y pie, proyecto con CUI, título, especialidad, tabla con encabezado
repetido y filas de grupo.

Fuente: el índice en Excel de la usuaria, versión vigente con el plano topográfico PT-01
(insumos/solange/INDICE_DE_LAMINAS_ARQUITECTURA_HOGAR_REFUGIO_v2.xlsx, hoja ARQUITECTURA; la anterior queda sin el PT-01).
Se toma cada fila tal cual (grupo, descripción, lámina, formato, cantidad), con estos arreglos de forma:
  - grupos numerados en orden (I a VI) y láminas numeradas correlativas (el Excel salta el 51);
  - códigos sin espacios sueltos ("A-02 /01-03" -> "A-02/01-03", "CG- 01" -> "CG-01");
  - tildes en las descripciones (EDIFICACIÓN, ZONIFICACIÓN, CIELORRASO, ...);
  - lámina repetida por copia en el Excel: se toma la hoja de su descripción (EDIFICACIÓN 06 / 03 - 04 -> A-06/03-04).
Las notas sueltas de la columna S del Excel (cartas y anexos de obra) no son láminas y no se incluyen.

Uso: python3 herramientas/indice_arquitectura_refugio.py
Salida: entregables/INDICE_DE_PLANOS_ARQUITECTURA_HOGAR_REFUGIO.docx
"""
import os
import re
import sys

import docx
import openpyxl
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import indice_planos_refugio as I  # noqa: E402  (mismas piezas de formato)

XLSX = os.path.join(RAIZ, 'insumos/solange/INDICE_DE_LAMINAS_ARQUITECTURA_HOGAR_REFUGIO_v2.xlsx')
SALIDA = os.path.join(RAIZ, 'entregables/INDICE_DE_PLANOS_ARQUITECTURA_HOGAR_REFUGIO.docx')
ROMANOS = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X']
TILDES = [('EDIFICACION', 'EDIFICACIÓN'), ('PERIMETRICO', 'PERIMÉTRICO'), ('ZONIFICACION', 'ZONIFICACIÓN'),
          ('UBICACION', 'UBICACIÓN'), ('MODULOS', 'MÓDULOS'), ('EVACUACION', 'EVACUACIÓN'), ('SEÑALIZACION', 'SEÑALIZACIÓN'),
          ('CIELORASO', 'CIELORRASO'), ('METALICAS', 'METÁLICAS'), ('TIPICOS', 'TÍPICOS')]


def _limpio(t):
    t = re.sub(r'\s+', ' ', str(t)).strip()
    for a, b in TILDES:
        t = re.sub(r'\b%s\b' % a, b, t)
    t = re.sub(r'(\d)\s*/\s*(\d)', r'\1 / \2', t)            # "02/ 02" -> "02 / 02"
    t = re.sub(r'OC-\s+(\d)', r'OC-\1', t)
    t = re.sub(r'(\d)-\s+(\d)', r'\1-\2', t)                # "01- 02" -> "01-02"
    t = re.sub(r'DETALLE -(\S)', r'DETALLE - \1', t)
    t = re.sub(r'\s*/\s*(\d\d-\d\d)\b', r'/\1', t)          # "DEPORTIVA /01-05", "JUEGOS/ 01-03" -> ".../01-05"
    t = re.sub(r'\(\s+', '(', re.sub(r'\s+\)', ')', t))           # "(PATIO DE DESCARGA )" -> "(PATIO DE DESCARGA)"
    return t


def _codigo(t):
    return re.sub(r'\s+', '', str(t)).strip()                  # "A-02 /01-03" -> "A-02/01-03"; "CG- 01" -> "CG-01"


def leer():
    ws = openpyxl.load_workbook(XLSX, data_only=True)['ARQUITECTURA']
    grupos, act, cambios = [], None, []
    for r in range(6, ws.max_row + 1):
        a, b, g, h, c = (ws.cell(r, k).value for k in (1, 2, 7, 8, 9))
        if b is None and g is None:
            continue
        if str(g).strip().upper().startswith('LÁMINA'):        # fila de grupo (I ARQUITECTURA, II DESARROLLO DE MÓDULOS, ...)
            act = [_limpio(b), []]; grupos.append(act); continue
        cod, desc = _codigo(g), _limpio(b)
        m = re.match(r'(A-\d\d)/\d\d-(\d\d)$', cod)                # hoja de la descripción "EDIFICACIÓN 06 / 03 - 04"
        d = re.search(r'EDIFICACIÓN (\d\d) / (\d\d) - (\d\d)', desc)
        if m and d and cod != 'A-%s/%s-%s' % d.groups():
            nuevo = 'A-%s/%s-%s' % d.groups(); cambios.append((cod, nuevo, desc)); cod = nuevo
        m = re.match(r'(OC-\d\d)/(\d\d)-(\d\d)$', cod); d = re.search(r'/(\d\d)-(\d\d)$', desc.replace(' ', ''))
        if m and d and (m.group(2), m.group(3)) != d.groups():
            nuevo = '%s/%s-%s' % (m.group(1), d.group(1), d.group(2)); cambios.append((cod, nuevo, desc)); cod = nuevo
        act[1].append(dict(codigo=cod, descripcion=desc, formato=(str(h).strip() if h else '—'), cantidad=c or 1, n_excel=a))
    return grupos, cambios


def generar(grupos):
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7); s.orientation = WD_ORIENT.PORTRAIT
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.6), Cm(1.7), Cm(1.4)
    s.header_distance, s.footer_distance = Cm(0.9), Cm(0.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
    ancho = (s.page_width - s.left_margin - s.right_margin) / 360000.0
    G = I.GRANATE
    h = s.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    I._run(h, I.CORTO, 7.5, True, G); I._run(h, '   |   CUI N.° %s   |   ARQUITECTURA' % I.CUI, 7.5, False, G)
    I._borde_inferior(h, G, 8)
    f = s.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    I._run(f, 'Página ', 7.5, color=I.GRIS); I._campo(f, 'PAGE', 7.5); I._run(f, ' de ', 7.5, color=I.GRIS); I._campo(f, 'NUMPAGES', 7.5)

    def par(txt, tam, neg=False, color=None, antes=0, despues=2, centro=True):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.0
        I._run(p, txt, tam, neg, color)
        return p

    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    par('GERENCIA TERRITORIAL BAJO MAYO - DIRECCIÓN DE INFRAESTRUCTURA', 9, True, I.GRIS, despues=4)
    par('PROYECTO: ' + I.PROYECTO + ', CON CUI N.° ' + I.CUI, 7.5, False, None, despues=6)
    p = par('ÍNDICE DE PLANOS', 15, True, G, despues=2); I._borde_inferior(p, G)
    par('ESPECIALIDAD: ARQUITECTURA', 9.5, True, despues=2, antes=4)
    par(I.UBIC, 8, False, I.GRIS, despues=8)

    anchos = (1.0, 2.6, ancho - 1.0 - 2.6 - 1.9 - 1.4, 1.9, 1.4)
    tb = d.add_table(rows=1, cols=5); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(int(w * 567)))
    for c, w, t in zip(tb.rows[0].cells, anchos, ('N.°', 'LÁMINA', 'DESCRIPCIÓN', 'FORMATO', 'CANT.')):
        c.width = Cm(w); I._celda(c, [(t, 7.5, True, 'FFFFFF')]); I._sombra(c, G)
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(e)
    n, cuenta, total = 0, {}, 0
    for k, (nombre, filas) in enumerate(grupos):
        fila = tb.add_row(); fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        c = fila.cells[0].merge(fila.cells[4])
        I._celda(c, [('%s.  %s  (%d)' % (ROMANOS[k], nombre, len(filas)), 7.5, True, G)], centro=False)
        I._sombra(c, I.GRIS_CLARO)
        for p_ in filas:
            n += 1; total += p_['cantidad']; cuenta[p_['formato']] = cuenta.get(p_['formato'], 0) + 1
            fila = tb.add_row(); fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            cel = fila.cells
            for c, w in zip(cel, anchos):
                c.width = Cm(w)
            I._celda(cel[0], [('%02d' % n if n < 100 else '%d' % n, 8, False, None)])
            I._celda(cel[1], [(p_['codigo'], 8, True, G)])
            I._celda(cel[2], [(p_['descripcion'], 8, True, None)], centro=False)
            I._celda(cel[3], [(p_['formato'], 8, False, None)])
            I._celda(cel[4], [('%d' % p_['cantidad'], 8, False, None)])
    resumen = ', '.join('%d en %s' % (cuenta[f], f) for f in sorted(cuenta) if f != '—')
    sin = cuenta.get('—', 0)
    par('Total: %d láminas (%s%s).' % (total, resumen, ', %d sin formato indicado' % sin if sin else ''), 9, True, antes=8, despues=2, centro=False)
    par('Láminas agrupadas por tema y numeradas de la primera a la última; códigos y formatos según el índice de láminas de '
        'arquitectura del proyecto.', 8, False, I.GRIS, despues=2, centro=False)
    d.core_properties.author = ''
    d.core_properties.title = 'Índice de planos - Arquitectura - Hogar de Refugio Temporal'
    d.save(SALIDA)
    return SALIDA, n, total


if __name__ == '__main__':
    grupos, cambios = leer()
    for nombre, filas in grupos:
        print('%-45s %d' % (nombre, len(filas)))
    for a, b, desc in cambios:
        print('  código repetido corregido: %s -> %s (%s)' % (a, b, desc))
    print(generar(grupos))
