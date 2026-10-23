# -*- coding: utf-8 -*-
"""Índice de planos de drenaje pluvial - Hogar de Refugio Temporal "Mujeres Violentadas" (CUI 2675514), en Word.

Fuente: el PDF de láminas del usuario (insumos/solange/PLANOS_DRENAJE_PLUVIAL_MUJERES_VIOLENTADAS_LAMINAS.pdf, 40 láminas
DPL, setiembre 2026). El índice sigue el orden del PDF, que está organizado por formato de hoja (A0, A1, A2 y A3), y cada
lámina lleva el código y el título de su rótulo, tal como están. De cada página se lee, sin modificarla: código y título
del rótulo, formato del papel y escalas de los dibujos.

Uso:  python3 herramientas/indice_planos_refugio.py
Salida: entregables/INDICE_DE_PLANOS_DRENAJE_PLUVIAL_HOGAR_REFUGIO.docx
        entregables/_calc/indice_planos.json
"""
import json
import os
import re
import subprocess
import unicodedata

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PDF = os.path.join(RAIZ, 'insumos/solange/PLANOS_DRENAJE_PLUVIAL_MUJERES_VIOLENTADAS_LAMINAS.pdf')
DATOS = os.path.join(RAIZ, 'entregables/_calc/indice_planos.json')
SALIDA = os.path.join(RAIZ, 'entregables/INDICE_DE_PLANOS_DRENAJE_PLUVIAL_HOGAR_REFUGIO.docx')
GRANATE, GRIS_CLARO, GRIS = '7B1E2B', 'F2E6E8', '595959'
PROYECTO = ('“CREACIÓN DEL SERVICIO DE PROTECCIÓN A MUJERES VÍCTIMAS DE VIOLENCIA EN SITUACIÓN DE RIESGO DE FEMINICIDIO O '
            'PELIGRE SU INTEGRIDAD Y/O SALUD FÍSICA Y/O MENTAL EN HOGAR DE REFUGIO TEMPORAL DISTRITO DE MORALES DE LA '
            'PROVINCIA DE SAN MARTÍN DEL DEPARTAMENTO DE SAN MARTÍN”')
CUI = '2675514'
CORTO = 'HOGAR DE REFUGIO TEMPORAL - MORALES, SAN MARTÍN'
UBIC = 'Predio rural Los Mangos, distrito de Morales, provincia y departamento de San Martín'

# Títulos de los rótulos (con tildes), por código; se comprueban contra el texto del PDF. DPL-36 figura en dos láminas.
TITULOS = {
    'DPL-01': 'RED GENERAL DE DRENAJE PLUVIAL',
    'DPL-02': 'PERFILES LONGITUDINALES',
    'DPL-03': 'SECCIONES TRANSVERSALES 01', 'DPL-04': 'SECCIONES TRANSVERSALES 02', 'DPL-05': 'SECCIONES TRANSVERSALES 03',
    'DPL-06': 'RED GENERAL COLECTOR PLUVIAL',
    'DPL-07': 'PERFIL LONGITUDINAL GENERAL DEL COLECTOR',
    'DPL-08': 'PERFIL LONGITUDINAL DETALLADO DEL COLECTOR', 'DPL-09': 'PERFIL LONGITUDINAL DETALLADO DEL COLECTOR',
    'DPL-10': 'SECCIONES TRANSVERSALES DEL COLECTOR', 'DPL-11': 'SECCIONES TRANSVERSALES DEL COLECTOR',
    'DPL-12': 'ISOMÉTRICOS DE LAS SECCIONES TRANSVERSALES', 'DPL-13': 'ISOMÉTRICOS DE LAS SECCIONES TRANSVERSALES',
    'DPL-14': 'ISOMÉTRICOS DE LAS SECCIONES TRANSVERSALES',
    **{'DPL-%02d' % (14 + k): 'DETALLE MONTANTE - EDIF%02d' % k for k in range(1, 10)},
    'DPL-24': 'DETALLE DE REGISTRO DE LIMPIEZA - TAPA CIEGA DE CONCRETO ARMADO 0.68x0.68x0.08',
    'DPL-25': 'EMPALME DE CUNETAS A COLECTOR',
    'DPL-26': 'CAJA DE LLEGADA (0+000) Y CAJA DE CAÍDA CON POZA DE DISIPACIÓN (ENTREGA AL CAR MUJERES)',
    'DPL-27': 'ISOMÉTRICOS CONSTRUCTIVOS',
    'DPL-28': 'DETALLE CONSTRUCTIVO: JUNTAS E=1" CON TECNOPOR Y SELLADOR ASFÁLTICO (COLECTOR Y CERCO)',
    'DPL-29': 'REGISTRO DE LIMPIEZA CON TAPA CIEGA DE CONCRETO ARMADO 0.68x0.68x0.08 - COMPONENTES',
    'DPL-30': 'DETALLES DE CANALETA DE TECHO, SOPORTE Y ACCESORIOS',
    'DPL-31': 'ISOMÉTRICO DE CANALETA, SOPORTES, PERNOS, TAPA Y BOQUILLA',
    'DPL-32': 'DETALLES 3D DE LOS ELEMENTOS DEL MONTANTE',
    'DPL-33': 'ELEVACIÓN DE CANALETA Y CORTE DEL MONTANTE PLUVIAL',
    'DPL-34': 'ISOMÉTRICO DEL MONTANTE PLUVIAL TÍPICO',
    'DPL-35': 'ISOMÉTRICO DE LA CAJA SUMIDERO DE PISO',
    'DPL-36': ('DETALLE DE CAJA SUMIDERO DE PISO: CORTE, PLANTA Y DESAGREGADO', 'ISOMÉTRICO DE INSTALACIÓN DE LLORADEROS EN CUNETAS'),
    'DPL-37': 'DETALLE DE LLORADERO: CORTE, TUBO PERFORADO Y DESAGREGADO',
    'DPL-38': 'DETALLE CONSTRUCTIVO FALSA COLUMNA PARA MONTANTE PLUVIAL EN EDIFICACIONES',
    'DPL-39': 'PLANTA DE UBICACIÓN DE LLORADEROS EN CUNETAS COLINDANTES CON ÁREA VERDE',
}
def _llano(t):
    """Mayúsculas sin tildes ni espacios (para comparar con el texto del PDF)."""
    t = unicodedata.normalize('NFD', t.upper())
    return re.sub(r'[^A-Z0-9.,:()+]', '', ''.join(c for c in t if unicodedata.category(c) != 'Mn'))


def _rango(nums):
    nums = sorted(set(nums))
    return '%02d A %02d' % (nums[0], nums[-1]) if len(nums) > 1 else '%02d' % nums[0]


def contenido(cod, t):
    """Lo que distingue a la lámina, leído de sus dibujos."""
    n = int(cod[4:])
    if n in (2, 3, 4, 5):
        s = [int(x) for x in re.findall(r'SECCION (\d+)\b', t)]
        return ('CUNETAS, CON LAS SECCIONES %s' if n == 2 else 'SECCIONES %s') % _rango(s)
    if n in (7, 8, 9):
        p = sorted(set(re.findall(r'\d\+\d{3}\.\d\d', t)))
        return 'PROGRESIVAS %s A %s' % (p[0], p[-1])
    if n in (10, 11, 12, 13):
        m = re.findall(r'PROG\. (\d\+\d{3}\.\d\d) - (?:CF [\d.]+ |interior [^-]+)- ([A-Z ]+?)(?: - ESC|$)', t, re.M)
        return 'PROG. ' + ', '.join('%s (%s)' % (a, b.strip()) for a, b in m)
    if n == 14:
        return 'SECCIONES TÍPICAS: TRAMO NORMAL (0+010.00), CRUCE DE MOTOS (0+040.00) Y CRUCE DE CAMIONES (0+028.80)'
    if 15 <= n <= 23:
        c = sorted(set(re.findall(r'CORTE ([A-Z])-\1', t)))
        return ('CORTE %s-%s' % (c[0], c[0]) if len(c) == 1 else 'CORTES ' + ' Y '.join('%s-%s' % (x, x) for x in c)) if c else ''
    return ''


def leer_pdf():
    info = subprocess.run(['pdfinfo', '-f', '1', '-l', '999', PDF], capture_output=True, text=True, check=True).stdout
    formatos = {int(a): b for a, b in re.findall(r'Page\s+(\d+) size:.*\((A\d)\)', info)}
    planos = []
    for pag in sorted(formatos):
        t = subprocess.run(['pdftotext', '-raw', '-f', str(pag), '-l', str(pag), PDF, '-'],
                           capture_output=True, text=True, check=True).stdout
        lin = t.splitlines()
        # linea del proyecto (con el CUI) seguida del titulo y de COORDINADOR: la del rotulo
        i, j = max((k, m) for k, l in enumerate(lin) if CUI in l
                   for m in range(k + 1, min(k + 6, len(lin))) if lin[m].startswith('COORDINADOR'))
        k = next(k for k, l in enumerate(lin) if k > j and re.search(r'20\d\d$', l.strip()))
        escala, rotulo = lin[k + 1].strip(), lin[k + 2].strip()
        tit_pdf = ' '.join(lin[i + 1:j])
        cod = rotulo
        opciones = TITULOS[cod] if isinstance(TITULOS[cod], tuple) else (TITULOS[cod],)   # noqa
        tit = next((x for x in opciones if _llano(x) == _llano(tit_pdf)), None)
        if not tit:
            raise SystemExit('Título distinto en %s: rótulo "%s" / índice "%s"' % (cod, tit_pdf, opciones))
        esc_dib = sorted(set(re.sub(r'\s', '', e) for e in re.findall(r'ESC(?:ALA)?\.?\s*:?\s*(1\s*/\s*\d+)', t)),
                         key=lambda e: int(e.split('/')[1]))
        planos.append(dict(codigo=cod, titulo=tit, contenido=contenido(cod, t), formato=formatos[pag], escala=escala,
                           escalas_dibujo=esc_dib, pagina=pag, rotulo=rotulo))
    os.makedirs(os.path.dirname(DATOS), exist_ok=True)
    json.dump(planos, open(DATOS, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    return planos


# ----------------------------------------------------------------------------------------------- Word
def _sombra(c, color):
    tcPr = c._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    va = tcPr.find(qn('w:vAlign'))
    if va is not None:
        va.addprevious(sh)
    else:
        tcPr.append(sh)


def _run(p, txt, tam, neg=False, color=None, cursiva=False):
    r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.italic = cursiva; r.font.name = 'Arial'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _celda(c, lineas, centro=True):
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for i, (txt, tam, neg, color) in enumerate(lineas):
        q = c.paragraphs[0] if i == 0 else c.add_paragraph()
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
        pf = q.paragraph_format; pf.space_before = Pt(1 if i == 0 else 0); pf.space_after = Pt(1); pf.line_spacing = 1.0
        _run(q, txt, tam, neg, color)


def _borde_inferior(p, color, tam=12):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', str(tam)), ('space', '2'), ('color', color)):
        e.set(qn('w:' + k), v)
    b.append(e)
    sig = [x for x in pPr if x.tag in (qn('w:shd'), qn('w:tabs'), qn('w:spacing'), qn('w:ind'), qn('w:jc'))]
    sig[0].addprevious(b) if sig else pPr.append(b)


def _campo(p, instr, tam):
    for tipo, txt in (('begin', None), (None, instr), ('separate', None), (None, '1'), ('end', None)):
        r = _run(p, '', tam, color=GRIS)
        if tipo:
            f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), tipo); r._r.append(f)
        elif txt == instr:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r._r.append(it)
        else:
            r.text = txt


def generar(planos):
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7); s.orientation = WD_ORIENT.PORTRAIT
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.6), Cm(2.0), Cm(1.6)
    s.header_distance, s.footer_distance = Cm(0.9), Cm(0.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
    ancho = (s.page_width - s.left_margin - s.right_margin) / 360000.0
    # encabezado y pie
    h = s.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(h, CORTO, 7.5, True, GRANATE)
    _run(h, '   |   CUI N.° %s   |   DRENAJE PLUVIAL' % CUI, 7.5, False, GRANATE)
    _borde_inferior(h, GRANATE, 8)
    f = s.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(f, 'Página ', 7.5, color=GRIS); _campo(f, 'PAGE', 7.5); _run(f, ' de ', 7.5, color=GRIS); _campo(f, 'NUMPAGES', 7.5)

    def par(txt, tam, neg=False, color=None, antes=0, despues=2, centro=True, cursiva=False):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.0
        _run(p, txt, tam, neg, color, cursiva)
        return p

    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    par('GERENCIA TERRITORIAL BAJO MAYO - DIRECCIÓN DE INFRAESTRUCTURA', 9, True, GRIS, despues=4)
    par('PROYECTO: ' + PROYECTO + ', CON CUI N.° ' + CUI, 7.5, False, None, despues=6)
    p = par('ÍNDICE DE PLANOS', 15, True, GRANATE, despues=2); _borde_inferior(p, GRANATE)
    par('ESPECIALIDAD: INSTALACIONES SANITARIAS - DRENAJE PLUVIAL', 9.5, True, despues=2, antes=4)
    par(UBIC, 8, False, GRIS, despues=8)

    anchos = (0.8, 1.6, ancho - 0.8 - 1.6 - 1.8 - 1.9, 1.8, 1.9)
    tb = d.add_table(rows=1, cols=5); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(int(w * 567)))
    for c, w, t in zip(tb.rows[0].cells, anchos, ('N.°', 'LÁMINA', 'DESCRIPCIÓN', 'FORMATO', 'ESCALA')):
        c.width = Cm(w); _celda(c, [(t, 7.5, True, 'FFFFFF')]); _sombra(c, GRANATE)
    trPr = tb.rows[0]._tr.get_or_add_trPr(); e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)
    cuenta = {}
    for p in planos:
        cuenta[p['formato']] = cuenta.get(p['formato'], 0) + 1
    previo = None
    for n, p in enumerate(planos, 1):
        if p['formato'] != previo:                     # el PDF va por formato de hoja: A0, A1, A2 y A3
            previo = p['formato']
            fila = tb.add_row(); fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            c = fila.cells[0].merge(fila.cells[4])
            _celda(c, [('LÁMINAS EN FORMATO %s  (%d)' % (previo, cuenta[previo]), 7.5, True, GRANATE)], centro=False)
            _sombra(c, GRIS_CLARO)
        fila = tb.add_row(); fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        cel = fila.cells
        for c, w in zip(cel, anchos):
            c.width = Cm(w)
        _celda(cel[0], [('%02d' % n, 8, False, None)])
        _celda(cel[1], [(p['codigo'], 8, True, GRANATE)])
        lineas = [(p['titulo'], 8, True, None)]
        if p['contenido']:
            lineas.append((p['contenido'], 6.5, False, GRIS))
        _celda(cel[2], lineas, centro=False)
        _celda(cel[3], [(p['formato'], 8, False, None)])
        esc = [(p['escala'], 8, False, None)]
        if p['escalas_dibujo']:
            esc.append((', '.join(p['escalas_dibujo']), 6.5, False, GRIS))
        _celda(cel[4], esc)
    resumen = ', '.join('%d en %s' % (cuenta[f], f) for f in sorted(cuenta))
    par('Total: %d láminas (%s).' % (len(planos), resumen), 9, True, antes=8, despues=2, centro=False)
    for t in ('Escala del rótulo: INDICADA; debajo, en gris, las escalas de los dibujos de cada lámina. Fecha de las láminas: setiembre 2026.',
              'El N.° es el orden de la lámina en el juego de planos, organizado por formato de hoja (A0, A1, A2 y A3).'):
        par(t, 8, False, GRIS, despues=2, centro=False)
    d.core_properties.author = ''
    d.core_properties.title = 'Índice de planos - Drenaje pluvial - Hogar de Refugio Temporal'
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    planos = leer_pdf()
    for p in planos:
        print('%s p%-3s %-3s %-9s %-62s | %s | %s' % (p['codigo'], p['pagina'], p['formato'], p['escala'], p['titulo'][:62],
                                                  p['contenido'][:70], ','.join(p['escalas_dibujo'])))
    print(generar(planos))
