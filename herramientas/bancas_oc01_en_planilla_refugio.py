"""Hogar de Refugio Temporal (CUI 2675514): agrega a la planilla de metrados del usuario la canaleta de los techos de las bancas
con pergola OC-01 (solo canaleta, soporte, tapa lateral y boquilla), con su hoja de metrado.

Las bancas son 9 = 3 juegos OC-01; cada juego tiene tres techos de 3.79, 4.21 y 3.79 m (planos DPL-41 y DPL-42,
planos_oc01_refugio.py). Canaleta de plancha galvanizada e = 0.50 mm, 0.18 x 0.21; 2 soportes, 2 tapas y 1 boquilla por canaleta.
- Hoja nueva METRADO CANALETA BANCAS (con los estilos de METRADO CANALETAS), despues de METRADO CANALETAS.
- PLANILLA GENERAL DE METRADOS: subtitulo 01.04.03.05.04 despues de ACCESORIOS, con los 4 items; lo demas baja 6 filas.
- RESUMEN: el mismo subtitulo y los 4 items antes de 01.04.04 COLECTOR PLUVIAL FRONTAL; lo demas baja 5 filas.
Ningun otro valor cambia (se comprueba). Uso: python3 herramientas/bancas_oc01_en_planilla_refugio.py
"""
import sys, os, re, html, shutil, zipfile
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI); sys.path.insert(0, os.path.join(AQUI, "mujeres")); sys.path.insert(0, os.path.join(AQUI, "techo"))
from angulo_contramarco_refugio import Libro, recalcular
from bancas_en_planilla_m import agregar_hoja_despues, sin_calcchain, altura
import planos_oc01_refugio as OC

ENT = os.path.join(RAIZ, "insumos", "metrados_vigentes", "METRADO_DRENAJE_MUJERES_VIOLENTADAS.xlsx")
SAL = os.path.join(RAIZ, "entregables", "METRADO_DRENAJE_MUJERES_VIOLENTADAS_CON_BANCAS_OC-01.xlsx")
RES, PG, MC, NB = "RESUMEN", "PLANILLA GENERAL DE METRADOS", "METRADO CANALETAS", "METRADO CANALETA BANCAS"
N_JUEGOS = 3                                   # 9 bancas = 3 juegos OC-01 de 3 techos cada uno
SUB = "RED DE RECOLECCION TECHO DE BANCAS CON PERGOLA (OC-01)"
# (descripcion, und, columna del total en la hoja nueva, fila de PLANILLA que se copia, fila de RESUMEN que se copia)
ITEMS = [("CANALETA DE PLANCHA DE ACERO GALVANIZADO e=0.50mm, SECCIÓN 0.18m x 0.21m", "m", "E", 149, 40),
         ("SOPORTE DE PLATINA F°G° 1\" x 1/8\" P/CANALETA", "und", "G", 150, 41),
         ("TAPA LATERAL DE PLANCHA GALVANIZADA P/CANALETA", "und", "I", 151, 42),
         ("BOQUILLA DE PLANCHA GALVANIZADA P/CONEXIÓN DE CANALETA A TUBERIA PVC Ø4\"", "und", "K", 152, 43)]
DP, DR = 6, 5


def hoja_bancas(L):
    """Hoja nueva armada con filas de METRADO CANALETAS (mismos estilos), renumeradas."""
    mc = L.hoja(MC); filas = []
    def copia(o, n):
        row = re.search(r'<row r="%d"[ >].*?</row>' % o, mc, re.S).group(0)
        row = re.sub(r'(<row r=")%d"' % o, r'\g<1>%d"' % n, row)
        row = re.sub(r'(<c r="[A-Z]+)%d"' % o, r'\g<1>%d"' % n, row)
        filas.append(row)
    plan = [(2, 2), (3, 3), (4, 4), (5, 5), (6, 6), (8, 8), (9, 9), (10, 10), (11, 11), (12, 12), (12, 13), (12, 14), (41, 15),
            (43, 17), (44, 18)] + [(45 + k, 19 + k) for k in range(7)]
    for o, n in plan: copia(o, n)
    merges = ["B2:K2", "C3:K3", "C4:K4", "C5:K5", "C6:K6", "B8:K8", "B9:B10", "C9:C10", "D9:E9", "F9:G9", "H9:I9", "J9:K9", "B17:K17"]
    merges += ["B%d:E%d" % (r, r) for r in range(18, 26)] + ["H%d:K%d" % (r, r) for r in range(18, 26)]
    i, j = mc.index("<sheetData>"), mc.index("</sheetData>")
    cab = re.sub(r' xr:uid="\{[^}]+\}"', "", mc[:i])
    cab = re.sub(r'<dimension ref="[^"]+"/>', '<dimension ref="B2:K25"/>', cab)
    cab = re.sub(r"<sheetViews>.*?</sheetViews>", '<sheetViews><sheetView showGridLines="0" view="pageBreakPreview" zoomScaleNormal="100" '
                 'zoomScaleSheetLayoutView="100" workbookViewId="0"><selection activeCell="F21" sqref="F21"/></sheetView></sheetViews>', cab, flags=re.S)
    cola = mc[j:]
    cola = re.sub(r"<mergeCells.*?</mergeCells>", '<mergeCells count="%d">%s</mergeCells>' % (len(merges), "".join('<mergeCell ref="%s"/>' % m for m in merges)), cola, flags=re.S)
    cola = re.sub(r"<conditionalFormatting.*?</conditionalFormatting>", "", cola, flags=re.S)
    cola = re.sub(r"<dataValidations.*?</dataValidations>", "", cola, flags=re.S)
    cola = re.sub(r' r:id="rId\d+"', "", cola).replace("<legacyDrawing/>", "").replace("<drawing/>", "")
    cola = re.sub(r"<rowBreaks.*?</rowBreaks>", "", cola, flags=re.S)
    return cab + "<sheetData>" + "".join(filas) + cola


def llenar_hoja(L):
    T = OC.TECHOS
    L.texto(NB, "B2", "METRADO DE CANALETA DE TECHO - BANCAS CON PERGOLA OC-01 (JARDINERA + BANCA + PERGOLA)")
    L.texto(NB, "B8", "CANALETA DE PLANCHA DE ACERO GALVANIZADO e=0.50mm - SECCIÓN INTERIOR 0.18 m (ANCHO) x 0.21 m (ALTO) - "
                      "TECHOS A UNA AGUA: UNA CANALETA POR TECHO, SIN UNIONES")
    L.texto(NB, "C9", "N° DE CANALETAS"); L.texto(NB, "D9", "CANALETA 0.18 x 0.21 m")
    L.texto(NB, "B11", "BANCAS OC-01 (%d JUEGOS)" % N_JUEGOS)
    L.texto(NB, "J9", "BOQUILLA A TUBERÍA PVC Ø4\"")
    for r, R in zip((12, 13, 14), T):
        L.texto(NB, "B%d" % r, "Techo %s (L = %.2f m)" % (R["n"], R["L"]))
        L.formula(NB, "C%d" % r, "$F$22"); L.num(NB, "D%d" % r, round(R["L"], 2)); L.formula(NB, "E%d" % r, "C%d*D%d" % (r, r))
        L.formula(NB, "F%d" % r, "$F$24"); L.formula(NB, "G%d" % r, "C%d*F%d" % (r, r))
        L.formula(NB, "H%d" % r, "$F$25"); L.formula(NB, "I%d" % r, "C%d*H%d" % (r, r))
        L.formula(NB, "J%d" % r, "$F$23"); L.formula(NB, "K%d" % r, "C%d*J%d" % (r, r))
    for c in "CEGIK":
        L.formula(NB, "%s15" % c, "SUM(%s11:%s14)" % (c, c))
    L.texto(NB, "B17", "PARÁMETROS DE CÁLCULO - CANALETA DE BANCAS OC-01")
    par = [("Ancho interior de canaleta", 0.18, "m", "Indicación del proyectista (plano DPL-42)"),
           ("Altura interior de canaleta", 0.21, "m", "Indicación del proyectista (plano DPL-42)"),
           ("Espesor de plancha galvanizada", 0.5, "mm", "Canaleta, tapas y boquilla"),
           ("N° de juegos OC-01 (3 techos por juego)", N_JUEGOS, "und", "9 bancas en total; editable"),
           ("Boquillas por canaleta", 1, "und", "Una salida a la tubería PVC Ø4\" (sin sumidero tipo cúpula)"),
           ("Soportes por canaleta", 2, "und", "A 0.80 m (T-1, T-3) y 0.85 m (T-2) de cada extremo"),
           ("Tapas laterales por canaleta", 2, "und", "Una en cada extremo")]
    for r, (t, v, u, cr) in enumerate(par, 19):
        L.texto(NB, "B%d" % r, t); L.num(NB, "F%d" % r, v); L.texto(NB, "G%d" % r, u); L.texto(NB, "H%d" % r, cr)


def valores(f):
    import openpyxl
    wv = openpyxl.load_workbook(f, data_only=True)
    out = {}
    for nom in (RES, PG):
        ws = wv[nom]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, (int, float)): out[(nom, c.coordinate)] = c.value
    return out


VIGENTE = os.path.join(RAIZ, "insumos", "metrados_vigentes", "METRADO_DRENAJE_MUJERES_VIOLENTADAS_CON_BANCAS_OC-01.xlsx")


def main():
    # 22/10/2026: el usuario reordeno la hoja METRADO CANALETA BANCAS (sin la fila 8); su archivo es el vigente y no se pisa
    if os.path.exists(VIGENTE) and "--forzar" not in sys.argv:
        raise SystemExit("Vigente del usuario: %s (no se regenera; usar --forzar solo si se pide)" % os.path.relpath(VIGENTE, RAIZ))
    os.makedirs(os.path.dirname(SAL), exist_ok=True)
    shutil.copy(ENT, SAL); L = Libro(SAL)
    assert L.leer(PG, "B160") == "ACCESORIOS" and L.leer(PG, "A165") == "01.04.04" and L.leer(RES, "A54") == "01.04.04", "la planilla no es la esperada"
    antes = valores(ENT)
    sin_calcchain(L)
    agregar_hoja_despues(L, NB, hoja_bancas(L), MC)
    wbx = L.F["xl/workbook.xml"].decode().replace("'%s'!$B$2:$K$21" % NB, "'%s'!$B$2:$K$25" % NB)
    L.F["xl/workbook.xml"] = wbx.encode()
    llenar_hoja(L)
    # ---- PLANILLA GENERAL: lo de la 165 en adelante baja 6 filas; 164 vacia, 165 subtitulo, 166-169 items, 170 vacia
    L.desplazar(PG, 165, 322, DP)
    L.insertar_copia(PG, 159, 164); L.insertar_copia(PG, 160, 165)
    L.texto(PG, "A165", "01.04.03.05.04"); L.texto(PG, "B165", SUB)
    for k, (d, u, col, rp, _) in enumerate(ITEMS):
        r = 166 + k; L.insertar_copia(PG, rp, r)
        L.texto(PG, "A%d" % r, "01.04.03.05.04.%02d" % (k + 1)); L.texto(PG, "B%d" % r, d); L.texto(PG, "L%d" % r, u)
        L.formula(PG, "K%d" % r, "+'%s'!%s15" % (NB, col))
    L.insertar_copia(PG, 159, 170)
    altura(L, PG, 165, 30.75)
    L.x[PG] = L.x[PG].replace('<dimension ref="A2:L322"/>', '<dimension ref="A2:L%d"/>' % (322 + DP))
    # ---- RESUMEN: lo de la 54 en adelante baja 5 filas; 53 subtitulo, 54-57 items, 58 vacia
    L.desplazar(RES, 54, 75, DR)
    L.insertar_copia(RES, 49, 53); L.texto(RES, "A53", "01.04.03.05.04"); L.texto(RES, "B53", SUB)
    for k, (d, u, col, _, rr) in enumerate(ITEMS):
        r = 54 + k; L.insertar_copia(RES, rr, r)
        L.texto(RES, "A%d" % r, "01.04.03.05.04.%02d" % (k + 1)); L.texto(RES, "B%d" % r, d); L.texto(RES, "C%d" % r, u)
        L.formula(RES, "D%d" % r, "+'%s'!K%d" % (PG, 166 + k))
    def mov(m):
        r = int(m.group(2)); return m.group(1) + str(r + DP if r >= 165 else r)
    partes = re.split(r'(<row r="5[4-7]"[ >].*?</row>)', L.x[RES], flags=re.S)
    for i, p in enumerate(partes):
        if not re.match(r'<row r="5[4-7]"', p):
            partes[i] = re.sub(r"((?:&apos;|')PLANILLA GENERAL DE METRADOS(?:&apos;|')!\$?[A-Z]+\$?)(\d+)", mov, p)
    L.x[RES] = "".join(partes).replace('<dimension ref="A1:G75"/>', '<dimension ref="A1:G%d"/>' % (75 + DR))
    wb = L.F["xl/workbook.xml"].decode()
    wb = wb.replace("'PLANILLA GENERAL DE METRADOS'!$A$1:$L$322", "'PLANILLA GENERAL DE METRADOS'!$A$1:$L$%d" % (322 + DP))
    wb = wb.replace("RESUMEN!$A$1:$D$75", "RESUMEN!$A$1:$D$%d" % (75 + DR))
    L.F["xl/workbook.xml"] = wb.encode()
    L.guardar(SAL)
    print("valores recalculados:", recalcular(SAL, [RES, PG, NB]))
    # ---- comprobacion: los valores anteriores quedan iguales (en su fila nueva)
    despues = valores(SAL); dif = []
    for (nom, ref), v in antes.items():
        col, r = re.match(r"([A-Z]+)(\d+)", ref).groups(); r = int(r)
        r2 = r + (DP if nom == PG and r >= 165 else DR if nom == RES and r >= 54 else 0)
        w = despues.get((nom, "%s%d" % (col, r2)))
        if w is None or abs(w - v) > 1e-6 * max(1, abs(v)): dif.append((nom, ref, v, w))
    print("valores anteriores comprobados:", len(antes), "diferencias:", dif[:10])
    import openpyxl
    wv = openpyxl.load_workbook(SAL, data_only=True)
    for r in range(53, 58): print([wv[RES].cell(r, c).value for c in range(1, 5)])
    print(SAL)


if __name__ == "__main__":
    main()
