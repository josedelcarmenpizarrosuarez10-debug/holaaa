# -*- coding: utf-8 -*-
"""Memoria descriptiva de drenaje pluvial - Hogar de Refugio Temporal "Mujeres Violentadas" (CUI 2675514).

Documento nuevo (el proyecto no tenía memoria descriptiva de la especialidad), con la misma estructura que la del CAR Mujeres
(memoria_descriptiva_m.py) y coherente con:
  - el metrado vigente del usuario con las bancas OC-01 (insumos/metrados_vigentes/..._CON_BANCAS_OC-01.xlsx; solo lectura),
  - la memoria hidrológica con el capítulo del colector (diseno.json y los parámetros del estudio),
  - el juego de láminas DPL (entregables/_calc/indice_planos.json, de indice_planos_refugio.py).
Las figuras son las láminas del PDF del usuario (y las vistas previas de DPL-41/42).

Uso: python3 herramientas/memoria_descriptiva_refugio.py
Salida: entregables/MEMORIA_DESCRIPTIVA_DRENAJE_PLUVIAL_HOGAR_REFUGIO.docx
"""
import json
import os
import re
import subprocess
import sys
import warnings

import docx
import openpyxl
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(RAIZ, 'herramientas', 'techo'))
VIG = os.path.join(RAIZ, 'insumos/metrados_vigentes/METRADO_DRENAJE_MUJERES_VIOLENTADAS_CON_BANCAS_OC-01.xlsx')
PDF = os.path.join(RAIZ, 'insumos/solange/PLANOS_DRENAJE_PLUVIAL_MUJERES_VIOLENTADAS_LAMINAS.pdf')
INDICE = os.path.join(RAIZ, 'entregables/_calc/indice_planos.json')
DJ = json.load(open(os.path.join(RAIZ, 'entregables/_calc/diseno.json')))
FIG = os.path.join(RAIZ, 'entregables/figuras_memoria_descriptiva')
PREV_OC = os.path.join(RAIZ, 'entregables/vistas_previas_oc01')
SALIDA = os.path.join(RAIZ, 'entregables/MEMORIA_DESCRIPTIVA_DRENAJE_PLUVIAL_HOGAR_REFUGIO.docx')
FUENTE, GRANATE, GRIS = 'Arial Narrow', '7B1E2B', '595959'
CUI = '2675514'
CORTO = 'HOGAR DE REFUGIO TEMPORAL'                                   # encabezado
UBIC = 'Predio rural Los Mangos, distrito de Morales, provincia y departamento de San Martín'
LAM_CAR = ('DPL-01', 'Red general de drenaje pluvial del Hogar de Refugio Temporal')   # imagen de la caratula
D = DJ['datos']
I10, I15, P24 = 210.98, 155.66, 121.81           # memoria hidrológica (intensidades de cunetas/canaletas y del colector)


# ----------------------------------------------------------------------------------------------- datos
def libro():
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        return openpyxl.load_workbook(VIG, data_only=True)


def resumen(wb):
    out = []; proyecto = None
    for r in wb['RESUMEN'].iter_rows(values_only=True):
        if r[0] == 'PROYECTO':
            proyecto = re.sub(r'[,\s]*con CUI N\.?\s*°?\s*\d+\s*$', '', str(r[1]).strip()).strip().strip('“”"')
        if r[0] and re.match(r'^\d+(\.\d+)+$', str(r[0]).strip()):
            out.append((str(r[0]).strip(), str(r[1]).strip(), r[2], r[3]))
    return proyecto, out


def cunetas(wb):
    g = [[c.value for c in r] for r in wb['CONCRETO EN CUNETAS'].iter_rows()]
    out = []
    for i, row in enumerate(g):
        for j, v in enumerate(row):
            if isinstance(v, str) and re.match(r'CUNETA EJE \d+', v):
                out.append(dict(eje=v.replace('CUNETA ', ''), L=g[i + 4][j + 3], tapa=g[i + 6][j + 3], conc=g[i][j + 5]))
    return sorted(out, key=lambda c: c['eje'])


def lloraderos(wb):
    filas = [r for r in wb['METRADO LLORADEROS'].iter_rows(values_only=True)]
    tramos = [r for r in filas if isinstance(r[0], (int, float)) and r[1]]
    tot = [v for v in next(r for r in filas if isinstance(r[0], str) and r[0].upper().startswith('TOTAL DE LLORADEROS')) if v not in (None, '')]
    return tramos, tot


# ----------------------------------------------------------------------------------------------- figuras
def figura_lamina(cod):
    """JPG de la lámina (página del PDF del usuario) o PNG de la vista previa de las láminas OC-01."""
    if os.path.exists(os.path.join(PREV_OC, cod + '.png')):
        return os.path.join(PREV_OC, cod + '.png')
    pag = {p['codigo']: p['pagina'] for p in json.load(open(INDICE, encoding='utf-8'))}[cod]
    os.makedirs(FIG, exist_ok=True); base = os.path.join(FIG, cod)
    if not os.path.exists(base + '.jpg'):
        subprocess.run(['pdftoppm', '-f', str(pag), '-l', str(pag), '-singlefile', '-jpeg', '-jpegopt', 'quality=85',
                        '-scale-to', '2200', PDF, base], check=True)
    return base + '.jpg'


# ----------------------------------------------------------------------------------------------- Word
class Memoria:
    def __init__(self):
        self.d = docx.Document(); self.nt = 0; self.nf = 0; self.h1 = 0; self.h2 = 0
        d = self.d; s = d.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(3.0), Cm(2.0), Cm(2.5), Cm(2.0)
        st = d.styles['Normal']; st.font.name = FUENTE; st.font.size = Pt(11)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FUENTE)
        for n, tam in (('Heading 1', 13), ('Heading 2', 11.5)):
            h = d.styles[n]; h.font.name = FUENTE; h.font.size = Pt(tam); h.font.bold = True
            h.font.color.rgb = RGBColor.from_string(GRANATE)
            rf = h.element.get_or_add_rPr().get_or_add_rFonts()
            for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
                rf.set(qn('w:' + k), FUENTE)
            h.paragraph_format.space_before = Pt(12 if n == 'Heading 1' else 8); h.paragraph_format.space_after = Pt(6)
        self.ancho = (s.page_width - s.left_margin - s.right_margin) / 360000.0
        if d.paragraphs:
            d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)

    @staticmethod
    def run(p, txt, tam=None, neg=False, cur=False, color=None):
        r = p.add_run(txt); r.bold = neg; r.italic = cur
        if tam: r.font.size = Pt(tam)
        if color: r.font.color.rgb = RGBColor.from_string(color)
        return r

    def par(self, txt='', tam=11, neg=False, cur=False, color=None, centro=False, antes=0, despues=6, sangria=0.0):
        p = self.d.add_paragraph(); f = p.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
        f.space_before = Pt(antes); f.space_after = Pt(despues); f.line_spacing = 1.15
        if sangria: f.left_indent = Cm(sangria)
        if txt: self.run(p, txt, tam, neg, cur, color)
        return p

    def vin(self, txt):
        p = self.par(despues=3, sangria=0.9); p.paragraph_format.first_line_indent = Cm(-0.45)
        self.run(p, '•  ' + txt); return p

    def t1(self, txt):
        self.h1 += 1; self.h2 = 0
        return self.d.add_heading('%d.  %s' % (self.h1, txt), 1)

    def t2(self, txt):
        self.h2 += 1
        return self.d.add_heading('%d.%d  %s' % (self.h1, self.h2, txt), 2)

    def leyenda(self, txt):
        p = self.par(txt, 10, True, centro=True, antes=4, despues=2); p.paragraph_format.keep_with_next = True

    def fuente(self, txt='Fuente: Elaboración propia.'):
        self.par(txt, 9, cur=True, color=GRIS, despues=8)

    def tabla(self, titulo, cab, filas, anchos, tam=9.5, fuente_txt='Fuente: Elaboración propia.', izq=(0,), total=False):
        self.nt += 1; self.leyenda('Tabla N.° %02d: %s' % (self.nt, titulo))
        t = self.d.add_table(rows=1, cols=len(cab)); t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
        for gc, w in zip(t._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
            gc.set(qn('w:w'), str(int(w * 567)))
        e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); t.rows[0]._tr.get_or_add_trPr().append(e)
        for k, fila in enumerate([cab] + filas):
            celdas = t.rows[0].cells if k == 0 else t.add_row().cells
            ultima = total and k == len(filas)
            for j, (c, v) in enumerate(zip(celdas, fila)):
                c.width = Cm(anchos[j]); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                q = c.paragraphs[0]; q.paragraph_format.space_after = Pt(1); q.paragraph_format.space_before = Pt(1)
                q.alignment = WD_ALIGN_PARAGRAPH.LEFT if (j in izq and k > 0) else WD_ALIGN_PARAGRAPH.CENTER
                self.run(q, str(v), tam, neg=(k == 0 or ultima), color='FFFFFF' if k == 0 else None)
                if k == 0 or ultima:
                    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto')
                    sh.set(qn('w:fill'), GRANATE if k == 0 else 'F2E6E8')
                    tcPr = c._tc.get_or_add_tcPr(); va = tcPr.find(qn('w:vAlign'))
                    va.addprevious(sh) if va is not None else tcPr.append(sh)
            if len(filas) <= 12 and k < len(filas):
                for c in celdas:
                    for q in c.paragraphs: q.paragraph_format.keep_with_next = True
        self.fuente(fuente_txt)

    def figura(self, cod, titulo, ancho=15.5):
        f = figura_lamina(cod)
        self.nf += 1; self.leyenda('Figura N.° %02d: %s (lámina %s)' % (self.nf, titulo, cod))
        p = self.par(centro=True, despues=2); p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(f, width=Cm(ancho))
        self.fuente('Fuente: Lámina %s del expediente de drenaje pluvial.' % cod)

    def salto(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def f2(x): return '%.2f' % x


def caratula(M, proyecto):
    d = M.d
    M.par('GERENCIA TERRITORIAL BAJO MAYO', 16, True, color=GRANATE, centro=True, antes=10, despues=0)
    p = M.par('DIRECCIÓN DE INFRAESTRUCTURA', 11, True, color=GRIS, centro=True, despues=14)
    b = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '14'), ('space', '6'), ('color', GRANATE)):
        e.set(qn('w:' + k), v)
    b.append(e); p._p.get_or_add_pPr().append(b)
    M.par('PROYECTO:', 10, True, centro=True, despues=2)
    M.par('“' + proyecto + '”', 11, True, centro=True, despues=4)
    M.par('CUI N.° ' + CUI, 12, True, color=GRANATE, centro=True, despues=14)
    p = M.par(centro=True, despues=2); p.add_run().add_picture(figura_lamina(LAM_CAR[0]), height=Cm(10.5))
    M.par('%s (lámina %s)' % (LAM_CAR[1], LAM_CAR[0]), 9, cur=True, color=GRIS, centro=True, despues=14)
    t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.width = Cm(M.ancho)
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), GRANATE)
    c._tc.get_or_add_tcPr().append(sh)
    q = c.paragraphs[0]; q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_before = Pt(6)
    M.run(q, 'MEMORIA DESCRIPTIVA', 22, True, color='FFFFFF')
    q = c.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER; q.paragraph_format.space_after = Pt(6)
    M.run(q, 'ESPECIALIDAD: INSTALACIONES SANITARIAS - DRENAJE PLUVIAL', 12, True, color='FFFFFF')
    M.par(UBIC, 10.5, centro=True, antes=14, despues=4)
    M.par('OCTUBRE 2026', 13, True, color=GRANATE, centro=True)


def encabezado(M):
    sec = M.d.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
    pg = OxmlElement('w:pgNumType'); pg.set(qn('w:start'), '1'); sec._sectPr.append(pg)
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    M.run(h, 'MEMORIA DESCRIPTIVA - DRENAJE PLUVIAL', 8.5, True, color=GRANATE)
    M.run(h, '   |   %s - CUI N.° %s' % (CORTO, CUI), 8.5, color=GRANATE)
    b = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '2'), ('color', GRANATE)):
        e.set(qn('w:' + k), v)
    b.append(e); pPr = h._p.get_or_add_pPr(); pPr.insert(1, b)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    M.run(f, 'Página ', 8.5, color=GRIS)
    for tipo in ('begin', 'instr', 'separate', 'txt', 'end'):
        r = M.run(f, '', 8.5, color=GRIS)
        if tipo == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)
        elif tipo == 'txt':
            r.text = '1'
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); r._r.append(fc)
    # primera pagina del cuerpo sin numerar en la caratula
    M.d.sections[0].header.is_linked_to_previous = False


def indice(M):
    M.par('CONTENIDO', 13, True, color=GRANATE, centro=True, despues=8)
    p = M.d.add_paragraph()
    for tipo, txt in (('begin', None), ('instr', 'TOC \\o "1-2" \\h \\z \\u'), ('separate', None),
                      ('txt', 'Actualice el índice (clic derecho > Actualizar campos) para ver la numeración de páginas.'), ('end', None)):
        r = p.add_run()
        if tipo == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; r._r.append(it)
        elif tipo == 'txt':
            r.text = txt; r.italic = True; r.font.color.rgb = RGBColor.from_string(GRIS)
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); r._r.append(fc)
    upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'), 'true'); M.d.settings.element.append(upd)
    M.salto()


# ----------------------------------------------------------------------------------------------- contenido
def generar():
    import planos_oc01_refugio as OC
    wb = libro(); proyecto, R = resumen(wb); res = {it: (de, un, to) for it, de, un, to in R}
    cun = cunetas(wb); tramos_ll, tot_ll = lloraderos(wb)
    planos = json.load(open(INDICE, encoding='utf-8'))
    g = DJ['geom']; perfil = DJ['perfil']; brink = DJ['brink']; poza = DJ['poza']; cl = DJ['caja_llegada']
    llen = max(p['y'] / p['h'] for p in perfil); bl = min(p['techo'] - p['NA'] for p in perfil)
    M = Memoria()
    caratula(M, proyecto)
    encabezado(M)
    indice(M)

    def q(it):            # cantidad del RESUMEN con su unidad
        de, un, to = res[it]
        return ('%s %s' % (f2(to), un)) if un in ('m', 'm²', 'm³', 'kg') else '%d %s' % (round(to), un)

    # ---------------------------------------------------------------- 1
    M.t1('INFORMACIÓN DEL PROYECTO')
    M.tabla('Datos generales del proyecto', ['DATO', 'DESCRIPCIÓN'], [
        ['Unidad ejecutora', 'Gerencia Territorial Bajo Mayo - Dirección de Infraestructura'],
        ['Proyecto', '“' + proyecto + '”'],
        ['Código único de inversiones', 'CUI N.° ' + CUI],
        ['Ubicación', 'Predio rural Los Mangos, distrito de Morales, provincia y departamento de San Martín'],
        ['Especialidad', 'Instalaciones sanitarias - drenaje pluvial'],
        ['Componentes de la especialidad', 'Red de drenaje pluvial del predio (partida 01.04.03) y colector pluvial frontal (partida 01.04.04)']],
        [4.5, M.ancho - 4.5], izq=(0, 1))

    # ---------------------------------------------------------------- 2
    M.t1('ANTECEDENTES DEL PROYECTO')
    M.t2('Introducción')
    M.par('La presente memoria describe el sistema de drenaje pluvial del Hogar de Refugio Temporal para mujeres víctimas de violencia '
          '(CUI %s), que capta las aguas de lluvia de los techos de las edificaciones, de las bancas con pérgola y de las superficies '
          'exteriores del predio, las conduce por canaletas, montantes, red de piso y cunetas de concreto armado, y las entrega a un '
          'colector pluvial frontal cubierto que recorre el frente del predio. El colector forma parte de un colector continuo compartido '
          'con los dos proyectos vecinos: recibe en su inicio el aporte del CAR Varones (CUI 2705619) y entrega el caudal total al '
          'registro R-01 del colector del CAR Mujeres (CUI 2717013), que lo conduce hasta la descarga final.' % CUI)
    M.par('El dimensionamiento hidráulico y estructural se sustenta en el Estudio Hidrológico y de Drenaje Pluvial del proyecto, con '
          'el capítulo del colector frontal, y en la memoria de cálculo del colector. Las cantidades citadas son las del metrado de la '
          'especialidad y los planos son los del juego de láminas DPL de drenaje pluvial (relación al final de esta memoria).')
    M.t2('Normas aplicables')
    for t in ('Reglamento Nacional de Edificaciones, Norma CE.040 Drenaje Pluvial (R.M. N.° 126-2021-VIVIENDA).',
              'Reglamento Nacional de Edificaciones, Norma IS.010 Instalaciones Sanitarias para Edificaciones.',
              'Reglamento Nacional de Edificaciones, Normas E.020 Cargas, E.050 Suelos y Cimentaciones y E.060 Concreto Armado.',
              'Manual de Hidrología, Hidráulica y Drenaje del Ministerio de Transportes y Comunicaciones.',
              'AASHTO LRFD Bridge Design Specifications y Manual de Puentes del MTC (carga de rueda en los cruces vehiculares del colector).',
              'Norma G.050 Seguridad durante la Construcción y D.S. N.° 011-2019-TR.'):
        M.vin(t)

    # ---------------------------------------------------------------- 3
    M.t1('UBICACIÓN DEL PROYECTO')
    M.par('El proyecto se ubica en el predio rural Los Mangos, distrito de Morales, provincia y departamento de San Martín, en el '
          'entorno urbano de Morales y Tarapoto. El frente del predio da a la carretera Oasis; a lo largo de ese frente, por dentro del '
          'cerco perimétrico, se desarrolla el colector pluvial frontal. Los predios colindantes del CAR Varones (aguas arriba) y del CAR '
          'Mujeres (aguas abajo) comparten el mismo colector.')
    M.tabla('Ubicación política', ['NIVEL', 'NOMBRE'], [['Departamento', 'San Martín'], ['Provincia', 'San Martín'], ['Distrito', 'Morales'],
                                                          ['Sector', 'Predio rural Los Mangos - frente a la carretera Oasis']], [5.0, M.ancho - 5.0], izq=(0, 1))

    # ---------------------------------------------------------------- 4
    M.t1('FINALIDAD Y OBJETIVOS')
    M.t2('Finalidad')
    M.par('Evacuar de forma ordenada y segura las aguas pluviales del Hogar de Refugio Temporal, evitando empozamientos, humedad en '
          'las edificaciones, erosión y deterioro de pisos y veredas, y entregarlas al colector compartido en condiciones compatibles con '
          'el proyecto receptor.')
    M.t2('Objetivos')
    for t in ('Captar el agua de los techos de las nueve edificaciones, la guardianía, la losa deportiva y las nueve bancas con pérgola mediante canaletas de plancha galvanizada.',
              'Conducir el agua de techo por montantes de PVC de Ø 4" protegidas en su tramo inferior por falsas columnas de concreto armado.',
              'Recoger el agua de pisos, veredas, estacionamiento y áreas verdes con sumideros, red de piso y cunetas de concreto armado con rejilla.',
              'Aliviar el agua del suelo de las áreas verdes que colindan con las cunetas mediante lloraderos de PVC.',
              'Conducir el caudal total del predio, más el aporte del CAR Varones, por el colector pluvial frontal cubierto hasta el registro R-01 del CAR Mujeres.'):
        M.vin(t)

    # ---------------------------------------------------------------- 5
    M.t1('PARÁMETROS Y CRITERIOS DE DISEÑO')
    M.par('Los caudales se calculan con el método racional, Q = C · I · A · FS, con las intensidades del estudio hidrológico (análisis de '
          'frecuencia de la precipitación máxima en 24 horas, período de retorno de %d años). Las canaletas, montantes y cunetas usan la '
          'intensidad para 10 minutos y el colector, que recoge toda el área del predio más el aporte aguas arriba, la de 15 minutos.' % D['TR'])
    M.tabla('Parámetros de diseño del sistema de drenaje pluvial', ['PARÁMETRO', 'VALOR', 'APLICACIÓN'], [
        ['Período de retorno', '%d años' % D['TR'], 'Todo el sistema'],
        ['Precipitación de diseño (10 min / 20 min)', '%.2f mm / %.2f mm' % (D['P10'], D['P20']), 'Estudio hidrológico'],
        ['Intensidad para tc = 10 min', '%.2f mm/h' % I10, 'Canaletas, montantes, sumideros y cunetas'],
        ['Intensidad para tc = %d min (Dick y Peschke)' % D['tc'], '%.2f mm/h' % I15, 'Colector pluvial frontal'],
        ['Factor de seguridad FS', '%.2f' % D['FS'], 'Todos los caudales'],
        ['Coeficiente de escorrentía', '0.88 techos y veredas; 0.86 estacionamiento; 0.40 patios; 0.29 áreas verdes', 'Método racional'],
        ['Rugosidad de Manning n', '0.013 canaletas; 0.015 cunetas y colector', 'Ecuación de Manning'],
        ['Pendiente longitudinal', '1.0 % canaletas de edificaciones; 0.5 % canaletas de bancas y cunetas; 0.3 % colector', 'Perfiles'],
        ['Diámetro de montantes', 'Ø 4" (Wyly-Eaton, rs = 0.25)', 'Montantes y bajadas'],
        ['Llenado máximo / borde libre mínimo', '85 % / 0.05 m', 'Colector']],
        [5.4, 5.0, M.ancho - 10.4], izq=(0, 1, 2), fuente_txt='Fuente: Estudio hidrológico y de drenaje pluvial del proyecto.')
    M.tabla('Área de drenaje del predio y coeficiente ponderado', ['CONCEPTO', 'VALOR'], [
        ['Área total de drenaje del predio', '%s m²' % f2(D['A_refugio'])],
        ['Σ C · A', '%s m²' % f2(D['CA_refugio'])],
        ['Coeficiente de escorrentía ponderado', '%.3f' % (D['CA_refugio'] / D['A_refugio'])],
        ['Caudal propio que llega al colector (I = %.2f mm/h, FS = %.2f)' % (I15, D['FS']), '%.1f L/s' % D['Q_refugio']]],
        [M.ancho - 4.0, 4.0], izq=(0,), fuente_txt='Fuente: Estudio hidrológico y de drenaje pluvial (capítulo del colector).')

    # ---------------------------------------------------------------- 6
    M.t1('DESCRIPCIÓN DEL SISTEMA DE DRENAJE PLUVIAL')
    M.t2('Esquema general')
    M.par('El agua de los techos cae a canaletas de plancha galvanizada y baja por montantes de PVC de Ø 4" hasta las cunetas o la red de '
          'piso; el agua de pisos y áreas exteriores entra a las cunetas por sus rejillas y por sumideros conectados a la red de piso. Las '
          'cunetas de concreto armado, con pendiente de 0.5 %, recorren el predio por ejes y llegan al frente, donde descargan al colector '
          'pluvial frontal. La red general está en la lámina DPL-01, los perfiles de cunetas en DPL-02 y sus secciones en DPL-03 a DPL-05.')
    M.figura('DPL-01', 'Red general de drenaje pluvial del Hogar de Refugio Temporal')
    M.t2('Canaletas de techo')
    M.par('Los techos de las edificaciones descargan en canaletas de plancha de acero galvanizado, con soportes de platina de F°G° de '
          '1" x 1/8", tapas laterales en los extremos y boquillas de conexión a las montantes (láminas DPL-30, DPL-31 y DPL-33). '
          'Metrado: %s de canaleta (%s), %s de soportes, %s de tapas laterales y %s de boquillas.'
          % (q('01.04.03.05.01.03'), res['01.04.03.05.01.03'][0].split(',', 1)[-1].strip().lower(), q('01.04.03.05.01.04'),
             q('01.04.03.05.01.05'), q('01.04.03.05.01.06')))
    M.t2('Montantes pluviales y falsas columnas')
    M.par('Las montantes son de tubería PVC de Ø 4" (%s), fijadas con abrazaderas de F°G° (%s) y colgadores en los desvíos bajo el alero '
          '(%s); su tramo inferior se aloja en una falsa columna de concreto armado de 0.20 x 0.15 x 1.30 m (%s), que la protege de golpes. '
          'Los detalles por edificación están en las láminas DPL-15 a DPL-23, el montante típico en DPL-32 y DPL-34 y la falsa columna en DPL-38.'
          % (q('01.04.03.05.01.01'), q('01.04.03.05.03.05'), q('01.04.03.05.03.04'), q('01.04.03.05.01.02')))
    M.t2('Techos de las bancas con pérgola (OC-01)')
    T = OC.TECHOS
    M.par('La obra complementaria OC-01 (jardinera + banca + pérgola) se repite tres veces; cada juego tiene tres techos a una agua de %.2f, '
          '%.2f y %.2f m de largo por %.2f m de fondo, con lo que las nueve bancas suman nueve techos. Cada techo descarga en una canaleta '
          'de plancha galvanizada de e = 0.50 mm, de sección interior %.2f x %.2f m, con dos soportes, dos tapas y una boquilla que entrega '
          'a la bajada de PVC de Ø 4" adosada a la columna (el sumidero tipo cúpula del plano se reemplaza por la boquilla). El caudal de '
          'cada techo es menor de 0.6 L/s y la canaleta trabaja con un tirante del orden de 0.01 m. Metrado (partida 01.04.03.05.04): %s de '
          'canaleta, %s de soportes, %s de tapas y %s de boquillas.'
          % (T[0]['L'], T[1]['L'], T[2]['L'], OC.FONDO, OC.CB, OC.CH, q('01.04.03.05.04.01'), q('01.04.03.05.04.02'),
             q('01.04.03.05.04.03'), q('01.04.03.05.04.04')))
    M.figura('DPL-41', 'Isométrico del drenaje pluvial de la OC-01 y planta de techos')
    M.t2('Sumideros y red de piso')
    M.par('Las superficies interiores que no llegan directamente a una cuneta drenan a sumideros metálicos con platina de 1" x 3/16" '
          '(%s), en caja sumidero de piso (láminas DPL-35 y DPL-36), conectados por tubería PVC-CP de Ø 4" (%s) con sus codos y tees '
          'hasta la cuneta más próxima.' % (q('01.04.03.05.02.02'), q('01.04.03.05.02.01')))
    M.t2('Cunetas de concreto armado')
    M.par('Las cunetas son de concreto armado f\'c = 175 kg/cm², de sección interior 0.40 m de ancho y altura variable según su perfil '
          '(muros y losa de 0.10 m, ancho exterior 0.60 m), con pendiente de 0.5 %%, rejilla metálica de platina en los tramos abiertos '
          '(%s), tapas de inspección de 0.60 x 0.60 m (%s) y juntas de dilatación de 1" con sellador asfáltico (%s). En los cruces y '
          'pasos de circulación la cuneta se cubre con losa de concreto armado.' % (q('01.04.03.04.03.01'), q('01.04.03.04.03.02'), q('01.04.03.04.03.03')))
    filas = [[c['eje'].title(), f2(c['L']), f2(c['tapa']) if c['tapa'] else '-', f2(c['conc'])] for c in cun]
    filas.append(['TOTAL', f2(sum(c['L'] for c in cun)), f2(sum(c['tapa'] for c in cun)), f2(sum(c['conc'] for c in cun))])
    M.tabla('Cunetas de concreto armado por eje', ['EJE', 'LONGITUD (m)', 'TRAMO CUBIERTO (m)', 'CONCRETO (m³)'], filas, [3.5, 3.5, 4.0, 3.5],
            izq=(), total=True, fuente_txt='Fuente: Metrado de la especialidad, hoja CONCRETO EN CUNETAS.')
    M.t2('Lloraderos en cunetas colindantes con área verde')
    M.par('En los tramos en que el muro de la cuneta colinda con un área verde se colocan lloraderos de tubería PVC de Ø 3" cada 1.50 m, '
          'con filtro de grava envuelto en geotextil en su boca interior, para aliviar el empuje del agua infiltrada en el jardín sin '
          'arrastrar finos: %d lloraderos en %.2f m de muro (lámina DPL-39 de ubicación, DPL-37 de detalle y DPL-36 de instalación).'
          % (tot_ll[2], tot_ll[1]))
    M.tabla('Tramos de cuneta con lloraderos', ['EJE', 'PROGRESIVAS', 'LADO', 'LONGITUD (m)', 'N.°'],
            [[r[1], r[2], r[3], f2(r[4]), '%d' % r[6]] for r in tramos_ll] + [['TOTAL', '', '', f2(tot_ll[1]), '%d' % tot_ll[2]]],
            [2.4, 4.4, 3.2, 2.8, 1.8], izq=(), total=True, fuente_txt='Fuente: Metrado de la especialidad, hoja METRADO LLORADEROS.')
    M.figura('DPL-39', 'Planta de ubicación de lloraderos en cunetas colindantes con área verde')

    # ---------------------------------------------------------------- 7
    M.t1('COLECTOR PLUVIAL FRONTAL')
    M.t2('Descripción general')
    M.par('El colector es un canal rectangular cubierto de concreto armado f\'c = 210 kg/cm², ubicado dentro del predio, junto al cerco '
          'frontal (eje a %.3f m del paramento exterior del cerco), con la losa superior a nivel de vereda (NPT +%.2f). El tramo recto mide '
          '%.2f m (0+000.00 a 0+%06.2f); al final gira dos veces 45° para alinearse con el registro R-01 del CAR Mujeres y termina en la '
          'caja de caída CC (canal hasta 0+%06.2f, estructura hasta 0+%06.2f). Planta general en la lámina DPL-06 y perfiles en DPL-07 a DPL-09.'
          % (D['eje_desde_cerco'], D['NPT'], g['P_B1'], g['P_B1'], g['P_BRINK'], g['P_FIN']))
    for t in ('Sección: ancho interior b = %.2f m, muros de %.2f m, losa de fondo de %.2f m, losa superior de %.2f m (%.2f m en el cruce de camiones).' % (D['b'], D['e_muro'], D['e_fondo'], D['e_losa'], D['e_losa_camion']),
              'Pendiente S = %.1f ‰; cota de fondo %.2f msnm en el inicio y %.2f msnm en la caída.' % (D['S'] * 1000, D['CF0'], brink['z']),
              'Altura interior variable de %.2f a %.2f m (losa superior a nivel fijo).' % (min(p['h'] for p in perfil), max(p['h'] for p in perfil)),
              'Cruces: ingreso vehicular OE-01 (0+026.32 a 0+031.33, camiones) y portón de motos (0+039.04 a 0+041.04).'):
        M.vin(t)
    M.figura('DPL-06', 'Red general del colector pluvial frontal')
    M.t2('Secciones del colector')
    M.tabla('Secciones del colector pluvial frontal', ['TRAMO', 'MUROS / FONDO (m)', 'LOSA SUPERIOR (m)', 'ARMADURA'], [
        ['Tramo normal (vereda)', '%.2f / %.2f' % (D['e_muro'], D['e_fondo']), f2(D['e_losa']), 'Marco cerrado de 3/8" @0.20 m; longitudinales 3/8" @0.25 m'],
        ['Cruce de motos (0+039.04 a 0+041.04)', '%.2f / %.2f' % (D['e_muro'], D['e_fondo']), f2(D['e_losa']), 'Marco de 3/8" @0.15 m'],
        ['Cruce de camiones (0+026.32 a 0+031.33)', '%.2f / %.2f' % (D['e_muro'], D['e_fondo']), f2(D['e_losa_camion']), 'Doble marco de 1/2" @0.15 m; longitudinales 3/8" @0.20 m']],
        [4.6, 2.8, 2.5, M.ancho - 9.9], izq=(0, 3), fuente_txt='Fuente: Láminas DPL-10 a DPL-14 y memoria de cálculo del colector.')
    M.figura('DPL-14', 'Secciones típicas del colector: tramo normal, cruce de motos y cruce de camiones')
    M.t2('Registros de limpieza')
    M.par('Los registros se ubican a no más de 12 m entre sí, en cada llegada de cuneta y en los quiebres, fuera de los cruces '
          'vehiculares, con tapa ciega de concreto armado de 0.68 x 0.68 x 0.08 m en rebaje de la losa (láminas DPL-24 y DPL-29): '
          '%s en total (partida 01.04.04.05.01).' % q('01.04.04.05.01'))
    M.tabla('Registros, cajas y puntos singulares del colector', ['ELEMENTO', 'PROGRESIVA', 'FUNCIÓN'],
            [[r['nombre'], '0+%06.2f' % r['prog'], (r['tipo'] + ': ' + r['nota']).replace('caida', 'caída').replace('grados', '°').replace('disipacion', 'disipación')]
             for r in DJ['registros']], [2.4, 2.6, M.ancho - 5.0], izq=(2,), fuente_txt='Fuente: Lámina DPL-06 y memoria de cálculo del colector.')
    M.figura('DPL-24', 'Detalle de registro de limpieza con tapa ciega de concreto armado')
    M.t2('Empalmes de cunetas al colector')
    M.par('Seis cunetas llegan al colector por una ventana de 0.40 x H en el muro del lado del predio, con caída libre al fondo, junta '
          'de tecnopor de 1" y registro sobre el empalme (lámina DPL-25). En todos los casos la cota de fondo de la cuneta queda sobre el '
          'nivel de agua del colector, de modo que descargan libremente.')
    M.tabla('Empalmes de cunetas al colector', ['CUNETA', 'PROG.', 'COTA DE FONDO (msnm)', 'H (m)', 'CAÍDA LIBRE (m)', 'PROLONG. (m)'],
            [[c['nombre'], '0+%06.2f' % c['prog'], f2(c['NCF']), f2(c['H']), f2(c['caida_libre']), f2(c['prolong']) if c['prolong'] > 0.1 else '-']
             for c in DJ['cunetas']], [4.6, 1.9, 2.6, 1.6, 2.2, 2.1], izq=(0,), fuente_txt='Fuente: Memoria de cálculo del colector.')
    M.t2('Caja de llegada y caja de caída')
    M.par('En 0+000 la caja de llegada CL (%.2f x %.2f m interiores, poza de %.2f m) recibe el colector del CAR Varones y la cuneta de '
          'su Eje 01; su tapa queda a +%.2f, al nivel del piso de ese proyecto. Al final, la caja de caída CC con poza de disipación de '
          '%.2f x %.2f m y %.2f m de profundidad bajo el fondo del receptor recibe la caída libre del canal (%.2f m) y forma un resalto '
          'ahogado; el agua sale sobre el umbral al registro R-01 del CAR Mujeres en régimen subcrítico (lámina DPL-26).'
          % (D['CL_largo'], D['CL_ancho'], D['CL_poza'], D['NPT_CL'], D['CC_poza_largo'], D['CC_ancho'], D['CC_poza_prof'], poza['caida']))
    M.figura('DPL-26', 'Caja de llegada CL y caja de caída CC con poza de disipación')
    M.t2('Juntas')
    M.par('Juntas de dilatación cada 4.00 m con tecnopor de 1" y sellador asfáltico de 1" x 2" (%s) y junta longitudinal de tecnopor de 1" '
          'con sellador entre el colector y el cimiento del cerco (%s); al otro lado del colector hay terreno natural, sin junta (lámina DPL-28).'
          % (q('01.04.04.06.01'), q('01.04.04.06.02')))
    M.t2('Funcionamiento hidráulico')
    M.par('Con el caudal de diseño de %.1f L/s el tirante normal es %.3f m, la velocidad %.2f m/s y el número de Froude %.2f: régimen '
          'subcrítico en todo el tramo, con perfil M2 controlado por la caída libre al final del canal (tirante crítico %.3f m). El llenado '
          'máximo es %.0f %% y el borde libre mínimo %.2f m. La velocidad está entre 0.90 y 3.0 m/s (CE.040) y el colector se autolimpia '
          'aun con caudales parciales.' % (DJ['Q'], DJ['yn'], DJ['V_n'], DJ['F_n'], DJ['yc'], llen * 100, bl))
    M.t2('Verificación estructural')
    M.par('El colector se verifica como marco cerrado de concreto armado con el empuje del relleno, las sobrecargas peatonales (Norma E.020) '
          'y, en los cruces, la carga de rueda del camión HL-93 con impacto (AASHTO LRFD) y una rueda de 1 t en el portón de motos; el diseño '
          'por resistencia sigue la Norma E.060. Todas las verificaciones de flexión y cortante cumplen (memoria de cálculo del colector).')

    # ---------------------------------------------------------------- 8
    M.t1('APORTES DE PROYECTOS COLINDANTES')
    M.tabla('Caudales del colector pluvial frontal', ['APORTE', 'ÁREA (m²)', 'Σ C·A (m²)', 'Q (L/s)'], [
        ['CAR Varones (CUI 2705619), dato de su memoria', f2(D['A_varones']), f2(D['CA_varones']), '%.1f' % D['Q_varones']],
        ['Hogar de Refugio Temporal (CUI %s)' % CUI, f2(D['A_refugio']), f2(D['CA_refugio']), '%.1f' % D['Q_refugio']],
        ['Caudal entregado al R-01 del CAR Mujeres (CUI 2717013)', '', '', '%.1f' % DJ['Q']]],
        [M.ancho - 8.4, 2.8, 2.8, 2.8], izq=(0,), total=True, fuente_txt='Fuente: Memoria hidrológica (capítulo del colector).')
    M.par('El caudal entregado no excede los 560.7 L/s previstos por el CAR Mujeres en su registro R-01 y la cota de fondo de llegada '
          '(%.2f msnm) es superior a la del R-01 (%.2f msnm).' % (brink['z'], D['CF_R01']))

    # ---------------------------------------------------------------- 9
    M.t1('METAS FÍSICAS')
    for raiz, titulo in (('01.04.03.', 'Metas físicas del sistema de drenaje pluvial del predio (partida 01.04.03)'),
                         ('01.04.04.', 'Metas físicas del colector pluvial frontal (partida 01.04.04)')):
        filas = []
        for it, de, un, to in R:
            if not it.startswith(raiz):
                continue
            filas.append([it, de, un or '', (f2(to) if un in ('m', 'm²', 'm³', 'kg') else '%d' % round(to)) if un else ''])
        M.tabla(titulo, ['ÍTEM', 'DESCRIPCIÓN', 'UND', 'METRADO'], filas, [3.6, M.ancho - 7.2, 1.4, 2.2], tam=8.5, izq=(0, 1),
                fuente_txt='Fuente: Metrado de la especialidad (hoja RESUMEN).')

    # ---------------------------------------------------------------- 10-12
    M.t1('OPERACIÓN Y MANTENIMIENTO')
    for t in ('Antes de la temporada de lluvias (setiembre a abril) y después de cada lluvia intensa, limpiar canaletas, boquillas y rejillas de hojas y basura, en especial las canaletas de las bancas bajo la pérgola.',
              'Revisar el estado de soportes, abrazaderas y colgadores; reponer los sueltos o corroídos y repintar con pintura rica en zinc.',
              'Retirar sedimentos de cunetas, cajas sumidero y del fondo del colector a través de los registros, al menos dos veces al año.',
              'Verificar que los lloraderos no estén obstruidos y que el área verde no se riegue en exceso contra los muros de las cunetas.',
              'Revisar y reponer el sellador de las juntas cuando presente fisuras o desprendimientos.',
              'Mantener las tapas de los registros en su sitio y no colocar cargas vehiculares fuera de los cruces diseñados.'):
        M.vin(t)
    M.t1('CONCLUSIONES')
    for t in ('El sistema de drenaje pluvial capta y conduce el agua de techos, bancas con pérgola y superficies exteriores del predio con canaletas, montantes de Ø 4", red de piso y cunetas de concreto armado de 0.40 m de ancho interior.',
              'El colector pluvial frontal de %.2f m de ancho conduce %.1f L/s en régimen subcrítico, con llenado máximo de %.0f %% y borde libre mínimo de %.2f m, y entrega al R-01 del CAR Mujeres en condiciones compatibles.' % (D['b'], DJ['Q'], llen * 100, bl),
              'Las canaletas de las nueve bancas con pérgola (0.18 x 0.21 m) tienen capacidad muy superior al caudal de sus techos.',
              'Las cantidades de esta memoria son las del metrado de la especialidad y los planos son los del juego de láminas DPL de drenaje pluvial.'):
        M.vin(t)
    M.t1('RECOMENDACIONES')
    for t in ('Verificar en el replanteo la ubicación del colector respecto del cerco y del R-01 del CAR Mujeres, y las cotas de fondo de llegada del CAR Varones.',
              'Coordinar la ejecución del colector con los proyectos colindantes para que los empalmes CL y CC se construyan con las cotas del expediente.',
              'Comparar la presión transmitida al suelo con la capacidad portante del estudio de mecánica de suelos del proyecto.',
              'Ejecutar las pruebas con agua de canaletas, montantes y red de piso antes de cerrar falsas columnas y rellenos.'):
        M.vin(t)

    # ---------------------------------------------------------------- 13
    M.t1('RELACIÓN DE PLANOS')
    M.tabla('Relación de planos de drenaje pluvial', ['N.°', 'LÁMINA', 'DESCRIPCIÓN', 'FORMATO'],
            [['%02d' % i, p['codigo'], p['titulo'], p['formato']] for i, p in enumerate(planos, 1)], [1.0, 1.8, M.ancho - 4.6, 1.8],
            tam=8.5, izq=(2,), fuente_txt='Fuente: Índice de planos de drenaje pluvial (orden del juego de láminas, por formato de hoja).')
    M.par('Las bancas con pérgola OC-01 se detallan además en las láminas DPL-41 (isométrico del drenaje pluvial y planta de techos) y '
          'DPL-42 (detalle isométrico de la salida de la canaleta), entregadas en DXF.', 10)
    M.d.core_properties.author = ''
    M.d.core_properties.title = 'Memoria descriptiva - Drenaje pluvial - Hogar de Refugio Temporal'
    M.d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
