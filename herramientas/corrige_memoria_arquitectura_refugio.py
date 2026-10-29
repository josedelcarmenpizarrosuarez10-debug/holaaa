# -*- coding: utf-8 -*-
"""Corrección de la memoria descriptiva de arquitectura - Hogar de Refugio Temporal (CUI 2675514).

Parte del Word "culminado" de la arquitecta (insumos/solange/arquitectura/, no se modifica) y escribe una copia corregida con el
mismo formato e imágenes. Mismo método que la memoria del CAR Varones (varones/corrige_memoria_arquitectura_v.py, de donde se
toman las piezas de edición, el índice y la paginación):
  - texto copiado de un local educativo ("Institución Educativa", "aulas", "aprendizaje", "infraestructura educativa", NTIE 001-2017
    derogada, R.V.M. 307-2019 de racionalización magisterial) y erratas;
  - datos del terreno (titular, "0.085562 ha", "8,556.24. m²", área sin unidad, leyendas de cuadros e imágenes repetidas, usos dentro
    de 100 m);
  - análisis de incompatibilidades de ubicación que estaba vacío (solo el título del cuadro): 2.3.1 con el método y el Cuadro N° 4 de
    18 incompatibilidades, y 2.3.2 con el Cuadro N° 5 de condiciones de ubicación del hogar (incompatibilidades_refugio.py);
  - fuente del terreno tipo II (es la R.V.M. N° 208-2019-MINEDU, cuadro 2, no los lineamientos), normativa completa;
  - drenaje pluvial del proyecto (no "cunetas existentes");
  - sustento normativo pegado: numeración dentro del numeral 3.2 y tablas renumeradas;
  - programa arquitectónico: citas normativas iguales a las del sustento (sin cambiar áreas) y erratas;
  - listado de planos con las láminas del índice de arquitectura de la arquitecta (indice_arquitectura_refugio.py);
  - índice (campo TOC de Word) con el número de página de cada título.

Uso: python3 herramientas/corrige_memoria_arquitectura_refugio.py
Salida: entregables/arquitectura/MEMORIA_DESCRIPTIVA_ARQUITECTURA_HOGAR_REFUGIO_CORREGIDA.docx
        entregables/revision_interna/REVISION_MEMORIA_ARQUITECTURA_HOGAR_REFUGIO.docx (uso interno)
"""
import copy
import os
import re
import sys

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI); sys.path.insert(0, os.path.join(AQUI, 'varones'))
import corrige_memoria_arquitectura_v as CV  # noqa: E402  (edición de runs, índice y paginación)
import incompatibilidades_refugio as IR  # noqa: E402  (Cuadro N° 1 de 18 y condiciones del hogar)
import indice_arquitectura_refugio as IAR  # noqa: E402  (láminas del índice de arquitectura)

ENT = os.path.join(RAIZ, 'insumos/solange/arquitectura/MEMORIA_DESCRIPTIVA_ARQUITECTURA_HOGAR_REFUGIO_CULMINADO.docx')
SALIDA = os.path.join(RAIZ, 'entregables/arquitectura/MEMORIA_DESCRIPTIVA_ARQUITECTURA_HOGAR_REFUGIO_CORREGIDA.docx')
SALIDA_N = os.path.join(RAIZ, 'entregables/revision_interna/REVISION_MEMORIA_ARQUITECTURA_HOGAR_REFUGIO.docx')
CV.MEM_C = SALIDA
CAMBIOS = CV.CAMBIOS
poner, reemplazar, celda, _copia = CV.poner, CV.reemplazar, CV.celda, CV._copia
P_ = docx.text.paragraph.Paragraph

# textos que se cambian tal cual (viejo -> nuevo); cada uno debe existir en un solo párrafo
TEXTOS = [
    ('El proyecto busca protege temporalmente', 'El proyecto busca proteger temporalmente'),
    ('limitaciones están relaciones con', 'limitaciones están relacionadas con'),
    ('se ha planteado construir la Institución Educativa, se encuentra ubicado:',
     'se ha planteado construir el Hogar de Refugio Temporal se encuentra ubicado en:'),
    ('se ingresa por la carretera F.B.T, altura del sector el Tumi, margen derecha aproximadamente 340m por una trocha carrozable que '
     'une con la vía de evitamiento.',
     'se ingresa por la carretera Fernando Belaunde Terry (F.B.T.), a la altura del sector El Tumi, margen derecha, y luego '
     'aproximadamente 340 m por la carretera Oasis, vía carrozable que une con la vía de evitamiento.'),
    ('(septiembre,2026)', '(septiembre, 2026)'),
    ('A La zona de estudio es un área que pertenece a la Aldea Infantil Virgen Del Pilar.',
     'La zona de estudio es un área que perteneció a la Aldea Infantil Virgen del Pilar y que desde el año 2025 es de propiedad del '
     'Gobierno Regional de San Martín (partida registral N.° 05003469).'),
    ('1. % y una pendiente transversal de 0.20%.', '1 % y una pendiente transversal de 0.20 %.'),
    ('La zona de estudio por la margen derecha, discurren aguas de drenaje pluvial, el cual deberá tomar en cuenta para el diseño de '
     'distribución arquitectónico.',
     'Por la margen derecha de la zona de estudio discurren aguas de drenaje pluvial, que se tomaron en cuenta en la distribución '
     'arquitectónica y que se captan y conducen con el sistema de drenaje pluvial del proyecto.'),
    ('m, esta pertenece a la Aldea Infantil Virgen del Pilar; este estudio servirá',
     'm, de propiedad del Gobierno Regional de San Martín; este estudio servirá'),
    ('topografía realizada en la zona de estudio es de 35,547.00 m2',
     'topografía realizada, la zona de estudio tiene un área de 35,547.00 m²'),
    ('TERRENO MATRÍZ', 'TERRENO MATRIZ'),
    ('predio matríz', 'predio matriz'),
    ('CUADRO DE TOLERANCIAS CATAS', 'CUADRO DE TOLERANCIAS CATASTRALES'),
    ('es de 8,556.24 y el perímetro', 'es de 8,556.24 m² y el perímetro'),
    ('Cuadro N°3: Superposición de polígonos.', 'Cuadro N°3: Datos técnicos de la poligonal resultante.'),
    ('Área: 0.085562 ha', 'Área: 0.85562 ha'),
    ('COLINDANCIA E INCOMPATILIDAD', 'COLINDANCIA E INCOMPATIBILIDAD'),
    ('uso comercio vecinal, uso vivienda taller y uso hogar de refugio temporal.',
     'uso comercio vecinal, uso vivienda taller y los Centros de Acogida Residencial colindantes (CAR Varones y CAR Mujeres), '
     'servicios de protección del Estado compatibles con el hogar.'),
    ('disminución de la reincidencia, INCREMENTAR las denuncias', 'disminución de la reincidencia e incremento de las denuncias'),
    ('Cuadro N°04: Casos con características de feminicidio', 'Cuadro N°06: Casos con características de feminicidio'),
    ('A continuación, veremos cómo se manifiesta este problema, sus consecuencias y las vías de ayuda disponibles. Según el Programa '
     'Nacional Warmi Ñan.',
     'A continuación se presentan las estadísticas del Programa Nacional Warmi Ñan sobre la manifestación de este problema y sus '
     'consecuencias.'),
    ('confort bioclimático de la infraestructura educativa', 'confort bioclimático de la infraestructura'),
    ('Se toma como referencia lo establecido en los  “LINEAMIENTOS ARQUITECTÓNICOS PARA EL DISEÑO Y CONSTRUCCIÓN DE REFUGIOS: '
     'especializados para mujeres, sus hijas e hijos, víctimas de violencia extrema, familiar y/o de género”',
     'Se toma como referencia la clasificación de terrenos del Cuadro N° 2 de la R.V.M. N° 208-2019-MINEDU, adaptada al Hogar de '
     'Refugio Temporal y a los “Lineamientos arquitectónicos para el diseño y construcción de refugios especializados para mujeres, '
     'sus hijas e hijos, víctimas de violencia extrema, familiar y/o de género”.'),
    ('Fuente: “LINEAMIENTOS ARQUITECTÓNICOS PARA EL DISEÑO Y CONSTRUCCIÓN DE REFUGIOS: especializados para mujeres, sus hijas e hijos, '
     'víctimas de violencia extrema, familiar y/o de género”',
     'Fuente: R.V.M. N° 208-2019-MINEDU, Cuadro N° 2 (condiciones del terreno tipo II), adaptado al Hogar de Refugio Temporal.'),
    ('víctimas de violencia extrema, familiar y/o de género”. TOMO I, II Y III',
     'víctimas de violencia extrema, familiar y/o de género” (ONU-Habitat / CONAVIM, 2016), tomos I, II y III.'),
    ('Acorde con el Artículo 27.- de la NTIE 001 – 2017, el objetivo principal',
     'Acorde con los criterios de diseño bioclimático de la Norma Técnica “Criterios Generales de Diseño para Infraestructura '
     'Educativa” (R.V.M. N° 010-2022-MINEDU) y de la Guía de aplicación de arquitectura bioclimática en locales educativos, el '
     'objetivo principal'),
    ('en los ambientes públicas haciendo', 'en los ambientes públicos haciendo'),
    ('siendo esencial para el aprendizaje y la productividad.',
     'esenciales para el descanso, la recuperación y las actividades de las residentes.'),
    ('hasta ser trasladados finalmente a las cunetas existentes en el exterior',
     'hasta ser conducidas a las cunetas y al colector pluvial frontal del proyecto'),
    ('ESTRATÉGIAS DE CONTROL TÉRMICO', 'ESTRATEGIAS DE CONTROL TÉRMICO'),
    ('El diseño arquitectónico de hogar de refugio temporal debe responder al RNE - NORMA A.030 - Hospedaje art.5. y a los '
     'lineamientos arquitectónicos para el diseño y construcción de refugios: especializados para mujeres, sus hijas e hijos, '
     'víctimas de violencia extrema, familiar y/o de género”. tomo i, ii y iii los Con este concepto muy presente,',
     'El diseño arquitectónico del Hogar de Refugio Temporal debe responder al RNE (Norma EM.110 Confort térmico y lumínico con '
     'eficiencia energética) y a los Lineamientos arquitectónicos para el diseño y construcción de refugios especializados para '
     'mujeres, sus hijas e hijos, víctimas de violencia extrema, familiar y/o de género (tomos I, II y III). Con este concepto muy '
     'presente,'),
    ('dentro de las aulas y demás ambientes', 'dentro de las habitaciones y demás ambientes'),
    ('acorde al Artículo 23.3. Medidas contra la radiación UV-B, de la NTIE 001 – 2017.',
     'en concordancia con las medidas de protección contra la radiación UV de la Norma Técnica “Criterios Generales de Diseño para '
     'Infraestructura Educativa” (R.V.M. N° 010-2022-MINEDU).'),
    ('por medio de los resultados del estudio de demanda y por la normativa correspondiente: Reglamento Nacional de Edificaciones '
     '(RNE) y los Lineamientos arquitectónicos-Refugios para mujeres sus hijas e hijos tomo I, II, II Norma Técnica de '
     'Infraestructura Educativa RVM Nº 208-2019 MINEDU - Criterios de Diseño para Locales Educativos',
     'por medio de los resultados del estudio de demanda y por la normativa correspondiente: el Reglamento Nacional de Edificaciones '
     '(RNE), los Lineamientos arquitectónicos para refugios de mujeres, sus hijas e hijos (tomos I, II y III) y las Normas Técnicas '
     'R.V.M. N° 104-2019-MINEDU y R.V.M. N° 208-2019-MINEDU.'),
    ('y R.V.M. N.° 104-2019-MINEDU.', 'y R.V.M. N.° 104-2019-MINEDU y N.° 208-2019-MINEDU.'),
    ('se tomó en consideración el Cuadro N°40 de la Norma Técnica R.V.N N° 208-2019-MINEDU y',
     'se tomaron en consideración el Cuadro N°40 de la Norma Técnica R.V.M. N° 208-2019-MINEDU, las normas del RNE y los '
     'Lineamientos arquitectónicos para refugios, según el sustento normativo del numeral anterior.'),
]
LISTA = [('Area de bienvenida', 'Área de bienvenida'), ('Ss.hh mujeres', 'SS.HH. mujeres'), ('Ss.hh varones', 'SS.HH. varones'),
         ('Ss.hh discapacitados', 'SS.HH. discapacitados'), ('Ss.hh hab. Simple', 'SS.HH. hab. simple'),
         ('Ss.hh habitación', 'SS.HH. habitación'), ('Almacen', 'Almacén'), ('Cuarto de maquinas', 'Cuarto de máquinas'),
         ('Sub estación', 'Subestación'), ('Garita + ss.hh', 'Garita + SS.HH.'), ('Comedor/sum', 'Comedor/SUM'),
         ('Ss.hh habitación triple-tipo II', 'SS.HH. habitación triple-tipo II')]

NORMAS_NUEVAS = [   # se agregan a la normativa (citadas en la memoria o en el sustento)
    'Ley N° 30364, Ley para prevenir, sancionar y erradicar la violencia contra las mujeres y los integrantes del grupo familiar, y '
    'su Reglamento (D.S. N° 009-2016-MIMP).',
    'R.M. N° 150-2016-MIMP, Lineamientos para la atención y funcionamiento de los Hogares de Refugio Temporal.',
    'R.V.M. N° 104-2019-MINEDU, Norma Técnica “Criterios de diseño para locales educativos del nivel de Educación Inicial”.',
    'R.V.M. N° 208-2019-MINEDU, Norma Técnica “Criterios de diseño para locales educativos de primaria y secundaria”.',
    'Ley N° 30102, Ley que dispone medidas preventivas contra los efectos nocivos para la salud por la exposición prolongada a la '
    'radiación solar.',
    'Guía de aplicación de arquitectura bioclimática en locales educativos (MINEDU).',
]

# programa arquitectónico: descripción -> (índice de ocupación o None si no cambia, cita normativa) = las del sustento
R208, R104 = 'R.V.M. N° 208-2019-MINEDU ', 'R.V.M. N° 104-2019-MINEDU '
CITAS = {
    'CISTERNA Y TANQUE ELEVADO': (None, 'RNE-NORMA IS.010'), 'PTAR': (None, 'RNE-NORMA IS.020'),
    'CUNETAS DE DRENAJE PLUVIAL': (None, 'RNE-NORMA OS.060'), 'CERCO PERIMÉTRICO': (None, 'Lineamientos ONU-Habitat 2.4'),
    'GARITA + SS.HH': (None, R208 + 'cuadro 40'),
    'SALA DE REUNIONES': (None, 'RNE-NORMA A.090-Servicios comunales art.11'),
    'CUARTO DE LIMPIEZA': (None, R208 + 'num. 12.3.6'),
    'CUARTO DE CÁMARAS DE VIGILANCIA': (None, R208 + 'num. 12.3.3 y cuadro 40'),
    'DATA CENTER': (None, 'RNE-NORMA EM.010'), 'ARCHIVO': (None, R208 + 'cuadro 29'), 'DEPÓSITO': (None, R208 + 'cuadro 29'),
    'TALLER DE BISUTERÍA': (None, R208 + 'cuadro 20'), 'TALLER DE REPOSTERÍA': (None, R208 + 'cuadro 20'),
    'TALLER DE COSTURA': (None, R208 + 'cuadro 20'), 'MÓDULO INFANTIL': (None, R104 + 'cuadro 18'),
    'MÓDULO DE ENSEÑANZA': (None, R208 + 'cuadro 12'),
    'HABITACIONES TRIPLE-TIPO II': ('4.00 m2 por persona', 'RNE-NORMA A.030-Hospedaje art.5'),
    'SS.HH HAB.TRIPLE-TIPO III': (None, 'RNE-NORMA A.030-Hospedaje Anexo 3'),
    'COCINA': ('no aplica', 'Lineamientos ONU-Habitat 3.2 / ' + R208 + 'anexo 4'),
    'DESPENSA': ('no aplica', 'Lineamientos ONU-Habitat 3.2 / ' + R208 + 'anexo 4'),
    'ALACENA': ('no aplica', 'Lineamientos ONU-Habitat 3.2 / ' + R208 + 'anexo 4'),
    'COMEDOR/SUM': ('1.00 m2 por persona', R208 + 'cuadro 22'), 'LAVANDERÍA GENERAL': (None, R208 + 'num. 12.2.1 i'),
    'BIOHUERTO': (None, R104 + '/ Lineamientos ONU-Habitat 3.9'), 'PATIO CÍVICO': (None, R208 + 'num. 11.6.3'),
    'PATIO DE JUEGOS': (None, R104 + 'num. 11.3'),
    'AREA DE DESCANSO (BANCA + PÉRGOLA)': ('1.00 m2 por persona', 'RNE-NORMA A.090-Servicios comunales art.11'),
    'LOSA MULTIUSO': (None, R208 + 'cuadro 23'), 'VESTIDORES': (None, 'RNE-NORMA A.040-Educación art.20'),
    'ALMACEN DE IMPLEMENTOS DEPORTIVOS': (None, R208 + 'num. 11.5.8'), 'ALMACEN DE DONACIONES': (None, 'Lineamientos ONU-Habitat 3.4'),
    'ALMACEN GENERAL Y DE LIMPIEZA': ('mín. 10.00 m2', R208 + 'num. 12.3.1'),
    'MAESTRANZA': ('no aplica', R208 + 'num. 12.3.2'), 'ALMACEN DE RESIDUOS SÓLIDOS': (None, R208 + 'num. 12.3.5'),
    'TABLERO ELÉCTRICO': (None, R208 + 'num. 12.3.7'), 'CUARTO DE MAQUINAS': (None, R208 + 'num. 12.3.4'),
    'SUB ESTACIÓN': (None, 'RNE-NORMA EM.010'), 'GRUPO ELECTRÓGENO': (None, 'RNE-NORMA EM.010'),
}
METODO = ('El análisis verifica que el terreno no presenta incompatibilidades de uso por cercanía a actividades, instalaciones o '
          'servidumbres reguladas por la normativa vigente (RNE, Norma A.040, art. 7 b; R.V.M. N° 208-2019-MINEDU, art. 8 c) y que '
          'reúne las condiciones de ubicación propias de un hogar de refugio (RNE, Norma A.090, arts. 3 y 4; lineamientos para '
          'refugios, numerales 2.1 a 2.8). Para cada dispositivo legal se define el área de influencia como la franja medida desde el '
          'lindero con la distancia que fija: 50 m (estaciones de servicio), 100 m (establecimientos de salud, GLP, bebidas '
          'alcohólicas y plantas de combustibles), 150 m (velatorios y cementerios) y 200 m (ductos de hidrocarburos); para los que no '
          'fijan distancia se verifica que la restricción no alcance al terreno. Se reconoció el entorno y se revisó la cartografía y '
          'las imágenes satelitales de la zona.')
CONDIC_INTRO = ('Además de las incompatibilidades por cercanía, el terreno y el proyecto cumplen las condiciones de ubicación propias de '
                'un Hogar de Refugio Temporal, que establecen el Reglamento Nacional de Edificaciones y los lineamientos para refugios:')
CONCLUSION = ('Conclusión: el terreno cumple las %d condiciones del Cuadro N° 4 y las %d del Cuadro N° 5. No existen dentro de su área de '
              'influencia actividades, instalaciones ni servidumbres incompatibles con la permanencia de mujeres, niñas y niños; la '
              'ubicación cumple la Norma A.090 (arts. 3 y 4) y la Norma A.040 (art. 7 b), y responde a los lineamientos para refugios. '
              'Por lo tanto, el terreno es apto para el Hogar de Refugio Temporal.')


def _tabla(modelo, ancla, cab, filas, anchos):
    """Tabla nueva con el formato de 'modelo' (encabezado y primera fila), anchos en dxa, insertada después de 'ancla'."""
    t = copy.deepcopy(modelo._tbl); ancla.addnext(t)
    T = docx.table.Table(t, modelo._parent)
    plantilla = copy.deepcopy(T.rows[1]._tr)
    for r in T.rows[1:]:
        t.remove(r._tr)
    for c, txt in zip(T.rows[0].cells, cab):
        celda(c, txt)
    for f in filas:
        tr = copy.deepcopy(plantilla); t.append(tr)
        for c, txt in zip(docx.table._Row(tr, T).cells, f):
            celda(c, txt)
    for gc, w in zip(t.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(w))
    for r in T.rows:
        r._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c, w in zip(r.cells, anchos):
            tcW = c._tc.get_or_add_tcPr().get_or_add_tcW(); tcW.set(qn('w:w'), str(w)); tcW.set(qn('w:type'), 'dxa')
    for c in T.rows[0].cells:
        for q in c.paragraphs:
            q.paragraph_format.keep_with_next = True
    tw = t.tblPr.find(qn('w:tblW'))
    if tw is not None:
        tw.set(qn('w:w'), str(sum(anchos))); tw.set(qn('w:type'), 'dxa')
    return t


def analisis_ubicacion(d):
    P = d.paragraphs
    titulo = next(p for p in P if p.text.strip() == 'ANÁLISIS DE INCOMPATIBILIDADES DE UBICACIÓN EN EL PROYECTO')
    sub = next(p for p in P if p.text.strip() == 'TOLERANCIAS CATASTRALES')
    cuerpo = next(p for p in P if p.text.startswith('Dentro del radio de 100 metros'))
    leyenda = next(p for p in P if p.text.startswith('Tabla N.° 01:'))
    fuente = next(p for p in P if p.text.strip() == 'Fuente: Elaboración propia.')
    vacia = next(t for t in d.tables if 'CONDICIONES DE UBICACI' in t.rows[0].cells[0].text)
    modelo = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == 'GRUPO')
    e = _copia(sub, titulo._p, 'ANÁLISIS DE INCOMPATIBILIDADES DE UBICACIÓN EN EL PROYECTO')
    titulo._p.getparent().remove(titulo._p)
    a = _copia(cuerpo, e, ''); a = _copia(cuerpo, a, METODO); a = _copia(cuerpo, a, '')
    a = _copia(leyenda, a, 'Cuadro N° 4: Incompatibilidades de uso del terreno por cercanía - Hogar de Refugio Temporal (relación de la '
                           'R.V.M. N° 010-2022-MINEDU)')
    filas = [['%d. %s' % (k, x[0]), x[1], IR.dist(x), 'CUMPLE: ' + x[4]] for k, x in enumerate(IR.CUADRO, 1)]
    a = _tabla(modelo, a, ['INCOMPATIBILIDAD POR CERCANÍA', 'DISPOSITIVO LEGAL', 'DISTANCIA O CRITERIO', 'CONDICIÓN DEL PREDIO Y ESTADO'],
               filas, [2500, 3500, 1450, 2499])
    a = _copia(fuente, a, 'Fuente: Elaboración propia.')
    a = _copia(cuerpo, a, '')
    a = _copia(sub, a, 'CONDICIONES DE UBICACIÓN DEL HOGAR DE REFUGIO TEMPORAL')
    a = _copia(cuerpo, a, ''); a = _copia(cuerpo, a, CONDIC_INTRO); a = _copia(cuerpo, a, '')
    a = _copia(leyenda, a, 'Cuadro N° 5: Condiciones de ubicación del Hogar de Refugio Temporal (RNE y lineamientos para refugios)')
    filas = [['%d. %s' % (k, c[0]), c[1], c[2], 'CUMPLE'] for k, c in enumerate(IR.CONDICIONES, 1)]
    a = _tabla(modelo, a, ['CONDICIÓN DE UBICACIÓN', 'FUENTE', 'CÓMO LA CUMPLE EL PREDIO O EL PROYECTO', 'ESTADO'], filas,
               [2700, 2300, 3949, 1000])
    a = _copia(fuente, a, 'Fuente: Elaboración propia.')
    a = _copia(cuerpo, a, '')
    _copia(cuerpo, a, CONCLUSION % (len(IR.CUADRO), len(IR.CONDICIONES)))
    vacia._tbl.getparent().remove(vacia._tbl)
    CAMBIOS.append(('2.3 análisis vacío', '2.3.1 método + Cuadro N° 4 (18) y 2.3.2 Cuadro N° 5 (%d) con conclusión' % len(IR.CONDICIONES)))


def normativa(d):
    P = d.paragraphs
    ini = next(p for p in P if p.text.startswith('LINEAMIENTOS ARQUITECTÓNICOS PARA EL DISEÑO') and p.style.name == 'List Paragraph')
    rvm307 = next(p for p in P if p.text.startswith('R.V.M. N°307-2019'))
    rvm010 = next(p for p in P if p.text.startswith('R.V.M. N°010-2022'))
    reemplazar(rvm010, 'R.V.M. N°010-2022- MINEDU', 'R.V.M. N° 010-2022-MINEDU')
    a = rvm010._p
    for t in NORMAS_NUEVAS:
        a = _copia(rvm010, a, t)
    rvm307._p.getparent().remove(rvm307._p)
    CAMBIOS.append(('normativa', 'sin R.V.M. 307-2019 (racionalización magisterial); + Ley 30364, R.M. 150-2016-MIMP, R.V.M. 104 y 208, '
                                 'Ley 30102, guía bioclimática'))
    assert ini is not None


def sustento_pegado(d):
    """El sustento normativo pegado en el numeral 3.2 lleva su propia numeración (1. a 7.); pasa a 3.2.1 ... y las tablas siguen a
    la Tabla N° 1 de la memoria (accesos)."""
    P = d.paragraphs
    i0 = next(i for i, p in enumerate(P) if p.text.strip().startswith('SUSTENTO NORMATIVO – CÁLCULO DE AMBIENTES'))
    i1 = next(i for i, p in enumerate(P) if i > i0 and p.text.strip() == 'PROGRAMACIÓN ARQUITECTÓNICA')
    mapa = {'1.1': '3.2.1', '1.2': '3.2.2', '1.3': '3.2.3', '1.4': '3.2.4', '2.': '3.2.5', '2.1': '3.2.6', '3.': '3.2.7', '4.': '3.2.8',
            '5.': '3.2.9', '6.': '3.2.10', '7.': '3.2.11'}
    sub = []
    for p in P[i0 + 1:i1]:
        t = p.text.strip()
        if t == 'GENERALIDADES':
            p._p.getparent().remove(p._p); continue
        m = re.match(r'^(\d\.\d|\d\.)\s+(.*)$', t)
        if m and len(t) < 80 and not t.startswith('•'):
            num = m.group(1)
            nuevo = mapa.get(num) or ('3.2.9.' + num.split('.')[1] if num.startswith('5.') else None)
            if nuevo:
                txt = m.group(2)
                txt = txt[0] + txt[1:].lower() if txt.isupper() else txt
                poner(p.runs, ['%s  %s' % (nuevo, txt)]); sub.append(p)
        m = re.match(r'^Tabla N\.° (\d\d):(.*)$', t)
        if m:
            poner(p.runs, ['Tabla N.° %02d:%s' % (int(m.group(1)) + 1, m.group(2))])
        if t == 'SERVICIOS HIGIÉNICOS DE LA ZONA RESIDENCIAL':
            poner(p.runs, ['Servicios higiénicos de la zona residencial'])
    CAMBIOS.append(('sustento pegado (1. a 7.; Tabla N.° 01 a 11)', '3.2.1 a 3.2.11; Tabla N.° 02 a 12'))
    return sub


def programa(d):
    t = next(t for t in d.tables if 'PROGRAMACIÓN DE AREAS' in t.rows[0].cells[0].text)
    hechos = set()
    for r in t.rows:
        u = []
        for c in r.cells:
            if not u or c._tc is not u[-1]._tc:
                u.append(c)
        txt = [' '.join(c.text.split()) for c in u]
        if len(u) < 6:
            continue
        desc = txt[1].upper()
        if desc == 'SS.HH HAB.TRIPLE-TIPO III':
            celda(u[1], 'SS.HH HAB.TRIPLE-TIPO II')
        if desc in CITAS:
            idx, cita = CITAS[desc]
            norma = next(k for k in range(2, len(u)) if re.search(r'R\.V\.M|RNE|no aplica|Según diseño|NORMA', u[k].text)
                         and k >= 3)
            celda(u[norma], cita)
            if idx is not None:
                celda(u[norma - 1], idx)
            hechos.add(desc)
        for c in u:
            if 'Arquitectonico' in c.text or 'EXTERIRORES' in c.text:
                celda(c, ' '.join(c.text.split()).replace('Arquitectonico', 'Arquitectónico').replace('EXTERIRORES', 'EXTERIORES'))
    faltan = set(CITAS) - hechos
    if faltan:
        raise ValueError('no encontradas en el programa: %s' % sorted(faltan))
    CAMBIOS.append(('programa: citas normativas', '%d ambientes con la cita del sustento (áreas sin cambio)' % len(hechos)))


def listado(d):
    P = d.paragraphs
    tit = next(p for p in P if p.text.strip() == 'LISTADO DE PLANOS')
    viejos = [p for p in P[P.index(tit) + 1:] if p.style.name == 'List Paragraph']
    grupo_t = next(p for p in viejos if p.text.strip() == 'ARQUITECTURA GENERAL')._p
    item_t = next(p for p in viejos if p.text.startswith('PP-01'))._p
    ancla = viejos[0]._p.getprevious()
    grupos, _ = IAR.leer(); n = 0
    for g, filas in grupos:
        e = copy.deepcopy(grupo_t); ancla.addnext(e); ancla = e
        poner(P_(e, tit._parent).runs, ['%s (%d)' % (g, len(filas))])
        for f in filas:
            e = copy.deepcopy(item_t); ancla.addnext(e); ancla = e
            poner(P_(e, tit._parent).runs, [f['codigo'], '  ', '', f['descripcion']]); n += 1
    for p in viejos:
        p._p.getparent().remove(p._p)
    CAMBIOS.append(('listado de planos (U-01, PZ-01 "TOPOGRÁFICO", A1...)', '%d láminas del índice de arquitectura' % n))


def imagenes(d):
    P = d.paragraphs
    ley = [p for p in P if re.match(r'^\s*Imagen (N°)?\d+\s*:', p.text)]
    for k, p in enumerate(ley, 1):
        r = next(r for r in p.runs if re.search(r'Imagen|N°\d', r.text))
        txt = p.text
        nuevo = re.sub(r'Imagen\s*(N°)?\s*\d+', 'Imagen %d' % k, txt, count=1)
        if nuevo != txt:
            reemplazar(p, txt, nuevo)
    assert r is not None
    CAMBIOS.append(('imágenes (dos "4" y dos "6")', 'Imagen 1 a %d' % len(ley)))


def basura(d):
    for p in d.paragraphs:
        if p.text.strip() in ('co', '|'):
            for r in p.runs:
                if r.text.strip() in ('co', '|'):
                    r.text = ''
    CAMBIOS.append(('textos sueltos "co" y "|" junto a imágenes', 'eliminados'))


def titulos_refugio(d):
    out = CV_titulos(d)
    extra = [p for p in d.paragraphs if re.match(r'^3\.2\.\d+(\.\d+)?\s', p.text.strip())
             and not p.text.endswith('#') and '\t' not in p.text.strip()]
    res = []
    for p, lv in out:
        res.append((p, lv))
        if p.text.strip().startswith('SUSTENTO NORMATIVO'):
            res += [(q, 2) for q in extra if q.text.strip().count('.') <= 3 and not re.match(r'^3\.2\.9\.\d', q.text.strip())]
    return res


def numeros_refugio(tit):
    c = [0, 0, 0]; res = []
    for p, lv in tit:
        t = p.text.strip()
        if t == 'LISTADO DE PLANOS' or re.match(r'^3\.2\.\d', t):
            res.append(''); continue
        c[lv] += 1
        for k in range(lv + 1, 3):
            c[k] = 0
        res.append('%d.0.' % c[0] if lv == 0 else '.'.join(str(x) for x in c[:lv + 1]) + '.')
    return res


CV_titulos = CV._titulos
VARIOS = {'(septiembre,2026)'}   # textos que se repiten a propósito


def corregir():
    d = docx.Document(ENT)
    P = d.paragraphs
    for viejo, nuevo in TEXTOS:
        ps = [p for p in P if viejo in p.text]
        if len(ps) != 1 and not (ps and viejo in VARIOS):
            raise ValueError('%d párrafos con: %s' % (len(ps), viejo[:60]))
        for p in ps:
            reemplazar(p, viejo, nuevo)
    for p in P:                                                  # lista de ambientes
        if p.style.name == 'List Paragraph':
            for viejo, nuevo in LISTA:
                if p.text.strip().startswith(viejo):
                    reemplazar(p, viejo, nuevo)
        for viejo, nuevo in (('SEMAMHI', 'SENAMHI'),):
            if viejo in p.text:
                reemplazar(p, viejo, nuevo)
    for t in d.tables:                                           # "8,556.24. m²"
        for r in t.rows:
            for c in r.cells:
                for q in c.paragraphs:
                    if '8,556.24. m²' in q.text:
                        reemplazar(q, '8,556.24. m²', '8,556.24 m²')
    # "pendiente longitudinal de" y "1 %" estaban en dos párrafos: se unen
    p1 = next(p for p in P if p.text.strip().endswith('pendiente longitudinal de'))
    p2 = next(p for p in P if p.text.startswith('1 % y una pendiente transversal'))
    poner(p1.runs, [p1.text.rstrip() + ' ' + p2.text]); p2._p.getparent().remove(p2._p)
    imagenes(d); basura(d)
    analisis_ubicacion(d)
    normativa(d)
    sustento_pegado(d)
    programa(d)
    listado(d)
    CV._titulos, CV._numeros = titulos_refugio, numeros_refugio
    CV.preparar_indice(d)
    d.core_properties.author = ''
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    d.save(SALIDA)
    CV.paginar_indice()
    return SALIDA


def generar_nota():
    """Nota interna (no va al expediente): lo que está bien, lo corregido en la copia y lo que debe confirmar la arquitecta."""
    from docx.shared import Cm, Pt
    R = IR.R
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('MEMORIA DESCRIPTIVA DE ARQUITECTURA: REVISIÓN, CORRECCIONES Y PUNTOS PARA CONFIRMAR', 12.5, True, color=R.GRANATE,
          centro=True, despues=2)
    M.par('Hogar de Refugio Temporal - Mujeres Violentadas - CUI N.° 2675514  |  Documento de trabajo, no forma parte del expediente',
          9, color=R.GRIS, centro=True, despues=10)
    M.par('Se revisó la memoria descriptiva "culminada" párrafo por párrafo, con sus cuadros, imágenes, programa arquitectónico y '
          'listado de planos, y se cotejó con el índice de láminas de arquitectura, el análisis de incompatibilidades y el sustento '
          'normativo del programa. La copia corregida conserva el formato, las 32 imágenes y todas las áreas del programa; solo se '
          'corrigió lo que estaba errado o incompleto. Lo que necesita un dato o una decisión de la arquitecta queda en el punto 3.')

    def sec(t):
        M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4)

    sec('1. LO QUE ESTÁ BIEN')
    M.tabla('Comprobaciones sin observación', ['TEMA', 'COMPROBACIÓN'], [
        ['Saneamiento', 'Partida registral N.° 05003469 a favor del Gobierno Regional de San Martín (aclaración de dominio, asiento '
         'C00006, 2025); terreno matriz de 35,547.00 m² y 914.06 m.'],
        ['Poligonal del proyecto', '8,556.24 m² y 391.50 m, iguales en el texto y en el cuadro de datos técnicos.'],
        ['Programa: sumas por zona', 'Obras exteriores 1,235.45; administrativa 365.00; especializados 330.50; residencial 541.00; '
         'complementarios 1,843.40; generales 125.60: suman 4,440.95 m², igual al total del cuadro.'],
        ['Área del proyecto', '4,440.95 + 2,803.64 + 728.80 = 7,973.39 m², menor que la poligonal de 8,556.24 m².'],
        ['Índices del programa', 'Cada área de proyecto cumple o supera el mínimo que resulta de su índice y capacidad.'],
        ['Demanda', '55 víctimas como demanda efectiva, con el cuadro de casos por año y región.'],
        ['Terreno tipo II', 'La clasificación es correcta para el programa propuesto (solo estaba mal citada la fuente).'],
        ['Bioclimática', 'Zona climática, asoleamiento, radiación, iluminación, control solar y pluvial desarrollados con gráficos.'],
    ], [3.6, W - 3.6], tam=8.5, izq=(0, 1))

    sec('2. CORREGIDO EN LA COPIA')
    M.tabla('Correcciones', ['TEMA', 'EN LA MEMORIA CULMINADA', 'EN LA COPIA CORREGIDA'], [
        ['Texto de un colegio', '"Institución Educativa", "aulas", "aprendizaje y productividad", "infraestructura educativa"',
         'Hogar de Refugio Temporal, habitaciones, descanso y recuperación de las residentes'],
        ['Normas derogadas o ajenas', 'NTIE 001-2017 (arts. 23.3 y 27) y R.V.M. N.° 307-2019-MINEDU (racionalización magisterial)',
         'R.V.M. N.° 010-2022-MINEDU y guía de arquitectura bioclimática; se retira la R.V.M. 307-2019'],
        ['Normativa incompleta', 'Faltaban las normas del propio servicio',
         'Ley N.° 30364 y D.S. 009-2016-MIMP, R.M. N.° 150-2016-MIMP, R.V.M. 104 y 208-2019, Ley N.° 30102, guía bioclimática'],
        ['Titular del terreno', 'Dos párrafos decían que el predio pertenece a la Aldea Infantil Virgen del Pilar',
         'Gobierno Regional de San Martín (partida 05003469), como en el saneamiento'],
        ['Área en hectáreas', '0.085562 ha', '0.85562 ha (8,556.24 m²)'],
        ['Unidades y erratas de datos', '"8,556.24" sin unidad; "8,556.24. m²"; "1. %"', '8,556.24 m²; 1 % y 0.20 %'],
        ['Leyendas', 'Imágenes con número repetido (dos 4 y dos 6); Cuadro N°3 "Superposición de polígonos" sobre los datos de la '
         'poligonal', 'Imágenes 1 a 9 en orden; Cuadro N°3 "Datos técnicos de la poligonal resultante"'],
        ['Usos a menos de 100 m', 'Incluía "uso hogar de refugio temporal" (el propio proyecto)',
         'CAR Varones y CAR Mujeres colindantes, servicios de protección compatibles'],
        ['Análisis de ubicación', 'Solo el título del cuadro, sin contenido',
         '2.3.1 método y Cuadro N° 4 (18 incompatibilidades, CUMPLE); 2.3.2 Cuadro N° 5 (8 condiciones del hogar); conclusión'],
        ['Fuente del terreno tipo II', 'Los lineamientos para refugios', 'R.V.M. N.° 208-2019-MINEDU, Cuadro N° 2'],
        ['Drenaje pluvial', '"cunetas existentes en el exterior"', 'Cunetas y colector pluvial frontal del proyecto'],
        ['Control térmico', 'A.030 (hospedaje) como norma de confort', 'EM.110 y lineamientos para refugios'],
        ['Demanda', 'Cita al "Cuadro N°04" (número de otro cuadro)', 'Cuadro N°06'],
        ['Sustento normativo pegado', 'Título "GENERALIDADES" repetido; numerales 1 a 7 y Tabla N.° 01 repetida',
         'Numerales 3.2.1 a 3.2.11 dentro de 3.2; Tablas N.° 02 a 12'],
        ['Programa arquitectónico', 'Citas distintas a las del sustento del mismo documento (p. ej. sala de reuniones con A.080, '
         'comedor/SUM con 3.00 m²/pers., "SS.HH HAB.TRIPLE-TIPO III")',
         'Citas iguales al sustento (A.090 art. 11, R.V.M. 208 cuadro 22, etc.) y "TIPO II"; áreas sin cambio'],
        ['Listado de planos', 'U-01, PZ-01 rotulado "PLANO TOPOGRÁFICO" (repetido con PT-01), A1 a A9 y DT sin el índice',
         'Las %d láminas del índice de arquitectura, con sus códigos y títulos' % sum(len(f) for _, f in IAR.leer()[0])],
        ['Índice', 'No tenía', 'Índice automático de Word con número de página (se actualiza al abrir)'],
        ['Redacción', 'protege, relaciones, MATRÍZ, CATAS, INCOMPATILIDAD, ESTRATÉGIAS, SEMAMHI, ambientes públicas, Ss.hh, '
         'Almacen, maquinas, R.V.N', 'proteger, relacionadas, MATRIZ, CATASTRALES, INCOMPATIBILIDAD, ESTRATEGIAS, SENAMHI, '
         'ambientes públicos, SS.HH., Almacén, máquinas, R.V.M.'],
    ], [3.2, 6.4, W - 9.6], tam=8.5, izq=(0, 1, 2))

    sec('3. PARA CONFIRMAR CON LA ARQUITECTA')
    M.tabla('Puntos para confirmar', ['TEMA', 'QUÉ CONFIRMAR', 'POR QUÉ'], [
        ['Garita + SS.HH.', 'Figura en obras exteriores y en servicios generales (10.00 m² cada una). ¿Son dos garitas o una '
         'repetida?', 'Si es una sola, el total baja en 10.00 m² y debe corregirse en el programa y en el Excel.'],
        ['Estacionamientos y patio de descarga', 'Autos 6 × 12.50 = 75.00 pero total 244.00; motos 12 × 2.52 = 30.24 pero total 94.60; '
         'patio 3 × 16.00 = 48.00 pero total 256.00.', 'Si el total incluye pasillos y maniobra, conviene decirlo en la fila '
         '(p. ej. "incluye circulación vehicular"); así no parece un error de multiplicación.'],
        ['Circulación + 15 % muros', '2,267.80 × 1.15 = 2,607.97, pero el cuadro dice 2,803.64 (valor escrito, sin fórmula).',
         'Indicar cómo se obtuvo o reemplazar por la fórmula; cambia el área del proyecto (7,973.39 m²).'],
        ['"Área construida (interior)"', 'El valor 4,440.95 incluye la zona de obras exteriores (estacionamientos, PTAR, cunetas, '
         'cerco).', 'Cambiar el rótulo a "áreas por zonas" o separar las obras exteriores.'],
        ['SS.HH. de habitaciones dobles', 'Hay 10 habitaciones dobles (3 tipo I + 7 tipo II) y 7 SS.HH. "compartidos".',
         'Confirmar cuáles comparten el baño; los lineamientos admiten un baño completo compartido por no más de tres habitaciones.'],
        ['Norma MINEDU vigente', 'Se citó la R.V.M. N.° 010-2022-MINEDU; la memoria del CAR Varones cita la R.V.M. N.° 068-2025-MINEDU.',
         'Usar la misma en los tres proyectos del predio Los Mangos.'],
        ['Nombre de la vía', '"Carretera Polvoraico" en el saneamiento.', 'La memoria del CAR Varones dice "Polvorico"; usar el mismo nombre en los tres proyectos.'],
        ['Zonificación', 'Certificado de zonificación y vías o de parámetros urbanísticos de la Municipalidad Distrital de Morales.',
         'Acredita la compatibilidad de uso del Cuadro N° 5 (A.090 art. 3).'],
        ['Servicios de apoyo', 'Nombre y distancia de la comisaría, el Centro Emergencia Mujer, la fiscalía y el centro de salud '
         'más cercanos.', 'Con nombres y distancias el CUMPLE del Cuadro N° 5 queda mejor sustentado.'],
        ['Fecha de imágenes', 'Las imágenes satelitales dicen "septiembre, 2026".', 'Confirmar la fecha real de la captura.'],
    ], [3.2, 7.4, W - 10.6], tam=8.5, izq=(0, 1, 2))
    d.core_properties.author = ''
    d.core_properties.title = 'Revisión de la memoria descriptiva de arquitectura - Hogar de Refugio Temporal'
    os.makedirs(os.path.dirname(SALIDA_N), exist_ok=True)
    d.save(SALIDA_N)
    return SALIDA_N


if __name__ == '__main__':
    print(corregir())
    print(generar_nota())
    for a, b in CAMBIOS:
        print('  %s -> %s' % (a[:70], b[:80]))
