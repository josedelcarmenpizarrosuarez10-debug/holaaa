# -*- coding: utf-8 -*-
"""Datos comunes de evacuación y señalización del Hogar de Refugio Temporal "Mujeres Violentadas" (CUI 2675514): una sola
fuente para el metrado, la memoria descriptiva, las especificaciones técnicas y la nota de observaciones.

Cantidades: del plano de señalización de la arquitecta (lee_plano_sr.py -> entregables/_calc/senalizacion_plano.json):
  - señal con bloque en la planta: la cantidad dibujada dentro del terreno, por edificación (E-01 a E-09 y exteriores);
    las copias sueltas fuera del terreno no se metran;
  - unidad de luces de emergencia: los bloques "LE" y "A$C7A2016C2" son el mismo símbolo ("LE" contiene a "A$C7A2016C2",
    que es el símbolo de la leyenda); se suman los dos;
  - atención riesgo eléctrico: en el DXF es una imagen pegada (OLE) sin ubicación; se metran las 39 imágenes de igual tamaño
    (0.90 x 1.36 m) del espacio modelo (la leyenda indica 38); su reparto por edificación es el del metrado de evacuación de
    arquitectura (hoja HD), que suma lo mismo; si no sumara igual, van a la columna "sin ubicar".
Datos que el plano no define y se adoptan aquí (nota de observaciones): extintor PQS tipo ABC de 9 kg (memoria de evacuación
de arquitectura: 33 extintores PQS de 9 kg), señal de extintor sobre el equipo a 1.80 m, punto ecológico a 1.80 m.
"""
import json
import math
import os
from collections import OrderedDict

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
PLANO = os.path.join(RAIZ, 'entregables/_calc/senalizacion_plano.json')
EVAC = os.path.join(RAIZ, 'entregables/_calc/evacuacion_plano.json')      # plano de evacuación (lee_plano_ev.py)
# metrado de evacuación de arquitectura (Excel de la usuaria): solo para el reparto por edificación de la señal de riesgo eléctrico
METRADO_ARQ = os.path.join(RAIZ, 'insumos/solange/senalizacion/METRADOS_EVACUACION_REFUGIO_MUJERES_ok.xlsx')
SAL = os.path.join(RAIZ, 'entregables')

PROYECTO = ('CREACIÓN DEL SERVICIO DE PROTECCIÓN A MUJERES VÍCTIMAS DE VIOLENCIA EN SITUACIÓN DE RIESGO DE FEMINICIDIO O PELIGRE '
            'SU INTEGRIDAD Y/O SALUD FÍSICA Y/O MENTAL EN HOGAR DE REFUGIO TEMPORAL DISTRITO DE MORALES DE LA PROVINCIA DE SAN '
            'MARTÍN DEL DEPARTAMENTO DE SAN MARTÍN')
CUI = '2675514'
CORTO = 'HOGAR DE REFUGIO TEMPORAL'
ENTIDAD = 'GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO'
UBICACION = 'Distrito de Morales, provincia y departamento de San Martín'
ESPECIALIDAD = 'ARQUITECTURA - EVACUACIÓN Y SEÑALIZACIÓN'
FECHA = 'OCTUBRE 2026'                         # la del rótulo del plano de señalización (Octubre-2026)
# códigos del índice de planos de arquitectura del hogar (grupo VI, evacuación y señalización)
LAMINA_S = ('SE-S-02', 'Planta general de señalización', '1/150')
LAMINA_E = ('SE-E-01', 'Planta general de evacuación', '1/150')
LAMINAS = [LAMINA_E, LAMINA_S]
LS = LAMINA_S[0]

EDIFICACIONES = OrderedDict([
    ('E-01', 'Administración'), ('E-02', 'Almacén de donaciones y tópico'), ('E-03', 'Mantenimiento'),
    ('E-04', 'Talleres'), ('E-05', 'SUM-comedor, cocina, lavandería, módulo de enseñanza y módulo infantil'),
    ('E-06', 'Habitaciones (1)'), ('E-07', 'Habitaciones (2)'), ('E-08', 'Habitaciones (3)'),
    ('E-09', 'Vestidores y almacén de implementos deportivos'), ('EXT', 'Exteriores')])

# población: residentes = plazas de las 23 habitaciones (memoria de arquitectura, tabla de plazas; A.030 art. 5: ocupantes =
# camas); personal y usuarios temporales = memoria de evacuación de arquitectura del hogar
POBLACION = dict(residentes=113, personal=62, temporales=130)
POBLACION['permanentes'] = POBLACION['residentes'] + POBLACION['personal']
POBLACION['maxima'] = POBLACION['permanentes'] + POBLACION['temporales']
PLAZAS = [('Habitaciones simples - tipo I', 3, 3), ('Habitaciones simples - tipo II', 3, 3), ('Habitaciones doble - tipo I', 3, 5),
          ('Habitaciones doble - tipo II', 7, 5), ('Habitaciones triple - tipo I', 3, 7), ('Habitaciones triple - tipo II', 3, 7),
          ('Habitación especial', 1, 3)]
D_MAX = 45.0                                   # m: recorrido máximo hasta un lugar seguro sin rociadores (RNE A.130)

# materiales de las señales (NTP 399.010-1; RNE A.120 para el braille)
FOTOLUM = 'Lámina de PVC celtex de 3 mm con vinil fotoluminiscente impreso y placa de acrílico con texto en braille'
FOTOLUM_SB = 'Lámina de PVC celtex de 3 mm con vinil fotoluminiscente impreso'
VINIL = 'Lámina de PVC celtex de 3 mm con vinil adhesivo impreso y placa de acrílico con texto en braille'

GRUPOS = [('01', 'SEÑALES DE EVACUACIÓN Y EMERGENCIA'), ('02', 'SEÑALES DE EQUIPOS CONTRA INCENDIO'),
          ('03', 'SEÑALES DE ADVERTENCIA'), ('04', 'SEÑALES INFORMATIVAS'),
          ('05', 'EQUIPOS DE PROTECCIÓN CONTRA INCENDIO Y EMERGENCIA')]
# clave = inicio de la descripción de la fila de la leyenda; bloques de la planta que la representan (None = imagen pegada)
#   clave, grupo, descripción de la partida, und, dimensiones, material, altura adoptada, función / ubicación, bloques
CATALOGO = [
    ('SALIDA CON', '01', 'SEÑAL DE SALIDA CON BRAILLE', 'und', '0.30 x 0.20 m', FOTOLUM, '2.10 m',
     'Indica la puerta de salida de cada ambiente y de cada edificación hacia la circulación exterior; sobre el dintel.', ['salida']),
    ('DIRECCIÓN DE EVACUACIÓN', '01', 'SEÑAL DE DIRECCIÓN DE EVACUACIÓN', 'und', '0.20 x 0.30 m', FOTOLUM_SB, '2.10 m',
     'Indica con una flecha el sentido de la ruta de evacuación en pasadizos y circulaciones, hasta la salida.', ['SALIDA 2']),
    ('ZONA SEGURA EN', '01', 'SEÑAL DE ZONA SEGURA EN CASO DE SISMO CON BRAILLE', 'und', '0.20 x 0.30 m', FOTOLUM, '1.80 m',
     'Identifica los lugares de protección dentro de los ambientes (columnas y muros estructurales) durante un sismo.', ['A$C3C91529B']),
    ('LETRERO DE AFORO', '01', 'LETRERO DE AFORO CON BRAILLE', 'und', '0.30 x 0.20 m', VINIL, '2.10 m',
     'Indica la capacidad máxima de personas del ambiente; junto al ingreso, en el lado de la cerradura.', ['A$C6B5B2E50']),
    ('ZONA DE REUNI', '01', 'PINTURA DE ZONA DE EVACUACIÓN EN CASO DE EMERGENCIA EN PISO (CÍRCULO DE 4.30 m, LÍNEA DE 0.15 m)', 'm2',
     'anillo de 4.30 m de diámetro exterior y 0.15 m de ancho', 'Pintura de tránsito color amarillo sobre piso de concreto', 'en piso',
     'Zona exterior de concentración de los ocupantes al final de las rutas de evacuación (zonas 01 a 13 de la lámina SE-E-01).', 'ZONAS'),
    ('SEÑAL DE UBICACIÓN DE EXTINTOR PQS', '02', 'SEÑAL DE UBICACIÓN DE EXTINTOR PQS CON BRAILLE', 'und', '0.20 x 0.30 m', FOTOLUM,
     '1.80 m', 'Identifica la ubicación del extintor de polvo químico seco; sobre el equipo.', ['extintor1']),
    ('ATENCIÓN RIESGO', '03', 'SEÑAL DE ATENCIÓN RIESGO ELÉCTRICO CON BRAILLE', 'und', '0.20 x 0.30 m', VINIL, '1.10 m',
     'Advierte la presencia de tableros y equipos eléctricos; en la puerta o el frente del tablero.', None),
    ('SEÑALES DE SERVICIOS HIGIENICOS PHD', '04', 'SEÑAL DE SERVICIOS HIGIÉNICOS PARA PERSONAS CON DISCAPACIDAD CON BRAILLE', 'und',
     '0.20 x 0.30 m', VINIL, '1.80 m', 'Identifica el servicio higiénico accesible.', ['sdsadsa']),
    ('SEÑALES DE SERVICIOS HIGIENICOS MIXTO', '04', 'SEÑAL DE SERVICIOS HIGIÉNICOS MIXTO CON BRAILLE', 'und', '0.20 x 0.30 m', VINIL,
     '1.80 m', 'Identifica el servicio higiénico de uso mixto (habitaciones, tópico y vestidores).', ['*U2820']),
    ('SEÑALES DE SERVICIOS HIGIENICOS MUJER', '04', 'SEÑAL DE SERVICIOS HIGIÉNICOS MUJER CON BRAILLE', 'und', '0.20 x 0.30 m', VINIL,
     '1.80 m', 'Identifica el servicio higiénico de mujeres.', ['sshhmujeres']),
    ('SEÑALES DE SERVICIOS HIGIENICOS VAR', '04', 'SEÑAL DE SERVICIOS HIGIÉNICOS VARÓN CON BRAILLE', 'und', '0.20 x 0.30 m', VINIL,
     '1.80 m', 'Identifica el servicio higiénico de varones.', ['sshhhombres']),
    ('PUNTO ECOL', '04', 'SEÑAL DE PUNTO ECOLÓGICO CON BRAILLE', 'und', '0.30 x 0.20 m', VINIL, '1.80 m',
     'Identifica el punto de acopio selectivo de residuos sólidos (cuatro contenedores de colores).', ['punto.ecologico']),
    ('UNIDAD DE LUCES', '05', 'UNIDAD DE LUCES DE EMERGENCIA', 'und', 'equipo de dos faros LED, autonomía mínima 90 min',
     'Equipo autónomo con batería recargable, 220 V', '2.10 m',
     'Ilumina las salidas, las circulaciones y los ambientes al cortarse la energía.', ['LE', 'A$C7A2016C2']),
    ('DETECTOR DE HUMO', '05', 'DETECTOR DE HUMO', 'und', 'detector fotoeléctrico de techo', 'Equipo fotoeléctrico con base de montaje',
     'techo', 'Detecta el humo en los ambientes en su etapa inicial.', ['sensorhumo']),
]
# equipo que la leyenda solo señala: un extintor por cada señal de ubicación de extintor de la planta
EQUIPOS = [
    ('SEÑAL DE UBICACIÓN DE EXTINTOR PQS', 'EXTINTOR DE POLVO QUÍMICO SECO (PQS) TIPO ABC DE 9 kg', 'und',
     'extintor portátil de 9 kg, agente PQS multipropósito ABC',
     'Cilindro de acero con manómetro, manguera, soporte de pared y tarjeta de control', '1.50 m (parte superior)',
     'Combate fuegos de clase A, B y C en las nueve edificaciones.', ['extintor1']),
]


def plano():
    return json.load(open(PLANO, encoding='utf-8'))


def edificaciones(P=None):
    return list(EDIFICACIONES.items())


def _fila(P, clave):
    return next(f for f in P['leyenda'] if f['descripcion'].upper().startswith(clave.upper()))


def ole_senal(P=None):
    """Imágenes pegadas del espacio modelo (propietario 1F) con el tamaño de la señal de riesgo eléctrico (0.90 x 1.36 m)."""
    P = P or plano()
    return [o for o in P['ole'] if o['propietario'] == '1F' and abs((o['ancho'] or 0) - 0.904) < 0.01 and abs((o['alto'] or 0) - 1.356) < 0.01]


def riesgo_por_edificacion():
    """Reparto por edificación de la señal de riesgo eléctrico en la hoja HD del metrado de arquitectura (partida 03.14.01.04)."""
    import openpyxl
    import re
    if not os.path.exists(METRADO_ARQ):
        return {}
    ws = openpyxl.load_workbook(METRADO_ARQ, data_only=True)['HD']; out = {}; dentro_p = False
    for r in range(14, ws.max_row + 1):
        b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
        if b:
            dentro_p = 'RIESGO ELECTRICO' in str(c).upper()
        elif dentro_p and isinstance(c, str) and c.upper().startswith('EDIFICACION'):
            out['E-%02d' % int(re.search(r'\d+', c).group())] = (ws.cell(r, 5).value or 1) * (ws.cell(r, 6).value or 0)
    return out


def partidas(P=None):
    """Partidas en orden, con cantidad por edificación (o sin ubicar), cantidad de la leyenda, diferencia y fuente."""
    P = P or plano()
    eds = list(EDIFICACIONES)
    out = []
    for g, titulo in GRUPOS:
        out.append(dict(item=g, desc=titulo, und=None))
        filas = [c for c in CATALOGO if c[1] == g]
        if g == '05':
            filas = [(e[0], g) + e[1:] for e in EQUIPOS] + filas
        for k, (clave, _, desc, und, dim, mat, alt, uso, bloques) in enumerate(filas, 1):
            f = _fila(P, clave)
            reg = dict(item='%s.%02d' % (g, k), desc=desc, und=und, dim=dim, material=mat, altura=alt, uso=uso, clave=clave,
                       leyenda=f['cantidad'] or 0, altura_leyenda=f['altura'], desc_leyenda=f['descripcion'], bloques=bloques)
            if g == '05' and clave.startswith('SEÑAL DE UBICACIÓN'):
                reg['altura_leyenda'] = None                 # la altura de la leyenda es la de la señal, no la del equipo
            if bloques == 'ZONAS':
                zs = zonas(); a = area_zona()
                reg.update(por={'EXT': round(len(zs) * a, 2)}, sin_ubicar=0, total=round(len(zs) * a, 2), n_zonas=len(zs), area_zona=a,
                           leyenda=0, fuente='Lámina %s: %d zonas de evacuación de 4.30 m de diámetro con anillo de 0.15 m '
                                             '(%d x %.4f m²); leyenda de %s: 00' % (LAMINA_E[0], len(zs), len(zs), a, LS))
            elif bloques is None:
                n = len(ole_senal(P)); rp = riesgo_por_edificacion()
                if rp and sum(rp.values()) == n:                 # reparto del metrado de arquitectura con el mismo total
                    reg.update(por={e: rp.get(e, 0) for e in eds}, sin_ubicar=0, total=n,
                               fuente='Imagen pegada en el DXF (sin ubicación): %d imágenes en el espacio modelo; reparto por '
                                      'edificación del metrado de evacuación de arquitectura; leyenda: %d' % (n, f['cantidad']))
                else:
                    reg.update(por={}, sin_ubicar=n, total=n,
                               fuente='Imagen pegada en el DXF (sin ubicación): %d imágenes en el espacio modelo; leyenda: %d' % (n, f['cantidad']))
            else:
                por = {e: sum(P['bloques'][b]['por_edificacion'].get(e, 0) for b in bloques) for e in eds}
                tot = sum(P['bloques'][b]['cantidad'] for b in bloques)
                fuera = sum(len(P['bloques'][b]['fuera_del_terreno']) for b in bloques)
                txt = 'Planta de %s: símbolos dibujados' % LS
                if len(bloques) > 1:
                    txt += ' (%s)' % ' + '.join('%d "%s"' % (P['bloques'][b]['cantidad'], b) for b in bloques)
                if fuera:
                    txt += '; %d copia fuera del terreno no se metra' % fuera
                reg.update(por=por, sin_ubicar=0, total=tot, fuente=txt + '; leyenda: %d' % (f['cantidad'] or 0))
            out.append(reg)
    return out


def evac():
    return json.load(open(EVAC, encoding='utf-8'))


def zonas(E=None):
    """Zonas de evacuación del plano SE-E-01, en el orden de su número (el número 10 está repetido en el plano)."""
    E = E or evac()
    zs = [dict(z) for z in E['zonas']]
    vistos = set(); n_max = max(int(z['numero']) for z in zs)
    for z in zs:                                    # el número repetido (10) pasa al siguiente libre: 13
        if z['numero'] in vistos:
            n_max += 1; z['numero_plano'] = z['numero']; z['numero'] = '%02d' % n_max
        vistos.add(z['numero'])
    return sorted(zs, key=lambda z: int(z['numero']))


def area_zona(E=None):
    """Área pintada del anillo: radios exterior e interior de los círculos del plano (2.15 y 2.00 m)."""
    E = E or evac()
    rs = next(z['radios'] for z in E['zonas'] if z['radios'])
    return math.pi * (max(rs) ** 2 - min(rs) ** 2)


def radios(E=None):
    E = E or evac()
    rs = next(z['radios'] for z in E['zonas'] if z['radios'])
    return max(rs), min(rs)


V_EVAC = (0.60, 1.00)                         # m/s: grupos con niñas y niños pequeños o movilidad reducida / adultos


def rutas(E=None):
    E = E or evac()
    return [dict(n=r['n'], L=r['L'], ed=r['edificacion'], t1=r['L'] / V_EVAC[0], t2=r['L'] / V_EVAC[1], cumple=r['L'] <= D_MAX)
            for r in E['rutas']]


def leyenda_corregida(P=None):
    """Filas de la leyenda con cantidad distinta de lo dibujado: (descripción, leyenda, planta, motivo)."""
    P = P or plano(); out = []
    for p in partidas(P):
        if not p['und'] or p['desc'].startswith('EXTINTOR DE'):
            continue
        if p['leyenda'] != p['total']:
            out.append((p['desc_leyenda'], p['leyenda'], p['total'], p['fuente']))
    return out


def comprobacion(P=None):
    """Controles de la lectura: suma por edificación = total; sin duplicados; luces = 'LE' + símbolo de la leyenda."""
    P = P or plano(); err = []
    for p in partidas(P):
        if p['und'] and sum(p['por'].values()) + p['sin_ubicar'] != p['total']:
            err.append('%s: suma por edificación distinta del total' % p['desc'])
    for b, v in P['bloques'].items():
        if v['duplicados']:
            err.append('%s: %d símbolos superpuestos' % (b, v['duplicados']))
    if P['anidados'].get('LE', [None, []])[1] != ['A$C7A2016C2']:
        err.append('el bloque LE ya no contiene al símbolo de luces de la leyenda')
    if P['anidados'].get('extintor1', [None, []])[1] != ['*U2804']:
        err.append('el bloque extintor1 ya no contiene al símbolo de extintor PQS de la leyenda')
    return err


if __name__ == '__main__':
    P = plano()
    for p in partidas(P):
        print(p['item'], p['desc'], p.get('total', ''), p.get('leyenda', ''), p.get('fuente', ''))
    print('leyenda a corregir:', leyenda_corregida(P))
    print('comprobación:', comprobacion(P) or 'conforme')
    print(POBLACION)
