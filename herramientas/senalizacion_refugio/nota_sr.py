# -*- coding: utf-8 -*-
"""Nota interna (no va al expediente) de evacuación y señalización del Hogar de Refugio Temporal (CUI 2675514): lectura del
plano de la arquitecta, cotejo leyenda / planta / metrado, criterio de cada cantidad adoptada, observaciones que se deben
corregir en el plano y puntos para confirmar.

Uso: python3 herramientas/senalizacion_refugio/nota_sr.py
Salida: entregables/revision_interna/REVISION_EVACUACION_Y_SENALIZACION_HOGAR_REFUGIO.docx
"""
import math
import os
import sys
from collections import Counter

import docx
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); HER = os.path.dirname(AQUI)
for d in (AQUI, HER, os.path.join(HER, 'varones', 'senalizacion')):
    sys.path.insert(0, d)
import datos_sr as D  # noqa: E402
import incompatibilidades_refugio as IR  # noqa: E402  (formato de las notas internas)
import memoria_se as MS  # noqa: E402

SALIDA = os.path.join(D.SAL, 'revision_interna/REVISION_EVACUACION_Y_SENALIZACION_HOGAR_REFUGIO.docx')
FIG = os.path.join(D.SAL, 'figuras_senalizacion')
LS = D.LS


def observaciones(P):
    rot = P['rotulo']
    proy = next((t for t in rot if 'I.E. 292' in t), '')
    cui_otro = proy.split('CUI')[-1].strip() if 'CUI' in proy else ''
    pts = {p['clave']: p for p in D.partidas(P) if p['und']}
    ole = D.ole_senal(P)
    dt = P['bloques']['sensorhumo']['fuera_del_terreno']
    ev_z = D.zonas()
    return [
        ('Rótulo', 'El rótulo de la lámina es de otro proyecto: “Mejoramiento y ampliación del servicio de educación inicial y '
         'primaria en la I.E. 292 ... distrito de La Banda de Shilcayo” (CUI %s), lámina “S-01” y “01-02”, distrito Banda de '
         'Shilcayo, con el coordinador y la especialista de ese proyecto.' % cui_otro,
         'Poner el nombre del proyecto del hogar con el CUI 2675514, distrito de Morales, y el código de la lámina del índice de '
         'planos de arquitectura: %s %s.' % (LS, D.LAMINA_S[1].upper())),
        ('Código de lámina', 'La memoria de evacuación de arquitectura cita “lámina SE-S-01” para evacuación y “lámina SE-E-02” para '
         'señalización; el índice de planos de arquitectura dice %s %s y %s %s.' % (D.LAMINA_E[0], D.LAMINA_E[1], LS, D.LAMINA_S[1]),
         'Usar en el rótulo y en la memoria los códigos del índice (%s evacuación, %s señalización).' % (D.LAMINA_E[0], LS)),
        ('Zona segura en caso de sismo', 'La leyenda indica 210; en la planta hay %d señales.' % pts['ZONA SEGURA EN']['total'],
         'Corregir la leyenda a %d.' % pts['ZONA SEGURA EN']['total']),
        ('Extintor PQS', 'La leyenda indica 44 señales de ubicación de extintor PQS; en la planta hay %d (bloque “extintor1”, que '
         'contiene al símbolo de la leyenda). La memoria de evacuación de arquitectura indica 33 extintores PQS de 9 kg.' % pts[
             'SEÑAL DE UBICACIÓN DE EXTINTOR PQS']['total'],
         'Corregir la leyenda a 33 e indicar en ella el equipo: extintor PQS tipo ABC de 9 kg.'),
        ('Altura de la señal de extintor', 'La leyenda da 1.20 m a las señales de extintor: el extintor de 9 kg se cuelga con su parte '
         'superior a no más de 1.50 m (NTP 350.043-1) y la señal va encima del equipo, de modo que a 1.20 m quedaría tapada.',
         'Corregir la altura a 1.80 m (borde inferior de la señal), la adoptada en el metrado y las especificaciones.'),
        ('Unidad de luces de emergencia', 'La leyenda indica 77, que son los bloques “LE”; además hay 38 bloques “A$C7A2016C2” con el '
         'mismo dibujo (es el símbolo de la leyenda y el bloque “LE” lo contiene), casi todos en las habitaciones (%s). Total '
         'dibujado: %d.' % (', '.join('%s: %d' % kv for kv in sorted(P['bloques']['A$C7A2016C2']['por_edificacion'].items())),
                            pts['UNIDAD DE LUCES']['total']),
         'Corregir la leyenda a %d (se metran todas las dibujadas) y usar un solo nombre de bloque.' % pts['UNIDAD DE LUCES']['total']),
        ('Detector de humo', 'La leyenda indica 68; dentro del terreno hay %d. La copia restante está suelta fuera del dibujo, en '
         'E %.1f N %.1f (coordenadas del doble de las del terreno), y no se metra.' % (pts['DETECTOR DE HUMO']['total'], dt[0][0], dt[0][1]),
         'Borrar la copia suelta y corregir la leyenda a %d.' % pts['DETECTOR DE HUMO']['total']),
        ('Punto ecológico', 'La leyenda indica 00, pero en los exteriores hay %d puntos ecológicos dibujados (cuatro contenedores cada '
         'uno).' % pts['PUNTO ECOL']['total'], 'Corregir la leyenda a %d.' % pts['PUNTO ECOL']['total']),
        ('Atención riesgo eléctrico', 'La señal no es un bloque: es una imagen pegada (OLE). El DXF guarda solo su tamaño, no su '
         'posición, por lo que no se puede ubicar ni asignar a cada edificación. En el espacio modelo hay %d imágenes iguales '
         '(%.2f x %.2f m); la leyenda indica 38.' % (len(ole), ole[0]['ancho'], ole[0]['alto']),
         'Reemplazar las imágenes por un bloque de señal (así se cuentan y ubican) y confirmar la cantidad: el metrado toma las %d '
         'imágenes; si una es copia sobrante, bajar a 38 en el metrado y en la leyenda.' % len(ole)),
        ('Letreros de aforo', 'Los %d letreros de aforo de la planta tienen el texto “__ PERSONAS”, sin el número de personas.' % pts[
            'LETRERO DE AFORO']['total'], 'Indicar en la lámina el aforo de cada ambiente (cálculo de aforo de arquitectura) para '
         'imprimir cada letrero con su número.'),
        ('Luces de emergencia “con braille”', 'La leyenda describe la unidad de luces de emergencia “con braille”; es un equipo de '
         'iluminación y no lleva texto en braille.', 'Quitar “con braille” de esa fila (la partida del metrado es “unidad de luces de '
         'emergencia”).'),
        ('Detección y alarma', 'Se dibujan %d detectores de humo, pero las filas de central de alarma, panel de alarma, pulsador '
         '(alarma contra incendios) y avisador sonoro están en 00; la memoria de arquitectura menciona “detección y alarma contra '
         'incendio”.' % pts['DETECTOR DE HUMO']['total'],
         'Definir con instalaciones eléctricas si los detectores son autónomos con alarma propia (como se especificó) o van a una '
         'central con pulsadores y avisadores; en ese caso dibujarlos y metrarlos.'),
        ('Bloques con otro nombre', 'Varios símbolos de la planta no son el bloque de la leyenda: salida (“salida” / “swkiswkis”), '
         'SS.HH. PHD (“sdsadsa” / “15”), extintor (“extintor1”, que contiene a “*U2804”), luces (“LE” y “A$C7A2016C2”), aforo '
         '(“A$C6B5B2E50”, que contiene a “dtjrtj”).', 'Usar en la planta el mismo bloque de la leyenda, para que la selección rápida '
         'de AutoCAD dé las cantidades de la leyenda.'),
        ('Ortografía', 'Leyenda: “ALARMA CONTRA INCEDIOS”; ambientes: “TALLER DE RESPOSTERÍA”, “DISPENSA”.',
         'Corregir a “INCENDIOS”, “REPOSTERÍA” y “DESPENSA”.'),
        ('Rótulo de la lámina de evacuación', 'El plano de evacuación tiene el mismo rótulo de otro proyecto (I.E. 292, Banda de '
         'Shilcayo), con “SEGURIDAD: EVACUACIÓN”, lámina “S-01” y “02-02”.',
         'Poner el proyecto del hogar, CUI 2675514, Morales, y el código del índice: %s %s.' % (D.LAMINA_E[0], D.LAMINA_E[1].upper())),
        ('Numeración de zonas', 'Hay %d zonas de evacuación, pero el número 10 está dos veces (una al norte, junto a E-08, y otra '
         'junto a E-06); los rótulos van del 01 al 12.' % len(ev_z),
         'Numerar de 01 a %d: en el metrado y la memoria la zona 10 junto a E-06 es la zona %s.' % (len(ev_z), next(z['numero'] for z in ev_z if z.get('numero_plano')))),
        ('Aforo de las zonas', 'Las %d zonas tienen el mismo tamaño (anillo de 4.30 m; %.2f m² libres dentro), pero el plano les da '
         'aforos distintos: %s. La zona 05, con 40 personas, deja %.2f m² por persona. En total reciben %d personas, menos que la '
         'capacidad máxima del hogar (%d) aunque más que la población permanente (%d).' % (
             len(ev_z), math.pi * D.radios()[1] ** 2, ', '.join('%d zonas de %d' % (n, a) for a, n in sorted(Counter(z['aforo'] for z in ev_z).items())),
             math.pi * D.radios()[1] ** 2 / 40, sum(z['aforo'] for z in ev_z), D.POBLACION['maxima'], D.POBLACION['permanentes']),
         'Unificar el criterio del aforo de las zonas (área por persona) y verificar que el total cubra la capacidad máxima, '
         'ampliando o agregando zonas si hace falta.'),
        ('Leyenda de evacuación', 'La lámina de evacuación no tiene leyenda ni instrucciones en caso de evacuación: la presentación '
         'solo tiene el rótulo.', 'Agregar la leyenda (ruta, punto de partida, zona de evacuación y aforo) y las instrucciones.'),
        ('Zona de reuniones en la leyenda de señalización', 'La fila “ZONA DE REUNIONES CON BRAILLE” de la leyenda de %s está en 00, '
         'pero la lámina %s tiene %d zonas pintadas en el piso (que no llevan braille).' % (LS, D.LAMINA_E[0], len(ev_z)),
         'Poner %d en esa fila (sin “con braille”) o retirarla y dejar las zonas solo en la lámina de evacuación.' % len(ev_z)),
        ('Población', 'La memoria de evacuación de arquitectura indica 41 residentes; la tabla de plazas de la memoria descriptiva de '
         'arquitectura suma 113 plazas en 23 habitaciones (ocupantes para la evacuación = camas, A.030 art. 5).',
         'Usar 113 residentes (capacidad máxima 305 personas con 62 de personal y 130 temporales), como en la memoria de señalización.'),
    ]


def generar():
    P = D.plano(); pts = D.partidas(P)
    R = IR.R
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('EVACUACIÓN Y SEÑALIZACIÓN: LECTURA DE LOS PLANOS, COTEJO Y OBSERVACIONES', 12.5, True, color=R.GRANATE,
          centro=True, despues=2)
    M.par('Hogar de Refugio Temporal - Mujeres Violentadas - CUI N.° %s  |  Documento de trabajo, no forma parte del expediente' % D.CUI,
          9, color=R.GRIS, centro=True, despues=10)
    n_ole = len(D.ole_senal(P))
    M.par('Se leyó el DXF del plano de señalización de la arquitecta (292 MB): la leyenda de 27 filas, cada símbolo de señal dibujado '
          'en la planta (posición = centro de su dibujo sin textos) y la edificación donde cae (contorno del bloque de cada '
          'edificación de la planta general, con 0.60 m de margen; donde los contornos se traslapan en los aleros, la que lo '
          'contiene con más holgura), los símbolos fuera del terreno, las imágenes pegadas (%d de la señal de riesgo '
          'eléctrico) y el rótulo. Se revisó que no hubiera símbolos superpuestos (ninguno a menos de 0.10 m de otro igual) y que la '
          'presentación A1 muestre la planta a escala 1/150. El metrado toma lo dibujado en la planta; donde la leyenda no coincide, '
          'se corrige la leyenda. Del plano de evacuación (125 MB, misma planta y coordenadas) se leyeron las %d rutas con su longitud y '
          'su punto de partida, las %d zonas de evacuación con su aforo y el rótulo.' % (n_ole, len(D.rutas()), len(D.zonas())))

    def sec(t):
        M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4)

    sec('1. COTEJO DE LA LEYENDA CON LA PLANTA Y EL METRADO')
    filas = []
    for p in pts:
        if not p['und']:
            continue
        dif = p['total'] - p['leyenda']; m2 = p['und'] == 'm2'
        filas.append([p['item'], p['desc'].replace(' CON BRAILLE', '').capitalize().replace('pqs', 'PQS').replace('(pqs)', '(PQS)')
                      .replace(' abc', ' ABC').replace('4.30 m, línea de 0.15 m', '4.30 m y línea de 0.15 m, m²'), p['leyenda'],
                      ('%.2f' % p['total']) if m2 else p['total'], (('%+.2f' if m2 else '%+d') % dif) if dif else 'conforme'])
    filas.append(['', 'TOTAL (und)', sum(f[2] for f in filas if not isinstance(f[3], str)),
                  sum(f[3] for f in filas if not isinstance(f[3], str)), ''])
    MS.tabla(M, 'Cantidades: leyenda de la lámina, planta y metrado', ['ITEM', 'PARTIDA', 'LEYENDA', 'METRADO', 'DIFERENCIA'], filas,
             [1.4, W - 8.2, 2.0, 2.2, 2.6], tam=8.5, izq=(1,))
    M.par('El extintor (05.01) es uno por cada señal de ubicación de extintor. Las demás filas de la leyenda (avisador sonoro, escalera '
          'de salida, extintores CO2 y K, alarma, barra antipánico, botiquín, riesgo de descarga, prohibido el ingreso, sensor de '
          'temperatura, pozo a tierra y central y panel de alarma) están en 00 y no tienen símbolos en la planta: no se metran. La '
          'pintura de las zonas de evacuación (01.05) sale de la lámina %s (su fila de la leyenda, “zona de reuniones”, está en 00).' % D.LAMINA_E[0], despues=4)

    sec('2. COMPARACIÓN CON EL METRADO DE EVACUACIÓN DE ARQUITECTURA')
    import compara_metrado_sr as CM
    _, res, det, errs = CM.generar()
    pts_c = {p['clave']: p for p in pts if p['und']}
    pts_c['EXTINTOR'] = next(p for p in pts if p['und'] and p['desc'].startswith('EXTINTOR DE'))
    filas = []
    for x in res:
        k = CM.clave(x['desc'])
        pl = pts_c[k]['total'] if k else 0
        hd = det.get(x['item'], {}).get('total')
        ok = abs(x['total'] - pl) < 0.1 and hd is not None and abs(hd - pl) < 0.1
        t = x['desc'].replace('SUMINISTRO Y COLOCACION DE ', '').capitalize()
        for a, b in (('ss.hh', 'SS.HH.'), ('pqs', 'PQS'), ('kg', 'kg'), ('e=10cm', 'e = 10 cm'), ('electrico', 'eléctrico'),
                     ('botiquin', 'botiquín')):
            t = t.replace(a, b)
        reparto = k == 'SEÑALES DE SERVICIOS HIGIENICOS PHD'
        filas.append([x['item'], t, x['total'], hd, ('%.2f' % pl) if x['und'].upper() == 'M2' else pl,
                      'revisar reparto' if (ok and reparto) else ('bien' if ok else 'revisar')])
    filas.append(['—', '(no está) Señal de punto ecológico', 0, 0, pts_c['PUNTO ECOL']['total'], 'revisar'])
    MS.tabla(M, 'Metrado de evacuación de arquitectura (HR y HD) y plano', ['ITEM', 'PARTIDA', 'HR', 'HD', 'PLANO', 'RESULTADO'], filas,
             [2.2, W - 10.2, 1.8, 1.8, 2.0, 2.4], tam=8, izq=(1,),
             colores={(i, 5): ('E2EFDA' if f_[5] == 'bien' else 'F8CBAD') for i, f_ in enumerate(filas)})
    M.par('Están bien (total y reparto por edificación iguales a la planta): salida, zona segura, dirección de evacuación, aforo, '
          'pintura de zonas de evacuación, '
          'señal de extintor y extintor PQS de 9 kg, luces de emergencia, SS.HH. hombres y mujeres, y riesgo eléctrico (total igual a '
          'las 39 imágenes; su reparto no se puede comprobar en el DXF y se toma el de ese metrado). Hay que corregir:', despues=3)
    MS.vines(M, ['SS.HH. mixtos: en HR sale 0 porque la celda E26 enlaza a HD!K107 (vacía) en vez de HD!K112; además el desagregado '
                 'suma 19 y en el plano hay 20 (falta la de la garita de control, junto al ingreso).',
                 'Detector de humo: 66; en el plano hay 67 (falta el de la garita de control).',
                 'SS.HH. discapacitados: el total (3) está bien, pero la tercera señal está en E-06 (junto al baño accesible de las '
                 'habitaciones), no en E-05.',
                 'Señal de botiquín: 9 (una por edificación), pero el plano no tiene ninguna (fila de la leyenda en 00): o se dibujan en la '
                 'planta o se retira la partida.',
                 'Pintura de zona segura: bien. Las 13 zonas están en la lámina %s (anillo de 4.30 m y línea de 0.15 m); el área exacta es '
                 '13 x 1.9556 = 25.42 m² y con el área redondeada a 1.96 m² sale 25.48 m² (0.06 m² de diferencia, por redondeo). Solo la '
                 'descripción está mal: dice “E=10CM” y la línea del plano (y la del área usada) es de 0.15 m.' % D.LAMINA_E[0],
                 'Falta la partida de la señal de punto ecológico: 4 en los exteriores.',
                 'Del archivo: título “PINTADO DE ZONA SEGURA” repetido (HD filas 134 y 135); fechas distintas (HR 01/09/2026, HD '
                 '01/09/2025); “SEÑAL DEDIRECCIÓN”; la hoja “Hoja1” es un cuadro de equipamiento de aulas de otro proyecto.'])
    M.par('El detalle por partida y por edificación, con fórmulas, está en el Excel de comparación de esta misma carpeta.', despues=4)

    sec('3. OBSERVACIONES QUE SE DEBEN CORREGIR EN LOS PLANOS')
    MS.tabla(M, 'Observaciones a los planos de señalización y evacuación', ['N.°', 'TEMA', 'OBSERVACIÓN', 'CORRECCIÓN'],
             [[k, a, b, c] for k, (a, b, c) in enumerate(observaciones(P), 1)], [0.9, 2.8, (W - 3.7) * 0.55, (W - 3.7) * 0.45],
             tam=8, izq=(1, 2, 3))

    sec('4. PARA CONFIRMAR CON LA ARQUITECTA')
    MS.tabla(M, 'Puntos para confirmar', ['TEMA', 'QUÉ CONFIRMAR', 'POR QUÉ'], [
        ['Extintor de la cocina', 'La cocina (E-05) solo tiene extintores PQS; la leyenda tiene la fila de extintor clase K en 00.',
         'Para fuegos de aceites y grasas de cocina conviene un extintor clase K; si se agrega, dibujarlo y metrarlo.'],
        ['Botiquín', 'La fila de botiquín de primeros auxilios está en 00.', 'El tópico (E-02) tiene botiquín en el mobiliario; '
         'conviene señalizarlo.'],
        ['Braille en señales altas', 'Salida, dirección de evacuación y aforo van a 2.10 m, fuera del alcance de la mano.',
         'Para que el braille sea útil, su placa puede ir en el muro junto a la puerta, a una altura accesible.'],
        ['Luces de emergencia', 'El metrado de evacuación de arquitectura las mide como “señal de luces de emergencia”; el metrado '
         'de señalización, como unidad de luces de emergencia (equipo).', 'Definir si el equipo va en instalaciones eléctricas, para '
         'no duplicarlo ni dejarlo sin metrar.'],
        ['Habitaciones', 'Los rótulos de ambientes de la planta suman 24 habitaciones (4 “simple tipo 2”); el programa tiene 23 (3 '
         'simples tipo II).', 'Revisar si un rótulo está repetido; la población se tomó del programa (113 plazas).'],
    ], [3.0, 7.2, W - 10.2], tam=8.5, izq=(0, 1, 2))

    d.add_page_break()
    sec('5. SÍMBOLOS CONTADOS EN LA PLANTA')
    M.par('Cada círculo de color marca un símbolo contado en el metrado (las señales de riesgo eléctrico no se pueden marcar porque '
          'el DXF no trae su posición):', despues=4)
    MS.figura(M, os.path.join(FIG, 'conteo.png'), 'Símbolos contados en la planta de la lámina %s' % LS, 14.5, 'Fuente: plano de señalización de la arquitecta.')
    d.core_properties.author = ''
    d.core_properties.title = 'Evacuación y señalización: lectura del plano y observaciones - Hogar de Refugio'
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
