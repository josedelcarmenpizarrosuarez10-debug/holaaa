# -*- coding: utf-8 -*-
"""Comparación del metrado de evacuación y señalización de arquitectura (Excel de la usuaria, hojas HR y HD) con lo dibujado en
el plano de señalización del Hogar de Refugio Temporal (CUI 2675514) y con el metrado propio (datos_sr.py).

Lee el Excel de la usuaria sin modificarlo (insumos/solange/senalizacion/METRADOS_EVACUACION_REFUGIO_MUJERES_ok.xlsx): partidas
de la hoja HR (total que pasa al presupuesto, con la celda que enlaza), desagregado por edificación de la hoja HD y fórmulas.
Salida (documento de trabajo): entregables/revision_interna/COMPARACION_METRADO_EVACUACION_Y_SENALIZACION_HOGAR_REFUGIO.xlsx con las hojas
COMPARACIÓN (por partida, con fórmulas de diferencia), POR EDIFICACIÓN y ERRORES DEL ARCHIVO.

Uso: python3 herramientas/senalizacion_refugio/compara_metrado_sr.py
"""
import os
import re
import sys

import openpyxl
from openpyxl.utils import get_column_letter as L

AQUI = os.path.dirname(os.path.abspath(__file__)); HER = os.path.dirname(AQUI); RAIZ = os.path.dirname(HER)
for d in (AQUI, os.path.join(HER, 'varones', 'senalizacion')):
    sys.path.insert(0, d)
import datos_sr as D  # noqa: E402
import metrado_se as MS  # noqa: E402

MS.D = D
celda, cabecera, impresion = MS.celda, MS.cabecera, MS.impresion
GRANATE, ROSA, GRIS, IZQ = MS.GRANATE, MS.ROSA, MS.GRIS, MS.IZQ
ENT = os.path.join(RAIZ, 'insumos/solange/senalizacion/METRADOS_EVACUACION_REFUGIO_MUJERES_ok.xlsx')
SALIDA = os.path.join(D.SAL, 'revision_interna/COMPARACION_METRADO_EVACUACION_Y_SENALIZACION_HOGAR_REFUGIO.xlsx')
VERDE, ROJO = 'E2EFDA', 'F8CBAD'

# palabra de la partida de la usuaria -> clave de la partida propia (None = no está en el plano)
EQUIV = [('PINTURA ZONA SEGURA', 'ZONA DE REUNI'), ('SEÑAL DE SALIDA', 'SALIDA CON'), ('ZONA SEGURA', 'ZONA SEGURA EN'), ('DIRECCIÓN', 'DIRECCIÓN DE EVACUACIÓN'),
         ('RIESGO ELECTRICO', 'ATENCIÓN RIESGO'), ('LUCES DE EMERGENCIA', 'UNIDAD DE LUCES'), ('AFORO', 'LETRERO DE AFORO'),
         ('SEÑAL DE EXTINTOR', 'SEÑAL DE UBICACIÓN DE EXTINTOR PQS'), ('BOTIQUIN', None),
         ('SS.HH HOMBRES', 'SEÑALES DE SERVICIOS HIGIENICOS VAR'), ('SS.HH MUJERES', 'SEÑALES DE SERVICIOS HIGIENICOS MUJER'),
         ('SS.HH MIXTOS', 'MIXTO'), ('SS.HH DISCAPACITADOS', 'SEÑALES DE SERVICIOS HIGIENICOS PHD'),
         ('DETECTOR DE HUMO', 'DETECTOR DE HUMO'), ('EXTINTOR PQS 9', 'EXTINTOR')]


def clave(desc):
    d = desc.upper()
    for k, v in EQUIV:
        if k in d:
            return v
    return '?'


def leer_usuaria():
    v = openpyxl.load_workbook(ENT, data_only=True); f = openpyxl.load_workbook(ENT)
    hr, hd, hrf = v['HR'], v['HD'], f['HR']
    resumen = []
    for r in range(13, hr.max_row + 1):
        it, de, un = hr.cell(r, 2).value, hr.cell(r, 3).value, hr.cell(r, 4).value
        if it and de and un:
            resumen.append(dict(item=str(it), desc=re.sub(r'\s+', ' ', de).strip(), und=un, total=hr.cell(r, 5).value or 0,
                                formula=hrf.cell(r, 5).value, fila=r))
    det = {}; actual = None
    for r in range(14, hd.max_row + 1):
        b, c, k = hd.cell(r, 2).value, hd.cell(r, 3).value, hd.cell(r, 11).value
        if b and c and hd.cell(r, 4).value:
            actual = str(b); det[actual] = dict(total=k or 0, fila=r, por={})
        elif actual and isinstance(c, str) and c.upper().startswith('EDIFICACION'):
            n = int(re.search(r'\d+', c).group())
            det[actual]['por']['E-%02d' % n] = (hd.cell(r, 5).value or 1) * (hd.cell(r, 6).value or 0)
    return resumen, det, v, f


def errores(resumen, det, v, f):
    out = []
    hd, hrf, hdf = v['HD'], f['HR'], f['HD']
    for x in resumen:
        m = re.match(r'=HD!K(\d+)', str(x['formula']))
        if m and x['item'] in det and int(m.group(1)) != det[x['item']]['fila']:
            out.append(('HR!E%d' % x['fila'], 'La partida %s %s enlaza a HD!K%s (vacía) en lugar de HD!K%d: en el resumen sale %s y en el '
                        'desagregado %s.' % (x['item'], x['desc'], m.group(1), det[x['item']]['fila'], x['total'], det[x['item']]['total']),
                        'Cambiar la fórmula a =HD!K%d.' % det[x['item']]['fila']))
    titulos = [(r, hd.cell(r, 3).value) for r in range(14, hd.max_row + 1) if hd.cell(r, 3).value and not hd.cell(r, 4).value
               and hd.cell(r, 2).value]
    for (r1, t1), (r2, t2) in zip(titulos, titulos[1:]):
        if t1 == t2 and r2 == r1 + 1:
            out.append(('HD!B%d:C%d' % (r1, r2), 'El título “%s” está repetido en dos filas seguidas.' % t1, 'Borrar una de las dos filas.'))
    fhr, fhd = v['HR']['C8'].value, hd['C8'].value
    if fhr != fhd:
        out.append(('HR!C8 / HD!C8', 'Fechas distintas: HR %s y HD %s.' % (fhr.strftime('%d/%m/%Y'), fhd.strftime('%d/%m/%Y')),
                    'Poner la misma fecha en las dos hojas.'))
    for ws in v.worksheets:
        if ws.title not in ('HR', 'HD'):
            t = next((c.value for row in ws.iter_rows(max_row=10) for c in row if isinstance(c.value, str) and len(c.value) > 20), '')
            out.append(('Hoja %s' % ws.title, 'Hoja de otro proyecto: “%s...”.' % t[:90], 'Retirar la hoja del metrado del hogar.'))
    for x in resumen:
        if re.search(r'E\s*=\s*10\s*CM', x['desc'].upper()):
            out.append(('HR!C%d' % x['fila'], 'La partida “%s” indica una línea de 10 cm, pero el plano de evacuación tiene una línea de '
                        '0.15 m (anillo de 2.15 a 2.00 m de radio) y el área usada, 1.96 m² por zona, es la de 0.15 m.' % x['desc'],
                        'Corregir la descripción a E = 15 cm.'))
        if 'DEDIRECCI' in x['desc'].replace(' ', '') and 'DE DIRECCI' not in x['desc']:
            out.append(('HR!C%d' % x['fila'], 'Texto “%s”.' % x['desc'], 'Corregir a “SEÑAL DE DIRECCIÓN DE EVACUACIÓN”.'))
    return out


def generar():
    P = D.plano(); pts = {p['clave']: p for p in D.partidas(P) if p['und']}
    pts['EXTINTOR'] = next(p for p in D.partidas(P) if p['und'] and p['desc'].startswith('EXTINTOR DE'))
    resumen, det, v, f = leer_usuaria()
    obs_txt = {
        'SALIDA CON': 'Conforme con el plano.',
        'ZONA SEGURA EN': 'Conforme con el plano (la leyenda dice 210, pero hay 209 dibujadas).',
        'DIRECCIÓN DE EVACUACIÓN': 'Conforme con el plano.',
        'ATENCIÓN RIESGO': 'Total conforme con las 39 imágenes del plano (la leyenda dice 38). El reparto por edificación no se puede '
                           'comprobar en el DXF (imágenes sin posición): se toma el de este metrado.',
        'UNIDAD DE LUCES': 'Total conforme (77 "LE" + 38 "A$C7A2016C2"). En el metrado propio es la unidad de luces de emergencia '
                           '(equipo); en este, "señal de luces de emergencia": definir si el equipo va en instalaciones eléctricas.',
        'LETRERO DE AFORO': 'Conforme con el plano.',
        'SEÑAL DE UBICACIÓN DE EXTINTOR PQS': 'Conforme con el plano.',
        'SEÑALES DE SERVICIOS HIGIENICOS VAR': 'Conforme con el plano.',
        'SEÑALES DE SERVICIOS HIGIENICOS MUJER': 'El hogar es solo para mujeres: las 20 señales de SS.HH. "mixto" del plano (19 en HD + '
                                                 'la de la garita de control) son de SS.HH. mujeres. Total: 3 + 20 = 23.',
        'SEÑALES DE SERVICIOS HIGIENICOS MIXTO': 'Falta la señal de la garita de control (SS.HH. del ingreso): son 20. Además, en HR '
                                                 'sale 0 por un error de fórmula (ver hoja ERRORES DEL ARCHIVO).',
        'SEÑALES DE SERVICIOS HIGIENICOS PHD': 'Total conforme, pero la tercera señal está en E-06 (junto al baño accesible de las '
                                               'habitaciones), no en E-05 (hoja POR EDIFICACIÓN).',
        'DETECTOR DE HUMO': 'Falta el detector de la garita de control: son 67 (la leyenda dice 68 por una copia suelta fuera del dibujo).',
        'EXTINTOR': 'Conforme: un extintor PQS de 9 kg por cada señal de extintor.',
        'ZONA DE REUNI': 'Conforme con la lámina SE-E-01: 13 zonas de 4.30 m de diámetro con anillo de 0.15 m. El área exacta de cada '
                         'anillo es π x (2.15² - 2.00²) = 1.9556 m² y 13 x 1.9556 = 25.42 m²; con el área redondeada a 1.96 m² sale '
                         '25.48 m² (0.06 m² más, solo por redondeo). La descripción dice "E=10CM", pero el plano y el área de 1.96 m² '
                         'son de una línea de 0.15 m (con 0.10 m serían 1.32 m² por zona): corregir a E = 15 cm.',
        None: '',
    }
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'COMPARACIÓN'
    r = cabecera(ws, 'COMPARACIÓN DEL METRADO DE EVACUACIÓN DE ARQUITECTURA CON EL PLANO DE SEÑALIZACIÓN', 9)
    for k, t in enumerate(['ITEM (HR)', 'PARTIDA DEL METRADO DE ARQUITECTURA', 'UND', 'METRADO HR', 'DESAGREGADO HD', 'PLANO %s' % D.LS,
                           'DIFERENCIA (PLANO - HR)', 'RESULTADO', 'OBSERVACIÓN'], 1):
        celda(ws, r, k, t, True, GRANATE, 'FFFFFF')
    r += 1; r0 = r; usadas = set()
    for x in resumen:
        k = clave(x['desc']); usadas.add(k)
        if k == 'MIXTO':
            plano, obs = 0, ('No corresponde: el hogar es solo para mujeres. Las 20 señales "mixto" del plano (en HD 19; falta la de la '
                             'garita) pasan a la partida de SS.HH. mujeres. Además, en HR sale 0 por un error de fórmula (HR!E26).')
        elif k is None:
            plano, obs = 0, ('El plano no tiene señales de botiquín (la fila de la leyenda está en 00). Si se quieren las 9 (una por '
                             'edificación), hay que dibujarlas en la planta; si no, retirar la partida.')
        else:
            plano, obs = pts[k]['total'], obs_txt.get(k, '')
        celda(ws, r, 1, x['item']); celda(ws, r, 2, x['desc'], al=IZQ); celda(ws, r, 3, x['und'])
        celda(ws, r, 4, x['total'], fmt='0.00' if x['und'].upper() == 'M2' else '0')
        celda(ws, r, 5, det.get(x['item'], {}).get('total'), fmt='0.00' if x['und'].upper() == 'M2' else '0')
        celda(ws, r, 6, plano, True, ROSA, fmt='0.00' if x['und'].upper() == 'M2' else '0')
        celda(ws, r, 7, '=F%d-D%d' % (r, r), fmt='+0.00;-0.00;0' if x['und'].upper() == 'M2' else '+0;-0;0')
        celda(ws, r, 8, '=IF(AND(ABS(D%d-F%d)<0.1,ABS(E%d-F%d)<0.1),"BIEN","REVISAR")' % (r, r, r, r), True)
        celda(ws, r, 9, obs, al=IZQ, size=8)
        ws.row_dimensions[r].height = max(24, 12 * (len(obs) // 70 + 1))
        r += 1
    for k in ('PUNTO ECOL',):
        if k not in usadas:
            celda(ws, r, 1, '—'); celda(ws, r, 2, '(no está) ' + pts[k]['desc'], al=IZQ); celda(ws, r, 3, 'Und')
            celda(ws, r, 4, 0, fmt='0'); celda(ws, r, 5, 0, fmt='0'); celda(ws, r, 6, pts[k]['total'], True, ROSA, fmt='0')
            celda(ws, r, 7, '=F%d-D%d' % (r, r), fmt='+0;-0;0'); celda(ws, r, 8, '=IF(AND(D%d=F%d,E%d=F%d),"BIEN","REVISAR")' % (r, r, r, r), True)
            celda(ws, r, 9, 'En los exteriores hay 4 puntos ecológicos dibujados (la leyenda dice 00): falta la partida de su señal.',
                  al=IZQ, size=8); r += 1
    from openpyxl.formatting.rule import CellIsRule
    ws.conditional_formatting.add('H%d:H%d' % (r0, r - 1), CellIsRule(operator='equal', formula=['"BIEN"'], fill=openpyxl.styles.PatternFill('solid', fgColor=VERDE)))
    ws.conditional_formatting.add('H%d:H%d' % (r0, r - 1), CellIsRule(operator='equal', formula=['"REVISAR"'], fill=openpyxl.styles.PatternFill('solid', fgColor=ROJO)))
    for c, w in zip(range(1, 10), [12, 46, 6, 10, 11, 10, 11, 10, 62]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 9, r)

    # ---------------------------------------------------------------- por edificación
    w2 = wb.create_sheet('POR EDIFICACIÓN')
    eds = [c for c, _ in D.edificaciones(P)]
    r = cabecera(w2, 'COMPARACIÓN POR EDIFICACIÓN (HD DEL METRADO DE ARQUITECTURA / PLANO %s)' % D.LS, 3 + len(eds) + 2)
    for k, t in enumerate(['ITEM', 'PARTIDA', 'FUENTE'] + eds + ['TOTAL', 'RESULTADO'], 1):
        celda(w2, r, k, t, True, GRANATE, 'FFFFFF')
    r += 1
    for x in resumen:
        k = clave(x['desc'])
        if k in (None, 'ZONA DE REUNI', 'MIXTO') or x['item'] not in det:
            continue
        p = pts[k]
        fu = det[x['item']]['por']
        por = p['por'] if not p['sin_ubicar'] else {}
        for fuente, val in (('HD arquitectura', fu), ('Plano %s' % D.LS, por)):
            celda(w2, r, 1, x['item'] if fuente.startswith('HD') else ''); celda(w2, r, 2, x['desc'] if fuente.startswith('HD') else '', al=IZQ, size=8)
            celda(w2, r, 3, fuente, fondo=None if fuente.startswith('HD') else ROSA, size=8)
            for j, e in enumerate(eds, 4):
                celda(w2, r, j, val.get(e, 0) or None, fmt='0', fondo=None if fuente.startswith('HD') else ROSA)
            celda(w2, r, 4 + len(eds), '=SUM(%s%d:%s%d)' % (L(4), r, L(3 + len(eds)), r), True, fmt='0')
            if fuente.startswith('Plano'):
                if p['sin_ubicar']:
                    celda(w2, r, 5 + len(eds), 'sin posición en el DXF', size=7)
                else:
                    cond = ','.join('%s%d=%s%d' % (L(j), r - 1, L(j), r) for j in range(4, 5 + len(eds)))
                    celda(w2, r, 5 + len(eds), '=IF(AND(%s),"BIEN","REVISAR")' % cond, True)
            else:
                celda(w2, r, 5 + len(eds), None)
            r += 1
    w2.conditional_formatting.add('%s7:%s%d' % (L(5 + len(eds)), L(5 + len(eds)), r), CellIsRule(operator='equal', formula=['"BIEN"'], fill=openpyxl.styles.PatternFill('solid', fgColor=VERDE)))
    w2.conditional_formatting.add('%s7:%s%d' % (L(5 + len(eds)), L(5 + len(eds)), r), CellIsRule(operator='equal', formula=['"REVISAR"'], fill=openpyxl.styles.PatternFill('solid', fgColor=ROJO)))
    for c, w in zip(range(1, 6 + len(eds)), [11, 40, 14] + [6] * len(eds) + [7, 11]):
        w2.column_dimensions[L(c)].width = w
    impresion(w2, 5 + len(eds), r)

    # ---------------------------------------------------------------- errores del archivo
    w3 = wb.create_sheet('ERRORES DEL ARCHIVO')
    r = cabecera(w3, 'ERRORES DEL ARCHIVO DEL METRADO DE ARQUITECTURA (FÓRMULAS Y FORMATO)', 4)
    for k, t in enumerate(['N.°', 'CELDA / HOJA', 'ERROR', 'CORRECCIÓN'], 1):
        celda(w3, r, k, t, True, GRANATE, 'FFFFFF')
    r += 1
    errs = errores(resumen, det, v, f)
    for n, (a, b, c) in enumerate(errs, 1):
        celda(w3, r, 1, n); celda(w3, r, 2, a); celda(w3, r, 3, b, al=IZQ); celda(w3, r, 4, c, al=IZQ); w3.row_dimensions[r].height = 30; r += 1
    for c, w in zip(range(1, 5), [5, 18, 70, 40]):
        w3.column_dimensions[L(c)].width = w
    impresion(w3, 4, r)
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    wb.properties.creator = ''; wb.properties.title = 'Comparación del metrado de señalización - Hogar de Refugio Temporal'
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    wb.save(SALIDA)
    return SALIDA, resumen, det, errs


if __name__ == '__main__':
    s, res, det, errs = generar()
    print(s)
    for e in errs: print(e)
