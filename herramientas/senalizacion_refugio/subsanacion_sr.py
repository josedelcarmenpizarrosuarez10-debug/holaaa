# -*- coding: utf-8 -*-
"""Evacuación y señalización del Hogar de Refugio Temporal (CUI 2675514): observaciones ya subsanadas en el metrado, la memoria
descriptiva y las especificaciones técnicas (en todo manda lo dibujado en los planos), lo que falta corregir en las láminas SE-S-02
y SE-E-01 para que cuadren con esos documentos, las correcciones del metrado de evacuación de arquitectura (hojas HR y HD) y las
decisiones que quedan a la arquitecta. Documento de trabajo (no va al expediente).

Fuentes: datos_sr.py (partidas y cantidades del plano), nota_sr.observaciones (observaciones a los planos), compara_metrado_sr.py
(cotejo con el Excel HR/HD de la usuaria).

Uso: python3 herramientas/senalizacion_refugio/subsanacion_sr.py
Salida: entregables/revision_interna/OBSERVACIONES_SUBSANADAS_Y_PENDIENTES_EVACUACION_Y_SENALIZACION_HOGAR_REFUGIO.docx
"""
import os
import sys

import docx
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); HER = os.path.dirname(AQUI)
for d in (AQUI, HER, os.path.join(HER, 'varones', 'senalizacion')):
    sys.path.insert(0, d)
import datos_sr as D  # noqa: E402
import nota_sr as NS  # noqa: E402
import compara_metrado_sr as CM  # noqa: E402
import incompatibilidades_refugio as IR  # noqa: E402  (formato de las notas internas)
import memoria_se as MS  # noqa: E402  (tablas)

SALIDA = os.path.join(D.SAL, 'revision_interna/OBSERVACIONES_SUBSANADAS_Y_PENDIENTES_EVACUACION_Y_SENALIZACION_HOGAR_REFUGIO.docx')
VERDE, AMBAR = 'E2EFDA', 'FFF2CC'
LS, LE = D.LS, D.LAMINA_E[0]


def cap(t):
    t = t.replace(' CON BRAILLE', '').capitalize()
    for a, b in (('pqs', 'PQS'), (' abc', ' ABC'), ('(pqs)', '(PQS)'), ('ss.hh', 'SS.HH.'), ('4.30 m, línea de 0.15 m', '4.30 m y línea de 0.15 m')):
        t = t.replace(a, b)
    return t


def generar():
    P = D.plano(); pts = D.partidas(P); und = [p for p in pts if p['und']]
    R = IR.R
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('EVACUACIÓN Y SEÑALIZACIÓN: OBSERVACIONES SUBSANADAS Y CORRECCIONES PENDIENTES EN LOS PLANOS', 12.5, True, color=R.GRANATE,
          centro=True, despues=2)
    M.par('Hogar de Refugio Temporal - Mujeres Violentadas - CUI N.° %s  |  Documento de trabajo, no forma parte del expediente' % D.CUI,
          9, color=R.GRIS, centro=True, despues=10)
    M.par('El metrado, la memoria descriptiva y las especificaciones técnicas de evacuación y señalización dicen lo mismo y siguen lo '
          'dibujado en los planos: %d partidas en la planta de señalización %s y en la planta de evacuación %s. Se leyeron todos los '
          'símbolos: los bloques de la planta (también los que están dentro de otros bloques), los dibujos sueltos (líneas, sombreados y '
          'textos) y las imágenes pegadas; no hay señales dibujadas fuera de los bloques. Donde la leyenda o los cuadros dicen otra cosa, '
          'los documentos ya están corregidos (parte 1); para que los planos cuadren, la arquitecta debe hacer las correcciones de la '
          'parte 2; la parte 3 corrige el metrado de evacuación de arquitectura (HR y HD) y la parte 4 son decisiones de la arquitecta.'
          % (len(und), LS, LE))

    def sec(t):
        p = M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4); p.paragraph_format.keep_with_next = True

    # ------------------------------------------------------------------ 1 subsanadas
    sec('1. OBSERVACIONES YA SUBSANADAS EN EL METRADO, LA MEMORIA Y LAS ESPECIFICACIONES')
    filas = []
    for p in und:
        if p['desc'].startswith('EXTINTOR DE'):
            continue
        m2 = p['und'] == 'm2'
        if abs(p['total'] - p['leyenda']) > 1e-9:
            filas.append([p['item'], cap(p['desc']), p['leyenda'], ('%.2f' % p['total']) if m2 else p['total'], p['fuente']])
    MS.tabla(M, 'Cantidades de los documentos distintas de la leyenda de %s (manda lo dibujado)' % LS,
             ['ITEM', 'PARTIDA', 'LEYENDA', 'DOCUMENTOS', 'DE DÓNDE SALE'], filas, [1.3, 4.6, 1.6, 2.0, W - 9.5], tam=8, izq=(1, 4),
             colores={(i, 3): VERDE for i in range(len(filas))})
    mz = P['bloques']['*U2820']
    otros = [
        ['SS.HH. mixto', 'El hogar es solo para mujeres: las %d señales “mixto” de la planta se metran, describen y especifican como '
         'SS.HH. mujeres (partida 04.02: 3 + %d = %d und). No hay partida de SS.HH. mixto.' % (mz['cantidad'], mz['cantidad'],
                                                                                              3 + mz['cantidad'])],
        ['Luces de emergencia', 'Se cuentan los dos bloques del mismo símbolo (“LE” y “A$C7A2016C2”): 115 und.'],
        ['Detector de humo', 'No se metra la copia suelta fuera del dibujo (coordenadas del doble de las del terreno): 67 und.'],
        ['Garita de control', 'Su señal de SS.HH. y su detector de humo se metran en exteriores (EXT).'],
        ['Extintores', '33 extintores PQS tipo ABC de 9 kg (uno por cada señal), con la señal a 1.80 m sobre el equipo.'],
        ['Riesgo eléctrico', '39 imágenes del plano (la leyenda dice 38), con el reparto por edificación del metrado de evacuación de '
         'arquitectura (el DXF no trae la posición de las imágenes).'],
        ['Zonas de evacuación', 'Las 13 zonas de %s (anillo de 4.30 m y línea de 0.15 m): 25.42 m² de pintura; la zona 10 repetida '
         '(junto a E-06) es la zona 13.' % LE],
        ['Rutas de evacuación', 'Las 84 rutas de %s con su longitud (máximo 39.67 m ≤ 45 m) y su edificación de origen, en el metrado y '
         'la memoria.' % LE],
        ['Reparto por edificación', 'Cada símbolo se asigna a su edificación (donde los aleros de dos edificaciones se traslapan, a la que '
         'lo contiene); coincide con el metrado de arquitectura en salida, zona segura, dirección, aforo, luces y extintores.'],
        ['Láminas y nombres', 'Se citan con los códigos del índice de arquitectura: %s planta general de señalización y %s planta general '
         'de evacuación; E-06 a E-08 se llaman “Habitaciones (1) a (3)”.' % (LS, LE)],
        ['Población', '113 residentes (plazas de las 23 habitaciones), 62 de personal y 130 temporales: 305 personas.'],
    ]
    MS.tabla(M, 'Otras observaciones subsanadas en los documentos', ['TEMA', 'CÓMO QUEDÓ EN LOS DOCUMENTOS'], otros, [3.4, W - 3.4],
             tam=8.5, izq=(0, 1))

    # ------------------------------------------------------------------ 2 pendientes en los planos
    sec('2. LO QUE FALTA QUE CORRIJA LA ARQUITECTA EN LOS PLANOS %s Y %s' % (LS, LE))
    M.par('Con estas correcciones, la leyenda y los rótulos dicen lo mismo que la planta y que el metrado:', despues=3)
    obs = NS.observaciones(P)
    MS.tabla(M, 'Correcciones pendientes en los planos', ['N.°', 'TEMA', 'OBSERVACIÓN', 'CORRECCIÓN'],
             [[k, a, b, c] for k, (a, b, c) in enumerate(obs, 1)], [0.9, 2.8, (W - 3.7) * 0.55, (W - 3.7) * 0.45], tam=8,
             izq=(1, 2, 3))

    # ------------------------------------------------------------------ 3 metrado de arquitectura
    sec('3. CORRECCIONES DEL METRADO DE EVACUACIÓN DE ARQUITECTURA (HOJAS HR Y HD)')
    _, res, det, errs = CM.generar()
    ptc = {p['clave']: p for p in und}
    ptc['EXTINTOR'] = next(p for p in und if p['desc'].startswith('EXTINTOR DE'))
    filas = []
    for x in res:
        k = CM.clave(x['desc'])
        pl = ptc[k]['total'] if k in ptc else 0
        hd = det.get(x['item'], {}).get('total')
        if abs(x['total'] - pl) < 0.1 and hd is not None and abs(hd - pl) < 0.1 and k != 'SEÑALES DE SERVICIOS HIGIENICOS PHD':
            continue
        if k == 'MIXTO':
            acc = 'Retirar la partida (o dejarla en 0): sus 20 señales van en SS.HH. mujeres. Además HR!E26 enlaza a HD!K107 (vacía).'
        elif k is None:
            acc = 'El plano no tiene señales de botiquín: retirar la partida o dibujarlas en la planta.'
        elif k == 'SEÑALES DE SERVICIOS HIGIENICOS MUJER':
            acc = 'Cambiar a 23 (3 de mujeres + 20 de “mixto”, incluida la de la garita).'
        elif k == 'DETECTOR DE HUMO':
            acc = 'Cambiar a 67: falta el detector de la garita de control.'
        elif k == 'SEÑALES DE SERVICIOS HIGIENICOS PHD':
            acc = 'Total bien (3); la tercera señal va en la edificación 6, no en la 5.'
        elif k == 'ZONA DE REUNI':
            acc = 'Cantidad bien (13 zonas; 25.48 m² con el área redondeada a 1.96 m²); la descripción debe decir E = 15 cm, no 10 cm.'
        else:
            acc = 'Cambiar a %s.' % pl
        filas.append([x['item'], x['desc'].replace('SUMINISTRO Y COLOCACION DE ', '').capitalize().replace('ss.hh', 'SS.HH.'),
                      x['total'], hd, ('%.2f' % pl) if x['und'].upper() == 'M2' else pl, acc])
    filas.append(['—', '(falta) Señal de punto ecológico', 0, 0, ptc['PUNTO ECOL']['total'], 'Agregar la partida: 4 und en exteriores.'])
    MS.tabla(M, 'Partidas del metrado de arquitectura que no cuadran con los planos', ['ITEM', 'PARTIDA', 'HR', 'HD', 'PLANO', 'CORRECCIÓN'],
             filas, [2.0, 3.6, 1.2, 1.2, 1.5, W - 9.5], tam=8, izq=(1, 5))
    MS.tabla(M, 'Errores del archivo', ['CELDA / HOJA', 'ERROR', 'CORRECCIÓN'], [list(e) for e in errs], [2.6, (W - 2.6) * 0.6, (W - 2.6) * 0.4],
             tam=8, izq=(1, 2))
    M.par('Las partidas que no figuran en estas tablas (salida, zona segura, dirección de evacuación, riesgo eléctrico, luces de '
          'emergencia, aforo, señal de extintor, extintor PQS de 9 kg y SS.HH. hombres) están bien en total y por edificación.', despues=4)

    # ------------------------------------------------------------------ 4 decisiones
    sec('4. DECISIONES DE LA ARQUITECTA (CAMBIARÍAN LOS DOCUMENTOS)')
    zs = D.zonas()
    MS.tabla(M, 'Decisiones por confirmar', ['TEMA', 'SITUACIÓN EN LOS DOCUMENTOS', 'SI LA ARQUITECTA LO CAMBIA'], [
        ['Aforo de las zonas de evacuación', 'Se pintan las 13 zonas del plano; sus aforos (12, 20 y 40 personas, mismo tamaño) suman %d '
         'personas, menos que la capacidad máxima (%d).' % (sum(z['aforo'] for z in zs), D.POBLACION['maxima']),
         'Si se agregan o agrandan zonas, cambian el área de pintura y la tabla de zonas.'],
        ['Riesgo eléctrico', '39 señales con el reparto del metrado de arquitectura.', 'Si se convierten las imágenes a bloque y son 38, '
         'o cambian de edificación, se actualiza la partida.'],
        ['Extintor de la cocina', 'Solo extintores PQS; la fila de extintor clase K está en 00.', 'Si se agrega un extintor K, se dibuja y '
         'se metra (señal y equipo).'],
        ['Botiquín', 'Sin señal de botiquín (la fila está en 00).', 'Si se dibuja (por ejemplo en el tópico), se agrega la partida.'],
        ['Detección y alarma', 'Detectores de humo con alarma propia; sin central, pulsadores ni avisadores (filas en 00).',
         'Si van a una central, se dibujan y metran la central, los pulsadores y los avisadores.'],
        ['Luces de emergencia', 'Se metran como unidad de luces de emergencia (equipo) en evacuación y señalización.',
         'Si el equipo va en instalaciones eléctricas, aquí queda solo la señal.'],
        ['Braille en señales altas', 'Las señales con braille llevan su placa debajo, a la altura de la señal (2.10 m en salida y aforo).',
         'Si la placa va en el muro a altura accesible, cambia la especificación (no la cantidad).'],
    ], [3.2, 6.4, W - 9.6], tam=8.5, izq=(0, 1, 2), colores={(i, 0): AMBAR for i in range(7)})
    M.par('Cuando la arquitecta entregue los planos corregidos (DXF), se leen de nuevo y se comparan con el metrado: si coinciden, los '
          'documentos quedan como están; si no, se regeneran con lo que digan los planos corregidos.', despues=4)
    d.core_properties.author = ''
    d.core_properties.title = 'Evacuación y señalización: observaciones subsanadas y pendientes - Hogar de Refugio'
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
