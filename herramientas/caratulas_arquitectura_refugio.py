# -*- coding: utf-8 -*-
"""Carátulas de arquitectura del Hogar de Refugio Temporal "Mujeres Violentadas" (CUI 2675514), a partir del Word de carátulas
que entregó el usuario (18 carátulas armadas para el CAR Varones: memorias, especificaciones, metrados y planos de arquitectura,
evacuación y señalización, y equipamiento y mobiliario).

Se conserva todo el formato, las tablas, el escudo y las medidas de cada imagen; solo se cambia lo propio del proyecto:
  - nombre del proyecto (el de las carátulas del Hogar de Refugio, caratulas_refugio.py);
  - CUI N° 2705619 -> CUI N° 2675514;
  - leyenda "Vista del Centro de Acogida Residencial - Varones" -> "Vista del Hogar de Refugio Temporal";
  - foto del proyecto: la del Hogar de Refugio que entregó la usuaria (insumos/solange/FOTO_PROYECTO_HOGAR_REFUGIO.jpg, del mismo
    tamaño 1600 x 900 que la anterior, así que cada carátula conserva su encuadre).

Uso: python3 herramientas/caratulas_arquitectura_refugio.py
Entrada: insumos/solange/caratulas/PARA_MODIFICAR_CARATULAS_ARQUITECTURA.docx (no se modifica)
Salida: entregables/arquitectura/CARATULAS_ARQUITECTURA_HOGAR_REFUGIO.docx
"""
import io
import os
import sys
import zipfile

import docx
from PIL import Image

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import caratulas_refugio as CR  # noqa: E402  (nombre, CUI y leyenda del Hogar de Refugio)

ENT = os.path.join(RAIZ, 'insumos/solange/caratulas/PARA_MODIFICAR_CARATULAS_ARQUITECTURA.docx')
FOTO = os.path.join(RAIZ, 'insumos/solange/FOTO_PROYECTO_HOGAR_REFUGIO.jpg')
SALIDA = os.path.join(RAIZ, 'entregables/arquitectura/CARATULAS_ARQUITECTURA_HOGAR_REFUGIO.docx')
PROY_V = ('“CREACIÓN DEL SERVICIO DE PROTECCIÓN INTEGRAL A NIÑAS, NIÑOS Y ADOLESCENTES SIN CUIDADOS PARENTALES O EN RIESGO DE '
          'PERDERLOS EN CENTRO DE ACOGIDA RESIDENCIAL - VARONES DISTRITO DE MORALES DE LA PROVINCIA DE SAN MARTÍN DEL DEPARTAMENTO '
          'DE SAN MARTÍN”')
CAMBIOS = [(PROY_V, CR.V.PROYECTO), ('CUI N° 2705619', CR.V.CUI), ('Vista del Centro de Acogida Residencial - Varones', CR.V.LEYENDA)]
FOTO_DOCX = 'word/media/image2.jpeg'                            # foto del proyecto en el Word del usuario (la otra imagen es el escudo)


def textos(d):
    n = {a: 0 for a, _ in CAMBIOS}
    for p in d.paragraphs:
        for r in p.runs:
            for a, b in CAMBIOS:
                if a in r.text:
                    r.text = r.text.replace(a, b); n[a] += 1
    resto = [p.text for p in d.paragraphs if 'VARONES' in p.text.upper() or '2705619' in p.text]
    resto += [c.text for t in d.tables for row in t.rows for c in row.cells if 'VARONES' in c.text.upper() or '2705619' in c.text]
    if resto:
        raise ValueError('quedan textos del CAR Varones: %s' % resto[:3])
    return n


def foto_jpeg(tam):
    im = Image.open(FOTO).convert('RGB')
    if im.size != tam:                                           # mismo encuadre que la foto anterior
        w, h = im.size; k = max(tam[0] / w, tam[1] / h)
        im = im.resize((round(w * k), round(h * k)), Image.LANCZOS)
        x0 = (im.width - tam[0]) // 2; y0 = (im.height - tam[1]) // 2
        im = im.crop((x0, y0, x0 + tam[0], y0 + tam[1]))
    b = io.BytesIO(); im.save(b, 'JPEG', quality=90, optimize=True)
    return b.getvalue()


def generar():
    d = docx.Document(ENT)
    n = textos(d)
    d.core_properties.author = ''; d.core_properties.title = 'Carátulas de arquitectura - Hogar de Refugio Temporal'
    tmp = io.BytesIO(); d.save(tmp)
    with zipfile.ZipFile(io.BytesIO(tmp.getvalue())) as zin:
        tam = Image.open(io.BytesIO(zin.read(FOTO_DOCX))).size
        os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
        with zipfile.ZipFile(SALIDA, 'w', zipfile.ZIP_DEFLATED) as zout:
            for it in zin.infolist():
                zout.writestr(it, foto_jpeg(tam) if it.filename == FOTO_DOCX else zin.read(it.filename))
    return SALIDA, n


if __name__ == '__main__':
    s, n = generar()
    print(s); print(n)
