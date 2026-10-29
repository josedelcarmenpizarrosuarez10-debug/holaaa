# -*- coding: utf-8 -*-
"""Metrado de equipamiento y mobiliario - Hogar de Refugio Temporal "Mujeres Violentadas" (CUI 2675514), en Excel con fórmulas.

Hojas:
  RESUMEN               partidas 03.01 (equipos, mobiliario, implementos por ambiente) con el metrado enlazado por SUMIF;
  METRADO POR AMBIENTE  cada ítem de cada ambiente de las láminas EM-01 a EM-07 con su cantidad (datos_mr.ambientes);
  POR EDIFICACIÓN       cuadro de partidas (equipos y mobiliario) por edificación, con SUMIFS y totales que se cruzan con el RESUMEN.

Uso: python3 herramientas/mobiliario_refugio/metrado_mr.py
Salida: entregables/METRADO_EQUIPAMIENTO_Y_MOBILIARIO_HOGAR_REFUGIO.xlsx
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import datos_mr as D  # noqa: E402

SALIDA = os.path.join(D.SAL, 'METRADO_EQUIPAMIENTO_Y_MOBILIARIO_HOGAR_REFUGIO.xlsx')
GUINDA, GRIS = '7B1E2B', 'F2F2F2'
FINO = Side(style='thin', color='808080'); BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
IZQ = Alignment(horizontal='left', vertical='center', wrap_text=True)


def celda(ws, r, c, v, neg=False, fondo=None, color=None, al=CENTRO, fmt=None, borde=True, size=9):
    x = ws.cell(r, c, v); x.font = Font(name='Arial', size=size, bold=neg, color=color); x.alignment = al
    if fondo:
        x.fill = PatternFill('solid', fgColor=fondo)
    if borde:
        x.border = BORDE
    if fmt:
        x.number_format = fmt
    return x


def cabecera(ws, titulo, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    celda(ws, 1, 1, titulo, True, GUINDA, 'FFFFFF', size=12, borde=False); ws.row_dimensions[1].height = 24
    filas = [('PROYECTO', '“%s”' % D.PROYECTO), ('CUI', 'N.° ' + D.CUI),
             ('PROPIETARIO', 'Gobierno Regional San Martín - Gerencia Territorial Bajo Mayo'),
             ('UBICACIÓN', 'Morales / San Martín / San Martín'), ('ESPECIALIDAD', D.ESPECIALIDAD),
             ('LÁMINAS', 'EM-01 a EM-07 (equipamiento y mobiliario, escala 1/50)'), ('FECHA', D.FECHA)]
    for k, (a, b) in enumerate(filas, 2):
        celda(ws, k, 1, a, True, GRIS, al=IZQ)
        ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=ncol); celda(ws, k, 2, b, k == 2, al=IZQ)
        for c in range(3, ncol + 1):
            ws.cell(k, c).border = BORDE
    ws.row_dimensions[2].height = 36
    return len(filas) + 3


def impresion(ws, ncol, nfil, horizontal=False):
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_area = 'A1:%s%d' % (L(ncol), nfil); ws.sheet_view.showGridLines = False
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_title_rows = '10:10'


def hoja_ambientes(wb, A):
    ws = wb.create_sheet('METRADO POR AMBIENTE')
    r = cabecera(ws, 'METRADO DE EQUIPAMIENTO Y MOBILIARIO POR AMBIENTE (LÁMINAS EM-01 A EM-07)', 8)
    for k, v in enumerate(['EDIF.', 'AMBIENTE', 'LÁMINA', 'CÓD.', 'DESCRIPCIÓN', 'UND', 'CANTIDAD', 'CLAVE'], 1):
        celda(ws, r, k, v, True, GUINDA, 'FFFFFF')
    r += 1; r0 = r
    for ed, amb, items in A:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        celda(ws, r, 1, '%s  %s - %s' % (ed, D.EDIFICACIONES[ed], amb), True, GRIS, al=IZQ)
        for c in range(2, 8):
            ws.cell(r, c).border = BORDE
        r += 1
        for it in items:
            for k, v in enumerate([ed, amb, D.LAMINA[ed], it['cod'], it['des'], 'und', it['cant'], it['clave']], 1):
                x = celda(ws, r, k, v, al=IZQ if k in (2, 5) else CENTRO, fmt='0' if k == 7 else None)
                if k == 8:
                    x.font = Font(name='Arial', size=7, color='808080')
            r += 1
    celda(ws, r, 5, 'TOTAL DE UNIDADES', True, GRIS, al=IZQ); celda(ws, r, 6, 'und', True, GRIS)
    celda(ws, r, 7, '=SUM(G%d:G%d)' % (r0, r - 1), True, GRIS, fmt='0')
    for c, w in zip(range(1, 9), [6, 30, 7, 6, 60, 5, 9, 16]):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = ws.cell(r0, 1)
    impresion(ws, 7, r, True)
    rng = lambda col: "'METRADO POR AMBIENTE'!$%s$%d:$%s$%d" % (col, r0, col, r - 1)   # noqa: E731
    return rng('H'), rng('G'), rng('A'), r


def hoja_resumen(ws, cat, rc, rq):
    r = cabecera(ws, 'HOJA RESUMEN DE METRADOS - ARQUITECTURA: EQUIPAMIENTO Y MOBILIARIO', 5)
    for k, v in enumerate(['PART. N.°', 'CÓD.', 'DESCRIPCIÓN', 'UND', 'METRADO'], 1):
        celda(ws, r, k, v, True, GUINDA, 'FFFFFF')
    r += 1
    fila = {}; grupos = []
    for f in D.numerar(cat):
        if f.get('tipo') == 'item':
            celda(ws, r, 1, f['item']); celda(ws, r, 2, f['cod']); celda(ws, r, 3, f['desc'], al=IZQ); celda(ws, r, 4, 'und')
            ws.cell(r, 6, f['clave']).font = Font(name='Arial', size=7, color='FFFFFF')
            celda(ws, r, 5, '=SUMIF(%s,F%d,%s)' % (rc, r, rq), True, fmt='0.00')
            fila[f['clave']] = r
        else:
            neg = f['nivel'] <= 3
            celda(ws, r, 1, f['item'], neg, GRIS if f['nivel'] <= 1 else None); celda(ws, r, 2, None, fondo=GRIS if f['nivel'] <= 1 else None)
            celda(ws, r, 3, f['desc'], neg, GRIS if f['nivel'] <= 1 else None, al=IZQ)
            celda(ws, r, 4, None, fondo=GRIS if f['nivel'] <= 1 else None); celda(ws, r, 5, None, fondo=GRIS if f['nivel'] <= 1 else None)
            if f['nivel'] == 2:
                grupos.append((f['desc'], r))
        r += 1
    r += 1
    celda(ws, r, 3, 'CONTROL: TOTAL DE UNIDADES DEL RESUMEN', True, GRIS, al=IZQ); celda(ws, r, 4, 'und', True, GRIS)
    celda(ws, r, 5, '=SUM(E11:E%d)' % (r - 2), True, GRIS, fmt='0'); r_tot = r
    for c, w in zip(range(1, 7), [15, 6, 72, 6, 11, 4]):
        ws.column_dimensions[L(c)].width = w
    ws.column_dimensions['F'].hidden = True
    impresion(ws, 5, r)
    return fila, r_tot


def hoja_edificaciones(wb, cat, rc, rq, re_):
    ws = wb.create_sheet('POR EDIFICACIÓN')
    eds = list(D.EDIFICACIONES); n = 4 + len(eds)
    r = cabecera(ws, 'EQUIPOS Y MOBILIARIO POR EDIFICACIÓN (UNIDADES)', n)
    for k, v in enumerate(['CÓD.', 'DESCRIPCIÓN'] + eds + ['TOTAL', 'CLAVE'], 1):
        celda(ws, r, k, v, True, GUINDA, 'FFFFFF')
    r += 1; r0 = r
    for f in cat:
        if f['tipo'] != 'item' or f['g'] == 'K':
            continue
        celda(ws, r, 1, f['cod']); celda(ws, r, 2, f['desc'], al=IZQ, size=8)
        for j, ed in enumerate(eds, 3):
            celda(ws, r, j, '=SUMIFS(%s,%s,$%s%d,%s,%s$%d)' % (rq, rc, L(n), r, re_, L(j), r0 - 1), fmt='0;-0;""')
        celda(ws, r, 3 + len(eds), '=SUM(C%d:%s%d)' % (r, L(2 + len(eds)), r), True, fmt='0')
        x = celda(ws, r, n, f['clave']); x.font = Font(name='Arial', size=7, color='808080')
        r += 1
    celda(ws, r, 2, 'TOTAL', True, GRIS, al=IZQ)
    for j in range(3, 4 + len(eds)):
        celda(ws, r, j, '=SUM(%s%d:%s%d)' % (L(j), r0, L(j), r - 1), True, GRIS, fmt='0')
    for c, w in zip(range(1, n + 1), [6, 52] + [7] * len(eds) + [8, 10]):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = ws.cell(r0, 3)
    impresion(ws, n - 1, r, True)


def generar():
    P = D.plano(); A = D.ambientes(P); cat = D.catalogo(P)
    errores = D.comprobacion(P)[1]
    if errores:
        raise ValueError('cantidades distintas a la planta: %s' % errores)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'RESUMEN'
    rc, rq, re_, _ = hoja_ambientes(wb, A)
    hoja_resumen(ws, cat, rc, rq)
    hoja_edificaciones(wb, cat, rc, rq, re_)
    wb.properties.creator = ''; wb.properties.title = 'Metrado de equipamiento y mobiliario - Hogar de Refugio Temporal'
    wb.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
