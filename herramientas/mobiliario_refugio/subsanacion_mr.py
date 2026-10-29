# -*- coding: utf-8 -*-
"""Equipamiento y mobiliario del Hogar de Refugio Temporal (CUI 2675514): observaciones ya subsanadas en el metrado, la memoria
descriptiva y las especificaciones técnicas, y lo que falta corregir en las láminas EM-01 a EM-07 para que cuadren con esos
documentos. Documento de trabajo para coordinar con la arquitecta (no va al expediente).

Fuente única: datos_mr.py (ajustes de cada ambiente, partidas agregadas o retiradas, observaciones a los cuadros y catálogo).

Uso: python3 herramientas/mobiliario_refugio/subsanacion_mr.py
Salida: entregables/revision_interna/OBSERVACIONES_SUBSANADAS_Y_PENDIENTES_MOBILIARIO_HOGAR_REFUGIO.docx
"""
import os
import sys

import docx
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); HER = os.path.dirname(AQUI)
for d in (AQUI, HER, os.path.join(HER, 'varones', 'senalizacion')):
    sys.path.insert(0, d)
import datos_mr as D  # noqa: E402
import incompatibilidades_refugio as IR  # noqa: E402  (formato de las notas internas)
import memoria_se as MS  # noqa: E402  (tablas)

SALIDA = os.path.join(D.SAL, 'revision_interna/OBSERVACIONES_SUBSANADAS_Y_PENDIENTES_MOBILIARIO_HOGAR_REFUGIO.docx')
VERDE, AMBAR = 'E2EFDA', 'FFF2CC'


def generar():
    P = D.plano(); A = D.ambientes(P); N = [f for f in D.numerar() if f.get('tipo') == 'item']
    R = IR.R
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('EQUIPAMIENTO Y MOBILIARIO: OBSERVACIONES SUBSANADAS Y CORRECCIONES PENDIENTES EN LAS LÁMINAS', 12.5, True, color=R.GRANATE,
          centro=True, despues=2)
    M.par('Hogar de Refugio Temporal - Mujeres Violentadas - CUI N.° %s  |  Documento de trabajo, no forma parte del expediente' % D.CUI,
          9, color=R.GRIS, centro=True, despues=10)
    M.par('El metrado, la memoria descriptiva y las especificaciones técnicas de equipamiento y mobiliario dicen lo mismo: %d partidas '
          'y %d unidades, contadas en la planta de las láminas EM-01 a EM-07 (lo rotulado en cada ambiente). Donde los cuadros de las '
          'láminas no coincidían con la planta, los documentos ya están corregidos (parte 1). Para que las láminas cuadren con esos '
          'documentos, la arquitecta debe hacer las correcciones de la parte 2; la parte 3 son decisiones que, si cambian, obligan a '
          'regenerar los tres documentos.' % (len(N), sum(f['total'] for f in N)))

    def sec(t):
        M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4)

    # ------------------------------------------------------------------ 1 subsanadas
    sec('1. OBSERVACIONES YA SUBSANADAS EN EL METRADO, LA MEMORIA Y LAS ESPECIFICACIONES')
    filas = []
    for ed, amb, its in A:
        for it in its:
            if it['nota']:
                filas.append([ed, D.cap(amb), it['cod'], it['cuadro'] or 0, it['cant'], it['nota']])
    for e, a, cod, q, motivo in D.CORRECCIONES:
        if q == 0:
            filas.append([e, D.cap(a), cod, 1, 0, motivo + ' No se metra.'])
    MS.tabla(M, 'Cantidades corregidas (el documento toma lo dibujado y rotulado en la planta)',
             ['EDIF.', 'AMBIENTE', 'CÓD.', 'CUADRO', 'DOCUMENTOS', 'MOTIVO'], filas, [1.2, 3.4, 1.2, 1.5, 1.9, W - 9.2], tam=8,
             izq=(1, 5), colores={(i, 4): VERDE for i in range(len(filas))})
    otros = [
        ['Descripciones', 'Cada código lleva la descripción de su elemento (catálogo de los cuadros), no la de la fila corrida: M31 '
         'escritorio con credenza, M32 sillón, M36 armario, M39 archivador (Dirección general), E9 impresora multifuncional y E7 '
         'proyector (Sala de reuniones), M37 armario para material (Taller de bisutería), M63 sofá de 3 cuerpos y M65 camilla (Tópico).'],
        ['Cuadros mal titulados', 'Cada cuadro se asignó a su ambiente real: "Sala de espera y sala de visita - Edificación 2" = almacén '
         'de donaciones, tópico y SS.HH. de E-02; "Pasadizo - Edificación 2" = E-04; "SS.HH. personal - Edificación 1" = SS.HH. de las '
         'habitaciones de E-06; "Módulo infantil - Edificación 5" = almacén de implementos deportivos de E-09; "Vestidores - '
         'Edificación 1" = vestidores de E-09.'],
        ['Cuadro repetido', 'El cuadro "Tópico - Edificación 3" repite los ítems del tópico de E-02: no se suma dos veces.'],
        ['Códigos dobles', 'M16 (mesa de trabajo del SUM-comedor y mesa de corte textil del taller de costura), M64 (camarote y canastilla '
         'para pelotas), K32 y K33 (tobogán y juego de teatro del módulo infantil; bidón de oxígeno y tensiómetro del tópico): cada '
         'uno se metra y especifica como partida separada, con su descripción.'],
        ['Extintores', 'Los extintores (S1) de los cuadros "Pasadizo" no se metran en mobiliario: están en evacuación y señalización.'],
        ['Implementos del almacén deportivo', 'K53, K73, K74 y K79, rotulados en la planta y ausentes del cuadro, se metran 1 de cada uno.'],
        ['Nombres', 'E-06, E-07 y E-08 se llaman "Habitaciones (1) a (3)", como en el plano y en el programa de arquitectura.'],
        ['Capacidad del SUM-comedor', 'Manda el plano: 24 mesas M16 y 192 sillas M3 (8 por mesa), contadas en la planta de EM-04. No se '
         'usan las 384 sillas del cuadro ni las 80 personas de la memoria de arquitectura.'],
    ]
    MS.tabla(M, 'Otras observaciones subsanadas en los documentos', ['TEMA', 'CÓMO QUEDÓ EN LOS DOCUMENTOS'], otros, [3.6, W - 3.6],
             tam=8.5, izq=(0, 1))

    # ------------------------------------------------------------------ 2 pendientes en las láminas
    sec('2. LO QUE FALTA QUE CORRIJA LA ARQUITECTA EN LAS LÁMINAS EM-01 A EM-07')
    M.par('Con estas correcciones, cada cuadro de las láminas queda igual a la planta y al metrado:', despues=3)
    filas = []
    lam = D.LAMINA
    for ed, amb, its in A:
        for it in its:
            if it['nota']:
                if (it['cuadro'] or 0) == 0:
                    acc = 'Agregar al cuadro %s: %s, %d und.' % (D.cap(amb), it['cod'], it['cant'])
                else:
                    acc = 'Cuadro %s: %s de %d a %d und.' % (D.cap(amb), it['cod'], it['cuadro'], it['cant'])
                filas.append([lam.get(ed, ''), ed, acc])
    for e, a, cod, q, motivo in D.CORRECCIONES:
        if q == 0:
            filas.append([lam.get(e, ''), e, 'Cuadro %s: retirar la fila %s (no está en la planta).' % (D.cap(a), cod)])
    for l, t in D.observaciones_plano():
        filas.append([l, '', t])
    filas.append(['Memoria de arquitectura', 'E-05', 'La memoria descriptiva de arquitectura considera el SUM-comedor para 80 personas; '
                  'la planta de EM-04 tiene 24 mesas y 192 sillas. Ajustar la capacidad de la memoria a la planta (192).'])
    MS.tabla(M, 'Correcciones pendientes en las láminas', ['N.°', 'LÁMINA', 'EDIF.', 'CORRECCIÓN'],
             [[k] + f for k, f in enumerate(filas, 1)], [0.9, 1.9, 1.3, W - 4.1], tam=8, izq=(3,))

    # ------------------------------------------------------------------ 3 decisiones
    sec('3. DECISIONES QUE CAMBIARÍAN LOS DOCUMENTOS (CONFIRMAR CON LA ARQUITECTA)')
    MS.tabla(M, 'Decisiones por confirmar', ['TEMA', 'SITUACIÓN EN LOS DOCUMENTOS', 'SI LA ARQUITECTA LO CAMBIA'], [
        ['Sillas del comedor', 'El comedor usa la silla regulable de estudiante M3.', 'Si se cambia por silla de adulto, cambia el código '
         'y la partida.'],
        ['Máquinas del taller de costura', 'E10 = "máquina de corte textil" (10 und), descripción del plano.',
         'Si son máquinas de coser, cambian la descripción y la especificación.'],
        ['Implementos K53, K73, K74 y K79', 'Se metra 1 de cada uno (rotulados en la planta, sin cantidad en el cuadro).',
         'Indicar la cantidad en el cuadro; si es otra, cambia el metrado.'],
        ['Nuevos códigos', 'Los códigos dobles (M16, M64, K32, K33) se mantienen como están en el plano.', 'Si se les da otro código '
         'en las láminas, hay que pasar los nuevos códigos al metrado, la memoria y las especificaciones.'],
        ['Mesas M19 y M22', 'Ambas son mesa de madera 1.00 x 1.80 x 0.75 m, en partidas separadas.', 'Si se unifican, se suman en una partida.'],
        ['Sala de reuniones', 'Se metran la refrigeradora E19 y el microondas E16 dibujados.', 'Si no van, se retiran.'],
    ], [3.2, 6.4, W - 9.6], tam=8.5, izq=(0, 1, 2), colores={(i, 0): AMBAR for i in range(6)})
    M.par('Cuando la arquitecta entregue las láminas corregidas (DXF), se leen de nuevo y se comparan con el metrado: si coinciden, los '
          'documentos quedan como están; si no, se regeneran con las cantidades del plano corregido.', despues=4)
    d.core_properties.author = ''
    d.core_properties.title = 'Equipamiento y mobiliario: observaciones subsanadas y pendientes - Hogar de Refugio'
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
