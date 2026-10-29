# -*- coding: utf-8 -*-
"""Nota interna (no va al expediente) del equipamiento y mobiliario del Hogar de Refugio Temporal (CUI 2675514): lectura del plano,
comprobación cuadros / planta / metrado por edificación, ajustes del metrado con su motivo, correcciones que la arquitecta debe hacer
en los cuadros de las láminas EM-01 a EM-07 y puntos para confirmar.

Uso: python3 herramientas/mobiliario_refugio/nota_mr.py
Salida: entregables/revision_interna/REVISION_EQUIPAMIENTO_Y_MOBILIARIO_HOGAR_REFUGIO.docx
"""
import os
import sys
from collections import defaultdict

import docx
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); HER = os.path.dirname(AQUI)
for d in (AQUI, HER):
    sys.path.insert(0, d)
import datos_mr as D  # noqa: E402
import incompatibilidades_refugio as IR  # noqa: E402  (formato de las notas internas)

SALIDA = os.path.join(D.SAL, 'revision_interna/REVISION_EQUIPAMIENTO_Y_MOBILIARIO_HOGAR_REFUGIO.docx')


def generar():
    P = D.plano(); A = D.ambientes(P); filas_c, err = D.comprobacion(P)
    R = IR.R
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('EQUIPAMIENTO Y MOBILIARIO: LECTURA DEL PLANO, AJUSTES Y CORRECCIONES A LOS CUADROS', 12.5, True, color=R.GRANATE,
          centro=True, despues=2)
    M.par('Hogar de Refugio Temporal - Mujeres Violentadas - CUI N.° %s  |  Documento de trabajo, no forma parte del expediente' % D.CUI,
          9, color=R.GRIS, centro=True, despues=10)
    n_f = sum(len(c['items']) for c in P['cuadros'])
    M.par('Se leyó el DXF de equipamiento y mobiliario (%d láminas EM-01 a EM-07 en el espacio modelo, %d cuadros de equipamiento con '
          '%d filas y %d etiquetas de código sobre el dibujo de las plantas, incluidas las de directrices y bloques) y se dibujó cada '
          'planta para revisarla ambiente por ambiente. El metrado toma lo dibujado y rotulado en la planta; los cuadros que no '
          'coinciden se ajustaron en el metrado y se deben corregir en las láminas, para que plano, metrado, memoria y '
          'especificaciones digan lo mismo.' % (len(P['laminas']), len(P['cuadros']), n_f, len(P['etiquetas'])))

    def sec(t):
        M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4)

    sec('1. COMPROBACIÓN POR EDIFICACIÓN (EQUIPOS Y MOBILIARIO)')
    tot = defaultdict(lambda: [0, 0, 0])
    for f in filas_c:
        if f['cod'][0] in 'EM':
            t = tot[f['ed']]; t[0] += f['cuadros']; t[1] += f['planta']; t[2] += f['adoptado']
    filas = [[ed, D.EDIFICACIONES[ed], D.LAMINA[ed], v[0], v[1], v[2], 'conforme' if v[0] == v[1] else 'cuadros corregidos']
             for ed, v in tot.items()]
    filas.append(['', 'TOTAL', '', sum(r[3] for r in filas), sum(r[4] for r in filas), sum(r[5] for r in filas), ''])
    M.tabla('Unidades de equipos y mobiliario: suma de los cuadros, rotulado en planta y metrado',
            ['EDIF.', 'EDIFICACIÓN', 'LÁMINA', 'CUADROS', 'PLANTA', 'METRADO', 'ESTADO'], filas,
            [1.2, W - 11.2, 1.6, 1.8, 1.8, 1.9, 2.9], tam=8, izq=(1,))
    M.par('El metrado coincide con la planta en todos los códigos de equipos y mobiliario de las nueve edificaciones%s. En los '
          'implementos rige el cuadro, salvo donde la planta rotula más piezas.' % ('' if not err else ' (revisar: %s)' % err), despues=4)

    sec('2. AJUSTES DEL METRADO RESPECTO DE LOS CUADROS')
    filas = []
    for ed, amb, its in A:
        for it in its:
            if it['nota']:
                filas.append([ed, D.cap(amb), it['cod'], it['cuadro'] or 0, it['cant'], it['nota']])
    for e, a, cod, q, motivo in D.CORRECCIONES:
        if q == 0:
            filas.append([e, D.cap(a), cod, 1, 0, motivo])
    filas.append(['E-02', 'Tópico (cuadro de EM-02)', 'todos', '13 filas', 0, 'Cuadro repetido: sus ítems ya están en el cuadro de E-02 '
                  '(almacén de donaciones, tópico y SS.HH.); no se suma.'])
    filas.append(['E-01, E-04', 'Pasadizos', 'S1', 2, 0, 'Extintores: partida de seguridad y señalización.'])
    M.tabla('Ajustes del metrado (cantidad del cuadro y cantidad del metrado)', ['EDIF.', 'AMBIENTE', 'CÓD.', 'CUADRO', 'METRADO', 'MOTIVO'],
            filas, [1.2, 3.4, 1.2, 1.5, 1.7, W - 9.0], tam=8, izq=(1, 5))

    sec('3. CORRECCIONES QUE SE DEBEN HACER EN LOS CUADROS DE LAS LÁMINAS')
    M.tabla('Correcciones a los cuadros de equipamiento', ['N.°', 'LÁMINA', 'CORRECCIÓN'],
            [[k, a, b] for k, (a, b) in enumerate(D.observaciones_plano(), 1)], [1.0, 2.4, W - 3.4], tam=8.5, izq=(2,))
    M.par('Además, en cada cuadro deben quedar las cantidades de la columna METRADO de la tabla 2 (las de la planta).', despues=4)

    sec('4. PARA CONFIRMAR CON LA ARQUITECTA')
    M.tabla('Puntos para confirmar', ['TEMA', 'QUÉ CONFIRMAR', 'POR QUÉ'], [
        ['Capacidad del SUM-comedor', 'La planta tiene 24 mesas M16 con 8 sillas cada una (192 asientos); la memoria de arquitectura '
         'considera el comedor/SUM para 80 personas.', 'Si la capacidad es 80, bastan unas 10 mesas y 80 sillas; el metrado sigue '
         'la planta hasta que se defina.'],
        ['Sillas del comedor', 'El comedor usa la silla regulable de estudiante M3 (0.30 x 0.27 x 0.32 m), pensada para niñas y niños.',
         'Para las residentes conviene una silla de adulto (por ejemplo M07 u otra de comedor).'],
        ['Máquinas del taller de costura', 'Las 10 máquinas E10 se describen como "máquina de corte textil"; en la planta están una '
         'por mesa, como estaciones de trabajo.', 'Si son máquinas de coser, cambiar la descripción (y la especificación).'],
        ['Implementos del almacén deportivo', 'Cantidad de K53 (colchoneta), K73 (kit de vóley), K74 (guantes de entrenamiento) y K79 '
         '(mica protectora), rotulados en la planta y ausentes del cuadro.', 'Se metró 1 de cada uno; la descripción se tomó del '
         'catálogo de códigos de la misma arquitectura (CAR Varones).'],
        ['Códigos repetidos', 'Dar otro código a la mesa de corte textil (M16 en costura) y a la canastilla de pelotas (M64 en el '
         'almacén deportivo).', 'Un código debe ser un solo elemento en plano, metrado y especificaciones.'],
        ['Mesas M19 y M22', 'Ambas son mesa de madera 1.00 x 1.80 x 0.75 m.', 'Unificar el código o diferenciar la descripción.'],
        ['Sala de reuniones', 'Tiene refrigeradora E19 y microondas E16 dibujados.', 'Confirmar que la sala tiene un área de apoyo (kitchenette).'],
    ], [3.2, 7.4, W - 10.6], tam=8.5, izq=(0, 1, 2))
    d.core_properties.author = ''
    d.core_properties.title = 'Equipamiento y mobiliario: lectura del plano y correcciones - Hogar de Refugio'
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True)
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
