# -*- coding: utf-8 -*-
"""Word del ESTUDIO DE GESTIÓN DE RIESGOS EN LA PLANIFICACIÓN DE LA EJECUCIÓN DE LA OBRA - SM-107 (CUI 2673852).
Uso: python3 estudio_docx.py   (antes: riesgos_sm107.py, matriz_xlsx.py, plano_riesgos_dxf.py y figuras_sm107.py)"""
import os, sys, json, collections
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import riesgos_sm107 as RS
RAIZ = RS.RAIZ; CALC = RS.CALC; FIG = os.path.join(CALC, "fig"); FOTO = os.path.join(CALC, "fotos")
SAL = os.path.join(RAIZ, "entregables_vial_sm107", "ESTUDIO_GESTION_RIESGOS_EJECUCION_OBRA_SM107.docx")
FT = "Arial"; AZUL = RGBColor(0x1F, 0x4E, 0x79)
D = json.load(open(os.path.join(CALC, "riesgos.json"))); MC = D["montecarlo"]; EJE = json.load(open(os.path.join(CALC, "eje_terreno.json")))
LL = json.load(open(os.path.join(CALC, "lluvia_ghcn.json")))
RS.calcular(); P = RS.P; R = RS.R
m = lambda v: "S/ {:,.2f}".format(v)
mm_ = lambda v: "S/ {:,.2f} M".format(v / 1e6)


# ------------------------------------------------------------------ utilidades
def fuente(r, tam=10, neg=False, cur=False, color=None):
    r.font.name = FT; r.font.size = Pt(tam); r.bold = neg; r.italic = cur
    if color: r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(k), FT)
    return r


def par(doc, texto="", tam=10, neg=False, al=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, despues=6, cur=False, color=None, sangria=0):
    p = doc.add_paragraph(); p.paragraph_format.alignment = al; p.paragraph_format.space_before = Pt(antes); p.paragraph_format.space_after = Pt(despues)
    p.paragraph_format.line_spacing = 1.15
    if sangria: p.paragraph_format.left_indent = Cm(sangria)
    if texto: fuente(p.add_run(texto), tam, neg, cur, color)
    return p


def mixto(doc, partes, tam=10, al=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = par(doc, al=al)
    for t, n in partes: fuente(p.add_run(t), tam, n)
    return p


def vin(doc, texto, nivel=0):
    p = par(doc, despues=3, sangria=0.9 + 0.6 * nivel); p.paragraph_format.first_line_indent = Cm(-0.45)
    fuente(p.add_run(("•  " if nivel == 0 else "–  ") + texto), 10); return p


CONT = collections.OrderedDict()


def tit(doc, num, texto, nivel=1):
    p = doc.add_paragraph(style="Heading %d" % nivel); p.paragraph_format.space_before = Pt(12 if nivel == 1 else 8); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    fuente(p.add_run(("%s. " % num if nivel == 1 else "%s " % num) + texto.upper() if nivel == 1 else "%s %s" % (num, texto)), 12 if nivel == 1 else 10.5, True, color=AZUL)
    CONT[num] = (texto, nivel)
    return p


def sombra(c, color):
    tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)


def tabla(doc, cab, filas, anchos, tam=8, colores=None, alinear=None):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(cab):
        c = t.rows[0].cells[j]; c.text = ""; fuente(c.paragraphs[0].add_run(h), tam, True, color=RGBColor(255, 255, 255)); sombra(c, "1F4E79")
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.rows[0]._tr; trPr = tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, fila in enumerate(filas):
        cells = t.add_row().cells
        for j, v in enumerate(fila):
            cells[j].text = ""; pp = cells[j].paragraphs[0]; fuente(pp.add_run(str(v)), tam)
            pp.alignment = (alinear[j] if alinear else WD_ALIGN_PARAGRAPH.LEFT)
            if colores and colores(i, j, v): sombra(cells[j], colores(i, j, v))
    t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(anchos) * 567))); tw.set(qn("w:type"), "dxa")
    for e in tblPr.findall(qn("w:tblW")): tblPr.remove(e)
    tblPr.append(tw)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos): gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for j, w in enumerate(anchos): row.cells[j].width = Cm(w)
    par(doc, despues=4)
    return t


NUMT = [0]; NUMF = [0]


def ctab(doc, texto):
    NUMT[0] += 1; p = par(doc, "Tabla N.° %02d: %s" % (NUMT[0], texto), 9, True, WD_ALIGN_PARAGRAPH.CENTER, antes=4, despues=3)
    p.paragraph_format.keep_with_next = True


def fig(doc, ruta, ancho, texto, fuente_txt):
    if not os.path.exists(ruta): return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(ruta, width=Cm(ancho))
    NUMF[0] += 1; par(doc, "Figura N.° %02d: %s" % (NUMF[0], texto), 9, True, WD_ALIGN_PARAGRAPH.CENTER, despues=1)
    par(doc, "Fuente: " + fuente_txt, 8, cur=True, al=WD_ALIGN_PARAGRAPH.CENTER, despues=8)


NIVCOL = {"ALTO": "FF7C80", "MODERADO": "FFD966", "BAJO": "A9D08E"}
colnivel = lambda i, j, v: NIVCOL.get(v)
C = WD_ALIGN_PARAGRAPH.CENTER; L_ = WD_ALIGN_PARAGRAPH.LEFT


def seccion(doc, horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = Cm(21.0), Cm(29.7)
    if horizontal: s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = h, w
    else: s.orientation = WD_ORIENT.PORTRAIT; s.page_width, s.page_height = w, h
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = Cm(2.0); s.bottom_margin = Cm(1.8)
    return s


def pie(sec, texto):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]; p.text = ""; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(p.add_run(texto + "   |   Página "), 7.5, color=RGBColor(0x59, 0x59, 0x59))
    r = p.add_run(); fuente(r, 7.5)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)


def main():
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = FT; st.font.size = Pt(10)
    s0 = doc.sections[0]; s0.page_width, s0.page_height = Cm(21.0), Cm(29.7)
    s0.left_margin = s0.right_margin = Cm(2.5); s0.top_margin = Cm(2.5); s0.bottom_margin = Cm(2.0)
    n_niv = collections.Counter(r["nivel"] for r in R); n_res = collections.Counter(r["nivel_r"] for r in R)
    n_as = collections.Counter(r["asig"] for r in R)
    I_, Rr = MC["inherente"], MC["residual"]
    reserva = Rr["entidad"]["80"]

    # ------------------------------------------------------------------ portada
    par(doc, "GOBIERNO REGIONAL SAN MARTÍN", 13, True, C, antes=10, despues=2, color=AZUL)
    par(doc, "DIRECCIÓN REGIONAL DE TRANSPORTES Y COMUNICACIONES", 11, True, C, despues=40, color=AZUL)
    par(doc, "ESTUDIO DE GESTIÓN DE RIESGOS EN LA PLANIFICACIÓN DE LA EJECUCIÓN DE LA OBRA", 17, True, C, despues=24)
    par(doc, "PROYECTO DE INVERSIÓN:", 10, True, C, despues=4)
    par(doc, "“" + P["nombre"] + "”", 12, True, C, despues=8)
    par(doc, "CUI N.° %s" % P["cui"], 12, True, C, despues=40)
    tabla(doc, ["DATO", "DESCRIPCIÓN"], [["Ubicación", "Distrito de Shamboyacu, provincia de Picota, departamento de San Martín"],
                                        ["Tramo", "Tingo de Ponasa (km 19+343.39) - Shamboyacu (km 44+693.39); 25.350 km según inventario vial"],
                                        ["Presupuesto de obra", m(P["PO"]) + " (costo al " + P["fecha_ppto"] + ")"],
                                        ["Costo total de inversión", m(P["TOTAL"])], ["Plazo de ejecución", "%d días calendario" % P["plazo"]],
                                        ["Modalidad", P["modalidad"]]], [4.5, 11.5], tam=9)
    par(doc, "OCTUBRE 2026", 11, True, C, antes=40)
    doc.add_page_break()

    # ------------------------------------------------------------------ contenido (se completa al final)
    par(doc, "CONTENIDO", 12, True, C, despues=8, color=AZUL)
    marca_contenido = doc.add_paragraph()
    doc.add_page_break()

    # ------------------------------------------------------------------ 1
    tit(doc, "1", "Aspectos generales")
    tit(doc, "1.1", "Objetivo", 2)
    par(doc, "Identificar, analizar y planificar la respuesta a los riesgos que pueden afectar el costo, el plazo y la calidad de la ejecución de la obra "
             "“%s”, CUI N.° %s, y asignar cada riesgo a la parte que mejor puede gestionarlo, de modo que la matriz de riesgos se incorpore al "
             "expediente de contratación y sirva para el seguimiento durante la obra." % (P["nombre"], P["cui"]))
    tit(doc, "1.2", "Alcance", 2)
    par(doc, "El estudio cubre la ejecución física de la obra (24 meses) y las acciones previas a la convocatoria que condicionan el inicio: libre disponibilidad "
             "del terreno, permisos, interferencias, canteras y depósitos de material excedente. Abarca los componentes del presupuesto: obras provisionales, "
             "demoliciones, movimiento de tierras, pavimento asfáltico, transporte, obras de arte y drenaje (alcantarillas, cunetas, muros y gaviones), "
             "señalización, cinco puentes y el plan de manejo ambiental. No reemplaza el plan de seguridad y salud en el trabajo, el plan de manejo ambiental ni "
             "el plan de contingencias, con los que se coordina.")
    tit(doc, "1.3", "Diferencia con la evaluación de riesgo de desastres (EVAR)", 2)
    par(doc, "Son dos estudios distintos y no se reemplazan entre sí:")
    tabla(doc, ["ASPECTO", "GESTIÓN DE RIESGOS EN LA EJECUCIÓN DE OBRA (este estudio)", "EVALUACIÓN DE RIESGO DE DESASTRES (EVAR)"],
          [["Norma", "Ley N.° 32069 y su Reglamento (D.S. N.° 009-2025-EF), artículo 156", "Ley N.° 29664 (SINAGERD) y Manual del CENEPRED (R.J. N.° 112-2014-CENEPRED/J)"],
           ["Pregunta", "¿Qué puede salir mal al construir la obra y quién lo gestiona?", "¿Qué peligros naturales amenazan a la población y a la infraestructura?"],
           ["Riesgos", "Técnicos, de diseño, prediales, permisos, sociales, económicos, clima y naturaleza durante la obra, seguridad", "Peligros naturales (inundación, movimientos en masa, sismo) y vulnerabilidad"],
           ["Método", "Probabilidad x impacto en costo y plazo; valor monetario esperado; Monte Carlo", "Peligro x vulnerabilidad con ponderación de parámetros (Saaty)"],
           ["Producto", "Matriz de riesgos con responsables, asignación y plan de respuesta", "Mapas de peligro, vulnerabilidad y riesgo, y medidas de control"],
           ["Momento", "Expediente técnico y de contratación; seguimiento en obra", "Formulación del proyecto y planificación territorial"],
           ["Elabora", "Área usuaria con la dependencia encargada de las contrataciones", "Evaluador de riesgo acreditado por el CENEPRED"]],
          [3.0, 6.6, 6.6], tam=8)
    par(doc, "Los peligros naturales que pueden ocurrir durante la obra (lluvias, crecidas, deslizamientos, sismo) sí se incluyen en este estudio, pero "
             "solo por su efecto en la ejecución.")
    tit(doc, "1.4", "Datos del proyecto", 2)
    ctab(doc, "Datos generales del proyecto")
    tabla(doc, ["CONCEPTO", "VALOR", "FUENTE"],
          [["CUI / código de idea", "%s / %s" % (P["cui"], P["idea"]), "Inventario vial / resumen ejecutivo de la FTE"],
           ["Entidad", P["entidad"], "FTE"], ["Unidad ejecutora presupuestal", P["ue"], "FTE"],
           ["Ruta y tramo", "SM-107, Tingo de Ponasa - Shamboyacu", "Inventario vial"],
           ["Progresivas", "19+343.39 a 44+693.39 (25.350 km)", "Inventario vial"],
           ["Situación actual", "Calzada de 4.50 a 6.20 m, tratamiento superficial tipo slurry deteriorado, material de préstamo y terreno natural", "Inventario vial"],
           ["Situación con proyecto", "Calzada de 6.00 m y bermas de 0.50 m con carpeta asfáltica en caliente", "FTE"],
           ["IMDA actual / proyectado", "%d / %d veh/día" % (P["imda_actual"], P["imda_proy"]), "FTE"],
           ["Población beneficiaria directa", "{:,} habitantes".format(P["poblacion"]), "FTE"],
           ["Costo directo", m(P["CD"]), "Presupuesto S10"], ["Presupuesto de obra (con IGV)", m(P["PO"]), "Presupuesto S10"],
           ["Costo total de inversión", m(P["TOTAL"]), "Presupuesto S10"], ["Plazo de ejecución", "%d días calendario" % P["plazo"], "Resumen ejecutivo"],
           ["Modalidad y financiamiento", "%s; %s" % (P["modalidad"], P["fuente_fin"].lower()), "FTE"]], [4.2, 7.2, 4.6], tam=8)

    # ------------------------------------------------------------------ 2 marco normativo
    tit(doc, "2", "Marco normativo")
    par(doc, "Normas de contratación pública (vigentes desde el 22 de abril de 2025):", neg=True)
    vin(doc, "Ley N.° 32069, Ley General de Contrataciones Públicas.")
    vin(doc, "Reglamento de la Ley N.° 32069, aprobado por D.S. N.° 009-2025-EF, y sus modificatorias. El artículo 156 regula la gestión de riesgos en obras: "
             "en las obras segmentadas como avanzadas la matriz de riesgos forma parte del expediente de contratación y contiene, como mínimo, la "
             "identificación de los riesgos, su responsable y su análisis cualitativo y cuantitativo; la entidad asigna cada riesgo a la parte que puede "
             "gestionarlo y mitigarlo; y, cuando la cuantía supera los S/ 50 000 000, conforma un comité de gestión de riesgos que acompaña las fases de la "
             "contratación, aprueba los riesgos y oportunidades nuevos y sus respuestas, y hace el seguimiento de las alertas tempranas.")
    vin(doc, "Comunicado N.° 016-2025-OECE: la aplicación de los artículos 128 y 156 del reglamento es obligatoria y no depende de la emisión de la directiva "
             "de la Dirección General de Abastecimiento.")
    vin(doc, "La Directiva N.° 012-2017-OSCE/CD, Gestión de riesgos en la planificación de la ejecución de obras, fue derogada por la Resolución "
             "N.° D000129-2025-OECE-PRE. En este estudio solo se usa como referencia de método (escalas de probabilidad e impacto y formatos), por ser "
             "compatible con el reglamento vigente. Solo rige para los procedimientos convocados antes del 22 de abril de 2025.")
    vin(doc, "Directiva de la Dirección General de Abastecimiento sobre gestión de riesgos anunciada en el artículo 156: a la fecha de este estudio no se ha "
             "identificado su publicación. Si se publica antes de la convocatoria, la matriz debe adecuarse a sus formatos.")
    par(doc, "Normas técnicas y sectoriales:", neg=True, antes=4)
    for t in ["Manual de Carreteras: Especificaciones Técnicas Generales para Construcción EG-2013 (R.D. N.° 22-2013-MTC/14).",
              "Manual de Carreteras: Diseño Geométrico DG-2018 (R.D. N.° 03-2018-MTC/14).",
              "Manual de Puentes (R.D. N.° 19-2018-MTC/14).",
              "Manual de Carreteras: Suelos, Geología, Geotecnia y Pavimentos (R.D. N.° 10-2014-MTC/14).",
              "Manual de Hidrología, Hidráulica y Drenaje (R.D. N.° 20-2011-MTC/14).",
              "Manual de Dispositivos de Control del Tránsito Automotor para Calles y Carreteras (R.D. N.° 26-2024-MTC/18).",
              "Manual de Inventarios Viales (R.D. N.° 09-2014-MTC/14 y R.D. N.° 022-2015-MTC/14)."]: vin(doc, t)
    par(doc, "Otras normas que intervienen en la respuesta a los riesgos:", neg=True, antes=4)
    for t in ["Ley N.° 29664, Sistema Nacional de Gestión del Riesgo de Desastres, y su reglamento (D.S. N.° 048-2011-PCM): respuesta ante emergencias durante la obra.",
              "Ley N.° 29783, Ley de Seguridad y Salud en el Trabajo, y D.S. N.° 011-2019-TR, Reglamento de Seguridad y Salud en el Trabajo para el Sector Construcción.",
              "Ley N.° 27446, Sistema Nacional de Evaluación del Impacto Ambiental, y su reglamento.",
              "D.S. N.° 011-2022-MC, Reglamento de Intervenciones Arqueológicas.",
              "Ley N.° 30025, que facilita la adquisición, expropiación y posesión de bienes inmuebles para obras de infraestructura.",
              "Ley N.° 29338, Ley de Recursos Hídricos: autorizaciones de uso de agua y de extracción de material de acarreo en cauces."]: vin(doc, t)
    par(doc, "Referencias de método: PMI, Guía del PMBOK y The Standard for Risk Management in Portfolios, Programs and Projects; ISO 31000:2018, Gestión del riesgo.",
        cur=True, antes=4)

    # ------------------------------------------------------------------ 3 descripcion
    tit(doc, "3", "Descripción del proyecto y de su entorno")
    tit(doc, "3.1", "Ubicación y trazo", 2)
    par(doc, "La carretera departamental SM-107 empalma con la ruta nacional PE-5N en Picota. El tramo del proyecto empieza en Tingo de Ponasa (km 19+343.39) y "
             "termina en Shamboyacu (km 44+693.39). Atraviesa Aypena, Nuevo América, Leoncio Prado, San Antonio, Huañipo, Paucar y Alfonso Ugarte, y se conecta "
             "con las rutas vecinales SM-803 y SM-806 en Tingo de Ponasa y con la SM-804 en Shamboyacu.")
    ctrl = EJE["control"]
    par(doc, "Para ubicar los riesgos se georreferenció el eje del plano clave PC-01 con su cuadrícula UTM (WGS 84, zona 18 S). Se comprobó con los puntos "
             "geodésicos y BM del inventario vial: %s de ellos quedan a menos de 12 m del eje. Las progresivas que resultan coinciden con las del inventario "
             "(BM Huañipo en 33+803 frente a 33+780 de la plaza; BM Aypena en 21+993 frente a 21+940). El eje del PC-01 mide %.2f km hasta el km 44+090; "
             "el tramo final hasta 44+693.39 se completó en línea recta (referencial)." %
        (sum(1 for c in ctrl if c[1] < 12), EJE["longitud_pc01_m"] / 1000))
    fig(doc, os.path.join(FIG, "f1_mapa.png"), 16.0, "Ubicación del tramo y de los sectores y puntos con riesgo identificado",
        "eje del plano clave PC-01 georreferenciado; sectores del estudio de inventario vial; relieve del modelo Copernicus GLO-30.")
    tit(doc, "3.2", "Componentes y costos de la obra", 2)
    comp = [("Obras provisionales", 626158.06), ("Seguridad y salud ocupacional", 268294.80), ("Demoliciones y desmontajes", 7286663.26),
            ("Obras preliminares", 1148377.50), ("Movimiento de tierras", 13693530.47), ("Pavimentos", 17951289.35), ("Transporte", 25602427.13),
            ("Obras de arte y drenaje", 9537008.29), ("Señalización y seguridad vial", 1089422.57), ("Puentes (5)", 9869727.35),
            ("Plan de manejo ambiental", 1177898.83), ("Plan arqueológico", 12000.00)]
    ctab(doc, "Componentes del costo directo (presupuesto S10)")
    tabla(doc, ["COMPONENTE", "COSTO DIRECTO (S/)", "% DEL CD"], [[c, "{:,.2f}".format(v), "%.1f %%" % (v / P["CD"] * 100)] for c, v in comp] +
          [["COSTO DIRECTO", "{:,.2f}".format(P["CD"]), "100.0 %"]], [8.5, 4.0, 2.5], tam=8, alinear=[L_, WD_ALIGN_PARAGRAPH.RIGHT, C])
    par(doc, "El transporte (29.0 %) y el movimiento de tierras (15.5 %) suman casi la mitad del costo directo: los riesgos de canteras, distancias de "
             "transporte, depósitos de material excedente y suelos de subrasante tienen el mayor peso económico.")
    ctab(doc, "Puentes a reemplazar")
    tabla(doc, ["PUENTE", "PROGRESIVA", "LUZ", "COSTO DIRECTO (S/)", "PUENTE EXISTENTE (INVENTARIO)"],
          [["Tingo de Ponasa", "20+300", "40.00 m", "3,619,664.38", "20+300 a 20+337, viga y losa, h = 8.00 m"],
           ["Barco", "29+300", "40.00 m", "3,687,055.05", "29+329 a 29+360, L = 31 m, señal de carga máx. 30 t"],
           ["Paucarillo", "37+870", "20.00 m", "1,092,279.51", "37+960 a 37+974, L = 14 m"],
           ["Apangurayco", "38+800", "15.00 m", "763,556.43", "38+919 a 38+929, L = 10 m"],
           ["Chanchamayco", "44+040", "12.00 m", "707,171.98", "44+149 a 44+158, pontón L = 9 m"]], [3.0, 2.2, 1.8, 3.0, 6.0], tam=8)
    tit(doc, "3.3", "Estado de la vía y puntos críticos", 2)
    par(doc, "El inventario vial registra, entre otros, 5 puentes, 50 alcantarillas (varias colmatadas o semicolmatadas), 28 badenes, 6 cruces de canal de "
             "regadío, 30 señales y 26 puntos o sectores críticos. La calzada tiene un solo carril de 4.50 a 6.20 m, con encalaminado, baches, pérdida del "
             "tratamiento superficial y empozamientos. Los puntos críticos que interesan a la ejecución son:")
    ctab(doc, "Puntos y sectores críticos del inventario vial")
    pcs = [["20+337 a 20+340", "Erosión h = 2.00 m (LI)"], ["20+360 a 21+170", "Inundación por cultivos de arroz"], ["22+270; 22+800 a 22+960", "Empozamiento de agua (LI)"],
           ["26+924; 34+954; 40+380; 43+660", "Empozamiento sin elemento de descarga"], ["28+800 a 29+060", "Material en la vía (LI)"],
           ["34+766 a 34+786", "Erosión h = 3 m (LD)"], ["36+060 a 36+100", "Inestabilidad de talud h = 20 m (LI)"], ["36+380 a 36+460", "Erosión h = 2.0 m (LD)"],
           ["37+520 a 37+780", "Inestabilidad de talud h = 30 m (LI)"], ["39+380 a 39+840", "Erosión h = 3 a 5 m (LD y LI)"],
           ["42+280", "Quebrada sin descarga, pendiente y contrapendiente, material de arrastre"], ["43+620 a 43+632", "Erosión h = 1.0 m (LI)"],
           ["34+400 a 44+600 (varios)", "Plataforma angosta"]]
    tabla(doc, ["PROGRESIVA", "DESCRIPCIÓN"], pcs, [5.0, 11.0], tam=8)
    for f_, t_ in (("t20_1.jpg", "Erosión del borde de la plataforma en el km 37+223 (UTM 369850 E, 9226904 N)"),
                   ("t24_1.jpg", "Empozamiento y pérdida de rodadura en el km 42+593 (UTM 373093 E, 9223667 N)")):
        fig(doc, os.path.join(FOTO, f_), 11.0, t_, "estudio de inventario vial, panel fotográfico (setiembre 2026); progresiva obtenida con el eje georreferenciado.")
    tit(doc, "3.4", "Relieve", 2)
    T = EJE["tramos_100m"]
    plano = [t for t in T if t["pi"] < 23000]; ondu = [t for t in T if t["pi"] >= 23000]
    par(doc, "El perfil del terreno se obtuvo del modelo digital Copernicus GLO-30 (30 m), que incluye la altura de la vegetación; sus valores son "
             "referenciales. Del inicio al km 23 el terreno es llano (cotas de %.0f a %.0f msnm, pendientes del entorno menores de 15°), con arrozales y "
             "zonas inundables. Desde el km 23 el terreno es ondulado: cotas de %.0f a %.0f msnm, relieve local de 20 a 90 m y laderas de hasta %.0f°, "
             "con los valores mayores entre los km 34 y 37 y entre los km 41 y 44. Allí están los taludes inestables y las erosiones del inventario. "
             "La FTE clasifica la orografía como plana, lo que no corresponde a la segunda mitad del tramo." %
        (min(t["z"] for t in plano), max(t["z"] for t in plano), min(t["z"] for t in ondu), max(t["z"] for t in ondu), max(t["pend_max_grados"] for t in ondu)))
    fig(doc, os.path.join(FIG, "f2_perfil.png"), 16.0, "Perfil del terreno y pendiente máxima del entorno (franja de 150 m a cada lado)",
        "modelo digital Copernicus GLO-30; sectores críticos del estudio de inventario vial.")
    tit(doc, "3.5", "Clima", 2)
    t_, j_ = LL["Tarapoto"], LL["Juanjuí"]
    par(doc, "No se cuenta con una serie pluviométrica en el tramo. Como referencia se usaron las estaciones más cercanas de la base NOAA GHCN-Daily: "
             "Tarapoto (56 km al norte; %d mm al año) y Juanjuí (60 km al suroeste; %d mm al año). Llueve todo el año, con un máximo en marzo y otro en "
             "octubre-noviembre; los meses más secos son julio, agosto y setiembre. Las series son cortas y con vacíos: el estudio hidrológico del "
             "expediente debe prevalecer para el diseño." % (round(sum(v[0] for v in t_["meses"].values())), round(sum(v[0] for v in j_["meses"].values()))))
    fig(doc, os.path.join(FIG, "f3_lluvia.png"), 15.0, "Lluvia mensual de referencia y temporadas para la programación de la obra",
        "NOAA GHCN-Daily, estaciones PE000084455 (Tarapoto) y PEM00084474 (Juanjuí).")

    # ------------------------------------------------------------------ 4 metodologia
    tit(doc, "4", "Metodología")
    par(doc, "Se siguió el proceso del artículo 156 del reglamento, con las escalas del PMBOK que usaban los formatos de la ex Directiva N.° 012-2017-OSCE/CD:")
    for t in ["Identificación: revisión del inventario vial, la FTE, su resumen ejecutivo, el presupuesto S10 y el plano clave; recorrido del eje georreferenciado "
              "con el relieve y la lluvia de la zona; lista de verificación por categorías (diseño, terreno, clima, predios, permisos, interferencias, sociales, "
              "económicos, suministros, ejecución, seguridad y ambiente).",
              "Análisis cualitativo: probabilidad (P) e impacto (I) con las escalas de la Tabla N.° %02d; el impacto se obtiene del valor más probable de costo y "
              "de plazo de cada riesgo (Tabla N.° %02d) y se toma el mayor. El nivel resulta de P x I." % (NUMT[0] + 1, NUMT[0] + 2),
              "Análisis cuantitativo: costo y retraso de cada riesgo con distribución triangular (mínimo, más probable, máximo); valor monetario esperado (VME); "
              "simulación de Monte Carlo para obtener la reserva de contingencia y el retraso probable.",
              "Respuesta: estrategia (evitar, mitigar, transferir o aceptar), acciones, responsable, partida que la financia, momento y alerta temprana.",
              "Asignación: a la Entidad, al contratista o compartido, según quién controla la causa.",
              "Seguimiento: comité de gestión de riesgos, revisión mensual y actualización de la matriz."]: vin(doc, t)
    ctab(doc, "Escala de probabilidad")
    tabla(doc, ["NIVEL", "VALOR", "DESCRIPCIÓN"], [[t, "%.2f" % v, d_] for (t, v), d_ in zip(RS.PROB, ["Solo en circunstancias excepcionales", "Podría ocurrir alguna vez durante la obra",
                                                                                           "Puede ocurrir durante la obra", "Es probable que ocurra", "Se espera que ocurra"])],
          [3.0, 2.0, 11.0], tam=8, alinear=[L_, C, L_])
    ctab(doc, "Escala de impacto (valor más probable)")
    lim_c = ["< %s" % mm_(RS.IMPACTO_COSTO[0] * P["PO"])] + ["%s a %s" % (mm_(RS.IMPACTO_COSTO[i - 1] * P["PO"]), mm_(RS.IMPACTO_COSTO[i] * P["PO"])) for i in range(1, 4)] + ["> %s" % mm_(RS.IMPACTO_COSTO[3] * P["PO"])]
    lim_p = ["< 7 días", "7 a 21 días", "21 a 45 días", "45 a 90 días", "> 90 días"]
    pct = ["< 0.1 %", "0.1 a 0.5 %", "0.5 a 1.5 %", "1.5 a 4.0 %", "> 4.0 %"]
    tabla(doc, ["NIVEL", "VALOR", "COSTO (% DEL PRESUPUESTO DE OBRA)", "COSTO (S/)", "PLAZO"],
          [[t, "%.2f" % v, pct[i], lim_c[i], lim_p[i]] for i, (t, v) in enumerate(RS.IMP)], [2.6, 1.6, 4.4, 4.6, 2.8], tam=8, alinear=[L_, C, C, C, C])
    ctab(doc, "Niveles de riesgo y tratamiento")
    tabla(doc, ["NIVEL", "P x I", "TRATAMIENTO"], [["ALTO", ">= 0.18", "Respuesta obligatoria antes del inicio o de la actividad afectada; seguimiento mensual del comité"],
                                                  ["MODERADO", "0.05 a 0.18", "Respuesta planificada; seguimiento del residente y del supervisor"],
                                                  ["BAJO", "< 0.05", "Se acepta; se vigila con el registro"]], [3.0, 2.5, 10.5], tam=8, colores=colnivel, alinear=[C, C, L_])
    par(doc, "Montos a precio de obra: el costo directo se multiplica por %.4f (gastos generales, utilidad e IGV). Cada día de ampliación atribuible a la "
             "Entidad se valoriza en %s (gastos generales con IGV y supervisión por día). En la simulación, los retrasos de la ruta crítica se agrupan "
             "por etapa (inicio, clima, terreno y gestión): dentro de cada grupo se suma el mayor retraso y el %d %% de los demás, porque actúan sobre los "
             "mismos frentes y períodos; los grupos se suman entre sí." % (P["K"], m(P["costo_dia"]), int(D["concurrencia"] * 100)))

    # ------------------------------------------------------------------ 5 identificacion
    tit(doc, "5", "Identificación de riesgos")
    cats = collections.Counter(r["cat"] for r in R)
    par(doc, "Se identificaron %d riesgos y %d oportunidades. Por categoría: %s. La descripción completa (causa, efecto, ubicación y evidencia en los documentos) "
             "está en la hoja IDENTIFICACION de la matriz en Excel y en el Anexo 1." % (len(R), len(RS.OPORT), "; ".join("%s (%d)" % (k.lower(), v) for k, v in cats.most_common())))
    ctab(doc, "Riesgos identificados")
    tabla(doc, ["COD.", "CATEGORÍA", "RIESGO", "UBICACIÓN"], [[r["cod"], r["cat"], r["riesgo"], r["ubic"]] for r in R], [1.4, 2.6, 7.0, 5.0], tam=7.5)

    # ------------------------------------------------------------------ 6 cualitativo
    tit(doc, "6", "Análisis cualitativo")
    par(doc, "Resultado inherente (sin respuesta): %d riesgos de nivel ALTO, %d MODERADO y %d BAJO. Con las respuestas del capítulo 9 ejecutadas en el momento "
             "indicado (residual): %d ALTO, %d MODERADO y %d BAJO." % (n_niv["ALTO"], n_niv["MODERADO"], n_niv["BAJO"], n_res["ALTO"], n_res["MODERADO"], n_res["BAJO"]))
    ctab(doc, "Análisis cualitativo P x I")
    filas = [[r["cod"], r["riesgo"], r["P_txt"], r["I_txt"], "%.3f" % r["PxI"], r["nivel"], r["Pr_txt"], "%.3f" % r["PxIr"], r["nivel_r"]] for r in sorted(R, key=lambda r: (-r["PxI"], r["cod"]))]
    tabla(doc, ["COD.", "RIESGO", "P", "I", "P x I", "NIVEL", "P RES.", "P x I RES.", "NIVEL RES."], filas, [1.1, 4.6, 1.6, 1.5, 1.0, 1.9, 1.6, 1.0, 1.9], tam=7,
          colores=lambda i, j, v: NIVCOL.get(v) if j in (5, 8) else None, alinear=[C, L_, C, C, C, C, C, C, C])
    fig(doc, os.path.join(FIG, "f4_calor.png"), 16.5, "Mapa de calor de los riesgos, inherente y residual", "matriz de gestión de riesgos (Excel).")
    altos = [r for r in R if r["nivel"] == "ALTO"]
    par(doc, "Riesgos de nivel ALTO:", neg=True)
    for r in altos: vin(doc, "%s %s. %s" % (r["cod"], r["riesgo"], r["causa"]))

    # ------------------------------------------------------------------ 7 cuantitativo
    tit(doc, "7", "Análisis cuantitativo")
    tit(doc, "7.1", "Valor monetario esperado", 2)
    par(doc, "El valor monetario esperado (VME) de cada riesgo es su probabilidad por el valor medio de su distribución triangular. La tabla muestra los "
             "riesgos con mayor VME de costo; la lista completa, con fórmulas, está en la hoja ANALISIS CUANTITATIVO.")
    top = sorted(R, key=lambda r: -r["EMV_costo"])[:15]
    ctab(doc, "Riesgos con mayor valor monetario esperado")
    tabla(doc, ["COD.", "RIESGO", "ASIGNADO", "P", "COSTO MÍN. - MÁS PROB. - MÁX. (S/ miles)", "VME COSTO (S/)", "VME PLAZO (días)"],
          [[r["cod"], r["riesgo"], r["asig"], "%.2f" % r["P"], "%s - %s - %s" % tuple("{:,.0f}".format(v / 1000) for v in r["costo"]), "{:,.0f}".format(r["EMV_costo"]),
            "%.1f" % r["EMV_plazo"]] for r in top], [1.2, 5.2, 1.9, 0.9, 3.2, 2.0, 1.6], tam=7, alinear=[C, L_, C, C, C, WD_ALIGN_PARAGRAPH.RIGHT, C])
    tit(doc, "7.2", "Simulación de Monte Carlo", 2)
    par(doc, "Se corrieron 20,000 iteraciones. En cada una, cada riesgo ocurre o no según su probabilidad y, si ocurre, su costo y su retraso se toman de su "
             "distribución triangular. Para la Entidad se suman los riesgos que le corresponden, la mitad de los compartidos (descontando lo que cubre la "
             "póliza CAR en R-02 y R-22) y el costo de los días de ampliación que le son atribuibles.")
    ctab(doc, "Resultados de la simulación")
    f2 = lambda q: "{:,.2f}".format(q / 1e6)
    tabla(doc, ["RESULTADO", "P50", "P80", "P90", "MEDIA"],
          [["Costo para la Entidad, inherente (S/ M)", f2(I_["entidad"]["50"]), f2(I_["entidad"]["80"]), f2(I_["entidad"]["90"]), f2(I_["entidad_media"])],
           ["Costo para la Entidad, residual (S/ M)", f2(Rr["entidad"]["50"]), f2(Rr["entidad"]["80"]), f2(Rr["entidad"]["90"]), f2(Rr["entidad_media"])],
           ["Retraso de la obra, inherente (días)", "%.0f" % I_["dias"]["50"], "%.0f" % I_["dias"]["80"], "%.0f" % I_["dias"]["90"], "%.0f" % I_["dias_media"]],
           ["Retraso de la obra, residual (días)", "%.0f" % Rr["dias"]["50"], "%.0f" % Rr["dias"]["80"], "%.0f" % Rr["dias"]["90"], "%.0f" % Rr["dias_media"]],
           ["Retraso atribuible a la Entidad, residual (días)", "%.0f" % Rr["diasE"]["50"], "%.0f" % Rr["diasE"]["80"], "%.0f" % Rr["diasE"]["90"], "—"],
           ["Costo de riesgos del contratista, residual (S/ M)", f2(Rr["contratista"]["50"]), f2(Rr["contratista"]["80"]), f2(Rr["contratista"]["90"]), f2(Rr["contratista_media"])]],
          [7.6, 2.0, 2.0, 2.0, 2.0], tam=8, alinear=[L_, C, C, C, C])
    fig(doc, os.path.join(FIG, "f5_montecarlo.png"), 16.5, "Curvas de probabilidad acumulada del costo para la Entidad y del retraso", "simulación de Monte Carlo, matriz de gestión de riesgos.")
    par(doc, "Lectura de los resultados:", neg=True)
    vin(doc, "Sin respuesta, la exposición de la Entidad al 80 %% de confianza es de %s (%.1f %% del presupuesto de obra) y la obra podría retrasarse unos %.0f días." %
        (mm_(I_["entidad"]["80"]), I_["entidad"]["80"] / P["PO"] * 100, I_["dias"]["80"]))
    vin(doc, "Con las respuestas ejecutadas, baja a %s (%.1f %%) y el retraso a unos %.0f días. La mayor reducción viene de las acciones previas a la convocatoria "
             "(canteras, DME, certificación ambiental, predios, interferencias y revisión del expediente)." % (mm_(Rr["entidad"]["80"]), Rr["entidad"]["80"] / P["PO"] * 100, Rr["dias"]["80"]))
    vin(doc, "Reserva de contingencia recomendada para la Entidad: %s (P80 residual, %.1f %% del presupuesto de obra), para cubrir prestaciones adicionales, "
             "mayores metrados, reajustes y ampliaciones de plazo derivados de los riesgos que asume." % (mm_(reserva), reserva / P["PO"] * 100))
    vin(doc, "El retraso residual (P80 de %.0f días) se debe sobre todo a las lluvias y a la gestión de la obra; para no superar los %d días hay que programar los "
             "puentes, terraplenes y carpeta en estiaje y abrir frentes paralelos (oportunidad O-02)." % (Rr["dias"]["80"], P["plazo"]))

    # ------------------------------------------------------------------ 8 observaciones
    tit(doc, "8", "Observaciones a los documentos del proyecto")
    par(doc, "La revisión cruzada del inventario vial, la FTE, su resumen ejecutivo y el presupuesto S10 encontró las siguientes incompatibilidades. Cada una es "
             "una causa del riesgo R-16 o de otro riesgo, y debe levantarse antes de la convocatoria:")
    ctab(doc, "Observaciones que generan riesgo")
    tabla(doc, ["N.°", "ASPECTO", "DATO EN LOS DOCUMENTOS", "OBSERVACIÓN", "RECOMENDACIÓN", "RIESGO"],
          [[str(i), o[0], o[1], o[2], o[3], o[4]] for i, o in enumerate(RS.OBS, 1)], [0.8, 2.4, 3.8, 4.2, 3.6, 1.4], tam=7, alinear=[C, L_, L_, L_, L_, C])

    # ------------------------------------------------------------------ 9 respuesta
    tit(doc, "9", "Plan de respuesta")
    tit(doc, "9.1", "Acciones previas a la convocatoria", 2)
    par(doc, "Son las acciones de mayor efecto, porque eliminan o reducen riesgos de la Entidad antes de que exista un contrato:")
    previas = [r for r in R if "convocatoria" in r["momento"].lower() or "antes del inicio" in r["momento"].lower()]
    for r in sorted(previas, key=lambda r: -r["PxI"]): vin(doc, "%s (%s): %s" % (r["cod"], r["nivel"], r["acc"]))
    tit(doc, "9.2", "Programación según la temporada de lluvias", 2)
    for t in ["Estiaje (junio a setiembre): cimentaciones y estribos de los 5 puentes, terraplenes en los sectores de arrozal (20+360 a 21+170 y 22+270 a 22+960), base y carpeta asfáltica.",
              "Transición (mayo, diciembre a febrero): cortes con drenaje de coronación, subbase y base con protección, desvíos de cauce y fabricación de vigas metálicas.",
              "Lluviosa (marzo, abril, octubre y noviembre): alcantarillas, cunetas, gaviones y muros, demoliciones menores, señalización y trabajos en taller."]: vin(doc, t)
    tit(doc, "9.3", "Plan de respuesta por riesgo", 2)
    ctab(doc, "Plan de respuesta")
    tabla(doc, ["COD.", "NIVEL", "ESTRATEGIA", "ACCIONES", "RESPONSABLE", "ALERTA TEMPRANA"],
          [[r["cod"], r["nivel"], r["resp"], r["acc"], r["quien"], r["alerta"]] for r in R], [1.1, 1.9, 1.8, 5.5, 2.5, 3.2], tam=7,
          colores=lambda i, j, v: NIVCOL.get(v) if j == 1 else None, alinear=[C, C, C, L_, L_, L_])

    # ------------------------------------------------------------------ 10 asignacion
    tit(doc, "10", "Asignación de riesgos")
    par(doc, "Se asigna cada riesgo a la parte que controla su causa y puede gestionarlo y mitigarlo a menor costo (artículo 156 del reglamento): "
             "%d a la Entidad, %d al contratista y %d compartidos." % (n_as["Entidad"], n_as["Contratista"], n_as["Compartido"]))
    for k, t in (("Entidad", "decisiones, permisos, predios, interferencias, financiamiento, canteras y DME definidos en el expediente, y el contenido del expediente "
                             "técnico; si ocurren, proceden las modificaciones contractuales o la ampliación de plazo que correspondan."),
                 ("Contratista", "planificación, métodos constructivos, equipos, suministros, calidad, seguridad, tránsito durante la obra y lluvias estacionales "
                                 "previsibles; se cubren con sus gastos generales, su programación y sus pólizas (CAR, responsabilidad civil y SCTR)."),
                 ("Compartido", "causas mixtas (taludes, saturación por riego, conflictos y canales): la Entidad asume la parte de diseño y gestión, y el "
                                "contratista la de ejecución, según la causa que se compruebe en obra.")):
        mixto(doc, [(k + ": ", True), (t, False)])
    ctab(doc, "Asignación de riesgos")
    tabla(doc, ["COD.", "RIESGO", "ASIGNADO A"], [[r["cod"], r["riesgo"], r["asig"]] for r in R], [1.4, 11.6, 3.0], tam=7.5, alinear=[C, L_, C])

    # ------------------------------------------------------------------ 11 oportunidades
    tit(doc, "11", "Oportunidades")
    ctab(doc, "Oportunidades identificadas")
    tabla(doc, ["COD.", "OPORTUNIDAD", "ESTRATEGIA", "ACCIONES", "VME AHORRO"],
          [[o["cod"], o["oport"], o["resp"], o["acc"], ("{:,.0f}".format(o["EMV"]) if o["EMV"] else "%.0f días" % o["EMV_plazo"])] for o in RS.OPORT],
          [1.2, 4.4, 1.8, 6.6, 2.0], tam=8)

    # ------------------------------------------------------------------ 12 seguimiento
    tit(doc, "12", "Organización y seguimiento")
    tit(doc, "12.1", "Comité de gestión de riesgos", 2)
    par(doc, "La cuantía de la obra (%s) supera los S/ 50 000 000, por lo que la Entidad debe conformar el comité de gestión de riesgos del artículo 156 del "
             "reglamento. Se propone integrarlo con:" % m(P["PO"]))
    for t in ["El director de caminos de la DRTC San Martín o su representante, que lo preside.", "Un representante del área usuaria y uno de la dependencia encargada de las contrataciones.",
              "El inspector o el supervisor de obra.", "Como invitados: el residente de obra y los especialistas que se requieran (geotecnia, puentes, ambiental, social)."]: vin(doc, t)
    par(doc, "Funciones: aprobar la matriz al inicio de la obra; revisar mensualmente los riesgos de nivel ALTO y MODERADO y las alertas tempranas; aprobar los riesgos "
             "y oportunidades nuevos y sus respuestas; decidir el uso de la reserva de contingencia; y registrar sus acuerdos en acta.")
    tit(doc, "12.2", "Ciclo de seguimiento", 2)
    for t in ["Cada semana: el residente revisa las alertas tempranas de los riesgos de su frente y las anota en el cuaderno de obra.",
              "Cada mes, con la valorización: el supervisor actualiza la hoja SEGUIMIENTO de la matriz (estado, probabilidad e impacto actuales, acción ejecutada) "
              "y la presenta al comité.",
              "Ante una alerta activada o un riesgo ocurrido: comunicación al supervisor en el día, aplicación del plan de respuesta y registro del costo y el plazo reales.",
              "Al cierre de cada temporada de lluvias y al término de la obra: lecciones aprendidas y cierre de los riesgos que ya no aplican."]: vin(doc, t)

    # ------------------------------------------------------------------ 13 conclusiones
    tit(doc, "13", "Conclusiones y recomendaciones")
    tit(doc, "13.1", "Conclusiones", 2)
    concl = ["Se identificaron %d riesgos y %d oportunidades para la ejecución de la obra. Sin respuesta, %d son de nivel ALTO: canteras y transporte (R-08), "
             "certificación ambiental (R-15), lluvias estacionales (R-01) y depósitos de material excedente (R-25)." % (len(R), len(RS.OPORT), n_niv["ALTO"]),
             "Con el plan de respuesta, ningún riesgo queda en nivel ALTO; %d quedan en MODERADO." % n_res["MODERADO"],
             "La exposición de la Entidad al 80 %% de confianza baja de %s a %s si las acciones previas a la convocatoria se ejecutan a tiempo." % (mm_(I_["entidad"]["80"]), mm_(Rr["entidad"]["80"])),
             "La FTE declara que no hay deslizamientos ni inundaciones y que la orografía es plana. El inventario vial y el relieve muestran lo contrario: inundación por "
             "arrozales, empozamientos, taludes de 20 y 30 m y terreno ondulado desde el km 23.",
             "Los documentos tienen %d incompatibilidades que generan riesgo (capítulo 8). Entre ellas: tres longitudes distintas del tramo, plazo de 8 meses "
             "frente a 720 días, DME de 0.38 ha para 297,756 m³ y metrados de puentes repetidos." % len(RS.OBS)]
    for t in concl: vin(doc, t)
    tit(doc, "13.2", "Recomendaciones", 2)
    rec = ["Levantar las observaciones del capítulo 8 y compatibilizar el expediente técnico antes de la convocatoria.",
           "Obtener antes de convocar la certificación ambiental, el instrumento arqueológico, las autorizaciones de agua y de extracción en cauce, las actas de "
           "libre disponibilidad y la factibilidad de las concesionarias. También confirmar la potencia y calidad de canteras y la capacidad de los DME.",
           "Verificar con estudios por sitio la estabilidad de los cortes en 36+060 y 37+520 y la hidráulica, la socavación y la cimentación de cada puente.",
           "Prever una reserva de contingencia de %s (P80 residual) y la certificación y previsión presupuestal multianual del proyecto." % mm_(reserva),
           "Conformar el comité de gestión de riesgos e incorporar la matriz (Excel) al expediente de contratación, con la asignación propuesta.",
           "Exigir al contratista la póliza Todo Riesgo Construcción, el seguro de responsabilidad civil, el plan de manejo del tránsito y el cronograma con calendario de lluvias.",
           "Adecuar la matriz a la directiva de gestión de riesgos de la Dirección General de Abastecimiento si se publica antes de la convocatoria."]
    for t in rec: vin(doc, t)
    par(doc, "Fuentes: estudio de inventario vial SM-107 (2026), FTE del MTC (registro 01/04/2025) y su resumen ejecutivo (diciembre 2024), presupuesto S10 "
             "(costo al 02/12/2024), plano clave PC-01 (noviembre 2024), modelo Copernicus GLO-30 y NOAA GHCN-Daily.", 8.5, cur=True, antes=6)
    par(doc, "\n\n_______________________________\nEspecialista en gestión de riesgos\nCIP: __________", 9, al=C, antes=18)

    # ------------------------------------------------------------------ anexos (horizontal)
    sa = seccion(doc, True)
    tit(doc, "14", "Anexos")
    tit(doc, "14.1", "Anexo 1: Registro de riesgos", 2)
    tabla(doc, ["COD.", "RIESGO", "CAUSA", "EFECTO", "EVIDENCIA", "MOMENTO", "PARTIDA O RECURSO"],
          [[r["cod"], r["riesgo"], r["causa"], r["efecto"], r["evid"], r["momento"], r["partida"]] for r in R], [1.2, 4.2, 5.0, 4.0, 4.6, 2.2, 4.4], tam=6.5)
    tit(doc, "14.2", "Anexo 2: Documentos que acompañan el estudio", 2)
    for t in ["MATRIZ_GESTION_RIESGOS_EJECUCION_OBRA_SM107.xlsx: datos, escalas, identificación, análisis cualitativo y cuantitativo con fórmulas, "
              "Monte Carlo, mapa de calor, plan de respuesta, asignación, oportunidades, formato de seguimiento, observaciones y calendario.",
              "PLANOS_GESTION_RIESGOS_SM107.dxf: PR-01 Plano de ubicación de riesgos (1/25,000, UTM WGS 84 zona 18 S) y PR-02 Diagrama lineal de riesgos por "
              "progresiva, con presentaciones para imprimir; vistas previas en PNG."]: vin(doc, t)
    pr1 = os.path.join(RAIZ, "entregables_vial_sm107", "vistas_previas", "PR-01_UBICACION_DE_RIESGOS.png")
    pr2 = os.path.join(RAIZ, "entregables_vial_sm107", "vistas_previas", "PR-02_DIAGRAMA_LINEAL_DE_RIESGOS.png")
    fig(doc, pr1, 21.0, "Lámina PR-01: plano de ubicación de riesgos (vista previa)", "PLANOS_GESTION_RIESGOS_SM107.dxf.")
    fig(doc, pr2, 21.0, "Lámina PR-02: diagrama lineal de riesgos por progresiva (vista previa)", "PLANOS_GESTION_RIESGOS_SM107.dxf.")

    # ------------------------------------------------------------------ contenido y pies
    for num, (t, nv) in CONT.items():
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0 if nv == 1 else 0.8); p.paragraph_format.space_after = Pt(2)
        fuente(p.add_run("%s  %s" % (num, t.upper() if nv == 1 else t)), 10 if nv == 1 else 9.5, nv == 1)
        marca_contenido._p.addprevious(p._p)
    marca_contenido._p.getparent().remove(marca_contenido._p)
    for s in doc.sections: pie(s, "Estudio de gestión de riesgos en la ejecución de obra - SM-107 Tingo de Ponasa - Shamboyacu - CUI N.° %s" % P["cui"])
    doc.save(SAL); print(SAL, "tablas", NUMT[0], "figuras", NUMF[0])


if __name__ == "__main__":
    main()
