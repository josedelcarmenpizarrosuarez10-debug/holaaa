# -*- coding: utf-8 -*-
"""Excel de la matriz de gestión de riesgos de la ejecución de obra - SM-107 Tingo de Ponasa - Shamboyacu (CUI 2673852).
Celda amarilla = dato; celda verde = fórmula. Los niveles P x I y el valor monetario esperado se calculan con fórmulas;
la simulación de Monte Carlo se corre en riesgos_sm107.py y se copia como valores.
Uso: python3 matriz_xlsx.py"""
import os, sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as COL
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, Reference

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import riesgos_sm107 as RS
RAIZ = RS.RAIZ
SAL = os.path.join(RAIZ, "entregables_vial_sm107", "MATRIZ_GESTION_RIESGOS_EJECUCION_OBRA_SM107.xlsx")

AZUL = PatternFill("solid", fgColor="1F4E79"); AMA = PatternFill("solid", fgColor="FFF2CC"); VER = PatternFill("solid", fgColor="E2EFDA")
GRIS = PatternFill("solid", fgColor="F2F2F2")
ROJO = PatternFill("solid", fgColor="FF7C80"); NAR = PatternFill("solid", fgColor="FFD966"); VERDE = PatternFill("solid", fgColor="A9D08E")
B = Side(style="thin", color="808080"); BORDE = Border(left=B, right=B, top=B, bottom=B)
FT = "Arial"


def hoja(wb, nombre, titulo, anchos):
    ws = wb.create_sheet(nombre)
    ws["A1"] = titulo; ws["A1"].font = Font(name=FT, bold=True, size=12, color="1F4E79")
    ws["A2"] = "PROYECTO: " + RS.P["nombre"] + " - CUI N.° " + RS.P["cui"]; ws["A2"].font = Font(name=FT, size=8, italic=True)
    for i, w in enumerate(anchos, 1): ws.column_dimensions[COL(i)].width = w
    ws.sheet_view.showGridLines = False
    return ws


def cab(ws, fila, textos, alto=30):
    for j, t in enumerate(textos, 1):
        c = ws.cell(fila, j, t); c.fill = AZUL; c.font = Font(name=FT, bold=True, color="FFFFFF", size=9)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BORDE
    ws.row_dimensions[fila].height = alto


def celda(ws, f, c, v, fill=None, fmt=None, neg=False, centro=False, wrap=True, tam=9):
    x = ws.cell(f, c, v); x.font = Font(name=FT, size=tam, bold=neg); x.border = BORDE
    x.alignment = Alignment(horizontal="center" if centro else "left", vertical="center", wrap_text=wrap)
    if fill: x.fill = fill
    if fmt: x.number_format = fmt
    return x


def colores_nivel(ws, rango):
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"ALTO"'], fill=ROJO))
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"MODERADO"'], fill=NAR))
    ws.conditional_formatting.add(rango, CellIsRule(operator="equal", formula=['"BAJO"'], fill=VERDE))


def main():
    RS.calcular(); d = json.load(open(os.path.join(RS.CALC, "riesgos.json"))); MC = d["montecarlo"]
    P = RS.P; R = RS.R; n = len(R)
    wb = Workbook(); wb.remove(wb.active)

    # ------------------------------------------------------------------ INDICE
    ws = hoja(wb, "INDICE", "MATRIZ DE GESTIÓN DE RIESGOS EN LA EJECUCIÓN DE LA OBRA", [6, 34, 90])
    filas = [("DATOS", "Datos del proyecto, presupuesto y costo diario de ampliación"),
             ("ESCALAS", "Escalas de probabilidad e impacto, umbrales y matriz P x I"),
             ("IDENTIFICACION", "Registro de riesgos: evento, causa, efecto, ubicación y evidencia"),
             ("ANALISIS CUALITATIVO", "Probabilidad, impacto en costo y plazo, P x I y nivel (inherente y residual)"),
             ("ANALISIS CUANTITATIVO", "Costo y plazo triangulares, valor monetario esperado y exposición de la Entidad"),
             ("MONTE CARLO", "Resultados de la simulación (20,000 iteraciones), inherente y residual"),
             ("MAPA DE CALOR", "Ubicación de cada riesgo en la matriz P x I, antes y después de la respuesta"),
             ("PLAN DE RESPUESTA", "Estrategia, acciones, responsable, partida, momento y alerta temprana"),
             ("ASIGNACION", "Parte que asume cada riesgo y criterio"),
             ("OPORTUNIDADES", "Oportunidades y su estrategia"),
             ("SEGUIMIENTO", "Formato de seguimiento mensual para el comité de gestión de riesgos"),
             ("OBSERVACIONES", "Observaciones a los documentos que generan riesgo"),
             ("CALENDARIO", "Lluvias mensuales de referencia y programación recomendada")]
    cab(ws, 4, ["N°", "HOJA", "CONTENIDO"])
    for i, (h, t) in enumerate(filas, 1):
        celda(ws, 4 + i, 1, i, centro=True); celda(ws, 4 + i, 2, h, neg=True); celda(ws, 4 + i, 3, t)
    ws.cell(19, 2, "Celda amarilla = dato; celda verde = fórmula.").font = Font(name=FT, size=9, italic=True)
    ws.cell(20, 2, "Montos en soles a precio de obra (incluye gastos generales, utilidad e IGV).").font = Font(name=FT, size=9, italic=True)

    # ------------------------------------------------------------------ DATOS
    ws = hoja(wb, "DATOS", "DATOS DEL PROYECTO", [52, 22, 60])
    cab(ws, 4, ["CONCEPTO", "VALOR", "FUENTE"])
    datos = [("Código único de inversiones (CUI)", P["cui"], None, "Inventario vial / FTE"),
             ("Entidad", P["entidad"], None, "FTE"), ("Unidad ejecutora presupuestal", P["ue"], None, "FTE"),
             ("Ubicación", P["ubicacion"], None, "FTE"), ("Progresiva inicial", P["prog_ini"], None, "Inventario vial"),
             ("Progresiva final", P["prog_fin"], None, "Inventario vial"), ("Longitud según inventario vial (km)", P["L_inv"], "0.000", "Inventario vial"),
             ("Longitud según FTE (km)", P["L_fte"], "0.000", "FTE"), ("Longitud en partida de topografía (km)", P["L_topo"], "0.000", "S10 1.04.01"),
             ("Costo directo (S/)", P["CD"], "#,##0.00", "S10 al " + P["fecha_ppto"]), ("Gastos generales 10.50 % (S/)", P["GG"], "#,##0.00", "S10"),
             ("Utilidad 7.00 % (S/)", P["UT"], "#,##0.00", "S10"), ("Subtotal (S/)", P["ST"], "#,##0.00", "S10"), ("IGV 18 % (S/)", P["IGV"], "#,##0.00", "S10"),
             ("Presupuesto de obra (S/)", P["PO"], "#,##0.00", "S10"), ("Supervisión 3.95 % (S/)", P["SUP"], "#,##0.00", "S10"),
             ("Costo total de inversión (S/)", P["TOTAL"], "#,##0.00", "S10"), ("Plazo de ejecución (días calendario)", P["plazo"], "0", "Resumen ejecutivo"),
             ("Modalidad de ejecución", P["modalidad"], None, "FTE 19"), ("Fuente de financiamiento", P["fuente_fin"], None, "FTE 20"),
             ("IMDA actual (veh/día)", P["imda_actual"], "0", "FTE 9.1"), ("Población beneficiaria directa (hab)", P["poblacion"], "#,##0", "FTE 6")]
    ref = {}
    for i, (a, v, fmt, fu) in enumerate(datos, 5):
        celda(ws, i, 1, a); celda(ws, i, 2, v, AMA, fmt, centro=isinstance(v, (int, float))); celda(ws, i, 3, fu); ref[a] = "DATOS!$B$%d" % i
    f = 5 + len(datos)
    celda(ws, f, 1, "Factor costo directo -> precio de obra", neg=True); celda(ws, f, 2, "=B17*1.18/B14", VER, "0.0000", centro=True); celda(ws, f, 3, "Subtotal x 1.18 / costo directo")
    celda(ws, f + 1, 1, "Costo por día de ampliación: gastos generales con IGV + supervisión (S/ por día)", neg=True)
    celda(ws, f + 1, 2, "=(B15*1.18+B20)/B22", VER, "#,##0.00", centro=True); celda(ws, f + 1, 3, "Se reconoce cuando el retraso es atribuible a la Entidad")
    CDIA = "DATOS!$B$%d" % (f + 1); PO = "DATOS!$B$19"
    assert ws["A19"].value.startswith("Presupuesto de obra") and ws["A22"].value.startswith("Plazo")

    # ------------------------------------------------------------------ ESCALAS
    ws = hoja(wb, "ESCALAS", "ESCALAS DE PROBABILIDAD E IMPACTO Y MATRIZ P x I", [16, 10, 58, 18, 18, 16, 16])
    cab(ws, 4, ["PROBABILIDAD", "VALOR", "DESCRIPCIÓN"])
    desc_p = ["Puede ocurrir solo en circunstancias excepcionales", "Podría ocurrir alguna vez durante la obra", "Puede ocurrir durante la obra",
              "Es probable que ocurra durante la obra", "Se espera que ocurra, casi con certeza"]
    for i, ((t, v), ds) in enumerate(zip(RS.PROB, desc_p)):
        celda(ws, 5 + i, 1, t, neg=True); celda(ws, 5 + i, 2, v, AMA, "0.00", centro=True); celda(ws, 5 + i, 3, ds)
    cab(ws, 11, ["IMPACTO", "VALOR", "DESCRIPCIÓN", "COSTO HASTA (% PO)", "COSTO HASTA (S/)", "PLAZO HASTA (días)"])
    desc_i = ["Efecto mínimo; se absorbe con la holgura del frente", "Afecta un frente; se recupera dentro del mes",
              "Afecta varios frentes o un hito; requiere reprogramar", "Afecta la ruta crítica o el presupuesto de forma significativa",
              "Compromete el plazo total, el financiamiento o la viabilidad"]
    for i, ((t, v), ds) in enumerate(zip(RS.IMP, desc_i)):
        r = 12 + i
        celda(ws, r, 1, t, neg=True); celda(ws, r, 2, v, AMA, "0.00", centro=True); celda(ws, r, 3, ds)
        if i < 4:
            celda(ws, r, 4, RS.IMPACTO_COSTO[i], AMA, "0.0%", centro=True); celda(ws, r, 5, "=D%d*%s" % (r, PO), VER, "#,##0", centro=True)
            celda(ws, r, 6, RS.IMPACTO_PLAZO[i], AMA, "0", centro=True)
        else:
            celda(ws, r, 4, "mayor", centro=True); celda(ws, r, 5, "mayor", centro=True); celda(ws, r, 6, "mayor", centro=True)
    ws.cell(17, 1, "El nivel de impacto es el mayor entre el de costo y el de plazo, evaluados con el valor más probable.").font = Font(name=FT, size=8, italic=True)
    cab(ws, 19, ["UMBRAL", "P x I", "INTERPRETACIÓN"])
    celda(ws, 20, 1, "ALTO", ROJO, neg=True, centro=True); celda(ws, 20, 2, RS.UMBRAL_ALTO, AMA, "0.00", centro=True)
    celda(ws, 20, 3, "P x I >= 0.18: respuesta obligatoria antes del inicio, seguimiento mensual del comité")
    celda(ws, 21, 1, "MODERADO", NAR, neg=True, centro=True); celda(ws, 21, 2, RS.UMBRAL_MOD, AMA, "0.00", centro=True)
    celda(ws, 21, 3, "0.05 <= P x I < 0.18: respuesta planificada y seguimiento por el residente y el supervisor")
    celda(ws, 22, 1, "BAJO", VERDE, neg=True, centro=True); celda(ws, 22, 2, "< 0.05", centro=True); celda(ws, 22, 3, "Se acepta y se vigila con el registro")
    # matriz
    ws.cell(24, 1, "MATRIZ P x I").font = Font(name=FT, bold=True, size=10, color="1F4E79")
    celda(ws, 25, 1, "P \\ I", GRIS, neg=True, centro=True)
    for j in range(5): celda(ws, 25, 2 + j, "=A%d&\" (\"&TEXT(B%d,\"0.00\")&\")\"" % (12 + j, 12 + j), GRIS, neg=True, centro=True)
    for i in range(5):
        r = 30 - i
        celda(ws, r, 1, "=A%d&\" (\"&TEXT(B%d,\"0.00\")&\")\"" % (5 + i, 5 + i), GRIS, neg=True, centro=True)
        for j in range(5): celda(ws, r, 2 + j, "=$B%d*$B%d" % (5 + i, 12 + j), None, "0.000", centro=True)
    ws.conditional_formatting.add("B26:F30", CellIsRule(operator="greaterThanOrEqual", formula=["$B$20"], fill=ROJO))
    ws.conditional_formatting.add("B26:F30", CellIsRule(operator="between", formula=["$B$21", "$B$20-0.000001"], fill=NAR))
    ws.conditional_formatting.add("B26:F30", CellIsRule(operator="lessThan", formula=["$B$21"], fill=VERDE))
    for c in "BCDEF": ws.column_dimensions[c].width = max(ws.column_dimensions[c].width, 17)

    # ------------------------------------------------------------------ IDENTIFICACION
    ws = hoja(wb, "IDENTIFICACION", "IDENTIFICACIÓN DE RIESGOS", [7, 16, 40, 42, 36, 32, 40])
    cab(ws, 4, ["COD.", "CATEGORÍA", "RIESGO (EVENTO)", "CAUSA", "EFECTO", "UBICACIÓN / PROGRESIVA", "EVIDENCIA EN LOS DOCUMENTOS"])
    for i, r in enumerate(R, 5):
        for j, v in enumerate([r["cod"], r["cat"], r["riesgo"], r["causa"], r["efecto"], r["ubic"], r["evid"]], 1):
            celda(ws, i, j, v, neg=(j == 1), centro=(j == 1))
        ws.row_dimensions[i].height = 75
    ws.freeze_panes = "C5"

    # ------------------------------------------------------------------ ANALISIS CUALITATIVO
    ws = hoja(wb, "ANALISIS CUALITATIVO", "ANÁLISIS CUALITATIVO (P x I)", [7, 44, 12, 7, 14, 8, 10, 8, 8, 11, 7, 8, 12, 12, 7, 8, 12])
    cab(ws, 4, ["COD.", "RIESGO", "PROBABILIDAD", "P", "COSTO MÁS PROBABLE (S/)", "NIVEL COSTO", "PLAZO MÁS PROBABLE (días)", "NIVEL PLAZO",
                "NIVEL IMPACTO", "IMPACTO", "I", "P x I", "NIVEL INHERENTE", "PROBABILIDAD RESIDUAL", "P RES.", "P x I RES.", "NIVEL RESIDUAL"], 42)
    dv = DataValidation(type="list", formula1="=ESCALAS!$A$5:$A$9", allow_blank=False); ws.add_data_validation(dv)
    for i, r in enumerate(R, 5):
        celda(ws, i, 1, "=IDENTIFICACION!A%d" % i, VER, neg=True, centro=True); celda(ws, i, 2, "=IDENTIFICACION!C%d" % i, VER)
        celda(ws, i, 3, r["P_txt"], AMA, centro=True); dv.add(ws.cell(i, 3))
        celda(ws, i, 4, "=VLOOKUP(C%d,ESCALAS!$A$5:$B$9,2,FALSE)" % i, VER, "0.00", centro=True)
        celda(ws, i, 5, r["costo"][1], AMA, "#,##0", centro=True)
        celda(ws, i, 6, "=1+(E{0}>ESCALAS!$E$12)+(E{0}>ESCALAS!$E$13)+(E{0}>ESCALAS!$E$14)+(E{0}>ESCALAS!$E$15)".format(i), VER, "0", centro=True)
        celda(ws, i, 7, r["plazo"][1], AMA, "0", centro=True)
        celda(ws, i, 8, "=1+(G{0}>ESCALAS!$F$12)+(G{0}>ESCALAS!$F$13)+(G{0}>ESCALAS!$F$14)+(G{0}>ESCALAS!$F$15)".format(i), VER, "0", centro=True)
        celda(ws, i, 9, "=MAX(F%d,H%d)" % (i, i), VER, "0", centro=True)
        celda(ws, i, 10, "=INDEX(ESCALAS!$A$12:$A$16,I%d)" % i, VER, centro=True)
        celda(ws, i, 11, "=INDEX(ESCALAS!$B$12:$B$16,I%d)" % i, VER, "0.00", centro=True)
        celda(ws, i, 12, "=D%d*K%d" % (i, i), VER, "0.000", centro=True)
        celda(ws, i, 13, "=IF(L{0}>=ESCALAS!$B$20,\"ALTO\",IF(L{0}>=ESCALAS!$B$21,\"MODERADO\",\"BAJO\"))".format(i), VER, neg=True, centro=True)
        celda(ws, i, 14, r["Pr_txt"], AMA, centro=True); dv.add(ws.cell(i, 14))
        celda(ws, i, 15, "=VLOOKUP(N%d,ESCALAS!$A$5:$B$9,2,FALSE)" % i, VER, "0.00", centro=True)
        celda(ws, i, 16, "=O%d*K%d" % (i, i), VER, "0.000", centro=True)
        celda(ws, i, 17, "=IF(P{0}>=ESCALAS!$B$20,\"ALTO\",IF(P{0}>=ESCALAS!$B$21,\"MODERADO\",\"BAJO\"))".format(i), VER, neg=True, centro=True)
        ws.row_dimensions[i].height = 36
    u = 4 + n
    colores_nivel(ws, "M5:M%d" % u); colores_nivel(ws, "Q5:Q%d" % u)
    r0 = u + 2
    celda(ws, r0, 2, "RESUMEN", GRIS, neg=True); celda(ws, r0, 13, "INHERENTE", GRIS, neg=True, centro=True); celda(ws, r0, 17, "RESIDUAL", GRIS, neg=True, centro=True)
    for k, niv in enumerate(["ALTO", "MODERADO", "BAJO"], 1):
        celda(ws, r0 + k, 2, "Riesgos de nivel " + niv); celda(ws, r0 + k, 13, '=COUNTIF(M5:M%d,"%s")' % (u, niv), VER, "0", centro=True)
        celda(ws, r0 + k, 17, '=COUNTIF(Q5:Q%d,"%s")' % (u, niv), VER, "0", centro=True)
    ws.freeze_panes = "C5"

    # ------------------------------------------------------------------ ANALISIS CUANTITATIVO
    ws = hoja(wb, "ANALISIS CUANTITATIVO", "ANÁLISIS CUANTITATIVO - VALOR MONETARIO ESPERADO", [7, 40, 12, 8, 8, 8, 13, 13, 13, 13, 13, 7, 7, 7, 8, 8, 9, 14, 8, 14])
    cab(ws, 4, ["COD.", "RIESGO", "ASIGNADO A", "FACTOR ENTIDAD", "FACTOR PÓLIZA", "P", "COSTO MÍN. (S/)", "COSTO MÁS PROB. (S/)", "COSTO MÁX. (S/)",
                "COSTO MEDIO (S/)", "VME COSTO (S/)", "PLAZO MÍN.", "PLAZO MÁS PROB.", "PLAZO MÁX.", "PLAZO MEDIO", "VME PLAZO (días)", "RUTA CRÍTICA",
                "EXPOSICIÓN ENTIDAD INHERENTE (S/)", "P RES.", "EXPOSICIÓN ENTIDAD RESIDUAL (S/)"], 54)
    dva = DataValidation(type="list", formula1='"Entidad,Contratista,Compartido"'); ws.add_data_validation(dva)
    for i, r in enumerate(R, 5):
        q = i
        celda(ws, i, 1, "=IDENTIFICACION!A%d" % q, VER, neg=True, centro=True); celda(ws, i, 2, "=IDENTIFICACION!C%d" % q, VER)
        celda(ws, i, 3, r["asig"], AMA, centro=True); dva.add(ws.cell(i, 3))
        celda(ws, i, 4, '=IF(C{0}="Entidad",1,IF(C{0}="Compartido",0.5,0))'.format(i), VER, "0.00", centro=True)
        celda(ws, i, 5, r["fr"], AMA, "0.00", centro=True)
        celda(ws, i, 6, "='ANALISIS CUALITATIVO'!D%d" % q, VER, "0.00", centro=True)
        for j, v in enumerate(r["costo"]): celda(ws, i, 7 + j, v, AMA, "#,##0", centro=True)
        ws.cell(i, 8).value = "='ANALISIS CUALITATIVO'!E%d" % q; ws.cell(i, 8).fill = VER
        celda(ws, i, 10, "=AVERAGE(G%d:I%d)" % (i, i), VER, "#,##0", centro=True)
        celda(ws, i, 11, "=F%d*J%d" % (i, i), VER, "#,##0", centro=True)
        for j, v in enumerate(r["plazo"]): celda(ws, i, 12 + j, v, AMA, "0", centro=True)
        ws.cell(i, 13).value = "='ANALISIS CUALITATIVO'!G%d" % q; ws.cell(i, 13).fill = VER
        celda(ws, i, 15, "=AVERAGE(L%d:N%d)" % (i, i), VER, "0.0", centro=True)
        celda(ws, i, 16, "=F%d*O%d" % (i, i), VER, "0.0", centro=True)
        celda(ws, i, 17, "SI" if r["rc"] else "NO", AMA, centro=True)
        celda(ws, i, 18, '=F{0}*J{0}*D{0}*E{0}+IF(Q{0}="SI",F{0}*O{0}*D{0}*{1},0)'.format(i, CDIA), VER, "#,##0", centro=True)
        celda(ws, i, 19, "='ANALISIS CUALITATIVO'!O%d" % q, VER, "0.00", centro=True)
        celda(ws, i, 20, '=S{0}*J{0}*D{0}*E{0}+IF(Q{0}="SI",S{0}*O{0}*D{0}*{1},0)'.format(i, CDIA), VER, "#,##0", centro=True)
        ws.row_dimensions[i].height = 30
    u = 4 + n; t = u + 1
    celda(ws, t, 2, "TOTAL (suma simple, sin concurrencia de plazos)", GRIS, neg=True)
    for c in (11, 16, 18, 20): celda(ws, t, c, "=SUM(%s5:%s%d)" % (COL(c), COL(c), u), GRIS, "#,##0", neg=True, centro=True)
    celda(ws, t + 1, 2, "Exposición de la Entidad / presupuesto de obra", GRIS, neg=True)
    celda(ws, t + 1, 18, "=R%d/%s" % (t, PO), GRIS, "0.0%", neg=True, centro=True); celda(ws, t + 1, 20, "=T%d/%s" % (t, PO), GRIS, "0.0%", neg=True, centro=True)
    ws.cell(t + 3, 2, "VME = P x valor medio de la distribución triangular. Factor Entidad: 1 (Entidad), 0.5 (compartido), 0 (contratista); factor póliza: "
                      "fracción que no cubre la póliza CAR. Los días de ruta crítica de la Entidad se valorizan con el costo diario de ampliación (hoja DATOS). "
                      "La suma simple no considera que los retrasos se superponen; la hoja MONTE CARLO sí lo hace.").font = Font(name=FT, size=8, italic=True)
    ws.cell(t + 3, 2).alignment = Alignment(wrap_text=True); ws.merge_cells(start_row=t + 3, start_column=2, end_row=t + 3, end_column=20); ws.row_dimensions[t + 3].height = 36
    ws.freeze_panes = "C5"

    # ------------------------------------------------------------------ MONTE CARLO
    ws = hoja(wb, "MONTE CARLO", "SIMULACIÓN DE MONTE CARLO (20,000 ITERACIONES)", [44, 16, 16, 16, 16, 16])
    cab(ws, 4, ["RESULTADO", "P10", "P50", "P80", "P90", "MEDIA"])
    fil = [("Costo para la Entidad - inherente (S/)", MC["inherente"]["entidad"], MC["inherente"]["entidad_media"], "#,##0"),
           ("Costo para la Entidad - residual (S/)", MC["residual"]["entidad"], MC["residual"]["entidad_media"], "#,##0"),
           ("Retraso total de la obra - inherente (días)", MC["inherente"]["dias"], MC["inherente"]["dias_media"], "0"),
           ("Retraso total de la obra - residual (días)", MC["residual"]["dias"], MC["residual"]["dias_media"], "0"),
           ("Retraso atribuible a la Entidad - inherente (días)", MC["inherente"]["diasE"], None, "0"),
           ("Retraso atribuible a la Entidad - residual (días)", MC["residual"]["diasE"], None, "0"),
           ("Costo de riesgos del contratista - inherente (S/)", MC["inherente"]["contratista"], MC["inherente"]["contratista_media"], "#,##0"),
           ("Costo de riesgos del contratista - residual (S/)", MC["residual"]["contratista"], MC["residual"]["contratista_media"], "#,##0")]
    for i, (t_, q, m, fmt) in enumerate(fil, 5):
        celda(ws, i, 1, t_, neg=True)
        for j, k in enumerate(["10", "50", "80", "90"]): celda(ws, i, 2 + j, q[k], None, fmt, centro=True)
        celda(ws, i, 6, m if m is not None else "—", None, fmt, centro=True)
    celda(ws, 14, 1, "Reserva de contingencia recomendada (P80 residual, S/)", ROJO, neg=True)
    celda(ws, 14, 2, MC["residual"]["entidad"]["80"], ROJO, "#,##0", neg=True, centro=True)
    celda(ws, 14, 3, "=B14/%s" % PO, ROJO, "0.0%", neg=True, centro=True)
    met = ["Método:",
           "1. Cada riesgo ocurre en una iteración con su probabilidad (inherente o residual, hoja ANALISIS CUALITATIVO).",
           "2. Si ocurre, su costo y su retraso se toman de una distribución triangular (mínimo, más probable, máximo).",
           "3. Costo de la Entidad: riesgos de la Entidad, 50 % de los compartidos (descontando la póliza CAR donde aplica) y los días de ampliación que le corresponden por el costo diario.",
           "4. Plazo: los riesgos de ruta crítica se agrupan por etapa (inicio, clima, terreno, gestión); en cada grupo se suma el mayor retraso y el 25 % de los demás, y los grupos se suman.",
           "5. Residual: con las respuestas del PLAN DE RESPUESTA ejecutadas en el momento indicado."]
    for k, t_ in enumerate(met): ws.cell(16 + k, 1, t_).font = Font(name=FT, size=9, bold=(k == 0))
    # histogramas
    ws.cell(24, 1, "Histograma del costo para la Entidad (S/ millones)").font = Font(name=FT, bold=True, size=10, color="1F4E79")
    cab(ws, 25, ["CLASE (S/ M)", "INHERENTE (%)", "RESIDUAL (%)"])
    hi, bi = MC["inherente"]["hist_entidad"]; hr, br = MC["residual"]["hist_entidad"]
    lo = min(bi[0], br[0]); hi_ = max(bi[-1], br[-1]); import numpy as np
    sI = np.array(MC["inherente"]["muestras_entidad"]) / 1e6; sR = np.array(MC["residual"]["muestras_entidad"]) / 1e6
    bins = np.linspace(lo, hi_, 31)
    cI = np.histogram(sI, bins)[0] / len(sI) * 100; cR = np.histogram(sR, bins)[0] / len(sR) * 100
    for k in range(30):
        celda(ws, 26 + k, 1, round((bins[k] + bins[k + 1]) / 2, 2), None, "0.00", centro=True)
        celda(ws, 26 + k, 2, round(float(cI[k]), 2), None, "0.0", centro=True); celda(ws, 26 + k, 3, round(float(cR[k]), 2), None, "0.0", centro=True)
    ch = BarChart(); ch.title = "Costo para la Entidad"; ch.y_axis.title = "% de iteraciones"; ch.x_axis.title = "S/ millones"
    ch.add_data(Reference(ws, min_col=2, min_row=25, max_col=3, max_row=55), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=26, max_row=55)); ch.height = 8; ch.width = 18; ch.gapWidth = 10
    ws.add_chart(ch, "E24")

    # ------------------------------------------------------------------ MAPA DE CALOR
    ws = hoja(wb, "MAPA DE CALOR", "MAPA DE CALOR - UBICACIÓN DE LOS RIESGOS EN LA MATRIZ P x I", [22, 24, 24, 24, 24, 24])
    for bloque, (tit, kp, f0) in enumerate([("INHERENTE (antes de la respuesta)", "p", 5), ("RESIDUAL (con la respuesta ejecutada)", "pr", 14)]):
        ws.cell(f0 - 1, 1, tit).font = Font(name=FT, bold=True, size=10, color="1F4E79")
        celda(ws, f0, 1, "PROBABILIDAD \\ IMPACTO", GRIS, neg=True, centro=True)
        for j in range(5): celda(ws, f0, 2 + j, RS.IMP[j][0], GRIS, neg=True, centro=True)
        for i in range(5):
            fila = f0 + 5 - i
            celda(ws, fila, 1, RS.PROB[i][0], GRIS, neg=True, centro=True)
            for j in range(5):
                cods = [r["cod"] for r in R if r[kp] == i and r["i"] == j]
                v = RS.PROB[i][1] * RS.IMP[j][1]
                fill = ROJO if v >= RS.UMBRAL_ALTO - 1e-9 else (NAR if v >= RS.UMBRAL_MOD - 1e-9 else VERDE)
                celda(ws, fila, 2 + j, ", ".join(cods), fill, centro=True); ws.row_dimensions[fila].height = 34

    # ------------------------------------------------------------------ PLAN DE RESPUESTA
    ws = hoja(wb, "PLAN DE RESPUESTA", "PLAN DE RESPUESTA A LOS RIESGOS", [7, 34, 11, 14, 52, 22, 30, 18, 34, 11])
    cab(ws, 4, ["COD.", "RIESGO", "NIVEL INHERENTE", "ESTRATEGIA", "ACCIONES DE RESPUESTA", "RESPONSABLE", "PARTIDA O RECURSO", "MOMENTO",
                "ALERTA TEMPRANA (DISPARADOR)", "NIVEL RESIDUAL"], 34)
    for i, r in enumerate(R, 5):
        celda(ws, i, 1, "=IDENTIFICACION!A%d" % i, VER, neg=True, centro=True); celda(ws, i, 2, "=IDENTIFICACION!C%d" % i, VER)
        celda(ws, i, 3, "='ANALISIS CUALITATIVO'!M%d" % i, VER, neg=True, centro=True)
        for j, v in enumerate([r["resp"], r["acc"], r["quien"], r["partida"], r["momento"], r["alerta"]], 4): celda(ws, i, j, v, centro=(j == 4))
        celda(ws, i, 10, "='ANALISIS CUALITATIVO'!Q%d" % i, VER, neg=True, centro=True)
        ws.row_dimensions[i].height = 80
    colores_nivel(ws, "C5:C%d" % (4 + n)); colores_nivel(ws, "J5:J%d" % (4 + n)); ws.freeze_panes = "C5"

    # ------------------------------------------------------------------ ASIGNACION
    ws = hoja(wb, "ASIGNACION", "ASIGNACIÓN DE RIESGOS ENTRE LAS PARTES", [7, 44, 14, 80])
    cab(ws, 4, ["COD.", "RIESGO", "ASIGNADO A", "CRITERIO (se asigna a la parte que mejor puede gestionarlo y mitigarlo)"])
    crit = {"Entidad": "Depende de decisiones, permisos, predios, financiamiento o del contenido del expediente técnico, que controla la Entidad; "
                       "si ocurre, procede la modificación contractual o la ampliación de plazo.",
            "Contratista": "Depende de la planificación, los métodos, los recursos y la seguridad de la obra, que controla el contratista; "
                           "se cubre con sus gastos generales, sus pólizas y su programación.",
            "Compartido": "La causa tiene una parte de diseño o gestión de la Entidad y otra de ejecución del contratista; se reparte según la causa "
                          "que se compruebe en obra (en el análisis se toma 50 %)."}
    for i, r in enumerate(R, 5):
        celda(ws, i, 1, "=IDENTIFICACION!A%d" % i, VER, neg=True, centro=True); celda(ws, i, 2, "=IDENTIFICACION!C%d" % i, VER)
        celda(ws, i, 3, "='ANALISIS CUANTITATIVO'!C%d" % i, VER, neg=True, centro=True); celda(ws, i, 4, crit[r["asig"]]); ws.row_dimensions[i].height = 40
    u = 4 + n
    for k, a in enumerate(["Entidad", "Contratista", "Compartido"]):
        celda(ws, u + 2 + k, 2, "Riesgos asignados a: " + a, GRIS, neg=True); celda(ws, u + 2 + k, 3, '=COUNTIF(C5:C%d,"%s")' % (u, a), VER, "0", centro=True)

    # ------------------------------------------------------------------ OPORTUNIDADES
    ws = hoja(wb, "OPORTUNIDADES", "OPORTUNIDADES", [7, 40, 40, 12, 7, 14, 14, 14, 14, 12, 50, 20])
    cab(ws, 4, ["COD.", "OPORTUNIDAD", "CAUSA", "PROBABILIDAD", "P", "AHORRO MÍN. (S/)", "AHORRO MÁS PROB. (S/)", "AHORRO MÁX. (S/)", "VME AHORRO (S/)",
                "ESTRATEGIA", "ACCIONES", "RESPONSABLE"], 34)
    for i, o in enumerate(RS.OPORT, 5):
        celda(ws, i, 1, o["cod"], neg=True, centro=True); celda(ws, i, 2, o["oport"]); celda(ws, i, 3, o["causa"])
        celda(ws, i, 4, RS.PROB[o["p"]][0], AMA, centro=True); celda(ws, i, 5, "=VLOOKUP(D%d,ESCALAS!$A$5:$B$9,2,FALSE)" % i, VER, "0.00", centro=True)
        for j, v in enumerate(o["beneficio"]): celda(ws, i, 6 + j, v, AMA, "#,##0", centro=True)
        celda(ws, i, 9, "=E%d*AVERAGE(F%d:H%d)" % (i, i, i), VER, "#,##0", centro=True)
        celda(ws, i, 10, o["resp"], centro=True); celda(ws, i, 11, o["acc"]); celda(ws, i, 12, o["quien"]); ws.row_dimensions[i].height = 45

    # ------------------------------------------------------------------ SEGUIMIENTO
    ws = hoja(wb, "SEGUIMIENTO", "FORMATO DE SEGUIMIENTO MENSUAL DE RIESGOS (COMITÉ DE GESTIÓN DE RIESGOS)", [7, 40, 12, 12, 12, 12, 12, 40, 18, 30])
    ws["A3"] = "Mes de obra: ______     Fecha de revisión: ____/____/______     Valorización N.°: ______"; ws["A3"].font = Font(name=FT, size=9)
    cab(ws, 5, ["COD.", "RIESGO", "ESTADO", "PROBABILIDAD ACTUAL", "IMPACTO ACTUAL", "NIVEL ACTUAL", "¿ALERTA ACTIVADA?", "ACCIÓN EJECUTADA / PENDIENTE",
                "RESPONSABLE", "OBSERVACIONES"], 34)
    dve = DataValidation(type="list", formula1='"Activo,Ocurrido,Cerrado,Nuevo"'); dvs = DataValidation(type="list", formula1='"SI,NO"')
    dvp = DataValidation(type="list", formula1="=ESCALAS!$A$5:$A$9"); dvi = DataValidation(type="list", formula1="=ESCALAS!$A$12:$A$16")
    for x in (dve, dvs, dvp, dvi): ws.add_data_validation(x)
    for i, r in enumerate(R, 6):
        celda(ws, i, 1, r["cod"], neg=True, centro=True); celda(ws, i, 2, r["riesgo"])
        for j in range(3, 11): celda(ws, i, j, None, AMA if j != 6 else VER, centro=True)
        celda(ws, i, 3, "Activo", AMA, centro=True); dve.add(ws.cell(i, 3)); dvp.add(ws.cell(i, 4)); dvi.add(ws.cell(i, 5)); dvs.add(ws.cell(i, 7))
        ws.cell(i, 6).value = ('=IF(OR(D{0}="",E{0}=""),"",IF(VLOOKUP(D{0},ESCALAS!$A$5:$B$9,2,FALSE)*VLOOKUP(E{0},ESCALAS!$A$12:$B$16,2,FALSE)>=ESCALAS!$B$20,"ALTO",'
                               'IF(VLOOKUP(D{0},ESCALAS!$A$5:$B$9,2,FALSE)*VLOOKUP(E{0},ESCALAS!$A$12:$B$16,2,FALSE)>=ESCALAS!$B$21,"MODERADO","BAJO")))').format(i)
        celda(ws, i, 9, r["quien"]); ws.row_dimensions[i].height = 30
    colores_nivel(ws, "F6:F%d" % (5 + n))
    for k in range(5):
        fi = 6 + n + k; celda(ws, fi, 1, "R-%02d" % (35 + k), AMA, neg=True, centro=True)
        for j in range(2, 11): celda(ws, fi, j, None, AMA)
        celda(ws, fi, 3, "Nuevo", AMA, centro=True); dve.add(ws.cell(fi, 3))
    ws.freeze_panes = "C6"

    # ------------------------------------------------------------------ OBSERVACIONES
    ws = hoja(wb, "OBSERVACIONES", "OBSERVACIONES A LOS DOCUMENTOS DEL PROYECTO QUE GENERAN RIESGO", [5, 26, 48, 52, 46, 12])
    cab(ws, 4, ["N°", "ASPECTO", "DOCUMENTO Y DATO", "OBSERVACIÓN", "RECOMENDACIÓN", "RIESGO"])
    for i, o in enumerate(RS.OBS, 1):
        celda(ws, 4 + i, 1, i, centro=True)
        for j, v in enumerate(o, 2): celda(ws, 4 + i, j, v, neg=(j == 2), centro=(j == 6))
        ws.row_dimensions[4 + i].height = 60

    # ------------------------------------------------------------------ CALENDARIO
    ws = hoja(wb, "CALENDARIO", "LLUVIA MENSUAL DE REFERENCIA Y PROGRAMACIÓN RECOMENDADA", [12, 14, 14, 14, 14, 16, 70])
    L = json.load(open(os.path.join(RS.CALC, "lluvia_ghcn.json")))
    t, j_ = L["Tarapoto"], L["Juanjuí"]
    ws["A3"] = ("Fuente: NOAA GHCN-Daily, estaciones TARAPOTO (PE000084455, 282 msnm, %d-%d) y JUANJUI (PEM00084474, 350 msnm, %d-%d); meses con al menos 90 %% "
                "de datos. El tramo está entre ambas estaciones (56 y 60 km). Series cortas: usar el estudio hidrológico del expediente para el diseño." %
                (t["periodo"][0], t["periodo"][1], j_["periodo"][0], j_["periodo"][1])); ws["A3"].font = Font(name=FT, size=8, italic=True)
    cab(ws, 5, ["MES", "TARAPOTO (mm)", "JUANJUÍ (mm)", "PROMEDIO (mm)", "DÍAS >= 10 mm (prom.)", "TEMPORADA", "PROGRAMACIÓN RECOMENDADA"], 34)
    MES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]
    for m in range(1, 13):
        f = 5 + m; a = t["meses"][str(m)]; b = j_["meses"][str(m)]
        celda(ws, f, 1, MES[m - 1], neg=True); celda(ws, f, 2, a[0], AMA, "0", centro=True); celda(ws, f, 3, b[0], AMA, "0", centro=True)
        celda(ws, f, 4, "=AVERAGE(B%d:C%d)" % (f, f), VER, "0", centro=True); celda(ws, f, 5, round((a[2] + b[2]) / 2, 1), AMA, "0.0", centro=True)
        celda(ws, f, 6, '=IF(D{0}>=110,"LLUVIOSA",IF(D{0}>=75,"TRANSICIÓN","ESTIAJE"))'.format(f), VER, neg=True, centro=True)
        celda(ws, f, 7, '=IF(F{0}="LLUVIOSA","Alcantarillas, cunetas, gaviones, fabricación de vigas, campamento; evitar terraplenes en zonas de arrozal y cimentación de puentes",'
                        'IF(F{0}="TRANSICIÓN","Cortes con drenaje de coronación, subbase y base con protección; preparar desvíos de cauce",'
                        '"Cimentación y estribos de puentes, terraplenes en zonas bajas, base y carpeta asfáltica"))'.format(f), VER)
    celda(ws, 18, 1, "ANUAL", GRIS, neg=True); [celda(ws, 18, c, "=SUM(%s6:%s17)" % (COL(c), COL(c)), GRIS, "0", neg=True, centro=True) for c in (2, 3, 4)]
    ws.conditional_formatting.add("F6:F17", CellIsRule(operator="equal", formula=['"LLUVIOSA"'], fill=ROJO))
    ws.conditional_formatting.add("F6:F17", CellIsRule(operator="equal", formula=['"TRANSICIÓN"'], fill=NAR))
    ws.conditional_formatting.add("F6:F17", CellIsRule(operator="equal", formula=['"ESTIAJE"'], fill=VERDE))
    ch = BarChart(); ch.title = "Lluvia mensual de referencia (mm)"; ch.add_data(Reference(ws, min_col=2, min_row=5, max_col=4, max_row=17), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=6, max_row=17)); ch.height = 8; ch.width = 20; ws.add_chart(ch, "A21")

    for w in wb.worksheets:
        w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.page_setup.fitToHeight = 0; w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SAL); print(SAL)


if __name__ == "__main__":
    main()
