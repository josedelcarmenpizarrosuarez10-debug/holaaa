# -*- coding: utf-8 -*-
"""Sustento normativo del programa arquitectónico - Hogar de Refugio Temporal para mujeres víctimas de violencia (CUI 2675514).

Lee la programación de áreas de la usuaria (insumos/solange/sustento_normativo/PROGRAMACION_DE_AREAS_HOGAR_REFUGIO_27-10-2026.xlsx,
sin modificarla) y verifica cada ambiente contra:
  - las normas peruanas citadas en la programación, con el texto de los artículos entregados por la usuaria: RNE A.030
    Hospedaje (R.M. N.° 005-2019-VIVIENDA), A.040 Educación (R.M. N.° 068-2020-VIVIENDA), A.080 Oficinas y A.090 Servicios
    Comunales (D.S. N.° 011-2006-VIVIENDA) y la R.V.M. N.° 104-2019-MINEDU (nivel Inicial);
  - los Lineamientos arquitectónicos para refugios de ONU-Habitat / CONAVIM (2016), referencia de buena práctica (superficies
    sugeridas de las fichas técnicas del capítulo III y criterios de los capítulos II y IV).
La R.V.M. N.° 208-2019-MINEDU (primaria y secundaria) se cita como la programación; su texto no fue entregado y sus índices se
dejan indicados para cotejo.

Uso: python3 herramientas/sustento_normativo_refugio.py
Salida: entregables/SUSTENTO_NORMATIVO_PROGRAMA_ARQUITECTONICO_HOGAR_REFUGIO.docx
        entregables/MATRIZ_SUSTENTO_NORMATIVO_PROGRAMA_ARQUITECTONICO_HOGAR_REFUGIO.xlsx
"""
import math
import os
import re
import sys
from collections import OrderedDict

import openpyxl
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
for d in (AQUI, os.path.join(AQUI, 'varones')):
    sys.path.insert(0, d)
import memoria_descriptiva_refugio as R  # noqa: E402  (formato de las memorias del Refugio)
import caratulas_refugio as CR  # noqa: E402,F401  (configura la carátula del Refugio en caratulas_v)
import caratulas_v as V  # noqa: E402

ENTRADA = os.path.join(RAIZ, 'insumos/solange/sustento_normativo/PROGRAMACION_DE_AREAS_HOGAR_REFUGIO_27-10-2026.xlsx')
SALIDA = os.path.join(RAIZ, 'entregables/SUSTENTO_NORMATIVO_PROGRAMA_ARQUITECTONICO_HOGAR_REFUGIO.docx')
SALIDA_X = os.path.join(RAIZ, 'entregables/MATRIZ_SUSTENTO_NORMATIVO_PROGRAMA_ARQUITECTONICO_HOGAR_REFUGIO.xlsx')
CUI = R.CUI
FECHA = 'OCTUBRE 2026'
LIN = 'Lineamientos ONU-Habitat/CONAVIM'

# ------------------------------------------------------------------------------------------------ normas (texto verificado)
NORMAS = [
    # (grupo, norma, dispositivo, aplicación)
    ('Marco legal del servicio', 'Ley N.° 30364, Ley para prevenir, sancionar y erradicar la violencia contra las mujeres y los '
     'integrantes del grupo familiar, y su Reglamento (D.S. N.° 009-2016-MIMP y modificatorias)', 'Ley y reglamento',
     'Marco del servicio de Hogares de Refugio Temporal: creación, implementación, registro y acreditación.'),
    ('Marco legal del servicio', 'Ley N.° 31621 y su Reglamento (D.S. N.° 005-2023-MIMP)', 'Ley y reglamento',
     'Promueve la implementación y el funcionamiento de los Hogares de Refugio Temporal por los gobiernos regionales y locales.'),
    ('Marco legal del servicio', 'R.M. N.° 150-2016-MIMP: Lineamientos para la atención y funcionamiento de los Hogares de Refugio '
     'Temporal, criterios de derivación y modelo de reglamento interno', 'Resolución ministerial',
     'Organización del servicio y de la atención; base de la programación funcional.'),
    ('Reglamento Nacional de Edificaciones', 'Norma Técnica A.030 Hospedaje (R.M. N.° 005-2019-VIVIENDA)', 'Arts. 4, 5 y 6; Anexos 1, 3 y 4',
     'Zona residencial: cálculo de ocupantes (albergue 4.0 m²/persona), área mínima de habitación (6.00 m²), servicios '
     'higiénicos (2.50 m² mínimo, 50 % de habitaciones con baño privado), pasajes de 1.20 m.'),
    ('Reglamento Nacional de Edificaciones', 'Norma Técnica A.040 Educación (R.M. N.° 068-2020-VIVIENDA)', 'Arts. 13 y 20',
     'Talleres, módulo de enseñanza, SUM y sus servicios higiénicos: índices de ocupación y dotación de aparatos sanitarios.'),
    ('Reglamento Nacional de Edificaciones', 'Norma A.080 Oficinas (D.S. N.° 011-2006-VIVIENDA)', 'Arts. 3, 6, 7, 9, 15 y 23',
     'Zona administrativa: 9.5 m² por persona, altura libre 2.40 m, servicios para empleados, ambiente para basura.'),
    ('Reglamento Nacional de Edificaciones', 'Norma A.090 Servicios Comunales (D.S. N.° 011-2006-VIVIENDA)', 'Arts. 2, 6, 10, 11, 15, 16 y 17',
     'Uso de protección social; tabla de ocupación (oficinas 10.0, asilos y orfanatos 6.0, reunión 1.0 m²/persona), servicios '
     'higiénicos para empleados y público, estacionamientos.'),
    ('Reglamento Nacional de Edificaciones', 'Normas A.010 Condiciones generales de diseño, A.120 Accesibilidad universal y A.130 '
     'Requisitos de seguridad', 'Exigidas por A.080 (arts. 3 y 9), A.090 (arts. 6 y 10) y A.030 (art. 4)',
     'Ventilación, iluminación, accesibilidad para personas con discapacidad, evacuación y seguridad.'),
    ('Normas del MINEDU', 'R.V.M. N.° 104-2019-MINEDU, Norma Técnica "Criterios de diseño para locales educativos del nivel de '
     'Educación Inicial"', 'Arts. 10, 12 y 13; cuadros 7, 16, 18, 19, 21 y 22',
     'Módulo infantil, módulo de enseñanza, ambientes administrativos, tópico, archivo, almacén y servicios generales.'),
    ('Normas del MINEDU', 'R.V.M. N.° 208-2019-MINEDU, Norma Técnica "Criterios de diseño para locales educativos de primaria y '
     'secundaria"', 'Programación arquitectónica', 'Índices citados en la programación para cocina, comedor/SUM, lavandería, '
     'patios, losa y servicios generales (texto no entregado: cotejar).'),
    ('Referencia internacional', 'Lineamientos arquitectónicos para el diseño y construcción de refugios especializados para '
     'mujeres, sus hijas e hijos, víctimas de violencia extrema, familiar y/o de género (ONU-Habitat / CONAVIM, México, 2016)',
     'Capítulos I a IV', 'Buena práctica: programa funcional en ocho áreas, superficies sugeridas por ambiente, seguridad, '
     'accesos, accesibilidad, diseño bioclimático y ecotecnias.'),
]

# índices verificados en el texto de las normas: clave -> (m² por persona o valor, norma corta, texto)
IDX = {
    'A080': (9.5, 'A.080 art. 6', 'El número de ocupantes de una edificación de oficinas se calculará a razón de una persona cada 9.5 m2.'),
    'A090of': (10.0, 'A.090 art. 11', 'Ambientes para oficinas administrativas: 10.0 m2 por persona.'),
    'A090reu': (1.0, 'A.090 art. 11', 'Ambientes de reunión: 1.0 m2 por persona.'),
    'A090asi': (6.0, 'A.090 art. 11', 'Asilos y orfanatos: 6.0 m2 por persona.'),
    'A090est': (16.0, 'A.090 art. 11', 'Estacionamientos de uso general: 16.0 m2 por persona.'),
    'A030alb': (4.0, 'A.030 art. 5', 'Albergue: 04.0 m² por persona.'),
    'A030eh': (12.0, 'A.030 art. 5', 'Establecimientos de hospedaje: 12.0 m² por persona.'),
    'A040tal': (3.0, 'A.040 art. 13', 'Talleres y laboratorios: 3.0 m² por persona.'),
    'A040sum': (1.0, 'A.040 art. 13', 'Salas de usos múltiples: 1.0 m² por persona.'),
    'A040aula': (1.5, 'A.040 art. 13', 'Aulas: 1.5 m² por persona.'),
    'R104aula': (2.40, 'R.V.M. 104-2019 cuadro 16', 'Aula de Ciclo II: capacidad 25 niños(as), I.O. 2.40 m², área 60.00 m².'),
    'R104psi': (2.00, 'R.V.M. 104-2019 cuadro 18', 'Sala de psicomotricidad: capacidad 25 niños(as), I.O. 2.00 m², área 50.00 m².'),
    'R104sum': (1.00, 'R.V.M. 104-2019 num. 12.2.1', 'SUM: I.O. 1.00 m² por niño(a); área no menor de 60.00 m² ni mayor de 120.00 m².'),
    'R104ext': (2.00, 'R.V.M. 104-2019 num. 11.3', 'Espacio exterior: I.O. de 2.00 m² por niño(a) del aula o aulas a la que sirve.'),
    'R104adm': (9.50, 'R.V.M. 104-2019 cuadro 21', 'Espacios independientes para el personal: I.O. por usuario = 9.50 m² por persona.'),
    'R104reu': (1.50, 'R.V.M. 104-2019 num. 13.1.4', 'Sala de reuniones: I.O. de 1.50 m² por personal docente; no menor de 12.00 m².'),
}


def L(a, b=None, txt=''):
    return dict(min=a, max=b, txt=txt)


# análisis de cada fila de la programación (fila del Excel de la usuaria -> verificación)
# idx: índice que se verifica (None: sin índice); lin: superficie sugerida por los lineamientos; ver: CUMPLE / OBSERVAR / NO APLICA
FILAS = {
    6: dict(idx=None, ref='A.090 art. 17; A.080 art. 21', lin=None, ver='OBSERVAR', obs='A.090 art. 17: 1 estacionamiento cada 6 personas del '
         'personal y 1 cada 10 del público, con 1 accesible de 3.80 x 5.00 m cada 50. Las 6 plazas cubren hasta 36 personas de personal. '
         'Los lineamientos (3.8) piden al menos 1 plaza interior para el vehículo institucional, para que las usuarias no desciendan '
         'fuera del predio. El total de la programación (244.00 m²) no es 6 x 12.50 = 75.00 m²: indicar si incluye el patio de maniobras.'),
    7: dict(idx=None, ref='Lineamientos 2.6 (zaguán de servicio)', lin=None, ver='OBSERVAR', obs='Los lineamientos piden un zaguán o patio '
         'de servicio con acceso solo a almacenes, despensa y basura, separado de los espacios de las residentes. El total de la '
         'programación (256.00 m²) no es 3 x 16.00 = 48.00 m²: indicar la base del cálculo.'),
    8: dict(idx=None, ref='Diseño arquitectónico', lin=None, ver='OBSERVAR', obs='El total de la programación (94.60 m²) no es '
         '12 x 2.52 = 30.24 m²: indicar si incluye la circulación de las motos.'),
    9: dict(idx=None, ref='Lineamientos 2.4; R.V.M. 104-2019 num. 13.3.2', lin=None, ver='OBSERVAR', obs='La caseta de vigilancia va en la '
         'entrada principal, con comunicación interna y línea telefónica externa, y vigila los accesos peatonal, vehicular y de '
         'proveedores (lineamientos 2.4 y 2.6). La GARITA + SS.HH. se repite en la zona de servicios generales (fila 78): se cuenta dos '
         'veces (10.00 m²).'),
    10: dict(idx=None, ref='IS.010; lineamientos 2.8', lin=None, ver='NO APLICA', obs='Sin índice de ocupación. Infraestructura básica '
          'exigida por los lineamientos (2.8): agua, cisterna y tanque elevado.'),
    11: dict(idx=None, ref='IS.020; lineamientos 2.8', lin=None, ver='NO APLICA', obs='Sin índice de ocupación. A.030 (anexos): sin red de '
          'alcantarillado, el proyecto contempla el tratamiento y la evacuación con un sistema de depuración.'),
    12: dict(idx=None, ref='Drenaje pluvial del proyecto', lin=None, ver='NO APLICA', obs='Sin índice de ocupación. Obra del sistema de '
          'drenaje pluvial del proyecto.'),
    13: dict(idx=None, ref='Lineamientos 2.4', lin=None, ver='NO APLICA', obs='Los lineamientos piden una barda perimetral de al menos '
          '2.00 m, con acabados acordes al contexto y sin árboles ni postes adosados que permitan escalarla.'),
    15: dict(idx='A080', ref='A.080 art. 6; R.V.M. 104-2019 cuadro 21', lin=L(30, 40, 'Área de bienvenida/lobby/sala, 20/30 usuarias'),
          ver='CUMPLE', obs='Supera la superficie sugerida (30/40 m²). Debe abrirse desde el zaguán de entrada por una puerta con mirilla '
          '(lineamientos 2.6).'),
    16: dict(idx='A080', ref='A.080 art. 6; R.V.M. 104-2019 cuadro 21', lin=L(6, 10, 'Oficina de la coordinadora/directora'),
          ver='CUMPLE', obs='Supera la superficie sugerida (6/10 m²). Los lineamientos piden que las pantallas de las cámaras de las áreas '
          'comunes estén también en esta oficina.'),
    17: dict(idx='A090reu', ref='A.090 art. 11 (ambientes de reunión)', lin=L(12, 20, 'Sala de juntas/comedor del personal, 7/8 personas'),
          ver='CUMPLE', obs='La programación cita la A.080 art. 6 (9.5 m²/persona), pero el índice de 1.00 m²/persona es el de "ambientes '
          'de reunión" de la A.090 art. 11: corregir la cita. Con la R.V.M. 104-2019 (1.50 m²/persona) 25 personas requerirían 37.50 m².'),
    18: dict(idx='A080', ref='A.080 art. 6', lin=None, ver='CUMPLE', obs='Igual al índice de oficinas.'),
    19: dict(idx='A080', ref='A.080 art. 6', lin=None, ver='CUMPLE', obs='2 x 9.50 = 19.00 m²; la programación pone 19.50 m².'),
    20: dict(idx='A080', ref='A.080 art. 6; R.V.M. 104-2019 num. 13.1.1', lin=None, ver='CUMPLE', obs='El índice de 9.5 m²/persona es de '
          'oficinas; para una sala de espera el uso más parecido es "ambientes de reunión" (1.0 m²/persona, A.090 art. 11), con lo que '
          '38.00 m² admite hasta 38 personas en el cálculo de evacuación.'),
    21: dict(idx=None, ref='A.090 art. 15', lin=None, ver='CUMPLE', obs='A.090 art. 15, de 7 a 25 empleados: mujeres 1 lavatorio y 1 inodoro. '
          'La A.080 art. 15 tiene una tabla semejante (de 7 a 20 empleados).'),
    22: dict(idx=None, ref='A.090 art. 15', lin=None, ver='CUMPLE', obs='A.090 art. 15, de 7 a 25 empleados: hombres 1 lavatorio, 1 urinario '
          'y 1 inodoro.'),
    23: dict(idx=None, ref='A.090 art. 16; A.120', lin=None, ver='CUMPLE', obs='A.090 art. 16: el servicio exclusivo para personas con '
          'discapacidad, sin diferenciación de sexo, es adicional a los aparatos exigibles; dimensiones según la A.120.'),
    24: dict(idx=None, ref='R.V.M. 104-2019 num. 13.3.5', lin=None, ver='CUMPLE', obs='La A.090 art. 15 trata de servicios higiénicos; el '
          'cuarto de limpieza está en la R.V.M. 104-2019 num. 13.3.5 (guardado de implementos, lavadero exclusivo y sumidero): corregir la cita.'),
    25: dict(idx='A080', ref='A.080 art. 6', lin=None, ver='CUMPLE', obs='Por confidencialidad (lineamientos 2.4 y 2.6), la sala de visitas '
          'debe estar junto al ingreso, separada de las áreas privadas y de atención, con registro previo de cada visitante.'),
    26: dict(idx='A090of', ref='A.090 art. 11 (oficinas administrativas)', lin=L(9, 13, 'Sala de atención y trabajo de psicología (opción B)'),
          ver='CUMPLE', obs='Supera la superficie sugerida. Los lineamientos piden divisiones fijas de piso a techo y puerta (aislamiento '
          'acústico), vista a un jardín silencioso y sin escritorio entre la profesional y la usuaria en la zona de atención.'),
    27: dict(idx='A090of', ref='A.090 art. 11 (oficinas administrativas)', lin=L(6, 10, 'Oficina de atención de asesoría jurídica'),
          ver='CUMPLE', obs='Supera la superficie sugerida (6/10 m²).'),
    28: dict(idx='A090of', ref='A.090 art. 11 (oficinas administrativas)', lin=L(6, 10, 'Oficina de trabajo social'),
          ver='CUMPLE', obs='Supera la superficie sugerida. Los lineamientos piden casilleros con candado para las pertenencias de valor '
          'de cada usuaria.'),
    29: dict(idx='A090of', ref='A.090 art. 11; R.V.M. 104-2019 cuadro 22', lin=L(9, 14, 'Consultorio de atención médica/enfermería'),
          ver='CUMPLE', obs='Supera el tópico de la R.V.M. 104-2019 (7.00 m²) y la superficie sugerida (9/14 m²). Medicamentos bajo llave.'),
    30: dict(idx=None, ref='R.V.M. 208-2019 (cotejar); lineamientos 2.4 y 2.5', lin=None, ver='CUMPLE', obs='Índice de 3.00 m²/persona de la '
          'programación (R.V.M. 208-2019, por cotejar): 4 x 3.00 = 12.00 m² < 15.00 m². Los lineamientos piden cámaras en accesos, '
          'perímetro y áreas comunes, con pantallas en la caseta y en la oficina de la dirección.'),
    31: dict(idx=None, ref='EM.010; lineamientos 2.8', lin=None, ver='NO APLICA', obs='Sin índice de ocupación.'),
    32: dict(idx=None, ref='R.V.M. 104-2019 cuadro 21', lin=None, ver='CUMPLE', obs='R.V.M. 104-2019: archivo de 6.00 m², contiguo al área '
          'de oficinas.', area_min=6.0),
    33: dict(idx=None, ref='R.V.M. 104-2019 num. 13.3.1', lin=None, ver='CUMPLE', obs='R.V.M. 104-2019: almacén o depósito general de 1.50 m² '
          'por aula, no menor de 10.00 m².', area_min=10.0),
    35: dict(idx='A040tal', ref='A.040 art. 13 (talleres)', lin=L(45, None, 'Taller productivo (ejemplo de 45 m²)'), ver='CUMPLE',
          obs='Los lineamientos (3.7) piden que los talleres productivos respondan a la oferta de empleo de la región.'),
    36: dict(idx='A040tal', ref='A.040 art. 13 (talleres)', lin=L(45, None, 'Taller productivo (ejemplo de 45 m²)'), ver='CUMPLE',
          obs='Requiere instalaciones de agua, desagüe, gas o electricidad de cocina, y punzocortantes bajo llave (lineamientos 2.5).'),
    37: dict(idx='A040tal', ref='A.040 art. 13 (talleres)', lin=L(45, None, 'Taller productivo (ejemplo de 45 m²)'), ver='CUMPLE', obs=''),
    38: dict(idx='A090of', ref='A.090 art. 11; lineamientos 3.5', lin=L(15, 21, 'Espacio para terapia grupal, hasta 10 participantes (mín. 20 m² en obra nueva)'),
          ver='CUMPLE', obs='Supera la superficie sugerida. El índice de oficinas (10.0 m²/persona) no es el de una sala grupal: con '
          '50.00 m² y el de reunión (1.0 m²/persona, A.090 art. 11) la sala admite grupos de 10 o más; se sugiere declarar una capacidad '
          'de 10 participantes. Los lineamientos piden puertas corredizas hacia una sala multiuso.'),
    39: dict(idx='R104psi', ref='R.V.M. 104-2019 cuadro 18', lin=L(25, 30, 'Ludoteca, 20/25 niñas y niños'), ver='CUMPLE',
          obs='El índice de 2.00 m²/niño(a) es el de la sala de psicomotricidad de la R.V.M. 104-2019 (cuadro 18). Los lineamientos '
          'piden piso de hule, vista y acceso directo a un jardín, y ubicación alejada de las salas de atención.'),
    40: dict(idx='R104aula', ref='R.V.M. 104-2019 cuadro 16; A.040 art. 13', lin=L(16, 21, 'Sala de educación multigrado, hasta 10/15 niñas y niños'),
          ver='CUMPLE', obs='La programación usa 2.00 m²/persona; el aula de la R.V.M. 104-2019 tiene I.O. de 2.40 m² (15 x 2.40 = '
          '36.00 m²) y el área de 39.00 m² también la cumple.'),
    41: dict(idx=None, ref='A.040 art. 20 (cuadro 8)', lin=None, ver='CUMPLE', obs='A.040 cuadro 8: mujeres 1 inodoro cada 30 y 1 lavatorio '
          'cada 30.'),
    42: dict(idx=None, ref='A.040 art. 20 (cuadro 8)', lin=None, ver='CUMPLE', obs='A.040 cuadro 8: hombres 1 inodoro cada 60, 1 lavatorio '
          'cada 30 y 1 urinario cada 60.'),
    43: dict(idx=None, ref='A.040 art. 20.2; A.120', lin=None, ver='CUMPLE', obs='A.040 art. 20.2: al menos un lavatorio, un inodoro y un '
          'urinario por piso para personas con discapacidad, pudiendo ser de uso mixto.'),
    45: dict(idx='A030alb', ref='A.030 arts. 5 y 6', lin=L(11, 25, 'Habitación para núcleo familiar, 4 personas'), ver='CUMPLE',
          obs='Mayor que el área mínima de habitación de la A.030 art. 6 (6.00 m²) y dentro del rango sugerido (11/25 m²).'),
    46: dict(idx=None, ref='A.030 art. 6 y Anexo 3', lin=L(2.7, None, 'Baño completo (ejemplos de 2.70 m²)'), ver='CUMPLE',
          obs='A.030: servicio higiénico de 2.50 m² mínimo (art. 6 y Anexo 3, hostal de 1 y 2 estrellas) con inodoro, lavatorio y '
          'ducha; su área no cuenta como parte de la habitación.', area_min=2.5),
    47: dict(idx='A030alb', ref='A.030 arts. 5 y 6', lin=L(11, 25, 'Habitación para núcleo familiar, 4 personas'), ver='CUMPLE', obs=''),
    48: dict(idx=None, ref='A.030 art. 6 y Anexo 3', lin=L(2.7, None, 'Baño completo (ejemplos de 2.70 m²)'), ver='CUMPLE', obs='', area_min=2.5),
    49: dict(idx='A030alb', ref='A.030 arts. 5 y 6', lin=L(11, 25, 'Habitación para núcleo familiar, 4 personas (puertas comunicantes para familias numerosas)'),
          ver='CUMPLE', obs='Dentro del rango sugerido. Esta habitación no tiene SS.HH. propio en la programación: indicar qué baño '
          'compartido la sirve.'),
    50: dict(idx='A030alb', ref='A.030 arts. 5 y 6', lin=L(11, 25, 'Habitación para núcleo familiar, 4 personas'), ver='CUMPLE', obs=''),
    51: dict(idx=None, ref='A.030 art. 6 (num. 6 a 8); Anexo 3', lin=L(2.7, None, 'Baño completo compartido por un máximo de tres habitaciones'),
          ver='OBSERVAR', obs='Ver la verificación de servicios higiénicos de la zona residencial.', area_min=2.5),
    52: dict(idx='A030alb', ref='A.030 arts. 5 y 6', lin=L(11, 25, 'Habitación para núcleo familiar, 4 personas'), ver='CUMPLE',
          obs='Mayor que el máximo sugerido, coherente con familias de hasta 7 integrantes (lineamientos 1.6: 1 mujer + 5 o más hijas/os).'),
    53: dict(idx=None, ref='A.030 art. 6 y Anexo 3', lin=L(2.7, None, 'Baño completo'), ver='CUMPLE', obs='', area_min=2.5),
    54: dict(idx='A030alb', ref='A.030 arts. 5 y 6', lin=L(11, 25, 'Habitación para núcleo familiar, 4 personas'), ver='CUMPLE',
          obs='La programación cita el Anexo 1 (hotel) y el tipo I el art. 5: unificar la cita en el art. 5 (albergue).'),
    55: dict(idx=None, ref='A.030 art. 6 y Anexo 3', lin=L(2.7, None, 'Baño completo'), ver='CUMPLE', obs='El nombre "TIPO III" no corresponde '
          'a la habitación triple tipo II a la que sirve: corregir.', area_min=2.5),
    56: dict(idx='A030alb', ref='A.030 arts. 4 y 5; A.120', lin=L(28, None, 'Habitación para casos especiales y con equipo de atención médica'),
          ver='OBSERVAR', obs='Cumple el índice de la A.030, pero los lineamientos sugieren 28 m² más baño de 6 m² (cama de hospital y '
          'dos camas para acompañantes, puertas de 1.00 m libres, 1.50 m entre muebles). La A.030 art. 4 pide ubicar las habitaciones '
          'accesibles en el piso de ingreso.'),
    57: dict(idx=None, ref='A.120; lineamientos 3.3.3', lin=L(6, None, 'Baño accesible de la habitación especial'), ver='CUMPLE', obs='', area_min=2.5),
    59: dict(idx=None, ref='R.V.M. 208-2019 (cotejar); lineamientos 3.2', lin=L(10, 16, 'Cocina para 30/40 residentes (0.40 m² por residente)'),
          ver='OBSERVAR', obs='Lineamientos: 0.40 m² por residente: 80 comensales x 0.40 = 32.00 m² (cumple); 113 residentes x 0.40 = 45.20 m² '
          '(no cumple si comen todas a la vez). Sin vidrio, punzocortantes bajo llave y reja para impedir el ingreso de niñas y niños.'),
    60: dict(idx=None, ref='Lineamientos 3.2', lin=L(9, 12, 'Despensa'), ver='OBSERVAR', obs='Menor que la superficie sugerida (9/12 m²); '
          'junto con la alacena (9.00 m²) suma 15.00 m². Acceso directo desde el zaguán de servicio.'),
    61: dict(idx=None, ref='Lineamientos 3.2', lin=None, ver='CUMPLE', obs='Ver despensa.'),
    62: dict(idx='A040sum', ref='A.040 art. 13; R.V.M. 104-2019 num. 12.2.1', lin=L(113, None, 'Comedor: 1.00 m² por residente, con todas las residentes a la vez'),
          ver='CUMPLE', obs='El área (240.00 m²) cubre a las 113 plazas de la zona residencial (1.00 m² por residente), pero la capacidad '
          'declarada (80) es menor: declarar la capacidad del comedor igual al total de residentes.'),
    63: dict(idx=None, ref='R.V.M. 208-2019 (cotejar); lineamientos 3.4', lin=None, ver='CUMPLE', obs='Los lineamientos piden al menos 2 '
          'lavadoras, 2 secadoras, 4 fregaderos y 2 mesas de planchado, y tendido al aire libre y cubierto (temporada de lluvias).'),
    64: dict(idx=None, ref='R.V.M. 104-2019 num. 12.4.1; lineamientos 3.9', lin=None, ver='CUMPLE', obs='Huertos previstos en los '
          'lineamientos (3.9) y como ecotecnia (4.4).'),
    65: dict(idx=None, ref='R.V.M. 208-2019 (cotejar)', lin=None, ver='CUMPLE', obs='Área exterior de recreación; no debe ser visible '
          'desde el exterior (lineamientos 3.9).'),
    66: dict(idx=None, ref='R.V.M. 104-2019 num. 11.3; lineamientos 3.9', lin=None, ver='CUMPLE', obs='Con unas 75 niñas y niños (2/3 de '
          '113 residentes, lineamientos 1.6) y 2.00 m² por niño(a) se requieren 150.00 m²; el patio de juegos (138.95 m²) con el patio '
          'cívico (222.75 m²) los cubre. Piso de hule en los juegos.'),
    67: dict(idx=None, ref='Lineamientos 3.9 (área sombreada con pérgola y bancos)', lin=None, ver='CUMPLE', obs='La programación cita la '
          'A.030 art. 5, que no trata este ambiente (su índice es de 4.0 m²/persona en albergues): corregir la cita a los lineamientos '
          '3.9. 20 personas x 2.00 m² = 40.00 m² < 44.00 m².', area_min=40.0),
    68: dict(idx=None, ref='Lineamientos 3.9', lin=L(420, None, 'Cancha de básquet 15 x 28 m (voleibol 9 x 18 m)'), ver='CUMPLE', obs=''),
    69: dict(idx=None, ref='A.040 art. 20', lin=None, ver='CUMPLE', obs='La programación cita "capítulo 4" de la A.040 (dotación de servicios, '
          'art. 20).'),
    70: dict(idx=None, ref='R.V.M. 208-2019 (cotejar)', lin=None, ver='CUMPLE', obs=''),
    72: dict(idx=None, ref='Lineamientos 3.4', lin=None, ver='CUMPLE', obs='Los lineamientos piden separar artículos nuevos y usados, '
          'anaqueles por tallas y acceso solo del personal.'),
    73: dict(idx=None, ref='R.V.M. 104-2019 num. 13.3.1; lineamientos 3.4', lin=None, ver='OBSERVAR', obs='El índice de 8.80 m²/persona '
          'de la R.V.M. 104-2019 es el del dormitorio del alojamiento temporal docente (num. 13.2.2), no de un almacén: corregir. El '
          'almacén general pide 1.50 m² por aula y no menos de 10.00 m² (cumple). Los lineamientos piden separar los productos de '
          'limpieza, bajo llave, de los artículos para las familias.', area_min=10.0),
    74: dict(idx=None, ref='R.V.M. 208-2019 (cotejar)', lin=None, ver='CUMPLE', obs='4 x 3.00 = 12.00 m² < 15.00 m².', area_min=12.0),
    75: dict(idx=None, ref='A.080 art. 23; R.V.M. 104-2019 num. 13.3.4', lin=None, ver='CUMPLE', obs='Ambiente con punto de agua y sumidero, '
          'sin contenedores en espacios exteriores (R.V.M. 104-2019); en el zaguán de servicio, con acceso desde el exterior y tres '
          'contenedores (lineamientos 3.4).', area_min=6.0),
    76: dict(idx=None, ref='EM.010; R.V.M. 104-2019 num. 13.3.6', lin=None, ver='NO APLICA', obs='Lejos del alcance de niñas y niños.'),
    77: dict(idx=None, ref='R.V.M. 104-2019 num. 13.3.3', lin=None, ver='NO APLICA', obs='Acceso solo del personal autorizado.'),
    78: dict(idx=None, ref='Lineamientos 2.4', lin=None, ver='OBSERVAR', obs='Repite la GARITA + SS.HH. de la zona de obras exteriores '
          '(fila 9): dejar una sola.'),
    79: dict(idx=None, ref='EM.010; Código Nacional de Electricidad', lin=None, ver='NO APLICA', obs='A.030 (anexos): ambiente separado para '
          'el equipo de generación de energía.'),
    80: dict(idx=None, ref='EM.010; A.030 anexos', lin=None, ver='NO APLICA', obs=''),
    82: dict(idx=None, ref='A.030 art. 4; A.130', lin=None, ver='CUMPLE', obs='Pasajes y escaleras de 1.20 m de ancho libre mínimo (A.030 '
          'art. 4) y anchos de evacuación según la A.130.'),
}

# ambientes de los lineamientos que no figuran en la programación
FALTANTES = [
    ('Cuarto de descanso para el personal en guardia nocturna', '4 m²', '3.3.4', 'Junto a las habitaciones, con timbre de alarma desde '
     'la habitación especial si está alejada.'),
    ('Sala de juntas / comedor del personal', '12 m² / 20 m²', '3.6.1', 'Separada de las áreas de las usuarias, con cocineta.'),
    ('Sala para adolescentes', '6 m² / 10 m²', '3.2.3', 'Estudio y recreación de las hijas adolescentes, cerca del patio.'),
    ('Taller de cómputo', '10.5 m² / 14 m²', '3.2.6', 'Siempre supervisado (evita que se publique la ubicación del refugio).'),
    ('Sala multiuso', '15 m² / 23 m²', '3.5.4', 'Puede integrarse con la terapia grupal mediante puertas corredizas.'),
    ('Bodega de artículos para las familias', 'según demanda', '3.4', 'Blancos, ropa interior, ropa y calzado, higiene personal; puede '
     'integrarse al almacén de donaciones.'),
    ('Zaguanes de entrada y de servicio; salida de emergencia', '—', '2.6', 'Cuatro accesos: peatonal, vehículos oficiales, proveedores '
     'y puerta trasera de emergencia.'),
]


# ------------------------------------------------------------------------------------------------ lectura de la programación
def leer():
    wb = openpyxl.load_workbook(ENTRADA, data_only=True); ws = wb.active
    zonas = OrderedDict(); zona = None
    for r in range(6, ws.max_row + 1):
        b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
        if b and str(b).strip() and str(b).strip().upper() not in ('TOTAL',):
            if str(b).strip().upper().startswith(('AREA', 'ÁREA', 'CIRCULACION +')):
                break
            zona = str(b).strip(); zonas.setdefault(zona, [])
        if not c or str(c).strip().upper() == 'TOTAL' or zona is None:
            continue
        v = [ws.cell(r, k).value for k in range(3, 10)]
        cant, cap, ind, norma, area, total = v[1], v[2], v[3], v[4], v[5], v[6]
        zonas[zona].append(dict(fila=r, amb=' '.join(str(c).split()), cant=cant, cap=cap, ind=(str(ind).strip() if ind else ''),
                                norma=(' '.join(str(norma).split()) if norma else ''), area=area, total=total))
    tot = dict(construida=ws['I85'].value, circ_muros=ws['I86'].value, verde=ws['I87'].value, proyecto=ws['I88'].value)
    return zonas, tot


def nombre(t):
    """Nombre de ambiente o zona en minúsculas con siglas y tipos en mayúscula (SS.HH., PTAR, SUM, tipo I/II/III)."""
    t = t.split(':')[0].strip().capitalize()
    t = re.sub(r'(?i)\bss\.?\s*hh\.?', 'SS.HH.', t)
    t = re.sub(r'(?i)\bhab\.\s*', 'hab. ', t)
    t = re.sub(r'\s*-\s*tipo', ' - tipo', t)
    t = re.sub(r'(?i)\btipo (i{1,3})\b', lambda m: 'tipo ' + m.group(1).upper(), t)
    t = re.sub(r'\bCirculacion\b', 'Circulación', t)
    for a, b in (('Ptar', 'PTAR'), ('/sum', '/SUM'), ('doble-compartido', 'doble compartido'), ('Area ', 'Área '), ('Almacen ', 'Almacén '),
                 ('maquinas', 'máquinas'), ('Sub estación', 'Subestación'), ('electrico', 'eléctrico')):
        t = t.replace(a, b)
    return t


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def analizar(zonas):
    out = OrderedDict()
    for z, filas in zonas.items():
        lista = []
        for f in filas:
            a = dict(f); A = FILAS.get(f['fila'], dict(idx=None, ref='', lin=None, ver='—', obs=''))
            a.update(ref=A['ref'], lin=A['lin'], ver=A['ver'], obs=A['obs'])
            k = A['idx']
            a['idx_val'] = IDX[k][0] if k else None
            a['idx_ref'] = IDX[k][1] if k else ''
            a['req'] = (f['cap'] * IDX[k][0]) if (k and num(f['cap'])) else A.get('area_min')
            a['calc'] = (f['cant'] or 0) * (f['area'] or 0) if num(f['area']) else None
            a['dif'] = (f['total'] - a['calc']) if (num(f['total']) and a['calc'] is not None) else None
            lin = A['lin']
            if lin and num(f['area']):
                a['lin_est'] = ('MENOR' if f['area'] < lin['min'] - 1e-6 else
                                'SUPERA' if (lin['max'] and f['area'] > lin['max'] + 1e-6) else 'EN RANGO')
            else:
                a['lin_est'] = '—'
            if a['req'] is not None and num(f['area']) and f['area'] < a['req'] - 1e-6:
                a['ver'] = 'NO CUMPLE'
            lista.append(a)
        out[z] = lista
    return out


def residencial(A):
    """Verificaciones de la zona residencial: población, % de habitaciones con baño, baños compartidos (A.030 y lineamientos)."""
    filas = next(v for k, v in A.items() if 'RESIDENCIAL' in k.upper())
    hab = [f for f in filas if f['amb'].upper().startswith('HABITACI')]
    ssh = [f for f in filas if f['amb'].upper().startswith('SS.HH')]
    plazas = sum(f['cant'] * f['cap'] for f in hab)
    n_hab = sum(f['cant'] for f in hab)
    dobles = [f for f in hab if 'DOBLE' in f['amb'].upper()]
    n_dob = sum(f['cant'] for f in dobles); p_dob = sum(f['cant'] * f['cap'] for f in dobles)
    priv = n_hab - n_dob
    comp = sum(f['cant'] for f in ssh if 'COMPARTIDO' in f['amb'].upper())
    return dict(plazas=plazas, n_hab=n_hab, priv=priv, pct=100.0 * priv / n_hab, n_dob=n_dob, p_dob=p_dob, comp=comp,
                req_a030=2 * math.ceil(n_dob / 4), req_lin=math.ceil(n_dob / 3), req_alb=math.ceil(p_dob / 4),
                ninos=round(plazas * 2 / 3), hab=hab)


# ------------------------------------------------------------------------------------------------ Word
def f2(x):
    return '%.2f' % x if num(x) else ('—' if x in (None, '') else str(x))


def encabezado(M):
    sec = M.d.add_section(WD_SECTION.NEW_PAGE)
    sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
    pg = OxmlElement('w:pgNumType'); pg.set(qn('w:start'), '1'); sec._sectPr.append(pg)
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    M.run(h, 'SUSTENTO NORMATIVO - ARQUITECTURA', 8, True, color=R.GRANATE)
    M.run(h, '   |   %s - CUI N.° %s' % (R.CORTO, CUI), 8, color=R.GRANATE)
    b = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '2'), ('color', R.GRANATE)):
        e.set(qn('w:' + k), v)
    b.append(e); h._p.get_or_add_pPr().insert(1, b)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    M.run(f, 'Página ', 8.5, color=R.GRIS)
    for tipo in ('begin', 'instr', 'separate', 'txt', 'end'):
        r = M.run(f, '', 8.5, color=R.GRIS)
        if tipo == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)
        elif tipo == 'txt':
            r.text = '1'
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); r._r.append(fc)


def etiq(M, e, t):
    p = M.par(despues=3, sangria=0.9); p.paragraph_format.first_line_indent = Cm(-0.45)
    M.run(p, '•  ' + e + ': ', neg=True); M.run(p, t)


def generar_docx(A, tot, RS):
    M = R.Memoria()
    V.C.SERVICIO = 'EXPEDIENTE TÉCNICO DEL PROYECTO DE INVERSIÓN:'; V.FECHA = 'OCTUBRE - 2026'
    foto = V.foto(); V.escudo()
    V.caratula(M.d, foto, 'SUSTENTO NORMATIVO', 'PROGRAMA ARQUITECTÓNICO - ESPECIALIDAD: ARQUITECTURA')
    encabezado(M); R.indice(M)
    W = M.ancho
    todas = [f for z in A.values() for f in z]
    obs = [f for f in todas if f['ver'] in ('OBSERVAR', 'NO CUMPLE') or 'corregir' in f['obs'] or 'indicar' in f['obs']]

    # ======================================================================================== 1
    M.t1('GENERALIDADES')
    M.t2('Objeto')
    M.par('El presente documento sustenta, ambiente por ambiente, el programa arquitectónico del proyecto “%s”, con CUI N.° %s, '
          'frente a la normativa peruana vigente citada en la programación de áreas y a los lineamientos arquitectónicos '
          'internacionales para refugios especializados de mujeres, sus hijas e hijos, víctimas de violencia.' % (R.CR_PROY, CUI))
    M.t2('Alcance')
    M.par('Comprende las siete zonas de la programación de áreas del 27 de octubre de 2026 (obras exteriores, administrativa, '
          'ambientes especializados, residencial, servicios complementarios, servicios generales y circulación): para cada ambiente '
          'se indica la norma y el artículo que lo sustentan, el índice de ocupación, el área requerida, el área del proyecto, la '
          'superficie sugerida por los lineamientos y el resultado de la verificación. Se añaden los criterios cualitativos de '
          'diseño (ubicación, seguridad, accesos, accesibilidad e infraestructura) y los ambientes recomendados que no figuran en '
          'la programación.')
    M.t2('Documentos de base')
    for t in ('Programación de áreas del Hogar de Refugio Temporal para mujeres víctimas de violencia (27-10-2026).',
              'Textos de las normas técnicas A.030, A.040, A.080 y A.090 del Reglamento Nacional de Edificaciones y de la R.V.M. '
              'N.° 104-2019-MINEDU.',
              'Lineamientos arquitectónicos para el diseño y construcción de refugios especializados para mujeres, sus hijas e hijos, '
              'víctimas de violencia extrema, familiar y/o de género (ONU-Habitat / CONAVIM, 2016).'):
        M.vin(t)
    M.t2('Ubicación')
    M.par(R.UBIC + '.')

    # ======================================================================================== 2
    M.t1('MARCO NORMATIVO')
    M.par('La normativa se organiza en cuatro grupos. Las normas peruanas son de cumplimiento obligatorio; los lineamientos de '
          'ONU-Habitat y CONAVIM son una referencia de buena práctica que complementa los vacíos de la normativa nacional para '
          'este tipo de equipamiento.')
    M.tabla('Normas aplicables al programa arquitectónico', ['GRUPO', 'NORMA', 'DISPOSITIVOS', 'APLICACIÓN EN EL PROGRAMA'],
            [[g, n, d, a] for g, n, d, a in NORMAS], [2.6, 5.4, 2.6, W - 10.6], tam=8, izq=(0, 1, 2, 3))
    M.t2('Clasificación del uso')
    M.par('El Hogar de Refugio Temporal es una edificación de protección social: la Norma A.090 Servicios Comunales (art. 2) '
          'comprende dentro de la protección social a los asilos y orfanatos, y su art. 11 dispone que "los casos no expresamente '
          'mencionados considerarán el uso más parecido". Por ello, cada zona se dimensiona con la norma del uso más parecido: '
          'oficinas (A.080 y A.090) para la zona administrativa y de atención, educación (A.040 y normas del MINEDU) para los '
          'ambientes especializados, y hospedaje en su clase albergue (A.030) para la zona residencial.')

    # ======================================================================================== 3
    M.t1('ÍNDICES DE OCUPACIÓN Y CRITERIOS DE DIMENSIONAMIENTO')
    M.par('Los índices siguientes se transcriben de los artículos de cada norma. El área requerida de un ambiente es su capacidad por '
          'el índice; los índices de la A.030, A.080, A.090 y A.040 se dan para calcular el número de ocupantes (salidas de '
          'emergencia y pasajes), y su uso para dimensionar ambientes es el criterio adoptado en la programación.')
    M.tabla('Índices verificados en el texto de las normas', ['NORMA', 'TEXTO DEL ARTÍCULO', 'ÍNDICE'],
            [[r, t, '%.2f' % v] for k, (v, r, t) in IDX.items()], [3.4, W - 5.2, 1.8], tam=8, izq=(0, 1))
    M.t2('Superficies sugeridas por los lineamientos')
    M.par('Las fichas técnicas del capítulo III de los lineamientos indican, para cada ambiente, el número de usuarias y una '
          'superficie sugerida (mínima y máxima). Se usan para comprobar que cada ambiente del programa tenga al menos la '
          'superficie mínima sugerida.')

    # ======================================================================================== 4
    M.t1('POBLACIÓN Y CAPACIDAD DEL HOGAR')
    filas = [[nombre(h['amb']), h['cant'], h['cap'], h['cant'] * h['cap']] for h in RS['hab']]
    filas.append(['TOTAL', RS['n_hab'], '', RS['plazas']])
    M.tabla('Plazas de la zona residencial', ['HABITACIÓN', 'CANT.', 'CAPACIDAD', 'PLAZAS'], filas, [W - 6.0, 1.8, 2.2, 2.0],
            izq=(0,), total=True)
    M.par('La zona residencial suma %d habitaciones y %d plazas. Los lineamientos (1.6) estiman que dos tercios de la población '
          'son hijas e hijos menores de edad, es decir unas %d niñas y niños y %d mujeres. Recomiendan un máximo de unos 10 núcleos '
          'familiares (alrededor de 35 residentes) por refugio, porque un número mayor dificulta la gestión y la convivencia; el '
          'programa los supera, por lo que conviene organizar las habitaciones en pabellones o grupos de convivencia (tipología de '
          'residencias o villas de los lineamientos, 2.2), cada uno con su personal de guardia nocturna.' %
          (RS['n_hab'], RS['plazas'], RS['ninos'], RS['plazas'] - RS['ninos']))
    M.par('Para la evacuación (A.030 art. 5) el número máximo de ocupantes de las habitaciones es el de sus camas; en las áreas de '
          'servicios y de usos diferentes al alojamiento se calcula con la norma de cada uso.')

    # ======================================================================================== 5
    M.t1('SUSTENTO NORMATIVO POR ZONA Y AMBIENTE')
    M.par('En cada tabla: CANT es la cantidad y CAP la capacidad de la programación; I.O., el índice de ocupación verificado (m² por persona); REQ., el área unitaria requerida '
          '(capacidad x índice o área mínima de la norma); PROY., el área unitaria del proyecto; LIN., la superficie sugerida por '
          'los lineamientos. El resultado es CUMPLE, OBSERVAR (cumple la norma pero hay que corregir o completar algo) o NO APLICA '
          '(ambiente sin índice de ocupación).')
    for z, fs in A.items():
        M.t2(nombre(z))
        filas = []
        for f in fs:
            lin = f['lin']
            ltx = ('%s%s' % (f2(lin['min']), (' / ' + f2(lin['max'])) if lin['max'] else '')) if lin else '—'
            filas.append([nombre(f['amb']), f['cant'] if f['cant'] is not None else '—',
                          f['cap'] if num(f['cap']) else '—', f2(f['idx_val']), f['idx_ref'] or f['ref'], f2(f['req']),
                          f2(f['area']), ltx, f['ver']])
        M.tabla('Sustento de la %s' % nombre(z).lower(), ['AMBIENTE', 'CANT', 'CAP', 'I.O.', 'NORMA', 'REQ.', 'PROY.', 'LIN.', 'RESULTADO'],
                filas, [2.9, 1.3, 1.0, 1.05, 2.7, 1.35, 1.45, 1.65, W - 13.4], tam=7.5, izq=(0, 4))
        for f in fs:
            if f['obs']:
                etiq(M, nombre(f['amb']), f['obs'])
        if 'RESIDENCIAL' in z.upper():
            M.par('Verificación de servicios higiénicos de la zona residencial', 10.5, True, color=R.GRANATE, antes=6, despues=3)
            M.tabla('Servicios higiénicos de las habitaciones', ['CRITERIO', 'REQUERIDO', 'PROGRAMA', 'RESULTADO'], [
                ['A.030 art. 6 num. 6: al menos el 50 % de las habitaciones con baño privado', '50 %',
                 '%d de %d (%.0f %%)' % (RS['priv'], RS['n_hab'], RS['pct']), 'CUMPLE' if RS['pct'] >= 50 else 'NO CUMPLE'],
                ['A.030 art. 6 num. 7: 2 SS.HH. comunes por cada 4 habitaciones sin baño (%d habitaciones dobles)' % RS['n_dob'],
                 '%d' % RS['req_a030'], '%d' % RS['comp'], 'CUMPLE' if RS['comp'] >= RS['req_a030'] else 'NO CUMPLE'],
                ['Lineamientos 3.4: un baño completo compartido por un máximo de tres habitaciones', '%d' % RS['req_lin'],
                 '%d' % RS['comp'], 'CUMPLE' if RS['comp'] >= RS['req_lin'] else 'NO CUMPLE'],
                ['A.030 Anexo 4 (albergue): 1 lavatorio, 1 inodoro y 1 ducha por cada 4 personas (%d plazas en dobles)' % RS['p_dob'],
                 '%d juegos' % RS['req_alb'], '%d' % RS['comp'], 'CUMPLE' if RS['comp'] >= RS['req_alb'] else 'OBSERVAR'],
            ], [W - 7.4, 2.2, 2.6, 2.6], tam=8, izq=(0,))
            M.par('Los SS.HH. compartidos cumplen la A.030 art. 6 y los lineamientos. Como el índice de ocupación adoptado es el del '
                  'albergue (4.0 m²/persona), conviene comprobar también el criterio del Anexo 4 de la A.030 (un juego por cada 4 '
                  'personas): las %d plazas de las habitaciones dobles requieren %d juegos y la programación tiene %d baños compartidos. '
                  'Se recomienda que cada baño compartido tenga dos inodoros y dos duchas en cubículos, o aumentar su número. En un '
                  'hogar solo de mujeres, sus hijas e hijos, la diferenciación por sexo de los baños comunes de la A.030 no es '
                  'necesaria.' % (RS['p_dob'], RS['req_alb'], RS['comp']))

    # ======================================================================================== 6
    M.t1('CRITERIOS CUALITATIVOS DE DISEÑO')
    M.par('Además del dimensionamiento, los lineamientos y el RNE fijan condiciones que el proyecto debe atender.')
    crit = [
        ('Ubicación (lineamientos 2.1)', 'zona tranquila no aislada, con servicios e instituciones de apoyo cerca, alumbrado público sin '
         'postes adosados a la barda, transporte público, calles que permitan rutas distintas y accesos alternativos, sin vegetación '
         'que pueda escalarse y lejos de edificios altos o cerros desde donde se observe el interior.'),
        ('Tipología y fachada (2.2 y 2.3)', 'en área suburbana o de baja densidad, casa individual o conjunto de residencias con jardín, '
         'todo dentro de un recinto bardado; fachada proporcionada al entorno, sin anuncios ni símbolos que revelen su uso, sin '
         'ventanas o balcones hacia la calle en el primer piso y con medidas de seguridad discretas.'),
        ('Seguridad frente a agresiones externas (2.4)', 'caseta de vigilancia en la entrada con comunicación interna y externa, vigilancia '
         'las 24 horas, circuito cerrado en accesos y perímetro, protecciones en ventanas y puertas exteriores, barda de al menos 2.00 m, '
         'puertas con mirilla, salida de emergencia opuesta al ingreso, alarma contra intrusión, estacionamiento interno y registro '
         'de todas las personas que ingresan.'),
        ('Seguridad de la población residente (2.5)', 'cámaras en áreas comunes, alarma y detectores contra incendio, señalización de '
         'evacuación (A.130), barandas y cinta antideslizante en escaleras y rampas, rejas para niñas y niños en cocina, lavandería y '
         'almacenes, cerraduras en muebles con punzocortantes y medicamentos, vidrios de seguridad o irrompibles, duchas y barras '
         'antivandálicas, sin tuberías ni barras altas y mobiliario sólido.'),
        ('Accesos (2.6)', 'acceso peatonal, acceso de vehículos oficiales, acceso de proveedores y salida de emergencia; zaguán de entrada '
         'que solo conduce al área de bienvenida y zaguán de servicio que solo conduce a almacenes, despensa y basura.'),
        ('Accesibilidad (2.7; A.120)', 'todas las áreas comunes y de atención accesibles, y al menos una habitación y un baño completos '
         'accesibles en el piso de ingreso (también A.030 art. 4); rampas según la A.010 y la A.120.'),
        ('Infraestructura básica (2.8)', 'agua potable, desagüe (PTAR del proyecto), electricidad, cisterna y tanque elevado, ventilación, '
         'telefonía e internet controlados y contenedores de residuos diferenciados. En Morales (clima cálido húmedo) la temperatura '
         'media mensual es del orden de 25 a 27 °C, por lo que la A.030 (anexos, nota 7) exige sistemas de ventilación o climatización en las habitaciones.'),
        ('Diseño bioclimático y ecotecnias (capítulo IV)', 'clima cálido húmedo: ventilación cruzada, protección solar de las ventanas, '
         'servicios hacia el norte, captación de agua de lluvia, separación de residuos, huertos y sistemas ahorradores.'),
        ('Organización funcional (3.1)', 'separación física entre áreas privadas, de atención y comunes recreativas, para que el ruido y '
         'los olores no lleguen a las sesiones terapéuticas y el personal pueda cerrar áreas según el horario.'),
        ('Alturas y vanos (A.080 arts. 5, 7 y 10; A.090 art. 9; A.040 art. 9)', 'altura libre mínima de 2.40 m en oficinas y de 2.50 m en '
         'ambientes educativos; vanos que abren de al menos el 10 % del área del ambiente; puertas de 1.00 m en el ingreso, 0.90 m en '
         'ambientes interiores y 0.80 m en servicios higiénicos.'),
    ]
    for e, t in crit:
        etiq(M, e, t)

    # ======================================================================================== 7
    M.t1('AMBIENTES RECOMENDADOS NO INCLUIDOS EN LA PROGRAMACIÓN')
    M.par('Los lineamientos proponen, para proyectos de nueva construcción, los siguientes ambientes que no figuran en la '
          'programación; se recomienda incorporarlos o justificar en qué ambiente existente se resuelven.')
    M.tabla('Ambientes recomendados por los lineamientos', ['AMBIENTE', 'SUPERFICIE SUGERIDA', 'LINEAMIENTO', 'NOTA'],
            [list(x) for x in FALTANTES], [4.6, 2.6, 2.0, W - 9.2], tam=8, izq=(0, 3))

    # ======================================================================================== 8
    M.t1('OBSERVACIONES A LA PROGRAMACIÓN DE ÁREAS')
    M.par('No alteran la viabilidad del programa; se recomienda levantarlas antes de presentar el expediente.')
    for f in obs:
        etiq(M, nombre(f['amb']) + ' (fila %d)' % f['fila'], f['obs'])
    dif = [f for f in todas if f['dif'] is not None and abs(f['dif']) > 0.01]
    if dif:
        M.par('Totales que no son cantidad x área unitaria en la programación:', 10.5, True, color=R.GRANATE, antes=6, despues=3)
        M.tabla('Totales digitados', ['AMBIENTE', 'CANT. x ÁREA', 'TOTAL PROGRAMA', 'DIFERENCIA'],
                [[nombre(f['amb']), f2(f['calc']), f2(f['total']), f2(f['dif'])] for f in dif], [W - 7.5, 2.5, 2.5, 2.5], izq=(0,))
    M.par('El "ÁREA CONSTRUIDA (INTERIOR)" de la programación (%.2f m²) suma también las obras exteriores, los patios, el biohuerto '
          'y la losa multiuso, que no son área techada; y el valor de "CIRCULACIÓN + 15 %% MUROS" (%.2f m²) está digitado sin fórmula: '
          'conviene separar el área techada del área libre e indicar la base del 15 %% de muros.' % (tot['construida'], tot['circ_muros']))

    # ======================================================================================== 9
    M.t1('RESUMEN Y CONCLUSIONES')
    filas = []
    for z, fs in A.items():
        n = len(fs); c = sum(1 for f in fs if f['ver'] == 'CUMPLE'); o = sum(1 for f in fs if f['ver'] in ('OBSERVAR', 'NO CUMPLE'))
        filas.append([nombre(z), n, c, o, n - c - o, f2(sum(f['total'] for f in fs if num(f['total'])))])
    filas.append(['TOTAL', sum(r[1] for r in filas), sum(r[2] for r in filas), sum(r[3] for r in filas), sum(r[4] for r in filas),
                  f2(sum(float(r[5]) for r in filas))])
    M.tabla('Resultado por zona', ['ZONA', 'AMBIENTES', 'CUMPLE', 'OBSERVAR', 'NO APLICA', 'ÁREA (m²)'], filas,
            [W - 11.6, 2.2, 2.0, 2.3, 2.3, 2.8], tam=8.5, izq=(0,), total=True)
    nc = sum(1 for f in todas if f['ver'] == 'NO CUMPLE')
    for t in (
        'Ningún ambiente con índice de ocupación queda por debajo del área que exige su norma%s.' % (
            '' if not nc else ' salvo %d' % nc),
        'Los ambientes de atención, administración, talleres y módulos educativos igualan o superan las superficies sugeridas por '
        'los lineamientos internacionales.',
        'Deben corregirse las citas normativas señaladas (sala de reuniones, cuarto de limpieza, área de descanso, almacén general y '
        'habitación triple tipo II), la garita contada dos veces y los totales digitados.',
        'Se recomienda ampliar la habitación especial accesible, completar la despensa, revisar los baños compartidos con el criterio '
        'de albergue e incorporar el cuarto del personal nocturno y la sala del personal.',
        'El número de plazas (%d) supera la escala recomendada por los lineamientos para un refugio: la organización en pabellones '
        'con su propio personal de guardia es la forma de atenderlo.' % RS['plazas'],
    ):
        M.vin(t)
    M.d.core_properties.author = ''
    M.d.core_properties.title = 'Sustento normativo del programa arquitectónico - Hogar de Refugio Temporal'
    M.d.save(SALIDA)
    return SALIDA


# ------------------------------------------------------------------------------------------------ Excel con fórmulas
def generar_xlsx(A, RS):
    wb = openpyxl.Workbook()
    G = R.GRANATE; th = Side(style='thin', color='999999'); B = Border(left=th, right=th, top=th, bottom=th)
    cab_f = PatternFill('solid', fgColor=G); cab_t = Font(bold=True, color='FFFFFF', name='Arial', size=9)
    zf = PatternFill('solid', fgColor='F2E6E8'); nf = Font(name='Arial', size=9); bf = Font(name='Arial', size=9, bold=True)
    al = Alignment(wrap_text=True, vertical='center'); ac = Alignment(wrap_text=True, vertical='center', horizontal='center')

    def cab(ws, fila, textos, anchos):
        for j, (t, w) in enumerate(zip(textos, anchos), 1):
            c = ws.cell(fila, j, t); c.fill = cab_f; c.font = cab_t; c.alignment = ac; c.border = B
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[fila].height = 32

    # ---- MATRIZ
    ws = wb.active; ws.title = 'MATRIZ'
    ws['A1'] = 'MATRIZ DE SUSTENTO NORMATIVO DEL PROGRAMA ARQUITECTÓNICO'; ws['A1'].font = Font(name='Arial', size=12, bold=True, color=G)
    ws['A2'] = 'Proyecto: “%s”, CUI N.° %s' % (R.CR_PROY, CUI); ws['A2'].font = nf
    T = ['FILA PROG.', 'ZONA', 'AMBIENTE', 'CANT.', 'CAPACIDAD', 'ÍNDICE (m²/pers.)', 'NORMA DEL ÍNDICE', 'NORMA CITADA EN LA PROGRAMACIÓN',
         'ÁREA MÍN. NORMA (m²)', 'ÁREA REQ. UNIT. (m²)', 'ÁREA PROY. UNIT. (m²)', 'CUMPLE NORMA', 'LIN. MÍN. (m²)', 'LIN. MÁX. (m²)',
         'LINEAMIENTOS', 'ÁREA CANT. x UNIT. (m²)', 'TOTAL PROGRAMA (m²)', 'DIFERENCIA (m²)', 'RESULTADO', 'OBSERVACIÓN']
    cab(ws, 4, T, [6, 16, 28, 6, 9, 9, 18, 22, 9, 10, 10, 11, 8, 8, 11, 11, 11, 10, 12, 60])
    r = 5
    for z, fs in A.items():
        ws.cell(r, 1, z).font = bf
        for j in range(1, len(T) + 1):
            ws.cell(r, j).fill = zf; ws.cell(r, j).border = B
        r += 1
        for f in fs:
            v = [f['fila'], z, f['amb'], f['cant'], f['cap'] if num(f['cap']) else None, f['idx_val'], f['idx_ref'] or f['ref'],
                 f['norma'], FILAS.get(f['fila'], {}).get('area_min')]
            for j, x in enumerate(v, 1):
                ws.cell(r, j, x)
            ws.cell(r, 10, '=IF(AND(ISNUMBER(E{0}),ISNUMBER(F{0})),E{0}*F{0},IF(ISNUMBER(I{0}),I{0},"—"))'.format(r))
            ws.cell(r, 11, f['area'] if num(f['area']) else None)
            ws.cell(r, 12, '=IF(AND(ISNUMBER(J{0}),ISNUMBER(K{0})),IF(K{0}>=J{0}-0.001,"CUMPLE","NO CUMPLE"),"—")'.format(r))
            lin = f['lin']
            ws.cell(r, 13, lin['min'] if lin else None); ws.cell(r, 14, lin['max'] if lin else None)
            ws.cell(r, 15, '=IF(AND(ISNUMBER(M{0}),ISNUMBER(K{0})),IF(K{0}<M{0},"MENOR",IF(AND(ISNUMBER(N{0}),K{0}>N{0}),"SUPERA","EN RANGO")),"—")'.format(r))
            ws.cell(r, 16, '=IF(AND(ISNUMBER(D{0}),ISNUMBER(K{0})),D{0}*K{0},"—")'.format(r))
            ws.cell(r, 17, f['total'] if num(f['total']) else None)
            ws.cell(r, 18, '=IF(AND(ISNUMBER(P{0}),ISNUMBER(Q{0})),Q{0}-P{0},"—")'.format(r))
            ws.cell(r, 19, f['ver']); ws.cell(r, 20, f['obs'])
            for j in range(1, len(T) + 1):
                c = ws.cell(r, j); c.font = nf; c.border = B; c.alignment = al if j in (2, 3, 7, 8, 20) else ac
                if j in (6, 9, 10, 11, 13, 14, 16, 17, 18):
                    c.number_format = '0.00'
            ws.row_dimensions[r].height = max(24, 12 * (1 + len(f['obs']) // 95))
            r += 1
    fin = r - 1
    ws.freeze_panes = 'D5'; ws.auto_filter.ref = 'A4:T%d' % fin

    # ---- RESIDENCIAL
    w2 = wb.create_sheet('RESIDENCIAL')
    w2['A1'] = 'VERIFICACIÓN DE LA ZONA RESIDENCIAL (A.030 Y LINEAMIENTOS)'; w2['A1'].font = Font(name='Arial', size=12, bold=True, color=G)
    cab(w2, 3, ['HABITACIÓN', 'CANT.', 'CAPACIDAD', 'PLAZAS', 'ÁREA UNIT. (m²)', 'm² POR PERSONA', 'BAÑO PRIVADO'], [34, 8, 10, 9, 12, 12, 12])
    r = 4
    for h in RS['hab']:
        priv = 'NO' if 'DOBLE' in h['amb'].upper() else 'SÍ'
        for j, x in enumerate([h['amb'], h['cant'], h['cap'], '=B%d*C%d' % (r, r), h['area'], '=E%d/C%d' % (r, r), priv], 1):
            c = w2.cell(r, j, x); c.font = nf; c.border = B; c.alignment = al if j == 1 else ac
            if j in (5, 6):
                c.number_format = '0.00'
        r += 1
    h1, h2 = 4, r - 1
    w2.cell(r, 1, 'TOTAL').font = bf
    w2.cell(r, 2, '=SUM(B%d:B%d)' % (h1, h2)); w2.cell(r, 4, '=SUM(D%d:D%d)' % (h1, h2)); tr = r
    r += 2
    cab(w2, r, ['CRITERIO', 'REQUERIDO', 'PROGRAMA', 'RESULTADO'], [34, 12, 12, 12]); r += 1
    comp = RS['comp']
    w2.cell(r, 1, 'Habitaciones con baño privado (%) - A.030 art. 6 num. 6')
    w2.cell(r, 2, 50)
    w2.cell(r, 3, '=SUMIF(G%d:G%d,"SÍ",B%d:B%d)/B%d*100' % (h1, h2, h1, h2, tr))
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","NO CUMPLE")' % (r, r)); r += 1
    w2.cell(r, 1, 'SS.HH. comunes: 2 por cada 4 habitaciones sin baño - A.030 art. 6 num. 7')
    w2.cell(r, 2, '=2*ROUNDUP(SUMIF(G%d:G%d,"NO",B%d:B%d)/4,0)' % (h1, h2, h1, h2)); w2.cell(r, 3, comp)
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","NO CUMPLE")' % (r, r)); r += 1
    w2.cell(r, 1, 'Baño compartido por máximo 3 habitaciones - lineamientos 3.4')
    w2.cell(r, 2, '=ROUNDUP(SUMIF(G%d:G%d,"NO",B%d:B%d)/3,0)' % (h1, h2, h1, h2)); w2.cell(r, 3, comp)
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","NO CUMPLE")' % (r, r)); r += 1
    w2.cell(r, 1, 'Albergue: 1 juego (lavatorio, inodoro, ducha) por cada 4 personas - A.030 Anexo 4')
    w2.cell(r, 2, '=ROUNDUP(SUMIF(G%d:G%d,"NO",D%d:D%d)/4,0)' % (h1, h2, h1, h2)); w2.cell(r, 3, comp)
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","OBSERVAR")' % (r, r)); r += 1
    w2.cell(r, 1, 'Niñas y niños estimados (2/3 de las plazas) - lineamientos 1.6'); w2.cell(r, 3, '=ROUND(D%d*2/3,0)' % tr); r += 1
    w2.cell(r, 1, 'Patio exterior para niñas y niños: 2.00 m² por niño(a) - R.V.M. 104-2019 num. 11.3')
    w2.cell(r, 2, '=C%d*2' % (r - 1)); w2.cell(r, 3, 138.95 + 222.75)
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","NO CUMPLE")' % (r, r)); r += 1
    w2.cell(r, 1, 'Comedor: 1.00 m² por residente - lineamientos 3.2'); w2.cell(r, 2, '=D%d*1' % tr); w2.cell(r, 3, 240)
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","NO CUMPLE")' % (r, r)); r += 1
    w2.cell(r, 1, 'Cocina: 0.40 m² por residente - lineamientos 3.2'); w2.cell(r, 2, '=D%d*0.4' % tr); w2.cell(r, 3, 36)
    w2.cell(r, 4, '=IF(C%d>=B%d,"CUMPLE","OBSERVAR")' % (r, r))
    for rr in range(tr, r + 1):
        for j in range(1, 5):
            c = w2.cell(rr, j); c.font = nf; c.border = B; c.alignment = al if j == 1 else ac
            if j in (2, 3):
                c.number_format = '0.00'

    # ---- ÍNDICES
    w3 = wb.create_sheet('INDICES')
    w3['A1'] = 'ÍNDICES VERIFICADOS EN EL TEXTO DE LAS NORMAS'; w3['A1'].font = Font(name='Arial', size=12, bold=True, color=G)
    cab(w3, 3, ['CLAVE', 'NORMA', 'TEXTO DEL ARTÍCULO', 'ÍNDICE'], [10, 26, 90, 10])
    for i, (k, (v, rf, t)) in enumerate(IDX.items(), 4):
        for j, x in enumerate([k, rf, t, v], 1):
            c = w3.cell(i, j, x); c.font = nf; c.border = B; c.alignment = al
        w3.cell(i, 4).number_format = '0.00'

    # ---- NORMAS
    w4 = wb.create_sheet('NORMAS')
    w4['A1'] = 'MARCO NORMATIVO'; w4['A1'].font = Font(name='Arial', size=12, bold=True, color=G)
    cab(w4, 3, ['GRUPO', 'NORMA', 'DISPOSITIVOS', 'APLICACIÓN'], [24, 60, 28, 70])
    for i, x in enumerate(NORMAS, 4):
        for j, t in enumerate(x, 1):
            c = w4.cell(i, j, t); c.font = nf; c.border = B; c.alignment = al
        w4.row_dimensions[i].height = 42

    # ---- FALTANTES y RESUMEN
    w5 = wb.create_sheet('AMBIENTES FALTANTES')
    cab(w5, 1, ['AMBIENTE', 'SUPERFICIE SUGERIDA', 'LINEAMIENTO', 'NOTA'], [44, 18, 12, 70])
    for i, x in enumerate(FALTANTES, 2):
        for j, t in enumerate(x, 1):
            c = w5.cell(i, j, t); c.font = nf; c.border = B; c.alignment = al
    w6 = wb.create_sheet('RESUMEN')
    w6['A1'] = 'RESULTADO POR ZONA'; w6['A1'].font = Font(name='Arial', size=12, bold=True, color=G)
    cab(w6, 3, ['ZONA', 'AMBIENTES', 'CUMPLE', 'OBSERVAR', 'NO APLICA', 'NO CUMPLE', 'ÁREA PROGRAMA (m²)'], [42, 11, 10, 10, 11, 11, 16])
    for i, z in enumerate(A, 4):
        rg = "MATRIZ!$B$5:$B$%d" % fin
        vals = [z, '=COUNTIF(%s,A%d)' % (rg, i)] + ['=COUNTIFS(%s,A%d,MATRIZ!$S$5:$S$%d,"%s")' % (rg, i, fin, k)
                                                     for k in ('CUMPLE', 'OBSERVAR', 'NO APLICA', 'NO CUMPLE')]
        vals.append('=SUMIF(%s,A%d,MATRIZ!$Q$5:$Q$%d)' % (rg, i, fin))
        for j, t in enumerate(vals, 1):
            c = w6.cell(i, j, t); c.font = nf; c.border = B; c.alignment = al if j == 1 else ac
        w6.cell(i, 7).number_format = '#,##0.00'
    n = 4 + len(A)
    w6.cell(n, 1, 'TOTAL').font = bf
    for j in range(2, 8):
        L_ = get_column_letter(j); c = w6.cell(n, j, '=SUM(%s4:%s%d)' % (L_, L_, n - 1)); c.font = bf; c.border = B; c.alignment = ac
    w6.cell(n, 7).number_format = '#,##0.00'
    wb.properties.creator = ''; wb.properties.title = 'Matriz de sustento normativo - Hogar de Refugio Temporal'
    wb.save(SALIDA_X)
    return SALIDA_X


R.CR_PROY = V.PROYECTO.strip('“”')


def generar():
    zonas, tot = leer()
    A = analizar(zonas); RS = residencial(A)
    return generar_docx(A, tot, RS), generar_xlsx(A, RS), A, RS


if __name__ == '__main__':
    d, x, A, RS = generar()
    for z, fs in A.items():
        print('%-40s %2d ambientes, observar %d' % (z, len(fs), sum(1 for f in fs if f['ver'] in ('OBSERVAR', 'NO CUMPLE'))))
    print('plazas', RS['plazas'], 'habitaciones', RS['n_hab'])
    print(d); print(x)
