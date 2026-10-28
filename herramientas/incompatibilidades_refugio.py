# -*- coding: utf-8 -*-
"""Análisis de incompatibilidades de ubicación - Hogar de Refugio Temporal para mujeres víctimas de violencia (CUI 2675514).

Base: el cuadro de incompatibilidades de la arquitecta (insumos/solange/sustento_normativo/ANALISIS_INCOMPATIBILIDADES_UBICACION_BORRADOR.docx,
18 incompatibilidades del cuadro de la Norma Técnica "Criterios Generales de Diseño para Infraestructura Educativa",
R.V.M. N.° 010-2022-MINEDU), con su resultado de campo (CUMPLE en las 18). Se conserva cada resultado y se ordena el documento:
marco normativo del proyecto (A.090 art. 3, A.040 art. 7 b), método (área de influencia medida desde el lindero), cuadro con la
distancia o el criterio de cada dispositivo, usos colindantes y conclusiones. Arreglos de forma: "I.E." y "UP" pasan a "predio
del Hogar de Refugio"; Decreto Ley N.° 25844 (Ley de Concesiones Eléctricas); Ley Orgánica de Municipalidades; aeródromo con la
distancia al aeropuerto de Tarapoto.
Lo que la arquitecta debe confirmar va en una nota interna aparte (no forma parte del expediente).

Uso: python3 herramientas/incompatibilidades_refugio.py
Salida: entregables/ANALISIS_INCOMPATIBILIDADES_UBICACION_HOGAR_REFUGIO.docx (expediente)
        entregables/MATRIZ_INCOMPATIBILIDADES_UBICACION_HOGAR_REFUGIO.xlsx (expediente)
        entregables/revision_interna/INCOMPATIBILIDADES_CORREGIDO_Y_POR_REVISAR_HOGAR_REFUGIO.docx (uso interno)
"""
import os
import subprocess
import sys
import tempfile

import openpyxl
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
for d in (AQUI, os.path.join(AQUI, 'varones')):
    sys.path.insert(0, d)
import sustento_normativo_refugio as SN  # noqa: E402  (encabezado, pie, viñetas con etiqueta, carátula configurada)
R, V = SN.R, SN.V

SALIDA = os.path.join(RAIZ, 'entregables/ANALISIS_INCOMPATIBILIDADES_UBICACION_HOGAR_REFUGIO.docx')
SALIDA_X = os.path.join(RAIZ, 'entregables/MATRIZ_INCOMPATIBILIDADES_UBICACION_HOGAR_REFUGIO.xlsx')
SALIDA_N = os.path.join(RAIZ, 'entregables/revision_interna/INCOMPATIBILIDADES_CORREGIDO_Y_POR_REVISAR_HOGAR_REFUGIO.docx')
CUI = R.CUI
RVM = 'R.V.M. N.° 010-2022-MINEDU'

# Aeropuerto de Tarapoto: distancia desde la ubicación UTM referencial del predio (E 345 749, N 9 282 885, WGS84 18S)
AERO_KM, AERO_OESTE_KM = 3.5, 2.4

# (incompatibilidad, dispositivo legal, distancia mínima en m o None, criterio, condición encontrada)
CUADRO = [
    ('Velatorios y cementerios', 'D.S. N.° 003-94-SA, Reglamento de la Ley de Cementerios y Servicios Funerarios', 150, '',
     'No existen velatorios ni cementerios a menos de 150 m del predio.'),
    ('Establecimientos de salud', 'R.M. N.° 045-2015/MINSA, NTS N.° 113-MINSA/DGIEM-V.01 "Infraestructura y equipamiento de los '
     'establecimientos de salud del Primer Nivel de Atención", y R.M. N.° 862-2015/MINSA, NTS N.° 119-MINSA/DGIEM-V.01 '
     '"Infraestructura y equipamiento de los establecimientos de salud del Tercer Nivel de Atención"', 100, '',
     'El establecimiento de salud de la localidad se encuentra a más de 100 m del predio.'),
    ('Plantas envasadoras de gas licuado de petróleo (GLP)', 'D.S. N.° 027-94-EM, Reglamento de Seguridad para Instalaciones y '
     'Transporte de Gas Licuado de Petróleo', 100, '', 'No existen plantas envasadoras de GLP en el entorno del predio.'),
    ('Estaciones de servicio, puestos de venta de combustibles (grifos), gasocentros y establecimientos de venta al público de GNV',
     'D.S. N.° 054-93-EM (modificado por el D.S. N.° 037-2007-EM), Reglamento de Seguridad para Establecimientos de Venta al '
     'Público de Combustibles Derivados de Hidrocarburos', 50, '',
     'No existen estaciones de servicio, grifos ni gasocentros a menos de 50 m del predio.'),
    ('Locales de comercialización y consumo de bebidas alcohólicas', 'Ley N.° 28681, Ley que regula la comercialización, consumo y '
     'publicidad de bebidas alcohólicas, y su Reglamento (D.S. N.° 012-2009-SA)', 100, '',
     'No existen locales de venta y consumo de bebidas alcohólicas a menos de 100 m del predio.'),
    ('Plantas de abastecimiento de combustibles líquidos y otros productos derivados de los hidrocarburos', 'D.S. N.° 045-2001-EM, '
     'Reglamento para la Comercialización de Combustibles Líquidos y otros Productos Derivados de los Hidrocarburos', 100, '',
     'No existen plantas de abastecimiento de combustibles ni de derivados de hidrocarburos.'),
    ('Fajas marginales de las fuentes de agua, naturales o artificiales', 'D.S. N.° 001-2010-AG, Reglamento de la Ley N.° 29338, '
     'Ley de Recursos Hídricos', None, 'Faja marginal de cada fuente de agua',
     'No existen fuentes de agua naturales ni artificiales colindantes; el predio no ocupa fajas marginales.'),
    ('Sistema de transporte de hidrocarburos por ductos', 'D.S. N.° 081-2007-EM (modificado por el D.S. N.° 007-2012-EM), '
     'Reglamento de Transporte de Hidrocarburos por Ductos', 200, '', 'No existe sistema de transporte de hidrocarburos por ductos.'),
    ('Pozos para la exploración y explotación de hidrocarburos', 'D.S. N.° 032-2004-EM, Reglamento de las Actividades de '
     'Exploración y Explotación de Hidrocarburos', None, 'Área de las instalaciones de exploración y explotación',
     'No existen actividades de exploración ni de explotación de hidrocarburos.'),
    ('Aeródromos', 'D.S. N.° 050-2001-MTC, Reglamento de la Ley de Aeronáutica Civil, y sus modificatorias', None,
     'Superficies limitadoras de obstáculos',
     'El aeropuerto de Tarapoto está a unos %.1f km, al costado del eje de la pista 17/35 (unos %.1f km al oeste), fuera de las '
     'superficies de aproximación y despegue; las edificaciones del hogar son de baja altura.' % (AERO_KM, AERO_OESTE_KM)),
    ('Servidumbre de líneas aéreas de instalaciones eléctricas', 'R.M. N.° 214-2011-MEM/DM, Código Nacional de Electricidad '
     '(Suministro 2011)', None, 'Faja de servidumbre según la tensión de la línea',
     'No existen líneas aéreas cuya faja de servidumbre ocupe el predio.'),
    ('Servidumbre de electroductos', 'Decreto Ley N.° 25844, Ley de Concesiones Eléctricas, y sus modificatorias', None,
     'Servidumbre impuesta sobre el predio', 'No existen electroductos ni servidumbres eléctricas impuestas sobre el predio.'),
    ('Restricciones radioeléctricas en áreas de uso público próximas a una estación radioeléctrica', 'R.M. N.° 120-2005-MTC/03, '
     'Norma Técnica sobre Restricciones Radioeléctricas en Áreas de Uso Público', None, 'Límites de exposición de la norma',
     'No existen estaciones radioeléctricas que afecten el predio; se cumple la R.M. N.° 120-2005-MTC/03.'),
    ('Plantas de tratamiento de aguas residuales', 'D.S. N.° 011-2006-VIVIENDA, Norma OS.090 "Plantas de tratamiento de aguas '
     'residuales" del RNE', None, 'Distancia de la planta a la zona habitada',
     'No existen plantas de tratamiento de aguas residuales de terceros en el entorno del predio.'),
    ('Faja de terreno lateral y colindante al derecho de vía', 'D.S. N.° 034-2008-MTC, Reglamento Nacional de Gestión de '
     'Infraestructura Vial', None, 'Derecho de vía y faja de propiedad restringida',
     'El predio no está en área restringida; el cerco y las obras, incluido el colector pluvial frontal, quedan dentro del lindero, '
     'sin ocupar la faja de la carretera Oasis.'),
    ('Zonas restringidas colindantes a las vías ferroviarias', 'D.S. N.° 032-2005-MTC, Reglamento Nacional de Ferrocarriles', None,
     'Zona restringida de la vía férrea', 'No existe vía ferroviaria en la zona.'),
    ('Casinos y máquinas tragamonedas', 'Ley N.° 27153, Ley que regula la explotación de los juegos de casino y máquinas '
     'tragamonedas, y sus modificatorias', None, 'Distancia de la ley y de la ordenanza municipal',
     'No existen casinos ni salas de máquinas tragamonedas en el entorno del predio.'),
    ('Hostales, peñas, discotecas, video-pubs, bingos y salas de billar', 'Ley N.° 27972, Ley Orgánica de Municipalidades, art. 79 '
     '(organización del espacio físico y uso del suelo)', None, 'Zonificación y licencias municipales',
     'No existen hostales, peñas, discotecas, video-pubs, bingos ni salas de billar en el entorno del predio.'),
]

NORMAS = [
    ('RNE, Norma A.090 Servicios Comunales, art. 3', 'Las edificaciones de servicios comunales se ubicarán en los lugares señalados '
     'en los Planes de Desarrollo Urbano, o en zonas compatibles con la zonificación vigente.'),
    ('RNE, Norma A.040 Educación (R.M. N.° 068-2020-VIVIENDA), art. 7 b)', 'Ubicación evitando las incompatibilidades de uso '
     'establecidas en la normativa vigente y/o adoptar las alternativas de solución, respecto de su ubicación. Se aplica a los '
     'ambientes educativos del hogar (módulo infantil y módulo de enseñanza).'),
    ('%s, Norma Técnica "Criterios Generales de Diseño para Infraestructura Educativa"' % RVM, 'Cuadro de incompatibilidades de uso '
     'del terreno, que reúne los dispositivos legales sectoriales y sus distancias; se adopta como lista de verificación.'),
    ('Lineamientos ONU-Habitat / CONAVIM (2016), numeral 2.1', 'Ubicación en una zona tranquila, no aislada, con servicios e '
     'instituciones de apoyo cercanos, alumbrado público, transporte público y accesos alternativos.'),
]

COLINDANTES = [
    ('Frente', 'Carretera Oasis', 'Vía de acceso al predio; el cerco y el colector pluvial frontal quedan dentro del lindero.'),
    ('Lateral (aguas arriba)', 'CAR Varones (CUI 2705619)', 'Servicio de protección a niñas, niños y adolescentes; uso compatible.'),
    ('Lateral (aguas abajo)', 'CAR Mujeres (CUI 2717013)', 'Servicio de protección a niñas, niños y adolescentes; uso compatible.'),
]


def seccion(M, apaisada):
    """Nueva sección (horizontal o vertical) que sigue la numeración y el encabezado de la anterior."""
    ant = M.d.sections[-1]
    sec = M.d.add_section(WD_SECTION.NEW_PAGE)
    for pg in sec._sectPr.findall(qn('w:pgNumType')):
        sec._sectPr.remove(pg)
    cor, lar = sorted((ant.page_width, ant.page_height))
    sec.orientation = WD_ORIENT.LANDSCAPE if apaisada else WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = (lar, cor) if apaisada else (cor, lar)
    if apaisada:
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0), Cm(2.5), Cm(2.0)
    else:
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.0), Cm(2.5), Cm(2.0)
    M.ancho = (sec.page_width - sec.left_margin - sec.right_margin) / 360000.0


def dist(x):
    return 'No menor de %d m' % x[2] if x[2] else x[3]


# --------------------------------------------------------------------------------------------------------------- Word (expediente)
def generar_docx():
    M = R.Memoria()
    V.C.SERVICIO = 'EXPEDIENTE TÉCNICO DEL PROYECTO DE INVERSIÓN:'; V.FECHA = 'OCTUBRE - 2026'
    foto = V.foto(); V.escudo()
    V.caratula(M.d, foto, 'INCOMPATIBILIDADES', 'DE UBICACIÓN DEL PROYECTO - ESPECIALIDAD: ARQUITECTURA')
    SN.encabezado(M, 'INCOMPATIBILIDADES DE UBICACIÓN - ARQUITECTURA')
    W = M.ancho

    M.t1('GENERALIDADES')
    M.t2('Objeto')
    M.par('El presente análisis verifica que el terreno del proyecto “%s”, con CUI N.° %s, no presenta incompatibilidades de uso por '
          'cercanía a actividades, instalaciones o servidumbres reguladas por la normativa vigente, de modo que su ubicación es apta '
          'para el Hogar de Refugio Temporal.' % (R.CR_PROY, CUI))
    M.t2('Ubicación del predio')
    M.par(R.UBIC + ', con frente a la carretera Oasis y en el entorno urbano de Morales y Tarapoto. Coordenadas UTM referenciales '
          '(WGS84, zona 18 Sur): E 345 749, N 9 282 885.')
    M.tabla('Usos colindantes del predio', ['LADO', 'COLINDANTE', 'USO Y COMPATIBILIDAD'], [list(c) for c in COLINDANTES],
            [3.2, 4.4, W - 7.6], tam=8.5, izq=(0, 1, 2))
    M.par('El predio forma parte de un conjunto de servicios de protección del Estado (CAR Varones, Hogar de Refugio Temporal y CAR '
          'Mujeres), usos de la misma naturaleza y compatibles entre sí.')

    M.t1('MARCO NORMATIVO')
    M.par('El Hogar de Refugio Temporal es una edificación de protección social (Norma A.090) que incluye ambientes educativos; por ello '
          'se verifica su ubicación con las disposiciones siguientes:')
    M.tabla('Normas que rigen la ubicación del proyecto', ['NORMA', 'DISPOSICIÓN'], [list(n) for n in NORMAS], [5.6, W - 5.6], tam=8.5,
            izq=(0, 1))
    M.par('El cuadro de incompatibilidades de la %s es el instrumento sectorial que reúne, para una edificación con permanencia de '
          'niñas y niños, las distancias y restricciones de los dispositivos legales de cada sector (salud, energía y minas, '
          'transportes, recursos hídricos y gobiernos locales); por ello se adopta para el Hogar de Refugio Temporal, que alberga a '
          'mujeres con sus hijas e hijos.' % RVM)

    M.t1('METODOLOGÍA')
    for t in ('Se define el área de influencia del predio como la franja medida desde el lindero con la distancia que fija cada '
              'dispositivo legal: 50 m (estaciones de servicio), 100 m (establecimientos de salud, GLP, bebidas alcohólicas y plantas '
              'de combustibles), 150 m (velatorios y cementerios) y 200 m (ductos de hidrocarburos).',
              'Para las incompatibilidades sin distancia fija (fajas marginales, servidumbres eléctricas, aeródromos, vías, '
              'restricciones radioeléctricas y usos regulados por la municipalidad) se verifica que la restricción no alcance al predio.',
              'Se reconoce el entorno del predio y se revisa la cartografía y las imágenes satelitales de la zona.',
              'Cada incompatibilidad se califica como CUMPLE cuando la actividad o restricción no existe dentro del área de influencia.'):
        M.vin(t)

    seccion(M, apaisada=True); W2 = M.ancho                        # cuadro en hoja horizontal
    M.t1('CUADRO DE INCOMPATIBILIDADES DE UBICACIÓN')
    filas = [[str(k), x[0], x[1], dist(x), x[4], 'CUMPLE'] for k, x in enumerate(CUADRO, 1)]
    M.tabla('Cuadro N.° 1: incompatibilidades de uso del terreno, según la %s - Hogar de Refugio Temporal' % RVM,
            ['N.°', 'INCOMPATIBILIDAD POR CERCANÍA', 'DISPOSITIVO LEGAL', 'DISTANCIA O CRITERIO', 'CONDICIÓN DEL PREDIO', 'ESTADO'],
            filas, [0.9, 5.0, 7.6, 3.0, W2 - 18.2, 1.7], tam=8, izq=(1, 2, 3, 4))
    for fila in M.d.tables[-1].rows:
        fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    seccion(M, apaisada=False)

    M.t1('CONCLUSIONES')
    for t in ('El predio cumple las %d condiciones del cuadro de incompatibilidades de uso: no existen dentro de su área de influencia '
              'las actividades, instalaciones ni servidumbres que restringen la ubicación del Hogar de Refugio Temporal.' % len(CUADRO),
              'La ubicación cumple la Norma A.090 (art. 3) y la Norma A.040 (art. 7 b): el hogar se emplaza en una zona compatible y sin '
              'incompatibilidades de uso.',
              'Los colindantes (CAR Varones y CAR Mujeres) son servicios de protección de la misma naturaleza, y la carretera Oasis da '
              'acceso al predio sin que las obras ocupen su faja.',
              'La ubicación responde además a los lineamientos para refugios: zona tranquila, no aislada, con servicios cercanos y '
              'acceso vehicular.'):
        M.vin(t)
    M.d.core_properties.author = ''
    M.d.core_properties.title = 'Análisis de incompatibilidades de ubicación - Hogar de Refugio Temporal'
    M.d.save(SALIDA)
    return SALIDA


# --------------------------------------------------------------------------------------------------------------- Excel (expediente)
def generar_xlsx():
    wb = openpyxl.Workbook()
    G = R.GRANATE; th = Side(style='thin', color='999999'); B = Border(left=th, right=th, top=th, bottom=th)
    cab_f = PatternFill('solid', fgColor=G); cab_t = Font(bold=True, color='FFFFFF', name='Arial', size=9)
    nf = Font(name='Arial', size=9); bf = Font(name='Arial', size=9, bold=True)
    al = Alignment(wrap_text=True, vertical='center'); ac = Alignment(wrap_text=True, vertical='center', horizontal='center')

    def cab(ws, fila, textos, anchos):
        for j, (t, w) in enumerate(zip(textos, anchos), 1):
            c = ws.cell(fila, j, t); c.fill = cab_f; c.font = cab_t; c.alignment = ac; c.border = B
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[fila].height = 32

    def titulo(ws, t):
        ws['A1'] = t; ws['A1'].font = Font(name='Arial', size=12, bold=True, color=G)
        ws['A2'] = 'Proyecto: “%s”, CUI N.° %s' % (R.CR_PROY, CUI); ws['A2'].font = nf

    ws = wb.active; ws.title = 'MATRIZ'
    titulo(ws, 'MATRIZ DE INCOMPATIBILIDADES DE UBICACIÓN (CUADRO DE LA %s)' % RVM)
    cab(ws, 4, ['N.°', 'INCOMPATIBILIDAD POR CERCANÍA', 'DISPOSITIVO LEGAL', 'DISTANCIA MÍNIMA (m)', 'CRITERIO (SIN DISTANCIA FIJA)',
                '¿EXISTE DENTRO DEL ÁREA DE INFLUENCIA?', 'CONDICIÓN DEL PREDIO', 'ESTADO'], [5, 32, 48, 11, 24, 14, 50, 11])
    r0 = 5
    for k, x in enumerate(CUADRO):
        r = r0 + k
        for j, v in enumerate([k + 1, x[0], x[1], x[2], x[3] or None, 'NO', x[4]], 1):
            c = ws.cell(r, j, v); c.font = nf; c.border = B; c.alignment = ac if j in (1, 4, 6) else al
        c = ws.cell(r, 8, '=IF(F%d="NO","CUMPLE","NO CUMPLE")' % r); c.font = bf; c.border = B; c.alignment = ac
    r1 = r0 + len(CUADRO) - 1
    ws.freeze_panes = 'C5'

    w2 = wb.create_sheet('AREA DE INFLUENCIA')
    titulo(w2, 'ÁREA DE INFLUENCIA MEDIDA DESDE EL LINDERO')
    cab(w2, 4, ['DISTANCIA (m)', 'N.° DE INCOMPATIBILIDADES', 'INCOMPATIBILIDADES'], [14, 16, 90])
    radios = sorted({x[2] for x in CUADRO if x[2]})
    for i, d in enumerate(radios):
        r = 5 + i
        w2.cell(r, 1, d); w2.cell(r, 2, '=COUNTIF(MATRIZ!D%d:D%d,A%d)' % (r0, r1, r))
        w2.cell(r, 3, '; '.join(x[0] for x in CUADRO if x[2] == d))
        for j in (1, 2, 3):
            w2.cell(r, j).font = nf; w2.cell(r, j).border = B; w2.cell(r, j).alignment = ac if j < 3 else al
    r = 5 + len(radios)
    w2.cell(r, 1, 'Sin distancia fija'); w2.cell(r, 2, '=COUNTBLANK(MATRIZ!D%d:D%d)' % (r0, r1))
    w2.cell(r, 3, '; '.join(x[0] for x in CUADRO if not x[2]))
    for j in (1, 2, 3):
        w2.cell(r, j).font = nf; w2.cell(r, j).border = B; w2.cell(r, j).alignment = ac if j < 3 else al

    w3 = wb.create_sheet('RESUMEN')
    titulo(w3, 'RESUMEN DEL ANÁLISIS DE INCOMPATIBILIDADES')
    cab(w3, 4, ['CONCEPTO', 'CANTIDAD'], [50, 14])
    for i, (t, f) in enumerate((('Incompatibilidades evaluadas', '=COUNTA(MATRIZ!B%d:B%d)' % (r0, r1)),
                                ('CUMPLE', '=COUNTIF(MATRIZ!H%d:H%d,"CUMPLE")' % (r0, r1)),
                                ('NO CUMPLE', '=COUNTIF(MATRIZ!H%d:H%d,"NO CUMPLE")' % (r0, r1)))):
        w3.cell(5 + i, 1, t).font = nf; c = w3.cell(5 + i, 2, f); c.font = bf; c.alignment = ac
        w3.cell(5 + i, 1).border = B; c.border = B
    w3.cell(9, 1, 'RESULTADO').font = bf
    c = w3.cell(9, 2, '=IF(B7=0,"UBICACIÓN APTA","REVISAR")'); c.font = bf; c.alignment = ac
    w3.cell(9, 1).border = B; c.border = B
    wb.properties.creator = ''; wb.properties.title = 'Matriz de incompatibilidades de ubicación - Hogar de Refugio Temporal'
    wb.save(SALIDA_X)
    with tempfile.TemporaryDirectory() as tmp:                     # recalcular las fórmulas con LibreOffice
        subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', tmp, SALIDA_X],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        out = os.path.join(tmp, os.path.basename(SALIDA_X))
        if os.path.exists(out):
            os.replace(out, SALIDA_X)
    return SALIDA_X


# --------------------------------------------------------------------------------------------- nota interna para la arquitecta
def generar_interno():
    import docx
    from docx.shared import Pt
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('INCOMPATIBILIDADES DE UBICACIÓN: CORREGIDO Y PUNTOS PARA CONFIRMAR', 12.5, True, color=R.GRANATE, centro=True, despues=2)
    M.par('Hogar de Refugio Temporal - CUI N.° %s  |  Documento de trabajo, no forma parte del expediente' % CUI, 9, color=R.GRIS,
          centro=True, despues=10)
    M.par('Las 18 incompatibilidades del cuadro quedan en CUMPLE, como en el borrador. Este documento indica lo que ya quedó ajustado '
          'en el análisis y lo que conviene confirmar antes de cerrar el expediente.')

    def sec(t):
        M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4)

    sec('1. YA AJUSTADO EN EL ANÁLISIS')
    M.tabla('Ajustes de redacción y citas', ['N.°', 'EN EL BORRADOR', 'EN EL ANÁLISIS'], [
        ['2, 4, 5, 10, 13, 14, 15, 17, 18', '"I.E.", "UP" y "área definida para la UP" (texto de un proyecto educativo)',
         '"predio" y "entorno del predio" del Hogar de Refugio'],
        ['12', 'Decreto Ley N.° 25884', 'Decreto Ley N.° 25844, Ley de Concesiones Eléctricas'],
        ['18', 'Ley Orgánica de Municipales', 'Ley N.° 27972, Ley Orgánica de Municipalidades'],
        ['13', '"cuando una Institución Educativa se encuentre próximo"', 'restricciones en áreas de uso público próximas a una estación'],
        ['10', 'No se encontró dentro del área definida', 'Aeropuerto de Tarapoto a unos %.1f km, al costado de la pista 17/35; '
         'edificaciones de baja altura (sigue CUMPLE)' % AERO_KM],
        ['15', 'El área no se encuentra en área restringida', 'Se agrega que el cerco y el colector pluvial frontal quedan dentro del '
         'lindero, sin ocupar la faja de la carretera Oasis'],
        ['Todos', 'Distancia dentro del texto del dispositivo', 'Columna aparte "DISTANCIA O CRITERIO"'],
    ], [2.6, 6.0, W - 8.6], tam=8.5, izq=(1, 2))
    M.par('Se añadieron el marco normativo del proyecto (A.090 art. 3 y A.040 art. 7 b), la metodología del área de influencia, los '
          'usos colindantes (CAR Varones y CAR Mujeres) y las conclusiones.', despues=4)

    sec('2. PARA CONFIRMAR')
    M.tabla('Puntos para confirmar', ['TEMA', 'QUÉ CONFIRMAR', 'POR QUÉ'], [
        ['Versión de la norma MINEDU', 'Si el proyecto debe citar la R.V.M. N.° 010-2022-MINEDU (con su modificatoria R.V.M. N.° '
         '020-2023-MINEDU) o la R.V.M. N.° 068-2025-MINEDU del 27-06-2025, que aprueba una norma técnica con el mismo nombre.',
         'Un evaluador puede observar la cita de una norma reemplazada; si cambia, el número del cuadro se ajusta en un solo lugar.'],
        ['Aeródromo (fila 10)', 'La distancia al aeropuerto (unos %.1f km) sale de la ubicación UTM referencial; confirmar con la '
         'ubicación real y, si la entidad lo pide, con las superficies limitadoras de CORPAC.' % AERO_KM,
         'El predio queda dentro del radio de la superficie horizontal interna (4 km); con edificaciones bajas no hay restricción.'],
        ['Líneas eléctricas (filas 11 y 12)', 'Si por la carretera Oasis pasa una red de media tensión junto al frente del predio.',
         'Si existe, indicar que las edificaciones guardan la distancia de seguridad del CNE Suministro 2011.'],
        ['PTAR propia (fila 14)', 'La PTAR del programa (152 m²) es del propio hogar; confirmar su distancia a dormitorios y '
         'ambientes según el sistema de tratamiento que defina la especialidad sanitaria.',
         'El cuadro trata de plantas de terceros; un evaluador puede preguntar por la propia.'],
        ['Establecimiento de salud (fila 2)', 'Nombre y distancia del puesto o centro de salud más cercano.',
         'Dar la distancia medida refuerza el CUMPLE.'],
        ['Sustento gráfico', 'Adjuntar una imagen satelital o un plano con el área de influencia de 50, 100, 150 y 200 m desde el lindero.',
         'Es el respaldo habitual del cuadro; se puede preparar en DXF sobre la planta general.'],
    ], [3.2, 7.4, W - 10.6], tam=8.5, izq=(0, 1, 2))
    d.core_properties.author = ''
    d.core_properties.title = 'Incompatibilidades de ubicación: corregido y puntos para confirmar - Hogar de Refugio'
    os.makedirs(os.path.dirname(SALIDA_N), exist_ok=True)
    d.save(SALIDA_N)
    return SALIDA_N


if __name__ == '__main__':
    print(generar_docx()); print(generar_xlsx()); print(generar_interno())
