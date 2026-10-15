"""Nota (15/10/2026): la planilla ya tiene la ACU nueva; no volver a correr este script (el comparativo "actual" leeria la ACU nueva).
La mano de obra 2026 se cambio directamente en el Excel entregado.
Cotizacion de insumos y propuesta de angulos para la ACU TAPA REGISTRO del Hogar de Refugio Temporal
(Mujeres Violentadas, CUI 2675514). Libro aparte con formulas; no modifica la planilla vigente.
Hojas: 1 COTIZACION (precio publicado, fuente y enlace de cada insumo), 2 ANGULOS (alternativas y verificacion
del contramarco), 3 ACU COMPARATIVO (ACU actual frente a la propuesta con precios cotizados).
"""
import os, warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLAN = os.path.join(RAIZ, "insumos", "metrados_vigentes", "METRADO_DRENAJE_MUJERES_VIOLENTADAS.xlsx")
SAL = os.path.join(RAIZ, "entregables", "COTIZACION_PRECIOS_ACU_TAPA_REGISTRO_HOGAR_REFUGIO.xlsx")
AZUL, CELESTE, AMARILLO, VERDE, ROJO, NARANJA = "1F4E79", "DDEBF7", "FFF2CC", "E2EFDA", "F8CBAD", "FCE4D6"
PROY = "PROYECTO: HOGAR DE REFUGIO TEMPORAL - MUJERES VIOLENTADAS, DISTRITO DE MORALES, SAN MARTIN - CUI N.° 2675514"
FECHA = "Busqueda de precios publicados: 14/10/2026"

# ---------------------------------------------------------------- precios publicados (con IGV salvo indicacion)
# clave, insumo ACU, und, producto, presentacion, precio publicado, contenido en und ACU, factor IGV, tienda, url, estado
COT = [
    ("ang15", 'Angulo L 1 1/2"x1 1/2"x1/8"', "kg", 'Angulo de Acero 1 1/2 x 1 1/2 x 1/8 x 6M (Tupemesa)', "barra 6 m",
     39.90, "=6*'2 ANGULOS'!C7", 1.18, "Sodimac Peru", "https://sodimac.falabella.com.pe/sodimac-pe/product/126996349/Angulo-de-Acero-1-1-2-x-1-1-2-x-1-8-x-6M/126996350",
     "COMPROBADO (precio nacional)"),
    ("ang1", 'Angulo L 1"x1"x1/8" (alternativa)', "kg", 'Angulo de Acero 25 x 25 x 3.0 x 6M (Tupemesa)', "barra 6 m",
     24.90, "=6*'2 ANGULOS'!C8", 1.18, "Sodimac Peru", "https://www.sodimac.com.pe/sodimac-pe/articulo/126996351/Angulo-de-Acero-25-x-25-x-3.0-X-6M/126996352",
     "COMPROBADO (sin stock en linea)"),
    ("ang2", 'Angulo L 2"x2"x3/16" (actual)', "kg", "Ninguna tienda peruana publica precio en soles", "barra 6 m",
     None, "=6*'2 ANGULOS'!C10", 1.18, "-", "", "SIN PRECIO PUBLICADO"),
    ("sold", 'Soldadura cellocord E6011 1/8"', "kg", 'Soldadura Cellocord 1/8" (E-6011)', "caja 5 kg",
     109.90, 5, 1.18, "Sodimac Peru", "https://www.sodimac.com.pe/sodimac-pe/articulo/113780372/Soldadura-Cellocord-1-8-Oerlikon-5KG/113780373",
     "COMPROBADO (precio nacional)"),
    ("cem", "Cemento Portland tipo I", "bol", "Cemento Pacasmayo Tipo I Portland 42.5 kg", "bolsa 42.5 kg",
     36.50, 1, 1.18, "Ferreteria Dinova", "https://ferreteriadinova.com/producto/cemento-tipo-i-portland-x-42-5-kg-pacasmayo/",
     "COMPROBADO (marca que se distribuye en San Martin)"),
    ("are", "Arena gruesa", "m3", "Arena gruesa - Agregados Tarapoto (Jr. Felipe Yap 220, Morales)", "m3 (delivery desde 3 m3)",
     48.00, 1, 1.18, "Agregados Tarapoto (Morales)", "https://www.deperu.com/comercios/arquitectura-inmobiliarias-e-ingenieria-de-la-construccion/agregados-tarapoto-1376751",
     "COMPROBADO LOCAL (Morales)"),
    ("pie", 'Piedra chancada 1/2"', "m3", 'En Tarapoto solo se publica la piedra chancada 3/4" a S/ 70.00 el m3 (mismo proveedor)', "m3",
     None, 1, 1.18, "Agregados Tarapoto (Morales)", "https://www.deperu.com/comercios/arquitectura-inmobiliarias-e-ingenieria-de-la-construccion/agregados-tarapoto-1376751",
     "SIN PRECIO PUBLICADO DE 1/2\" (se mantiene el actual)"),
    ("agua", "Agua", "m3", "EMAPA San Martin, categoria comercial y otros, tramo de 30 m3 a mas: agua 6.80 + alcantarillado 1.53", "m3",
     8.33, 1, 1.00, "Res. 00077-2024-SUNASS-CD (El Peruano 18/12/2024); tarifa sin IGV",
     "https://www.emapasanmartin.com/uploads/Gobernanza%20Gobernabilidad/Gobernabilidad/transparencia-informacion/indicador-8/2025/Publicacion%20Oficial%20-%20Diario%20Oficial%20El%20Peruano%20ESTRUCTURA%20TARIFARIA.pdf",
     "COMPROBADO (tarifa oficial Tarapoto)"),
    ("mad", "Madera tornillo", "p2", "Madera tornillo por pie", "p2",
     5.20, 1, 1.18, "Maderas del Norte", "https://maderas-del-norte.webnode.pe/productos/", "REFERENCIA (la pagina no indica ciudad)"),
    ("a12", 'Acero corrugado 1/2" (borde de abertura)', "kg", 'Barras de Acero 1/2" x 9 m Aceros Arequipa (8.946 kg)', "varilla 9 m",
     36.50, 8.946, 1.18, "Sodimac Peru", "https://www.sodimac.com.pe/sodimac-pe/articulo/113523810/barras-de-acero-1-2/113523812",
     "COMPROBADO (precio nacional)"),
    ("a38", 'Acero corrugado 3/8" (parrilla, asas, anclajes)', "kg", 'Barras de Acero 3/8" x 9 m Aceros Arequipa (5.04 kg)', "varilla 9 m",
     20.35, 5.04, 1.18, "Sodimac Peru", "https://www.sodimac.com.pe/sodimac-pe/articulo/113528077/barras-de-acero-3-8/113528084",
     "COMPROBADO (precio nacional)"),
    ("ala", "Alambre negro N 16", "kg", "Alambre Recocido N° 16 x 10 kg (Prodac)", "rollo 10 kg",
     59.90, 10, 1.18, "Sodimac Peru", "https://sodimac.falabella.com.pe/sodimac-pe/product/113533012/Alambre-Recocido-N%C2%B0-16-10-kg/113533017",
     "COMPROBADO (precio nacional)"),
    ("anti", "Pintura anticorrosiva", "gal", "Anticorrosivo Vencedor Gris 1 gl", "galon",
     53.00, 1, 1.18, "Promart", "https://www.promart.pe/anticorrosivo-vencedor-gris-1-gl/p", "COMPROBADO (precio nacional)"),
    ("esm", "Esmalte sintetico", "gal", "Esmalte sintetico CPP negro 1 gl (cod. 30635)", "galon",
     57.00, 1, 1.18, "Promart", "https://www.promart.pe/esmalte-sintetico-pintor-negro-1-gl-30635/p", "COMPROBADO (precio nacional)"),
    ("lija", "Lija para fierro N° 80", "und", "Lija Fierro 3M ASA grano 80", "pliego",
     2.50, 1, 1.18, "Sodimac Peru", "https://www.sodimac.com.pe/sodimac-pe/b/lija-para-fierro", "COMPROBADO (precio nacional)"),
    ("thin", "Thinner acrilico", "gal", "TOPEX Thinner Acrilico 3.5 L", "envase 3.5 L",
     38.90, "=3.5/3.785", 1.18, "Sodimac Peru", "https://www.sodimac.com.pe/sodimac-pe/b/thinner-acrilico", "COMPROBADO (precio nacional)"),
    ("mez", "Mezcladora 9-11 p3", "hm", "Alquiler de mezcladora de concreto 9-11 p3", "hora-maquina",
     15.00, 1, 1.18, "Konstrutecnia (Juliaca)", "https://konstrutecnia.com/p/alquiler-de-mezcladora-de-concreto-9-11-pie3",
     "REFERENCIA (otra ciudad; no hay tarifa publicada en Tarapoto)"),
    ("vib", "Vibrador de concreto", "hm", "Alquiler de vibradora de concreto 5.5 HP", "hora-maquina",
     13.00, 1, 1.18, "Konstrutecnia (Juliaca)", "https://konstrutecnia.com/p/alquiler-de-vibradora-de-concreto-5-5-hp",
     "REFERENCIA (otra ciudad; no hay tarifa publicada en Tarapoto)"),
    ("sol", "Maquina de soldar 250 A", "hm", "Ninguna empresa publica tarifa por hora o por dia (solo S/ 30 y S/ 55 sin periodo)", "hora-maquina",
     None, 1, 1.18, "Ferropolis / ST America", "https://stamerica.com.pe/alquiler_maquina_soldar", "SIN TARIFA PUBLICADA (se mantiene el actual)"),
]
ACTUAL = {"ang15": 4.66, "ang1": None, "ang2": 4.66, "sold": 13.56, "cem": 30.51, "are": 50.85, "pie": 101.69, "agua": 5.00,
          "mad": 5.08, "a12": 3.81, "a38": 3.81, "ala": 5.08, "anti": 59.32, "esm": 55.08, "lija": None, "thin": None,
          "mez": 15.00, "vib": 10.00, "sol": 8.00}


def datos():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(PLAN, data_only=True)
    a = wb["ACU TAPA REGISTRO"]; ac = wb["COLECTOR ACERO"]
    filas = {}
    for r in range(69, 93):
        if a["C%d" % r].value:
            filas[r] = (a["B%d" % r].value, a["C%d" % r].value, a["E%d" % r].value, a["F%d" % r].value)
    return dict(filas=filas, cu=a["G94"].value, n=a["E11"].value, kg12=ac["U33"].value,
                kg38=ac["U34"].value + ac["U35"].value + ac["U36"].value, acero=a["F16"].value, conc=a["F15"].value)


def construir():
    D = datos()
    wb = openpyxl.Workbook()
    fino = Side(style="thin", color="808080"); borde = Border(left=fino, right=fino, top=fino, bottom=fino)

    def hoja(nombre, anchos):
        ws = wb.create_sheet(nombre)
        for i, w in enumerate(anchos): ws.column_dimensions[chr(65 + i)].width = w
        return ws

    def c(ws, ref, v, fill=None, b=False, h="center", fmt=None, color="000000", wrap=True, sz=10, link=None):
        x = ws[ref]; x.value = v; x.font = Font(name="Arial", size=sz, bold=b, color="0563C1" if link else color, underline="single" if link else None)
        x.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap); x.border = borde
        if fill: x.fill = PatternFill("solid", fgColor=fill)
        if fmt: x.number_format = fmt
        if link: x.hyperlink = link

    def fila(ws, r, texto, ncol, fill=CELESTE, b=True, color="000000", alto=None):
        ult = chr(64 + ncol); ws.merge_cells("A%d:%s%d" % (r, ult, r)); c(ws, "A%d" % r, texto, fill, b, "left", color=color)
        for k in range(2, ncol + 1): ws["%s%d" % (chr(64 + k), r)].border = borde
        if alto: ws.row_dimensions[r].height = alto

    def cab(ws, r, titulos, alto=30):
        for i, t in enumerate(titulos): c(ws, "%s%d" % (chr(65 + i), r), t, AZUL, True, color="FFFFFF")
        ws.row_dimensions[r].height = alto

    # =============================================================== 2 ANGULOS (se crea primero por las referencias)
    wa = hoja("2 ANGULOS", (5, 34, 11, 13, 13, 13, 16, 52))
    # =============================================================== 1 COTIZACION
    ws = wb.active; ws.title = "1 COTIZACION"; wb.move_sheet(ws, -1)
    for i, w in enumerate((5, 30, 6, 40, 14, 11, 11, 7, 11, 11, 10, 26, 14, 30)): ws.column_dimensions[chr(65 + i)].width = w
    N = 14
    fila(ws, 1, "COTIZACION DE INSUMOS - ACU TAPA DE REGISTRO 0.68 x 0.68 C/MARCO Y CONTRAMARCO (Delphin 01.04.04.06.05.01; metrado 01.04.04.05.01)", N, AZUL, color="FFFFFF")
    fila(ws, 2, PROY, N, None, False)
    fila(ws, 3, FECHA + ". Precio publicado tal como lo muestra la tienda (con IGV, salvo la tarifa de agua). Precio sin IGV = precio publicado / contenido / factor IGV. "
             "Los precios de Sodimac y Promart son nacionales (ambas tienen tienda en Tarapoto); el precio de la tienda de Tarapoto puede variar.", N, None, False, alto=42)
    cab(ws, 5, ("N°", "INSUMO DE LA ACU", "UND", "PRODUCTO COTIZADO", "PRESENTACION", "PRECIO PUBLICADO S/", "CONTENIDO (UND ACU)",
                "IGV", "PRECIO S/ SIN IGV", "PRECIO ACTUAL ACU", "VARIACION", "TIENDA / FUENTE", "ENLACE", "ESTADO"), 42)
    REF = {}
    for i, (k, ins, u, prod, pres, p, cont, igv, tienda, url, est) in enumerate(COT):
        r = 6 + i; REF[k] = "'1 COTIZACION'!I%d" % r
        sin = est.startswith("SIN")
        c(ws, "A%d" % r, i + 1); c(ws, "B%d" % r, ins, h="left"); c(ws, "C%d" % r, u); c(ws, "D%d" % r, prod, h="left"); c(ws, "E%d" % r, pres)
        c(ws, "F%d" % r, p, AMARILLO, fmt="0.00"); c(ws, "G%d" % r, cont, VERDE if isinstance(cont, str) else AMARILLO, fmt="0.000")
        c(ws, "H%d" % r, igv, AMARILLO, fmt="0.00")
        act = ACTUAL[k]
        c(ws, "J%d" % r, act, fmt="0.00")
        if sin:
            c(ws, "I%d" % r, "=J%d" % r, ROJO, True, fmt="0.00")
        else:
            c(ws, "I%d" % r, "=F%d/G%d/H%d" % (r, r, r), VERDE, True, fmt="0.00")
        c(ws, "K%d" % r, '=IF(OR(J{0}="",F{0}=""),"",I{0}/J{0}-1)'.format(r), fmt="+0%;-0%")
        c(ws, "L%d" % r, tienda, h="left"); c(ws, "M%d" % r, "abrir" if url else "-", link=url or None)
        c(ws, "N%d" % r, est, ROJO if sin else (NARANJA if est.startswith("REF") else VERDE), h="left")
        ws.row_dimensions[r].height = 30
    r0 = 6 + len(COT)
    # acero ponderado y mano de obra
    r = r0
    c(ws, "A%d" % r, len(COT) + 1); c(ws, "B%d" % r, "Acero corrugado fy=4200 (ponderado de la tapa)", h="left", b=True); c(ws, "C%d" % r, "kg")
    c(ws, "D%d" % r, "1/2\": %.2f kg + 3/8\": %.2f kg por las 10 tapas (hoja COLECTOR ACERO filas 33 a 36)" % (D["kg12"], D["kg38"]), h="left")
    for col in "EFGH": c(ws, "%s%d" % (col, r), None)
    c(ws, "I%d" % r, "=(%s*%.3f+%s*%.3f)/%.3f" % ("I15", D["kg12"], "I16", D["kg38"], D["kg12"] + D["kg38"]), VERDE, True, fmt="0.00")
    c(ws, "J%d" % r, 3.81, fmt="0.00"); c(ws, "K%d" % r, "=I%d/J%d-1" % (r, r), fmt="+0%;-0%"); c(ws, "L%d" % r, "Sodimac Peru"); c(ws, "M%d" % r, "-")
    c(ws, "N%d" % r, "COMPROBADO (precio nacional)", VERDE, h="left"); REF["acero"] = "'1 COTIZACION'!I%d" % r
    mo = [("Operario", 29.30, 89.30), ("Oficial", 22.95, 69.75), ("Peon", 20.79, 62.80)]
    for j, (cat, hh, jor) in enumerate(mo):
        r = r0 + 1 + j
        c(ws, "A%d" % r, len(COT) + 2 + j); c(ws, "B%d" % r, cat, h="left"); c(ws, "C%d" % r, "hh")
        c(ws, "D%d" % r, "Jornal basico 2026 S/ %.2f por dia (convenio CAPECO-FTCCP 2026, R.M. 197-2025-TR); costo hora-hombre 2026 de la hoja COSTO HORA HOMBRE 2026 de la planilla" % jor, h="left")
        c(ws, "E%d" % r, "hora-hombre"); c(ws, "F%d" % r, hh, AMARILLO, fmt="0.00"); c(ws, "G%d" % r, 1, fmt="0.000"); c(ws, "H%d" % r, 1.00, fmt="0.00")
        c(ws, "I%d" % r, "=F%d/G%d/H%d" % (r, r, r), VERDE, True, fmt="0.00"); c(ws, "J%d" % r, hh, fmt="0.00"); c(ws, "K%d" % r, "=I%d/J%d-1" % (r, r), fmt="+0%;-0%")
        c(ws, "L%d" % r, "FTCCP - tablas salariales 2026"); c(ws, "M%d" % r, "abrir", link="https://www.ftccperu.com/media/attachments/2025/12/30/tablas-salariales-2026-construccion-civil.pdf")
        c(ws, "N%d" % r, "JORNAL OFICIAL 2026; costo hh 2026 = costo CAPECO 2024-2025 x jornal 2026 / jornal 2024-2025", VERDE, h="left"); REF[cat] = "'1 COTIZACION'!I%d" % r
        ws.row_dimensions[r].height = 30
    rn = r0 + 5
    fila(ws, rn, "NOTAS", N)
    notas = ["1. Verde = precio publicado y encontrado; naranja = referencia (otra ciudad o sin ciudad); rojo = sin precio publicado: se mantiene el precio actual y debe cotizarse en Tarapoto.",
             "2. Angulo L 2\"x2\"x3/16\": ninguna tienda en linea del Peru lo publica en soles; solo se vende por pedido a distribuidores. Por eso se propone el L 1 1/2\"x1 1/2\"x1/8\" (hoja 2 ANGULOS).",
             "3. Piedra chancada 1/2\": en Tarapoto se publica la de 3/4\" (S/ 70.00 el m3, Agregados Tarapoto, Morales); la de 1/2\" debe cotizarse. En la tapa (8 cm) su incidencia es menor de S/ 2.00.",
             "4. Arena (Agregados Tarapoto): se tomo el precio publicado como precio con IGV; si el proveedor no emite factura, usar S/ 48.00 directo (celda H11 = 1.00).",
             "5. Lija y thinner reemplazan al global \"Lija y thinner\" de la ACU actual (ver DESAGREGADO_LIJA_THINNER_TAPA_REGISTRO_HOGAR_REFUGIO.xlsx).",
             "6. Mano de obra 2026: jornal basico oficial (R.M. 197-2025-TR) y costo hora-hombre 29.30 / 22.95 / 20.79 de la hoja COSTO HORA HOMBRE 2026 de la planilla (costo CAPECO 2024-2025 llevado al jornal 2026)."]
    for i, t in enumerate(notas): fila(ws, rn + 1 + i, t, N, None, False, alto=28)
    ws.freeze_panes = "C6"

    # =============================================================== 2 ANGULOS
    NA = 8
    fila(wa, 1, "PROPUESTA DE ANGULOS PARA MARCO Y CONTRAMARCO - TAPA DE REGISTRO 0.68 x 0.68 x 0.08 (DD-01, DD-02, DP-06B)", NA, AZUL, color="FFFFFF")
    fila(wa, 2, PROY, NA, None, False)
    fila(wa, 3, "Los 10 registros (RS-01 a RS-07, CL y 2 en CC) estan fuera del ingreso vehicular OE-01 (prog. 0+026.32 a 0+031.33) y del porton de motos "
               "(0+039.04 a 0+041.04), bajo vereda y jardin: carga peatonal (E.020). El RS-03 (0+032.50) y el RS-04 (0+042.68) son los mas cercanos a los cruces.", NA, None, False, alto=42)
    fila(wa, 5, "1. ANGULOS COMERCIALES (barra de 6 m, ASTM A36)", NA)
    cab(wa, 6, ("N°", "ANGULO", "PESO kg/m", "ALA mm", "ESPESOR mm", "PRECIO BARRA S/ (con IGV)", "S/ POR kg SIN IGV", "OBSERVACION"))
    angs = [('L 1 1/2"x1 1/2"x1/8"', 1.83, 38.1, 3.18, "=F%d" % 0, "PROPUESTO para marco y contramarco: precio publicado (Sodimac Peru), 1 barra por tapa"),
            ('L 1"x1"x1/8"', 1.19, 25.4, 3.18, None, "Ala de 25 mm: asiento util de 15 mm con la holgura de 1 cm; no se recomienda"),
            ('L 2"x2"x1/8"', 2.46, 50.8, 3.18, None, "Sin precio publicado en tiendas"),
            ('L 2"x2"x3/16" (actual contramarco)', 3.63, 50.8, 4.76, None, "Sin precio publicado en tiendas; sobredimensionado para el asiento"),
            ('L 1 1/2"x1 1/2"x3/16"', 2.68, 38.1, 4.76, None, "Sin precio publicado en tiendas")]
    for i, (n_, kgm, ala, e, _, obs) in enumerate(angs):
        r = 7 + i
        c(wa, "A%d" % r, i + 1); c(wa, "B%d" % r, n_, h="left"); c(wa, "C%d" % r, kgm, AMARILLO, fmt="0.00"); c(wa, "D%d" % r, ala, fmt="0.0"); c(wa, "E%d" % r, e, fmt="0.00")
        if i == 0:
            c(wa, "F%d" % r, "='1 COTIZACION'!F6", VERDE, fmt="0.00"); c(wa, "G%d" % r, "=F%d/1.18/(6*C%d)" % (r, r), VERDE, True, fmt="0.00")
        elif i == 1:
            c(wa, "F%d" % r, "='1 COTIZACION'!F7", VERDE, fmt="0.00"); c(wa, "G%d" % r, "=F%d/1.18/(6*C%d)" % (r, r), VERDE, fmt="0.00")
        else:
            c(wa, "F%d" % r, "sin precio", ROJO); c(wa, "G%d" % r, "-", ROJO)
        c(wa, "H%d" % r, obs, VERDE if i == 0 else None, h="left"); wa.row_dimensions[r].height = 28
    fila(wa, 13, "2. VERIFICACION DEL CONTRAMARCO CON L 1 1/2\"x1 1/2\"x1/8\" (carga peatonal)", NA)
    cab(wa, 14, ("N°", "CONCEPTO", "VALOR", "UND", "", "", "", "SUSTENTO"))
    ver = [("Lado de la tapa", 0.68, "m", "DP-06B"),
           ("Espesor de la tapa", 0.08, "m", "DP-06B"),
           ("Peso de la tapa de concreto armado", "=C15^2*C16*2400", "kg", "peso unitario 2400 kg/m3 (E.020)"),
           ("Peso del marco L 1 1/2\"x1/8\"", "=4*C15*C7", "kg", "4 x 0.68 x 1.83"),
           ("Carga muerta D", "=C17+C18", "kg", ""),
           ("Sobrecarga peatonal", 500, "kg/m2", "E.020, veredas y pasadizos"),
           ("Carga viva L sobre la tapa", "=C20*C15^2", "kg", ""),
           ("Carga ultima Wu = 1.4 D + 1.7 L", "=1.4*C19+1.7*C21", "kg", "E.060"),
           ("Holgura entre tapa y abertura", 0.01, "m", "DD-01, 1 cm por lado"),
           ("Ancho de asiento util = ala - holgura", "=D7/1000-C23", "m", "el ala horizontal apoya entera sobre el borde engrosado de concreto (corte B de DD-01)"),
           ("Area de apoyo = 4 x 0.68 x ancho util", "=4*C15*C24*10000", "cm2", ""),
           ("Esfuerzo de aplastamiento sobre el concreto", "=C22/C25", "kg/cm2", ""),
           ("Resistencia al aplastamiento 0.65 x 0.85 x f'c (f'c = 210)", "=0.65*0.85*210", "kg/cm2", "E.060 10.17"),
           ("Verificacion", '=IF(C26<C27,"CUMPLE","NO CUMPLE")', "-", ""),
           ("Luz libre del registro = 0.70 - 2 x ala", "=0.70-2*D7/1000", "m", "con L 2\" era 0.60 m: mejora el acceso para limpieza"),
           ("Longitud por tapa: contramarco 4 x 0.706 + marco 4 x 0.68", "=4*(0.70+2*E7/1000)+4*0.68", "m", "cabe en 1 barra de 6.00 m (un solo producto para marco y contramarco)"),
           ("Barras de 6 m por tapa", "=ROUNDUP(C30/6,0)", "und", "")]
    for i, (t, v, u, s) in enumerate(ver):
        r = 15 + i
        c(wa, "A%d" % r, i + 1); c(wa, "B%d" % r, t, h="left"); c(wa, "C%d" % r, v, AMARILLO if not isinstance(v, str) else VERDE, b=(r == 28), fmt="0.00")
        c(wa, "D%d" % r, u)
        for col in "EFG": c(wa, "%s%d" % (col, r), None)
        c(wa, "H%d" % r, s, h="left"); wa.row_dimensions[r].height = 26
    fila(wa, 33, "3. CONCLUSION", NA)
    concl = ["El ala horizontal del contramarco descansa entera sobre el concreto del borde engrosado, de modo que el angulo no trabaja a flexion: protege la arista y sirve de asiento. "
             "El espesor de 3/16\" no aporta resistencia necesaria; con 1/8\" el esfuerzo de apoyo es muy inferior al admisible.",
             "Se propone un solo angulo L 1 1/2\"x1 1/2\"x1/8\" para marco y contramarco: precio publicado, una barra de 6 m por tapa, menos peso (10.10 kg en lugar de 15.14 kg por tapa) y mayor luz libre.",
             "Al usar ala de 1 1/2\", el asiento de concreto bajo el ala horizontal pasa de 50 mm a 38 mm. Los anclajes 3/8\" se sueldan al ala vertical con cordon de 1/8\" (espesor del angulo 3.2 mm).",
             "Si se aprueba, cambian: COLECTOR PARAMETROS (fila 36, 38 y 40), COLECTOR REGISTROS Y TAPAS (fila 23), DESAGREGADO TAPA REGISTRO, ACU TAPA REGISTRO, la descripcion de la "
             "planilla general (fila 317) y las laminas DD-01, DD-02, DP-06B y cuadros. Angulos de la partida: de 151.42 kg a 101.02 kg (10 tapas)."]
    for i, t in enumerate(concl): fila(wa, 34 + i, t, NA, None, False, alto=40)

    # =============================================================== 3 ACU COMPARATIVO
    wc = hoja("3 ACU COMPARATIVO", (5, 34, 7, 12, 11, 12, 12, 11, 12, 34))
    NC = 10
    fila(wc, 1, "ACU TAPA DE REGISTRO 0.68 x 0.68 C/MARCO Y CONTRAMARCO - ACTUAL FRENTE A PROPUESTA CON PRECIOS COTIZADOS (por und de tapa)", NC, AZUL, color="FFFFFF")
    fila(wc, 2, PROY, NC, None, False)
    fila(wc, 3, "Actual = ACU TAPA REGISTRO de la planilla vigente. Propuesta = marco y contramarco con L 1 1/2\"x1 1/2\"x1/8\", lija y thinner desagregados y precios de la hoja 1 COTIZACION. "
               "Ratios de COLECTOR INSUMOS sin cambio. Mano de obra con el costo hora-hombre de la planilla.", NC, None, False, alto=42)
    fila(wc, 5, "PARAMETROS DE LA PROPUESTA (por tapa)", NC)
    par = [("Perimetro del contramarco (4 x 0.70)", 2.80, "m"), ("Perimetro del marco (4 x 0.68)", 2.72, "m"),
           ("Peso del angulo L 1 1/2\"x1 1/2\"x1/8\"", "='2 ANGULOS'!C7", "kg/m"), ("Angulos por tapa", "=(D6+D7)*D8", "kg"),
           ("Desarrollo pintado del angulo 1 1/2\" (4 caras)", 0.152, "m2/m"), ("Area pintada por tapa", "=(D6+D7)*D10", "m2"),
           ("Concreto por tapa", D["conc"], "m3"), ("Acero corrugado por tapa", D["acero"], "kg")]
    for i, (t, v, u) in enumerate(par):
        r = 6 + i
        c(wc, "A%d" % r, i + 1); c(wc, "B%d" % r, t, h="left"); c(wc, "C%d" % r, u); c(wc, "D%d" % r, v, AMARILLO if not isinstance(v, str) else VERDE, fmt="0.0000")
        for col in "EFGHIJ": c(wc, "%s%d" % (col, r), None)
    cab(wc, 15, ("N°", "RECURSO", "UND", "CANT. ACTUAL", "PRECIO ACTUAL", "PARCIAL ACTUAL", "CANT. PROPUESTA", "PRECIO COTIZADO", "PARCIAL PROPUESTA", "OBSERVACION"), 32)
    F = D["filas"]
    # cantidades propuestas (por tapa), mismas reglas de la ACU actual
    A_KG, AREA, CONC, AC = "$D$9", "$D$11", "$D$12", "$D$13"
    rec = [
        ("MANO DE OBRA", None),
        ("Operario", 69, "=%s*0.1+%s*0.8+%s*0.032+%s*0.4+1" % (A_KG, CONC, AC, AREA), REF["Operario"], "angulos 0.1 + concreto 0.8 + acero 0.032 + pintura 0.4 + colocacion 1"),
        ("Oficial", 70, "=%s*0.1+%s*0.8+%s*0.032" % (A_KG, CONC, AC), REF["Oficial"], ""),
        ("Peon", 71, "=%s*6.4+1" % CONC, REF["Peon"], ""),
        ("MATERIALES", None),
        ('Angulo L 2"x2"x3/16"', 74, 0, None, "se elimina"),
        ('Angulo L 1 1/2"x1 1/2"x1/8"', 75, "=%s*1.05" % A_KG, REF["ang15"], "marco + contramarco, 5 % de desperdicio"),
        ('Soldadura cellocord E6011 1/8"', 76, "=%s*0.03" % A_KG, REF["sold"], ""),
        ("Cemento Portland tipo I", 77, "=%s*9.73" % CONC, REF["cem"], ""),
        ("Arena gruesa", 78, "=%s*0.52" % CONC, REF["are"], ""),
        ('Piedra chancada 1/2"', 79, "=%s*0.53" % CONC, REF["pie"], "precio actual (sin precio publicado de 1/2\")"),
        ("Agua", 80, "=%s*0.186" % CONC, REF["agua"], ""),
        ("Madera tornillo", 81, 1.15, REF["mad"], ""),
        ("Acero corrugado fy=4200", 82, "=%s*1.05" % AC, REF["acero"], "precio ponderado 1/2\" y 3/8\""),
        ("Alambre negro N 16", 83, "=%s*0.06" % AC, REF["ala"], ""),
        ("Pintura anticorrosiva", 84, "=%s*0.08" % AREA, REF["anti"], ""),
        ("Esmalte sintetico", 85, "=%s*0.08" % AREA, REF["esm"], ""),
        ("Lija y thinner", 86, 0, None, "se reemplaza por lija y thinner desagregados"),
        ("Lija para fierro N° 80", None, "=%s*0.25" % AREA, REF["lija"], "0.25 und/m2"),
        ("Thinner acrilico", None, "=%s*0.024" % AREA, REF["thin"], "0.024 gal/m2 (15 % de la pintura)"),
        ("EQUIPOS", None),
        ("Herramientas manuales", 89, None, None, "3 % de la mano de obra"),
        ("Mezcladora 9-11 p3", 90, "=%s*0.8" % CONC, REF["mez"], ""),
        ("Vibrador de concreto", 91, "=%s*0.8" % CONC, REF["vib"], ""),
        ("Maquina de soldar 250 A", 92, "=%s*0.1" % A_KG, REF["sol"], "precio actual (sin tarifa publicada)"),
    ]
    r = 16; n = 0; sub = {}; grupo = None; ini = {}
    for item in rec:
        if item[1] is None and len(item) == 2:
            if grupo:
                sub[grupo] = (ini[grupo], r - 1)
            grupo = item[0]; ini[grupo] = r
            fila(wc, r, grupo, NC); r += 1; continue
        nom, fa, qp, pref, obs = item
        n += 1
        c(wc, "A%d" % r, n); c(wc, "B%d" % r, nom, h="left")
        if nom == "Herramientas manuales":
            c(wc, "C%d" % r, "%mo"); c(wc, "D%d" % r, 0.03, fmt="0.00%"); c(wc, "E%d" % r, "=F%d" % 0, fmt="0.00")
            hm_row = r
        else:
            und = F[fa][1] if fa else ("und" if "Lija" in nom else "gal")
            c(wc, "C%d" % r, und)
            c(wc, "D%d" % r, F[fa][2] if fa else 0, fmt="0.0000"); c(wc, "E%d" % r, F[fa][3] if fa else None, fmt="0.00")
            c(wc, "F%d" % r, "=D%d*N(E%d)" % (r, r), fmt="0.00")
            c(wc, "G%d" % r, qp, VERDE if isinstance(qp, str) else AMARILLO, fmt="0.0000")
            c(wc, "H%d" % r, ("=" + pref) if pref else None, VERDE, fmt="0.00")
            c(wc, "I%d" % r, "=G%d*N(H%d)" % (r, r), fmt="0.00")
        c(wc, "J%d" % r, obs, h="left"); r += 1
    sub[grupo] = (ini[grupo], r - 1)
    # subtotales
    tot = {}
    for g in ("MANO DE OBRA", "MATERIALES", "EQUIPOS"):
        a_, b_ = sub[g]
        c(wc, "B%d" % r, "Subtotal " + g.lower(), CELESTE, True, "left")
        for col in "ACDEGHJ": c(wc, "%s%d" % (col, r), None, CELESTE)
        c(wc, "F%d" % r, "=SUM(F%d:F%d)" % (a_, b_), CELESTE, True, fmt="0.00"); c(wc, "I%d" % r, "=SUM(I%d:I%d)" % (a_, b_), CELESTE, True, fmt="0.00")
        tot[g] = r; r += 1
    # herramientas: 3 % mo
    mo = tot["MANO DE OBRA"]
    wc["E%d" % hm_row] = None
    c(wc, "F%d" % hm_row, "=D%d*F%d" % (hm_row, mo), fmt="0.00"); c(wc, "G%d" % hm_row, 0.03, fmt="0.00%"); c(wc, "I%d" % hm_row, "=G%d*I%d" % (hm_row, mo), fmt="0.00")
    c(wc, "E%d" % hm_row, "=F%d" % mo, fmt="0.00"); c(wc, "H%d" % hm_row, "=I%d" % mo, fmt="0.00")
    c(wc, "B%d" % r, "COSTO UNITARIO DIRECTO POR TAPA (S/ por und)", AZUL, True, "left", color="FFFFFF")
    for col in "ACDEGHJ": c(wc, "%s%d" % (col, r), None, AZUL)
    c(wc, "F%d" % r, "=F%d+F%d+F%d" % (tot["MANO DE OBRA"], tot["MATERIALES"], tot["EQUIPOS"]), AZUL, True, fmt="0.00", color="FFFFFF")
    c(wc, "I%d" % r, "=I%d+I%d+I%d" % (tot["MANO DE OBRA"], tot["MATERIALES"], tot["EQUIPOS"]), AZUL, True, fmt="0.00", color="FFFFFF")
    rcu = r; r += 1
    c(wc, "B%d" % r, "Comprobacion: costo unitario de la ACU vigente", h="left"); c(wc, "F%d" % r, D["cu"], fmt="0.00")
    for col in "ACDEGHIJ": c(wc, "%s%d" % (col, r), None)
    r += 1
    c(wc, "B%d" % r, "Diferencia por tapa (propuesta - actual)", h="left", b=True); c(wc, "I%d" % r, "=I%d-F%d" % (rcu, rcu), VERDE, True, fmt="+0.00;-0.00")
    for col in "ACDEFGHJ": c(wc, "%s%d" % (col, r), None)
    r += 1
    c(wc, "B%d" % r, "Partida: %d tapas" % D["n"], h="left", b=True); c(wc, "F%d" % r, "=F%d*%d" % (rcu, D["n"]), fmt="#,##0.00"); c(wc, "I%d" % r, "=I%d*%d" % (rcu, D["n"]), VERDE, True, fmt="#,##0.00")
    c(wc, "J%d" % r, "=\"ahorro de S/ \"&TEXT(F%d-I%d,\"#,##0.00\")" % (r, r), h="left")
    for col in "ACDEGH": c(wc, "%s%d" % (col, r), None)
    wc.freeze_panes = "C16"
    for w in (ws, wa, wc):
        w.page_setup.orientation = "landscape"; w.page_setup.fitToWidth = 1; w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.calculation.fullCalcOnLoad = True
    wb.save(SAL)
    return SAL


if __name__ == "__main__":
    print(construir())
