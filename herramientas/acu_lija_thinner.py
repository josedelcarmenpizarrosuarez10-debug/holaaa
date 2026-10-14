"""Desagregado del insumo global "Lija y thinner" de la hoja ACU TAPA REGISTRO del Hogar de Refugio Temporal
(Mujeres Violentadas, CUI 2675514) en lija para fierro (und) y thinner (gal). Excel aparte, con formulas.
Datos de la planilla vigente: area pintada por tapa, pintura anticorrosiva y esmalte por m2, numero de tapas.
"""
import os, warnings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLAN = os.path.join(RAIZ, "insumos", "metrados_vigentes", "METRADO_DRENAJE_MUJERES_VIOLENTADAS.xlsx")
SAL = os.path.join(RAIZ, "entregables", "DESAGREGADO_LIJA_THINNER_TAPA_REGISTRO_HOGAR_REFUGIO.xlsx")
AZUL, CELESTE, AMARILLO, VERDE = "1F4E79", "DDEBF7", "FFF2CC", "E2EFDA"


def datos():
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(PLAN, data_only=True)
    a = wb["ACU TAPA REGISTRO"]; ci = wb["COLECTOR INSUMOS"]
    return dict(area=a["F17"].value, tapas=a["E17"].value, anti=ci["G64"].value, esm=ci["G65"].value,
                glb=ci["G66"].value, p_glb=a["F86"].value, proyecto=wb["RESUMEN"]["C5"].value or "")


def construir():
    D = datos()
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "LIJA Y THINNER"
    fino = Side(style="thin", color="808080"); borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    for col, w in zip("ABCDEFGH", (6, 46, 9, 14, 14, 12, 12, 62)): ws.column_dimensions[col].width = w

    def c(ref, v, fill=None, b=False, h="center", fmt=None, color="000000", wrap=True, sz=10):
        x = ws[ref]; x.value = v; x.font = Font(name="Arial", size=sz, bold=b, color=color)
        x.alignment = Alignment(horizontal=h, vertical="center", wrap_text=wrap); x.border = borde
        if fill: x.fill = PatternFill("solid", fgColor=fill)
        if fmt: x.number_format = fmt

    def fila(r, texto, fill=CELESTE, b=True, color="000000"):
        ws.merge_cells("A%d:H%d" % (r, r)); c("A%d" % r, texto, fill, b, "left", color=color)
        for col in "BCDEFGH": ws["%s%d" % (col, r)].border = borde

    fila(1, "DESAGREGADO DEL INSUMO \"LIJA Y THINNER\" - ACU TAPA REGISTRO (01.04.04.05.04 TAPA DE CONCRETO ARMADO 0.68 x 0.68 x 0.08 m)", AZUL, color="FFFFFF")
    fila(2, "PROYECTO: HOGAR DE REFUGIO TEMPORAL - MUJERES VIOLENTADAS, DISTRITO DE MORALES, SAN MARTIN - CUI N.° 2675514", None, False)
    fila(3, "Reemplaza la fila 16 \"Lija y thinner\" (glb) de la ACU TAPA REGISTRO y la fila 66 de COLECTOR INSUMOS. "
            "Celda amarilla = dato (se puede cambiar); verde = formula. Precios sin IGV, referenciales: confirmar con cotizacion en Tarapoto.", None, False)
    ws.row_dimensions[3].height = 30

    # ---- 1. datos de la planilla
    fila(5, "1. DATOS TOMADOS DE LA PLANILLA (ACU TAPA REGISTRO y COLECTOR INSUMOS)")
    for col, t in zip("ABCDH", ("N°", "DATO", "UND", "VALOR", "FUENTE")): c("%s6" % col, t, AZUL, True, color="FFFFFF")
    for col in "EFG": c("%s6" % col, None, AZUL)
    datos_f = [("Area pintada de angulos por tapa (marco + contramarco)", "m2", D["area"], "ACU TAPA REGISTRO F17: 2.80 x 0.203 + 2.72 x 0.152 = 0.98 m2"),
               ("Pintura anticorrosiva por m2", "gal/m2", D["anti"], "COLECTOR INSUMOS G64"),
               ("Esmalte sintetico por m2", "gal/m2", D["esm"], "COLECTOR INSUMOS G65"),
               ("Numero de tapas de registro", "und", D["tapas"], "ACU TAPA REGISTRO E17 (metrado de tapas)"),
               ("Global actual \"Lija y thinner\" por m2", "glb/m2", D["glb"], "COLECTOR INSUMOS G66 (se reemplaza)"),
               ("Precio actual del global", "S/ glb", D["p_glb"], "ACU TAPA REGISTRO F86 (se reemplaza)")]
    for i, (t, u, v, s) in enumerate(datos_f):
        r = 7 + i
        c("A%d" % r, i + 1); c("B%d" % r, t, h="left"); c("C%d" % r, u); c("D%d" % r, v, AMARILLO, fmt="0.0000"); c("H%d" % r, s, h="left")
        for col in "EFG": c("%s%d" % (col, r), None)

    # ---- 2. criterios de consumo
    fila(14, "2. CRITERIO DE CONSUMO")
    for col, t in zip("ABCDH", ("N°", "CRITERIO", "UND", "VALOR", "SUSTENTO")): c("%s15" % col, t, AZUL, True, color="FFFFFF")
    for col in "EFG": c("%s15" % col, None, AZUL)
    crit = [("Thinner: porcentaje del volumen de pintura (dilucion + limpieza)", "%", 0.15,
             "10 % de dilucion (pintura alquidica, segun ficha del fabricante) + 5 % de limpieza de brochas y equipos"),
            ("Thinner por m2 = % x (anticorrosivo + esmalte)", "gal/m2", "=D16*(D8+D9)", "formula"),
            ("Lija para fierro N° 80 por m2", "und/m2", 0.25, "1 pliego rinde ~4 m2 en la limpieza y desoxidado del angulo antes del anticorrosivo")]
    for i, (t, u, v, s) in enumerate(crit):
        r = 16 + i
        c("A%d" % r, i + 1); c("B%d" % r, t, h="left"); c("C%d" % r, u)
        c("D%d" % r, v, VERDE if isinstance(v, str) else AMARILLO, fmt="0.00%" if u == "%" else "0.0000"); c("H%d" % r, s, h="left")
        for col in "EFG": c("%s%d" % (col, r), None)

    # ---- 3. desagregado por tapa (para el cuadro de recursos de la ACU)
    fila(20, "3. DESAGREGADO POR TAPA (filas para el cuadro de recursos de la ACU TAPA REGISTRO)")
    for col, t in zip("ABCDEFGH", ("N°", "RECURSO", "UND", "RATIO POR m2", "CANTIDAD POR TAPA", "PRECIO S/", "PARCIAL S/", "FUENTE DEL PRECIO")):
        c("%s21" % col, t, AZUL, True, color="FFFFFF")
    rec = [("16", "Lija para fierro N° 80", "und", "=D18", 2.12, "Sodimac Peru, lija 3M fierro ASA grano 80: S/ 2.50 con IGV -> 2.50 / 1.18 = 2.12"),
           ("17", "Thinner acrilico", "gal", "=D17", 35.65, "Sodimac Peru, thinner acrilico TOPEX 3.5 L a S/ 38.90 con IGV -> x 3.785 / 3.5 / 1.18 = 35.65 por galon")]
    for i, (n, t, u, ratio, p, s) in enumerate(rec):
        r = 22 + i
        c("A%d" % r, n); c("B%d" % r, t, h="left"); c("C%d" % r, u); c("D%d" % r, ratio, VERDE, fmt="0.0000")
        c("E%d" % r, "=D%d*$D$7" % r, VERDE, fmt="0.0000"); c("F%d" % r, p, AMARILLO, fmt="0.00"); c("G%d" % r, "=E%d*F%d" % (r, r), VERDE, fmt="0.00")
        c("H%d" % r, s, h="left")
    c("A24", None, CELESTE); c("B24", "Total lija + thinner por tapa", CELESTE, True, "left")
    for col in "CDEF": c("%s24" % col, None, CELESTE)
    c("G24", "=G22+G23", CELESTE, True, fmt="0.00"); c("H24", None, CELESTE)
    c("A25", None); c("B25", "Antes: \"Lija y thinner\" global por tapa", h="left"); c("C25", "glb"); c("D25", "=D11", fmt="0.0000")
    c("E25", "=D25*D7", fmt="0.0000"); c("F25", "=D12", fmt="0.00"); c("G25", "=E25*F25", fmt="0.00"); c("H25", "valor actual de la ACU", h="left")
    c("A26", None); c("B26", "Diferencia en el costo unitario de la tapa", h="left", b=True)
    for col in "CDEF": c("%s26" % col, None)
    c("G26", "=G24-G25", VERDE, True, fmt="+0.00;-0.00"); c("H26", "S/ por tapa (sube o baja el costo unitario directo)", h="left")

    # ---- 4. total de obra
    fila(28, "4. TOTAL DE OBRA (para la lista de insumos)")
    for col, t in zip("ABCDEFGH", ("N°", "INSUMO", "UND", "POR TAPA", "N° TAPAS", "TOTAL", "A COMPRAR", "OBSERVACION")):
        c("%s29" % col, t, AZUL, True, color="FFFFFF")
    for i, (t, u, compra, obs) in enumerate((("Lija para fierro N° 80", "und", "=ROUNDUP(F30,0)", "pliegos enteros"),
                                              ("Thinner acrilico", "gal", "=ROUNDUP(F31*4,0)/4", "se compra por 1/4 de galon como minimo"))):
        r = 30 + i
        c("A%d" % r, i + 1); c("B%d" % r, t, h="left"); c("C%d" % r, u); c("D%d" % r, "=E%d" % (22 + i), VERDE, fmt="0.0000")
        c("E%d" % r, "=$D$10", VERDE, fmt="0"); c("F%d" % r, "=D%d*E%d" % (r, r), VERDE, fmt="0.000"); c("G%d" % r, compra, VERDE, fmt="0.00")
        c("H%d" % r, obs, h="left")

    fila(33, "NOTAS", CELESTE)
    notas = ["1. En la seccion E del desglose de la ACU (fila 52) el global se divide en dos filas: lija = area por tapa x 0.25 und/m2; thinner = area por tapa x ratio de la fila 17.",
             "2. En COLECTOR INSUMOS (fila 66) el global 0.05 glb/m2 se reemplaza por lija 0.25 und/m2 y thinner 0.0240 gal/m2.",
             "3. Los ratios y precios son criterio y referencia de mercado (Sodimac Peru); el ingeniero de presupuesto puede cambiar las celdas amarillas y todo se recalcula."]
    for i, t in enumerate(notas):
        fila(34 + i, t, None, False); ws.row_dimensions[34 + i].height = 28
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(SAL)
    return SAL


if __name__ == "__main__":
    print(construir())
