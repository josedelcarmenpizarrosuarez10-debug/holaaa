# -*- coding: utf-8 -*-
"""Manual técnico de drenaje pluvial (Word) a partir de la base de conocimiento en Markdown (conocimiento/NN_*.md).

Convierte el Markdown simple de los capítulos (encabezados #/##/###/####, párrafos, listas con - o 1., tablas con |,
bloques ``` de pseudocódigo, **negrita** y `código`) a un Word con carátula, índice automático (se actualiza al abrir),
encabezado, pie con número de página y un capítulo por archivo.

Uso: python3 herramientas/conocimiento_docx.py
Salida: conocimiento/MANUAL_TECNICO_DRENAJE_PLUVIAL.docx
"""
import glob
import os
import re

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CARPETA = os.path.join(RAIZ, 'conocimiento')
SALIDA = os.path.join(CARPETA, 'MANUAL_TECNICO_DRENAJE_PLUVIAL.docx')
AZUL, GRIS, CELESTE = '1F3864', '595959', 'DCE6F2'


def _sombra(c, color):
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), color)
    c._tc.get_or_add_tcPr().append(sh)


def _campo(p, instr):
    r = p.add_run(); a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin'); r._r.append(a)
    r = p.add_run(); t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r._r.append(t)
    r = p.add_run(); a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'separate'); r._r.append(a)
    p.add_run('(clic derecho > Actualizar campos)' if 'TOC' in instr else '1')
    r = p.add_run(); a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'end'); r._r.append(a)


def _runs(p, texto, tam=None, color=None, neg=False):
    """Texto con **negrita** y `código`."""
    for parte in re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', texto):
        if not parte:
            continue
        if parte.startswith('**') and parte.endswith('**'):
            r = p.add_run(parte[2:-2]); r.bold = True
        elif parte.startswith('`') and parte.endswith('`'):
            r = p.add_run(parte[1:-1]); r.font.name = 'Consolas'; r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            r.font.color.rgb = RGBColor.from_string('1F4E79')
        else:
            r = p.add_run(parte)
        if tam: r.font.size = Pt(tam)
        if neg: r.bold = True
        if color and not (parte.startswith('`')): r.font.color.rgb = RGBColor.from_string(color)


def _tabla(d, filas, ancho):
    filas = [f for f in filas if not re.match(r'^\s*\|?\s*:?-{2,}', f)]
    celdas = [[c.strip() for c in f.strip().strip('|').split('|')] for f in filas]
    n = max(len(c) for c in celdas)
    tb = d.add_table(rows=0, cols=n); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(celdas):
        cel = tb.add_row().cells
        for j in range(n):
            txt = fila[j] if j < len(fila) else ''
            p = cel[j].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            _runs(p, txt, 8, 'FFFFFF' if i == 0 else None, neg=(i == 0))
            if i == 0: _sombra(cel[j], AZUL)
        if i == 0:
            e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(e)
        tb.rows[-1]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    # anchos de columna según el largo del texto (entre 9 y 60 caracteres por columna)
    pesos = [min(60, max(9, max(len(f[j]) if j < len(f) else 0 for f in celdas))) for j in range(n)]
    tb.autofit = False
    for gc, p in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), pesos):
        gc.set(qn('w:w'), str(int(ancho * p / sum(pesos) * 567)))
    for fila in tb.rows:
        for c, p in zip(fila.cells, pesos):
            c.width = Cm(ancho * p / sum(pesos))
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def _codigo(d, lineas):
    tb = d.add_table(rows=1, cols=1); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tb.rows[0].cells[0]; _sombra(c, 'F2F2F2')
    p = c.paragraphs[0]
    for k, l in enumerate(lineas):
        r = p.add_run(l); r.font.name = 'Consolas'; r.font.size = Pt(7.5); r._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
        if k < len(lineas) - 1: r.add_break()
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def capitulo(d, ruta, ancho):
    lineas = open(ruta, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lineas):
        l = lineas[i]
        if l.startswith('```'):
            j = i + 1
            while j < len(lineas) and not lineas[j].startswith('```'): j += 1
            _codigo(d, lineas[i + 1:j]); i = j + 1; continue
        if l.strip().startswith('|'):
            j = i
            while j < len(lineas) and lineas[j].strip().startswith('|'): j += 1
            _tabla(d, lineas[i:j], ancho); i = j; continue
        m = re.match(r'^(#{1,4})\s+(.*)', l)
        if m:
            nivel = len(m.group(1)); txt = re.sub(r'\*\*', '', m.group(2))
            h = d.add_heading(txt, level=nivel)
            if nivel == 1: h.paragraph_format.page_break_before = True
            i += 1; continue
        m = re.match(r'^(\s*)([-*]|\d+[.)])\s+(.*)', l)
        if m:
            sangria = len(m.group(1).replace('\t', '  ')) // 2
            estilo = 'List Number' if m.group(2)[0].isdigit() else 'List Bullet'
            if sangria: estilo += ' 2'
            txt = m.group(3); j = i + 1                    # líneas de continuación del mismo punto (sangradas)
            while j < len(lineas) and re.match(r'^\s+\S', lineas[j]) and not re.match(r'^\s*([-*]|\d+[.)])\s|^\s*\||^\s*```', lineas[j]):
                txt += ' ' + lineas[j].strip(); j += 1
            p = d.add_paragraph(style=estilo); p.paragraph_format.space_after = Pt(1)
            _runs(p, txt); i = j; continue
        if l.strip() in ('', '---'):
            i += 1; continue
        txt = l.strip(); j = i + 1                        # párrafo: une líneas seguidas
        while j < len(lineas) and lineas[j].strip() and not re.match(r'^(#|```|\s*\||\s*([-*]|\d+[.)])\s)', lineas[j]):
            txt += ' ' + lineas[j].strip(); j += 1
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(4); _runs(p, txt)
        i = j


def generar():
    d = docx.Document()
    s = d.sections[0]; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = Cm(2.0); s.bottom_margin = Cm(1.8)
    ancho = 17.0
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    for k, tam in ((1, 16), (2, 13), (3, 11), (4, 10)):
        h = d.styles['Heading %d' % k]; h.font.name = 'Arial'; h.font.size = Pt(tam); h.font.color.rgb = RGBColor.from_string(AZUL)
        h.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); h.element.rPr.rFonts.set(qn('w:ascii'), 'Arial'); h.element.rPr.rFonts.set(qn('w:hAnsi'), 'Arial')
    if d.paragraphs:
        d.element.body.remove(d.paragraphs[0]._p)
    # carátula
    for txt, tam, color, antes in (('MANUAL TÉCNICO DE DRENAJE PLUVIAL', 24, AZUL, 150),
                                   ('Criterios, métodos, formatos y algoritmos', 14, GRIS, 10),
                                   ('Colector pluvial frontal, cunetas sobre terreno, planos, metrado, especificaciones y documentos del expediente', 11, GRIS, 10),
                                   ('Base de conocimiento de los proyectos de Morales y Barranquita (San Martín)', 10, GRIS, 40)):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(antes)
        r = p.add_run(txt); r.bold = tam > 12; r.font.size = Pt(tam); r.font.color.rgb = RGBColor.from_string(color)
    # índice
    sec = d.add_section(WD_SECTION.NEW_PAGE)
    hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT; sec.header.is_linked_to_previous = False
    r = hp.add_run('Manual técnico de drenaje pluvial'); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(GRIS)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; sec.footer.is_linked_to_previous = False
    _campo(fp, 'PAGE')
    p = d.add_paragraph(); r = p.add_run('ÍNDICE'); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(AZUL)
    _campo(d.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    for ruta in sorted(glob.glob(os.path.join(CARPETA, '[0-9][0-9]_*.md'))):
        capitulo(d, ruta, ancho)
    upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'), 'true'); d.settings.element.append(upd)
    d.core_properties.author = ''; d.core_properties.title = 'Manual técnico de drenaje pluvial'
    d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
