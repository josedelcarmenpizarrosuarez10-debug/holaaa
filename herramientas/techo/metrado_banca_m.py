"""Metrado de canaleta y accesorios pluviales de las bancas de jardineria tipo 01 - CAR MUJERES (CUI 2717013).
Cantidades por banca tomadas de las laminas DB-01 a DB-03 (planos_banca_m.py); el numero de bancas es un dato (celda amarilla).
No incluye el tubo PVC ni la falsa columna.
Uso: python3 herramientas/techo/metrado_banca_m.py
"""
import os, sys
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import planos_banca_m as BM

SAL = os.path.join(RAIZ, "entregables_mujeres", "METRADO_CANALETA_Y_ACCESORIOS_BANCAS_CAR_MUJERES.xlsx")
N_BANCAS = 16
AZUL, AMA, VER, GRIS = "1F4E79", "FFF2CC", "E2EFDA", "F2F2F2"
_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
F = lambda **k: Font(name="Arial", size=k.pop("size", 9), **k)

# (item, descripcion, und, cantidad por banca, sustento)
FILAS = [
    ("01", "CANALETA DE PLANCHA DE ALUZINC e=0.50mm, SECCION INTERIOR 0.18 m (ANCHO) x 0.20 m (ALTO)", "m", BM.AX,
     "1 tramo de 2.00 m por techo (techo a una agua, sin uniones de canaleta)"),
    ("02", "TAPA LATERAL DE PLANCHA DE ALUZINC e=0.50mm P/CANALETA", "und", 2, "Una en cada extremo del tramo"),
    ("03", "SOPORTE DE PLATINA F°G° 1\" x 1/8\" P/CANALETA", "und", len(BM.SOPS), "A 0.40 m de cada extremo, 1.20 m entre ellos (DB-03)"),
    ("04", "BOQUILLA DE PLANCHA DE ALUZINC e=0.50mm P/CONEXION DE CANALETA A TUBERIA PVC 4\"", "und", 1,
     "Brida 0.16 x 0.16 con 8 remaches y sellador, cuello Ø 0.10 x 0.15; a 0.30 m del extremo"),
    ("05", "CODO PVC 4\" x 90°", "und", 2, "Bajo la boquilla y sobre la bajada (corte A-A)"),
    ("06", "COLGADOR F°G° C/VARILLA ROSCADA 3/8\" P/TUBERIA PVC 4\"", "und", 1, "Sobre el tramo horizontal, colgado de la viga metalica 2\" x 2\""),
    ("07", "ABRAZADERA DE PLATINA F°G° 1\" x 1/8\" P/TUBERIA PVC 4\"", "und", 1, "En la bajada, fijada a la columna metalica (+1.77)"),
]


def main():
    wb = Workbook(); ws = wb.active; ws.title = "METRADO"
    for col, w in zip("ABCDEFG", (8, 62, 7, 13, 12, 13, 62)): ws.column_dimensions[col].width = w
    ws.merge_cells("A1:G1"); ws["A1"] = "METRADO DE CANALETA Y ACCESORIOS PLUVIALES - BANCAS DE JARDINERIA TIPO 01"
    ws["A1"].font = F(bold=True, size=12, color=AZUL); ws["A1"].alignment = Alignment(horizontal="center")
    datos = [("ENTIDAD", "GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO"),
             ("PROYECTO", "“CREACIÓN DEL SERVICIO DE PROTECCIÓN INTEGRAL A NIÑAS, NIÑOS Y ADOLESCENTES SIN CUIDADOS PARENTALES O EN RIESGO DE "
                          "PERDERLOS EN CENTRO DE ACOGIDA RESIDENCIAL - MUJERES DISTRITO DE MORALES DE LA PROVINCIA DE SAN MARTÍN DEL DEPARTAMENTO "
                          "DE SAN MARTÍN”, con CUI N.° 2717013."),
             ("UBICACIÓN", "SAN MARTIN - SAN MARTIN - MORALES"), ("FECHA", "OCTUBRE 2026")]
    for i, (a, b) in enumerate(datos, 3):
        ws.cell(i, 1, a).font = F(bold=True); ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
        c = ws.cell(i, 2, b); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 36
    ws.cell(8, 1, "N.° DE TECHOS IGUALES (BANCAS)").font = F(bold=True)
    ws.merge_cells("A8:C8")
    c = ws.cell(8, 4, N_BANCAS); c.fill = PatternFill("solid", fgColor=AMA); c.font = F(bold=True, size=11); c.border = BOR
    c.alignment = Alignment(horizontal="center")
    cab = ["ITEM", "DESCRIPCION", "UND.", "METRADO POR TECHO", "N.° DE TECHOS", "TOTAL", "SUSTENTO (LAMINAS DB-01 A DB-03)"]
    for j, t in enumerate(cab, 1):
        c = ws.cell(10, j, t); c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOR
    ws.row_dimensions[10].height = 28
    for i, (it, desc, und, cant, sust) in enumerate(FILAS, 11):
        vals = [it, desc, und, cant, "=$D$8", "=D%d*E%d" % (i, i), sust]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.font = F(); c.border = BOR
            c.alignment = Alignment(horizontal="center" if j in (1, 3, 4, 5, 6) else "left", vertical="center", wrap_text=True)
        ws.cell(i, 4).fill = PatternFill("solid", fgColor=AMA)
        for j in (5, 6): ws.cell(i, j).fill = PatternFill("solid", fgColor=VER)
        fmt = "0.00" if und == "m" else "0"
        for j in (4, 6): ws.cell(i, j).number_format = fmt
        ws.cell(i, 6).font = F(bold=True)
        ws.row_dimensions[i].height = 26
    n = 11 + len(FILAS) + 1
    ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=7)
    c = ws.cell(n, 1, "Notas: no incluye el tubo PVC 4\" ni la falsa columna. Un solo tramo de canaleta por techo (a una agua), sin uniones. Celda amarilla = dato; celda verde = formula. "
                      "Cantidades por banca segun el detalle de banca tipo 01 y las laminas DB-01 (isometrico), DB-02 (detalles) y DB-03 (cortes A-A y B-B).")
    c.font = F(italic=True, size=8); c.alignment = Alignment(wrap_text=True); c.fill = PatternFill("solid", fgColor=GRIS)
    ws.row_dimensions[n].height = 30
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    wb.save(SAL); print(SAL)


if __name__ == "__main__":
    main()
