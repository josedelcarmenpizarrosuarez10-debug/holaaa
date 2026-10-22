"""Metrado de canaleta, tapas, sujetadores y boquillas de la OC-01 (jardinera + banca + pergola) - HOGAR DE REFUGIO TEMPORAL
(CUI 2675514). Cantidades por techo de las laminas DO-01 y DO-02 (planos_oc01_refugio.py); el numero de estructuras OC-01 es un
dato (celda amarilla). Solo los cuatro items pedidos por el usuario: no incluye tubo PVC, codos, colgador, abrazadera ni falsa columna.
Uso: python3 herramientas/techo/metrado_oc01_refugio.py
"""
import os, sys
AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, AQUI); sys.path.insert(0, os.path.dirname(AQUI))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import planos_oc01_refugio as OC
import dxf_base as B

SAL = os.path.join(RAIZ, "entregables", "METRADO_CANALETA_OC-01_BANCA_PERGOLA_HOGAR_REFUGIO.xlsx")
N_OC = 1
AZUL, AMA, VER, GRIS = "1F4E79", "FFF2CC", "E2EFDA", "F2F2F2"
_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
F = lambda **k: Font(name="Arial", size=k.pop("size", 9), **k)
T1, T2, T3 = OC.TECHOS

# (item, descripcion, und, cantidad en T-1, T-2, T-3, sustento)
FILAS = [
    ("01", "CANALETA DE PLANCHA GALVANIZADA e=0.50mm, SECCION INTERIOR 0.18 m (ANCHO) x 0.21 m (ALTO)", "m",
     T1["L"], T2["L"], T3["L"], "Un tramo por techo, del largo del techo (planta de techos OC-01), S = 0.5 % hacia la boquilla"),
    ("02", "TAPA LATERAL DE PLANCHA GALVANIZADA e=0.50mm P/CANALETA", "und", 2, 2, 2, "Una en cada extremo de cada canaleta"),
    ("03", "SOPORTE (SUJETADOR) DE PLATINA F°G° 1\" x 1/8\" P/CANALETA", "und", len(T1["sop"]), len(T2["sop"]), len(T3["sop"]),
     "2 por canaleta: a 0.80 m (T-1, T-3) y 0.85 m (T-2) de cada extremo"),
    ("04", "BOQUILLA DE PLANCHA GALVANIZADA e=0.50mm P/CONEXION DE CANALETA A TUBERIA PVC 4\"", "und", 1, 1, 1,
     "Una por canaleta, en el punto del sumidero del plano (la cupula se reemplaza por boquilla)"),
]


def main():
    wb = Workbook(); ws = wb.active; ws.title = "METRADO OC-01"
    for col, w in zip("ABCDEFGHIJ", (7, 58, 6, 10, 10, 10, 12, 10, 12, 52)): ws.column_dimensions[col].width = w
    ws.merge_cells("A1:J1"); ws["A1"] = "METRADO DE CANALETA Y ACCESORIOS - OC-01 (JARDINERA + BANCA + PERGOLA)"
    ws["A1"].font = F(bold=True, size=12, color=AZUL); ws["A1"].alignment = Alignment(horizontal="center")
    proy = B.PROYECTO if isinstance(B.PROYECTO, str) else " ".join(B.PROYECTO)
    datos = [("ENTIDAD", "GERENCIA TERRITORIAL BAJO MAYO - TARAPOTO"), ("PROYECTO", proy),
             ("UBICACIÓN", B.UBICACION.replace("UBICACION: ", "")), ("FECHA", "OCTUBRE 2026")]
    for i, (a, b) in enumerate(datos, 3):
        ws.cell(i, 1, a).font = F(bold=True); ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=10)
        c = ws.cell(i, 2, b); c.font = F(); c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[4].height = 26
    ws.merge_cells("A8:C8"); ws.cell(8, 1, "N.° DE ESTRUCTURAS OC-01").font = F(bold=True)
    c = ws.cell(8, 4, N_OC); c.fill = PatternFill("solid", fgColor=AMA); c.font = F(bold=True, size=11); c.border = BOR
    c.alignment = Alignment(horizontal="center")
    cab = ["ITEM", "DESCRIPCION", "UND.", "TECHO T-1 (3.79)", "TECHO T-2 (4.21)", "TECHO T-3 (3.79)", "POR OC-01", "N.° OC-01",
           "TOTAL", "SUSTENTO (LAMINAS DO-01 Y DO-02)"]
    for j, t in enumerate(cab, 1):
        c = ws.cell(10, j, t); c.font = F(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOR
    ws.row_dimensions[10].height = 30
    for i, (it, desc, und, a, b, cc, sust) in enumerate(FILAS, 11):
        vals = [it, desc, und, a, b, cc, "=SUM(D%d:F%d)" % (i, i), "=$D$8", "=G%d*H%d" % (i, i), sust]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, v); c.font = F(bold=(j == 9)); c.border = BOR
            c.alignment = Alignment(horizontal="left" if j in (2, 10) else "center", vertical="center", wrap_text=True)
        fmt = "0.00" if und == "m" else "0"
        for j in range(4, 10): ws.cell(i, j).number_format = fmt
        for j in (4, 5, 6): ws.cell(i, j).fill = PatternFill("solid", fgColor=AMA)
        for j in (7, 8, 9): ws.cell(i, j).fill = PatternFill("solid", fgColor=VER)
        ws.row_dimensions[i].height = 28
    n = 11 + len(FILAS) + 1
    ws.merge_cells(start_row=n, start_column=1, end_row=n, end_column=10)
    c = ws.cell(n, 1, "Notas: canaleta, tapas y boquilla de plancha galvanizada e = 0.50 mm; seccion interior 0.18 de ancho x 0.21 de alto. "
                      "Un tramo de canaleta por techo, sin uniones. No incluye tubo PVC 4\", codos, colgador, abrazadera ni falsa columna. "
                      "Celda amarilla = dato; celda verde = formula.")
    c.font = F(italic=True, size=8); c.alignment = Alignment(wrap_text=True); c.fill = PatternFill("solid", fgColor=GRIS)
    ws.row_dimensions[n].height = 30
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False
    wb.save(SAL); print(SAL)


if __name__ == "__main__":
    main()
