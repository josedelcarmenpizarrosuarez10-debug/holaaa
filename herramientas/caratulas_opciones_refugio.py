# -*- coding: utf-8 -*-
"""Cuatro opciones de carátula en color guinda - Hogar de Refugio Temporal "Mujeres Violentadas" (CUI 2675514), memoria
descriptiva de arquitectura (Solange), en un solo Word con vista previa.

Datos del proyecto de caratulas_refugio.py (escudo, entidad, proyecto, CUI, foto de la usuaria, fecha) y servicio de la carátula
de la memoria de arquitectura. La página 1 muestra las cuatro opciones en miniatura; luego cada opción en su página.
  Opción 1 - Banda lateral guinda.   Opción 2 - Franjas superior e inferior.
  Opción 3 - Marco clásico.          Opción 4 - Bloques.

Uso: python3 herramientas/caratulas_opciones_refugio.py
Salida: entregables/arquitectura/OPCIONES_CARATULA_GUINDA_MEMORIA_ARQUITECTURA_HOGAR_REFUGIO.docx y vistas previas PNG en
        entregables/figuras_caratula/
"""
import os
import subprocess
import sys
import tempfile

import docx
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import caratulas_refugio as CR  # noqa: E402  (fija en caratulas_v los datos del Hogar de Refugio: proyecto, CUI, foto, fecha)

V = CR.V

C = V.C
G, G_CLARO, G_OSCURO, GRIS, ROSA = '7B1E2B', 'F3E6E8', '5A1520', '595959', 'F2D7DB'
TITULO, SUBTITULO, ESPEC_CORTA = 'MEMORIA DESCRIPTIVA', 'ESPECIALIDAD DE ARQUITECTURA', 'ARQUITECTURA'
SERVICIO = 'SERVICIO DE CONSULTORÍA PARA LA ELABORACIÓN DE LA ESPECIALIDAD DE ARQUITECTURA DEL PROYECTO:'
NOMBRE = 'Hogar de Refugio Temporal - Mujeres Violentadas'
SALIDA = os.path.join(RAIZ, 'entregables/arquitectura/OPCIONES_CARATULA_GUINDA_MEMORIA_ARQUITECTURA_HOGAR_REFUGIO.docx')
NOMBRES = ['Opción 1 - Banda lateral guinda', 'Opción 2 - Franjas', 'Opción 3 - Marco clásico',
           'Opción 4 - Bloques']
par, sombra, bordes, sin_bordes, anchos, linea = C._par, C._sombra, C._bordes, C._sin_bordes, C._anchos, C._linea


def margen(c, izq=0.0, der=0.0, arr=0.0, aba=0.0):
    """Relleno interior de una celda (cm)."""
    tcPr = c._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
    for lado, v in (('top', arr), ('left', izq), ('bottom', aba), ('right', der)):
        e = OxmlElement('w:' + lado); e.set(qn('w:w'), str(int(v * 567))); e.set(qn('w:type'), 'dxa'); m.append(e)
    C._poner_en_orden(tcPr, m, C.ORDEN_TCPR)


def tabla(cont, ancho_cols, alto=None):
    tb = cont.add_table(rows=1, cols=len(ancho_cols)); sin_bordes(tb); tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    anchos(tb, ancho_cols)
    if alto:
        tb.rows[0].height = Cm(alto); tb.rows[0].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    return tb, tb.rows[0].cells


def borde_par(p, lados, color=G, tam=12, espacio=6):
    pPr = p._p.get_or_add_pPr(); bdr = OxmlElement('w:pBdr')
    for lado in ('top', 'left', 'bottom', 'right'):
        if lado in lados:
            b = OxmlElement('w:' + lado)
            for k, v in (('val', 'single'), ('sz', str(tam)), ('space', str(espacio)), ('color', color)):
                b.set(qn('w:' + k), v)
            bdr.append(b)
    pPr.append(bdr)


def minimo(cont):
    """Párrafo de 1 pt (cierre de una tabla sin empujar la página)."""
    p = cont.add_paragraph(); pf = p.paragraph_format; pf.space_before = pf.space_after = Pt(0); pf.line_spacing = Pt(1)
    rpr = OxmlElement('w:rPr'); sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '2'); rpr.append(sz)
    p._p.get_or_add_pPr().append(rpr)
    return p


def seccion(d, primera, m_lat, m_sup, m_inf):
    s = d.sections[0] if primera else d.add_section(WD_SECTION.NEW_PAGE)
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(m_lat); s.top_margin = Cm(m_sup); s.bottom_margin = Cm(m_inf)
    s.header_distance = s.footer_distance = Cm(0.5)
    for pg in s._sectPr.findall(qn('w:pgBorders')):                 # la sección nueva no hereda el marco de la anterior
        s._sectPr.remove(pg)
    return s


def borde_pagina(s, color=G):
    pg = OxmlElement('w:pgBorders'); pg.set(qn('w:offsetFrom'), 'page')
    for lado in ('top', 'left', 'bottom', 'right'):
        e = OxmlElement('w:' + lado)
        for k, v in (('val', 'thinThickSmallGap'), ('sz', '24'), ('space', '24'), ('color', color)):
            e.set(qn('w:' + k), v)
        pg.append(e)
    sp = s._sectPr
    ant = sp.find(qn('w:pgMar'))
    ant.addnext(pg)


def ubicacion(cont, color_etq, color_val, tam=9.5, centro=False, sangria=''):
    for etq, val in (('Distrito', 'Morales'), ('Provincia', 'San Martín'), ('Departamento', 'San Martín')):
        p = par(cont, sangria + '%s:  ' % etq, tam, False, color_etq, centro=centro, despues=1)
        r = p.add_run(val); r.font.size = Pt(tam); r.bold = True; r.font.name = 'Arial'
        if color_val:
            from docx.shared import RGBColor
            r.font.color.rgb = RGBColor.from_string(color_val)


# ------------------------------------------------------------------------------------------------- opciones
def opcion1(d, img, primera=False):
    """Banda lateral guinda a toda la altura con escudo, CUI, ubicación y fecha; contenido a la derecha."""
    seccion(d, primera, 1.0, 1.0, 0.8)
    tb, (c0, c1) = tabla(d, (5.6, 13.4), alto=27.7)
    sombra(c0, G); margen(c0, 0.45, 0.45, 0.9, 0.6); margen(c1, 0.8, 0.2, 0.9, 0.4)
    c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP; c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    t, (e,) = tabla(c0, (3.4,)); sombra(e, 'FFFFFF'); margen(e, 0.2, 0.2, 0.25, 0.25)
    par(e.paragraphs[0]).add_run().add_picture(V.ESCUDO, height=Cm(3.0))
    par(c0, C.ENTIDAD, 12, True, 'FFFFFF', antes=12, despues=3, sep=4, esp=1.1)
    par(c0, C.GERENCIA, 8, True, ROSA, despues=2, esp=1.1)
    par(c0, '', antes=150)
    p = par(c0, V.CUI, 15, True, 'FFFFFF', despues=4, sep=4); borde_par(p, ('top', 'bottom'), 'FFFFFF', 8, 6)
    par(c0, '', antes=150)
    par(c0, 'UBICACIÓN', 9.5, True, ROSA, centro=False, despues=3, sep=6)
    ubicacion(c0, 'FFFFFF', 'FFFFFF', 9)
    par(c0, 'FECHA', 9.5, True, ROSA, centro=False, antes=14, despues=3, sep=6)
    par(c0, V.FECHA, 13, True, 'FFFFFF', centro=False, sep=4)
    par(c1.paragraphs[0], SERVICIO, 10.5, True, G, centro=False, despues=6)
    par(c1, V.PROYECTO, 9.5, True, centro=False, despues=30, esp=1.1)
    par(c1).add_run().add_picture(img, width=Cm(12.4))
    par(c1, V.LEYENDA, 7.5, False, GRIS, cursiva=True, antes=2, despues=110)
    p = par(c1, TITULO, 34, True, G, centro=False, despues=4, sep=4, esp=0.95); linea(p, G, 24)
    par(c1, SUBTITULO, 12, True, GRIS, centro=False, antes=6, sep=6)
    minimo(d)


def opcion2(d, img, primera=False):
    """Franja guinda superior con la entidad, foto a todo el ancho, título entre filetes y franja inferior con los datos."""
    seccion(d, primera, 1.5, 1.2, 0.8)
    A = 18.0
    tb, (c0, c1) = tabla(d, (3.6, A - 3.6), alto=3.4)
    for c in (c0, c1):
        sombra(c, G); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margen(c0, 0.4, 0.1); margen(c1, 0.3, 0.4)
    t, (e,) = tabla(c0, (2.6,)); sombra(e, 'FFFFFF'); margen(e, 0.15, 0.15, 0.15, 0.15)
    par(e.paragraphs[0]).add_run().add_picture(V.ESCUDO, height=Cm(2.3))
    par(c1.paragraphs[0], C.ENTIDAD, 19, True, 'FFFFFF', centro=False, sep=10)
    par(c1, C.GERENCIA, 10, True, ROSA, centro=False, antes=2, sep=6)
    par(d, SERVICIO, 10.5, True, G, antes=16, despues=6)
    par(d, V.PROYECTO, 9.5, True, despues=4, esp=1.1)
    par(d, V.CUI, 12, True, G, despues=18, sep=4)
    tb, (c,) = tabla(d, (A,)); par(c.paragraphs[0]).add_run().add_picture(img, width=Cm(A))
    p = par(d, TITULO, 30, True, G, antes=44, despues=2, sep=6); borde_par(p, ('top',), G, 18, 10)
    p = par(d, SUBTITULO, 12, True, GRIS, despues=58, sep=8); borde_par(p, ('bottom',), G, 18, 10)
    tb, celdas = tabla(d, (6.6, 5.0, 6.4), alto=2.6)
    for k, c in enumerate(celdas):
        sombra(c, G if k != 1 else G_OSCURO); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; margen(c, 0.5, 0.3)
    par(celdas[0].paragraphs[0], 'UBICACIÓN DEL PROYECTO', 9, True, ROSA, centro=False, despues=2, sep=4)
    ubicacion(celdas[0], 'FFFFFF', 'FFFFFF', 9)
    par(celdas[1].paragraphs[0], 'ESPECIALIDAD', 9, True, ROSA, despues=2, sep=4)
    par(celdas[1], ESPEC_CORTA, 11, True, 'FFFFFF', sep=4)
    par(celdas[2].paragraphs[0], 'FECHA', 9, True, ROSA, despues=2, sep=4)
    par(celdas[2], V.FECHA, 15, True, 'FFFFFF', sep=6)
    minimo(d)


def opcion3(d, img, primera=False):
    """Marco doble guinda en la página, composición centrada y título en recuadro."""
    s = seccion(d, primera, 2.4, 2.2, 1.6); borde_pagina(s)
    par(d).add_run().add_picture(V.ESCUDO, height=Cm(3.2))
    par(d, C.ENTIDAD, 19, True, G, antes=6, sep=10)
    p = par(d, C.GERENCIA, 10, True, GRIS, antes=2, despues=4, sep=6); linea(p, G, 12)
    par(d, SERVICIO, 10.5, True, G, antes=14, despues=6)
    par(d, V.PROYECTO, 9.5, True, despues=4, esp=1.1)
    par(d, V.CUI, 12, True, G, despues=16, sep=4)
    tb, (c,) = tabla(d, (13.2,)); bordes(c, top=(12, G), left=(12, G), bottom=(12, G), right=(12, G)); margen(c, 0.2, 0.2, 0.2, 0.2)
    par(c.paragraphs[0]).add_run().add_picture(img, width=Cm(12.8))
    par(d, '', antes=22)
    tb, (c,) = tabla(d, (13.2,)); bordes(c, top=(24, G), left=(24, G), bottom=(24, G), right=(24, G)); sombra(c, G_CLARO)
    par(c.paragraphs[0], TITULO, 26, True, G, antes=12, despues=2, sep=6)
    par(c, SUBTITULO, 11, True, GRIS, despues=12, sep=6)
    par(d, 'UBICACIÓN DEL PROYECTO', 9.5, True, G, antes=26, despues=2, sep=6)
    par(d, 'Distrito de Morales  ·  Provincia de San Martín  ·  Departamento de San Martín', 9.5, False, GRIS, despues=12)
    par(d, V.FECHA, 14, True, G, sep=8)


def opcion4(d, img, primera=False):
    """Bloques: título grande con barra guinda, foto, recuadro claro con el proyecto y fila de datos guinda."""
    seccion(d, primera, 2.0, 1.6, 1.0)
    A = 17.0
    tb, (c0, c1, c2) = tabla(d, (2.4, 9.6, 5.0))
    for c in (c0, c1, c2):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    par(c0.paragraphs[0], centro=False).add_run().add_picture(V.ESCUDO, height=Cm(2.0))
    par(c1.paragraphs[0], C.ENTIDAD, 13, True, G, centro=False, sep=6)
    par(c1, C.GERENCIA, 8.5, True, GRIS, centro=False, antes=1, sep=4)
    sombra(c2, G); margen(c2, 0.2, 0.2, 0.25, 0.25)
    par(c2.paragraphs[0], V.CUI, 12.5, True, 'FFFFFF', sep=4)
    tb, (c,) = tabla(d, (A,)); bordes(c, left=(72, G)); margen(c, 0.6)
    palabras = TITULO.split(' ', 1)
    par(c.paragraphs[0], palabras[0], 38, True, G, centro=False, antes=26, esp=0.9, sep=6)
    par(c, palabras[1] if len(palabras) > 1 else '', 38, True, '3A3A3A', centro=False, esp=0.9, sep=6)
    par(c, SUBTITULO, 12, True, GRIS, centro=False, antes=6, despues=4, sep=8)
    tb, (c,) = tabla(d, (A,)); par(c.paragraphs[0], antes=18).add_run().add_picture(img, width=Cm(A))
    tb, (c,) = tabla(d, (A,)); sombra(c, G_CLARO); margen(c, 0.5, 0.5, 0.3, 0.3); bordes(c, top=(24, G))
    par(c.paragraphs[0], SERVICIO, 10, True, G, centro=False, antes=4, despues=4)
    par(c, V.PROYECTO, 9.5, True, centro=False, despues=4, esp=1.1)
    par(d, '', antes=12)
    tb, celdas = tabla(d, (5.6, 5.8, 5.6), alto=1.9)
    for k, (etq, val) in enumerate((('DISTRITO', 'Morales'), ('PROVINCIA / DEPARTAMENTO', 'San Martín'), ('FECHA', V.FECHA))):
        c = celdas[k]; sombra(c, G); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        bordes(c, right=(36, 'FFFFFF')) if k < 2 else None
        par(c.paragraphs[0], etq, 8.5, True, ROSA, despues=2, sep=4)
        par(c, val, 13, True, 'FFFFFF', sep=4)
    minimo(d)


OPCIONES = [opcion1, opcion2, opcion3, opcion4]


def documento(img, mosaico=None):
    d = docx.Document()
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial'); st.paragraph_format.space_after = Pt(0)
    if d.paragraphs:
        d.element.body.remove(d.paragraphs[0]._p)
    if mosaico:                                                     # página de vista previa
        seccion(d, True, 1.5, 1.5, 1.2)
        par(d, 'OPCIONES DE CARÁTULA - COLOR GUINDA', 16, True, G, despues=2, sep=6)
        p = par(d, NOMBRE + ' - ' + V.CUI + '  |  vista previa de las cuatro opciones (cada una, a tamaño real, en las páginas '
                'siguientes)', 9, False, GRIS, despues=10); linea(p, G, 12)
        par(d).add_run().add_picture(mosaico, width=Cm(16.0))
    for k, f in enumerate(OPCIONES):
        f(d, img, primera=not mosaico and k == 0)
    d.core_properties.author = ''; d.core_properties.title = 'Opciones de carátula - ' + NOMBRE
    return d


def vistas(docx_path):
    """PDF temporal -> PNG de cada opción y mosaico 2 x 2 con su rótulo."""
    from PIL import Image, ImageDraw, ImageFont
    fig = V.FIG; os.makedirs(fig, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, docx_path],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=True)
        pdf = os.path.join(tmp, os.path.basename(docx_path)[:-5] + '.pdf')
        subprocess.run(['pdftoppm', '-png', '-r', '80', pdf, os.path.join(tmp, 'p')], check=True)
        pags = sorted(f for f in os.listdir(tmp) if f.startswith('p') and f.endswith('.png'))
        if len(pags) != len(OPCIONES):
            raise ValueError('se esperaban %d páginas y hay %d' % (len(OPCIONES), len(pags)))
        ims = []
        for k, f in enumerate(pags):
            im = Image.open(os.path.join(tmp, f)).convert('RGB')
            ruta = os.path.join(fig, 'opcion_caratula_%d.png' % (k + 1)); im.save(ruta); ims.append(im)
    w, h = ims[0].size; esc = 0.5; tw, th = int(w * esc), int(h * esc); pad, cab = 30, 46
    m = Image.new('RGB', (2 * tw + 3 * pad, 2 * (th + cab) + 3 * pad), 'white'); dr = ImageDraw.Draw(m)
    try:
        fnt = ImageFont.truetype('/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf', 22)
    except OSError:
        fnt = ImageFont.load_default()
    for k, im in enumerate(ims):
        x = pad + (k % 2) * (tw + pad); y = pad + (k // 2) * (th + cab + pad)
        dr.text((x, y), NOMBRES[k], fill='#7B1E2B', font=fnt)
        m.paste(im.resize((tw, th), Image.LANCZOS), (x, y + cab))
        dr.rectangle([x - 1, y + cab - 1, x + tw, y + cab + th], outline='#BBBBBB', width=2)
    ruta = os.path.join(fig, 'opciones_caratula_mosaico.png'); m.save(ruta)
    return ruta


def generar():
    img = V.foto(); V.escudo()
    with tempfile.TemporaryDirectory() as tmp:
        prueba = os.path.join(tmp, 'opciones.docx'); documento(img).save(prueba)
        mos = vistas(prueba)
    d = documento(img, mos); d.save(SALIDA)
    return SALIDA, mos


if __name__ == '__main__':
    print(*generar(), sep='\n')
