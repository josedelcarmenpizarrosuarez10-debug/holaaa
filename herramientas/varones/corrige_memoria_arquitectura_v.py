# -*- coding: utf-8 -*-
"""Corrección de la memoria descriptiva de arquitectura y de la programación arquitectónica - CAR Varones (CUI 2705619).

Parte de los archivos de la arquitecta (insumos/katiuska/arquitectura/, no se modifican) y escribe copias corregidas con el mismo
formato e imágenes. Solo se cambia el texto con error; todo lo demás queda igual.

Memoria descriptiva:
  - cuadro de incompatibilidades: D.L. N.° 25844, Ley Orgánica de Municipalidades, fila 13 sin texto de colegio, condiciones
    concretas del CAR en lugar de "área definida para la UP", aeródromo con el aeropuerto de Tarapoto, fajas marginales coherentes
    con las aguas pluviales que menciona la memoria; numerado como Cuadro N.° 4 (los N.° 1 a 3 son de topografía);
  - acceso (carretera F.B.T. y carretera Oasis), titular del terreno (Aldea Infantil -> Gobierno Regional), pendiente "1. %",
    "TOLERANCIAS CATAS", imágenes numeradas del 1 al 9, "R.V.N", "SEMAMHI", "el Locales", "infraestructura educativa", norma de
    confort (A.030 -> EM.110) y "aulas", drenaje pluvial del proyecto (no "cunetas existentes"), A.090 en la normativa;
  - programa: fila de subtotal mal rotulada ("OBRAS EXTERIORES") y total que no suma (4,205.30): el cerco se mide en ml y no suma área;
    se agrega el estacionamiento que la fórmula del Excel omitía;
  - listado de planos "(POR DEFINIR)" reemplazado por las 89 láminas del índice de la arquitecta (indice_arquitectura_v.py);
  - análisis de ubicación dentro del numeral 2.3: 2.3.1 (método + Cuadro N.° 4) y 2.3.2 (Cuadro N.° 5 con las condiciones de
    ubicación de incompatibilidades_v.py) con su conclusión;
  - índice (campo TOC de Word, se actualiza al abrir) con el número de página de cada título.
Programación (Excel): mismas correcciones del subtotal y del total, #REF! eliminado, unidad de Nutrición y erratas.

Uso: python3 herramientas/varones/corrige_memoria_arquitectura_v.py
Salida: entregables_varones/arquitectura/MEMORIA_DESCRIPTIVA_ARQUITECTURA_CAR_VARONES_CORREGIDA.docx
        entregables_varones/arquitectura/PROGRAMACION_ARQUITECTONICA_CAR_VARONES_CORREGIDA.xlsx
"""
import copy
import os
import subprocess
import sys
import tempfile

import docx
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, AQUI)
import indice_arquitectura_v as IA  # noqa: E402  (láminas del índice de la arquitecta)

ENT = os.path.join(RAIZ, 'insumos/katiuska/arquitectura')
MEM = os.path.join(ENT, 'MEMORIA_DESCRIPTIVA_ARQUITECTURA_CAR_VARONES.docx')
PRG = os.path.join(ENT, 'PROGRAMACION_ARQUITECTONICA_CAR_VARONES.xlsx')
SAL = os.path.join(RAIZ, 'entregables_varones/arquitectura')
MEM_C = os.path.join(SAL, 'MEMORIA_DESCRIPTIVA_ARQUITECTURA_CAR_VARONES_CORREGIDA.docx')
PRG_C = os.path.join(SAL, 'PROGRAMACION_ARQUITECTONICA_CAR_VARONES_CORREGIDA.xlsx')
CERCO = 101.95 + 24.00                               # cerco perimétrico en ml (no suma área)
CAMBIOS = []


def poner(runs, textos):
    """Escribe los textos en los primeros runs y vacía los demás (conserva el formato de cada run)."""
    for k, r in enumerate(runs):
        r.text = textos[k] if k < len(textos) else ''


def reemplazar(p, viejo, nuevo):
    """Reemplaza un texto dentro de un párrafo aunque esté repartido en varios runs (el resultado queda en el primer run)."""
    if viejo not in p.text:
        raise ValueError('no está: ' + viejo)
    t = p.text.replace(viejo, nuevo)
    poner(p.runs, [t]); CAMBIOS.append((viejo, nuevo))


def celda_estado(c, texto):
    """Celda de ESTADO: deja "CUMPLE" en su run y pone la condición en el run siguiente."""
    runs = [r for q in c.paragraphs for r in q.runs]
    mismo = runs[0]._r.getparent() is runs[1]._r.getparent()
    poner(runs, ['CUMPLE', (', ' + texto[0].lower() + texto[1:]) if mismo else texto])


def celda(c, texto):
    runs = [r for q in c.paragraphs for r in q.runs]
    poner(runs, [texto])


ESTADOS = {   # fila del cuadro (N.°) -> condición del CAR
    4: 'No existen estaciones de servicio, grifos ni gasocentros a menos de 50 m del CAR.',
    5: 'No existen locales de venta y consumo de bebidas alcohólicas a menos de 100 m del CAR.',
    7: 'No existen fuentes de agua naturales ni artificiales con faja marginal; las aguas pluviales que discurren por la margen '
       'derecha del terreno se captan con el sistema de drenaje pluvial del proyecto.',
    10: 'El aeropuerto de Tarapoto está a unos 3.5 km, al costado de la pista 17/35 y fuera de las superficies de aproximación y '
        'despegue; las edificaciones del CAR son de baja altura.',
    14: 'No existen plantas de tratamiento de aguas residuales de terceros en el entorno del CAR.',
    15: 'El área del CAR no se encuentra en un área restringida de construcción; su frente da a la carretera Oasis a través del área '
        'libre del predio matriz.',
    17: 'No existen casinos ni salas de máquinas tragamonedas en el entorno del CAR.',
    18: 'No existen hostales, peñas, discotecas, video-pubs, bingos ni salas de billar en el entorno del CAR.',
}


def corregir_memoria():
    d = docx.Document(MEM)
    P = d.paragraphs
    buscar = lambda txt: next(p for p in P if txt in p.text)        # noqa: E731

    reemplazar(buscar('se ingresa por la carretera F.B.T'),
               'se ingresa por la carretera F.B.T, altura del sector el Tumi, margen derecha aproximadamente 340m por una trocha '
               'carrozable que une con la vía de evitamiento.',
               'se ingresa por la carretera Fernando Belaunde Terry (F.B.T.), a la altura del sector El Tumi, margen derecha, y luego '
               'aproximadamente 340 m por la carretera Oasis, vía carrozable que une con la vía de evitamiento.')
    reemplazar(buscar('pertenece a la Aldea Infantil'), 'es un área que pertenece a la Aldea Infantil Virgen Del Pilar.',
               'es un área que perteneció a la Aldea Infantil Virgen del Pilar y que desde el año 2025 es de propiedad del Gobierno '
               'Regional de San Martín (partida registral N.° 05003469).')
    reemplazar(buscar('pendiente longitudinal de'), '1. % y una pendiente transversal de 0.20%.', '1 % y una pendiente transversal de 0.20 %.')
    reemplazar(buscar('por la margen derecha, discurren'),
               'La zona de estudio por la margen derecha, discurren aguas de drenaje pluvial, el cual deberá tomar en cuenta para el '
               'diseño de distribución arquitectónico.',
               'Por la margen derecha de la zona de estudio discurren aguas de drenaje pluvial, que se tomaron en cuenta en la '
               'distribución arquitectónica y que se captan y conducen con el sistema de drenaje pluvial del proyecto.')
    reemplazar(buscar('CUADRO DE TOLERANCIAS CATAS'), 'TOLERANCIAS CATAS', 'TOLERANCIAS CATASTRALES')
    # imágenes numeradas del 1 al 9 (había dos "Imagen 4" y dos "Imagen 6"); solo se cambia el run del número
    leyendas = [(buscar('Área de Influencia.'), 'N°4:', '4:'), (buscar('Zonas Climáticas del Perú'), 'Imagen 4', 'Imagen 5')]
    leyendas += [(p, 'Imagen %d' % a, 'Imagen %d' % b) for p, a, b in
                 zip([p for p in P if p.text.strip() in ('Imagen 5: Área de proyecto - Estudio Topográfico.',
                                                          'Imagen 6: Área de proyecto - Estudio Topográfico.',
                                                          'Imagen 6: Superposición de polígonos.', 'Imagen 7: Superposición de polígonos.')],
                     (5, 6, 6, 7), (6, 7, 8, 9))]
    for p, viejo, nuevo_ in leyendas:
        r = next(r for r in p.runs if viejo in r.text)
        r.text = r.text.replace(viejo, nuevo_); CAMBIOS.append((p.text.strip(), nuevo_))
    reemplazar(buscar('confort bioclimático de la infraestructura educativa'), 'de la infraestructura educativa', 'de la infraestructura')
    reemplazar(buscar('Bioclimática el Locales'), 'Bioclimática el Locales educativos', 'Bioclimática en Locales Educativos')
    reemplazar(buscar('a las cunetas existentes en el exterior'), 'hasta ser trasladados finalmente a las cunetas existentes en el '
               'exterior', 'hasta ser conducidas a las cunetas y al colector pluvial frontal del proyecto')
    p = buscar('debe responder al RNE - Norma A.030')
    reemplazar(p, 'RNE - Norma A.030', 'RNE - Norma EM.110 Confort térmico y lumínico con eficiencia energética')
    reemplazar(p, 'dentro de las aulas y demás ambientes', 'dentro de los dormitorios y demás ambientes')
    for p in P:
        if 'SEMAMHI' in p.text:
            reemplazar(p, 'SEMAMHI', 'SENAMHI')
        if 'R.V.N' in p.text:
            reemplazar(p, 'R.V.N', 'R.V.M.')
    # normativa: A.090 (ubicación de servicios comunales), después de la A.080
    a080 = buscar('A.080 - OFICINAS')
    nuevo = copy.deepcopy(a080._p); a080._p.addnext(nuevo)
    q = docx.text.paragraph.Paragraph(nuevo, a080._parent); poner(q.runs, ['A.090 - Servicios comunales.'])
    CAMBIOS.append(('(normativa)', 'A.090 - Servicios comunales'))

    # cuadro de incompatibilidades
    t = next(t for t in d.tables if 'incompatibilidades de uso de terreno' in t.rows[0].cells[0].text)
    cab = [r for q in t.rows[0].cells[0].paragraphs for r in q.runs]
    for r in cab:
        if r.text == ' 1:':
            r.text = ' 4:'
    for fila in t.rows[2:]:
        n = int(fila.cells[0].text.strip())
        if n in ESTADOS:
            celda_estado(fila.cells[3], ESTADOS[n])
        if n == 12:
            celda(fila.cells[2], 'Decreto Ley Nº 25844 y sus modificatorias. Ley de Concesiones Eléctricas.')
        if n == 13:
            celda(fila.cells[1], 'En relación a restricciones radioeléctricas en áreas de uso público próximas a una estación '
                                 'radioeléctrica.')
        if n == 18:
            celda(fila.cells[2], 'Artículo 79° de la Ley N° 27972 – Ley Orgánica de Municipalidades, en materia de organización del '
                                 'espacio físico y uso del suelo.')
    CAMBIOS.append(('Cuadro Nº 1 de incompatibilidades', 'Cuadro N° 4, filas 4, 5, 7, 10, 12, 13, 14, 15, 17 y 18'))

    # programa: subtotal mal rotulado y total
    t = next(t for t in d.tables if t.rows[0].cells[0].text.strip() == 'TIPO' and len(t.rows) > 100)
    sub = next(f for f in t.rows if f.cells[0].text.strip() == 'OBRAS EXTERIORES' and f.cells[-1].text.strip() == '1229.76')
    celda(sub.cells[1], 'SUB TOTAL (sin el cerco perimétrico, que se mide en ml)')
    celda(sub.cells[-1], '%.2f' % (1229.76 - CERCO))
    tot = next(f for f in t.rows if 'ÁREA TOTAL' in f.cells[0].text)
    total = 0.0
    for f in t.rows:
        lab = ' '.join(c.text for c in f.cells[:2])
        v = f.cells[-1].text.replace(',', '').strip()
        if 'SUB TOTAL' in lab or 'ESTACIONAMIENTO' in lab:
            total += float(v)
    celda(tot.cells[-1], '{:,.2f}'.format(total))
    CAMBIOS.append(('subtotal 1229.76 / total 4,205.30', '%.2f / %s' % (1229.76 - CERCO, '{:,.2f}'.format(total))))

    # listado de planos: las 89 láminas del índice de la arquitecta
    tit = buscar('LISTADO DE PLANOS')
    reemplazar(tit, 'LISTADO DE PLANOS (POR DEFINIR)', 'LISTADO DE PLANOS')
    viejos = [p for p in P[P.index(tit) + 1:] if p.style.name == 'List Paragraph']
    grupo_t = next(p for p in viejos if p.text.strip() == 'EDIFICACIONES')._p
    item_t = next(p for p in viejos if p.text.startswith('PU-01'))._p
    ancla = viejos[0]._p.getprevious()
    tomos, _ = IA.leer()
    for g, filas in IA._planos(tomos):
        e = copy.deepcopy(grupo_t); ancla.addnext(e); ancla = e
        poner(docx.text.paragraph.Paragraph(e, tit._parent).runs, ['%s (%d)' % (g, len(filas))])
        for f in filas:
            e = copy.deepcopy(item_t); ancla.addnext(e); ancla = e
            poner(docx.text.paragraph.Paragraph(e, tit._parent).runs, [f['codigo'], ' ' + f['descripcion']])
    for p in viejos:
        p._p.getparent().remove(p._p)
    CAMBIOS.append(('listado de planos (por definir)', '89 láminas del índice de arquitectura'))
    reemplazar(buscar('COLINDANCIA E INCOMPATILIDAD'), 'INCOMPATILIDAD', 'INCOMPATIBILIDAD')
    agregar_analisis(d)
    preparar_indice(d)
    d.core_properties.author = ''
    os.makedirs(SAL, exist_ok=True)
    d.save(MEM_C)
    paginar_indice()
    return MEM_C, total


# --------------------------------------------------------------------------- análisis de ubicación dentro del numeral 2.3
METODO = ('El análisis verifica que el terreno no presenta incompatibilidades de uso por cercanía a actividades, instalaciones o '
          'servidumbres reguladas por la normativa vigente (RNE, Norma A.040, art. 7 b; R.V.M. N° 208-2019-MINEDU, art. 8 c). Para '
          'cada dispositivo legal se define el área de influencia como la franja medida desde el lindero con la distancia que fija: '
          '50 m (estaciones de servicio), 100 m (establecimientos de salud, GLP, bebidas alcohólicas y plantas de combustibles), '
          '150 m (velatorios y cementerios) y 200 m (ductos de hidrocarburos); para los que no fijan distancia se verifica que la '
          'restricción no alcance al terreno. Se reconoció el entorno y se revisó la cartografía y las imágenes satelitales de la zona.')
CONDIC_INTRO = ('Además de las incompatibilidades por cercanía, el terreno y el proyecto cumplen las condiciones de ubicación que el '
                'Reglamento Nacional de Edificaciones y la R.V.M. N° 208-2019-MINEDU establecen para el Centro de Acogida Residencial:')
CONCLUSION = ('Conclusión: el terreno cumple las 18 condiciones del Cuadro N° 4 y las %d del Cuadro N° 5; no existen dentro de su área '
              'de influencia actividades, instalaciones ni servidumbres incompatibles con la permanencia de niños y adolescentes, y su '
              'ubicación cumple el RNE (A.090, arts. 3 y 4; A.040, art. 7) y la R.V.M. N° 208-2019-MINEDU (art. 8). Por lo tanto, el '
              'terreno es apto para el Centro de Acogida Residencial – Varones.')


def _copia(modelo, despues, texto):
    """Párrafo nuevo con el formato del modelo, insertado después de 'despues' (elemento XML)."""
    e = copy.deepcopy(modelo._p); despues.addnext(e)
    q = docx.text.paragraph.Paragraph(e, modelo._parent); poner(q.runs, [texto] if q.runs else [])
    if not q.runs:
        q.add_run(texto)
    return e


def agregar_analisis(d):
    import incompatibilidades_v as IV                      # condiciones de ubicación del CAR (Cuadro N.° 2 del análisis)
    P = d.paragraphs
    titulo = next(p for p in P if p.text.strip() == 'ANÁLISIS DE INCOMPATIBILIDADES DE UBICACIÓN EN EL PROYECTO')
    sub = next(p for p in P if p.text.strip() == 'TOLERANCIAS CATASTRALES')         # título numerado de tercer nivel (2.x.y)
    cuerpo = next(p for p in P if p.text.startswith('Dentro del radio de 100 metros'))
    # 2.3.1: el título del análisis pasa a ser numerado, con el método debajo (antes del Cuadro N.° 4)
    e = _copia(sub, titulo._p, 'ANÁLISIS DE INCOMPATIBILIDADES DE UBICACIÓN EN EL PROYECTO')
    titulo._p.getparent().remove(titulo._p)
    m = _copia(cuerpo, e, '')
    m = _copia(cuerpo, m, METODO)
    _copia(cuerpo, m, '')
    # después del Cuadro N.° 4: 2.3.2 condiciones de ubicación y Cuadro N.° 5
    t4 = next(t for t in d.tables if 'incompatibilidades de uso de terreno' in t.rows[0].cells[0].text)
    ancla = t4._tbl
    for txt, mod in (('', cuerpo), ('CONDICIONES DE UBICACIÓN DEL CENTRO DE ACOGIDA RESIDENCIAL', sub), ('', cuerpo),
                     (CONDIC_INTRO, cuerpo), ('', cuerpo)):
        ancla = _copia(mod, ancla, txt)
    t5 = copy.deepcopy(t4._tbl); ancla.addnext(t5)
    T = docx.table.Table(t5, t4._parent)
    for r in T.rows[2 + len(IV.CONDICIONES):]:
        t5.remove(r._tr)
    cab = [r for q in T.rows[0].cells[0].paragraphs for r in q.runs]
    poner(cab, ['Cuadro ', 'N°', ' 5:', ' Condiciones de ubicación del Centro de Acogida Residencial - Varones, según el Reglamento '
                'Nacional de Edificaciones y la R.V.M. ', 'N°', ' 208-2019-MINEDU.'])
    celda(T.rows[1].cells[0], 'CONDICIÓN DE UBICACIÓN DEL CAR')
    celda(T.rows[1].cells[2], 'FUENTE NORMATIVA')
    for k, (cond, fuente, como) in enumerate(IV.CONDICIONES):
        f = T.rows[2 + k]
        celda(f.cells[0], str(k + 1)); celda(f.cells[1], cond); celda(f.cells[2], fuente)
        celda_estado(f.cells[3], como)
    for f in T.rows[:2]:                                     # el título del cuadro no se separa de la tabla
        for c in f.cells:
            for q in c.paragraphs:
                q.paragraph_format.keep_with_next = True
    ancla = _copia(cuerpo, t5, '')
    _copia(cuerpo, ancla, CONCLUSION % len(IV.CONDICIONES))
    CAMBIOS.append(('2.3 Colindancia', '2.3.1 análisis (método + Cuadro N° 4), 2.3.2 condiciones (Cuadro N° 5) y conclusión'))


# --------------------------------------------------------------------------- índice (tabla de contenido)
def _titulos(d):
    """(párrafo, nivel) de los títulos numerados de la memoria y del listado de planos."""
    out = []
    for p in d.paragraphs:
        pPr = p._p.pPr; t = p.text.strip()
        if not t:
            continue
        num = pPr.numPr if pPr is not None else None
        if num is not None and num.numId is not None and num.numId.val in (1, 2, 8):
            lv = num.ilvl.val if num.ilvl is not None else 0
            if lv <= 2 and (num.numId.val != 1 or lv == 1):
                out.append((p, lv))
        elif t == 'LISTADO DE PLANOS':
            out.append((p, 0))
    return out


def _numeros(tit):
    c = [0, 0, 0]; res = []
    for p, lv in tit:
        if p.text.strip() == 'LISTADO DE PLANOS':
            res.append(''); continue
        c[lv] += 1
        for k in range(lv + 1, 3):
            c[k] = 0
        res.append('%d.0.' % c[0] if lv == 0 else '.'.join(str(x) for x in c[:lv + 1]) + '.')
    return res


def preparar_indice(d):
    """Índice con campo TOC de Word (se actualiza al abrir) y entradas ya escritas; las páginas se llenan en paginar_indice."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    tit = _titulos(d)
    for p, lv in tit:                                          # nivel de esquema para el campo TOC \u
        o = p._p.get_or_add_pPr().find(qn('w:outlineLvl'))
        if o is None:
            o = OxmlElement('w:outlineLvl'); p._p.get_or_add_pPr().append(o)
        o.set(qn('w:val'), str(lv))
    s = d.sections[0]; ancho = int((s.page_width - s.left_margin - s.right_margin) / 635)
    gen = tit[0][0]
    gen._p.get_or_add_pPr().insert(0, OxmlElement('w:pageBreakBefore'))
    cab = d.add_paragraph(); gen._p.addprevious(cab._p)
    cab.alignment = 1; cab.paragraph_format.page_break_before = True; cab.paragraph_format.space_after = Pt(12)
    r = cab.add_run('ÍNDICE'); r.bold = True; r.font.size = Pt(14); r.font.name = 'Arial Narrow'
    nums = _numeros(tit)
    for k, ((p, lv), n) in enumerate(zip(tit, nums)):
        q = d.add_paragraph(); gen._p.addprevious(q._p)
        pf = q.paragraph_format; pf.left_indent = Pt(14 * lv); pf.space_after = Pt(2); pf.space_before = Pt(6 if lv == 0 else 0)
        tabs = OxmlElement('w:tabs'); tb = OxmlElement('w:tab')
        tb.set(qn('w:val'), 'right'); tb.set(qn('w:leader'), 'dot'); tb.set(qn('w:pos'), str(ancho)); tabs.append(tb)
        q._p.get_or_add_pPr().append(tabs)
        if k == 0:
            for tipo, txt in (('begin', None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ('separate', None)):
                rr = q.add_run()
                if tipo:
                    fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); rr._r.append(fc)
                else:
                    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = txt; rr._r.append(it)
        texto = (n + ' ' if n else '') + ' '.join(p.text.split()).rstrip(':')
        r1 = q.add_run(texto); r1.bold = lv == 0; r1.font.size = Pt(11); r1.font.name = 'Arial Narrow'
        r2 = q.add_run('\t#'); r2.font.size = Pt(11); r2.font.name = 'Arial Narrow'
        if k == len(tit) - 1:
            rr = q.add_run(); fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), 'end'); rr._r.append(fc)
    upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'), 'true'); d.settings.element.append(upd)
    CAMBIOS.append(('(sin índice)', 'índice de %d títulos con número de página' % len(tit)))


def paginar_indice():
    """Escribe el número de página de cada entrada del índice según la paginación del documento (LibreOffice)."""
    import re
    if 'Arial Narrow' not in subprocess.run(['fc-match', 'Arial Narrow'], capture_output=True, text=True).stdout:
        subprocess.run([sys.executable, os.path.join(AQUI, 'fuente_paginacion.py')], check=False)   # paginar como Word
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, MEM_C],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        pdf = os.path.join(tmp, os.path.basename(MEM_C)[:-5] + '.pdf')
        txt = subprocess.run(['pdftotext', '-layout', pdf, '-'], capture_output=True, text=True).stdout.split('\f')
    norm = lambda x: re.sub(r'\s+', ' ', x).strip().upper()          # noqa: E731
    paginas = [norm(t) for t in txt]
    d = docx.Document(MEM_C)
    tit = _titulos(d)
    entradas = [p for p in d.paragraphs if p.text.endswith('\t#')]
    ini = next(i for i, t in enumerate(paginas) if 'ÍNDICE' in t) + 1
    while ini < len(paginas) and '#' in paginas[ini] and '1.0. GENERALIDADES' not in paginas[ini].split('#')[-1]:
        ini += 1
    pag = ini
    for (p, lv), q in zip(tit, entradas):
        clave = norm(p.text).rstrip(':')[:45]
        while pag < len(paginas) and clave not in paginas[pag]:
            pag += 1
        if pag >= len(paginas):
            raise ValueError('no se encontró en el PDF: ' + clave)
        q.runs[-2 if q.runs[-1].text == '' else -1].text = '\t%d' % (pag + 1)
    d.save(MEM_C)


def corregir_programacion():
    wb = openpyxl.load_workbook(PRG); ws = wb.active
    for fila in ws.iter_rows():
        for c in fila:
            if isinstance(c.value, str):
                v = (c.value.replace('Albaliñería', 'Albañilería').replace('AREA DE SEMRIOS', 'ÁREA DE SEMBRÍOS')
                     .replace('DEPOSTIO', 'DEPÓSITO'))
                if v.startswith('=') and '#REF!' in v:
                    v = None
                c.value = v
    nut = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or '').strip().upper().startswith('NUTRICI'))
    if ws.cell(nut, 6).value == 'm3':
        ws.cell(nut, 6).value = 'm2'
    ws['B103'] = 'SUB TOTAL (sin el cerco perimétrico, que se mide en ml)'
    ws['I103'] = '=I98+I99+I102'
    ws['I125'] = '=I12+I19+I36+I53+I70+I77+I88+I92+I97+I103+I111+I116+I123+I124'
    wb.save(PRG_C)
    with tempfile.TemporaryDirectory() as tmp:                        # recalcular con LibreOffice
        subprocess.run(['soffice', '--headless', '--convert-to', 'xlsx', '--outdir', tmp, PRG_C],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
        out = os.path.join(tmp, os.path.basename(PRG_C))
        if os.path.exists(out):
            os.replace(out, PRG_C)
    return PRG_C, openpyxl.load_workbook(PRG_C, data_only=True).active['I125'].value


if __name__ == '__main__':
    m, total = corregir_memoria()
    x, total_x = corregir_programacion()
    for a, b in CAMBIOS:
        print('  %s -> %s' % (a[:70], b[:70]))
    print(m, 'total %.2f' % total); print(x, 'total', total_x)
