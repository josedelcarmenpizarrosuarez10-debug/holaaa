# -*- coding: utf-8 -*-
"""Memoria descriptiva de evacuación y señalización - CAR Varones (CUI 2705619), en Word.

Redactada con los datos del plano final de la arquitecta (láminas SE-S-01 y SE-E-01, leídos por lee_plano_se.py) y con las
mismas cantidades del metrado (datos_se.py). Incorpora las correcciones de la memoria (M-01 a M-06): población de 22
personas de personal y capacidad máxima de 66, códigos de lámina SE-S-01 y SE-E-01, simbología igual a la leyenda del plano,
obras exteriores O.E. 01 a O.E. 07, título de la NTP 350.043-1 y redacción propia de un centro de acogida residencial.
Formato: el de las memorias descriptivas del expediente (memoria_descriptiva_refugio.Memoria) con la carátula del CAR Varones.

Uso: python3 herramientas/varones/senalizacion/memoria_se.py
Salida: entregables_varones/MEMORIA_DESCRIPTIVA_EVACUACION_Y_SENALIZACION_CAR_VARONES.docx
"""
import os
import sys

from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

AQUI = os.path.dirname(os.path.abspath(__file__)); VAR = os.path.dirname(AQUI); HER = os.path.dirname(VAR)
for d in (AQUI, VAR, HER):
    sys.path.insert(0, d)
import datos_se as D  # noqa: E402
import memoria_descriptiva_refugio as R  # noqa: E402
import caratulas_v as V  # noqa: E402

SALIDA = os.path.join(D.SAL, 'MEMORIA_DESCRIPTIVA_EVACUACION_Y_SENALIZACION_CAR_VARONES.docx')
FIG = os.path.join(D.SAL, 'figuras_senalizacion')          # vistas de las láminas (vistas_se.py)

# zonificación del plano (rótulos EDIFICACION, O.COMPLEMENTARIA y O.EXTERIOR de la planta general)
OBRAS = [('O.C. 01', 'Obra complementaria', 'Patio principal'), ('O.C. 02', 'Obra complementaria', 'Losa deportiva'),
         ('O.C. 03', 'Obra complementaria', 'Patio de maniobras'), ('O.C. 04', 'Obra complementaria', 'Patio de juegos'),
         ('O.E. 01', 'Obra exterior', 'Cerco perimétrico con ingreso principal e ingreso al patio de maniobras'),
         ('O.E. 02', 'Obra exterior', 'Tanque elevado y cisterna'), ('O.E. 03', 'Obra exterior', 'Ingreso principal'),
         ('O.E. 04', 'Obra exterior', 'Caseta de control'), ('O.E. 05', 'Obra exterior', 'Zona de pozo séptico'),
         ('O.E. 06', 'Obra exterior', 'Área de sembríos'), ('O.E. 07', 'Obra exterior', 'Veredas y circulación posterior de los módulos')]


def encabezado(M):
    sec = M.d.add_section(WD_SECTION.NEW_PAGE)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3.0), Cm(2.0), Cm(2.5), Cm(2.0)
    sec.header.is_linked_to_previous = False; sec.footer.is_linked_to_previous = False
    pg = OxmlElement('w:pgNumType'); pg.set(qn('w:start'), '1'); sec._sectPr.append(pg)
    h = sec.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    M.run(h, 'MEMORIA DESCRIPTIVA - EVACUACIÓN Y SEÑALIZACIÓN', 8.5, True, color=R.GRANATE)
    M.run(h, '   |   %s - CUI N.° %s' % (D.CORTO, D.CUI), 8.5, color=R.GRANATE)
    b = OxmlElement('w:pBdr'); e = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '2'), ('color', R.GRANATE)):
        e.set(qn('w:' + k), v)
    b.append(e); h._p.get_or_add_pPr().insert(1, b)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    M.run(f, 'Página ', 8.5, color=R.GRIS)
    for tipo in ('begin', 'instr', 'separate', 'txt', 'end'):
        r = M.run(f, '', 8.5, color=R.GRIS)
        if tipo == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'PAGE'; r._r.append(it)
        elif tipo == 'txt':
            r.text = '1'
        else:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), tipo); r._r.append(fc)


def figura(M, cod, titulo, ancho=16.0):
    ruta = os.path.join(FIG, cod + '.png')
    if not os.path.exists(ruta):
        return
    M.nf += 1; M.leyenda('Figura N.° %02d: %s (lámina %s)' % (M.nf, titulo, cod))
    p = M.par(centro=True, despues=2); p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(ruta, width=Cm(ancho))
    M.fuente('Fuente: Lámina %s del expediente de arquitectura (evacuación y señalización).' % cod)


def generar():
    P = D.plano(); pts = D.partidas(P); Pb = D.POBLACION; rutas = D.rutas(P); zs = D.zonas(P); z = D.zona_segura()
    q = {p['clave']: p['total'] for p in pts if p['und']}
    eds = D.edificaciones(P)
    M = R.Memoria()
    s0 = M.d.sections[0]; s0.left_margin = s0.right_margin = Cm(2.0); s0.top_margin = Cm(2.3); s0.bottom_margin = Cm(1.2)
    V.C.SERVICIO = 'EXPEDIENTE TÉCNICO DEL PROYECTO DE INVERSIÓN:'
    img = V.foto(); V.escudo()
    V.caratula(M.d, img, 'MEMORIA DESCRIPTIVA', D.ESPECIALIDAD)
    encabezado(M)
    R.indice(M)
    W = M.ancho

    # ------------------------------------------------------------------ 1
    M.t1('GENERALIDADES')
    M.t2('Objetivo')
    M.par('La presente memoria describe el planteamiento de evacuación y la señalización de seguridad del Centro de Acogida '
          'Residencial para niños y adolescentes varones (CAR Varones), con CUI N.° %s: las rutas de evacuación, las zonas seguras '
          'interiores y las zonas de reunión exteriores, la señalización de evacuación, de equipos contra incendio, de advertencia, '
          'de prohibición e informativa, y los equipos de protección contra incendio y de emergencia, según las láminas SE-E-01 '
          '(plano de evacuación) y SE-S-01 (plano de señalización).' % D.CUI)
    M.par('El objetivo es que, ante un sismo, un incendio u otra emergencia, los residentes, el personal y los visitantes se '
          'desplacen en forma ordenada y en el menor tiempo posible desde cualquier ambiente hasta un lugar seguro, guiados por la '
          'señalización, y que el personal identifique de inmediato la ubicación de los equipos de emergencia.')
    M.tabla('Datos generales del proyecto', ['DATO', 'DESCRIPCIÓN'], [
        ['Unidad ejecutora', 'Gerencia Territorial Bajo Mayo - Dirección de Infraestructura'],
        ['Proyecto', '“' + D.PROYECTO + '”'],
        ['Código único de inversiones', 'CUI N.° ' + D.CUI],
        ['Ubicación', D.UBICACION],
        ['Especialidad', 'Arquitectura - evacuación y señalización'],
        ['Láminas', 'SE-E-01 Plano de evacuación y SE-S-01 Plano de señalización (escala 1/150, setiembre de 2026)']],
        [4.5, W - 4.5], izq=(0, 1))
    M.t2('Alcance')
    M.par('Comprende las diez edificaciones del centro (E-01 a E-10), las obras complementarias (O.C. 01 a O.C. 04) y las obras '
          'exteriores (O.E. 01 a O.E. 07). Las cantidades citadas son las del metrado de evacuación y señalización y las '
          'características de cada partida se detallan en las especificaciones técnicas de la especialidad; los tres documentos y '
          'las láminas dicen lo mismo.')

    # ------------------------------------------------------------------ 2
    M.t1('NORMATIVIDAD')
    for t in ('Reglamento Nacional de Edificaciones, Norma A.010 Condiciones Generales de Diseño.',
              'Reglamento Nacional de Edificaciones, Norma A.120 Accesibilidad Universal en Edificaciones (señales con texto en braille).',
              'Reglamento Nacional de Edificaciones, Norma A.130 Requisitos de Seguridad (medios de evacuación, señalización, '
              'iluminación de emergencia, detección y alarma, y extintores portátiles).',
              'Reglamento Nacional de Edificaciones, Norma EM.010 Instalaciones Eléctricas Interiores.',
              'NTP 399.010-1 Señales de seguridad. Colores, símbolos, formas y dimensiones de señales de seguridad. Parte 1: Reglas '
              'para el diseño de las señales de seguridad y franjas de seguridad.',
              'NTP 350.043-1 Extintores portátiles. Selección, distribución, inspección, mantenimiento, recarga y prueba '
              'hidrostática (versión vigente).',
              'D.S. N.° 002-2018-PCM, Reglamento de Inspecciones Técnicas de Seguridad en Edificaciones.',
              'Normas NFPA de referencia: NFPA 10 (extintores portátiles), NFPA 72 (detección y alarma) y NFPA 101 (código de '
              'seguridad humana).'):
        M.vin(t)

    # ------------------------------------------------------------------ 3
    M.t1('DESCRIPCIÓN DEL CENTRO DE ACOGIDA')
    M.t2('Características generales')
    M.par('El CAR Varones es un centro de acogida residencial que brinda alojamiento, alimentación, cuidado y atención integral a '
          'niños y adolescentes varones sin cuidados parentales o en riesgo de perderlos. Funciona de manera permanente, las 24 horas '
          'del día, con residentes que duermen en el local; por ello la evacuación considera a niños pequeños (módulo de 3 a 5 años) '
          'que necesitan la conducción de un adulto y el horario nocturno, en el que la detección temprana del humo y la '
          'iluminación de emergencia son determinantes.')
    M.par('Todas las edificaciones son de un solo nivel, sin escaleras ni rampas en las rutas de evacuación: la salida de cada '
          'ambiente se hace directamente a las circulaciones exteriores, a los patios y a las zonas de reunión.')
    M.t2('Zonificación')
    M.tabla('Edificaciones del centro (planta general, láminas SE-S-01 y SE-E-01)', ['CÓDIGO', 'EDIFICACIÓN'],
            [[c, n] for c, n in eds if c != 'EXT'], [3.0, W - 3.0], izq=(1,))
    M.tabla('Obras complementarias y obras exteriores', ['CÓDIGO', 'TIPO', 'DESCRIPCIÓN'],
            [list(o) for o in OBRAS], [2.4, 3.6, W - 6.0], izq=(2,))
    M.t2('Población y aforo')
    M.par('La población permanente del centro es de %d personas: %d residentes y %d personas de personal (%d madres sustitutas, %d '
          'de personal administrativo y %d de personal de apoyo). A ella se suman los usuarios temporales (visitas, familiares y '
          'personal eventual), %d personas, con lo que la capacidad máxima simultánea es de %d personas, igual al aforo del SUM, '
          'el ambiente de mayor capacidad, donde puede reunirse toda la población del centro.' % (
              Pb['permanentes'], Pb['residentes'], Pb['personal'], Pb['madres'], Pb['administrativos'], Pb['apoyo'],
              Pb['temporales'], Pb['maxima']))
    M.tabla('Población del centro', ['POBLACIÓN', 'DETALLE', 'PERSONAS'], [
        ['Personal', 'Madres sustitutas', Pb['madres']], ['Personal', 'Personal administrativo', Pb['administrativos']],
        ['Personal', 'Personal de apoyo', Pb['apoyo']], ['Personal', 'Subtotal', Pb['personal']],
        ['Residentes', 'Niños y adolescentes varones (cuatro módulos de 10)', Pb['residentes']],
        ['Permanentes', 'Personal + residentes', Pb['permanentes']], ['Temporales', 'Visitas y personal eventual', Pb['temporales']],
        ['CAPACIDAD MÁXIMA', 'Permanentes + temporales', Pb['maxima']]], [3.6, W - 6.0, 2.4], izq=(0, 1), total=True)
    M.par('El aforo de cada ambiente es el indicado en la lámina SE-E-01 y se publica en el letrero de aforo colocado en su ingreso:')
    M.tabla('Aforo por edificación (lámina SE-E-01)', ['CÓDIGO', 'AMBIENTE', 'AFORO (personas)'],
            [[c, a, n] for c, a, n in D.aforos(P)], [2.4, W - 6.0, 3.6], izq=(1,))

    # ------------------------------------------------------------------ 4
    M.t1('PLANTEAMIENTO DE EVACUACIÓN')
    M.t2('Criterios')
    for t in ('Cada ambiente evacúa por su puerta directamente al exterior; las puertas de los ambientes de mayor aforo (SUM, '
              'módulos y administración) abren en el sentido de la evacuación.',
              'Las rutas se trazan por las circulaciones exteriores y los patios, sin cruzar zonas de riesgo (cocina, tanque elevado, '
              'servicios generales ni la zona del pozo séptico) y terminan en una zona de reunión exterior.',
              'El recorrido desde la puerta más alejada de cada edificación hasta su zona de reunión no supera %.0f m, que es la distancia '
              'máxima de recorrido hasta un lugar seguro en edificaciones sin rociadores (Norma A.130).' % D.D_MAX,
              'Las zonas de reunión están lejos de fachadas, cables aéreos, árboles y elementos que puedan caer, y tienen espacio para '
              'la población de las edificaciones que llegan a ellas.',
              'Las rutas y salidas se señalizan con señales fotoluminiscentes y se iluminan con unidades de luces de emergencia.'):
        M.vin(t)
    M.t2('Rutas de evacuación')
    M.par('La lámina SE-E-01 define trece rutas de evacuación (RUTA 01 a RUTA 13), con su longitud. El tiempo de desplazamiento se '
          'estima con una velocidad de %.2f m/s, adoptada para niños pequeños y personas con movilidad reducida; el mayor recorrido '
          'es el de la ruta 03 (SUM y cocina), de %.2f m, que se cubre en %.0f segundos.' % (
              D.V_EVAC, max(r['L'] for r in rutas), max(r['L'] for r in rutas) / D.V_EVAC))
    M.tabla('Rutas de evacuación (lámina SE-E-01)', ['RUTA', 'AMBIENTE QUE EVACÚA', 'LONGITUD (m)', 'TIEMPO (s)', 'L ≤ 45 m'],
            [['RUTA %02d' % r['n'], r['ambiente'], '%.2f' % r['L'], '%.0f' % r['t'], 'cumple' if r['cumple'] else 'no cumple'] for r in rutas],
            [2.2, W - 9.4, 2.6, 2.2, 2.4], izq=(1,))
    M.t2('Zonas de reunión exteriores')
    M.par('Al final de las rutas se ubican %d zonas de reunión (zonas seguras externas) S1 a S%d, pintadas en el piso de concreto '
          'como un círculo de %.2f m de radio, con una línea de %.2f m de ancho de %s color amarillo y la letra S con su número. '
          'Cada círculo tiene %.2f m de perímetro y %.2f m² de pintura.' % (
              len(zs), len(zs), z['radio'], z['ancho'], z['pintura'], z['perimetro'], z['area']))
    M.tabla('Zonas de reunión exteriores (lámina SE-E-01)', ['ZONA', 'UBICACIÓN (edificación más cercana)'],
            [['S%d' % int(x['nombre'][1:]), x['cerca']] for x in zs], [2.4, W - 2.4], izq=(1,))
    M.t2('Procedimiento de evacuación')
    M.par('Las instrucciones de la lámina SE-E-01, que se difunden al personal y a los residentes y se publican junto a los planos '
          'de evacuación de cada edificación, son:')
    for t in ('Mantener la calma y no entrar en pánico.', 'Avisar de inmediato la emergencia a la persona responsable más cercana.',
              'No entretenerse recogiendo objetos personales.', 'En caso de incendio y humo intenso, gatear pegado a la pared.',
              'Seguir la señalización de evacuación hasta las zonas seguras indicadas en el plano.',
              'Por ninguna razón regresar al interior de la edificación.'):
        M.vin(t)
    M.par('Las madres sustitutas y el personal de turno conducen a los residentes de su módulo, empezando por los más pequeños, y '
          'cuentan a los niños en la zona de reunión; el personal administrativo verifica que ningún ambiente quede ocupado.')
    figura(M, 'SE-E-01', 'Plano de evacuación')

    # ------------------------------------------------------------------ 5
    M.t1('SEÑALIZACIÓN DE SEGURIDAD')
    M.t2('Criterios de diseño')
    M.par('Las señales siguen la NTP 399.010-1 en colores, formas y símbolos, y llevan el texto en braille (Norma A.120). Las de '
          'evacuación y de equipos contra incendio son fotoluminiscentes, para que se vean sin energía eléctrica.')
    M.tabla('Colores y formas de las señales (NTP 399.010-1)', ['TIPO DE SEÑAL', 'FORMA', 'COLOR DE SEGURIDAD', 'SEÑALES DEL PROYECTO'], [
        ['Evacuación y emergencia', 'Rectangular o cuadrada', 'Verde con símbolo blanco', 'Salida, zona segura en caso de sismo, botiquín'],
        ['Equipos contra incendio', 'Rectangular o cuadrada', 'Rojo con símbolo blanco', 'Extintor PQS y CO2, alarma, avisador sonoro, central de alarma'],
        ['Advertencia', 'Triangular', 'Amarillo con borde y símbolo negro', 'Atención riesgo eléctrico, riesgo de descarga eléctrica'],
        ['Prohibición', 'Circular con banda diagonal', 'Rojo con fondo blanco', 'Prohibido el ingreso, solo personal autorizado'],
        ['Informativas', 'Rectangular o cuadrada', 'Azul o verde con símbolo blanco', 'Servicios higiénicos, punto ecológico, aforo']],
        [3.6, 3.2, 3.4, W - 10.2], izq=(0, 3))
    M.t2('Simbología utilizada')
    M.par('La simbología es la de la leyenda de la lámina SE-S-01; la cantidad de cada señal es la dibujada en la planta, que '
          'coincide con la leyenda, y la altura es la de instalación medida desde el piso terminado hasta el borde inferior de la señal.')
    filas = []
    for p in pts:
        if p['und'] is None or p['item'].startswith('05'):
            continue
        filas.append(['01.' + p['item'], p['desc'].replace(' CON BRAILLE', '').capitalize(), p['dim'], p['altura'], p['total']])
    M.tabla('Señales del proyecto (leyenda de la lámina SE-S-01)', ['ITEM', 'SEÑAL', 'DIMENSIONES', 'ALTURA', 'CANT.'], filas,
            [1.7, W - 8.3, 3.0, 2.0, 1.6], tam=8.5, izq=(1,))
    M.par('Las señales de avisador sonoro, atención riesgo eléctrico, riesgo de descarga eléctrica y prohibido el ingreso se ubican '
          'en los tableros eléctricos, en servicios generales (tablero general, subestación y grupo electrógeno) y junto a los '
          'equipos del sistema de alarma, en coordinación con el plano de instalaciones eléctricas. La señal de ubicación de extintor '
          'se coloca sobre el equipo, con su borde inferior a 1.80 m, de modo que quede visible por encima del extintor.')
    M.t2('Distribución por edificación')
    cols = [('SALIDA', 'Salida'), ('ZONA SEGURA EN', 'Zona segura'), ('LETRERO DE AFORO', 'Aforo'), ('ALARMA CONTRA', 'Alarma'),
            ('BOTIQUIN', 'Botiquín'), ('SEÑALES DE SERVICIOS HIGIENICOS VAR', 'SS.HH.')]
    por = {p['clave']: p for p in pts if p['und']}
    filas = []
    for cod, nombre in eds:
        fila = [cod]
        for clave, _ in cols:
            if clave.startswith('SEÑALES DE SERVICIOS'):
                n = sum(por[k]['por'].get(cod, 0) for k in por if k.startswith('SEÑALES DE SERVICIOS'))
            else:
                n = por[clave]['por'].get(cod, 0)
            fila.append(n or '—')
        filas.append(fila)
    filas.append(['TOTAL'] + [sum(f[k] for f in filas if f[k] != '—') for k in range(1, len(cols) + 1)])
    M.tabla('Señales principales por edificación (planta de SE-S-01)', ['EDIF.'] + [c[1] for c in cols], filas,
            [1.9] + [(W - 1.9) / len(cols)] * len(cols), tam=8.5, total=True)
    M.par('El detalle de todas las partidas por edificación está en la hoja METRADO del metrado de evacuación y señalización.')
    figura(M, 'SE-S-01', 'Plano de señalización')

    # ------------------------------------------------------------------ 6
    M.t1('EQUIPOS DE PROTECCIÓN CONTRA INCENDIO Y EMERGENCIA')
    M.t2('Extintores portátiles')
    pq = por['SEÑAL DE UBICACIÓN DE EXTINTOR PQS']; co = por['SEÑAL DE UBICACIÓN DE EXTINTOR CO2']
    lugares = ', '.join('%s (%d)' % (c, n) if n > 1 else c for c, n in pq['por'].items() if n)
    M.par('Se instalan %d extintores de polvo químico seco (PQS) multipropósito tipo ABC de 6 kg, en %s, y %d extintores de '
          'dióxido de carbono (CO2) de 10 lb en la cocina del SUM (E-08), para fuegos de clase B y C de sus equipos eléctricos. '
          'Cada extintor se cuelga en su soporte de pared con la parte superior a no más de 1.50 m del piso, en un lugar visible y '
          'de libre acceso, con su señal de ubicación encima, y se selecciona, distribuye, inspecciona y recarga según la NTP 350.043-1.' % (
              pq['total'], lugares, co['total']))
    M.t2('Detección y alarma')
    M.par('El sistema de detección y alarma contra incendio tiene %d detectores de humo de techo distribuidos en los ambientes, '
          '%d estaciones manuales de alarma (una por edificación principal), %d avisadores sonoros y una central de alarma en la '
          'administración (E-01). Su alimentación, cableado y pruebas corresponden a las instalaciones eléctricas del proyecto.' % (
              q['DETECTOR DE HUMO'], q['ALARMA CONTRA'], q['AVISADOR SONORO']))
    M.t2('Iluminación de emergencia')
    M.par('Se instalan %d unidades de luces de emergencia autónomas a 2.10 m del piso, en las salidas, circulaciones y ambientes, '
          'con autonomía mínima de 90 minutos, que se encienden automáticamente al cortarse la energía (Norma A.130).' % q['UNIDAD DE LUCES'])
    M.t2('Primeros auxilios y residuos')
    M.par('Se señalizan %d botiquines de primeros auxilios, distribuidos en todas las edificaciones, y %d puntos ecológicos para '
          'el acopio selectivo de residuos sólidos.' % (q['BOTIQUIN'], q['PUNTO ECOL']))

    # ------------------------------------------------------------------ 7
    M.t1('MEDIDAS DE MITIGACIÓN Y RECOMENDACIONES')
    for t in ('Conformar el comité y las brigadas de emergencia del centro (evacuación, primeros auxilios y lucha contra incendio) '
              'con el personal de cada turno, de modo que siempre haya brigadistas en el local, también de noche.',
              'Realizar simulacros de evacuación al menos dos veces al año, uno de ellos en horario nocturno, y registrar los tiempos.',
              'Capacitar al personal en el uso de extintores y en la conducción de niños pequeños durante la evacuación.',
              'Mantener las rutas de evacuación y las salidas libres de muebles, materiales y objetos, y las puertas sin llave mientras '
              'haya personas en el ambiente.',
              'Inspeccionar los extintores cada mes y recargarlos según la NTP 350.043-1; probar las luces de emergencia y la alarma '
              'cada mes; limpiar los detectores de humo según el fabricante.',
              'Reponer de inmediato toda señal dañada, retirada o que haya perdido su fotoluminiscencia, y repintar los círculos de las '
              'zonas de reunión cuando pierdan visibilidad.',
              'Publicar el plano de evacuación de cada edificación junto a su ingreso.'):
        M.vin(t)

    # ------------------------------------------------------------------ 8
    M.t1('LÁMINAS Y DOCUMENTOS')
    M.tabla('Láminas de evacuación y señalización', ['LÁMINA', 'DESCRIPCIÓN', 'ESCALA'], [[a, b, c] for a, b, c in D.LAMINAS],
            [3.0, W - 5.4, 2.4], izq=(1,))
    M.par('Forman parte de la especialidad, con las mismas cantidades y características: el metrado de evacuación y señalización '
          '(Excel) y las especificaciones técnicas de evacuación y señalización (Word).')
    M.d.core_properties.author = ''; M.d.core_properties.title = 'Memoria descriptiva - Evacuación y señalización - CAR Varones'
    M.d.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
