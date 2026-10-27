# -*- coding: utf-8 -*-
"""Metrado de evacuación y señalización - CAR Varones (CUI 2705619), en Excel con fórmulas.

Cantidades del plano final de la arquitecta (datos_se.py / lee_plano_se.py): cada señal contada en la planta de la lámina
SE-S-01 y repartida por edificación (E-01 a E-10 y exteriores), cotejada con la columna CANTIDAD de su leyenda; las zonas
de reunión pintadas en el piso, de la lámina SE-E-01 (S1 a S13).

Hojas: RESUMEN (partidas y metrado, enlazado), METRADO (por edificación, total = suma, cotejo con la leyenda),
ZONAS DE REUNIÓN (perímetro y área de pintura de cada círculo), RUTAS Y AFOROS (recorrido, tiempo y aforo, de SE-E-01),
SEÑALES (dimensiones, material y altura de cada partida).

Uso: python3 herramientas/varones/senalizacion/metrado_se.py
Salida: entregables_varones/METRADO_EVACUACION_Y_SENALIZACION_CAR_VARONES.xlsx
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import datos_se as D  # noqa: E402

SALIDA = os.path.join(D.SAL, 'METRADO_EVACUACION_Y_SENALIZACION_CAR_VARONES.xlsx')
GRANATE, ROSA, GRIS = '7B1E2B', 'F2E6E8', 'F2F2F2'
FINO = Side(style='thin', color='808080')
BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
F = lambda **k: Font(name='Arial', size=k.pop('size', 9), **k)
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
IZQ = Alignment(horizontal='left', vertical='center', wrap_text=True)


def celda(ws, r, c, v, neg=False, fondo=None, color=None, al=CENTRO, fmt=None, borde=True, size=9):
    x = ws.cell(r, c, v); x.font = F(bold=neg, color=color, size=size); x.alignment = al
    if fondo: x.fill = PatternFill('solid', fgColor=fondo)
    if borde: x.border = BORDE
    if fmt: x.number_format = fmt
    return x


def cabecera(ws, titulo, ncol):
    """Encabezado común: título, proyecto, CUI, especialidad, ubicación, láminas y fecha."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    celda(ws, 1, 1, titulo, True, GRANATE, 'FFFFFF', size=12, borde=False); ws.row_dimensions[1].height = 24
    filas = [('PROYECTO', '“%s”' % D.PROYECTO), ('CUI', 'N.° ' + D.CUI), ('ESPECIALIDAD', D.ESPECIALIDAD), ('UBICACIÓN', D.UBICACION),
             ('LÁMINAS', ' y '.join('%s %s (esc. %s)' % l for l in D.LAMINAS)), ('FECHA', D.FECHA)]
    for k, (a, b) in enumerate(filas, 2):
        celda(ws, k, 1, a, True, GRIS, al=IZQ)
        ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=ncol)
        celda(ws, k, 2, b, k == 2, al=IZQ)
        for c in range(3, ncol + 1):
            ws.cell(k, c).border = BORDE
    ws.row_dimensions[2].height = 36
    return len(filas) + 3


def impresion(ws, ultima_col, ultima_fila, horizontal=True):
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_area = 'A1:%s%d' % (L(ultima_col), ultima_fila)
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5


def hoja_metrado(wb, P, pts):
    ws = wb.create_sheet('METRADO')
    eds = D.edificaciones(P)
    n_e = len(eds)
    c_sin = 4 + n_e; c_tot = c_sin + 1; c_ley = c_tot + 1; c_dif = c_ley + 1; c_fue = c_dif + 1
    r = cabecera(ws, 'METRADO DE EVACUACIÓN Y SEÑALIZACIÓN POR EDIFICACIÓN', c_fue)
    cab = ['ITEM', 'DESCRIPCIÓN', 'UND'] + [c for c, _ in eds] + ['SIN UBICAR', 'TOTAL', 'LEYENDA SE-S-01', 'DIFERENCIA', 'FUENTE DE LA CANTIDAD']
    for k, v in enumerate(cab, 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    for k, (_, nombre) in enumerate(eds, 4):
        celda(ws, r + 1, k, nombre, False, ROSA, size=7)
    for k in list(range(1, 4)) + list(range(c_sin, c_fue + 1)):
        ws.merge_cells(start_row=r, start_column=k, end_row=r + 1, end_column=k)
    ws.row_dimensions[r + 1].height = 48
    r += 2; filas = {}
    for p in pts:
        if p['und'] is None:
            celda(ws, r, 1, p['item'], True, GRIS); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=c_fue)
            celda(ws, r, 2, p['desc'], True, GRIS, al=IZQ)
            for c in range(3, c_fue + 1): ws.cell(r, c).border = BORDE
            r += 1; continue
        celda(ws, r, 1, p['item']); celda(ws, r, 2, p['desc'], al=IZQ); celda(ws, r, 3, p['und'])
        for k, (cod, _) in enumerate(eds, 4):
            v = p['por'].get(cod, 0)
            celda(ws, r, k, v if v else None, fmt='0')
        celda(ws, r, c_sin, p['sin_ubicar'] or None, fmt='0')
        celda(ws, r, c_tot, '=SUM(%s%d:%s%d)' % (L(4), r, L(c_sin), r), True, ROSA, fmt='0')
        celda(ws, r, c_ley, p['leyenda'], fmt='0')
        celda(ws, r, c_dif, '=%s%d-%s%d' % (L(c_tot), r, L(c_ley), r), fmt='0;-0;"conforme"')
        celda(ws, r, c_fue, p['fuente'], al=IZQ, size=8)
        filas[p['item']] = r; r += 1
    celda(ws, r, 2, 'TOTAL DE SEÑALES Y EQUIPOS', True, ROSA, al=IZQ)
    for c in range(4, c_tot + 1):
        celda(ws, r, c, '=SUM(%s%d:%s%d)' % (L(c), r - len(pts), L(c), r - 1), True, ROSA, fmt='0')
    for c in (1, 3, c_ley, c_dif, c_fue): celda(ws, r, c, None, fondo=ROSA)
    r += 2
    notas = ['Notas:',
             '1. Cantidad por edificación: cada símbolo de la planta de la lámina SE-S-01 se asigna a la edificación donde está (contorno del bloque de la planta general, '
             'con 0.60 m de margen para las señales de fachada); EXT = obras exteriores y complementarias (patios, ingresos, caseta, losa deportiva y circulaciones).',
             '2. SIN UBICAR: señales que en el DXF son imágenes pegadas (avisador sonoro, atención riesgo eléctrico, riesgo de descarga eléctrica y prohibido el ingreso); '
             'su ubicación no se exporta y vale la cantidad de la leyenda. Se instalan donde indica la lámina SE-S-01 y en coordinación con el plano de instalaciones eléctricas.',
             '3. Zona de reunión pintada en el piso: los %d círculos de la lámina SE-E-01 (S1 a S%d); la leyenda de SE-S-01 indica 11 y debe corregirse a %d.' % ((P['evacuacion']['circulos_zona_segura'],) * 3),
             '4. Extintores: uno por cada señal de ubicación de extintor de la planta (PQS tipo ABC de 6 kg y CO2 de 10 lb).',
             '5. Unidades de luces de emergencia y detectores de humo: equipos de la leyenda de SE-S-01; su conexión y circuitos son de instalaciones eléctricas.']
    for t in notas:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=c_fue)
        celda(ws, r, 1, t, t == 'Notas:', al=IZQ, borde=False, size=8); ws.row_dimensions[r].height = 12 if t == 'Notas:' else 24; r += 1
    for c, w in zip(range(1, c_fue + 1), [7, 46, 5] + [6.5] * n_e + [7.5, 7.5, 8.5, 9, 40]):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = ws.cell(10, 4)
    impresion(ws, c_fue, r)
    return filas, c_tot


def hoja_zonas(wb, P):
    ws = wb.create_sheet('ZONAS DE REUNIÓN')
    r = cabecera(ws, 'ZONAS DE REUNIÓN (ZONAS SEGURAS EXTERNAS) PINTADAS EN EL PISO - LÁMINA SE-E-01', 8)
    z = D.ZONA_SEGURA
    celda(ws, r, 1, 'Radio (m)', True, GRIS); celda(ws, r, 2, z['radio'], fmt='0.00')
    celda(ws, r, 3, 'Ancho de línea (m)', True, GRIS); celda(ws, r, 4, z['ancho'], fmt='0.00')
    ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=8)
    celda(ws, r, 5, 'Círculo pintado con %s color %s (leyenda de SE-E-01)' % (z['pintura'], z['color']), al=IZQ)
    rr, ra = r, r; r += 2
    for k, v in enumerate(['N.°', 'ZONA', 'UBICACIÓN (edificación más cercana)', 'RADIO (m)', 'ANCHO (m)', 'PERÍMETRO (m)', 'ÁREA PINTADA (m²)', 'UND'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1; r0 = r
    for k, zz in enumerate(D.zonas(P), 1):
        nombre = 'S%d' % int(zz['nombre'][1:])
        celda(ws, r, 1, k); celda(ws, r, 2, nombre); celda(ws, r, 3, zz['cerca'], al=IZQ)
        celda(ws, r, 4, '=$B$%d' % rr, fmt='0.00'); celda(ws, r, 5, '=$D$%d' % ra, fmt='0.00')
        celda(ws, r, 6, '=2*PI()*D%d' % r, fmt='0.00'); celda(ws, r, 7, '=F%d*E%d' % (r, r), fmt='0.00'); celda(ws, r, 8, 1, fmt='0')
        r += 1
    celda(ws, r, 3, 'TOTAL', True, ROSA, al=IZQ)
    for c in (6, 7, 8):
        celda(ws, r, c, '=SUM(%s%d:%s%d)' % (L(c), r0, L(c), r - 1), True, ROSA, fmt='0.00' if c < 8 else '0')
    for c in (1, 2, 4, 5): celda(ws, r, c, None, fondo=ROSA)
    for c, w in zip(range(1, 9), [10, 9, 40, 10, 10, 13, 15, 7]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 8, r, horizontal=False)
    return r


def hoja_rutas(wb, P):
    ws = wb.create_sheet('RUTAS Y AFOROS')
    r = cabecera(ws, 'RUTAS DE EVACUACIÓN, AFOROS Y POBLACIÓN - LÁMINA SE-E-01', 6)
    celda(ws, r, 1, 'Velocidad de desplazamiento (m/s)', True, GRIS, al=IZQ); ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
    celda(ws, r, 4, D.V_EVAC, fmt='0.00'); rv = r; r += 1
    celda(ws, r, 1, 'Recorrido máximo hasta un lugar seguro, sin rociadores - RNE A.130 (m)', True, GRIS, al=IZQ)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3); celda(ws, r, 4, D.D_MAX, fmt='0.00'); rd = r; r += 2
    for k, v in enumerate(['RUTA', 'AMBIENTE QUE EVACÚA', 'LONGITUD (m)', 'TIEMPO DE DESPLAZAMIENTO (s)', 'VERIFICACIÓN', ''], 1):
        if v: celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    for ru in D.rutas(P):
        celda(ws, r, 1, 'RUTA %02d' % ru['n']); celda(ws, r, 2, ru['ambiente'], al=IZQ); celda(ws, r, 3, ru['L'], fmt='0.00')
        celda(ws, r, 4, '=C%d/$D$%d' % (r, rv), fmt='0.0'); celda(ws, r, 5, '=IF(C%d<=$D$%d,"cumple","no cumple")' % (r, rd))
        r += 1
    r += 1
    for k, v in enumerate(['EDIFICACIÓN', 'AMBIENTE', 'AFORO (personas)'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    for cod, amb, n in D.aforos(P):
        celda(ws, r, 1, cod); celda(ws, r, 2, amb, al=IZQ); celda(ws, r, 3, n, fmt='0'); r += 1
    r += 1
    Pb = D.POBLACION
    for k, v in enumerate(['POBLACIÓN', 'DETALLE', 'PERSONAS'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1; r0 = r
    for a, b, n in (('Personal', 'Madres sustitutas', Pb['madres']), ('Personal', 'Personal administrativo', Pb['administrativos']),
                    ('Personal', 'Personal de apoyo', Pb['apoyo'])):
        celda(ws, r, 1, a, al=IZQ); celda(ws, r, 2, b, al=IZQ); celda(ws, r, 3, n, fmt='0'); r += 1
    celda(ws, r, 1, 'Personal', True, ROSA, al=IZQ); celda(ws, r, 2, 'Subtotal personal', True, ROSA, al=IZQ)
    celda(ws, r, 3, '=SUM(C%d:C%d)' % (r0, r - 1), True, ROSA, fmt='0'); rp = r; r += 1
    celda(ws, r, 1, 'Residentes', al=IZQ); celda(ws, r, 2, 'Niños y adolescentes varones (4 módulos)', al=IZQ); celda(ws, r, 3, Pb['residentes'], fmt='0'); rr = r; r += 1
    celda(ws, r, 1, 'Permanentes', True, ROSA, al=IZQ); celda(ws, r, 2, 'Personal + residentes', True, ROSA, al=IZQ)
    celda(ws, r, 3, '=C%d+C%d' % (rp, rr), True, ROSA, fmt='0'); rpe = r; r += 1
    celda(ws, r, 1, 'Temporales', al=IZQ); celda(ws, r, 2, 'Visitas y personal eventual', al=IZQ); celda(ws, r, 3, Pb['temporales'], fmt='0'); rt = r; r += 1
    celda(ws, r, 1, 'CAPACIDAD MÁXIMA', True, ROSA, al=IZQ); celda(ws, r, 2, 'Permanentes + temporales (= aforo del SUM)', True, ROSA, al=IZQ)
    celda(ws, r, 3, '=C%d+C%d' % (rpe, rt), True, ROSA, fmt='0')
    for c, w in zip(range(1, 7), [30, 36, 14, 16, 14, 4]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 5, r, horizontal=False)


def hoja_senales(wb, pts):
    ws = wb.create_sheet('SEÑALES')
    r = cabecera(ws, 'CARACTERÍSTICAS DE LAS SEÑALES Y EQUIPOS (NTP 399.010-1, RNE A.120 Y A.130)', 7)
    for k, v in enumerate(['ITEM', 'PARTIDA', 'DIMENSIONES', 'MATERIAL', 'ALTURA DE INSTALACIÓN', 'ALTURA EN LA LEYENDA', 'FUNCIÓN Y UBICACIÓN'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    for p in pts:
        if p['und'] is None:
            celda(ws, r, 1, p['item'], True, GRIS); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
            celda(ws, r, 2, p['desc'], True, GRIS, al=IZQ)
            for c in range(3, 8): ws.cell(r, c).border = BORDE
        else:
            for k, v in enumerate([p['item'], p['desc'], p['dim'], p['material'], p['altura'], p['altura_leyenda'] or '—', p['uso']], 1):
                celda(ws, r, k, v, al=IZQ if k in (2, 4, 7) else CENTRO, size=8)
        r += 1
    for c, w in zip(range(1, 8), [7, 40, 18, 40, 12, 11, 48]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 7, r - 1)


def generar():
    P = D.plano(); pts = D.partidas(P)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'RESUMEN'
    filas, c_tot = hoja_metrado(wb, P, pts)
    hoja_zonas(wb, P); hoja_rutas(wb, P); hoja_senales(wb, pts)
    r = cabecera(ws, 'RESUMEN DE METRADOS - EVACUACIÓN Y SEÑALIZACIÓN', 4)
    for k, v in enumerate(['ITEM', 'DESCRIPCIÓN', 'UND', 'METRADO'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    celda(ws, r, 1, '01', True, GRIS); celda(ws, r, 2, 'SEÑALIZACIÓN, EVACUACIÓN Y SEGURIDAD', True, GRIS, al=IZQ)
    celda(ws, r, 3, None, fondo=GRIS); celda(ws, r, 4, None, fondo=GRIS); r += 1
    for p in pts:
        it = '01.' + p['item']
        if p['und'] is None:
            celda(ws, r, 1, it, True); celda(ws, r, 2, p['desc'], True, al=IZQ); celda(ws, r, 3, None); celda(ws, r, 4, None)
        else:
            celda(ws, r, 1, it); celda(ws, r, 2, p['desc'], al=IZQ); celda(ws, r, 3, p['und'])
            celda(ws, r, 4, "='METRADO'!%s%d" % (L(c_tot), filas[p['item']]), True, fmt='0.00')
        r += 1
    for c, w in zip(range(1, 5), [11, 70, 7, 12]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 4, r - 1, horizontal=False)
    for w in wb.worksheets:
        w.sheet_view.showGridLines = False
    wb.properties.creator = ''; wb.properties.title = 'Metrado de evacuación y señalización - CAR Varones'
    wb.save(SALIDA)
    return SALIDA, pts


if __name__ == '__main__':
    s, pts = generar()
    print(s)
