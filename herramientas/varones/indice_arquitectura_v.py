# -*- coding: utf-8 -*-
"""Índices de ARQUITECTURA - CAR Varones (CUI 2705619), en Word, con el mismo estilo de los índices del proyecto
(indice_planos_refugio.py / indice_arquitectura_refugio.py): encabezado y pie, proyecto con CUI, título, especialidad, tabla con
encabezado repetido y filas de grupo.

Fuente: el índice en Excel de la usuaria (insumos/katiuska/INDICE_DE_LAMINAS_ARQUITECTURA_CAR_VARONES.xlsx, hoja ARQUITECTURA):
Tomo I documentación (memorias, especificaciones y metrados) y Tomos II a V planos. Se toma cada fila tal cual, con estos arreglos
de forma:
  - tomos y grupos en orden; láminas numeradas correlativas (el Excel no numera UL-02, PT-01 ni DT-09/00-03);
  - códigos sin espacios sueltos ("PT - 01" -> "PT-01", "OC - 02- PZ - 01" -> "OC-02-PZ-01", "DT-01/ 07-8" -> "DT-01/07-08");
  - tildes y erratas en las descripciones (UBICACIÓN, EXCEDENTE, INTERIOR, SERVICIOS, CIELO RASO, MOBILIARIO, EXTERIORES, ...);
  - ED-02.2/2: la descripción toma la hoja de su código (EDIFICACIÓN 02 / 02 - 02);
  - SE-E-01 evacuación y SE-S-01 señalización, como en los rótulos del plano SE (entregables_varones/_calc/senalizacion_plano.json);
  - MB-01 a MB-09: título de cada rótulo (entregables_varones/_calc/mobiliario_plano.json).
Las notas sueltas de la columna S (cartas y anexos de obra) no son láminas y no se incluyen.

Uso: python3 herramientas/varones/indice_arquitectura_v.py
Salidas: entregables_varones/INDICE_GENERAL_EXPEDIENTE_ARQUITECTURA_CAR_VARONES.docx (tomos I a V)
         entregables_varones/INDICE_DE_PLANOS_ARQUITECTURA_CAR_VARONES.docx (solo planos)
"""
import json
import os
import re
import sys

import docx
import openpyxl
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, os.path.dirname(AQUI))
import indice_planos_refugio as I  # noqa: E402  (mismas piezas de formato)

XLSX = os.path.join(RAIZ, 'insumos/katiuska/INDICE_DE_LAMINAS_ARQUITECTURA_CAR_VARONES.xlsx')
MB_JSON = os.path.join(RAIZ, 'entregables_varones/_calc/mobiliario_plano.json')
SALIDA_G = os.path.join(RAIZ, 'entregables_varones/INDICE_GENERAL_EXPEDIENTE_ARQUITECTURA_CAR_VARONES.docx')
SALIDA_P = os.path.join(RAIZ, 'entregables_varones/INDICE_DE_PLANOS_ARQUITECTURA_CAR_VARONES.docx')
CUI = '2705619'
PROYECTO = ('“CREACIÓN DEL SERVICIO DE PROTECCIÓN INTEGRAL A NIÑAS, NIÑOS Y ADOLESCENTES SIN CUIDADOS PARENTALES O EN RIESGO DE '
            'PERDERLOS EN CENTRO DE ACOGIDA RESIDENCIAL - VARONES DISTRITO DE MORALES DE LA PROVINCIA DE SAN MARTÍN DEL '
            'DEPARTAMENTO DE SAN MARTÍN”')
CORTO = 'CAR VARONES - MORALES, SAN MARTÍN'
UBIC = 'Distrito de Morales, provincia y departamento de San Martín'
ROMANOS = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X']
TOMOS = {'TOMO I': 'DOCUMENTACIÓN', 'TOMO II': 'PLANOS DE ARQUITECTURA', 'TOMO III': 'DESARROLLO DE MÓDULOS Y EVACUACIÓN',
         'TOMO IV': 'DETALLES DE ARQUITECTURA', 'TOMO V': 'OBRAS COMPLEMENTARIAS, OBRAS EXTERIORES Y MOBILIARIO'}
DOCUMENTOS = {'MEMORIAS': 'MEMORIA DESCRIPTIVA', 'ESPECIFICACIONES': 'ESPECIFICACIONES TÉCNICAS', 'METRADOS': 'METRADO'}
TILDES = [('EDIFICACION', 'EDIFICACIÓN'), ('PERIMETRICO', 'PERIMÉTRICO'), ('ZONIFICACION', 'ZONIFICACIÓN'),
          ('UBICACION', 'UBICACIÓN'), ('MODULOS', 'MÓDULOS'), ('EVACUACION', 'EVACUACIÓN'), ('SEÑALIZACION', 'SEÑALIZACIÓN'),
          ('TOPOGRAFICO', 'TOPOGRÁFICO'), ('DEPOSITO', 'DEPÓSITO'), ('ESCEDENTE', 'EXCEDENTE'), ('PUBLICOS', 'PÚBLICOS'),
          ('ADMINISTRACION', 'ADMINISTRACIÓN'), ('SERIVICIOS', 'SERVICIOS'), ('CIVICO', 'CÍVICO'), ('INTERIO', 'INTERIOR'),
          ('RAZO', 'RASO'), ('AREA', 'ÁREA'), ('SEMBRIOS', 'SEMBRÍOS'), ('MOBILIRIO', 'MOBILIARIO'),
          ('EXTERIOREES', 'EXTERIORES'), ('DOCUMENTACIÒN', 'DOCUMENTACIÓN')]


def _limpio(t):
    t = re.sub(r'\s+', ' ', str(t)).strip()
    for a, b in TILDES:
        t = re.sub(r'\b%s\b' % a, b, t)
    t = re.sub(r'(\d)\s*/\s*(\d)', r'\1 / \2', t)              # "02/ 02" -> "02 / 02"
    t = re.sub(r'(\S)\(', r'\1 (', t)                         # "VENTANAS(" -> "VENTANAS ("
    t = re.sub(r'\(\s+', '(', re.sub(r'\s+\)', ')', t))        # "(... DEPORTIVA )" -> "(... DEPORTIVA)"
    t = re.sub(r'\)(\d)', r') \1', t)                         # "AÑOS)2" -> "AÑOS) 2"
    t = re.sub(r'\s+,', ',', t); t = re.sub(r',(\S)', r', \1', t)   # "MANIOBRAS , ESTAC." / "05,06" -> "05, 06"
    t = t.replace('EDIF, ', 'EDIF. ')
    t = re.sub(r'(\S) -(\S)', r'\1 - \2', t)                  # "DETALLES -ACCESORIOS" -> "DETALLES - ACCESORIOS"
    t = re.sub(r'^OC-\s*', 'OC - ', t)                        # "OC- DETALLE" -> "OC - DETALLE"
    t = re.sub(r'GENERAL-(\S)', r'GENERAL - \1', t)           # "PLANTA GENERAL-TECHOS"
    return t


def _codigo(t):
    c = re.sub(r'\s+', '', str(t)).strip()                   # "OC - 02- PZ - 01" -> "OC-02-PZ-01"
    return re.sub(r'/(\d\d)-(\d)$', r'/\1-0\2', c)            # "DT-01/07-8" -> "DT-01/07-08"


def _mb_titulos():
    try:
        lam = json.load(open(MB_JSON, encoding='utf-8'))['laminas']
    except (OSError, KeyError, ValueError):
        return {}
    out = {}
    for k, v in (lam.items() if isinstance(lam, dict) else ((x.get('codigo'), x) for x in lam)):
        t = v.get('titulo') if isinstance(v, dict) else v
        if k and t:
            out[re.sub(r'\s+', '', k)] = _limpio(t).replace('MOBILIARIO EDIFICACIÓN', 'MOBILIARIO - EDIFICACIÓN')
    return out


def leer():
    """[(tomo, [(grupo, [filas])])]; las filas del Tomo I son documentos y las demás láminas."""
    ws = openpyxl.load_workbook(XLSX, data_only=True)['ARQUITECTURA']
    mb = _mb_titulos()
    tomos, tomo, grupo, cambios = [], None, None, []
    for r in range(6, ws.max_row + 1):
        a, b, g, c = (ws.cell(r, k).value for k in (1, 2, 7, 9))
        if b is None:
            continue
        b_ = _limpio(b)
        if b_.startswith('TOMO '):
            tomo = [b_, []]; tomos.append(tomo); grupo = None; continue
        if g is not None and str(g).strip().upper().startswith('LÁMINA'):   # fila de grupo
            grupo = [b_, []]; tomo[1].append(grupo); continue
        if g is None:                                                     # Tomo I: bloque o documento
            if b_ in DOCUMENTOS:
                grupo = [b_, []]; tomo[1].append(grupo)
            else:
                grupo[1].append(dict(codigo='—', descripcion='%s - %s' % (DOCUMENTOS[grupo[0]], b_), formato='—', cantidad=None))
            continue
        cod, desc = _codigo(g), b_
        if cod != re.sub(r'\s+', '', str(g)).strip():
            cambios.append('código %s -> %s' % (str(g).strip(), cod))
        m = re.match(r'ED-(\d\d)\.(\d)/(\d)$', cod)
        d = re.search(r'EDIFICACIÓN (\d\d) / (\d\d) - (\d\d)', desc)
        if m and d and d.groups() != (m.group(1), '0' + m.group(2), '0' + m.group(3)):
            nuevo = 'EDIFICACIÓN %s / 0%s - 0%s' % m.groups(); cambios.append('%s: "%s" -> "%s"' % (cod, desc, nuevo)); desc = nuevo
        if cod in ('SE-S-01', 'SE-E-01'):                                  # rótulos: E = evacuación, S = señalización
            nuevo = 'PLANTA GENERAL DE ' + ('EVACUACIÓN' if cod == 'SE-E-01' else 'SEÑALIZACIÓN')
            if nuevo != desc:
                cambios.append('%s: "%s" -> "%s"' % (cod, desc, nuevo)); desc = nuevo
        if cod in mb and desc == 'MOBILIARIO':
            desc = mb[cod]
        grupo[1].append(dict(codigo=cod, descripcion=desc, formato='—', cantidad=c or 1, n_excel=a))
    for t in tomos:                                                   # SE-E-01 antes que SE-S-01
        for gr in t[1]:
            gr[1].sort(key=lambda f: (0 if f['codigo'] == 'SE-E-01' else 1) if f['codigo'].startswith('SE-') else 0)
    return tomos, cambios


def _documento(titulo, nombre):
    d = docx.Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7); s.orientation = WD_ORIENT.PORTRAIT
    s.left_margin, s.right_margin, s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.6), Cm(1.7), Cm(1.4)
    s.header_distance, s.footer_distance = Cm(0.9), Cm(0.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(9)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
    ancho = (s.page_width - s.left_margin - s.right_margin) / 360000.0
    G = I.GRANATE
    h = s.header.paragraphs[0]; h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    I._run(h, CORTO, 7.5, True, G); I._run(h, '   |   CUI N.° %s   |   ARQUITECTURA' % CUI, 7.5, False, G)
    I._borde_inferior(h, G, 8)
    f = s.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    I._run(f, 'Página ', 7.5, color=I.GRIS); I._campo(f, 'PAGE', 7.5); I._run(f, ' de ', 7.5, color=I.GRIS); I._campo(f, 'NUMPAGES', 7.5)

    def par(txt, tam, neg=False, color=None, antes=0, despues=2, centro=True):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.LEFT
        pf = p.paragraph_format; pf.space_before = Pt(antes); pf.space_after = Pt(despues); pf.line_spacing = 1.0
        I._run(p, txt, tam, neg, color)
        return p

    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    par('GERENCIA TERRITORIAL BAJO MAYO - DIRECCIÓN DE INFRAESTRUCTURA', 9, True, I.GRIS, despues=4)
    par('PROYECTO: ' + PROYECTO + ', CON CUI N.° ' + CUI, 7.5, False, None, despues=6)
    p = par(titulo, 15, True, G, despues=2); I._borde_inferior(p, G)
    par('ESPECIALIDAD: ARQUITECTURA', 9.5, True, despues=2, antes=4)
    par(UBIC, 8, False, I.GRIS, despues=8)
    d.core_properties.author = ''
    d.core_properties.title = nombre
    return d, ancho, par


def _tabla(d, anchos, titulos):
    tb = d.add_table(rows=1, cols=len(anchos)); tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.autofit = False
    for gc, w in zip(tb._tbl.tblGrid.findall(qn('w:gridCol')), anchos):
        gc.set(qn('w:w'), str(int(w * 567)))
    for c, w, t in zip(tb.rows[0].cells, anchos, titulos):
        c.width = Cm(w); I._celda(c, [(t, 7.5, True, 'FFFFFF')]); I._sombra(c, I.GRANATE)
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); tb.rows[0]._tr.get_or_add_trPr().append(e)
    return tb


def _fila_banda(tb, texto, fondo, color, tam=7.5):
    fila = tb.add_row(); fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    c = fila.cells[0].merge(fila.cells[-1])
    I._celda(c, [(texto, tam, True, color)], centro=False); I._sombra(c, fondo)


def _fila(tb, anchos, valores):
    fila = tb.add_row(); fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    for c, w in zip(fila.cells, anchos):
        c.width = Cm(w)
    for k, (c, v) in enumerate(zip(fila.cells, valores)):
        I._celda(c, [(v, 8, k in (1, 2), I.GRANATE if k == 1 else None)], centro=(k != 2))


def _planos(tomos):
    return [(gr, filas) for t, grs in tomos if t != 'TOMO I' for gr, filas in grs]


def _totales(par, n, copias):
    par('Total: %d láminas de arquitectura, %d copias de cada una (%d láminas impresas); formato no indicado en el índice de láminas.'
        % (n, copias, n * copias), 9, True, antes=8, despues=2, centro=False)


def generar_planos(tomos):
    d, ancho, par = _documento('ÍNDICE DE PLANOS', 'Índice de planos - Arquitectura - CAR Varones')
    anchos = (1.0, 3.2, ancho - 1.0 - 3.2 - 1.9 - 1.4, 1.9, 1.4)
    tb = _tabla(d, anchos, ('N.°', 'LÁMINA', 'DESCRIPCIÓN', 'FORMATO', 'CANT.'))
    n = 0
    for k, (gr, filas) in enumerate(_planos(tomos)):
        _fila_banda(tb, '%s.  %s  (%d)' % (ROMANOS[k], gr, len(filas)), I.GRIS_CLARO, I.GRANATE)
        for f in filas:
            n += 1; _fila(tb, anchos, ('%02d' % n, f['codigo'], f['descripcion'], f['formato'], '%d' % f['cantidad']))
    _totales(par, n, filas[-1]['cantidad'])
    par('Láminas agrupadas por tema y numeradas de la primera a la última; códigos según el índice de láminas de arquitectura '
        'del proyecto.', 8, False, I.GRIS, despues=2, centro=False)
    d.save(SALIDA_P)
    return SALIDA_P, n


def generar_general(tomos):
    d, ancho, par = _documento('ÍNDICE GENERAL DEL EXPEDIENTE', 'Índice general del expediente - Arquitectura - CAR Varones')
    G = I.GRANATE
    # Resumen por tomo
    par('CONTENIDO POR TOMO', 9.5, True, G, despues=3, centro=False)
    an = (1.8, ancho - 1.8 - 2.6, 2.6)
    tb = _tabla(d, an, ('TOMO', 'CONTENIDO', 'CANTIDAD'))
    for t, grs in tomos:
        cant = sum(len(f) for _, f in grs)
        cont = ('DOCUMENTACIÓN: MEMORIAS DESCRIPTIVAS, ESPECIFICACIONES TÉCNICAS Y METRADOS' if t == 'TOMO I'
                else ' / '.join(g for g, _ in grs))
        _fila(tb, an, (t.replace('TOMO ', ''), cont,
                       '%d %s' % (cant, 'documentos' if t == 'TOMO I' else 'láminas')))
    # Detalle
    par('DETALLE DEL CONTENIDO', 9.5, True, G, antes=10, despues=3, centro=False)
    anchos = (1.0, 3.2, ancho - 1.0 - 3.2 - 1.9 - 1.4, 1.9, 1.4)
    tb = _tabla(d, anchos, ('N.°', 'LÁMINA', 'DESCRIPCIÓN', 'FORMATO', 'CANT.'))
    nd = n = k = 0
    for t, grs in tomos:
        _fila_banda(tb, '%s  -  %s' % (t, TOMOS.get(t, '')), G, 'FFFFFF', 8)
        for gr, filas in grs:
            if t == 'TOMO I':
                _fila_banda(tb, '%s  (%d)' % ({'ESPECIFICACIONES': 'ESPECIFICACIONES TÉCNICAS'}.get(gr, gr), len(filas)),
                            I.GRIS_CLARO, G)
                for f in filas:
                    nd += 1; _fila(tb, anchos, ('D-%02d' % nd, '—', f['descripcion'], '—', '—'))
                continue
            _fila_banda(tb, '%s.  %s  (%d)' % (ROMANOS[k], gr, len(filas)), I.GRIS_CLARO, G); k += 1
            for f in filas:
                n += 1; _fila(tb, anchos, ('%02d' % n, f['codigo'], f['descripcion'], f['formato'], '%d' % f['cantidad']))
    par('Total: %d documentos en el Tomo I y %d láminas de arquitectura en los Tomos II al V, %d copias de cada lámina (%d láminas '
        'impresas); formato no indicado en el índice de láminas.' % (nd, n, filas[-1]['cantidad'], n * filas[-1]['cantidad']),
        9, True, antes=8, despues=2, centro=False)
    par('Documentos del Tomo I numerados D-01 a D-%02d; láminas agrupadas por tema y numeradas de la primera a la última, con los '
        'códigos del índice de láminas de arquitectura del proyecto.' % nd, 8, False, I.GRIS, despues=2, centro=False)
    d.save(SALIDA_G)
    return SALIDA_G, nd, n


if __name__ == '__main__':
    tomos, cambios = leer()
    for t, grs in tomos:
        print(t, [(g, len(f)) for g, f in grs])
    for c in cambios:
        print('  ', c)
    print(generar_general(tomos))
    print(generar_planos(tomos))
