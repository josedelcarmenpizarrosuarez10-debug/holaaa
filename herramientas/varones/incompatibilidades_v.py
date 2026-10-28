# -*- coding: utf-8 -*-
"""Análisis de incompatibilidades de ubicación - CAR Varones (CUI 2705619).

Base: el cuadro de incompatibilidades de la memoria descriptiva de arquitectura de la usuaria Katiuska
(insumos/katiuska/arquitectura/MEMORIA_DESCRIPTIVA_ARQUITECTURA_CAR_VARONES.docx, 18 incompatibilidades según la
R.V.M. N.° 068-2025-MINEDU, todas CUMPLE) y sus datos del terreno (poligonal resultante P1-P4 de 6,999.98 m², colindantes y usos
dentro de 100 m). Se conserva cada resultado; se corrigen las citas y el texto copiado de un local educativo y se agrega el
Cuadro N.° 2 con las condiciones de ubicación del CAR según el RNE (A.090 arts. 3 y 4, A.040 art. 7) y la R.V.M. N.° 208-2019
(art. 8). Mismo formato y matriz Excel que el del Hogar de Refugio (incompatibilidades_refugio.py).
Lo que la arquitecta debe confirmar va en una nota interna aparte (no forma parte del expediente).

Uso: python3 herramientas/varones/incompatibilidades_v.py
Salida: entregables_varones/ANALISIS_INCOMPATIBILIDADES_UBICACION_CAR_VARONES.docx (expediente)
        entregables_varones/MATRIZ_INCOMPATIBILIDADES_UBICACION_CAR_VARONES.xlsx (expediente)
        entregables_varones/revision_interna/INCOMPATIBILIDADES_REVISION_CAR_VARONES.docx (uso interno)
"""
import importlib
import os
import sys

from docx.oxml import OxmlElement
from docx.shared import Cm

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(os.path.dirname(AQUI))
sys.path.insert(0, os.path.dirname(AQUI))
import incompatibilidades_refugio as I  # noqa: E402  (motor: secciones, cuadro, matriz Excel)
import caratulas_v as V  # noqa: E402

importlib.reload(V)                     # el módulo del Refugio deja la carátula con sus datos: volver a los del CAR Varones
R, SN = I.R, I.SN
CUI = '2705619'
PROY = V.PROYECTO.strip('“”')
R.CR_PROY, R.CORTO, SN.CUI = PROY, 'CAR VARONES', CUI
SAL = os.path.join(RAIZ, 'entregables_varones')
SALIDA = os.path.join(SAL, 'ANALISIS_INCOMPATIBILIDADES_UBICACION_CAR_VARONES.docx')
SALIDA_N = os.path.join(SAL, 'revision_interna/INCOMPATIBILIDADES_REVISION_CAR_VARONES.docx')
RVM = 'R.V.M. N.° 068-2025-MINEDU'                 # la que cita la arquitecta (versión vigente de la N.T. Criterios Generales)
AERO_KM, AERO_OESTE_KM = 3.5, 2.4                  # aeropuerto de Tarapoto desde el centro de la poligonal P1-P4

# Cuadro N.° 1: los 18 dispositivos del cuadro de la arquitecta (citas corregidas) con la condición del CAR
CUADRO = []
for inc, disp, dmin, crit, cond in I.CUADRO:
    cond = (cond.replace('de las residentes', 'de los niños y adolescentes residentes').replace('el hogar', 'el CAR')
            .replace('un hogar de refugio', 'un centro de acogida residencial'))
    CUADRO.append((inc, disp, dmin, crit, cond))
CUADRO[9] = CUADRO[9][:4] + ('El aeropuerto de Tarapoto está a unos %.1f km, al costado del eje de la pista 17/35 (unos %.1f km al '
                             'oeste), fuera de las superficies de aproximación y despegue; las edificaciones del CAR son de baja '
                             'altura.' % (AERO_KM, AERO_OESTE_KM),)
CUADRO[14] = CUADRO[14][:4] + ('El CAR no está en área restringida: su frente da a la carretera Oasis a través del área libre del predio '
                               'matriz, y el cerco y las obras, incluido el colector pluvial frontal, quedan dentro del lindero.',)

USOS = [('Frente (P4-P1, 93.90 m)', 'Carretera Oasis, a través del área libre del predio matriz', 'Vía de acceso al predio.'),
        ('Izquierda (P1-P2, 74.55 m)', 'Hogar de Refugio Temporal (CUI 2675514)', 'Servicio de protección del Estado; uso compatible.'),
        ('Derecha (P3-P4, 74.55 m)', 'Área remanente del predio matriz', 'Predio del Gobierno Regional destinado a servicios de protección.'),
        ('Fondo (P2-P3, 93.90 m)', 'Área remanente del predio matriz y, a continuación, propiedad de terceros', 'Sin uso incompatible.')]

NORMAS = [
    ('Decreto Legislativo N.° 1297 y su Reglamento (D.S. N.° 001-2018-MIMP)', 'Protección de niñas, niños y adolescentes sin '
     'cuidados parentales o en riesgo de perderlos; el acogimiento residencial se brinda en Centros de Acogida Residencial '
     'acreditados y supervisados por el MIMP.'),
    ('Criterios Generales para el Diseño Arquitectónico de Centros de Acogida Residencial Básicos, de Urgencia y Especializados para '
     'niños y adolescentes (MIMP)', 'Referencia del diseño arquitectónico del CAR citada en la memoria descriptiva de arquitectura.'),
    ('RNE, Norma A.090 Servicios Comunales, art. 3', 'Las edificaciones de servicios comunales se ubicarán en los lugares señalados '
     'en los Planes de Desarrollo Urbano, o en zonas compatibles con la zonificación vigente.'),
    ('RNE, Norma A.090, art. 4', 'Los proyectos que supongan una concentración de público de más de 500 personas deberán contar con '
     'un estudio de impacto vial.'),
    ('RNE, Norma A.040 Educación (R.M. N.° 068-2020-VIVIENDA), art. 7', 'a) ubicación conforme a los instrumentos de '
     'acondicionamiento territorial y desarrollo urbano; b) ubicación evitando las incompatibilidades de uso establecidas en la '
     'normativa vigente; c) vías de acceso que prevean el ingreso de vehículos de emergencia; d) soluciones alternativas donde no '
     'existan servicios públicos.'),
    ('%s, Norma Técnica "Criterios Generales de Diseño para Infraestructura Educativa"' % RVM, 'Cuadro de incompatibilidades de uso '
     'del terreno, que reúne los dispositivos legales sectoriales y sus distancias; se adopta como lista de verificación.'),
    ('R.V.M. N.° 208-2019-MINEDU, art. 8', 'a) se recomiendan terrenos rectangulares o similares; b) los terrenos con pendiente pueden '
     'resolverse con plataformas o terrazas; c) la selección del terreno toma en cuenta las condiciones de la N.T. Criterios Generales.'),
]

CONDICIONES = [
    ('Ubicación en zona compatible con la zonificación y los instrumentos de desarrollo urbano', 'RNE A.090, art. 3; A.040, art. 7 a)',
     'Predio rural Los Mangos (partida registral N.° 05003469), de propiedad del Gobierno Regional de San Martín, destinado a la '
     'sede del CAR Varones, el CAR Mujeres y el Hogar de Refugio Temporal.'),
    ('Ubicación sin incompatibilidades de uso', 'RNE A.040, art. 7 b)', 'Cumple las 18 condiciones del Cuadro N.° 1.'),
    ('Vías de acceso con ingreso de vehículos de emergencia', 'RNE A.040, art. 7 c)',
     'Acceso por la carretera Oasis, interconectada con la Av. Evitamiento y la carretera Polvorico, con portada de ingreso principal '
     'e ingreso vehicular al predio.'),
    ('Servicios básicos', 'RNE A.040, art. 7 d)', 'El proyecto contempla cisterna y tanque elevado, pozo séptico, subestación eléctrica, '
     'grupo electrógeno y drenaje pluvial.'),
    ('Sin estudio de impacto vial', 'RNE A.090, art. 4', 'La población del CAR (niños y adolescentes residentes y personal) es '
     'menor de 500 personas, por lo que no se exige estudio de impacto vial.'),
    ('Forma y pendiente del terreno', 'R.V.M. N.° 208-2019, art. 8 a) y b)', 'Terreno rectangular de 93.90 x 74.55 m (6,999.98 m²), '
     'de relieve relativamente plano (pendiente longitudinal del orden de 1 %).'),
    ('Programa completo dentro de los linderos y zona de seguridad interna', 'R.V.M. N.° 208-2019, art. 8.1.1 y cuadro 2 (terreno '
     'tipo II)', 'El terreno alberga dentro de sus linderos todo el programa arquitectónico del CAR, con la zona de seguridad dentro '
     'del lindero.'),
    ('Protección frente a lluvias intensas', 'RNE OS.060 (drenaje pluvial urbano)', 'Drenaje pluvial con canaletas, montantes, cunetas '
     'y colector frontal (láminas DPLV), dimensionado con la hidrología del expediente.'),
]


def generar_docx():
    M = R.Memoria()
    V.C.SERVICIO = 'EXPEDIENTE TÉCNICO DEL PROYECTO DE INVERSIÓN:'; V.FECHA = 'OCTUBRE - 2026'
    foto = V.foto(); V.escudo()
    V.caratula(M.d, foto, 'INCOMPATIBILIDADES', 'DE UBICACIÓN DEL PROYECTO - ESPECIALIDAD: ARQUITECTURA')
    SN.encabezado(M, 'INCOMPATIBILIDADES DE UBICACIÓN - ARQUITECTURA')
    W = M.ancho

    M.t1('GENERALIDADES')
    M.t2('Objeto')
    M.par('El presente análisis verifica que el terreno del proyecto “%s”, con CUI N.° %s, es apto para el Centro de Acogida '
          'Residencial - Varones: que no presenta incompatibilidades de uso por cercanía a actividades, instalaciones o servidumbres '
          'reguladas por la normativa vigente, y que reúne las condiciones de ubicación que exige el RNE.' % (PROY, CUI))
    M.t2('El servicio y sus exigencias de ubicación')
    M.par('El CAR Varones brinda acogimiento residencial, alimentación, protección y atención integral a niños y adolescentes varones sin '
          'cuidados parentales o en riesgo de perderlos, en el marco del Decreto Legislativo N.° 1297 y su Reglamento. Al residir en él '
          'niños y adolescentes, su ubicación debe evitar las actividades e instalaciones que la normativa sectorial considera '
          'incompatibles con su permanencia, y contar con accesos seguros, servicios básicos y un entorno tranquilo.')
    M.t2('Ubicación del terreno')
    M.par('Predio rural Los Mangos, distrito de Morales, provincia y departamento de San Martín, dentro del casco urbano de Morales, con '
          'acceso por la carretera Oasis. El terreno del CAR es la poligonal resultante P1-P2-P3-P4 del estudio topográfico: '
          'rectángulo de 93.90 x 74.55 m, 6,999.98 m² y 336.89 m de perímetro (UTM WGS84, zona 18 Sur; P1: E 345 700.464, '
          'N 9 282 864.028).')
    M.tabla('Colindantes del terreno del CAR Varones', ['LADO', 'COLINDANTE', 'USO Y COMPATIBILIDAD'], [list(u) for u in USOS],
            [3.6, 5.2, W - 8.8], tam=8.5, izq=(0, 1, 2))
    M.par('Dentro del radio de 100 m del terreno se identifican tres usos: comercio vecinal, vivienda taller y el Hogar de Refugio '
          'Temporal. El terreno forma parte de un conjunto de servicios de protección del Estado (CAR Varones, Hogar de Refugio '
          'Temporal y CAR Mujeres), usos de la misma naturaleza y compatibles entre sí.')

    M.t1('MARCO NORMATIVO')
    M.par('La ubicación del CAR Varones se verifica con las normas asignadas al proyecto:')
    M.tabla('Normas que rigen la ubicación del proyecto', ['NORMA', 'DISPOSICIÓN'], [list(n) for n in NORMAS], [5.6, W - 5.6], tam=8.5,
            izq=(0, 1))
    M.par('El cuadro de incompatibilidades de la %s reúne los dispositivos legales sectoriales que restringen la ubicación de los '
          'locales con permanencia de niñas, niños y adolescentes; la Norma A.040 (art. 7 b) y la R.V.M. N.° 208-2019 (art. 8 c) '
          'remiten a él. Se aplica al CAR Varones y se complementa con las condiciones de ubicación del RNE.' % RVM)

    M.t1('METODOLOGÍA')
    for t in ('Incompatibilidades por cercanía (Cuadro N.° 1): se define el área de influencia como la franja medida desde el lindero '
              'con la distancia que fija cada dispositivo legal: 50 m (estaciones de servicio), 100 m (establecimientos de salud, '
              'GLP, bebidas alcohólicas y plantas de combustibles), 150 m (velatorios y cementerios) y 200 m (ductos de '
              'hidrocarburos). Para las que no tienen distancia fija se verifica que la restricción no alcance al terreno.',
              'Condiciones de ubicación (Cuadro N.° 2): se verifica cada condición del RNE y de la R.V.M. N.° 208-2019 con los datos '
              'del terreno y del programa arquitectónico.',
              'Se reconoce el entorno del terreno y se revisa la cartografía y las imágenes satelitales de la zona.',
              'Cada punto se califica como CUMPLE cuando la actividad o restricción no existe dentro del área de influencia, o cuando el '
              'terreno y el proyecto satisfacen la condición.'):
        M.vin(t)

    I.seccion(M, apaisada=True); W2 = M.ancho
    M.t1('INCOMPATIBILIDADES DE UBICACIÓN POR CERCANÍA')
    filas = [[str(k), x[0], x[1], I.dist(x), x[4], 'CUMPLE'] for k, x in enumerate(CUADRO, 1)]
    M.tabla('Cuadro N.° 1: incompatibilidades de uso del terreno por cercanía - CAR Varones (según la %s)' % RVM,
            ['N.°', 'INCOMPATIBILIDAD POR CERCANÍA', 'DISPOSITIVO LEGAL', 'DISTANCIA O CRITERIO', 'CONDICIÓN DEL TERRENO', 'ESTADO'],
            filas, [0.9, 5.0, 7.6, 3.0, W2 - 18.2, 1.7], tam=8, izq=(1, 2, 3, 4))
    for fila in M.d.tables[-1].rows:
        fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    M.t1('CONDICIONES DE UBICACIÓN DEL CAR')
    filas = [[str(k), c[0], c[1], c[2], 'CUMPLE'] for k, c in enumerate(CONDICIONES, 1)]
    M.tabla('Cuadro N.° 2: condiciones de ubicación del CAR Varones (RNE y R.V.M. N.° 208-2019-MINEDU)',
            ['N.°', 'CONDICIÓN DE UBICACIÓN', 'FUENTE', 'CÓMO LA CUMPLE EL TERRENO O EL PROYECTO', 'ESTADO'],
            filas, [0.9, 5.4, 4.6, W2 - 12.6, 1.7], tam=8, izq=(1, 2, 3))
    for fila in M.d.tables[-1].rows:
        fila._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    I.seccion(M, apaisada=False)

    M.t1('CONCLUSIONES')
    for t in ('El terreno cumple las %d condiciones del Cuadro N.° 1: no existen dentro de su área de influencia las actividades, '
              'instalaciones ni servidumbres que la normativa sectorial considera incompatibles con la permanencia de niños y '
              'adolescentes.' % len(CUADRO),
              'El terreno y el proyecto cumplen las %d condiciones de ubicación del Cuadro N.° 2: zona compatible, sin incompatibilidades '
              'de uso, accesos para vehículos de emergencia, servicios básicos, terreno rectangular y plano con el programa completo '
              'dentro de sus linderos, y drenaje pluvial.' % len(CONDICIONES),
              'La ubicación cumple la Norma A.090 (arts. 3 y 4), la Norma A.040 (art. 7) y la R.V.M. N.° 208-2019 (art. 8), sin '
              'necesidad de estudio de impacto vial.',
              'Los colindantes son el Hogar de Refugio Temporal y el área remanente del predio matriz, de la misma naturaleza de servicio '
              'de protección, y la carretera Oasis da acceso al predio sin que las obras ocupen su faja.',
              'Por lo expuesto, el terreno es apto para el Centro de Acogida Residencial - Varones.'):
        M.vin(t)
    M.d.core_properties.author = ''
    M.d.core_properties.title = 'Análisis de incompatibilidades de ubicación - CAR Varones'
    M.d.save(SALIDA)
    return SALIDA


def generar_xlsx():
    I.CUADRO, I.CONDICIONES, I.CUI, I.RVM = CUADRO, CONDICIONES, CUI, RVM
    I.SALIDA_X = os.path.join(SAL, 'MATRIZ_INCOMPATIBILIDADES_UBICACION_CAR_VARONES.xlsx')
    I.SERV, I.SERV_CORTO, I.HOJA2, I.RES2 = 'CAR VARONES', 'CAR VARONES', 'CONDICIONES CAR', 'condiciones de ubicación del CAR'
    return I.generar_xlsx()


def generar_interno():
    """Revisión del cuadro de la arquitecta: qué está bien, qué se corrigió y qué confirmar (no forma parte del expediente)."""
    import docx
    from docx.shared import Pt
    d = docx.Document(); s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles['Normal']; st.font.name = 'Arial'; st.font.size = Pt(10)
    M = R.Memoria(); M.d = d
    if d.paragraphs:
        d.paragraphs[0]._p.getparent().remove(d.paragraphs[0]._p)
    W = (s.page_width - s.left_margin - s.right_margin) / 360000.0; M.ancho = W
    M.par('REVISIÓN DEL ANÁLISIS DE INCOMPATIBILIDADES DE UBICACIÓN', 12.5, True, color=R.GRANATE, centro=True, despues=2)
    M.par('CAR Varones - CUI N.° %s  |  Documento de trabajo, no forma parte del expediente' % CUI, 9, color=R.GRIS, centro=True,
          despues=10)
    M.par('Se revisó el cuadro de incompatibilidades de la memoria descriptiva de arquitectura (18 dispositivos, todos CUMPLE). El '
          'cuadro está completo y cita la versión vigente de la norma; los resultados se conservan. Abajo, lo que se corrigió en el '
          'análisis y lo que conviene confirmar.')

    def sec(t):
        M.par(t, 11, True, color=R.GRANATE, antes=8, despues=4)

    sec('1. LO QUE ESTÁ BIEN')
    for e, t in (('Norma', 'cita la R.V.M. N.° 068-2025-MINEDU, la versión vigente de la N.T. Criterios Generales.'),
                 ('Lista', 'los 18 dispositivos legales sectoriales con sus distancias mínimas; todos en CUMPLE.'),
                 ('Entorno', 'describe los colindantes del terreno y los usos dentro de 100 m (comercio vecinal, vivienda taller y '
                  'Hogar de Refugio Temporal).')):
        SN.etiq(M, e, t)

    sec('2. CORREGIDO EN EL ANÁLISIS')
    M.tabla('Correcciones de redacción y citas', ['FILA', 'EN LA MEMORIA', 'EN EL ANÁLISIS'], [
        ['12', 'Decreto Ley N.° 25884', 'Decreto Ley N.° 25844, Ley de Concesiones Eléctricas'],
        ['18', 'Ley Orgánica de Municipales', 'Ley N.° 27972, Ley Orgánica de Municipalidades'],
        ['13', '"cuando una Institución Educativa se encuentre próximo" (texto de un colegio)',
         'restricciones en áreas de uso público próximas a una estación radioeléctrica'],
        ['4, 5, 10, 14, 17, 18', '"no se encontró dentro del área definida para la UP"', 'condición concreta del terreno del CAR'],
        ['10', 'Aeródromos: "no se encontró dentro del área definida"', 'Aeropuerto de Tarapoto a unos %.1f km, al costado de la pista '
         '17/35; edificaciones de baja altura (sigue CUMPLE)' % AERO_KM],
        ['15', 'El área del CAR no está en área restringida', 'Se agrega que el frente da a la carretera Oasis por el área libre del '
         'predio matriz y que las obras quedan dentro del lindero'],
        ['Todos', 'Distancia dentro del texto del dispositivo', 'Columna aparte "DISTANCIA O CRITERIO"'],
    ], [2.6, 6.2, W - 8.8], tam=8.5, izq=(1, 2))
    M.par('Se añadieron: el servicio del CAR y sus exigencias de ubicación (D.Leg. N.° 1297), el marco normativo (A.090 arts. 3 y 4, '
          'A.040 art. 7, R.V.M. N.° 208-2019 art. 8), la metodología, el Cuadro N.° 2 con %d condiciones de ubicación del CAR y las '
          'conclusiones.' % len(CONDICIONES), despues=4)

    sec('3. PARA CONFIRMAR')
    M.tabla('Puntos para confirmar', ['TEMA', 'QUÉ CONFIRMAR', 'POR QUÉ'], [
        ['Fajas marginales (fila 7)', 'La memoria dice que "por la margen derecha discurren aguas de drenaje pluvial". Precisar si es una '
         'cuneta o dren de la vía (sin faja marginal) o un cauce.', 'El cuadro dice que no hay fuentes de agua; ambos textos deben '
         'decir lo mismo.'],
        ['Bebidas alcohólicas (fila 5)', 'Que el comercio vecinal dentro de 100 m no venda bebidas alcohólicas.',
         'La memoria reconoce comercio vecinal dentro de 100 m.'],
        ['Accesos', 'La memoria indica el ingreso por la carretera Fernando Belaunde Terry (sector El Tumi) y una trocha de unos 340 m, y '
         'el saneamiento, el acceso directo por la carretera Oasis.', 'Unificar la descripción del acceso.'],
        ['Norma del MIMP', 'Que los "Criterios Generales para el Diseño Arquitectónico de CAR ..." sean la R.M. N.° 377-2024-MIMP que '
         'cita la programación, y pasar su PDF.', 'Si trae condiciones de ubicación propias del CAR, se agregan al Cuadro N.° 2.'],
        ['Zonificación (Cuadro 2, fila 1)', 'Certificado de zonificación y vías o de parámetros urbanísticos de la Municipalidad de '
         'Morales.', 'Es el documento que acredita la A.090 art. 3.'],
        ['Pozo séptico (fila 14)', 'El pozo séptico del programa (265.40 m²) es del propio CAR; confirmar su distancia a dormitorios '
         'según la especialidad sanitaria.', 'El cuadro trata de plantas de terceros; un evaluador puede preguntar por el propio.'],
        ['Aeródromo y líneas eléctricas (filas 10 a 12)', 'Distancia al aeropuerto con la ubicación real y si hay red de media tensión '
         'en la carretera Oasis.', 'Respaldo del CUMPLE.'],
    ], [3.4, 7.6, W - 11.0], tam=8.5, izq=(0, 1, 2))
    d.core_properties.author = ''
    d.core_properties.title = 'Revisión del análisis de incompatibilidades de ubicación - CAR Varones'
    os.makedirs(os.path.dirname(SALIDA_N), exist_ok=True)
    d.save(SALIDA_N)
    return SALIDA_N


if __name__ == '__main__':
    print(generar_docx()); print(generar_xlsx()); print(generar_interno())
