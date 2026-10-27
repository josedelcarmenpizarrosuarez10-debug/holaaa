# -*- coding: utf-8 -*-
"""Metrado de equipamiento y mobiliario - CAR Varones (CUI 2705619), en Excel con fórmulas y verificación.

Hojas:
  RESUMEN               partidas 03.01 (equipos, mobiliario, implementos por ambiente) con el metrado enlazado por SUMIF;
  METRADO POR AMBIENTE  cada ítem de cada cuadro de equipamiento del plano (láminas MB-01 a MB-09), con la cantidad del cuadro
                        y la adoptada (corregida donde el cuadro contradice lo dibujado, con su nota);
  VERIFICACIÓN          por código: metrado original del usuario, suma de cuadros, etiquetas en planta, adoptado y observación;
  VERIF. IMPLEMENTOS    implementos del metrado original frente a los cuadros del plano;
  OBSERVACIONES         errores encontrados en el plano y en el metrado original.

Uso: python3 herramientas/varones/mobiliario/metrado_mb.py
Salida: entregables_varones/METRADO_EQUIPAMIENTO_Y_MOBILIARIO_CAR_VARONES.xlsx
"""
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import datos_mb as D  # noqa: E402

SALIDA = os.path.join(D.SAL, 'METRADO_EQUIPAMIENTO_Y_MOBILIARIO_CAR_VARONES.xlsx')
GRANATE, ROSA, GRIS, AMARILLO = '7B1E2B', 'F2E6E8', 'F2F2F2', 'FFF2CC'
FINO = Side(style='thin', color='808080'); BORDE = Border(left=FINO, right=FINO, top=FINO, bottom=FINO)
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
IZQ = Alignment(horizontal='left', vertical='center', wrap_text=True)
LAMINA = {'E-01': 'MB-01', 'E-02': 'MB-02', 'E-10': 'MB-02', 'E-03': 'MB-03', 'E-04': 'MB-04', 'E-05': 'MB-05', 'E-06': 'MB-06',
          'E-07': 'MB-07', 'E-08': 'MB-08', 'E-09': 'MB-09'}


def celda(ws, r, c, v, neg=False, fondo=None, color=None, al=CENTRO, fmt=None, borde=True, size=9):
    x = ws.cell(r, c, v); x.font = Font(name='Arial', size=size, bold=neg, color=color); x.alignment = al
    if fondo: x.fill = PatternFill('solid', fgColor=fondo)
    if borde: x.border = BORDE
    if fmt: x.number_format = fmt
    return x


def cabecera(ws, titulo, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    celda(ws, 1, 1, titulo, True, GRANATE, 'FFFFFF', size=12, borde=False); ws.row_dimensions[1].height = 24
    filas = [('PROYECTO', '“%s”' % D.PROYECTO), ('CUI', 'N.° ' + D.CUI), ('PROPIETARIO', 'Gobierno Regional San Martín - Gerencia Territorial Bajo Mayo'),
             ('UBICACIÓN', 'Morales / San Martín / San Martín'), ('ESPECIALIDAD', D.ESPECIALIDAD),
             ('LÁMINAS', 'MB-01 a MB-09 (equipamiento y mobiliario, escala 1/150)'), ('FECHA', D.FECHA)]
    for k, (a, b) in enumerate(filas, 2):
        celda(ws, k, 1, a, True, GRIS, al=IZQ)
        ws.merge_cells(start_row=k, start_column=2, end_row=k, end_column=ncol); celda(ws, k, 2, b, k == 2, al=IZQ)
        for c in range(3, ncol + 1): ws.cell(k, c).border = BORDE
    ws.row_dimensions[2].height = 36
    return len(filas) + 3


def impresion(ws, ncol, nfil, horizontal=False):
    ws.page_setup.orientation = 'landscape' if horizontal else 'portrait'; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_area = 'A1:%s%d' % (L(ncol), nfil); ws.sheet_view.showGridLines = False
    ws.page_margins.left = ws.page_margins.right = 0.4; ws.page_margins.top = ws.page_margins.bottom = 0.5


def hoja_ambientes(wb, A):
    ws = wb.create_sheet('METRADO POR AMBIENTE')
    r = cabecera(ws, 'METRADO DE EQUIPAMIENTO Y MOBILIARIO POR AMBIENTE (CUADROS DE LAS LÁMINAS MB-01 A MB-09)', 10)
    for k, v in enumerate(['EDIF.', 'AMBIENTE', 'LÁMINA', 'CÓD.', 'DESCRIPCIÓN', 'UND', 'CANT. CUADRO', 'CANT. ADOPTADA', 'CLAVE', 'NOTA'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1; r0 = r
    for ed, amb, items in A:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        celda(ws, r, 1, '%s  %s - %s' % (ed, D.EDIFICACIONES.get(ed, ''), amb), True, GRIS, al=IZQ)
        for c in range(2, 11): ws.cell(r, c).border = BORDE
        r += 1
        for it in items:
            if it['grupo'] == 'S':
                continue
            nota = bool(it['nota'])
            for k, v in enumerate([ed, amb, LAMINA.get(ed, ''), it['cod'] if it['grupo'] in 'EM' else '', it['des'], 'und', it['cuadro'],
                                   it['cant'], it['clave'], it['nota']], 1):
                celda(ws, r, k, v, al=IZQ if k in (2, 5, 9, 10) else CENTRO, size=8 if k in (9, 10) else 9,
                      fondo=AMARILLO if nota and k in (7, 8, 10) else None, fmt='0' if k in (7, 8) else None)
            r += 1
    for c, w in zip(range(1, 11), [6, 26, 7, 6, 52, 5, 8, 9, 22, 48]):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = ws.cell(r0, 1)
    impresion(ws, 10, r - 1, True)
    return "'METRADO POR AMBIENTE'!$I$%d:$I$%d" % (r0, r - 1), "'METRADO POR AMBIENTE'!$H$%d:$H$%d" % (r0, r - 1)


def hoja_resumen(ws, cat, rango_clave, rango_cant):
    r = cabecera(ws, 'HOJA RESUMEN DE METRADOS - ARQUITECTURA: EQUIPAMIENTO Y MOBILIARIO', 5)
    for k, v in enumerate(['PART. N.°', 'CÓD.', 'DESCRIPCIÓN', 'UND', 'METRADO'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    for it, de in (('03', 'ARQUITECTURA'), ('03.01', 'MOBILIARIO Y EQUIPAMIENTO')):
        celda(ws, r, 1, it, True, GRIS); celda(ws, r, 2, None, fondo=GRIS); celda(ws, r, 3, de, True, GRIS, al=IZQ)
        celda(ws, r, 4, None, fondo=GRIS); celda(ws, r, 5, None, fondo=GRIS); r += 1
    fila = {}
    for f in D.numerar(cat)[2:]:
        if f['tipo'] in ('grupo', 'sub'):
            celda(ws, r, 1, f['item'], True); celda(ws, r, 2, None); celda(ws, r, 3, f['desc'], True, al=IZQ); celda(ws, r, 4, None); celda(ws, r, 5, None)
        else:
            celda(ws, r, 1, f['item']); celda(ws, r, 2, f['cod']); celda(ws, r, 3, f['desc'], al=IZQ); celda(ws, r, 4, 'und')
            ws.cell(r, 6, f['clave']).font = Font(name='Arial', size=7, color='FFFFFF')
            celda(ws, r, 5, '=SUMIF(%s,F%d,%s)' % (rango_clave, r, rango_cant), True, fmt='0.00')
            fila[f['clave']] = r
        r += 1
    for c, w in zip(range(1, 7), [15, 6, 70, 6, 11, 4]):
        ws.column_dimensions[L(c)].width = w
    ws.column_dimensions['F'].hidden = True
    impresion(ws, 5, r - 1)
    return fila


def hoja_verificacion(wb, fila):
    ws = wb.create_sheet('VERIFICACIÓN')
    r = cabecera(ws, 'VERIFICACIÓN DEL METRADO ORIGINAL CON EL PLANO (CUADROS DE EQUIPAMIENTO Y ETIQUETAS DE LA PLANTA)', 9)
    for k, v in enumerate(['CÓD.', 'DESCRIPCIÓN', 'METRADO ORIGINAL', 'SUMA DE CUADROS', 'ETIQUETAS EN PLANTA', 'ADOPTADO', 'DIFERENCIA (ADOPTADO - ORIGINAL)',
                           'ESTADO', 'OBSERVACIÓN'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    ver, _, _ = D.verificacion()
    for v in ver:
        adop = "='RESUMEN'!E%d" % fila[v['clave']] if v['clave'] in fila else 0
        mal = v['original'] != v['adoptado']
        for k, x in enumerate([v['clave'], v['desc'], v['original'] if v['original'] is not None else '—', v['cuadros'], v['planta'], adop,
                               '=IF(ISNUMBER(C{0}),F{0}-C{0},F{0})'.format(r), '=IF(G%d=0,"conforme","corregir")' % r, v['obs']], 1):
            celda(ws, r, k, x, al=IZQ if k in (2, 9) else CENTRO, size=8 if k == 9 else 9, fmt='0' if k in (3, 4, 5, 6, 7) else None,
                  fondo=AMARILLO if mal and k in (3, 6, 7, 8) else None)
        r += 1
    for c, w in zip(range(1, 10), [6, 52, 10, 10, 10, 10, 12, 10, 70]):
        ws.column_dimensions[L(c)].width = w
    ws.freeze_panes = ws.cell(11, 3)
    impresion(ws, 9, r - 1, True)


def hoja_implementos(wb):
    ws = wb.create_sheet('VERIF. IMPLEMENTOS')
    r = cabecera(ws, 'VERIFICACIÓN DE IMPLEMENTOS: METRADO ORIGINAL FRENTE A LOS CUADROS DEL PLANO', 5)
    for k, v in enumerate(['IMPLEMENTO', 'METRADO ORIGINAL', 'CUADROS DEL PLANO (ADOPTADO)', 'ESTADO', 'OBSERVACIÓN'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    for v in D.verificacion_implementos():
        for k, x in enumerate([v['desc'], v['original'] if v['original'] is not None else '—', v['plano'],
                               'conforme' if not v['obs'] else 'corregir', v['obs']], 1):
            celda(ws, r, k, x, al=IZQ if k in (1, 5) else CENTRO, fondo=AMARILLO if v['obs'] and k in (2, 3, 4) else None)
        r += 1
    for c, w in zip(range(1, 6), [52, 12, 14, 10, 52]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 5, r - 1)


def hoja_observaciones(wb, hojas):
    ws = wb.create_sheet('OBSERVACIONES')
    r = cabecera(ws, 'OBSERVACIONES AL PLANO Y AL METRADO ORIGINAL', 3)
    for k, v in enumerate(['N.°', 'DOCUMENTO', 'OBSERVACIÓN'], 1):
        celda(ws, r, k, v, True, GRANATE, 'FFFFFF')
    r += 1
    obs = [('Lámina ' + a, b) for a, b in D.OBS_PLANO]
    obs += [('Metrado original', 'La hoja "Hoja1" es el cuadro de equipamiento de otro proyecto (institución educativa de cuatro niveles con aulas '
                                 '01 a 55, laboratorios y talleres); no corresponde al CAR Varones y no se usa.' if 'Hoja1' in hojas else ''),
            ('Metrado original', 'Las cantidades de la hoja HR están escritas a mano, sin vínculo con los cuadros del plano.'),
            ('Metrado original', 'Numeración: EQUIPOS y MOBILIARIO llevan el mismo código 03.01.02; se repite 03.01.02.04; M54 y M55 no tienen número; '
                                 'los ítems de la sala de psicomotricidad muestran #REF!; falta el 03.01.02.08 de equipos.'),
            ('Metrado original', 'Faltan las sillas regulables M03 (88), las mesas regulables M04 (16), los sillones M32 (7), los maceteros M14 (2) y '
                                 'las camillas M65 (2); M63 figura como camilla (6) y en el plano es el sofá de 3 cuerpos (4).'),
            ('Metrado original', 'Los implementos deportivos del original (pelotas de básquet #5 y #7 por 24, torres de básquet, futsal, ajedrez, '
                                 'etc.) no son los del cuadro del almacén deportivo de E-02; se adoptan los del plano.'),
            ('Metrado original', 'Los implementos de la caseta de vigilancia (linterna, megáfono, chaleco) no tienen cuadro en el plano; si se '
                                 'requieren, deben agregarse al plano.')]
    for k, (a, b) in enumerate([o for o in obs if o[1]], 1):
        celda(ws, r, 1, k); celda(ws, r, 2, a, al=IZQ); celda(ws, r, 3, b, al=IZQ); r += 1
    for c, w in zip(range(1, 4), [5, 18, 100]):
        ws.column_dimensions[L(c)].width = w
    impresion(ws, 3, r - 1)


def generar():
    P = D.plano(); A = D.ambientes(P); cat = D.catalogo(P)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'RESUMEN'
    rc, rq = hoja_ambientes(wb, A)
    fila = hoja_resumen(ws, cat, rc, rq)
    hoja_verificacion(wb, fila); hoja_implementos(wb)
    hoja_observaciones(wb, D.original()[2])
    wb.properties.creator = ''; wb.properties.title = 'Metrado de equipamiento y mobiliario - CAR Varones'
    wb.save(SALIDA)
    return SALIDA


if __name__ == '__main__':
    print(generar())
