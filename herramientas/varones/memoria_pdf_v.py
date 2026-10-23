"""CAR Varones (CUI 2705619): memoria de calculo del colector pluvial frontal en PDF, con el mismo estilo visual de la
memoria HIDRO-CE040 del proyecto y de las memorias del colector del Hogar de Refugio y del CAR Mujeres.

Fuente: entregables_varones/MEMORIA_CALCULO_COLECTOR_CAR_VARONES.xlsx (hojas MEMORIA, DATOS, CAUDALES, CUNETAS, EMPALME_CL,
PERFIL_FLUJO, ESTRUCTURAL, CUMPLIMIENTO). El libro se recalcula con LibreOffice en una copia temporal y se copian sus valores
sin cambiarlos; solo se da formato (tildes, superindices) y las referencias a hojas pasan a referencias a las secciones.
Uso: python3 herramientas/varones/memoria_pdf_v.py
"""
import os, sys, shutil, subprocess, tempfile
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__)); HERR = os.path.dirname(AQUI); RAIZ = os.path.dirname(HERR)
sys.path.insert(0, HERR)
import memoria_pdf_refugio as R  # noqa: E402  (mismo motor: tildes, perfil, tablas de parametros)
from memoria_pdf_refugio import MM, tabla, formulas, fmt, CSS, tilde, und, esc, perfil_svg, tabla_param  # noqa: E402

XLSX = os.path.join(RAIZ, "entregables_varones", "MEMORIA_CALCULO_COLECTOR_CAR_VARONES.xlsx")
SALIDA = os.path.join(RAIZ, "entregables_varones", "MEMORIA_CALCULO_COLECTOR_CAR_VARONES.pdf")
PROYECTO = ("“CREACIÓN DEL SERVICIO DE PROTECCIÓN INTEGRAL A NIÑAS, NIÑOS Y ADOLESCENTES SIN CUIDADOS PARENTALES O EN RIESGO "
            "DE PERDERLOS EN CENTRO DE ACOGIDA RESIDENCIAL - VARONES DISTRITO DE MORALES DE LA PROVINCIA DE SAN MARTÍN DEL "
            "DEPARTAMENTO DE SAN MARTÍN”, con CUI N.° 2705619 · SAN MARTÍN - SAN MARTÍN - MORALES · octubre 2026")

MM.TILDES.update({
    "perimetrico": "perimétrico", "alli": "allí", "caia": "caía", "mecanica": "mecánica", "proporcion": "proporción",
    "lamina": "lámina", "laminas": "láminas", "alfeizar": "alféizar", "asfaltico": "asfáltico", "practica": "práctica",
    "Coordinacion": "Coordinación", "practicable": "practicable", "transmitida": "transmitida", "mas": "más",
})
MM._RX = R.re.compile(r"\b(" + "|".join(sorted(map(R.re.escape, MM.TILDES), key=len, reverse=True)) + r")\b")
# frases propias primero (van antes que las generales del Refugio, que reemplazan "hoja PERFIL_FLUJO")
MM.FRASES = {
    "Memoria del Hogar de Refugio, hoja PERFIL_FLUJO (ultima fila)": "Memoria del Hogar de Refugio, perfil de flujo (última estación)",
    "hoja CUNETAS": "sección 3", "ver tau": "ver τ", "(hoja PERFIL_FLUJO).": "(sección 5).", **MM.FRASES}


def libro():
    """Copia temporal del libro recalculada con LibreOffice (el entregable no se toca)."""
    tmp = tempfile.mkdtemp()
    shutil.copy(XLSX, os.path.join(tmp, "m.xlsx"))
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "c"), os.path.join(tmp, "m.xlsx")],
                   check=True, capture_output=True)
    wb = openpyxl.load_workbook(os.path.join(tmp, "c", "m.xlsx"), data_only=True)
    shutil.rmtree(tmp)
    return wb


def sub(t):
    return t.replace("&lt;sub&gt;", "<sub>").replace("&lt;/sub&gt;", "</sub>")


def construir():
    wb = libro()
    M, DA, CA, CU, EM, PF, ES, CP = (wb[k] for k in ("MEMORIA", "DATOS", "CAUDALES", "CUNETAS", "EMPALME_CL", "PERFIL_FLUJO", "ESTRUCTURAL", "CUMPLIMIENTO"))
    v = lambda ws, ref: ws[ref].value
    out = []
    # ---------------------------------------------------------- 0. encabezado y memoria descriptiva
    out.append('<h1>Memoria de cálculo del colector pluvial frontal</h1><div class="proy">%s</div>' % esc(PROYECTO))
    out.append('<div class="meta"><b>TRAMO CAR VARONES</b> · 0+000 a 0+%06.2f · entrega a la caja de llegada CL del Hogar de Refugio Temporal '
               'Mujeres Violentadas (CUI 2675514) · RNE CE.040</div>' % v(DA, "B23"))
    out.append('<div class="kpis">%s</div>' % "".join(
        '<div class="kpi"><div class="v">%s</div><div class="l">%s</div></div>' % (a, b) for a, b in (
            ("%.1f L/s" % v(M, "B12"), "Caudal de diseño del colector"), ("%.3f" % v(M, "B17"), "Llenado máximo y/h (admisible 0.85)"),
            ("%.3f m" % v(M, "B18"), "Borde libre mínimo bajo losa (mín. 0.05)"), ("%.3f" % v(M, "B19"), "Froude máximo (régimen subcrítico)"))))
    out.append("<h2>Datos y metodología</h2>")
    for r in range(4, 11):
        out.append("<p><b>%s.</b> %s</p>" % (esc(v(M, "A%d" % r).split(". ", 1)[1].capitalize()), esc(v(M, "B%d" % r))))
    res = []
    for r in range(12, 26):
        b = v(M, "B%d" % r)
        if isinstance(b, float):
            b = "%.1f" % b if b > 100 else ("%.2f" % b if r in (14, 15) else "%.3f" % b)
        res.append([tilde(v(M, "A%d" % r)), b if isinstance(b, str) else fmt(b)])
    out.append('<div class="final"><h3>Resultados de la verificación</h3>%s</div>' % tabla(["Resultado", "Valor"], res, anchos=["68%", "32%"], alinear={1: "r"}))
    out.append('<p class="ref"><b>Notas.</b></p>' + "".join('<p class="ref">%s</p>' % esc(v(M, "B%d" % r)) for r in range(28, 31)))
    # ---------------------------------------------------------- 1. datos de diseno
    out.append('<h2><span class="n">1.</span>Datos de diseño</h2>')
    out.append(formulas(["Q = Σ(C·A) · I · FS / 3 600 000 &nbsp;[m<sup>3</sup>/s] &nbsp;(I en mm/h, A en m<sup>2</sup>); &nbsp;C ponderado = Σ(C·A) / ΣA",
                         "CF<sub>final</sub> = CF<sub>0</sub> − S · L; &nbsp;h = cota bajo losa − cota de fondo"]))
    out.append(tabla_param(DA, 5, 38, anchos=["33%", "11%", "8%", "40%", "8%"]))
    # ---------------------------------------------------------- 2. caudales
    out.append('<h2><span class="n">2.</span>Caudales de diseño por cuneta y caudal en el colector</h2>')
    out.append("<p>%s</p>" % esc(v(CA, "A3")))
    out.append(formulas(["Q<sub>cuneta</sub> = Q<sub>total</sub> · franja de frente / Σ franjas; &nbsp;Q<sub>colector</sub>(x) = Σ Q<sub>cuneta</sub> aguas arriba de x"]))
    rows = []
    for a, b, c, d, e, f, g, h in MM.filas(CA, 6, 11, 8):
        rows.append([tilde(a), "%.2f" % b, "%.4f" % c, "%.1f" % d, "0+%06.2f" % e if e else "", tilde(f or ""), tilde(g or ""), tilde(h or "")])
    out.append(tabla(["Cuneta", "Franja de frente (m)", "Fracción", "Q (L/s)", "Progresiva", "Entra en", "Registro / ventana", "Observación"], rows,
                     anchos=["13%", "10%", "9%", "8%", "10%", "10%", "24%", "16%"], alinear={1: "r", 2: "r", 3: "r", 4: "r"}))
    out.append("<p><b>%s:</b> %s</p>" % (esc(v(CA, "A12")), esc(v(CA, "B12"))))
    out.append("<h3>%s</h3>" % esc(v(CA, "A14")))
    rows = [[tilde(a), "0+%06.2f" % b, "%.1f" % c if c is not None else "", "%.1f" % d] for a, b, c, d in MM.filas(CA, 16, 21, 4)]
    out.append(tabla(["Punto", "Progresiva", "Aporte (L/s)", "Q acumulado (L/s)"], rows, anchos=["46%", "16%", "19%", "19%"], alinear={1: "r", 2: "r", 3: "r"}))
    rows = [[tilde(v(CA, "A%d" % r)), "%.4f" % v(CA, "D%d" % r) if r == 22 else "%.1f" % v(CA, "D%d" % r), "m³/s" if r == 22 else "L/s"] for r in (22, 23, 24)]
    out.append(tabla(["Resultado", "Caudal", "Und"], rows, anchos=["60%", "22%", "18%"], alinear={1: "r"}))
    out.append('<p class="small">%s</p>' % esc(v(CA, "A26")))
    # ---------------------------------------------------------- 3. cunetas
    out.append('<h2><span class="n">3.</span>Cunetas de arquitectura que llegan al frente: acortamiento y caída libre</h2>')
    out.append("<p>%s</p>" % esc(v(CU, "A3")))
    out.append(formulas(["NCF<sub>muro</sub> = NCF<sub>extremo</sub> + 0.005 · acortamiento; &nbsp;caída libre = NCF<sub>muro</sub> − NA<sub>receptor</sub> &gt; 0"]))
    rows = []
    for r in MM.filas(CU, 6, 10, 12):
        rows.append([tilde(r[0]), "%.2f" % r[1], "%.2f" % r[2], "%.2f" % r[3], "%.2f" % r[4], "%.2f" % r[5], "%.2f" % r[6], "%.3f" % r[7],
                     tilde(r[8]), "%.3f" % r[9], "%.3f" % r[10], tilde(r[11])])
    out.append(tabla(["Cuneta", "Prog. (m)", "Long. (m)", "NCT", "NCF extremo", "H (m)", "Acort. (m)", "NCF muro", "Entra en",
                      "NA receptor", "Caída libre (m)", "Verificación"], rows, cls="perf cun"))
    out.append('<p class="small">%s</p>' % esc(v(CU, "A12")))
    # ---------------------------------------------------------- 4. empalme con la CL
    out.append('<h2><span class="n">4.</span>Empalme con la caja de llegada CL del colector del Hogar de Refugio</h2>')
    out.append(formulas(["y<sub>c</sub> = (Q<sup>2</sup> / (g · b<sup>2</sup>))<sup>1/3</sup>; &nbsp;NA<sub>final</sub> = CF<sub>final</sub> + y<sub>c</sub>; "
                         "&nbsp;caída libre = CF<sub>final</sub> − NA<sub>CL</sub> ≥ 0.10 m",
                         "colchón = NA<sub>CL</sub> − piso de la poza ≥ 0.50 m; &nbsp;V = Q / (b · y<sub>c</sub>)"]))
    for r0, r1 in ((5, 17), (19, 27), (29, 34), (36, 39), (41, 42)):
        out.append('<div class="bloque"><h3>%s</h3>%s</div>' % (esc(v(EM, "A%d" % r0)), tabla_param(
            EM, r0 + 1, r1, cab=("Parámetro", "Valor", "Und", "Fórmula / criterio"), ncol=4, anchos=["36%", "16%", "8%", "40%"])))
    # ---------------------------------------------------------- 5. perfil de flujo
    out.append('<h2><span class="n">5.</span>Perfil de flujo gradualmente variado con caudal creciente: método del paso estándar</h2>')
    out.append("<p>Cálculo desde la caída libre a la CL (0+%06.2f, tirante crítico) hacia aguas arriba; el caudal de cada estación es la suma de "
               "las cunetas que entran aguas arriba (sección 2). El tirante de cada estación se resuelve por bisección en 30 iteraciones. "
               "En el tramo 0+000 a 0+%06.2f, antes de la primera cuneta, no circula caudal de diseño.</p>" % (v(DA, "B23"), v(CA, "B16")))
    out.append(formulas(["y<sub>2</sub> + V<sub>2</sub><sup>2</sup>/2g + z<sub>2</sub> = y<sub>1</sub> + V<sub>1</sub><sup>2</sup>/2g + z<sub>1</sub> + (S<sub>f1</sub> + S<sub>f2</sub>)/2 · Δx",
                         "S<sub>f</sub> = (n · V / R<sup>2/3</sup>)<sup>2</sup>; &nbsp;Fr = V / √(g · y); &nbsp;llenado = y / h; &nbsp;BL = cota bajo losa − NA"]))
    pr = [[PF.cell(r, c).value for c in range(1, 16)] for r in range(6, 116)]
    pr = [r for r in pr if isinstance(r[0], (int, float))]
    out.append(perfil_svg(sorted([(r[0], r[1], r[6], r[9]) for r in pr], key=lambda p: (p[0], p[2]))))
    out.append('<div class="ley"><span><i style="border-color:#5b6670"></i>Cota de fondo</span><span><i style="border-color:#2b8cbf"></i>Nivel de agua</span>'
               '<span><i style="border-color:#9aa6ad;border-top-style:dashed"></i>Cota bajo losa (cara inferior de la losa superior)</span></div>')
    rows = [["%.2f" % r[0], "%.3f" % r[1], "%.4f" % r[2], "%.3f" % r[4], "%.3f" % r[5], "%.3f" % r[6], "%.3f" % r[7], "%.3f" % r[8],
             "%.3f" % r[9], "%.3f" % r[10], "%.3f" % r[11], "%.3f" % r[12], r[13] or "", tilde(r[14] or "")] for r in pr]
    out.append(sub(tabla(["Prog. (m)", "Cota de fondo", "Q (m³/s)", "y<sub>c</sub> (m)", "y (m)", "Nivel agua", "V (m/s)", "Fr", "Cota bajo losa",
                          "h (m)", "y/h", "BL (m)", "Verif.", "Zona"], rows, cls="perf")))
    out.append('<div class="final"><h3>Resumen del perfil</h3>%s</div>' % tabla(
        ["Resultado", "Valor"], [[tilde(v(PF, "A%d" % r)), ("%d" % v(PF, "B%d" % r)) if r == 121 else "%.3f" % v(PF, "B%d" % r)] for r in range(118, 123)],
        anchos=["70%", "30%"], alinear={1: "r"}))
    out.append('<div class="bloque"><h3>%s</h3>' % esc(v(PF, "A124")))
    out.append('<p class="small">%s</p>' % esc(v(PF, "A125")))
    rows = []
    for r in MM.filas(PF, 127, 133, 12):
        rows.append([tilde(r[0])] + ["%.4f" % r[1], "%.3f" % r[2], "%.3f" % r[3], "%.3f" % r[4], "%.3f" % r[5], "%.2f" % r[6], "%.3f" % r[7]] +
                    [tilde(x) for x in r[8:12]])
    out.append(sub(tabla(["Caso", "Q (m³/s)", "y<sub>n</sub> (m)", "V (m/s)", "R (m)", "τ (kg/m²)", "τ (Pa)", "Fr", "V ≥ 0.90 m/s", "τ ≥ 0.15 kg/m²",
                          "V ≤ 3.0 m/s", "Autolimpieza del tramo"], rows, cls="perf")))
    out.append(tabla(["Resultado", "Valor"], [[tilde(v(PF, "A%d" % r)), "%.3f" % v(PF, "B%d" % r) if isinstance(v(PF, "B%d" % r), float) else tilde(v(PF, "B%d" % r))]
                                             for r in (135, 136, 137)], anchos=["70%", "30%"], alinear={1: "r"}))
    out.append('<p class="small">%s</p></div>' % esc(v(PF, "A138")))
    # ---------------------------------------------------------- 6. estructural
    out.append('<h2><span class="n">6.</span>Verificación estructural</h2>')
    out.append("<p>Concreto armado según RNE E.060; cargas según RNE E.020 y rueda de diseño AASHTO LRFD HL-93 en el cruce de camiones cisterna.</p>")
    out.append(formulas(["w<sub>u</sub> = 1.4 D + 1.7 L; &nbsp;M<sub>u</sub> = w<sub>u</sub> L<sup>2</sup>/8 &nbsp;(losa simplemente apoyada, conservador)",
                         "M<sub>rueda</sub> = P (1 + IM) (L/4 − c/8) / E; &nbsp;a = A<sub>s</sub> f<sub>y</sub> / (0.85 f'<sub>c</sub> b); &nbsp;φM<sub>n</sub> = 0.9 A<sub>s</sub> f<sub>y</sub> (d − a/2) ≥ M<sub>u</sub>",
                         "Muro (marco cerrado, extremos empotrados): M = p H<sup>2</sup>/20 + q H<sup>2</sup>/12; &nbsp;A<sub>s,mín</sub> = 0.0018 · b · e &nbsp;(E.060 9.7.2); &nbsp;φV<sub>c</sub> ≥ V<sub>u</sub>"]))
    bloques, act = [], None
    for r in range(5, 84):
        a, b, c, d, e = (ES.cell(r, k).value for k in range(1, 6))
        if a is None: continue
        if b is None and c is None and d is None and e is None:
            act = [tilde(a), []]; bloques.append(act); continue
        act[1].append([tilde(a), b, und(c), tilde(d or ""), tilde(e or "")])
    for t, rows in bloques:
        cab = "Parámetros generales" if t.upper().startswith("PAR") else t
        out.append('<div class="bloque"><h3>%s</h3>%s</div>' % (esc(cab), tabla(["Elemento / parámetro", "Valor", "Und", "Fórmula / criterio", "Resultado"], rows,
                                                                            anchos=["32%", "11%", "8%", "35%", "14%"], alinear={1: "r"})))
    out.append('<p class="small">%s</p>' % esc(v(ES, "A85")))
    filas_res = []
    for t, rows in bloques[1:]:
        ok = [r for r in rows if isinstance(r[4], str) and r[4]]
        if ok:
            filas_res.append([t, "%d verificaci%s" % (len(ok), "ón" if len(ok) == 1 else "ones"),
                              "CUMPLE" if all(r[4].upper().startswith("CUMPLE") for r in ok) else "NO CUMPLE"])
    out.append('<div class="final"><h3>Cuadro conjunto de verificaciones estructurales</h3>%s</div>' % tabla(["Elemento", "Comprobaciones", "Resultado"], filas_res,
                                                                                                         anchos=["68%", "18%", "14%"]))
    # ---------------------------------------------------------- 7. cumplimiento
    out.append('<h2><span class="n">7.</span>Cuadro de cumplimiento normativo y condiciones de compatibilidad entre los tres expedientes</h2>')
    rows = []
    for r in range(5, 40):
        a, b, c, d, e = (CP.cell(r, k).value for k in range(1, 6))
        if a is None: continue
        if b is None and c is None: rows.append(("__grupo__", a)); continue
        if isinstance(d, float): d = "%.2f" % d if d > 100 else "%.3f" % d
        if isinstance(d, str) and d.upper().startswith(("CUMPLE", "RESALTO")) and d == e: d = "—"
        rows.append([tilde(a), tilde(b), tilde(c), tilde(d) if isinstance(d, str) else fmt(d), tilde(e)])
    out.append(tabla(["Requisito", "Norma / referencia", "Criterio", "Valor del diseño", "Resultado"], rows,
                     anchos=["24%", "21%", "27%", "12%", "16%"], alinear={3: "r"}))
    out.append('<p><b>%s: %d</b></p>' % (esc(v(CP, "A41")), v(CP, "B41")))
    out.append('<div class="final"><h3>%s</h3>%s</div>' % (esc(v(CP, "A43").capitalize()), "".join("<p>%s</p>" % esc(v(CP, "A%d" % r)) for r in range(44, 50))))
    # ---------------------------------------------------------- diseno adoptado y firma
    out.append('<div class="final"><h3>Diseño adoptado</h3><p>%s</p><p>%s</p><p>%s</p></div>' % (esc(v(M, "B7")), esc(v(M, "B8")), esc(v(M, "B9"))))
    out.append('<p class="ref">Referencia: RNE CE.040 Drenaje Pluvial (RM N.° 126-2021-VIVIENDA); RNE E.020 Cargas; RNE E.060 Concreto Armado; '
               'RNE E.050 Suelos y Cimentaciones; AASHTO LRFD para la carga de rueda en el cruce de camiones; Chow, Hidráulica de canales abiertos. '
               'Dimensiones sujetas a la aprobación del proyectista.</p>')
    out.append('<div class="firma">Profesional responsable: ______________________________ · CIP ____________ · Firma: ______________________</div>')
    return ("<!doctype html><html lang='es'><head><meta charset='utf-8'><title>Memoria de cálculo - colector pluvial CAR Varones</title>"
            "<style>%s .badge{white-space:nowrap} table.cun td:first-child, table.cun th:first-child{text-align:left}</style></head><body>%s</body></html>" % (CSS, "".join(out)))


PDF_JS = MM.PDF_JS.replace("CAR Mujeres (CUI 2717013)", "CAR Varones (CUI 2705619)").replace(
    "tramo CAR Mujeres 0+000 a 0+139.08", "tramo CAR Varones 0+000 a 0+105.29")


if __name__ == "__main__":
    doc = construir()
    tmp = os.path.join(RAIZ, "entregables_varones", "_tmp", "pdf_calculo"); os.makedirs(tmp, exist_ok=True)
    fh = os.path.join(tmp, "memoria_calculo_colector.html"); open(fh, "w", encoding="utf-8").write(doc)
    fj = os.path.join(tmp, "pdf.js"); open(fj, "w").write(PDF_JS)
    subprocess.run(["node", fj, fh, SALIDA], check=True, cwd=RAIZ)
    shutil.rmtree(tmp, ignore_errors=True)
    print(os.path.relpath(SALIDA, RAIZ))
