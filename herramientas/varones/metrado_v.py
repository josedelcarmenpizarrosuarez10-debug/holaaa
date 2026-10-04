"""Agrega la partida 01.04.04 COLECTOR PLUVIAL FRONTAL a la planilla de metrados del CAR Varones (la de cunetas,
entregables/METRADO_DRENAJE_PLUVIAL_CAR_VARONES_CUNETAS.xlsx), sin tocar lo existente, con el mismo formato y las
mismas hojas que la partida del Hogar de Refugio: COLECTOR PARAMETROS, MOV. TIERRAS, CONCRETO, ENCOFRADO, ACERO,
REGISTROS Y TAPAS, JUNTAS Y EMPALMES e INSUMOS. Sin cajas propias: la caja CL es del Hogar de Refugio.
"""
import os, sys, re, copy
import openpyxl
from openpyxl.styles import Alignment
AQUI = os.path.dirname(os.path.abspath(__file__)); HERR = os.path.dirname(AQUI); RAIZ = os.path.dirname(HERR)
for p in (HERR, AQUI):
    if p not in sys.path: sys.path.insert(0, p)
import compat as C
import diseno_v as dz
import metrado_calc as MC
import metrado_xlsx as MX
from metrado_xlsx import celda, cabecera, encabezado, fila, barra_total, subtitulo, nota, prog, Planilla, copiar_estilo, reinyectar_vml, NOMBRES, Q, PG

D = dz.D
ORIG = os.path.join(RAIZ, "entregables", "METRADO_DRENAJE_PLUVIAL_CAR_VARONES_CUNETAS.xlsx")
SALIDA = os.path.join(dz.SAL, "METRADO_DRENAJE_PLUVIAL_CAR_VARONES_CON_COLECTOR.xlsx")
ITEM = "01.04.04"

PARAMETROS = [
    ("b", "Ancho interior del colector b", D["b"], "m", "Lamina DP-04 (secciones); memoria de calculo, hoja DATOS"),
    ("em", "Espesor de muros", D["e_muro"], "m", "Lamina DP-04; memoria, hoja ESTRUCTURAL"),
    ("ef", "Espesor de losa de fondo", D["e_fondo"], "m", "Lamina DP-04"),
    ("et", "Espesor de losa superior (tramo normal)", D["e_losa"], "m", "Lamina DP-04 (secciones S-02 a S-08)"),
    ("etc", "Espesor de losa superior (cruce de camiones cisterna)", D["e_losa_camion"], "m", "Lamina DP-04 (seccion S-01); memoria, hoja ESTRUCTURAL (carga HL-93)"),
    ("es", "Espesor de solado", D["e_solado"], "m", "Lamina DP-04 y DA-02"),
    ("NPT", "Nivel de piso terminado del frente (cara superior de la losa)", D["NPT"], "msnm", "Lamina DP-02 (perfil) y plano de arquitectura (+261.15)"),
    ("sob", "Sobreancho de excavacion a cada lado del colector", MC.SOBREEXC, "m", "Lamina DA-01: zanja de 1.40 m (0.90 + 2 x 0.25)"),
    ("fr", "Franja de relleno de nivelacion del retiro, lado via", MC.FRANJA_RELLENO, "m", "Lamina DA-01, nota 2"),
    ("rec", "Recubrimiento del acero", MC.RECUB, "m", "Norma E.060, tabla 7.7.1 (concreto en contacto con el suelo)"),
    ("p38", "Peso de la barra de 3/8\"", MC.PESO["3/8"], "kg/m", "Igual a PARAMETROS!B10 de esta planilla"),
    ("p12", "Peso de la barra de 1/2\"", MC.PESO["1/2"], "kg/m", "Catalogo del fabricante (0.994 kg/m)"),
    ("tr", "Traslape de barras longitudinales", MC.TRASLAPE, "m", "Norma E.060, 12.15 (clase B, 3/8\"): 0.40 m"),
    ("lb", "Longitud comercial de barra", MC.L_BARRA, "m", "Igual a PARAMETROS!B11 de esta planilla"),
    ("gan", "Ganchos de cierre del marco (2 x 0.15)", MC.GANCHO, "m", "Lamina DP-08, cuadro de doblado"),
    ("esp", "Factor de esponjamiento", "=PARAMETROS!$B$16", "-", "Igual a PARAMETROS!B16 de esta planilla"),
    ("s_n", "Espaciamiento de marcos en tramo normal", 0.20, "m", "Lamina DP-04 y DP-08"),
    ("s_c", "Espaciamiento de marcos en el cruce de camiones", 0.15, "m", "Lamina DP-04 (seccion S-01) y DP-08"),
    ("sl_n", "Espaciamiento de barras longitudinales (tramo normal)", 0.25, "m", "Lamina DP-04 y DP-08"),
    ("sl_c", "Espaciamiento de barras longitudinales (camiones)", 0.20, "m", "Lamina DP-04 (seccion S-01) y DP-08"),
    ("nreg", "Numero de registros de limpieza con tapa", MC.registros()["n"], "und", "Lamina DP-01 (RV-01 a RV-11), todos en la losa del colector; hoja COLECTOR REGISTROS Y TAPAS"),
    ("tl", "Lado de la tapa de concreto", 0.68, "m", "Lamina DP-06B"),
    ("te", "Espesor de la tapa", 0.08, "m", "Lamina DP-06B"),
    ("ab", "Lado de la abertura en la losa (luz del contramarco)", 0.70, "m", "Lamina DP-06B"),
    ("pcm", "Perimetro del contramarco L 2\"x2\"x3/16\" (4 x 0.70)", 2.80, "m", "Lamina DP-06B"),
    ("pm", "Perimetro del marco de tapa L 1 1/2\"x1 1/2\"x1/8\" (4 x 0.68)", 2.72, "m", "Lamina DP-06B"),
    ("ka", "Peso del angulo L 2\"x2\"x3/16\"", MC.ANG["2x2x3/16"], "kg/m", "Catalogo del fabricante"),
    ("kb", "Peso del angulo L 1 1/2\"x1 1/2\"x1/8\"", MC.ANG["1.5x1.5x1/8"], "kg/m", "Catalogo del fabricante"),
    ("da", "Desarrollo pintado del angulo 2\"x2\" (4 caras)", 0.203, "m2/m", "4 x 0.0508 m"),
    ("db", "Desarrollo pintado del angulo 1 1/2\"x1 1/2\" (4 caras)", 0.152, "m2/m", "4 x 0.0381 m"),
    ("bp", "Borde engrosado de la abertura: perimetro medio", 3.00, "m", "Lamina DP-06B (4 x 0.75)"),
    ("bs", "Borde engrosado: seccion (0.15 x 0.10)", 0.15 * 0.10, "m2", "Lamina DP-06B"),
    ("CF0", "Cota de fondo del colector en el arranque (0+000)", D["CF0"], "msnm", "Lamina DP-02; memoria, hoja PERFIL_FLUJO"),
    ("CFF", "Cota de fondo del colector en la llegada a la CL (0+%06.2f)" % dz.P_FIN, round(dz.fondo(dz.P_FIN), 3), "msnm", "Lamina DP-02 y DP-07; memoria, hoja EMPALME_CL"),
    ("jd", "Espaciamiento de juntas de dilatacion", 4.00, "m", "Lamina DP-01, nota; DP-02"),
    ("hmax", "Altura interior maxima del colector (para el perimetro de la junta)", round(dz.techo(dz.P_FIN) - dz.fondo(dz.P_FIN), 2), "m", "Lamina DP-02 (perfil): llegada a la CL"),
    ("PF", "Progresiva final (cara este de la caja CL)", round(dz.P_FIN, 2), "m", "Lamina DP-01"),
    ("PQ", "Progresiva del quiebre (inicio del tramo de empalme)", round(dz.P_QUIEBRE, 2), "m", "Lamina DP-01"),
    ("CLb", "Ancho exterior del muro este de la caja CL (contacto con el colector)", round(dz.D["CL_ancho"] + 2 * dz.D["e_muro"], 2), "m", "Lamina DP-07"),
]


def hoja_parametros(wb):
    ws = wb.create_sheet(NOMBRES["PAR"])
    cols = [("N°", 5, "txt", None), ("PARAMETRO", 62, "txt", None), ("VALOR", 12, "auto", "0.000"), ("UND", 8, "und", None), ("SUSTENTO (lamina / norma / criterio)", 70, "sust", None)]
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - PARAMETROS DE METRADO", len(cols),
                 "Tramo CAR Varones (CUI 2705619), empalme con la caja de llegada CL del colector del Hogar de Refugio Temporal (CUI 2675514). Valores tomados de los planos DP-01 a DP-08, DA-01 a DA-03 y DD-01 a DD-04 y de la memoria de calculo del colector. Celda amarilla = dato; verde = formula.")
    encabezado(ws, r, cols, 30)
    P = {}
    for i, (k, desc, v, u, sus) in enumerate(PARAMETROS):
        rr = r + 1 + i
        fila(ws, rr, cols, [i + 1, desc, v, u, sus]); ws.cell(row=rr, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        P[k] = "%s!$C$%d" % (Q["PAR"], rr)
    return ws, P


def hoja_mov_tierras(wb, segs, P):
    ws = wb.create_sheet(NOMBRES["MT"])
    cols = [("N°", 5, "txt", None), ("TRAMO", 28, "txt", None), ("ZONA", 11, "tipo", None), ("PROG. INICIO", 11, "txt", None), ("PROG. FIN", 11, "txt", None),
            ("LONGITUD (m)", 11, "auto", "0.00"), ("ALTURA INTERIOR h (m)", 11, "dato", "0.000"), ("e LOSA SUP. (m)", 10, "auto", "0.00"), ("ANCHO EXT. b + 2e (m)", 11, "form", "0.00"),
            ("ANCHO DE ZANJA (m)", 11, "form", "0.00"), ("Hz: TERRENO - FONDO DE SOLADO (m)", 13, "dato", "0.000"), ("NPT - TERRENO (m)", 11, "dato", "0.000"),
            ("LIMPIEZA Y TRAZO (m2)", 12, "form", "0.00"), ("EXCAVACION (m3)", 12, "form", "0.00"), ("REFINE Y NIVELACION (m2)", 12, "form", "0.00"),
            ("RELLENO LATERAL DE ZANJA (m3)", 12, "form", "0.00"), ("RELLENO FRANJA DE NIVELACION (m3)", 12, "form", "0.00"), ("SUSTENTO", 46, "sust", None)]
    n = len(cols)
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - METRADO DE MOVIMIENTO DE TIERRAS", n,
                 "Criterio (lamina DA-01): LIMPIEZA Y TRAZO = ancho de zanja x L.  EXCAVACION = ancho de zanja x Hz x L, con Hz = terreno existente (topografia) - fondo del solado.  "
                 "REFINE = (ancho exterior + 0.10) x L (fondo del solado).  RELLENO LATERAL = 2 x sobreancho x Hz x L.  RELLENO FRANJA = 1.00 x (NPT - terreno) x L (nivelacion del retiro lado via).  "
                 "h y Hz son valores medios ponderados del tramo (perfil DP-02 y topografia); el detalle cada 2.00 m esta en la memoria de calculo.")
    encabezado(ws, r, cols); r0 = r + 1
    b, em, es, sob, fr = P["b"], P["em"], P["es"], P["sob"], P["fr"]
    filas = []
    for i, s in enumerate(segs):
        rr = r0 + i; R = str(rr)
        vals = [i + 1, s["nombre"], s["zona"], prog(s["p1"]), prog(s["p2"]), round(s["p2"] - s["p1"], 2), s["h"], s["et"], "=%s+2*%s" % (b, em), "=I%s+2*%s" % (R, sob),
                s["Hz"], s["dnpt"], "=J%s*F%s" % (R, R), "=J%s*K%s*F%s" % (R, R, R), "=(I%s+0.10)*F%s" % (R, R), "=2*%s*K%s*F%s" % (sob, R, R), "=%s*L%s*F%s" % (fr, R, R),
                "DP-01 (planta), DP-02 (perfil: fondo y terreno), DP-04 (seccion), DA-01 (criterio)"]
        fila(ws, rr, cols, vals, 30); filas.append(rr)
    rt = r0 + len(segs)
    barra_total(ws, rt, n, "TOTAL COLECTOR", {c: "=SUM(%s%d:%s%d)" % (c, r0, c, rt - 1) for c in "FMNOPQ"})
    rr = rt + 2
    subtitulo(ws, rr, n, "RESUMEN PARA LA PLANILLA GENERAL (partidas %s.01 y %s.02)" % (ITEM, ITEM))
    cols_r = [("N°", 5, "txt", None), ("PARTIDA", 28, "txt", None), ("", 11, "txt", None), ("", 11, "txt", None), ("", 11, "txt", None), ("UND", 11, "und", None), ("TOTAL", 11, "tot", "0.00")]
    res = [("LIMPIEZA MANUAL DE TERRENO / TRAZO Y REPLANTEO", "m2", "=M%d" % rt), ("EXCAVACION MANUAL DE ZANJAS", "m3", "=N%d" % rt), ("REFINE, NIVELACION Y COMPACTACION", "m2", "=O%d" % rt), ("RELLENO COMPACTADO (lateral + franja)", "m3", "=P%d+Q%d" % (rt, rt))]
    for i, (t, u, f) in enumerate(res):
        r_ = rr + 1 + i
        fila(ws, r_, cols_r, [i + 1, t, None, None, None, u, f], 22); ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=5); ws.cell(row=r_, column=2).alignment = Alignment(horizontal="left", vertical="center")
    r_ = rr + 1 + len(res)
    fila(ws, r_, cols_r, [len(res) + 1, "ELIMINACION DE MATERIAL EXCEDENTE = (excavacion - relleno) x esponjamiento", None, None, None, "m3", "=(G%d-G%d)*%s" % (rr + 2, rr + 4, P["esp"])], 22)
    ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=5); ws.cell(row=r_, column=2).alignment = Alignment(horizontal="left", vertical="center")
    nota(ws, r_ + 2, n, "Nota: el acortamiento de las cunetas que llegan al colector (Ejes 09, 08, 06, 04 y 01) esta aplicado en las hojas de cunetas de esta planilla (hoja CUNETAS - LLEGADA AL COLECTOR); no forma parte de esta partida. La caja CL pertenece al expediente del Hogar de Refugio.")
    ws.freeze_panes = ws.cell(row=r0, column=3)
    return ws, dict(filas=filas, rt=rt, elim=r_)


def hoja_concreto(wb, segs, P, mt):
    ws = wb.create_sheet(NOMBRES["CO"])
    cols = [("N°", 5, "txt", None), ("TRAMO", 28, "txt", None), ("ZONA", 11, "tipo", None), ("LONGITUD (m)", 11, "form", "0.00"), ("ALTURA INTERIOR h (m)", 11, "form", "0.000"),
            ("e LOSA SUP. (m)", 10, "form", "0.00"), ("ANCHO EXT. (m)", 10, "form", "0.00"), ("SOLADO f'c=100 (m2)", 12, "form", "0.00"), ("LOSA DE FONDO (m3)", 12, "form", "0.00"),
            ("MUROS (m3)", 12, "form", "0.00"), ("LOSA SUPERIOR (m3)", 12, "form", "0.00"), ("ACABADO DE LOSA SUP. (m2)", 12, "form", "0.00"), ("CURADO (m2)", 12, "form", "0.00"), ("SUSTENTO", 46, "sust", None)]
    n = len(cols)
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - METRADO DE CONCRETO SIMPLE Y ARMADO", n,
                 "Criterio (lamina DA-02): SOLADO = (ancho exterior + 0.10) x L.  LOSA DE FONDO = ancho exterior x e fondo x L.  MUROS = 2 x e muro x h x L.  LOSA SUPERIOR = ancho exterior x e losa x L.  "
                 "ACABADO = ancho exterior x L (losa superior frotachada y brunada).  CURADO = 2 caras interiores de muros (2 x h x L) + fondo interior (b x L) + losa superior (ancho exterior x L), igual al criterio de la hoja METRADO DE CURADO.")
    encabezado(ws, r, cols); r0 = r + 1
    MT = Q["MT"]; b, em, ef = P["b"], P["em"], P["ef"]
    filas = []
    for i, s in enumerate(segs):
        rr = r0 + i; R = str(rr); rm = mt["filas"][i]
        vals = [i + 1, s["nombre"], s["zona"], "=%s!F%d" % (MT, rm), "=%s!G%d" % (MT, rm), "=%s!H%d" % (MT, rm), "=%s!I%d" % (MT, rm),
                "=(G%s+0.10)*D%s" % (R, R), "=G%s*%s*D%s" % (R, ef, R), "=2*%s*E%s*D%s" % (em, R, R), "=G%s*F%s*D%s" % (R, R, R), "=G%s*D%s" % (R, R), "=(2*E%s+%s+G%s)*D%s" % (R, b, R, R),
                "DP-04 (secciones), DP-02 (h por tramo), DA-02 (areas por metro)"]
        fila(ws, rr, cols, vals, 30); filas.append(rr)
    rt = r0 + len(segs)
    barra_total(ws, rt, n, "TOTAL COLECTOR (tramos)", {c: "=SUM(%s%d:%s%d)" % (c, r0, c, rt - 1) for c in "DHIJKLM"})
    rc = rt + 2
    subtitulo(ws, rc, n, "REGISTROS DE LIMPIEZA EN LA LOSA DEL COLECTOR - lamina DP-06B (borde engrosado 0.15 x 0.10 alrededor de la abertura; descuento de la abertura 0.70 x 0.70 en la losa superior)")
    cols_c = [("N°", 5, "txt", None), ("ELEMENTO", 28, "txt", None), ("DETALLE", 11, "txt", None), ("LARGO (m)", 11, "auto", "0.00"), ("ANCHO (m)", 11, "auto", "0.00"), ("ALTO / ESPESOR (m)", 10, "auto", "0.000"),
              ("VECES", 10, "auto", "0"), ("SOLADO (m2)", 12, "form", "0.00"), ("LOSA DE FONDO (m3)", 12, "form", "0.00"), ("MUROS (m3)", 12, "form", "0.00"), ("LOSA SUPERIOR (m3)", 12, "form", "0.00"),
              ("ACABADO (m2)", 12, "form", "0.00"), ("CURADO (m2)", 12, "form", "0.00"), ("SUSTENTO", 46, "sust", None)]
    encabezado(ws, rc + 1, cols_c)
    rr = rc + 2; rreg = rr
    fila(ws, rr, cols_c, [1, "Registros RV-01 a RV-11: borde engrosado de la abertura", "perimetro x seccion", "=%s" % P["bp"], "=%s" % P["bs"], 1, "=%s" % P["nreg"], None, None, None, "=D%d*E%d*F%d*G%d" % (rr, rr, rr, rr), None, None, "DP-06B, DD-01: borde engrosado en los 11 registros"], 30); rr += 1
    fila(ws, rr, cols_c, [2, "Registros RV-01 a RV-11: descuento de la abertura en la losa superior", "0.70 x 0.70 x e", "=%s" % P["ab"], "=%s" % P["ab"], "=%s" % P["et"], "=-%s" % P["nreg"], None, None, None, "=D%d*E%d*F%d*G%d" % (rr, rr, rr, rr), None, None, "DP-06B"], 30); rr += 1
    for r_ in (rreg, rreg + 1): ws.cell(row=r_, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    rtc = rr
    barra_total(ws, rtc, n, "TOTAL REGISTROS", {c: "=SUM(%s%d:%s%d)" % (c, rc + 2, c, rtc - 1) for c in "HIJKLM"})
    rr = rtc + 2
    subtitulo(ws, rr, n, "RESUMEN PARA LA PLANILLA GENERAL (partidas %s.03 y %s.04)" % (ITEM, ITEM))
    cols_r = [("N°", 5, "txt", None), ("PARTIDA", 28, "txt", None), ("", 11, "txt", None), ("", 11, "txt", None), ("", 11, "txt", None), ("UND", 10, "und", None), ("TRAMOS", 10, "form", "0.00"), ("REGISTROS", 12, "form", "0.00"), ("TOTAL", 12, "tot", "0.00")]
    res = [("CONCRETO f'c=100 PARA SOLADO e=0.05", "m2", "H"), ("CONCRETO f'c=210 EN LOSA DE FONDO", "m3", "I"), ("CONCRETO f'c=210 EN MUROS", "m3", "J"), ("CONCRETO f'c=210 EN LOSA SUPERIOR", "m3", "K"),
           ("ACABADO FROTACHADO Y BRUÑADO DE LOSA SUPERIOR", "m2", "L"), ("CURADO DE CONCRETO", "m2", "M")]
    R_ = {}
    for i, (t, u, col) in enumerate(res):
        r_ = rr + 1 + i; R = str(r_)
        fila(ws, r_, cols_r, [i + 1, t, None, None, None, u, "=%s%d" % (col, rt), "=%s%d" % (col, rtc), "=G%s+H%s" % (R, R)], 22)
        ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=5); ws.cell(row=r_, column=2).alignment = Alignment(horizontal="left", vertical="center"); R_[col] = r_
    ws.freeze_panes = ws.cell(row=r0, column=3)
    return ws, dict(filas=filas, rt=rt, rtc=rtc, reg=rreg, res=R_)


def hoja_encofrado(wb, segs, P, mt):
    ws = wb.create_sheet(NOMBRES["EN"])
    cols = [("N°", 5, "txt", None), ("TRAMO", 28, "txt", None), ("ZONA", 11, "tipo", None), ("LONGITUD (m)", 11, "form", "0.00"), ("ALTURA INTERIOR h (m)", 11, "form", "0.000"),
            ("e LOSA SUP. (m)", 10, "form", "0.00"), ("ANCHO INT. b (m)", 10, "form", "0.00"), ("CARAS INTERIORES DE MUROS 2 x h x L (m2)", 13, "form", "0.00"),
            ("CARAS EXTERIORES 2 x (e fondo + h + e losa) x L (m2)", 14, "form", "0.00"), ("FONDO DE LOSA SUPERIOR b x L (m2)", 13, "form", "0.00"), ("TOTAL ENCOFRADO (m2)", 12, "tot", "0.00"), ("SUSTENTO", 46, "sust", None)]
    n = len(cols)
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - METRADO DE ENCOFRADO Y DESENCOFRADO", n,
                 "Criterio (lamina DA-02): se encofran las dos caras interiores de los muros, las dos caras exteriores en toda la altura del cajon (losa de fondo + muro + losa superior) y el fondo de la losa superior (ancho interior b). "
                 "La losa de fondo se vacia sobre el solado y la cara superior de la losa queda libre (acabado).")
    encabezado(ws, r, cols); r0 = r + 1
    MT = Q["MT"]; b, ef = P["b"], P["ef"]
    filas = []
    for i, s in enumerate(segs):
        rr = r0 + i; R = str(rr); rm = mt["filas"][i]
        vals = [i + 1, s["nombre"], s["zona"], "=%s!F%d" % (MT, rm), "=%s!G%d" % (MT, rm), "=%s!H%d" % (MT, rm), "=%s" % b,
                "=2*E%s*D%s" % (R, R), "=2*(%s+E%s+F%s)*D%s" % (ef, R, R, R), "=G%s*D%s" % (R, R), "=SUM(H%s:J%s)" % (R, R), "DP-04 (secciones), DA-02 (areas por metro)"]
        fila(ws, rr, cols, vals, 30); filas.append(rr)
    rt = r0 + len(segs)
    barra_total(ws, rt, n, "TOTAL COLECTOR", {c: "=SUM(%s%d:%s%d)" % (c, r0, c, rt - 1) for c in "DHIJK"})
    rr = rt + 2
    subtitulo(ws, rr, n, "RESUMEN PARA LA PLANILLA GENERAL (partida %s.04.04)" % ITEM)
    cols_r = [("N°", 5, "txt", None), ("PARTIDA", 28, "txt", None), ("", 11, "txt", None), ("", 11, "txt", None), ("", 11, "txt", None), ("UND", 10, "und", None), ("TOTAL", 10, "tot", "0.00")]
    r_ = rr + 1
    fila(ws, r_, cols_r, [1, "ENCOFRADO Y DESENCOFRADO NORMAL EN COLECTOR", None, None, None, "m2", "=K%d" % rt], 22)
    ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=5); ws.cell(row=r_, column=2).alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = ws.cell(row=r0, column=3)
    return ws, dict(filas=filas, rt=rt, res=r_)


def hoja_acero(wb, segs, P, mt):
    ws = wb.create_sheet(NOMBRES["AC"])
    cols = [("N°", 5, "txt", None), ("TRAMO", 28, "txt", None), ("ZONA", 11, "tipo", None), ("LONGITUD (m)", 10, "form", "0.00"), ("ALTURA INTERIOR h (m)", 10, "form", "0.000"), ("e LOSA SUP. (m)", 9, "form", "0.00"), ("ANCHO EXT. (m)", 9, "form", "0.00"),
            ("ESPAC. MARCOS s (m)", 9, "form", "0.00"), ("CAPAS (1 = marco unico; 2 = ext. + int.)", 11, "dato", "0"), ("DIAM. MARCOS", 9, "tipo", None), ("N° DE JUEGOS (L / s, redondeado arriba)", 11, "form", "0"),
            ("PERIMETRO MARCO 1 (m)", 11, "form", "0.000"), ("PERIMETRO MARCO 2 (m)", 11, "form", "0.000"), ("LONGITUD DE MARCOS (m)", 11, "form", "0.00"), ("PESO (kg/m)", 9, "form", "0.000"), ("ACERO EN MARCOS (kg)", 11, "form", "0.00"),
            ("ESPAC. LONGIT. (m)", 9, "form", "0.00"), ("N° BARRAS LONGIT. (perimetro / espac.)", 11, "form", "0"), ("LONGITUD LONGIT. CON TRASLAPE (m)", 12, "form", "0.00"), ("ACERO LONGIT. 3/8\" (kg)", 11, "form", "0.00"), ("TOTAL (kg)", 11, "tot", "0.00"), ("SUSTENTO", 44, "sust", None)]
    n = len(cols)
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - METRADO DE ACERO DE REFUERZO fy=4200 kg/cm2", n,
                 "Criterio (laminas DP-04 y DP-08, norma E.060 14.3.4): el tramo normal lleva un solo marco cerrado en el eje de muros y losas (perimetro = 2 x (ancho ext. - e) + 2 x (e fondo/2 + h + e losa/2) + ganchos); "
                 "el cruce de camiones cisterna lleva marco exterior + marco interior de 1/2\" (perimetros con recubrimiento 0.04). Barras longitudinales 3/8\" en cada marco, con traslape de 0.40 cada 9.00 m. Peso: 3/8\" = 0.56 kg/m, 1/2\" = 0.994 kg/m.")
    encabezado(ws, r, cols, 60); r0 = r + 1
    MT = Q["MT"]; b, em, ef, rec, gan, p38, p12, tr, lb = P["b"], P["em"], P["ef"], P["rec"], P["gan"], P["p38"], P["p12"], P["tr"], P["lb"]
    filas = []
    for i, s in enumerate(segs):
        rr = r0 + i; R = str(rr); rm = mt["filas"][i]
        sp = "=%s" % (P["s_n"] if s["zona"] == "NORMAL" else P["s_c"]); sl = "=%s" % (P["sl_c"] if s["zona"] == "CAMION" else P["sl_n"])
        vals = [i + 1, s["nombre"], s["zona"], "=%s!F%d" % (MT, rm), "=%s!G%d" % (MT, rm), "=%s!H%d" % (MT, rm), "=%s!I%d" % (MT, rm), sp, s["capas"], s["marcos_dia"] + '"', "=ROUNDUP(D%s/H%s,0)" % (R, R),
                "=IF(I%s=2,2*(G%s-2*%s)+2*(%s+E%s+F%s-2*%s)+%s,2*(G%s-%s)+2*(%s/2+E%s+F%s/2)+%s)" % (R, R, rec, ef, R, R, rec, gan, R, em, ef, R, R, gan),
                "=IF(I%s=2,2*(%s+2*(%s-%s))+2*(E%s+2*(%s-%s))+%s,0)" % (R, b, em, rec, R, em, rec, gan),
                "=K%s*(L%s+M%s)" % (R, R, R), "=IF(J%s=\"1/2\"\"\",%s,%s)" % (R, p12, p38), "=N%s*O%s" % (R, R),
                sl, "=ROUND(L%s/Q%s,0)+IF(M%s>0,ROUND(M%s/Q%s,0),0)" % (R, R, R, R, R), "=R%s*D%s*(1+%s/%s)" % (R, R, tr, lb), "=S%s*%s" % (R, p38), "=P%s+T%s" % (R, R),
                "DP-04 (seccion %s), DP-08 (cuadro de doblado), DA-03" % ("S-01" if s["zona"] == "CAMION" else "S-02 a S-08")]
        fila(ws, rr, cols, vals, 30); filas.append(rr)
    rt = r0 + len(segs)
    barra_total(ws, rt, n, "TOTAL COLECTOR", {c: "=SUM(%s%d:%s%d)" % (c, r0, c, rt - 1) for c in "DNPSTU"})
    rg = rt + 2
    subtitulo(ws, rg, n, "REGISTROS DE LIMPIEZA Y TAPAS - laminas DP-06B, DD-01 y DD-02: refuerzo de borde de la abertura, parrilla de la tapa, asas y anclajes del contramarco (cantidad de registros de la hoja COLECTOR REGISTROS Y TAPAS)")
    cols_g = [("N°", 5, "txt", None), ("ELEMENTO", 28, "txt", None), ("DIAM.", 11, "tipo", None), ("FORMA", 10, "txt", None), ("PIEZAS POR REGISTRO", 10, "auto", "0"), ("REGISTROS", 9, "form", "0"), ("LONG. POR PIEZA (m)", 9, "auto", "0.00"), ("LONG. TOTAL (m)", 9, "form", "0.00"), ("PESO (kg/m)", 11, "form", "0.000")] + [("", 9, "txt", None)] * 11 + [("ACERO (kg)", 11, "tot", "0.00"), ("SUSTENTO", 44, "sust", None)]
    encabezado(ws, rg + 1, cols_g, 40)
    reg = [("Refuerzo de borde de la abertura", '1/2"', "recta", 8, 1.40, p12, "DP-06B, DD-01: 2 barras de 1/2\" por lado, L = 1.40"),
           ("Parrilla de la tapa de concreto", '3/8"', "recta", 14, 0.62, p38, "DD-02: 7 + 7 barras 3/8\" @0.10"),
           ("Asas de la tapa", '3/8" liso', "U", 2, 0.40, p38, "DD-02: 2 asas de 0.40 desarrollado"),
           ("Anclajes del contramarco", '3/8"', "L", 8, 0.20, p38, "DD-01: 8 anclajes de 0.20 soldados al angulo")]
    rr = rg + 2
    for i, (desc, dia, forma, piezas, lp, peso, sus) in enumerate(reg):
        R = str(rr)
        fila(ws, rr, cols_g, [i + 1, desc, dia, forma, piezas, "=%s" % P["nreg"], lp, "=E%s*F%s*G%s" % (R, R, R), "=%s" % peso] + [None] * 11 + ["=H%s*I%s" % (R, R), sus], 30)
        ws.cell(row=rr, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); rr += 1
    rtg = rr
    barra_total(ws, rtg, n, "TOTAL REGISTROS Y TAPAS", {"U": "=SUM(U%d:U%d)" % (rg + 2, rtg - 1)})
    rr = rtg + 2
    subtitulo(ws, rr, n, "RESUMEN PARA LA PLANILLA GENERAL (partidas %s.04.05 y %s.05.06)" % (ITEM, ITEM))
    cols_r = [("N°", 5, "txt", None), ("PARTIDA", 28, "txt", None), ("", 11, "txt", None), ("", 10, "txt", None), ("", 10, "txt", None), ("UND", 9, "und", None), ("3/8\" (kg)", 9, "form", "0.00"), ("1/2\" (kg)", 9, "form", "0.00"), ("TOTAL (kg)", 11, "tot", "0.00")]
    r1 = rr + 1
    fila(ws, r1, cols_r, [1, "ACERO DE REFUERZO EN COLECTOR", None, None, None, "kg", "=SUMIF(J%d:J%d,\"3/8\"\"\",P%d:P%d)+T%d" % (r0, rt - 1, r0, rt - 1, rt), "=SUMIF(J%d:J%d,\"1/2\"\"\",P%d:P%d)" % (r0, rt - 1, r0, rt - 1), "=U%d" % rt], 22)
    r2 = rr + 2
    fila(ws, r2, cols_r, [2, "ACERO DE REFUERZO EN TAPAS, BORDES Y ANCLAJES", None, None, None, "kg", "=SUMIF(C%d:C%d,\"3/8*\",U%d:U%d)" % (rg + 2, rtg - 1, rg + 2, rtg - 1), "=SUMIF(C%d:C%d,\"1/2*\",U%d:U%d)" % (rg + 2, rtg - 1, rg + 2, rtg - 1), "=U%d" % rtg], 22)
    for r_ in (r1, r2):
        ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=5); ws.cell(row=r_, column=2).alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = ws.cell(row=r0, column=3)
    return ws, dict(filas=filas, rt=rt, rtg=rtg, colector=r1, registros=r2)


def hoja_registros(wb, P):
    ws = wb.create_sheet(NOMBRES["RT"])
    cols = [("N°", 5, "txt", None), ("REGISTRO", 16, "txt", None), ("PROGRESIVA", 12, "txt", None), ("UBICACION", 46, "txt", None), ("TAPA 0.68 x 0.68 (und)", 12, "auto", "0"), ("SUSTENTO", 50, "sust", None)]
    n = len(cols)
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - REGISTROS DE LIMPIEZA Y TAPAS", n,
                 "Lamina DP-01 (ubicacion), DP-06B, DD-01 y DD-02 (detalle): contramarco L 2\"x2\"x3/16\" anclado a la losa, marco L 1 1/2\"x1 1/2\"x1/8\" soldado a la tapa, tapa de concreto armado 0.68 x 0.68 x 0.08 con parrilla 3/8\" @0.10 y asas; borde engrosado 0.15 x 0.10 alrededor de la abertura de 0.70 x 0.70. El registro de la caja CL pertenece al expediente del Hogar de Refugio.")
    encabezado(ws, r, cols, 30); r0 = r + 1
    rr = r0
    for i, x in enumerate([x for x in dz.REGISTROS if x["tipo"] == "registro"]):
        fila(ws, rr, cols, [i + 1, x["nombre"], prog(x["prog"]), "Registro en la losa del colector; " + x["nota"], 1, "DP-01, DP-06B"], 24); ws.cell(row=rr, column=4).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); rr += 1
    rt = rr
    barra_total(ws, rt, n, "TOTAL DE REGISTROS CON TAPA", {"E": "=SUM(E%d:E%d)" % (r0, rt - 1)}, nf="0")
    rr = rt + 2
    subtitulo(ws, rr, n, "METRADO DE LAS PARTIDAS DE REGISTROS (cantidad = total de registros; dimensiones de la hoja COLECTOR PARAMETROS)")
    cols_m = [("N°", 5, "txt", None), ("PARTIDA", 16, "txt", None), ("UND", 12, "und", None), ("CALCULO", 46, "txt", None), ("TOTAL", 12, "tot", "0.00"), ("SUSTENTO", 50, "sust", None)]
    encabezado(ws, rr + 1, cols_m, 30)
    items = [("contramarco", "CONTRAMARCO METALICO L 2\"x2\"x3/16\" CON ANCLAJES", "und", "1 por registro", "=E%d" % rt, "DP-06B, DD-01"),
             ("marco", "MARCO METALICO DE TAPA L 1 1/2\"x1 1/2\"x1/8\"", "und", "1 por tapa", "=E%d" % rt, "DP-06B, DD-02"),
             ("angulos", "ANGULOS METALICOS P/MARCO Y CONTRAMARCO", "kg", "registros x (2.80 x 3.63 + 2.72 x 1.83)", "=E%d*(%s*%s+%s*%s)" % (rt, P["pcm"], P["ka"], P["pm"], P["kb"]), "DP-06B; pesos de catalogo"),
             ("tapa", "TAPA DE CONCRETO ARMADO 0.68 x 0.68 x 0.08", "und", "1 por registro", "=E%d" % rt, "DD-02"),
             ("conc_tapa", "CONCRETO f'c=210 EN TAPAS", "m3", "registros x 0.68 x 0.68 x 0.08", "=E%d*%s*%s*%s" % (rt, P["tl"], P["tl"], P["te"]), "DD-02"),
             ("pintura", "PINTURA ANTICORROSIVA Y ESMALTE EN ANGULOS", "m2", "registros x (2.80 x 0.203 + 2.72 x 0.152)", "=E%d*(%s*%s+%s*%s)" % (rt, P["pcm"], P["da"], P["pm"], P["db"]), "4 caras de cada angulo")]
    refs = {}; rr = rr + 2
    for i, (k, desc, u, calc, f, sus) in enumerate(items):
        fila(ws, rr, cols_m, [i + 1, desc, u, calc, f, sus], 30)
        ws.cell(row=rr, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True); ws.column_dimensions["B"].width = 40
        refs[k] = rr; rr += 1
    return ws, dict(n=rt, refs=refs)


def hoja_juntas(wb, P, mt):
    ws = wb.create_sheet(NOMBRES["JE"])
    cols = [("N°", 5, "txt", None), ("ELEMENTO", 48, "txt", None), ("PROGRESIVA / DETALLE", 22, "txt", None), ("VECES", 9, "auto", "0"), ("LONGITUD (m)", 12, "auto", "0.00"), ("PARCIAL (m)", 12, "form", "0.00"), ("SUSTENTO", 52, "sust", None)]
    n = len(cols)
    r = cabecera(ws, "COLECTOR PLUVIAL FRONTAL - JUNTAS Y EMPALMES DE CUNETAS", n,
                 "Laminas DP-01, DP-02, DP-06C, DP-07 y DD-03. Junta de dilatacion de 1\" cada 4.00 m en todo el perimetro de la seccion (tecnopor + sellador elastomerico). Junta de tecnopor de 1\" entre el borde de la losa y el piso adyacente (no hay cerco perimetrico) y en el contacto con la caja CL. Empalme de cada cuneta por ventana en el muro lado predio con caida libre al fondo (sin dintel adicional).")
    encabezado(ws, r, cols, 30); r0 = r + 1
    rr = r0
    subtitulo(ws, rr, n, "JUNTA DE DILATACION e=1\" CON TECNOPOR Y SELLADOR ELASTOMERICO (perimetro = 2 x ancho exterior + 2 x (e fondo + h max + e losa))"); rr += 1
    be = "(%s+2*%s)" % (P["b"], P["em"]); per = "2*%s+2*(%s+%s+%s)" % (be, P["ef"], P["hmax"], P["et"])
    r_jd = rr
    fila(ws, rr, cols, [1, "Juntas de dilatacion cada 4.00 m entre 0+004.00 y 0+%06.2f" % (int(dz.P_FIN // 4) * 4), "=\"n = ENTERO(\"&TEXT(%s,\"0.00\")&\" / 4.00)\"" % P["PF"], "=INT(%s/%s)" % (P["PF"], P["jd"]), "=" + per, "=D%d*E%d" % (rr, rr), "DP-01; DP-02 (juntas en el perfil); DD-03; perimetro de la seccion"], 30); rr += 1
    rt1 = rr; barra_total(ws, rt1, n, "TOTAL JUNTA DE DILATACION (m)", {"F": "=SUM(F%d:F%d)" % (r_jd, rt1 - 1)}); rr += 2
    subtitulo(ws, rr, n, "JUNTA DE TECNOPOR e=1\" ENTRE COLECTOR, PISO ADYACENTE Y CAJA CL (el frente de Varones no tiene cerco perimetrico)"); rr += 1
    r_jt = rr
    fila(ws, rr, cols, [1, "Borde de la losa con el piso adyacente: colector (ambos lados)", "=\"2 x 0+\"&TEXT(%s,\"000.00\")" % P["PF"], 2, "=%s" % P["PF"], "=D%d*E%d" % (rr, rr), "DP-04, DD-03: junta de tecnopor en el borde de la losa superior"], 30); rr += 1
    fila(ws, rr, cols, [2, "Contacto del colector con el muro este de la caja CL (perimetro de la seccion)", "0+%06.2f" % dz.P_FIN, 1, "=" + per.replace("+%s)" % P["hmax"], "+%s)" % P["hmax"]), "=D%d*E%d" % (rr, rr), "DP-07: tecnopor en todo el contacto colector - CL"], 30); rr += 1
    rt2 = rr; barra_total(ws, rt2, n, "TOTAL JUNTA DE TECNOPOR (m)", {"F": "=SUM(F%d:F%d)" % (r_jt, rt2 - 1)}); rr += 2
    subtitulo(ws, rr, n, "EMPALME DE CUNETA AL COLECTOR (ventana 0.40 x H en el muro lado predio, caida libre y registro encima) - laminas DP-06C y DD-04"); rr += 1
    r_e = rr
    k = 0
    for c in dz.CUNETAS:
        if c["entra_en"] != "colector": continue
        k += 1
        fila(ws, rr, cols, [k, "Cuneta Eje %s (registro %s)" % (c["perfil"], [x["nombre"] for x in dz.REGISTROS if abs(x["prog"] - c["prog"]) < 1.5][0]), prog(c["prog"]), 1, None, "=D%d" % rr, "DP-06C (cuadro de empalmes); DP-01"], 24); ws.cell(row=rr, column=6).number_format = "0"; rr += 1
    rt3 = rr; barra_total(ws, rt3, n, "TOTAL EMPALMES (und)", {"F": "=SUM(F%d:F%d)" % (r_e, rt3 - 1)}, nf="0")
    nota(ws, rt3 + 2, n, "Nota: la cuneta del Eje 01 no entra al colector: cae directamente a la caja CL por una ventana en su muro norte (DP-07); esa ventana y la del muro este (llegada del colector) se dejan en la caja CL del expediente del Hogar de Refugio. Las cunetas se acortan hasta el muro (hoja CUNETAS - LLEGADA AL COLECTOR).")
    for rr_ in range(r0, rt3 + 1): ws.cell(row=rr_, column=2).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    return ws, dict(dilat=rt1, tecnopor=rt2, empalmes=rt3)


def construir():
    wb = openpyxl.load_workbook(ORIG)
    ws = wb[PG]
    segs = MC.segmentos()
    wsP, P = hoja_parametros(wb)
    wsM, mt = hoja_mov_tierras(wb, segs, P)
    wsC, co = hoja_concreto(wb, segs, P, mt)
    wsE, en = hoja_encofrado(wb, segs, P, mt)
    wsA, ac = hoja_acero(wb, segs, P, mt)
    wsR, rg = hoja_registros(wb, P)
    wsJ, je = hoja_juntas(wb, P, mt)
    MT, CO, EN, AC, RT, JE = Q["MT"], Q["CO"], Q["EN"], Q["AC"], Q["RT"], Q["JE"]
    ult = max(c.row for row in ws.iter_rows() for c in row if c.value is not None)
    pl = Planilla(ws, ult + 2)
    pl.madre(ITEM, "COLECTOR PLUVIAL FRONTAL (TRAMO CAR VARONES - EMPALME CON LA CAJA CL DEL HOGAR DE REFUGIO)")
    def det(hoja, filas, col, nombres, extra=""):
        return [("%s%s" % (nm, extra), 1, None, None, None, "=%s!%s%d" % (hoja, col, r)) for nm, r in zip(nombres, filas)]
    nom_t = ["%s (%s - %s)" % (s["nombre"].capitalize(), prog(s["p1"]), prog(s["p2"])) for s in segs]
    I = lambda s: ITEM + s
    pl.titulo(I(".01"), "TRABAJOS PRELIMINARES")
    pl.partida(I(".01.01"), "LIMPIEZA MANUAL DE TERRENO", "m²", det(MT, mt["filas"], "M", nom_t))
    pl.partida(I(".01.02"), "TRAZO, NIVELES Y REPLANTEO PRELIMINAR", "m²", det(MT, mt["filas"], "M", nom_t))
    pl.titulo(I(".02"), "MOVIMIENTO DE TIERRAS")
    pl.partida(I(".02.01"), "EXCAVACIÓN MANUAL DE ZANJAS P/COLECTOR (ancho 1.40 m)", "m³", det(MT, mt["filas"], "N", nom_t))
    pl.partida(I(".02.02"), "REFINE, NIVELACIÓN Y COMPACTACIÓN EN TERRENO NORMAL", "m²", det(MT, mt["filas"], "O", nom_t))
    r_rell = pl.partida(I(".02.03"), "RELLENO COMPACTADO CON MATERIAL PROPIO SELECCIONADO (incluye nivelacion del retiro)", "m³",
                        det(MT, mt["filas"], "P", nom_t, " - lateral de zanja") + det(MT, mt["filas"], "Q", nom_t, " - franja de nivelacion"))
    r_exc = pl.refs[I(".02.01")]
    pl.partida(I(".02.04"), "ELIMINACIÓN DE MATERIAL EXCEDENTE", "m³", [("Excavacion de zanjas", 1, None, None, None, "=K%d" % r_exc), ("Relleno compactado (se descuenta)", -1, None, None, None, "=-K%d" % r_rell)])
    ws.cell(row=pl.refs[I(".02.04")], column=11, value="=SUM(J%d:J%d)*PARAMETROS!$B$16" % (pl.refs[I(".02.04")] + 1, pl.refs[I(".02.04")] + 2))
    pl.titulo(I(".03"), "OBRAS DE CONCRETO SIMPLE")
    pl.partida(I(".03.01"), "CONCRETO f'c=100 kg/cm2 PARA SOLADO E=2\" (0.05 m)", "m²", det(CO, co["filas"], "H", nom_t))
    pl.titulo(I(".04"), "OBRAS DE CONCRETO ARMADO - COLECTOR")
    pl.partida(I(".04.01"), "CONCRETO f'c=210 kg/cm2 EN LOSA DE FONDO", "m³", det(CO, co["filas"], "I", nom_t))
    pl.partida(I(".04.02"), "CONCRETO f'c=210 kg/cm2 EN MUROS", "m³", det(CO, co["filas"], "J", nom_t))
    pl.partida(I(".04.03"), "CONCRETO f'c=210 kg/cm2 EN LOSA SUPERIOR (incluye bordes de registro)", "m³", det(CO, co["filas"], "K", nom_t)
                + [("Bordes engrosados de abertura de registros RV-01 a RV-11", 1, None, None, None, "=%s!K%d" % (CO, co["reg"])), ("Descuento de aberturas de registro 0.70 x 0.70 x 0.10", 1, None, None, None, "=%s!K%d" % (CO, co["reg"] + 1))])
    pl.partida(I(".04.04"), "ENCOFRADO Y DESENCOFRADO NORMAL EN COLECTOR", "m²", det(EN, en["filas"], "K", nom_t))
    pl.partida(I(".04.05"), "ACERO DE REFUERZO fy=4200 kg/cm2 EN COLECTOR", "kg", det(AC, ac["filas"], "U", nom_t))
    pl.partida(I(".04.06"), "ACABADO FROTACHADO Y BRUÑADO DE LOSA SUPERIOR", "m²", det(CO, co["filas"], "L", nom_t))
    pl.partida(I(".04.07"), "CURADO DE CONCRETO EN COLECTOR", "m²", det(CO, co["filas"], "M", nom_t))
    pl.titulo(I(".05"), "REGISTROS DE LIMPIEZA Y TAPAS")
    rf = rg["refs"]
    pl.partida(I(".05.01"), "CONTRAMARCO METALICO L 2\"x2\"x3/16\" CON ANCLAJES", "und", None, total="=+%s!E%d" % (RT, rf["contramarco"]))
    pl.partida(I(".05.02"), "MARCO METALICO DE TAPA L 1 1/2\"x1 1/2\"x1/8\"", "und", None, total="=+%s!E%d" % (RT, rf["marco"]))
    pl.partida(I(".05.03"), "ANGULOS METALICOS P/MARCO Y CONTRAMARCO (incluye soldadura)", "kg", None, total="=+%s!E%d" % (RT, rf["angulos"]))
    pl.partida(I(".05.04"), "TAPA DE CONCRETO ARMADO 0.68 x 0.68 x 0.08 m (incluye acero y asas)", "und", None, total="=+%s!E%d" % (RT, rf["tapa"]))
    pl.partida(I(".05.05"), "CONCRETO f'c=210 kg/cm2 EN TAPAS", "m³", None, total="=+%s!E%d" % (RT, rf["conc_tapa"]))
    pl.partida(I(".05.06"), "ACERO DE REFUERZO fy=4200 kg/cm2 EN TAPAS, BORDES Y ANCLAJES", "kg", None, total="=+%s!I%d" % (AC, ac["registros"]))
    pl.partida(I(".05.07"), "PINTURA ANTICORROSIVA Y ESMALTE EN ANGULOS", "m²", None, total="=+%s!E%d" % (RT, rf["pintura"]))
    pl.titulo(I(".06"), "JUNTAS Y EMPALMES")
    pl.partida(I(".06.01"), "JUNTA DE DILATACION E=1\" CON TECNOPOR Y SELLADOR ELASTOMERICO", "m", None, total="=+%s!F%d" % (JE, je["dilat"]))
    pl.partida(I(".06.02"), "JUNTA DE TECNOPOR E=1\" ENTRE COLECTOR, PISO ADYACENTE Y CAJA CL", "m", None, total="=+%s!F%d" % (JE, je["tecnopor"]))
    pl.partida(I(".06.03"), "EMPALME DE CUNETA AL COLECTOR (ventana en muro y caida)", "und", None, total="=+%s!F%d" % (JE, je["empalmes"]))
    MX.hoja_insumos(wb, pl.refs)
    # RESUMEN
    wr = wb["RESUMEN"]; ultr = max(c.row for row in wr.iter_rows() for c in row if c.value is not None)
    r = ultr + 2
    def res_row(item, desc, und, rp, ref_r):
        for c in range(1, 6): copiar_estilo(wr.cell(row=ref_r, column=c), wr.cell(row=r, column=c))
        wr.cell(row=r, column=1, value=item); wr.cell(row=r, column=2, value=desc)
        if und: wr.cell(row=r, column=3, value=und); wr.cell(row=r, column=4, value="=+'%s'!K%d" % (PG, rp))
    res_row(ITEM, "COLECTOR PLUVIAL FRONTAL (TRAMO CAR VARONES - EMPALME CON LA CAJA CL DEL HOGAR DE REFUGIO)", None, None, 9); r += 1
    for rr_ in range(ult + 2, pl.r):
        item = ws.cell(row=rr_, column=1).value; desc = ws.cell(row=rr_, column=2).value; und = ws.cell(row=rr_, column=12).value
        if not item or item == ITEM: continue
        if und: res_row(item, desc, und, rr_, 11)
        else: res_row(item, desc, None, None, 10)
        r += 1
    os.makedirs(os.path.dirname(SALIDA), exist_ok=True); wb.save(SALIDA)
    reinyectar_vml(ORIG, SALIDA)
    return SALIDA, (ult + 2, pl.r - 1), (ultr + 2, r - 1)


if __name__ == "__main__":
    fn, pg, rs = construir()
    print(fn, "planilla filas", pg, "resumen filas", rs)
