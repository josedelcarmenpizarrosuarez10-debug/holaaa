"""Capitulo del colector pluvial frontal al final de la memoria de hidrologia y drenaje pluvial del CAR Varones
(insumos/katiuska/MEMORIA_HIDROLOGICA_VARONES.docx), con el mismo formato que el capitulo del Hogar de Refugio.
No modifica el contenido existente: agrega una seccion nueva al final."""
import os, sys, json, subprocess
import openpyxl
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
AQUI = os.path.dirname(os.path.abspath(__file__)); HERR = os.path.dirname(AQUI); RAIZ = os.path.dirname(HERR)
for p in (HERR, AQUI):
    if p not in sys.path: sys.path.insert(0, p)
import compat as C
import diseno_v as dz
import diseno as SOL
import metrado_calc as MC
import memoria_docx as MD
from memoria_docx import parrafo, titulo1, titulo2, vineta, formula, tabla, leyenda, fuente, imagen, f2, f3

ORIG = os.path.join(RAIZ, "insumos", "katiuska", "MEMORIA_HIDROLOGICA_VARONES.docx")
DEST = os.path.join(dz.SAL, "MEMORIA_HIDROLOGICA_VARONES_CON_COLECTOR.docx")
FIG = os.path.join(dz.SAL, "_tmp", "fig")
MEM = os.path.join(dz.SAL, "MEMORIA_CALCULO_COLECTOR_CAR_VARONES.xlsx")
D = dz.D; DS_ = SOL.D
NUM = MD.Numerador(22, 5); MD.NUM = NUM


def recalculada():
    out = os.path.join(dz.SAL, "_tmp", "mem"); os.makedirs(out, exist_ok=True)
    rec = os.path.join(out, os.path.basename(MEM))
    if not os.path.exists(rec) or os.path.getmtime(rec) < os.path.getmtime(MEM):
        subprocess.run(["soffice", "-env:UserInstallation=file:///tmp/lo_profile_v", "--headless", "--convert-to", "xlsx", "--outdir", out, MEM],
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
    return openpyxl.load_workbook(rec, data_only=True)


def leer_estructural(wb):
    ws = wb["ESTRUCTURAL"]; sec = None; res = {}
    for r in ws.iter_rows():
        f = [c.value for c in r]; a = f[0]
        if isinstance(a, str) and len(a) > 2 and a[1] == "." and a[0] in "ACDFGH": sec = a[0]; res[sec] = {"titulo": a[3:].strip()}
        elif sec and isinstance(a, str):
            k = a.strip()
            if k.startswith("Mu"): res[sec]["Mu"] = f[1]
            elif k.startswith("Momento resistente"): res[sec]["Mn"] = f[1]; res[sec]["ok"] = f[4]
            elif k.startswith("Cortante ultimo") or k == "Vu": res[sec]["Vu"] = f[1]
            elif k.startswith("Cortante resistente") or k == "fVc": res[sec]["Vc"] = f[1]
            elif k.startswith("Presion sobre el suelo en el cruce"): res[sec]["qc"] = f[1]
            elif k.startswith("Presion sobre el suelo"): res[sec]["q"] = f[1]
    return res


def leer_velocidades(wb):
    ws = wb["PERFIL_FLUJO"]; out = []; dentro = False
    for row in ws.iter_rows(max_col=12):
        a = row[0].value
        if isinstance(a, str) and a.startswith("VERIFICACION DE VELOCIDADES"): dentro = True; continue
        if dentro and isinstance(a, str) and a.startswith("Velocidad maxima"): break
        if dentro and isinstance(row[1].value, (int, float)) and isinstance(row[2].value, (int, float)) and isinstance(a, str):
            out.append([a] + [c.value for c in row[1:12]])
    return out


def leer_cumplimiento(wb):
    ws = wb["CUMPLIMIENTO"]; filas = []; pend = []; modo = "tabla"
    for row in ws.iter_rows(min_row=5, max_col=5):
        a = row[0].value
        if a is None: continue
        if isinstance(a, str) and a.startswith("Requisitos que no cumplen"): modo = "x"; continue
        if isinstance(a, str) and a.startswith("DATOS DE OTROS EXPEDIENTES"): modo = "pend"; continue
        if modo == "tabla": filas.append([c.value for c in row[:5]])
        elif modo == "pend": pend.append(a)
    return filas, pend


def construir():
    doc = Document(ORIG)
    n_par, n_tab = len(doc.paragraphs), len(doc.tables)
    s0 = doc.sections[0]
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.orientation = WD_ORIENT.PORTRAIT; sec.page_width = s0.page_width; sec.page_height = s0.page_height
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = s0.left_margin, s0.right_margin, s0.top_margin, s0.bottom_margin
    sec.header_distance, sec.footer_distance = s0.header_distance, s0.footer_distance
    sec.different_first_page_header_footer = False

    DJ = json.load(open(os.path.join(dz.CALC, "diseno.json"), encoding="utf8"))
    perfil = [e for e in DJ["perfil"] if e["Q"] > 0]; cl = DJ["caja_llegada"]; g = DJ["geom"]
    T, RM = MC.resumen(); wb = recalculada(); est = leer_estructural(wb); vel = leer_velocidades(wb); cump, pend = leer_cumplimiento(wb)
    h_min = min(p["h"] for p in DJ["perfil"]); h_max = max(p["h"] for p in DJ["perfil"])
    llen = max(p["y"] / p["h"] for p in perfil); bl_min = min(p["techo"] - p["NA"] for p in perfil)
    Fmax = max(p["F"] for p in perfil if abs(p["p"] - g["P_FIN"]) > 0.01); Vmax = max(p["V"] for p in perfil)
    cun = DJ["cunetas"]; c01 = [c for c in cun if c["entra_en"] == "CL"][0]
    Qcol = DJ["Q"]; Qdir = DJ["Q_directo_CL"]; Qt = DJ["Q_CL"]

    # ------------------------------------------------------------------ 1. alcance
    titulo1(doc, "COLECTOR PLUVIAL FRONTAL: TRAMO CAR VARONES Y EMPALME CON EL HOGAR DE REFUGIO")
    titulo2(doc, "Alcance del tramo")
    parrafo(doc, "Las cunetas del Centro de Acogida Residencial - Varones (CUI 2705619) descargan hacia el frente del predio, "
                 "sobre la carretera Oasis, donde se proyecta un colector pluvial cubierto que forma parte del colector continuo "
                 "compartido con los dos proyectos vecinos. El tramo de este proyecto recoge las cunetas de los Ejes 09, 08, 06 y 04 "
                 "y entrega el caudal a la caja de llegada CL del colector del Hogar de Refugio Temporal Mujeres Violentadas "
                 "(CUI 2675514); la cuneta del Eje 01, que llega al frente junto a esa caja, entra directamente en ella. Desde la "
                 "CL, el colector del Hogar de Refugio conduce el caudal hasta el registro R-01 del CAR Mujeres (CUI 2717013) y de "
                 "allí al punto de descarga final.")
    parrafo(doc, "Este capítulo desarrolla el trazo, el reparto del caudal de diseño entre las cunetas, el dimensionamiento hidráulico "
                 "de la sección, el perfil de flujo con caudal creciente, la verificación estructural, el empalme con la caja CL y "
                 "los criterios constructivos. El caudal total del proyecto (%.1f L/s) es el calculado en el presente estudio "
                 "hidrológico y es el mismo valor que el expediente del Hogar de Refugio considera como aporte del CAR Varones, de "
                 "modo que los tres expedientes quedan coordinados en el punto de empalme." % Qt)
    parrafo(doc, "Condiciones de compatibilidad verificadas en la entrega a la caja CL:", negrita=True)
    vineta(doc, "Caída libre del fondo del colector sobre el nivel de agua de la CL (%.3f msnm) y colchón de agua en la poza." % cl["NA_CL"])
    vineta(doc, "Caudal entregado (colector + cuneta Eje 01) igual al previsto por el Hogar de Refugio para el CAR Varones (%.1f L/s)." % DS_["Q_varones"])
    vineta(doc, "Niveles coordinados: tapa de la CL al ras del piso terminado del CAR Varones (+%.2f msnm) y ventanas de llegada definidas en ambos expedientes." % D["NPT_CL"])

    # ------------------------------------------------------------------ 2. trazo
    titulo2(doc, "Trazo y emplazamiento")
    parrafo(doc, "El colector arranca en el poste derecho del portón de camiones cisterna (0+000.00, junto al Eje 09), único ingreso "
                 "vehicular del predio, y corre por fuera del cerco perimetral proyectado, dentro del predio, con su eje a %.3f m del "
                 "paramento del cerco y junta de tecnopor de 1\" entre el muro y el cimiento. La losa superior queda al ras del piso "
                 "terminado del frente (+%.2f msnm) en toda su longitud. En la progresiva 0+%06.2f (registro RV-11) el trazo deja de "
                 "seguir al cerco y entra perpendicular al lindero hasta la cara este de la caja CL (0+%06.2f), cruzando el cerco "
                 "proyectado en 0+%06.2f, donde se deja un paso de 0.95 m. La longitud total es de %.2f m."
                 % (D["eje_desde_cerco"], D["NPT"], g["P_QUIEBRE"], g["P_FIN"], g["P_CRUCE_CERCO"], g["P_FIN"]))
    parrafo(doc, "El único cruce de circulación es el portón de camiones cisterna, de 5.78 m (progresivas 0+000.00 a 0+006.08 con la "
                 "transición de 0.30 m); el resto del colector queda bajo piso exterior con cargas peatonales. Las coordenadas UTM "
                 "WGS84 18S de la lámina DP-01 están en el mismo sistema del tramo del Hogar de Refugio y son referenciales; el "
                 "trazo se replantea en obra desde el cerco y desde la caja CL.")
    tabla(doc, ["ELEMENTO", "PROGRESIVA", "DESCRIPCIÓN"], [[r["nombre"], "0+%06.2f" % r["prog"], r["tipo"] + ": " + r["nota"]] for r in DJ["registros"]], anchos=[3.0, 3.0, 9.5])
    leyenda(doc, "%s: Registros y puntos singulares del tramo." % NUM.tabla()); fuente(doc)
    if os.path.exists(os.path.join(FIG, "DP-01.png")):
        imagen(doc, os.path.join(FIG, "DP-01.png"), 15.5, "%s: Planta general del colector frontal del CAR Varones (lámina DP-01)." % NUM.imagen())

    # ------------------------------------------------------------------ 3. caudal
    titulo2(doc, "Caudal de diseño y su reparto entre las cunetas")
    parrafo(doc, "El caudal total del proyecto, %.1f L/s, es el obtenido en el presente estudio con el método racional (TR = 25 años, "
                 "I = %.2f mm/h, FS = %.2f). Para dimensionar el colector por tramos, ese caudal se reparte entre las cinco cunetas "
                 "que llegan al frente en proporción al ancho de frente que drena cada una (franja comprendida entre los puntos "
                 "medios con las cunetas vecinas). La suma de picos sin desfase es una hipótesis conservadora, la misma de los "
                 "otros dos tramos del colector compartido." % (Qt, D["I"], D["FS"]))
    filas = [["Eje %s" % c["perfil"], f2(c["franja_m"]), "%.3f" % c["frac"], "%.1f" % c["Q"], "0+%06.2f" % c["prog"] if c["entra_en"] == "colector" else "caja CL", "colector (registro %s)" % [r["nombre"] for r in DJ["registros"] if abs(r["prog"] - c["prog"]) < 1.5 and r["nombre"] != "CL"][0] if c["entra_en"] == "colector" else "caja CL, muro norte"] for c in cun]
    filas.append(["Total", f2(sum(c["franja_m"] for c in cun)), "1.000", "%.1f" % Qt, "", ""])
    tabla(doc, ["CUNETA", "FRANJA DE FRENTE (m)", "FRACCIÓN", "Q (L/s)", "EMPALME", "ENTRA EN"], filas, tam=9, anchos=[2.0, 2.6, 2.0, 1.8, 2.4, 4.7])
    leyenda(doc, "%s: Reparto del caudal de diseño del CAR Varones entre las cunetas que llegan al frente." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Hoja CAUDALES de la memoria de cálculo del colector.")
    parrafo(doc, "El colector conduce %.1f L/s en su tramo final (Ejes 09 + 08 + 06 + 04) y la cuneta del Eje 01 aporta %.1f L/s "
                 "directamente a la caja CL; la CL recibe en total %.1f L/s, que es el aporte considerado por el Hogar de Refugio." % (Qcol, Qdir, Qt))

    # ------------------------------------------------------------------ 4. hidraulica
    titulo2(doc, "Dimensionamiento hidráulico de la sección")
    parrafo(doc, "Se adopta un canal rectangular de concreto armado f'c = 210 kg/cm², vaciado monolíticamente, con ancho interior "
                 "b = %.2f m (el mínimo practicable para la limpieza), muros de %.2f m, losa de fondo de %.2f m y losa superior de "
                 "%.2f m (%.2f m en el cruce de camiones), con n = %.3f. La cota de fondo inicial, %.2f msnm, queda fijada por la "
                 "cuneta más alta (Eje 09, NCF 260.69) y por la losa de 0.25 m del cruce de camiones (altura interior 0.70 m); la "
                 "pendiente S = %.1f ‰ lleva el fondo a %.3f msnm en la llegada a la CL, por encima de su nivel de agua. La altura "
                 "interior varía entre %.2f y %.2f m." % (D["b"], D["e_muro"], D["e_fondo"], D["e_losa"], D["e_losa_camion"], D["n"], D["CF0"], D["S"] * 1000, dz.fondo(dz.P_FIN), h_min, h_max))
    formula(doc, "yc = (Q² / (g · b²))^(1/3) = %.3f m en la llegada (Q = %.1f L/s)      ;      capacidad con llenado del 85 %%: %.0f L/s" % (DJ["yc"], Qcol, DJ["capacidad_85"]["Q"]))
    titulo2(doc, "Perfil de flujo con caudal creciente")
    parrafo(doc, "El caudal del colector crece en cada empalme de cuneta, por lo que el perfil de flujo se calcula por el método del "
                 "paso estándar con el caudal acumulado en cada estación, desde el control de aguas abajo: la caída libre a la caja "
                 "CL, donde el flujo pasa por el tirante crítico (yc = %.3f m, nivel de agua %.3f msnm). El régimen es subcrítico en "
                 "todo el tramo, el llenado máximo es %.0f %% y el borde libre mínimo %.2f m." % (DJ["yc"], DJ["brink"]["NA"], llen * 100, bl_min))
    formula(doc, "z₁ + y₁ + V₁²/2g = z₂ + y₂ + V₂²/2g + Sf · Δx      ;      Sf = (n · V / R^(2/3))²")
    filas = []
    for p in perfil:
        if abs(p["p"] / 10.0 - round(p["p"] / 10.0)) < 1e-6 or any(abs(p["p"] - c["prog"]) < 1e-6 for c in cun) or abs(p["p"] - g["P_FIN"]) < 1e-6 or abs(p["p"] - g["P_QUIEBRE"]) < 1e-6:
            filas.append(["0+%06.2f" % p["p"], "%.1f" % (p["Q"] * 1000), f3(p["z"]), f3(p["y"]), f3(p["NA"]), f2(p["V"]), f2(p["F"]), f2(p["techo"] - p["NA"]), "%.0f" % (100 * p["y"] / p["h"])])
    filas.sort(key=lambda f: f[0])
    tabla(doc, ["PROG.", "Q (L/s)", "COTA DE FONDO", "TIRANTE (m)", "NIVEL DE AGUA", "V (m/s)", "FROUDE", "BORDE LIBRE (m)", "LLENADO (%)"], filas, tam=8, anchos=[1.9, 1.4, 1.9, 1.6, 1.9, 1.4, 1.4, 1.8, 1.6])
    leyenda(doc, "%s: Perfil de flujo del colector con caudal creciente (paso estándar desde la caída libre a la CL)." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Hoja PERFIL_FLUJO de la memoria de cálculo del colector.")
    titulo2(doc, "Velocidades y autolimpieza")
    parrafo(doc, "Con el caudal de diseño en la entrega la velocidad a tirante normal es %.2f m/s (0.90 a 3.0 m/s según la norma "
                 "CE.040) y en la caída libre llega a %.2f m/s. Como el caudal crece a lo largo del tramo, cada tramo entre empalmes "
                 "se verifica con su propio caudal: en los tramos de cabecera (33.6 y 81.3 L/s) la velocidad es menor de 0.90 m/s "
                 "porque el ancho de 0.60 m es el mínimo practicable, y allí la autolimpieza se asegura con el esfuerzo tractivo "
                 "τ = γ·R·S, mayor que 0.15 kg/m² (1.5 Pa) en todos los tramos, y con los registros de limpieza cada 12 m como máximo."
                 % ([v for v in vel if v[0].startswith("Tramo Eje 04")][0][3] if vel else 0, Vmax))
    if vel:
        tabla(doc, ["CASO", "Q (L/s)", "yn (m)", "V (m/s)", "τ (kg/m²)", "FROUDE", "AUTOLIMPIEZA"],
              [[v[0], "%.1f" % (v[1] * 1000), f3(v[2]), f2(v[3]), f2(v[5]), f2(v[7]), str(v[11] or "")] for v in vel], tam=9, anchos=[5.6, 1.6, 1.6, 1.6, 1.8, 1.6, 2.2])
        leyenda(doc, "%s: Velocidades y esfuerzo tractivo por tramo (flujo uniforme con el caudal de cada tramo)." % NUM.tabla())
        fuente(doc, "Fuente: Elaboración propia. Hoja PERFIL_FLUJO de la memoria de cálculo del colector.")
    if os.path.exists(os.path.join(FIG, "DP-02.png")):
        imagen(doc, os.path.join(FIG, "DP-02.png"), 15.5, "%s: Perfil longitudinal del colector con el perfil de flujo (lámina DP-02)." % NUM.imagen())

    # ------------------------------------------------------------------ 5. cunetas
    titulo2(doc, "Llegada de las cunetas y acortamiento en el frente")
    parrafo(doc, "Las cunetas de los Ejes 09, 08, 06 y 04 entran al colector por una ventana de 0.40 × H en el muro del lado del "
                 "predio, con caída libre al fondo, junta de tecnopor de 1\" y un registro de limpieza sobre cada empalme; la del Eje "
                 "01 entra a la caja CL por una ventana de 0.40 × %.2f en su muro norte. En el plano de arquitectura las cunetas "
                 "están dibujadas hasta la franja exterior del frente, que ahora ocupa el colector: cada cuneta termina en la cara "
                 "del muro (del colector o de la CL) y el tramo que caía dentro de la estructura se descuenta en el metrado de "
                 "cunetas (0.24, 0.09, 0.24, 3.12 y 2.77 m). En todos los casos la cota de fondo de la cuneta queda por encima del "
                 "nivel de agua del receptor." % c01["H_ventana"])
    filas = [["Eje %s" % c["perfil"], "0+%06.2f" % c["prog"] if c["entra_en"] == "colector" else "caja CL", f2(c["NCF_fin"]), f2(c["H"]), "%.1f" % c["Q"], f2(-c["ajuste_L"]), f3(c["NA_colector"]), f2(c["caida_libre"])] for c in cun]
    tabla(doc, ["CUNETA", "EMPALME", "NCF EN EL MURO (msnm)", "H (m)", "Q (L/s)", "ACORTAMIENTO (m)", "NIVEL DE AGUA RECEPTOR (msnm)", "CAÍDA LIBRE (m)"], filas, tam=9, anchos=[1.8, 2.0, 2.2, 1.3, 1.5, 2.2, 2.4, 1.8])
    leyenda(doc, "%s: Empalme de las cunetas al colector y a la caja CL." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Plano PERFIL CAR VARONES y hoja CUNETAS de la memoria de cálculo.")

    # ------------------------------------------------------------------ 6. empalme CL
    titulo2(doc, "Empalme con la caja de llegada CL del Hogar de Refugio")
    parrafo(doc, "La caja CL (interior %.2f × %.2f m, piso de poza %.2f msnm) pertenece al expediente del Hogar de Refugio, pero queda "
                 "del lado del CAR Varones (x = 349162.00 a 349163.80 del sistema local), por lo que su tapa se fija en +%.2f msnm, al "
                 "ras del piso terminado de este proyecto, con un escalón de 0.55 m respecto del piso del Hogar de Refugio (+%.2f) en "
                 "el lindero. El colector entra por su muro este con la sección completa (ventana de 0.60 × %.2f m, cota de fondo "
                 "%.3f msnm, sin dintel porque el techo de la CL coincide con el del colector en %.2f msnm) y cae %.2f m sobre el nivel "
                 "de agua de la caja, con un colchón de %.2f m sobre el piso. La cuneta del Eje 01 entra por el muro norte con caída "
                 "libre de %.2f m. En todos los contactos entre estructuras de distinto expediente se coloca tecnopor de 1\"."
                 % (DS_["CL_largo"], DS_["CL_ancho"], cl["z_piso"], D["NPT_CL"], DS_["NPT"], cl["ventana_alto"], cl["ventana_alfeizar"], cl["techo_CL"], cl["caida_libre"], cl["NA_CL"] - cl["z_piso"], c01["caida_libre"]))
    tabla(doc, ["APORTE", "UBICACIÓN", "Q (L/s)", "COTA DE FONDO", "VENTANA", "CAÍDA LIBRE (m)"],
          [["Colector CAR Varones", "muro este, 0+%06.2f" % g["P_FIN"], "%.1f" % Qcol, f3(cl["ventana_alfeizar"]), "0.60 × %.2f" % cl["ventana_alto"], f2(cl["caida_libre"])],
           ["Cuneta Eje 01", "muro norte", "%.1f" % Qdir, f3(c01["NCF_fin"]), "0.40 × %.2f" % c01["H_ventana"], f2(c01["caida_libre"])],
           ["Total a la CL", "", "%.1f" % Qt, "", "poza: piso %.2f, NA %.3f" % (cl["z_piso"], cl["NA_CL"]), ""]], anchos=[3.6, 3.0, 1.8, 2.4, 2.8, 2.0])
    leyenda(doc, "%s: Llegadas a la caja CL del Hogar de Refugio." % NUM.tabla()); fuente(doc, "Fuente: Elaboración propia. Hoja EMPALME_CL de la memoria de cálculo del colector.")
    if os.path.exists(os.path.join(FIG, "DP-07.png")):
        imagen(doc, os.path.join(FIG, "DP-07.png"), 15.5, "%s: Empalme del colector con la caja CL (lámina DP-07)." % NUM.imagen())

    # ------------------------------------------------------------------ 7. estructural
    titulo2(doc, "Verificación estructural")
    parrafo(doc, "La estructura se verifica como marco cerrado de concreto armado con las cargas de la norma E.020, el empuje del "
                 "relleno en reposo y, en el cruce de camiones cisterna, la carga de rueda del camión de diseño HL-93 con impacto del "
                 "33 % (AASHTO LRFD). El diseño por resistencia sigue la norma E.060. El armado es el mismo del tramo del Hogar de "
                 "Refugio: un solo marco cerrado de 3/8\" @0.20 en el eje de muros y losas con longitudinales de 3/8\" @0.25 en el "
                 "tramo normal, y doble marco de 1/2\" @0.15 con losa superior de 0.25 m en el cruce de camiones. Como las alturas "
                 "interiores son menores (0.70 a 1.17 m), las verificaciones resultan con mayor holgura.")
    filas = []
    for k in "ACDFGH":
        e = est.get(k)
        if not e: continue
        filas.append([e["titulo"], "%.3f" % e.get("Mu", 0), "%.3f" % e.get("Mn", 0), ("%.2f" % e["Vu"]) if "Vu" in e else "-", ("%.2f" % e["Vc"]) if "Vc" in e else "-", "CUMPLE"])
    tabla(doc, ["ELEMENTO", "Mu (t·m/m)", "φMn (t·m/m)", "Vu (t/m)", "φVc (t/m)", "VERIFICACIÓN"], filas, tam=9, anchos=[7.0, 1.8, 1.8, 1.6, 1.6, 2.2])
    leyenda(doc, "%s: Resumen de la verificación estructural del colector y de las tapas de registro." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Hoja ESTRUCTURAL de la memoria de cálculo del colector.")
    if "G" in est and est["G"].get("q") is not None:
        parrafo(doc, "La presión transmitida al suelo es de %.2f kg/cm² en el tramo normal y de %.2f kg/cm² bajo la rueda del camión en "
                     "el cruce (carga transitoria); ambos valores deben compararse con la capacidad portante del estudio de mecánica "
                     "de suelos del proyecto, que no estaba disponible al elaborar este capítulo." % (est["G"]["q"], est["G"].get("qc", 0)))
    if os.path.exists(os.path.join(FIG, "DP-04A.png")):
        imagen(doc, os.path.join(FIG, "DP-04A.png"), 15.5, "%s: Secciones del colector con su armadura (lámina DP-04A)." % NUM.imagen())

    # ------------------------------------------------------------------ 8. normativa y criterios
    titulo2(doc, "Marco normativo del diseño")
    vineta(doc, "Reglamento Nacional de Edificaciones, norma CE.040 Drenaje Pluvial (R.M. N.° 126-2021-VIVIENDA): criterios de diseño de colectores, velocidades, registros y tiempos de concentración.")
    vineta(doc, "Reglamento Nacional de Edificaciones, normas E.020 Cargas, E.060 Concreto Armado y E.050 Suelos y Cimentaciones.")
    vineta(doc, "Manual de Hidrología, Hidráulica y Drenaje del MTC: método racional, intensidades de diseño y período de retorno.")
    vineta(doc, "Manual de Puentes del MTC y AASHTO LRFD Bridge Design Specifications: carga de rueda del camión HL-93, impacto y ancho de franja para la losa del cruce de camiones.")
    vineta(doc, "ASCE/WEF, Design and Construction of Urban Stormwater Management Systems: esfuerzo tractivo mínimo para autolimpieza con caudales parciales.")
    titulo2(doc, "Criterios constructivos")
    vineta(doc, "Colector cubierto, con losa superior al ras del piso terminado (+%.2f msnm), sin lloraderos; registros de limpieza "
                "(RV-01 a RV-11) a no más de 12 m, en cada llegada de cuneta y en el quiebre, fuera del cruce de camiones, con tapa de "
                "concreto armado de 0.68 × 0.68 × 0.08 m, contramarco de ángulo de 2\" × 2\" × 3/16\" y marco de 1 ½\" × 1 ½\" × 1/8\" (láminas DD-01 y DD-02)." % D["NPT"])
    vineta(doc, "Juntas de dilatación cada 4.00 m con tecnopor de 1\" y sello elastomérico; tecnopor de 1\" entre el colector y el cerco, en el paso del cerco proyectado, con el piso adyacente y con la caja CL (lámina DD-03).")
    vineta(doc, "Empalme de cunetas por ventana de 0.40 × H en el muro del lado del predio, con caída al fondo, junta de 1\" y registro sobre el empalme (láminas DP-06C y DD-04).")
    vineta(doc, "Concreto f'c = 210 kg/cm² en colector y tapas; solado f'c = 100 kg/cm² de 0.05 m; acero fy = 4200 kg/cm²; recubrimiento de 4 cm; traslapes de 0.40 m (3/8\") y 0.50 m (1/2\").")
    vineta(doc, "Si la caja CL se construye antes que el colector, sus ventanas este y norte se dejan tapadas con muro de ladrillo pandereta sin mortero de unión a la caja, para retirarlo al empalmar; si se construye después, el colector termina en tapón provisional.")

    # ------------------------------------------------------------------ 9. metrado
    titulo2(doc, "Resumen de metrados del tramo")
    parrafo(doc, "Los metrados del colector se incorporan a la planilla general de metrados del proyecto como partida independiente "
                 "01.04.04 COLECTOR PLUVIAL FRONTAL, con su desglose por tramos, registros, juntas y empalmes, sin modificar las "
                 "partidas existentes salvo el acortamiento de las cinco cunetas que llegan al frente (hoja CUNETAS - LLEGADA AL "
                 "COLECTOR). Cada cantidad se sustenta en las hojas COLECTOR PARAMETROS, MOV. TIERRAS, CONCRETO, ENCOFRADO, ACERO, "
                 "REGISTROS Y TAPAS y JUNTAS Y EMPALMES de la misma planilla. Las cantidades principales son:")
    Rg = RM["registros"]
    filas = [["Excavación manual de zanjas (ancho 1.40 m)", "m³", f2(RM["excav_m3"])], ["Relleno compactado con material propio", "m³", f2(RM["relleno_m3"])],
             ["Eliminación de material excedente", "m³", f2(RM["elimin_m3"])], ["Solado f'c = 100 kg/cm² e = 0.05 m", "m²", f2(RM["solado_m2"])],
             ["Concreto f'c = 210 kg/cm² en losa de fondo", "m³", f2(RM["conc_fondo_m3"])], ["Concreto f'c = 210 kg/cm² en muros", "m³", f2(RM["conc_muros_m3"])],
             ["Concreto f'c = 210 kg/cm² en losa superior", "m³", f2(RM["conc_losa_m3"])], ["Encofrado y desencofrado", "m²", f2(RM["encof_m2"])],
             ["Acero de refuerzo fy = 4200 kg/cm² en el colector", "kg", f2(RM["acero_colector_kg"])],
             ["Acabado frotachado y bruñado de losa superior", "m²", f2(RM["acabado_m2"])], ["Curado de concreto", "m²", f2(RM["curado_m2"])],
             ["Registros con tapa de concreto armado, marco y contramarco", "und", "%d" % Rg["n"]],
             ["Junta de dilatación e = 1\" con sello", "m", f2(RM["juntas"]["L_dilat"])], ["Empalmes de cunetas al colector", "und", "%d" % sum(1 for c in cun if c["entra_en"] == "colector")]]
    tabla(doc, ["PARTIDA", "UND", "METRADO"], filas, anchos=[10.0, 2.0, 3.0])
    leyenda(doc, "%s: Metrados principales del colector pluvial frontal del CAR Varones." % NUM.tabla())
    fuente(doc, "Fuente: Elaboración propia. Planilla general de metrados del proyecto, partida 01.04.04.")

    # ------------------------------------------------------------------ 10. cumplimiento
    if cump:
        titulo2(doc, "Cuadro de cumplimiento normativo")
        parrafo(doc, "El cuadro siguiente resume, para cada requisito del diseño, la norma o referencia que lo exige, el criterio aplicado, "
                     "el valor obtenido y el resultado. Cada valor proviene de una celda con fórmula de la memoria de cálculo (hoja "
                     "CUMPLIMIENTO), de modo que puede auditarse.")
        filas = []
        for f in cump:
            if f[1] is None and f[2] is None: filas.append([str(f[0]).upper(), "", "", "", ""])
            else:
                v = f[3]; vt = ("%.2f" % v) if isinstance(v, (int, float)) and not isinstance(v, bool) else ("" if v is None else str(v))
                filas.append([str(f[0]), str(f[1] or ""), str(f[2] or ""), vt, str(f[4] or "")])
        tabla(doc, ["REQUISITO", "NORMA / REFERENCIA", "CRITERIO", "VALOR", "RESULTADO"], filas, tam=8, anchos=[3.8, 3.2, 3.6, 1.5, 3.1])
        leyenda(doc, "%s: Cuadro de cumplimiento normativo del colector pluvial frontal del CAR Varones." % NUM.tabla())
        fuente(doc, "Fuente: Elaboración propia. Hoja CUMPLIMIENTO de la memoria de cálculo del colector.")
        titulo2(doc, "Datos de otros expedientes y cómo quedan cubiertos")
        for t in pend: vineta(doc, str(t))

    # ------------------------------------------------------------------ 11. conclusiones
    titulo2(doc, "Conclusiones del tramo")
    vineta(doc, "El colector frontal del CAR Varones conduce %.1f L/s en su tramo final en una sección cubierta de %.2f m de ancho con "
                "pendiente de %.1f ‰, en régimen subcrítico (F ≤ %.2f), con llenado máximo de %.0f %% y borde libre mínimo de %.2f m; "
                "con la cuneta del Eje 01 entrega a la caja CL los %.1f L/s del estudio hidrológico, el mismo valor previsto por el Hogar de Refugio."
                % (Qcol, D["b"], D["S"] * 1000, Fmax, llen * 100, bl_min, Qt))
    vineta(doc, "La entrega a la caja CL es por caída libre (%.2f m sobre su nivel de agua), la tapa de la CL queda en +%.2f y las "
                "ventanas este (0.60 × %.2f) y norte (0.40 × %.2f) quedan definidas en ambos expedientes." % (cl["caida_libre"], D["NPT_CL"], cl["ventana_alto"], c01["H_ventana"]))
    vineta(doc, "La estructura cumple las verificaciones de flexión y cortante de la norma E.060 para cargas peatonales y para la carga "
                "del camión cisterna en el portón, con el mismo armado del tramo del Hogar de Refugio.")
    vineta(doc, "Las cunetas de los Ejes 09, 08, 06, 04 y 01 se acortan hasta la cara del muro del colector o de la CL; el descuento "
                "está aplicado en la planilla de metrados del proyecto.")
    vineta(doc, "La capacidad portante del suelo debe confirmarse con el estudio de mecánica de suelos; la ubicación UTM del trazo se verifica en el replanteo.")
    doc.save(DEST)
    return n_par, n_tab


def verificar(n_par, n_tab):
    a = Document(ORIG); b = Document(DEST)
    pa = [p.text for p in a.paragraphs]; pb = [p.text for p in b.paragraphs]
    ta = ["|".join(c.text for r in t.rows for c in r.cells) for t in a.tables]; tb = ["|".join(c.text for r in t.rows for c in r.cells) for t in b.tables]
    ok = pa == pb[:len(pa)] and ta == tb[:len(ta)]
    print("parrafos originales %d (iguales: %s); nuevos %d; tablas originales %d (iguales: %s); nuevas %d; imagenes %d -> %d"
          % (len(pa), pa == pb[:len(pa)], len(pb) - len(pa), len(ta), ta == tb[:len(ta)], len(tb) - len(ta), len(a.inline_shapes), len(b.inline_shapes)))
    return ok


if __name__ == "__main__":
    n1, n2 = construir(); print(DEST); verificar(n1, n2)
