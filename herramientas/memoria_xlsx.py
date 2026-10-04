"""Memoria de calculo en Excel (con formulas) del colector pluvial frontal, tramo
Hogar de Refugio Temporal Mujeres Violentadas (CUI 2675514).

Mismas hojas y convenciones que la memoria del tramo CAR Mujeres:
  MEMORIA, DATOS, CAUDALES, CUNETAS, EMPALME, PERFIL_FLUJO, ESTRUCTURAL.
Las formulas se guardan con nombres en ingles (SUM, IF, ROUND...) y separador ",":
Excel en espanol las muestra como SUMA, SI, REDONDEAR con ";".
Azul = dato editable; negro = formula; verde = vinculo a otra hoja.
"""
import os, sys, json, re
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter as L

AQUI = os.path.dirname(os.path.abspath(__file__)); RAIZ = os.path.dirname(AQUI)
sys.path.insert(0, AQUI)
import diseno as dz

AZUL = Font(name="Arial", size=9, color="0000FF")
NEGRO = Font(name="Arial", size=9)
VERDE = Font(name="Arial", size=9, color="008000")
NEG = Font(name="Arial", size=9, bold=True)
TIT = Font(name="Arial", size=12, bold=True)
SUB = Font(name="Arial", size=9, italic=True, color="555555")
GRIS = PatternFill("solid", fgColor="EEEEEE")
BORDE = Border(*(Side(style="thin", color="999999"),) * 4)
ROW = {}   # fila de cada simbolo en la hoja DATOS
SUBTITULO = "Colector pluvial frontal - tramo Hogar de Refugio Temporal Mujeres Violentadas (CUI 2675514), Morales - San Martin. Hidroconsult, Oct. 2026"


def celda(ws, ref, valor, font=NEGRO, fmt=None, bold=False, fill=None, al=None):
    c = ws[ref]; c.value = valor; c.font = NEG if bold else font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if al: c.alignment = Alignment(horizontal=al, vertical="center", wrap_text=True)
    return c


def fila(ws, r, valores, fonts=None, fmts=None):
    for j, v in enumerate(valores):
        if v is None: continue
        c = ws.cell(row=r, column=j + 1, value=v)
        f = (fonts[j] if fonts and j < len(fonts) and fonts[j] else None)
        c.font = f or (AZUL if isinstance(v, (int, float)) and not isinstance(v, bool) else NEGRO)
        if fmts and j < len(fmts) and fmts[j]: c.number_format = fmts[j]


def cabecera(ws, titulo, anchos):
    celda(ws, "A1", titulo, TIT); celda(ws, "A2", SUBTITULO, SUB)
    for i, a in enumerate(anchos): ws.column_dimensions[L(i + 1)].width = a


def encabezado_tabla(ws, r, textos):
    for j, t in enumerate(textos):
        c = ws.cell(row=r, column=j + 1, value=t); c.font = NEG; c.fill = GRIS; c.border = BORDE
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ----------------------------------------------------------------------------
def hoja_datos(wb, R):
    ws = wb.create_sheet("DATOS"); D = dz.D
    cabecera(ws, "1. DATOS DE DISENO", [46, 12, 8, 70, 8])
    encabezado_tabla(ws, 4, ["PARAMETRO", "VALOR", "UND", "FUENTE / CRITERIO", "SIMBOLO"])
    filas = [
        ("HIDROLOGIA",),
        ("Periodo de retorno", D["TR"], "anos", "Analisis estadistico de precipitaciones (memoria HIDRO-CE040 del proyecto)", "TR"),
        ("Precipitacion de diseno D=10 min", D["P10"], "mm", "Dick-Peschke, P24 = 121.81 mm", "P10"),
        ("Precipitacion de diseno D=20 min", D["P20"], "mm", "Dick-Peschke", "P20"),
        ("Tiempo de concentracion en el colector", D["tc"], "min", "10 min de entrada + recorrido en cunetas y colector", "tc"),
        ("Exponente de la curva P-D", "=LN(B8/B7)/LN(2)", "-", "b = ln(P20/P10)/ln 2", "bx"),
        ("Intensidad de diseno I", "=B7*(B9/10)^B10/(B9/60)", "mm/h", "I = P10 (tc/10)^b / (tc/60)", "I"),
        ("Factor de seguridad FS", D["FS"], "-", "Memorias de calculo HIDRO-CE040", "FS"),
        ("Intensidad para D = 10 min (cunetas y montantes del estudio hidrologico)", "=B7*60/10", "mm/h", "I10 = P10 x 60 / 10 = 210.98 mm/h: es la misma lluvia de 35.163 mm en 10 min, expresada como intensidad (igual al estudio hidrologico)", "I10"),
        ("Intensidad para D = 20 min", "=B8*60/20", "mm/h", "I20 = P20 x 60 / 20 = 125.45 mm/h (analisis estadistico, hoja C. P diseno)", "I20"),
        ("Nota: el colector tiene tc = 15 min (mayor que los 10 min de cada cuneta), por eso su intensidad es 155.66 mm/h, interpolada en la curva P-D entre 10 y 20 min; es el mismo criterio del colector del CAR Mujeres (3 proyectos).",),
        ("GEOMETRIA DEL COLECTOR (tramo Hogar de Refugio)",),
        ("Ancho interior b", D["b"], "m", "Adoptado: llenado <= 85 % y regimen subcritico (hoja PERFIL_FLUJO)", "b"),
        ("Espesor de muros", D["e_muro"], "m", "Muros de 1.40 a 1.65 m de altura interior (hoja ESTRUCTURAL)", "em"),
        ("Espesor de losa de fondo", D["e_fondo"], "m", "", "ef"),
        ("Espesor de losa superior (tramo normal y porton de motos)", D["e_losa"], "m", "Vaciada monoliticamente con los muros", "et"),
        ("Espesor de losa superior en cruce de camiones", D["e_losa_camion"], "m", "Ingreso vehicular OE-01 (cisterna al PTAR)", "ec"),
        ("Cota de fondo inicial (0+000, caja de llegada)", D["CF0"], "msnm", "Fijada por las cunetas de arquitectura: la mas baja llega con NCF 259.70 (Eje 02)", "CF0"),
        ("Pendiente S", D["S"], "m/m", "Adoptada 0.30 %: regimen subcritico (F = 0.60) y fondo sobre el R-01 del receptor", "S"),
        ("Progresiva del quiebre 1 (45 grados)", round(dz.P_B1, 2), "m", "Salida del frente hacia el lindero", "pq1"),
        ("Progresiva del quiebre 2 (45 grados)", round(dz.P_B2, 2), "m", "", "pq2"),
        ("Progresiva del brink de la caida (fin del colector b=0.80)", round(dz.P_BRINK, 2), "m", "Inicio de la poza de disipacion", "pb"),
        ("Progresiva de entrega (0+000 del tramo CAR Mujeres)", round(dz.P_FIN, 2), "m", "Cara aguas arriba del registro R-01 del receptor", "pf"),
        ("Longitud total del tramo", "=B23", "m", "", "L"),
        ("NPT del frente (cara superior de la losa)", D["NPT"], "msnm", "Arquitectura: NT 260.50 + NPT 0.10 (estacionamientos e ingresos)", "NPT"),
        ("Cruce de camiones: inicio", R["zonas"][0]["p1"], "m", R["zonas"][0]["nombre"] + " + 0.30 m a cada lado", "pc1"),
        ("Cruce de camiones: fin", R["zonas"][0]["p2"], "m", "", "pc2"),
        ("Cruce de motos: inicio", R["zonas"][1]["p1"], "m", R["zonas"][1]["nombre"] + ", tramo de 2.00 m", "pm1"),
        ("Cruce de motos: fin", R["zonas"][1]["p2"], "m", "", "pm2"),
        ("Coeficiente de Manning n", D["n"], "-", "Concreto vaciado in situ (Chow)", "n"),
        ("CRITERIOS DE VERIFICACION HIDRAULICA",),
        ("Llenado maximo admisible y/h (conducto cubierto)", D["llenado_max"], "-", "Criterio de conducto cerrado a superficie libre", "llen"),
        ("Borde libre minimo bajo losa", D["BL_min"], "m", "Complementario al criterio de llenado", "BLmin"),
        ("RECEPTOR: COLECTOR DEL CAR MUJERES (CUI 2717013)",),
        ("Cota de fondo del receptor en su 0+000 (R-01)", D["CF_R01"], "msnm", "Plano DP-01 del tramo CAR Mujeres", "CFr"),
        ("Nivel de agua del receptor en su 0+000", D["NA_R01"], "msnm", "Memoria del tramo CAR Mujeres (Q = 626 L/s en 0+000)", "NAr"),
        ("Ancho interior del receptor", D["b_wilma"], "m", "", "br"),
        ("APORTE EXTERNO: COLECTOR DEL CAR VARONES (CUI 2705619)",),
        ("Cota de fondo de llegada del colector del CAR Varones (cara este de la CL)", D["CF_varones_sup"], "msnm", "Expediente del CAR Varones (CUI 2705619): 260.20 - 0.3 % x 105.29 m; ventana este de la CL de 0.60 x 1.17", "CFv"),
        ("Cota de la tapa de la caja CL (piso terminado del CAR Varones)", D["NPT_CL"], "msnm", "La CL queda del lado del CAR Varones (x = 349162.00 a 349163.80); su tapa va al ras de ese piso (+261.15)", "CLt"),
        ("Cuneta Eje 01 del CAR Varones: cota de fondo en el muro norte de la CL", D["E01_NCF"], "msnm", "Entra directamente a la CL por ventana 0.40 x 0.66 (37.0 L/s de los 258.7)", "E01"),
        ("CAJAS",),
        ("Caja de llegada CL: profundidad de poza bajo el fondo", D["CL_poza"], "m", "Colchon de agua para la caida del aporte externo", "pcl"),
        ("Caja de caida CC: profundidad de poza bajo el fondo del receptor", D["CC_poza_prof"], "m", "Poza de disipacion deprimida", "pcc"),
        ("Caja de caida CC: longitud de poza", D["CC_poza_largo"], "m", "", "Lcc"),
        ("Caja de caida CC: ancho de poza", D["CC_ancho"], "m", "Igual al ancho del receptor", "bcc"),
    ]
    r = 5
    for f in filas:
        if len(f) == 1:
            celda(ws, f"A{r}", f[0], NEG, fill=GRIS)
        else:
            fila(ws, r, list(f), fonts=[NEGRO, None, NEGRO, NEGRO, NEGRO])
            if isinstance(f[1], str) and f[1].startswith("="): ws[f"B{r}"].font = NEGRO
            ROW[f[4]] = r
        r += 1
    ws[f"B{ROW['L']}"] = f"=B{ROW['pf']}"
    celda(ws, f"A{r+1}", "Azul = dato editable; negro = formula; verde = vinculo.", SUB)
    return ws


def hoja_caudales(wb, R):
    ws = wb.create_sheet("CAUDALES"); D = dz.D
    cabecera(ws, "2. CAUDALES DE DISENO - METODO RACIONAL Q = C I A FS", [44, 12, 12, 12, 10, 10, 34, 34])
    encabezado_tabla(ws, 4, ["PROYECTO", "AREA (m2)", "C x A (m2)", "C PONDERADO", "I (mm/h)", "Q (L/s)", "INGRESA AL COLECTOR EN", "FUENTE"])
    fila(ws, 5, ["CAR Varones (CUI 2705619) - aporte externo", D["A_varones"], D["CA_varones"], "=C5/B5", "=DATOS!$B$11", "=C5*E5/3600000*DATOS!$B$12*1000", "Caja de llegada CL (0+000)", "Memoria HIDRO-CE040 del proyecto, dato no validado"],
         fonts=[NEGRO, AZUL, AZUL, NEGRO, VERDE, NEGRO, NEGRO, NEGRO], fmts=[None, "0.00", "0.0000", "0.000", "0.00", "0.0"])
    fila(ws, 6, ["Hogar de Refugio Temporal M. Violentadas (CUI 2675514)", D["A_refugio"], D["CA_refugio"], "=C6/B6", "=DATOS!$B$11", "=C6*E6/3600000*DATOS!$B$12*1000", "Cunetas Ejes 01, 02, 06, 07, 11 y 12", "Memoria HIDRO-CE040, 12 fichas"],
         fonts=[NEGRO, AZUL, AZUL, NEGRO, VERDE, NEGRO, NEGRO, NEGRO], fmts=[None, "0.00", "0.0000", "0.000", "0.00", "0.0"])
    fila(ws, 7, ["TOTAL ENTREGADO AL CAR MUJERES", "=SUM(B5:B6)", "=SUM(C5:C6)", "=C7/B7", None, "=SUM(F5:F6)"], fonts=[NEG, NEGRO, NEGRO, NEGRO, None, NEG], fmts=[None, "0.00", "0.0000", "0.000", None, "0.0"])
    celda(ws, "A9", "Caudal de diseno adoptado (valores de las memorias de cada proyecto)", NEG)
    fila(ws, 10, ["CAR Varones", D["Q_varones"], "L/s"], fonts=[NEGRO, AZUL, NEGRO])
    fila(ws, 11, ["Hogar de Refugio", D["Q_refugio"], "L/s"], fonts=[NEGRO, AZUL, NEGRO])
    fila(ws, 12, ["Q de diseno del tramo", "=B10+B11", "L/s"], fonts=[NEG, NEG, NEGRO]); ws["B12"].number_format = "0.0"
    fila(ws, 13, ["Q de diseno del tramo", "=B12/1000", "m3/s"], fonts=[NEGRO, NEGRO, NEGRO]); ws["B13"].number_format = "0.0000"
    celda(ws, "A15", "Caudal que circula por el tramo (de aguas arriba hacia aguas abajo)", NEG)
    encabezado_tabla(ws, 16, ["PUNTO", "PROGRESIVA", "APORTE (L/s)", "Q ACUMULADO (L/s)"])
    pts = [("Caja de llegada CL (CAR Varones)", 0.0, "=B10")]
    for c in R["cunetas"]:
        pts.append((f"Cuneta {c['nombre']} - perfil {c['perfil']}", c["prog"], "=B11/6"))
    r = 17
    for i, (nm, p, ap) in enumerate(pts):
        fila(ws, r, [nm, p, ap, f"=SUM($C$17:C{r})"], fonts=[NEGRO, AZUL, NEGRO, NEGRO], fmts=[None, "0.00", "0.0", "0.0"]); r += 1
    celda(ws, f"A{r+1}", "Suma de picos sin desfase de tiempos y con FS = 1.15: hipotesis conservadora. El reparto del aporte del Refugio entre sus 6 cunetas es nominal; el perfil se verifica con el caudal total en todo el tramo.", SUB)
    return ws


def hoja_cunetas(wb, R):
    ws = wb.create_sheet("CUNETAS")
    cabecera(ws, "3. CUNETAS DE ARQUITECTURA QUE LLEGAN AL COLECTOR", [30, 10, 10, 10, 10, 10, 12, 12, 14, 14, 14, 30])
    encabezado_tabla(ws, 4, ["CUNETA (perfil de arquitectura)", "PERFIL", "PROGRESIVA (m)", "LONGITUD (m)", "NCT (msnm)", "NCF al cerco (msnm)", "H al cerco (m)", "PROLONG. (m)", "NCF en el muro (msnm)", "NA del colector (msnm)", "CAIDA LIBRE (m)", "VERIFICACION"])
    r = 5
    for c in R["cunetas"]:
        fila(ws, r, [c["nombre"], c["perfil"], c["prog"], c["L"], c["NCT"], c["NCF"], f"=E{r}-F{r}", c["prolong"], f"=F{r}-0.005*H{r}",
                     f"=IFERROR(INDEX(PERFIL_FLUJO!$G$6:$G$200,MATCH(C{r},PERFIL_FLUJO!$A$6:$A$200,0)),\"\")", f"=I{r}-J{r}",
                     f"=IF(K{r}>0,\"CAIDA LIBRE\",\"AHOGADA\")"],
             fonts=[NEGRO, AZUL, AZUL, AZUL, AZUL, AZUL, NEGRO, AZUL, NEGRO, VERDE, NEGRO, NEGRO],
             fmts=[None, None, "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.000", "0.000", "0.000"])
        r += 1
    celda(ws, f"A{r+1}", "Seccion de las cunetas: 0.40 x H, muros 0.10, pendiente 0.5 % (plano de arquitectura, perfiles 01 a 12). Las prolongaciones fuera del cerco mantienen la misma pendiente. El empalme es por ventana 0.40 x H en el muro lado predio, con caida al fondo del colector y registro encima (lamina DP-06).", SUB)
    return ws


def hoja_empalme(wb, R):
    ws = wb.create_sheet("EMPALME"); D = dz.D
    cabecera(ws, "4. CAJA DE LLEGADA (0+000) Y CAJA DE CAIDA EN LA ENTREGA AL CAR MUJERES", [52, 12, 8, 60])
    encabezado_tabla(ws, 4, ["PARAMETRO", "VALOR", "UND", "FORMULA / CRITERIO"])
    rows = [
        ("A. CAJA DE CAIDA CC - ENTREGA AL REGISTRO R-01 DEL CAR MUJERES",),
        ("Caudal Q", "=CAUDALES!B13", "m3/s", ""),
        ("Ancho del colector en el brink b", "=DATOS!B14", "m", ""),
        ("Cota de fondo en el brink", "=DATOS!B19-DATOS!B20*DATOS!B23", "msnm", "CF0 - S x pb"),
        ("Tirante critico en el brink yc", "=(B6^2/(9.81*B7^2))^(1/3)", "m", "yc = (Q2/(g b2))^(1/3): control de la caida libre"),
        ("Nivel de agua en el brink", "=B8+B9", "msnm", "nivel de control del perfil de flujo (hoja PERFIL_FLUJO)"),
        ("Piso de la poza", "=DATOS!B37-DATOS!B45", "msnm", "fondo del receptor - profundidad de poza"),
        ("Caida total (brink a piso de poza)", "=B8-B11", "m", ""),
        ("Energia al pie respecto al piso E1", "=B12+1.5*B9", "m", "E1 = caida + 1.5 yc (sin perdidas, conservador)"),
        ("Tirante al pie y1", 0.0, "m", "resuelto por iteracion (valor de la memoria en Python); verificar en la ultima fila"),
        ("Velocidad al pie V1", "=B6/(B7*B14)", "m/s", "V1 = Q/(b y1)"),
        ("Froude al pie F1", "=B15/SQRT(9.81*B14)", "-", ""),
        ("Tirante conjugado y2", "=B14/2*(SQRT(1+8*B16^2)-1)", "m", "resalto hidraulico en el ancho b (conservador: no considera la expansion a 1.50 m)"),
        ("Tirante disponible en la poza", "=DATOS!B38-B11", "m", "NA del receptor - piso de la poza"),
        ("Verificacion del resalto", "=IF(B18>=B17,\"RESALTO AHOGADO - CUMPLE\",\"REVISAR PROFUNDIDAD DE POZA\")", "", "el receptor ahoga el resalto dentro de la poza"),
        ("Longitud de resalto 6 y2", "=6*B17", "m", "referencia (resalto libre)"),
        ("Longitud de poza adoptada", "=DATOS!B46", "m", "con umbral de salida de 0.40 m hasta el fondo del receptor; resalto ahogado"),
        ("Verificacion de longitud", "=IF(B21>=0.8*B20,\"CUMPLE (resalto ahogado)\",\"AMPLIAR POZA\")", "", "poza >= 0.8 x 6 y2 para resalto ahogado (USBR)"),
        ("Nivel de agua en el receptor", "=DATOS!B38", "msnm", "regimen subcritico aguas abajo del umbral (F = 0.67 en el receptor)"),
        ("",),
        ("B. CAJA DE LLEGADA CL (0+000) - CAIDA DEL APORTE DE CAR VARONES",),
        ("Cota de fondo de llegada del aporte externo", "=DATOS!B41", "msnm", "expediente del CAR Varones (CUI 2705619): colector b = 0.60, fondo 259.884 en la cara este de la CL"),
        ("Piso de la poza de la caja", "=DATOS!B19-DATOS!B44", "msnm", "CF0 - profundidad de poza"),
        ("Caida del aporte externo", "=B26-B27", "m", ""),
        ("Nivel de agua en la caja (= inicio del colector)", "=PERFIL_FLUJO!G200", "msnm", "nivel de agua en 0+000 (hoja PERFIL_FLUJO, ultima fila)"),
        ("Colchon de agua en la poza", "=B29-B27", "m", "amortigua el chorro de 259 L/s"),
        ("Verificacion", "=IF(B30>=0.5,\"CUMPLE (colchon >= 0.50 m)\",\"REVISAR\")", "", ""),
        ("Dimensiones interiores de la caja", "1.50 x 1.00", "m", "largo x ancho; profundidad hasta el piso de poza"),
    ]
    r = 5
    for f in rows:
        if len(f) == 1:
            if f[0]: celda(ws, f"A{r}", f[0], NEG, fill=GRIS)
        else:
            fila(ws, r, list(f), fonts=[NEGRO, NEGRO, NEGRO, NEGRO])
            if isinstance(f[1], str) and f[1].startswith("=DATOS") or (isinstance(f[1], str) and f[1].startswith("=CAUDALES")) or (isinstance(f[1], str) and f[1].startswith("=PERFIL")):
                ws[f"B{r}"].font = VERDE
            ws[f"B{r}"].number_format = "0.000"
        r += 1
    # y1 por iteracion circular: B14 y B15 se referencian mutuamente; ponemos valor inicial por bisección en Python
    ws["B14"].value = round(R["poza"]["y1"], 4); ws["B14"].font = AZUL
    ws["D14"].value = "y1 resuelto en Python (y1 + (Q/(b y1))^2/2g = E1); verificar: B14 + B15^2/19.62 = B13"
    celda(ws, f"A{r}", "Verificacion de y1: y1 + V1^2/2g", NEGRO); celda(ws, f"B{r}", "=B14+B15^2/(2*9.81)", NEGRO, "0.000"); celda(ws, f"C{r}", "m", NEGRO); celda(ws, f"D{r}", "debe ser igual a E1 (B13)", NEGRO)
    return ws


def hoja_perfil(wb, R):
    ws = wb.create_sheet("PERFIL_FLUJO"); D = dz.D
    cabecera(ws, "5. PERFIL DE FLUJO GRADUALMENTE VARIADO - METODO DEL PASO ESTANDAR (de aguas abajo hacia aguas arriba)", [11, 10, 9, 8, 8, 9, 10, 9, 8, 10, 9, 9, 9, 13, 10])
    celda(ws, "A3", "Ecuacion de energia entre estaciones: y1 + V1^2/2g + z1 = y2 + V2^2/2g + z2 + dx (Sf1 + Sf2)/2 ; Sf = (n V / R^(2/3))^2 ; control: tirante critico en el brink de la caida (0+%06.2f)" % dz.P_BRINK, SUB)
    enc = ["PROGRESIVA (m)", "COTA FONDO", "Q (m3/s)", "b (m)", "yc (m)", "TIRANTE y (m)", "NIVEL DE AGUA", "VELOCIDAD (m/s)", "FROUDE", "COTA BAJO LOSA", "ALTURA INTERIOR h", "LLENADO y/h", "BORDE LIBRE (m)", "VERIFICACION", "ZONA"]
    encabezado_tabla(ws, 5, enc)
    NIT = 30
    perf = R["perfil"][::-1]            # de aguas abajo (brink) hacia aguas arriba
    r0 = 6
    for k, e in enumerate(perf):
        r = r0 + k
        ws.cell(row=r, column=1, value=round(e["p"], 2)).font = AZUL
        ws.cell(row=r, column=2, value=f"=DATOS!$B$19-DATOS!$B$20*A{r}").font = NEGRO
        ws.cell(row=r, column=3, value="=CAUDALES!$B$13").font = VERDE
        ws.cell(row=r, column=4, value="=DATOS!$B$14").font = VERDE
        ws.cell(row=r, column=5, value=f"=(C{r}^2/(9.81*D{r}^2))^(1/3)").font = NEGRO
        if k == 0:
            ws.cell(row=r, column=6, value=f"=E{r}").font = NEGRO
        else:
            ws.cell(row=r, column=6, value=f"=({L(16+2*NIT-2)}{r}+{L(16+2*NIT-1)}{r})/2").font = NEGRO
        ws.cell(row=r, column=7, value=f"=B{r}+F{r}").font = NEGRO
        ws.cell(row=r, column=8, value=f"=C{r}/(D{r}*F{r})").font = NEGRO
        ws.cell(row=r, column=9, value=f"=H{r}/SQRT(9.81*F{r})").font = NEGRO
        ws.cell(row=r, column=10, value=f"=DATOS!$B$25-IF(AND(A{r}>=DATOS!$B$26,A{r}<DATOS!$B$27),DATOS!$B$18,DATOS!$B$17)").font = NEGRO
        ws.cell(row=r, column=11, value=f"=J{r}-B{r}").font = NEGRO
        ws.cell(row=r, column=12, value=f"=F{r}/K{r}").font = NEGRO
        ws.cell(row=r, column=13, value=f"=J{r}-G{r}").font = NEGRO
        ws.cell(row=r, column=14, value=f"=IF(AND(L{r}<=DATOS!$B$33,M{r}>=DATOS!$B$34,I{r}<=1),\"CUMPLE\",\"NO CUMPLE\")").font = NEGRO
        ws.cell(row=r, column=15, value=f"=IF(AND(A{r}>=DATOS!$B$26,A{r}<DATOS!$B$27),\"cruce de camiones\",IF(AND(A{r}>=DATOS!$B$28,A{r}<DATOS!$B$29),\"cruce de motos\",\"tramo normal\"))").font = NEGRO
        for c, f in zip(range(2, 14), ["0.000", "0.0000", "0.00", "0.000", "0.000", "0.000", "0.00", "0.00", "0.000", "0.00", "0.00", "0.000"]):
            ws.cell(row=r, column=c).number_format = f
        if k > 0:
            rp = r - 1  # fila aguas abajo (ya resuelta)
            # columnas de iteracion: P=lo0, Q=hi0, R=lo1, S=hi1, ...
            dx = f"(A{rp}-A{r})"
            E2 = f"(G{rp}+H{rp}^2/(2*9.81))"
            Sf2 = f"((DATOS!$B$32*H{rp}/((D{rp}*F{rp})/(D{rp}+2*F{rp}))^(2/3))^2)"
            ws.cell(row=r, column=16, value=f"=E{r}*1.001").font = NEGRO
            ws.cell(row=r, column=17, value=3).font = AZUL
            for it in range(1, NIT):
                lo, hi = L(16 + 2 * (it - 1)), L(17 + 2 * (it - 1))
                mid = f"(({lo}{r}+{hi}{r})/2)"
                fmid = (f"({mid}+(C{r}/(D{r}*{mid}))^2/(2*9.81)+B{r}-{E2}-{dx}*(((DATOS!$B$32*(C{r}/(D{r}*{mid}))/((D{r}*{mid})/(D{r}+2*{mid}))^(2/3))^2+{Sf2})/2))")
                ws.cell(row=r, column=16 + 2 * it, value=f"=IF({fmid}>0,{lo}{r},{mid})").font = NEGRO
                ws.cell(row=r, column=17 + 2 * it, value=f"=IF({fmid}>0,{mid},{hi}{r})").font = NEGRO
    rlast = r0 + len(perf) - 1
    celda(ws, "P4", "Biseccion del tirante y (lo, hi) en %d iteraciones: f(y) = y + V^2/2g + z - E2 - dx (Sf + Sf2)/2" % NIT, SUB)
    # resumen
    rr = rlast + 2
    celda(ws, f"A{rr}", "RESUMEN", NEG)
    celda(ws, f"A{rr+1}", "Llenado maximo y/h"); celda(ws, f"B{rr+1}", f"=MAX(L{r0}:L{rlast})", NEGRO, "0.000")
    celda(ws, f"A{rr+2}", "Borde libre minimo (m)"); celda(ws, f"B{rr+2}", f"=MIN(M{r0}:M{rlast})", NEGRO, "0.000")
    celda(ws, f"A{rr+3}", "Froude maximo en el colector (sin el brink)"); celda(ws, f"B{rr+3}", f"=MAX(I{r0+1}:I{rlast})", NEGRO, "0.000")
    celda(ws, f"A{rr+4}", "Estaciones que no cumplen"); celda(ws, f"B{rr+4}", f"=COUNTIF(N{r0}:N{rlast},\"NO CUMPLE\")", NEGRO)
    celda(ws, f"A{rr+5}", "Nivel de agua en 0+000 (msnm)"); celda(ws, f"B{rr+5}", f"=G{rlast}", NEGRO, "0.000")
    # verificacion de velocidades y autolimpieza
    rv = rr + 7
    celda(ws, f"A{rv}", "VERIFICACION DE VELOCIDADES Y AUTOLIMPIEZA (RNE CE.040 Drenaje Pluvial; esfuerzo tractivo minimo)", NEG)
    celda(ws, f"A{rv+1}", "Tirante normal por punto fijo: y(k+1) = (Q n (b + 2 y(k))^(2/3) / (b^(5/3) S^(1/2)))^(3/5), 8 iteraciones (columnas P a W). Esfuerzo tractivo tau = gamma R S.", SUB)
    encabezado_tabla(ws, rv + 2, ["CASO", "Q (m3/s)", "yn (m)", "V (m/s)", "R (m)", "TAU (kg/m2)", "TAU (Pa)", "FROUDE", "V >= 0.90 m/s (CE.040)", "TAU >= 0.15 kg/m2", "V <= 3.0 m/s (concreto)"])
    casos = [("Caudal de diseno del tramo (Varones + Refugio)", "=CAUDALES!$B$13"), ("Solo aporte de CAR Varones (CUI 2705619)", "=CAUDALES!$B$10/1000"),
             ("Solo aporte del Hogar de Refugio (CUI 2675514)", "=CAUDALES!$B$11/1000"), ("50 % del caudal de diseno", "=0.50*CAUDALES!$B$13"),
             ("25 % del caudal de diseno", "=0.25*CAUDALES!$B$13"), ("10 % del caudal de diseno (lluvia menor)", "=0.10*CAUDALES!$B$13"), ("5 % del caudal de diseno", "=0.05*CAUDALES!$B$13")]
    for k, (nm, q) in enumerate(casos):
        r = rv + 3 + k
        ws.cell(row=r, column=1, value=nm).font = NEGRO
        ws.cell(row=r, column=2, value=q).font = VERDE; ws.cell(row=r, column=2).number_format = "0.0000"
        ws.cell(row=r, column=16, value=f"=(B{r}*DATOS!$B$32/(DATOS!$B$14*SQRT(DATOS!$B$20)))^(3/5)").font = NEGRO
        for it in range(1, 8):
            prev = L(16 + it - 1)
            ws.cell(row=r, column=16 + it, value=f"=(B{r}*DATOS!$B$32*(DATOS!$B$14+2*{prev}{r})^(2/3)/(DATOS!$B$14^(5/3)*SQRT(DATOS!$B$20)))^(3/5)").font = NEGRO
        for c in range(16, 24): ws.cell(row=r, column=c).number_format = "0.0000"
        ws.cell(row=r, column=3, value=f"={L(23)}{r}").font = NEGRO
        ws.cell(row=r, column=4, value=f"=B{r}/(DATOS!$B$14*C{r})").font = NEGRO
        ws.cell(row=r, column=5, value=f"=DATOS!$B$14*C{r}/(DATOS!$B$14+2*C{r})").font = NEGRO
        ws.cell(row=r, column=6, value=f"=1000*E{r}*DATOS!$B$20").font = NEGRO
        ws.cell(row=r, column=7, value=f"=F{r}*9.81").font = NEGRO
        ws.cell(row=r, column=8, value=f"=D{r}/SQRT(9.81*C{r})").font = NEGRO
        ws.cell(row=r, column=9, value=(f"=IF(D{r}>=0.9,\"CUMPLE\",\"NO CUMPLE\")" if k == 0 else f"=IF(D{r}>=0.9,\"CUMPLE\",\"caudal parcial: ver tau\")")).font = NEGRO
        ws.cell(row=r, column=10, value=f"=IF(F{r}>=0.15,\"CUMPLE\",\"NO CUMPLE\")").font = NEGRO
        ws.cell(row=r, column=11, value=f"=IF(D{r}<=3.0,\"CUMPLE\",\"NO CUMPLE\")").font = NEGRO
        for c, f in zip(range(3, 9), ["0.000", "0.00", "0.000", "0.000", "0.00", "0.00"]): ws.cell(row=r, column=c).number_format = f
    rf = rv + 3 + len(casos)
    celda(ws, f"A{rf+1}", "Velocidad maxima en el perfil de flujo (m/s)"); celda(ws, f"B{rf+1}", f"=MAX(H{r0}:H{rlast})", NEGRO, "0.00")
    celda(ws, f"A{rf+2}", "Velocidad minima en el perfil de flujo con el caudal de diseno (m/s)"); celda(ws, f"B{rf+2}", f"=MIN(H{r0}:H{rlast})", NEGRO, "0.00")
    celda(ws, f"A{rf+3}", "Verificacion global de velocidades"); celda(ws, f"B{rf+3}", f"=IF(AND(B{rf+1}<=3.0,B{rf+2}>=0.9,COUNTIF(J{rv+3}:J{rf-1},\"NO CUMPLE\")=0),\"CUMPLE\",\"NO CUMPLE\")", NEGRO)
    celda(ws, f"A{rf+4}", "Criterios: velocidad minima 0.90 m/s con el caudal de diseno (autolimpieza, RNE CE.040); velocidad maxima 3.0 m/s para revestimiento de concreto (valor conservador); esfuerzo tractivo minimo 0.15 kg/m2 (1.5 Pa) para caudales parciales (arrastre de arena fina en colectores pluviales).", SUB)
    ws["B200"] = f"=B{rf+3}"; ws["A200"] = "Velocidades:"; ws["A200"].font = SUB
    # G200: referencia fija usada por EMPALME -> copiamos
    ws["G200"] = f"=G{rlast}"; ws["G200"].font = NEGRO; ws["F200"] = "NA en 0+000:"; ws["F200"].font = SUB
    ws.freeze_panes = "B6"
    return ws, rlast


def hoja_estructural(wb, R):
    """Verificacion estructural; las formulas se escriben con claves {clave} que se
    traducen a la fila real al final (evita errores de numeracion)."""
    ws = wb.create_sheet("ESTRUCTURAL"); D = dz.D
    cabecera(ws, "6. VERIFICACION ESTRUCTURAL (RNE E.060; cargas E.020 y rueda AASHTO LRFD HL-93)", [56, 12, 9, 64, 16])
    encabezado_tabla(ws, 4, ["ELEMENTO / PARAMETRO", "VALOR", "UND", "FORMULA / CRITERIO", "RESULTADO"])
    rows = []
    def sec(t): rows.append((None, t))
    def it(k, *a): rows.append((k,) + a)
    sec("PARAMETROS")
    it("fc", "f'c", 210, "kg/cm2", "E.060"); it("fy", "fy", 4200, "kg/cm2", "Grado 60")
    it("gc", "Peso concreto armado", 2.4, "t/m3", "E.020"); it("gs", "Peso suelo / relleno", 1.8, "t/m3", "Verificar con EMS")
    it("Ko", "Ko (reposo)", 0.5, "-", "Relleno compactado contra muro rigido")
    it("sp", "Sobrecarga peatonal", 0.5, "t/m2", "E.020"); it("sl", "Sobrecarga lateral vehiculo liviano", 1.0, "t/m2", "")
    it("sc", "Sobrecarga lateral camion (heq 1.50 m)", 2.7, "t/m2", "AASHTO 3.11.6.4")
    it("Pl", "Rueda vehiculo liviano / moto", 1.0, "t", ""); it("Pc", "Rueda camion HL-93", 7.26, "t", "AASHTO")
    it("IM", "Impacto", 0.33, "-", "AASHTO 3.6.2"); it("c", "Huella de rueda", 0.25, "m", "AASHTO 3.6.1.2.5")
    it("Lc", "Luz de calculo de losas (b + e muro)", "=DATOS!$B$14+DATOS!$B$15", "m", "centro a centro de muros")
    it("H", "Altura interior maxima del muro H", "=MAX(PERFIL_FLUJO!K6:K200)", "m", "hoja PERFIL_FLUJO")
    it("Hmin", "Altura interior minima del muro", "=MIN(PERFIL_FLUJO!K6:K200)", "m", "")
    sec("A. LOSA SUPERIOR MONOLITICA e=0.10 - TRAMO NORMAL (marco 3/8\" @0.20)")
    it("eA", "Espesor", "=DATOS!$B$17", "m", "")
    it("wA", "Carga ultima", "=1.4*{gc}*{eA}+1.7*{sp}", "t/m2", "1.4 D + 1.7 L")
    it("MuA", "Mu (simplemente apoyada, conservador)", "={wA}*{Lc}^2/8", "t.m/m", "wu L2/8")
    it("dA", "Peralte efectivo d", "={eA}*100/2", "cm", "")
    it("AsA", "As colocado", "=0.71/0.20", "cm2/m", "3/8\" @0.20")
    it("aA", "a", "={AsA}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnA", "Momento resistente fMn", "=0.9*{AsA}*{fy}*({dA}-{aA}/2)/100000", "t.m/m", "0.9 As fy (d - a/2)", "=IF({MnA}>={MuA},\"CUMPLE\",\"NO CUMPLE\")")
    it("AmA", "Cuantia minima 0.0018 b e", "=0.0018*100*{eA}*100", "cm2/m", "E.060 9.7.2", "=IF({AsA}>={AmA},\"CUMPLE\",\"NO CUMPLE\")")
    sec("B. LOSA SUPERIOR MONOLITICA e=0.10 - CRUCE DE MOTOS (marco 3/8\" @0.15)")
    it("eB", "Espesor", "=DATOS!$B$17", "m", "")
    it("EB", "Ancho de franja E", "=0.66+0.55*{Lc}", "m", "AASHTO 4.6.2.1.3")
    it("MrB", "Momento por rueda", "={Pl}*(1+{IM})*({Lc}/4-{c}/8)/{EB}", "t.m/m", "P (1+IM) (L/4 - c/8) / E")
    it("MuB", "Mu", "=1.4*{gc}*{eB}*{Lc}^2/8+1.7*{MrB}", "t.m/m", "1.4 MD + 1.7 ML")
    it("dB", "Peralte efectivo d", "={eB}*100/2", "cm", "")
    it("AsB", "As colocado", "=0.71/0.15", "cm2/m", "3/8\" @0.15")
    it("aB", "a", "={AsB}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnB", "Momento resistente fMn", "=0.9*{AsB}*{fy}*({dB}-{aB}/2)/100000", "t.m/m", "", "=IF({MnB}>={MuB},\"CUMPLE\",\"NO CUMPLE\")")
    it("AmB", "Cuantia minima", "=0.0018*100*{eB}*100", "cm2/m", "E.060 9.7.2", "=IF({AsB}>={AmB},\"CUMPLE\",\"NO CUMPLE\")")
    sec("C. LOSA SUPERIOR e=0.25 - CRUCE DE CAMIONES (doble marco 1/2\" @0.15)")
    it("eC", "Espesor", "=DATOS!$B$18", "m", "")
    it("EC", "Ancho de franja E", "=0.66+0.55*{Lc}", "m", "AASHTO 4.6.2.1.3")
    it("MrC", "Momento por rueda", "={Pc}*(1+{IM})*({Lc}/4-{c}/8)/{EC}", "t.m/m", "simplemente apoyada (conservador)")
    it("MuC", "Mu", "=1.4*{gc}*{eC}*{Lc}^2/8+1.7*{MrC}", "t.m/m", "")
    it("dC", "Peralte efectivo d", "=({eC}-0.04)*100-0.635", "cm", "recubrimiento 4 cm")
    it("AsC", "As colocado", "=1.27/0.15", "cm2/m", "1/2\" @0.15 (una capa en traccion)")
    it("aC", "a", "={AsC}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnC", "Momento resistente fMn", "=0.9*{AsC}*{fy}*({dC}-{aC}/2)/100000", "t.m/m", "", "=IF({MnC}>={MuC},\"CUMPLE\",\"NO CUMPLE\")")
    it("AmC", "Cuantia minima", "=0.0018*100*{eC}*100", "cm2/m", "", "=IF({AsC}>={AmC},\"CUMPLE\",\"NO CUMPLE\")")
    it("VuC", "Cortante ultimo Vu", "=1.7*{Pc}*(1+{IM})*({Lc}-{dC}/100-{c}/2)/{Lc}/{EC}+1.4*{gc}*{eC}*{Lc}/2", "t/m", "rueda a d del apoyo")
    it("VcC", "Cortante resistente fVc", "=0.85*0.53*SQRT({fc})*100*{dC}/1000", "t/m", "", "=IF({VcC}>={VuC},\"CUMPLE\",\"NO CUMPLE\")")
    sec("D. MURO e=0.15 - TRAMO NORMAL (marco unico 3/8\" @0.20 en el eje de la seccion) - marco cerrado monolitico")
    it("HD", "Altura del muro H", "={H}", "m", "altura interior maxima")
    it("pD", "Empuje en reposo en la base p", "={Ko}*{gs}*{HD}", "t/m2", "Ko gamma H")
    it("qD", "Sobrecarga lateral q", "={Ko}*{sp}", "t/m2", "Ko x s/c peatonal")
    it("MD", "Momento de servicio", "={pD}*{HD}^2/20+{qD}*{HD}^2/12", "t.m/m", "miembro de marco cerrado con extremos empotrados: p H2/20 (triangular) + q H2/12 (uniforme)")
    it("MuD", "Mu", "=1.7*{MD}", "t.m/m", "1.7 empuje")
    it("dD", "Peralte efectivo d", "=DATOS!$B$15*100/2", "cm", "marco unico en el eje del muro (una capa; E.060 14.3.4 exige dos capas solo en muros de mas de 0.20 m)")
    it("AsD", "As colocado", "=0.71/0.20", "cm2/m", "3/8\" @0.20")
    it("aD", "a", "={AsD}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnD", "Momento resistente fMn", "=0.9*{AsD}*{fy}*({dD}-{aD}/2)/100000", "t.m/m", "", "=IF({MnD}>={MuD},\"CUMPLE\",\"NO CUMPLE\")")
    it("AmD", "Cuantia minima 0.0018 b e", "=0.0018*100*DATOS!$B$15*100", "cm2/m", "", "=IF({AsD}>={AmD},\"CUMPLE\",\"NO CUMPLE\")")
    it("VuD", "Vu", "=1.7*({pD}*{HD}/2*0.6+{qD}*{HD}/2)", "t/m", "reaccion en la base (0.6 de la triangular)")
    it("VcD", "fVc", "=0.85*0.53*SQRT({fc})*100*{dD}/1000", "t/m", "", "=IF({VcD}>={VuD},\"CUMPLE\",\"NO CUMPLE\")")
    sec("E. MURO e=0.15 - CRUCE DE MOTOS (marco unico 3/8\" @0.15)")
    it("HE", "Altura del muro H", "={H}", "m", "")
    it("pE", "Empuje en reposo en la base p", "={Ko}*{gs}*{HE}", "t/m2", "")
    it("qE", "Sobrecarga lateral q", "={Ko}*{sl}", "t/m2", "Ko x s/c vehiculo liviano")
    it("MuE", "Mu", "=1.7*({pE}*{HE}^2/20+{qE}*{HE}^2/12)", "t.m/m", "")
    it("AsE", "As colocado", "=0.71/0.15", "cm2/m", "")
    it("aE", "a", "={AsE}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnE", "Momento resistente fMn", "=0.9*{AsE}*{fy}*({dD}-{aE}/2)/100000", "t.m/m", "", "=IF({MnE}>={MuE},\"CUMPLE\",\"NO CUMPLE\")")
    sec("F. MURO e=0.15 - CRUCE DE CAMIONES (doble marco 1/2\" @0.15: una capa en cada cara)")
    it("HF", "Altura del muro H", "={H}", "m", "")
    it("pF", "Empuje en reposo en la base p", "={Ko}*{gs}*{HF}", "t/m2", "")
    it("qF", "Sobrecarga lateral q", "={Ko}*{sc}", "t/m2", "Ko x s/c camion")
    it("MuF", "Mu", "=1.7*({pF}*{HF}^2/20+{qF}*{HF}^2/12)", "t.m/m", "")
    it("dF", "Peralte efectivo d", "=(DATOS!$B$15-0.04)*100-0.635", "cm", "")
    it("AsF", "As colocado (una capa)", "=1.27/0.15", "cm2/m", "")
    it("aF", "a", "={AsF}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnF", "Momento resistente fMn", "=0.9*{AsF}*{fy}*({dF}-{aF}/2)/100000", "t.m/m", "", "=IF({MnF}>={MuF},\"CUMPLE\",\"NO CUMPLE\")")
    sec("G. LOSA DE FONDO e=0.15 (marco unico 3/8\" @0.20)")
    it("WG", "Peso de la estructura por metro", "={gc}*(DATOS!$B$14+2*DATOS!$B$15)*(DATOS!$B$17+DATOS!$B$16)+2*{gc}*DATOS!$B$15*{H}", "t/m", "losas + muros (altura maxima)")
    it("WwG", "Agua (colector lleno) y sobrecarga", "=1.0*DATOS!$B$14*{H}+{sp}*(DATOS!$B$14+2*DATOS!$B$15)", "t/m", "")
    it("qG", "Reaccion del suelo", "=({WG}+{WwG})/(DATOS!$B$14+2*DATOS!$B$15)", "t/m2", "uniforme")
    it("MuG", "Mu", "=1.5*{qG}*{Lc}^2/8", "t.m/m", "1.5 (D+L) wL2/8, simplemente apoyada (conservador)")
    it("dG", "Peralte efectivo d", "=DATOS!$B$16*100/2", "cm", "marco unico en el eje de la losa de fondo (una capa)")
    it("AsG", "As colocado", "=0.71/0.20", "cm2/m", "")
    it("aG", "a", "={AsG}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnG", "Momento resistente fMn", "=0.9*{AsG}*{fy}*({dG}-{aG}/2)/100000", "t.m/m", "", "=IF({MnG}>={MuG},\"CUMPLE\",\"NO CUMPLE\")")
    it("sG", "Presion sobre el suelo", "={qG}/10", "kg/cm2", "comparar con la capacidad portante del EMS", "=IF({sG}<=1.0,\"CUMPLE (<= 1.0 kg/cm2)\",\"VERIFICAR EMS\")")
    sec("H. TAPA DE REGISTRO 0.68 x 0.68 x 0.08 - 3/8\" @0.10 (rueda liviana; no hay registros en los cruces)")
    it("eH", "Espesor", 0.08, "m", ""); it("LH", "Luz de calculo", 0.64, "m", "luz libre 0.60 + apoyo 0.02 a cada lado")
    it("EH", "Ancho de franja", "=MIN(0.66+0.55*{LH},0.68)", "m", "limitado al ancho de la tapa")
    it("MrH", "Momento por rueda", "={Pl}*(1+{IM})*({LH}/4-{c}/8)/{EH}", "t.m/m", "")
    it("MuH", "Mu", "=1.4*{gc}*{eH}*{LH}^2/8+1.7*{MrH}", "t.m/m", "")
    it("dH", "Peralte efectivo d", "=({eH}-0.025)*100-0.48", "cm", "")
    it("AsH", "As colocado", "=0.71/0.10", "cm2/m", "")
    it("aH", "a", "={AsH}*{fy}/(0.85*{fc}*100)", "cm", "")
    it("MnH", "Momento resistente fMn", "=0.9*{AsH}*{fy}*({dH}-{aH}/2)/100000", "t.m/m", "", "=IF({MnH}>={MuH},\"CUMPLE\",\"NO CUMPLE\")")
    # asignar filas
    KEY = {}; r = 5
    for f in rows:
        if f[0] is not None: KEY[f[0]] = r
        r += 1
    def tr(txt):
        return re.sub(r"\{(\w+)\}", lambda m: "B%d" % KEY[m.group(1)], txt) if isinstance(txt, str) else txt
    r = 5
    for f in rows:
        if f[0] is None:
            celda(ws, f"A{r}", f[1], NEG, fill=GRIS)
        else:
            vals = [f[1], tr(f[2]), f[3], f[4], tr(f[5]) if len(f) > 5 else None]
            fila(ws, r, vals, fonts=[NEGRO, None, NEGRO, NEGRO, NEG])
            if isinstance(vals[1], str) and ("DATOS" in vals[1] or "PERFIL" in vals[1]): ws[f"B{r}"].font = VERDE
            ws[f"B{r}"].number_format = "0.000"
        r += 1
    celda(ws, f"A{r+1}", "Los muros se verifican como miembros del marco cerrado (losa superior y de fondo vaciadas monoliticamente con los muros). En los cruces el acero indicado va en ambas caras. Registros solo fuera de los cruces vehiculares.", SUB)
    hoja_estructural.KEY = KEY
    return ws


def hoja_cumplimiento(wb, rlast):
    """Cuadro de cumplimiento normativo: cada requisito con su norma, criterio, valor del diseno (formula) y resultado."""
    ws = wb.create_sheet("CUMPLIMIENTO"); D = dz.D; K = hoja_estructural.KEY
    cabecera(ws, "7. CUADRO DE CUMPLIMIENTO NORMATIVO Y CONDICIONES DE COMPATIBILIDAD", [58, 34, 40, 14, 18])
    encabezado_tabla(ws, 4, ["REQUISITO", "NORMA / REFERENCIA", "CRITERIO", "VALOR DEL DISENO", "RESULTADO"])
    E = lambda k: "ESTRUCTURAL!B%d" % K[k]
    filas = [
        ("HIDROLOGIA E HIDRAULICA",),
        ("Periodo de retorno del colector", "Estudio hidrologico del proyecto (RNE CE.040)", "TR = 25 anos, igual al estudio hidrologico y al colector receptor", "=DATOS!$B$6", "=IF(DATOS!$B$6=25,\"CUMPLE\",\"VERIFICAR\")"),
        ("Intensidad de diseno del colector", "Estudio hidrologico; MTC Manual de Hidrologia (Dick y Peschke)", "tc = 15 min sobre la curva P-D del analisis estadistico (P10 = 35.16, P20 = 41.82 mm)", "=DATOS!$B$11", "=IF(ABS(DATOS!$B$11-155.66)<0.5,\"CUMPLE\",\"VERIFICAR\")"),
        ("Caudal de diseno del tramo", "Metodo racional, FS = 1.15 (memorias HIDRO-CE040)", "Q = Q CAR Varones (dato) + Q Hogar de Refugio", "=CAUDALES!$B$12", "=IF(CAUDALES!$B$12<=560.7,\"CUMPLE (<= 560.7 L/s del receptor)\",\"NO CUMPLE\")"),
        ("Regimen de flujo en el colector", "RNE CE.040; Chow (flujo gradualmente variado)", "Subcritico en todo el tramo: Froude < 1", f"=PERFIL_FLUJO!B{rlast+5}", f"=IF(PERFIL_FLUJO!B{rlast+5}<1,\"CUMPLE\",\"NO CUMPLE\")"),
        ("Llenado maximo de la seccion", "Criterio de diseno adoptado (y/h <= 85 %)", "Tirante / altura interior", f"=PERFIL_FLUJO!B{rlast+3}", f"=IF(PERFIL_FLUJO!B{rlast+3}<=DATOS!$B$33,\"CUMPLE\",\"NO CUMPLE\")"),
        ("Borde libre bajo la losa", "Criterio de diseno adoptado (>= 0.05 m)", "Cota bajo losa - nivel de agua", f"=PERFIL_FLUJO!B{rlast+4}", f"=IF(PERFIL_FLUJO!B{rlast+4}>=DATOS!$B$34,\"CUMPLE\",\"NO CUMPLE\")"),
        ("Velocidad minima (autolimpieza)", "RNE CE.040 Drenaje Pluvial", "V >= 0.90 m/s con el caudal de diseno", "=PERFIL_FLUJO!B200", "=PERFIL_FLUJO!B200"),
        ("Velocidad maxima (revestimiento de concreto)", "RNE CE.040 Drenaje Pluvial", "V <= 3.0 m/s", "=PERFIL_FLUJO!B200", "=PERFIL_FLUJO!B200"),
        ("Esfuerzo tractivo con caudales parciales", "Criterio de autolimpieza (ASCE / WEF)", "tau = gamma R S >= 0.15 kg/m2 hasta el 5 % del caudal", "=PERFIL_FLUJO!B200", "=PERFIL_FLUJO!B200"),
        ("Caida libre de las cunetas al colector", "Criterio de diseno: NA del colector bajo el fondo de cada cuneta", "Cunetas Ejes 01, 02, 06, 07, 11 y 12 (cotas del plano de arquitectura vigente)", "=COUNTIF(CUNETAS!L5:L10,\"CAIDA LIBRE\")", "=IF(COUNTIF(CUNETAS!L5:L10,\"CAIDA LIBRE\")=6,\"CUMPLE (6 de 6)\",\"VERIFICAR\")"),
        ("Capacidad de la seccion con llenado del 85 % (margen frente al caudal de diseno)", "Manning; criterio de llenado maximo adoptado", "Q85 = (1/n) A R^(2/3) S^(1/2) con y = 0.85 h minima; margen = Q85 / Q diseno",
         "=1000*DATOS!$B$14*0.85*%s*(DATOS!$B$14*0.85*%s/(DATOS!$B$14+2*0.85*%s))^(2/3)*SQRT(DATOS!$B$20)/DATOS!$B$32" % (E("Hmin"), E("Hmin"), E("Hmin")),
         "=\"CUMPLE (margen \"&TEXT(1000*DATOS!$B$14*0.85*%s*(DATOS!$B$14*0.85*%s/(DATOS!$B$14+2*0.85*%s))^(2/3)*SQRT(DATOS!$B$20)/DATOS!$B$32/CAUDALES!$B$12,\"0.0\")&\" veces el caudal de diseno)\"" % (E("Hmin"), E("Hmin"), E("Hmin"))),
        ("Cota de llegada del colector del CAR Varones: caida libre sobre el NA de la CL y bajo el techo de la caja", "Compatibilidad con el expediente del CAR Varones (CUI 2705619)", "Entre el nivel de agua en 0+000 + 0.10 (caida libre) y el techo de la CL (tapa 261.15 - losa 0.10); ventana este 0.60 x 1.17",
         "=DATOS!$B$41", "=IF(AND(DATOS!$B$41>=PERFIL_FLUJO!G200+0.10,DATOS!$B$41<=DATOS!$B$50-DATOS!$B$17),\"CUMPLE (rango \"&TEXT(PERFIL_FLUJO!G200+0.10,\"0.00\")&\" a \"&TEXT(DATOS!$B$50-DATOS!$B$17,\"0.00\")&\" msnm)\",\"VERIFICAR\")"),
        ("Cuneta Eje 01 del CAR Varones: caida libre a la CL", "Compatibilidad con el expediente del CAR Varones", "NCF en el muro norte > NA de la CL", "=DATOS!$B$51", "=IF(DATOS!$B$51>PERFIL_FLUJO!G200,\"CUMPLE (caida libre \"&TEXT(DATOS!$B$51-PERFIL_FLUJO!G200,\"0.00\")&\" m)\",\"VERIFICAR\")"),
        ("ENTREGA AL COLECTOR RECEPTOR (CAR MUJERES, CUI 2717013)",),
        ("Cota de fondo de llegada >= cota de fondo del R-01", "Compatibilidad con el expediente del receptor", "CF llegada 258.89 >= 258.72", "=EMPALME!B8", "=IF(EMPALME!B8>=DATOS!$B$37,\"CUMPLE\",\"NO CUMPLE\")"),
        ("Caudal entregado <= caudal previsto por el receptor", "Compatibilidad con el expediente del receptor", "Q <= 560.7 L/s", "=CAUDALES!$B$12", "=IF(CAUDALES!$B$12<=560.7,\"CUMPLE\",\"NO CUMPLE\")"),
        ("Entrega en regimen subcritico (resalto ahogado en la poza)", "USBR, poza de disipacion; Chow", "y2 conjugado < tirante disponible en la poza", "=EMPALME!B19", "=EMPALME!B19"),
        ("Longitud de la poza de disipacion", "USBR (resalto ahogado: L >= 0.8 x 6 y2)", "L poza = 3.50 m", "=EMPALME!B22", "=EMPALME!B22"),
        ("ESTRUCTURAS (RNE E.060 CONCRETO ARMADO, E.020 CARGAS, AASHTO LRFD)",),
        ("Combinaciones de carga", "RNE E.060 art. 9.2", "U = 1.4 D + 1.7 L; factores phi = 0.90 flexion y 0.85 cortante (art. 9.3)", "aplicado", "CUMPLE"),
        ("Losa superior tramo normal: flexion", "RNE E.060 cap. 10", "phi Mn >= Mu", "=" + E("MnA"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnA"), E("MuA"))),
        ("Losa superior tramo normal: cuantia minima", "RNE E.060 art. 9.7.2 (0.0018 b h, grado 60)", "As >= As min", "=" + E("AsA"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("AsA"), E("AmA"))),
        ("Losa superior cruce de motos: flexion (rueda 1 t)", "RNE E.060; AASHTO LRFD 4.6.2.1.3 (ancho de franja)", "phi Mn >= Mu", "=" + E("MnB"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnB"), E("MuB"))),
        ("Losa superior cruce de camiones: flexion (rueda HL-93, IM 33 %)", "AASHTO LRFD 3.6.1.2 y 3.6.2; Manual de Puentes MTC", "phi Mn >= Mu", "=" + E("MnC"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnC"), E("MuC"))),
        ("Losa superior cruce de camiones: cortante", "RNE E.060 art. 11.3 (Vc = 0.53 raiz f'c b d)", "phi Vc >= Vu", "=" + E("VcC"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("VcC"), E("VuC"))),
        ("Muro tramo normal: flexion por empuje de suelo", "RNE E.060; E.020 (empuje en reposo Ko = 0.5)", "phi Mn >= Mu (marco cerrado)", "=" + E("MnD"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnD"), E("MuD"))),
        ("Muro tramo normal: cortante", "RNE E.060 art. 11.3", "phi Vc >= Vu", "=" + E("VcD"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("VcD"), E("VuD"))),
        ("Muro tramo normal: refuerzo en una capa", "RNE E.060 art. 14.3.4 (dos capas solo si e > 0.20 m)", "e = 0.15 m: una capa en el eje", "=DATOS!$B$15", "=IF(DATOS!$B$15<=0.20,\"CUMPLE\",\"VERIFICAR\")"),
        ("Muro cruce de motos: flexion", "RNE E.060", "phi Mn >= Mu", "=" + E("MnE"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnE"), E("MuE"))),
        ("Muro cruce de camiones: flexion (sobrecarga lateral camion)", "RNE E.060; AASHTO LRFD 3.11.6.4", "phi Mn >= Mu", "=" + E("MnF"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnF"), E("MuF"))),
        ("Losa de fondo: flexion", "RNE E.060", "phi Mn >= Mu", "=" + E("MnG"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnG"), E("MuG"))),
        ("Presion sobre el suelo", "RNE E.050 Suelos y Cimentaciones", "q <= 0.50 kg/cm2, valor minimo usual de capacidad admisible en suelos blandos; el EMS del proyecto lo confirma", "=" + E("sG"), "=IF(%s<=0.5,\"CUMPLE (q = \"&TEXT(%s,\"0.00\")&\" kg/cm2 < 0.50)\",\"VERIFICAR EMS\")" % (E("sG"), E("sG"))),
        ("Tapa de registro 0.68 x 0.68 x 0.08: flexion (rueda liviana)", "RNE E.060", "phi Mn >= Mu", "=" + E("MnH"), "=IF(%s>=%s,\"CUMPLE\",\"NO CUMPLE\")" % (E("MnH"), E("MuH"))),
        ("Recubrimientos", "RNE E.060 art. 7.7 (concreto sobre solado: 4 cm)", "4 cm muros y losa de fondo; 2.5 cm losa superior no expuesta; 4 cm cruce de camiones", "aplicado", "CUMPLE"),
        ("Traslapes y ganchos", "RNE E.060 cap. 12", "Traslape 0.40 m (3/8\"), 0.50 m (1/2\"); ganchos 0.30 m", "aplicado", "CUMPLE"),
        ("CONSTRUCTIVOS",),
        ("Registros de limpieza", "RNE CE.040 (accesibilidad para mantenimiento)", "Separacion <= 12 m y en cada llegada de cuneta; fuera de los cruces vehiculares", "7 + 3 registros", "CUMPLE"),
        ("Juntas de dilatacion", "Practica del colector receptor (CAR Mujeres)", "Cada 4.00 m con tecnopor 1\" y sello", "17 juntas", "CUMPLE"),
        ("Emplazamiento", "Lindero y faja de la carretera Oasis", "Dentro del predio, pegado por fuera del cerco, sin invadir la via", "eje a 0.575 m del cerco", "CUMPLE"),
    ]
    r = 5
    for f in filas:
        if len(f) == 1:
            celda(ws, f"A{r}", f[0], NEG, fill=GRIS)
        else:
            fila(ws, r, list(f), fonts=[NEGRO, NEGRO, NEGRO, VERDE, NEG])
            ws[f"D{r}"].number_format = "0.000"
        r += 1
    r += 1
    celda(ws, f"A{r}", "Requisitos que no cumplen o por verificar", NEG); celda(ws, f"B{r}", f'=COUNTIF(E5:E{r-2},"NO CUMPLE*")+COUNTIF(E5:E{r-2},"VERIFICAR*")', NEG)
    r += 2
    celda(ws, f"A{r}", "DATOS DE OTROS EXPEDIENTES Y COMO QUEDAN CUBIERTOS EN ESTE DISENO", NEG, fill=GRIS); r += 1
    for t in ["1. Llegada del colector del CAR Varones (CUI 2705619): cota de fondo 259.884 msnm en la cara este de la CL (ventana 0.60 x 1.17 hasta el techo de la CL) y cuneta Eje 01 de Varones por el muro norte (ventana 0.40 x 0.66, NCF 260.394). Como la CL queda del lado de Varones, su tapa se fija en +261.15 (piso de ese proyecto); el colector de Varones entrega 221.7 + 37.0 = 258.7 L/s.",
              "2. Caudal del CAR Varones (258.7 L/s): dato de su memoria y base del colector receptor (560.7 L/s). La seccion de este tramo tiene capacidad para mas del doble del caudal de diseno con llenado del 85 % (ver fila del cuadro), de modo que una variacion de ese dato no compromete el tramo.",
              "3. Capacidad portante: la presion transmitida (0.33 kg/cm2) es menor que 0.50 kg/cm2, valor minimo usual de suelos blandos; el EMS del proyecto, exigido por la norma E.050, la confirma.",
              "4. Coordenadas UTM: obtenidas del registro R-01 del CAR Mujeres (CUI 2717013) y del rumbo del lindero (azimut 49.54); el trazo queda definido por su posicion fisica (eje a 0.575 m del cerco) y se replantea en obra desde ese cerco.",
              "5. Cotas de fondo de las cunetas: tomadas del plano de arquitectura vigente (PLANTA GENERAL REFUGIO, 03-10-2026), perfiles 01 a 12; todas caen libremente al colector."]:
        celda(ws, f"A{r}", t, NEGRO); r += 1
    ws.freeze_panes = "A5"
    return ws


def hoja_memoria(wb, R, rlast):
    ws = wb.create_sheet("MEMORIA", 0); D = dz.D
    cabecera(ws, "MEMORIA DE CALCULO - COLECTOR PLUVIAL TRAMO HOGAR DE REFUGIO TEMPORAL MUJERES VIOLENTADAS", [38, 120, 10])
    txt = [
        ("1. OBJETO", "Dimensionar el tramo del colector pluvial frontal que recibe el aporte del CAR Varones (CUI 2705619, dato de su memoria) y las seis cunetas del Hogar de Refugio Temporal Mujeres Violentadas (CUI 2675514), y lo entrega al registro R-01 del colector del CAR Mujeres (CUI 2717013). Longitud %.2f m, de la caja de llegada (0+000, limite con CAR Varones) a la caja de caida en el limite con CAR Mujeres." % dz.P_FIN),
        ("2. NORMATIVA", "RNE CE.040 Drenaje Pluvial (RM 126-2021-VIVIENDA); RNE E.020 Cargas; RNE E.060 Concreto Armado; AASHTO LRFD (rueda HL-93 en el cruce de camiones); Chow, Hidraulica de canales abiertos (paso estandar)."),
        ("3. HIDROLOGIA", "Metodo Racional con los coeficientes y areas de las memorias HIDRO-CE040 de cada proyecto, TR = 25 anos, tc = 15 min, I = 155.66 mm/h, FS = 1.15. Q = 258.7 + 301.9 = 560.6 L/s."),
        ("4. HIDRAULICA", "Colector cubierto de concreto armado b = 0.80 m, S = 0.30 %, n = 0.015. El fondo inicial (259.10) queda bajo las cunetas de arquitectura (NCF 259.70 a 260.12) y el fondo final (258.89) sobre el R-01 del receptor (258.72). La entrega es por caja de caida con poza de disipacion deprimida 0.40 m, ahogada por el tirante del receptor; el control del perfil es el tirante critico en el brink y el flujo en el colector es subcritico (F <= 0.90). El aporte externo cae en una caja de llegada con colchon de agua."),
        ("5. ESTRUCTURAS", "Marco cerrado monolitico: losa superior e = 0.10, muros e = 0.15 (altura interior 1.40 a 1.65 m) y losa de fondo e = 0.15 con un solo marco 3/8\" @0.20 en el eje de la seccion y longitudinales 3/8\" @0.25 (una capa; E.060 14.3.4); cruce de motos marco 3/8\" @0.15; cruce de camiones losas e = 0.25 y doble marco 1/2\" @0.15 (una capa en cada cara). Tapas de registro 0.68 x 0.68 x 0.08."),
        ("6. RESULTADOS", "Ver cuadro resumen y la hoja CUMPLIMIENTO (requisito, norma, criterio, valor y resultado). Secciones, acero y detalles en las laminas DP-01 a DP-10 y DA-01 a DA-03."),
    ]
    r = 4
    for a, b in txt:
        celda(ws, f"A{r}", a, NEG); celda(ws, f"B{r}", b, NEGRO, al="left"); ws.row_dimensions[r].height = 42; r += 1
    r += 1
    res = [
        ("Caudal de diseno del tramo (L/s)", "=CAUDALES!B12", "0.0"),
        ("Intensidad de diseno (mm/h)", "=DATOS!B11", "0.00"),
        ("Ancho interior adoptado (m)", "=DATOS!B14", "0.00"),
        ("Pendiente del fondo (%)", "=DATOS!B20*100", "0.00"),
        ("Cota de fondo inicial / final del colector (msnm)", f"=TEXT(DATOS!B19,\"0.00\")&\" / \"&TEXT(EMPALME!B8,\"0.00\")", None),
        ("Llenado maximo y/h", f"=PERFIL_FLUJO!B{rlast+3}", "0.000"),
        ("Borde libre minimo bajo losa (m)", f"=PERFIL_FLUJO!B{rlast+4}", "0.000"),
        ("Froude maximo en el colector", f"=PERFIL_FLUJO!B{rlast+5}", "0.00"),
        ("Estaciones hidraulicas que no cumplen", f"=PERFIL_FLUJO!B{rlast+6}", "0"),
        ("Velocidades y autolimpieza (0.90 a 3.0 m/s; tau >= 0.15 kg/m2)", "=PERFIL_FLUJO!B200", None),
        ("Cunetas con caida libre al colector (de 6)", "=COUNTIF(CUNETAS!L5:L10,\"CAIDA LIBRE\")", "0"),
        ("Resalto en la poza de la caja de caida", "=EMPALME!B19", None),
        ("Elementos estructurales que no cumplen", "=COUNTIF(ESTRUCTURAL!E:E,\"NO CUMPLE\")", "0"),
        ("Requisitos normativos que no cumplen o por verificar (hoja CUMPLIMIENTO)", "=COUNTIF(CUMPLIMIENTO!E:E,\"NO CUMPLE*\")+COUNTIF(CUMPLIMIENTO!E:E,\"VERIFICAR*\")", "0"),
    ]
    for a, b, f in res:
        celda(ws, f"A{r}", a, NEGRO); c = celda(ws, f"B{r}", b, VERDE); c.alignment = Alignment(horizontal="left")
        if f: c.number_format = f
        r += 1
    r += 1
    celda(ws, f"A{r}", "NOTAS", NEG)
    notas = ["La llegada del CAR Varones (CUI 2705619) queda definida por su expediente: colector b = 0.60 con fondo 259.884 en la cara este de la CL y cuneta Eje 01 directa a la CL; la CL lleva su tapa en +261.15 (piso de Varones) y ventanas este 0.60 x 1.17 y norte 0.40 x 0.66.",
             "La ubicacion UTM es referencial: el plano de arquitectura se georreferencio haciendo coincidir el R-01 del CAR Mujeres con la esquina sur-oeste del frente y el rumbo del lindero (azimut 49.54). Verificar en campo.",
             "Las cunetas de arquitectura se toman como dato (perfiles 01 a 12 del plano PLANTA GENERAL REFUGIO, version 03-10-2026)."]
    for n in notas:
        r += 1; celda(ws, f"B{r}", n, NEGRO, al="left"); ws.row_dimensions[r].height = 30
    return ws


def construir(fn=os.path.join(RAIZ, "entregables", "MEMORIA_CALCULO_COLECTOR_HOGAR_REFUGIO.xlsx")):
    R = dz.disenar()
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    hoja_datos(wb, R); hoja_caudales(wb, R); hoja_cunetas(wb, R); hoja_empalme(wb, R)
    ws, rlast = hoja_perfil(wb, R); hoja_estructural(wb, R); hoja_cumplimiento(wb, rlast); hoja_memoria(wb, R, rlast)
    OLD = {6: "TR", 11: "I", 12: "FS", 14: "b", 15: "em", 16: "ef", 17: "et", 18: "ec", 19: "CF0", 20: "S", 23: "pb", 25: "NPT",
           26: "pc1", 27: "pc2", 28: "pm1", 29: "pm2", 32: "n", 33: "llen", 34: "BLmin", 37: "CFr", 38: "NAr", 41: "CFv",
           44: "pcl", 45: "pcc", 46: "Lcc", 50: "CLt", 51: "E01"}
    import re
    def tr(m):
        return "DATOS!" + m.group(1) + "B" + m.group(2) + str(ROW[OLD[int(m.group(3))]])
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        if ws.title == "DATOS": continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "DATOS!" in c.value:
                    c.value = re.sub(r"DATOS!(\$?)B(\$?)(\d+)", tr, c.value)
    os.makedirs(os.path.dirname(fn), exist_ok=True); wb.save(fn)
    return fn, R


if __name__ == "__main__":
    print(construir()[0])
