"""MetraP - aplicacion de escritorio (PyQt6) para el metrado de cunetas desde un plano DXF.

Paginas: Proyecto (archivos y reglas) - Perfiles (dibujo de cada eje) - Metrado (tramos, acero y totales) -
Planilla (todas las hojas del Excel generado, con sus colores y formulas evaluadas) - Exportar (guardar, abrir, imprimir).
"""
import os, sys, json, math, tempfile, shutil, subprocess, traceback

AQUI = os.path.dirname(os.path.abspath(sys.argv[0] if getattr(sys, "frozen", False) else __file__))
BASE = getattr(sys, "_MEIPASS", AQUI)
for p in (BASE, AQUI):
    if p not in sys.path: sys.path.insert(0, p)
import metrap_engine as ME

from PyQt6.QtCore import Qt, QThread, pyqtSignal, QSize, QMarginsF
from PyQt6.QtGui import QColor, QFont, QIcon, QPainter, QPageLayout, QPageSize, QTextDocument, QBrush
from PyQt6.QtWidgets import (QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QGridLayout, QLabel, QPushButton,
                             QLineEdit, QFileDialog, QStackedWidget, QFrame, QListWidget, QListWidgetItem, QTableWidget, QTableWidgetItem,
                             QTabWidget, QTextEdit, QMessageBox, QDoubleSpinBox, QSpinBox, QSizePolicy, QHeaderView, QScrollArea, QSplitter)
from PyQt6.QtPrintSupport import QPrinter, QPrintPreviewDialog
import matplotlib
matplotlib.use("QtAgg")
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure
import openpyxl

VERSION = "2.0"
PLANTILLA_DEF = os.path.join(BASE, "plantilla", "PLANTILLA_METRADO_CUNETAS.xlsx")
CONFIG_DEF = os.path.join(AQUI, "metrap_config.json")
if not os.path.exists(CONFIG_DEF): CONFIG_DEF = os.path.join(BASE, "metrap_config.json")

# ----------------------------------------------------------------------------- tema
C_BG, C_PANEL, C_CARD, C_LINE = "#0f172a", "#111c33", "#172342", "#24324f"
C_TXT, C_MUTED, C_ACC, C_ACC2, C_OK, C_WARN = "#e5e9f2", "#8b97b1", "#38bdf8", "#22c55e", "#4ade80", "#fbbf24"
QSS = f"""
QMainWindow, QWidget {{ background: {C_BG}; color: {C_TXT}; font-family: 'Segoe UI', 'Noto Sans', Arial; font-size: 13px; }}
#sidebar {{ background: {C_PANEL}; border-right: 1px solid {C_LINE}; }}
#logo {{ font-size: 24px; font-weight: 700; color: {C_ACC}; padding: 18px 16px 2px 16px; }}
#logosub {{ color: {C_MUTED}; padding: 0 16px 16px 16px; font-size: 11px; }}
QPushButton#nav {{ text-align: left; padding: 11px 16px; border: none; border-radius: 8px; margin: 2px 8px; color: {C_MUTED}; font-size: 14px; }}
QPushButton#nav:hover {{ background: {C_CARD}; color: {C_TXT}; }}
QPushButton#nav:checked {{ background: {C_CARD}; color: {C_ACC}; font-weight: 600; }}
QLabel#h1 {{ font-size: 22px; font-weight: 700; }}
QLabel#h2 {{ font-size: 15px; font-weight: 600; color: {C_ACC}; }}
QLabel#muted {{ color: {C_MUTED}; }}
QFrame#card {{ background: {C_CARD}; border: 1px solid {C_LINE}; border-radius: 12px; }}
QLabel#kpi {{ font-size: 24px; font-weight: 700; color: {C_TXT}; }}
QLabel#kpit {{ color: {C_MUTED}; font-size: 11px; letter-spacing: 1px; }}
QLineEdit, QDoubleSpinBox, QSpinBox {{ background: {C_PANEL}; border: 1px solid {C_LINE}; border-radius: 8px; padding: 7px 10px; color: {C_TXT}; selection-background-color: {C_ACC}; }}
QLineEdit:focus, QDoubleSpinBox:focus {{ border: 1px solid {C_ACC}; }}
QPushButton {{ background: {C_PANEL}; border: 1px solid {C_LINE}; border-radius: 8px; padding: 8px 14px; color: {C_TXT}; }}
QPushButton:hover {{ border-color: {C_ACC}; }}
QPushButton#primary {{ background: {C_ACC}; color: #06202c; font-weight: 700; border: none; padding: 10px 22px; font-size: 14px; }}
QPushButton#primary:hover {{ background: #5fd0ff; }}
QPushButton#primary:disabled {{ background: {C_LINE}; color: {C_MUTED}; }}
QPushButton#ok {{ background: {C_ACC2}; color: #052e16; font-weight: 700; border: none; padding: 10px 22px; }}
QTableWidget {{ background: {C_PANEL}; gridline-color: {C_LINE}; border: 1px solid {C_LINE}; border-radius: 8px; selection-background-color: #1e3a5f; }}
QHeaderView::section {{ background: {C_CARD}; color: {C_MUTED}; padding: 6px; border: none; border-bottom: 1px solid {C_LINE}; font-weight: 600; }}
QTableCornerButton::section {{ background: {C_CARD}; border: none; }}
QTabWidget::pane {{ border: 1px solid {C_LINE}; border-radius: 8px; top: -1px; }}
QTabBar::tab {{ background: {C_PANEL}; color: {C_MUTED}; padding: 8px 14px; border: 1px solid {C_LINE}; border-bottom: none; border-top-left-radius: 8px; border-top-right-radius: 8px; margin-right: 2px; }}
QTabBar::tab:selected {{ background: {C_CARD}; color: {C_ACC}; }}
QListWidget {{ background: {C_PANEL}; border: 1px solid {C_LINE}; border-radius: 8px; padding: 4px; }}
QListWidget::item {{ padding: 9px 10px; border-radius: 6px; }}
QListWidget::item:selected {{ background: {C_CARD}; color: {C_ACC}; }}
QTextEdit {{ background: {C_PANEL}; border: 1px solid {C_LINE}; border-radius: 8px; font-family: Consolas, 'Courier New', monospace; font-size: 12px; }}
QScrollBar:vertical {{ background: {C_PANEL}; width: 10px; }} QScrollBar::handle:vertical {{ background: {C_LINE}; border-radius: 5px; min-height: 30px; }}
QScrollBar:horizontal {{ background: {C_PANEL}; height: 10px; }} QScrollBar::handle:horizontal {{ background: {C_LINE}; border-radius: 5px; }}
QSplitter::handle {{ background: {C_LINE}; }}
"""


def card(titulo=None):
    f = QFrame(); f.setObjectName("card"); lay = QVBoxLayout(f); lay.setContentsMargins(16, 14, 16, 14); lay.setSpacing(8)
    if titulo:
        l = QLabel(titulo); l.setObjectName("h2"); lay.addWidget(l)
    return f, lay


def kpi(titulo, valor="-"):
    f, lay = card(); t = QLabel(titulo.upper()); t.setObjectName("kpit"); v = QLabel(valor); v.setObjectName("kpi")
    lay.addWidget(t); lay.addWidget(v); f.valor = v
    return f


def fmt(v, nd=2):
    if isinstance(v, (int, float)) and not isinstance(v, bool): return ("{:,.%df}" % nd).format(v)
    return "" if v is None else str(v)


# ----------------------------------------------------------------------------- hilo de calculo
class Trabajo(QThread):
    listo = pyqtSignal(object); error = pyqtSignal(str); aviso = pyqtSignal(str)

    def __init__(self, dxf, plantilla, cfg):
        super().__init__(); self.dxf, self.plantilla, self.cfg = dxf, plantilla, cfg

    def run(self):
        try:
            ME.CFG = dict(ME.CONFIG_DEFAULT); ME.CFG.update(self.cfg)
            D, E = ME.analizar(self.dxf)
            tmp = os.path.join(tempfile.gettempdir(), "MetraP_%d.xlsx" % os.getpid())
            ME.construir(self.dxf, self.plantilla, tmp, (D, E))
            filas = ME.filas_acero(D, E)
            self.aviso.emit("Evaluando formulas de la planilla...")
            valores = evaluar_libro(tmp)
            self.listo.emit(dict(D=D, E=E, filas=filas, tmp=tmp, valores=valores, pend=list(ME.PENDIENTES), log=list(ME.LOG)))
        except Exception as e:
            self.error.emit(str(e) + "\n" + traceback.format_exc())


def evaluar_libro(fn):
    """Valores de todas las celdas (formulas evaluadas con pycel; si falla, se muestra la formula)."""
    try:
        from pycel import ExcelCompiler
        xl = ExcelCompiler(filename=fn)
    except Exception:
        xl = None
    wb = openpyxl.load_workbook(fn); out = {}
    for ws in wb.worksheets:
        celdas = {}
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None: continue
                if isinstance(v, str) and v.startswith("=") and xl is not None:
                    try: v = xl.evaluate("'%s'!%s" % (ws.title, c.coordinate))
                    except Exception: pass
                celdas[c.coordinate] = v
        out[ws.title] = celdas
    return out


# ----------------------------------------------------------------------------- ventana principal
class Ventana(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("MetraP %s - Metrado de cunetas desde DXF" % VERSION); self.resize(1360, 840)
        self.res = None; self.cfg = self._cargar_cfg(CONFIG_DEF)
        raiz = QWidget(); self.setCentralWidget(raiz); H = QHBoxLayout(raiz); H.setContentsMargins(0, 0, 0, 0); H.setSpacing(0)
        # sidebar
        sb = QFrame(); sb.setObjectName("sidebar"); sb.setFixedWidth(210); S = QVBoxLayout(sb); S.setContentsMargins(0, 0, 0, 0); S.setSpacing(0)
        logo = QLabel("MetraP"); logo.setObjectName("logo"); sub = QLabel("Metrado de cunetas desde DXF"); sub.setObjectName("logosub")
        S.addWidget(logo); S.addWidget(sub)
        self.nav = []
        for i, (ic, t) in enumerate((("☰", "Proyecto"), ("∿", "Perfiles"), ("∑", "Metrado"), ("▦", "Planilla"), ("⤓", "Exportar"))):
            b = QPushButton("  %s   %s" % (ic, t)); b.setObjectName("nav"); b.setCheckable(True); b.clicked.connect(lambda _, k=i: self.ir(k)); S.addWidget(b); self.nav.append(b)
        S.addStretch()
        self.estado = QLabel("Listo"); self.estado.setObjectName("muted"); self.estado.setWordWrap(True); self.estado.setContentsMargins(16, 8, 16, 16); S.addWidget(self.estado)
        H.addWidget(sb)
        self.pila = QStackedWidget(); H.addWidget(self.pila, 1)
        self.pila.addWidget(self.pag_proyecto()); self.pila.addWidget(self.pag_perfiles()); self.pila.addWidget(self.pag_metrado())
        self.pila.addWidget(self.pag_planilla()); self.pila.addWidget(self.pag_exportar())
        self.ir(0)

    # ------------------------------------------------------------------ utilidades
    def _cargar_cfg(self, fn):
        cfg = dict(ME.CONFIG_DEFAULT)
        try:
            if fn and os.path.exists(fn): cfg.update(json.load(open(fn, encoding="utf8")))
        except Exception: pass
        return cfg

    def ir(self, k):
        for i, b in enumerate(self.nav): b.setChecked(i == k)
        self.pila.setCurrentIndex(k)
        if k == 1 and self.fig.axes: self.fig.tight_layout(); self.canvas.draw()

    def _fila_archivo(self, grid, r, etiqueta, valor, tipo, filtro, attr):
        grid.addWidget(QLabel(etiqueta), r, 0); e = QLineEdit(valor); grid.addWidget(e, r, 1); setattr(self, attr, e)
        b = QPushButton("Examinar"); grid.addWidget(b, r, 2)
        def elegir():
            if tipo == "abrir": f, _ = QFileDialog.getOpenFileName(self, etiqueta, "", filtro)
            else: f, _ = QFileDialog.getSaveFileName(self, etiqueta, e.text(), filtro)
            if f:
                e.setText(f)
                if attr == "e_dxf" and not self.e_sal.text(): self.e_sal.setText(os.path.splitext(f)[0] + "_METRADO.xlsx")
        b.clicked.connect(elegir)

    # ------------------------------------------------------------------ pagina 1: proyecto
    def pag_proyecto(self):
        w = QWidget(); V = QVBoxLayout(w); V.setContentsMargins(28, 24, 28, 24); V.setSpacing(16)
        t = QLabel("Proyecto"); t.setObjectName("h1"); V.addWidget(t)
        s = QLabel("Elija el plano DXF con los perfiles de las cunetas y la plantilla de metrados del proyecto. Las reglas de metrado se pueden ajustar aqui sin editar archivos."); s.setObjectName("muted"); s.setWordWrap(True); V.addWidget(s)
        f, lay = card("Archivos"); g = QGridLayout(); g.setHorizontalSpacing(10); g.setVerticalSpacing(8); lay.addLayout(g)
        self._fila_archivo(g, 0, "Plano DXF (perfiles de cunetas)", "", "abrir", "Planos DXF (*.dxf)", "e_dxf")
        self._fila_archivo(g, 1, "Plantilla Excel del proyecto", PLANTILLA_DEF, "abrir", "Excel (*.xlsx)", "e_pla")
        self._fila_archivo(g, 2, "Archivo Excel de salida", "", "guardar", "Excel (*.xlsx)", "e_sal")
        g.setColumnStretch(1, 1); V.addWidget(f)
        f, lay = card("Datos del proyecto y reglas de metrado"); g = QGridLayout(); g.setHorizontalSpacing(10); g.setVerticalSpacing(8); lay.addLayout(g)
        g.addWidget(QLabel("Nombre del proyecto (RESUMEN!B5)"), 0, 0); self.e_proy = QLineEdit(self.cfg.get("proyecto", "")); g.addWidget(self.e_proy, 0, 1, 1, 3)
        g.addWidget(QLabel("Altura minima con acero H > (m, desde el NCT)"), 1, 0); self.s_hmin = QDoubleSpinBox(); self.s_hmin.setDecimals(2); self.s_hmin.setSingleStep(0.05); self.s_hmin.setValue(float(self.cfg.get("h_minima_acero", 0.40))); g.addWidget(self.s_hmin, 1, 1)
        g.addWidget(QLabel("Tapas de registro por tramo tapado"), 1, 2); self.s_tapas = QSpinBox(); self.s_tapas.setValue(int(self.cfg.get("tapas_por_tramo_tapado", 1))); g.addWidget(self.s_tapas, 1, 3)
        g.addWidget(QLabel("Acero transversal de tapa: largo (m) / separacion (m)"), 2, 0); self.s_tl = QDoubleSpinBox(); self.s_tl.setDecimals(2); self.s_tl.setValue(float(self.cfg.get("ancho_transversal_tapa", 0.56))); g.addWidget(self.s_tl, 2, 1)
        self.s_ts = QDoubleSpinBox(); self.s_ts.setDecimals(2); self.s_ts.setValue(float(self.cfg.get("sep_transversal_tapa", 0.25))); g.addWidget(self.s_ts, 2, 3)
        g.addWidget(QLabel("Acortamientos (eje: metros que se descuentan del ultimo tramo; termina en colector / caja)"), 3, 0, 1, 4)
        self.t_acort = QTableWidget(0, 3); self.t_acort.setHorizontalHeaderLabels(["Eje (numero)", "Acortamiento (m)", "Termina en"]); self.t_acort.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch); self.t_acort.setMaximumHeight(170)
        for k, v in sorted(self.cfg.get("acortamientos", {}).items(), key=lambda kv: int(kv[0])):
            r = self.t_acort.rowCount(); self.t_acort.insertRow(r)
            for j, val in enumerate((k, "%.2f" % float(v[0]), v[1])): self.t_acort.setItem(r, j, QTableWidgetItem(str(val)))
        g.addWidget(self.t_acort, 4, 0, 1, 4)
        hb = QHBoxLayout(); b1 = QPushButton("Agregar fila"); b2 = QPushButton("Quitar fila"); hb.addWidget(b1); hb.addWidget(b2); hb.addStretch(); g.addLayout(hb, 5, 0, 1, 4)
        b1.clicked.connect(lambda: self.t_acort.insertRow(self.t_acort.rowCount())); b2.clicked.connect(lambda: self.t_acort.removeRow(self.t_acort.currentRow()))
        g.setColumnStretch(1, 1); g.setColumnStretch(3, 1); V.addWidget(f)
        hb = QHBoxLayout(); self.b_calc = QPushButton("LEER PLANO Y CALCULAR"); self.b_calc.setObjectName("primary"); self.b_calc.clicked.connect(self.calcular)
        bc = QPushButton("Cargar configuracion..."); bc.clicked.connect(self.cargar_config); bg = QPushButton("Guardar configuracion..."); bg.clicked.connect(self.guardar_config)
        hb.addWidget(self.b_calc); hb.addSpacing(12); hb.addWidget(bc); hb.addWidget(bg); hb.addStretch(); V.addLayout(hb)
        self.log = QTextEdit(); self.log.setReadOnly(True); self.log.setPlaceholderText("Aqui aparece el detalle de la lectura del plano: ejes, tramos, alturas y datos pendientes."); V.addWidget(self.log, 1)
        return w

    def cfg_actual(self):
        cfg = dict(self.cfg); cfg["proyecto"] = self.e_proy.text(); cfg["h_minima_acero"] = self.s_hmin.value(); cfg["tapas_por_tramo_tapado"] = self.s_tapas.value()
        cfg["ancho_transversal_tapa"] = self.s_tl.value(); cfg["sep_transversal_tapa"] = self.s_ts.value()
        ac = {}
        for r in range(self.t_acort.rowCount()):
            try:
                k = self.t_acort.item(r, 0).text().strip(); m = float(self.t_acort.item(r, 1).text().replace(",", ".")); d = (self.t_acort.item(r, 2).text() if self.t_acort.item(r, 2) else "colector") or "colector"
                if k: ac[str(int(k))] = [m, d]
            except Exception: pass
        cfg["acortamientos"] = ac
        return cfg

    def cargar_config(self):
        f, _ = QFileDialog.getOpenFileName(self, "Configuracion", "", "JSON (*.json)")
        if not f: return
        self.cfg = self._cargar_cfg(f); self.e_proy.setText(self.cfg.get("proyecto", "")); self.s_hmin.setValue(float(self.cfg.get("h_minima_acero", 0.4)))
        self.t_acort.setRowCount(0)
        for k, v in sorted(self.cfg.get("acortamientos", {}).items(), key=lambda kv: int(kv[0])):
            r = self.t_acort.rowCount(); self.t_acort.insertRow(r)
            for j, val in enumerate((k, "%.2f" % float(v[0]), v[1])): self.t_acort.setItem(r, j, QTableWidgetItem(str(val)))

    def guardar_config(self):
        f, _ = QFileDialog.getSaveFileName(self, "Guardar configuracion", "metrap_config.json", "JSON (*.json)")
        if f: json.dump(self.cfg_actual(), open(f, "w", encoding="utf8"), indent=2, ensure_ascii=False)

    def calcular(self):
        dxf, pla = self.e_dxf.text().strip(), self.e_pla.text().strip()
        if not os.path.exists(dxf): QMessageBox.warning(self, "MetraP", "Elija el plano DXF."); return
        if not os.path.exists(pla): QMessageBox.warning(self, "MetraP", "Elija la plantilla Excel."); return
        self.b_calc.setEnabled(False); self.estado.setText("Leyendo el plano..."); self.log.clear()
        self.hilo = Trabajo(dxf, pla, self.cfg_actual()); self.hilo.listo.connect(self.resultado); self.hilo.error.connect(self.fallo); self.hilo.aviso.connect(self.estado.setText); self.hilo.start()

    def fallo(self, msg):
        self.b_calc.setEnabled(True); self.estado.setText("Error"); self.log.setPlainText(msg); QMessageBox.critical(self, "MetraP", msg.split("\n")[0])

    def resultado(self, res):
        self.res = res; self.b_calc.setEnabled(True)
        self.log.setPlainText("\n".join(res["log"]) + ("\n\nDATOS QUE FALTAN (no estan en el DXF):\n - " + "\n - ".join(res["pend"]) if res["pend"] else ""))
        self.estado.setText("Plano leido: %d ejes. Revise Perfiles, Metrado y Planilla; luego Exportar." % len(res["E"]))
        self.llenar_perfiles(); self.llenar_metrado(); self.llenar_planilla(); self.ir(1)

    # ------------------------------------------------------------------ pagina 2: perfiles
    def pag_perfiles(self):
        w = QWidget(); V = QVBoxLayout(w); V.setContentsMargins(28, 24, 28, 24); V.setSpacing(12)
        t = QLabel("Perfiles leidos del plano"); t.setObjectName("h1"); V.addWidget(t)
        sp = QSplitter(); V.addWidget(sp, 1)
        self.l_ejes = QListWidget(); self.l_ejes.setMinimumWidth(170); self.l_ejes.setMaximumWidth(240); self.l_ejes.currentRowChanged.connect(self.dibujar_eje); sp.addWidget(self.l_ejes)
        der = QWidget(); D = QVBoxLayout(der); D.setContentsMargins(0, 0, 0, 0); D.setSpacing(10)
        kp = QHBoxLayout(); self.k_perf = [kpi("Longitud metrada", "-"), kpi("Tramo abierto", "-"), kpi("Tramo tapado", "-"), kpi("H promedio", "-"), kpi("Area de muro", "-"), kpi("NCF inicio / fin", "-")]
        for k in self.k_perf: kp.addWidget(k)
        D.addLayout(kp)
        self.fig = Figure(figsize=(9, 4.2), facecolor=C_CARD); self.canvas = FigureCanvas(self.fig); D.addWidget(self.canvas, 1)
        self.t_tramos = QTableWidget(0, 7); self.t_tramos.setHorizontalHeaderLabels(["Tramo", "Tipo", "Inicio", "Fin", "L (m)", "H inicio", "H fin"]); self.t_tramos.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch); self.t_tramos.setMaximumHeight(160)
        D.addWidget(self.t_tramos); sp.addWidget(der); sp.setStretchFactor(0, 0); sp.setStretchFactor(1, 1); sp.setSizes([190, 1100])
        self.canvas.mpl_connect("resize_event", lambda ev: (self.fig.tight_layout(), self.canvas.draw_idle()) if self.fig.axes else None)
        return w

    def llenar_perfiles(self):
        self.l_ejes.clear()
        for e in self.res["E"]:
            it = QListWidgetItem("%s   %.2f m" % (e["nombre"], e["L"])); self.l_ejes.addItem(it)
        self.l_ejes.setCurrentRow(0)

    def dibujar_eje(self, i):
        if not self.res or i < 0: return
        e = self.res["E"][i]; d = self.res["D"][e["nombre"]]; cfg = self.cfg_actual()
        vals = ["%.2f m" % e["L"], "%.2f m" % e["L_ab"], "%.2f m" % e["L_tap"], "%.3f m" % e["Hprom"], "%.2f m2" % e["area"], "%.2f / %.2f" % (e["NCF_ini"], e["NCF_fin"])]
        for k, v in zip(self.k_perf, vals): k.valor.setText(v)
        self.fig.clear(); ax = self.fig.add_subplot(111); ax.set_facecolor(C_PANEL)
        xs = [r["x"] for r in d["sec"]]; nct = [r["NCT"] for r in d["sec"]]; ncf = [r["NCF"] for r in d["sec"]]; hs = [r["H"] for r in d["sec"]]
        L = e["L"]
        import numpy as np
        xx = np.linspace(0, L, 300); hh = [ME.H_en(d, x) for x in xx]
        base = [np.interp(x, xs, nct) if all(v is not None for v in nct) else 0 for x in xx]
        ax.plot(xx, base, color=C_TXT, lw=1.6, label="NCT (piso terminado)")
        ax.plot(xx, [b - h for b, h in zip(base, hh)], color=C_WARN, lw=1.6, label="NCF (fondo de la cuneta)")
        for j, t in enumerate(e["tramos"]):
            col = "#38bdf8" if t["tipo"] == "ABIERTA" else "#22c55e"
            ax.axvspan(t["a"], t["b"], color=col, alpha=0.18, lw=0)
            ax.text((t["a"] + t["b"]) / 2, max(base) + 0.05 + (0.13 if j % 2 else 0), "%s  %.2f m" % (t["tipo"].lower(), t["L"]), ha="center", va="bottom", color=col, fontsize=8)
        hmin = cfg["h_minima_acero"]
        for (ee, k, t) in self.res["filas"]:
            if ee is e: ax.axvspan(t["a"], t["b"], ymin=0.0, ymax=0.04, color="#f87171", alpha=0.9, lw=0)
        x_ant = -1e9; nivel = 0
        for r in d["sec"]:
            if r["NCT"] is None or r["NCF"] is None: continue
            ax.plot([r["x"], r["x"]], [r["NCF"], r["NCT"]], color=C_MUTED, lw=0.8, ls=":")
            nivel = (nivel + 1) % 2 if r["x"] - x_ant < L * 0.07 else 0; x_ant = r["x"]   # etiquetas cercanas se alternan en altura
            ax.text(r["x"], r["NCF"] - 0.05 - 0.17 * nivel, "S%02d\nH %.2f" % (r["sec"], r["H"]), ha="center", va="top", color=C_MUTED, fontsize=7)
        ax.set_xlim(-1, L + 1); ax.set_ylim(min(ncf) - 0.55, max(nct) + 0.45)
        ax.set_xlabel("progresiva (m)", color=C_MUTED); ax.set_ylabel("cota (msnm)", color=C_MUTED); ax.tick_params(colors=C_MUTED, labelsize=8)
        for s in ax.spines.values(): s.set_color(C_LINE)
        ax.grid(color=C_LINE, lw=0.5, alpha=0.6)
        from matplotlib.patches import Patch
        h, lb = ax.get_legend_handles_labels()
        h += [Patch(color="#38bdf8", alpha=0.35), Patch(color="#22c55e", alpha=0.35), Patch(color="#f87171")]
        lb += ["tramo abierto", "tramo tapado", "lleva acero (H > %.2f m o tapado)" % hmin]
        ax.legend(h, lb, loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=5, fontsize=8, facecolor=C_CARD, edgecolor=C_LINE, labelcolor=C_TXT, frameon=False)
        ax.set_title(e["nombre"], color=C_TXT, fontsize=12, loc="left", fontweight="bold")
        self.fig.tight_layout(); self.canvas.draw()
        self.t_tramos.setRowCount(0)
        for j, t in enumerate(e["tramos"]):
            r = self.t_tramos.rowCount(); self.t_tramos.insertRow(r)
            for c, v in enumerate(("%02d" % (j + 1), t["tipo"], "0+%06.2f" % t["a"], "0+%06.2f" % t["b"], "%.2f" % t["L"], "%.2f" % t["Hi"], "%.2f" % t["Hf"])):
                it = QTableWidgetItem(v); it.setTextAlignment(Qt.AlignmentFlag.AlignCenter); self.t_tramos.setItem(r, c, it)

    # ------------------------------------------------------------------ pagina 3: metrado
    def pag_metrado(self):
        w = QWidget(); V = QVBoxLayout(w); V.setContentsMargins(28, 24, 28, 24); V.setSpacing(12)
        t = QLabel("Metrado"); t.setObjectName("h1"); V.addWidget(t)
        g = QGridLayout(); g.setSpacing(10); V.addLayout(g)
        self.k_met = {}
        for i, (k, tit) in enumerate((("conc", "Concreto en cunetas (m3)"), ("enc", "Encofrado (m2)"), ("ace", "Acero (kg)"), ("cur", "Curado (m2)"), ("rej", "Rejillas (ml)"), ("jun", "Juntas (ml)"), ("exc", "Excavacion (m3)"), ("sol", "Solado (m2)"))):
            c = kpi(tit); self.k_met[k] = c; g.addWidget(c, i // 4, i % 4)
        tabs = QTabWidget(); V.addWidget(tabs, 1)
        self.t_acero = QTableWidget(0, 9); self.t_acero.setHorizontalHeaderLabels(["N", "Eje", "Tramo", "Tipo", "Sec. ini", "Sec. fin", "L (m)", "H ini - fin", "Marco @ (m)"]); self.t_acero.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        tabs.addTab(self.t_acero, "Tramos con acero")
        self.t_res = QTableWidget(0, 4); self.t_res.setHorizontalHeaderLabels(["Item", "Partida", "Und", "Metrado"]); self.t_res.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        tabs.addTab(self.t_res, "Resumen de partidas")
        self.t_pend = QTextEdit(); self.t_pend.setReadOnly(True); tabs.addTab(self.t_pend, "Datos pendientes")
        return w

    def llenar_metrado(self):
        res = self.res; cfg = self.cfg_actual(); V = res["valores"].get("RESUMEN", {})
        # KPIs: busca las partidas por su descripcion en el RESUMEN evaluado
        def buscar(texto):
            for coord, v in V.items():
                if coord.startswith("B") and isinstance(v, str) and texto in v.upper():
                    r = coord[1:]; return V.get("D" + r)
            return None
        claves = dict(conc="CONCRETO P/CUNETAS", enc="ENCOFRADO Y DESENCOFRADO DE CUNETA", ace="ACERO F'Y=4200 KG/CM2 PARA CUNETA", cur="CURADO DE CONCRETO EN CUNETA", rej="REJILLA METALICA", jun="JUNTA DE DILATACION", exc="EXCAVACIÓN MANUAL DE ZANJAS P/CUNETAS", sol="SOLADO")
        for k, tx in claves.items():
            v = buscar(tx); self.k_met[k].valor.setText(fmt(v) if v is not None else "-")
        self.t_acero.setRowCount(0)
        for i, (e, j, t) in enumerate(res["filas"]):
            Hm = (t["Hi"] + t["Hf"]) / 2; sep = next(sp for lim, sp in cfg["separaciones_marco"] if Hm < lim)
            r = self.t_acero.rowCount(); self.t_acero.insertRow(r)
            for c, v in enumerate((str(i + 1), e["nombre"], "TRAMO %02d" % (j + 1), t["tipo"], "SEC %02d" % t["sec_i"], "SEC %02d" % t["sec_f"], "%.2f" % t["L"], "%.2f - %.2f" % (t["Hi"], t["Hf"]), "%.2f" % sep)):
                it = QTableWidgetItem(v); it.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
                if c == 3: it.setForeground(QBrush(QColor("#22c55e" if t["tipo"] == "TAPADA" else "#38bdf8")))
                self.t_acero.setItem(r, c, it)
        self.t_res.setRowCount(0)
        for r_ in range(9, 80):
            a, b, c, d = V.get("A%d" % r_), V.get("B%d" % r_), V.get("C%d" % r_), V.get("D%d" % r_)
            if a is None and b is None: continue
            r = self.t_res.rowCount(); self.t_res.insertRow(r)
            for cc, v in enumerate((fmt(a), fmt(b), fmt(c), fmt(d) if c else "")):
                it = QTableWidgetItem(v)
                if not c: it.setForeground(QBrush(QColor(C_ACC))); f = it.font(); f.setBold(True); it.setFont(f)
                if cc == 3: it.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
                self.t_res.setItem(r, cc, it)
        self.t_res.resizeColumnsToContents(); self.t_res.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Stretch)
        self.t_pend.setPlainText("\n".join("- " + p for p in res["pend"]) or "Sin pendientes.")

    # ------------------------------------------------------------------ pagina 4: planilla
    def pag_planilla(self):
        w = QWidget(); V = QVBoxLayout(w); V.setContentsMargins(28, 24, 28, 24); V.setSpacing(12)
        t = QLabel("Planilla generada"); t.setObjectName("h1"); V.addWidget(t)
        s = QLabel("Cada pestana es una hoja del Excel que se va a guardar, con sus colores; las formulas se muestran ya calculadas."); s.setObjectName("muted"); V.addWidget(s)
        self.tabs_pl = QTabWidget(); V.addWidget(self.tabs_pl, 1)
        return w

    def llenar_planilla(self):
        self.tabs_pl.clear(); wb = openpyxl.load_workbook(self.res["tmp"])
        for ws in wb.worksheets:
            if ws.title in ("Listas",): continue
            vals = self.res["valores"].get(ws.title, {})
            maxr = min(ws.max_row, 160); maxc = min(ws.max_column, 30)
            t = QTableWidget(maxr, maxc); t.setHorizontalHeaderLabels([openpyxl.utils.get_column_letter(c) for c in range(1, maxc + 1)])
            t.setVerticalHeaderLabels([str(r) for r in range(1, maxr + 1)])
            for c in range(1, maxc + 1):
                wcol = ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width
                t.setColumnWidth(c - 1, int((wcol or 9) * 7.5))
            for row in ws.iter_rows(min_row=1, max_row=maxr, max_col=maxc):
                for cell in row:
                    v = vals.get(cell.coordinate, cell.value)
                    if v is None and (cell.fill is None or cell.fill.fgColor is None or cell.fill.fgColor.rgb in (None, "00000000")): continue
                    nf = (cell.number_format or "").lower()
                    if hasattr(v, "strftime"): v = v.strftime("%d/%m/%Y")                                        # fecha real
                    elif isinstance(v, (int, float)) and not isinstance(v, bool) and ("yy" in nf or "dd" in nf):   # fecha calculada (serie Excel)
                        import datetime as _dt; v = (_dt.datetime(1899, 12, 30) + _dt.timedelta(days=float(v))).strftime("%d/%m/%Y")
                    it = QTableWidgetItem(fmt(v) if not isinstance(v, str) else v)
                    try:
                        rgb = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
                        if isinstance(rgb, str) and len(rgb) == 8 and rgb != "00000000":
                            it.setBackground(QBrush(QColor("#" + rgb[2:]))); it.setForeground(QBrush(QColor("#111111")))
                        elif cell.fill and cell.fill.fill_type == "solid" and cell.fill.fgColor.theme is not None:
                            it.setBackground(QBrush(QColor("#1f4e78"))); it.setForeground(QBrush(QColor("#ffffff")))
                    except Exception: pass
                    if cell.font and cell.font.bold:
                        f = it.font(); f.setBold(True); it.setFont(f)
                    al = cell.alignment.horizontal if cell.alignment else None
                    it.setTextAlignment((Qt.AlignmentFlag.AlignLeft if al == "left" else Qt.AlignmentFlag.AlignRight if al == "right" else Qt.AlignmentFlag.AlignCenter) | Qt.AlignmentFlag.AlignVCenter)
                    t.setItem(cell.row - 1, cell.column - 1, it)
            for rng in ws.merged_cells.ranges:
                if rng.min_row <= maxr and rng.min_col <= maxc:
                    t.setSpan(rng.min_row - 1, rng.min_col - 1, min(rng.max_row, maxr) - rng.min_row + 1, min(rng.max_col, maxc) - rng.min_col + 1)
            t.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
            self.tabs_pl.addTab(t, ws.title)

    # ------------------------------------------------------------------ pagina 5: exportar
    def pag_exportar(self):
        w = QWidget(); V = QVBoxLayout(w); V.setContentsMargins(28, 24, 28, 24); V.setSpacing(14)
        t = QLabel("Exportar"); t.setObjectName("h1"); V.addWidget(t)
        f, lay = card("Excel"); lay.addWidget(QLabel("Guarda la planilla completa (mismo formato y formulas de la plantilla, con las imagenes de encabezado) en el archivo de salida indicado en Proyecto."))
        hb = QHBoxLayout(); b1 = QPushButton("GUARDAR EXCEL"); b1.setObjectName("ok"); b1.clicked.connect(self.guardar_excel); b2 = QPushButton("Abrir en Excel"); b2.clicked.connect(self.abrir_excel)
        hb.addWidget(b1); hb.addWidget(b2); hb.addStretch(); lay.addLayout(hb); V.addWidget(f)
        f, lay = card("Imprimir"); lay.addWidget(QLabel("Imprime o guarda en PDF el resumen de partidas, los tramos con acero y las longitudes por eje (vista previa antes de imprimir)."))
        hb = QHBoxLayout(); b3 = QPushButton("VISTA PREVIA E IMPRIMIR"); b3.setObjectName("primary"); b3.clicked.connect(self.imprimir); b4 = QPushButton("Guardar resumen en PDF"); b4.clicked.connect(self.pdf)
        hb.addWidget(b3); hb.addWidget(b4); hb.addStretch(); lay.addLayout(hb); V.addWidget(f)
        self.l_exp = QLabel(""); self.l_exp.setObjectName("muted"); V.addWidget(self.l_exp); V.addStretch()
        return w

    def guardar_excel(self):
        if not self.res: QMessageBox.warning(self, "MetraP", "Primero lea el plano en Proyecto."); return
        sal = self.e_sal.text().strip()
        if not sal:
            sal, _ = QFileDialog.getSaveFileName(self, "Guardar Excel", "METRADO.xlsx", "Excel (*.xlsx)")
            if not sal: return
            self.e_sal.setText(sal)
        sal = os.path.abspath(sal if sal.lower().endswith(".xlsx") else sal + ".xlsx"); os.makedirs(os.path.dirname(sal), exist_ok=True)
        shutil.copyfile(self.res["tmp"], sal); self.ultimo = sal; self.l_exp.setText("Guardado: " + sal); self.estado.setText("Excel guardado.")
        QMessageBox.information(self, "MetraP", "Planilla guardada en:\n" + sal)

    def abrir_excel(self):
        fn = getattr(self, "ultimo", None) or (self.res and self.res["tmp"])
        if not fn: QMessageBox.warning(self, "MetraP", "Primero lea el plano."); return
        if sys.platform.startswith("win"): os.startfile(fn)
        else: subprocess.Popen(["xdg-open", fn])

    def html_resumen(self):
        res = self.res; V = res["valores"].get("RESUMEN", {}); cfg = self.cfg_actual()
        h = ["<html><body style='font-family:Arial;font-size:10pt'>", "<h2 style='margin:0'>MetraP - Metrado de cunetas</h2>", "<p>%s</p>" % cfg.get("proyecto", "")]
        h.append("<h3>Resumen de partidas</h3><table border='1' cellspacing='0' cellpadding='3' width='100%'><tr style='background:#1f4e78;color:#fff'><th>Item</th><th>Partida</th><th>Und</th><th>Metrado</th></tr>")
        for r_ in range(9, 80):
            a, b, c, d = V.get("A%d" % r_), V.get("B%d" % r_), V.get("C%d" % r_), V.get("D%d" % r_)
            if a is None and b is None: continue
            h.append("<tr%s><td>%s</td><td>%s</td><td>%s</td><td align='right'>%s</td></tr>" % ("" if c else " style='font-weight:bold;background:#e8eef7'", fmt(a), fmt(b), fmt(c), fmt(d) if c else ""))
        h.append("</table><h3>Longitudes por eje</h3><table border='1' cellspacing='0' cellpadding='3'><tr style='background:#1f4e78;color:#fff'><th>Eje</th><th>L dibujada</th><th>Acortamiento</th><th>L metrada</th><th>Abierta</th><th>Tapada</th><th>H prom.</th><th>Area muro</th></tr>")
        for e in res["E"]:
            h.append("<tr><td>%s</td><td align='right'>%.2f</td><td align='right'>%.2f</td><td align='right'>%.2f</td><td align='right'>%.2f</td><td align='right'>%.2f</td><td align='right'>%.3f</td><td align='right'>%.2f</td></tr>" % (e["nombre"], e["L_dib"], e["desc"], e["L"], e["L_ab"], e["L_tap"], e["Hprom"], e["area"]))
        h.append("</table><h3>Tramos con acero (H > %.2f m desde el NCT; tapados siempre)</h3><table border='1' cellspacing='0' cellpadding='3'><tr style='background:#1f4e78;color:#fff'><th>N</th><th>Eje</th><th>Tramo</th><th>Tipo</th><th>L (m)</th><th>H ini - fin</th><th>Marco @</th></tr>" % cfg["h_minima_acero"])
        for i, (e, j, t) in enumerate(res["filas"]):
            Hm = (t["Hi"] + t["Hf"]) / 2; sep = next(sp for lim, sp in cfg["separaciones_marco"] if Hm < lim)
            h.append("<tr><td>%d</td><td>%s</td><td>%02d</td><td>%s</td><td align='right'>%.2f</td><td>%.2f - %.2f</td><td>%.2f</td></tr>" % (i + 1, e["nombre"], j + 1, t["tipo"], t["L"], t["Hi"], t["Hf"], sep))
        h.append("</table>")
        if res["pend"]: h.append("<h3>Datos pendientes</h3><ul>" + "".join("<li>%s</li>" % p for p in res["pend"]) + "</ul>")
        h.append("</body></html>"); return "".join(h)

    def imprimir(self):
        if not self.res: QMessageBox.warning(self, "MetraP", "Primero lea el plano."); return
        doc = QTextDocument(); doc.setHtml(self.html_resumen())
        pr = QPrinter(QPrinter.PrinterMode.HighResolution); dlg = QPrintPreviewDialog(pr, self); dlg.resize(1000, 760)
        dlg.paintRequested.connect(lambda p: doc.print(p)); dlg.exec()

    def pdf(self):
        if not self.res: QMessageBox.warning(self, "MetraP", "Primero lea el plano."); return
        fn, _ = QFileDialog.getSaveFileName(self, "Guardar PDF", "METRADO_RESUMEN.pdf", "PDF (*.pdf)")
        if not fn: return
        doc = QTextDocument(); doc.setHtml(self.html_resumen()); pr = QPrinter(QPrinter.PrinterMode.HighResolution)
        pr.setOutputFormat(QPrinter.OutputFormat.PdfFormat); pr.setOutputFileName(fn); pr.setPageSize(QPageSize(QPageSize.PageSizeId.A4)); doc.print(pr)
        self.l_exp.setText("PDF guardado: " + fn)


def main():
    app = QApplication(sys.argv); app.setStyleSheet(QSS); app.setStyle("Fusion")
    v = Ventana(); v.show(); sys.exit(app.exec())


if __name__ == "__main__":
    main()
