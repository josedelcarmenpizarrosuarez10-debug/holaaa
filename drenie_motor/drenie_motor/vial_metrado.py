"""Metrado del drenaje vial (METRADO_DRENAJE_VIAL.xlsx): planilla con N.° de veces × largo × ancho × alto y fórmulas.

Las cantidades salen de la geometría del cálculo (vial.py) y de las medidas típicas de la biblioteca (cabezal, cama, protección,
juntas). Nombres de partidas de presupuestos viales S10. El acero de refuerzo de cabezales, marcos y badenes no se metra: depende
del diseño estructural.
"""
import math

from openpyxl import Workbook
from shapely.geometry import LineString, Polygon

from . import normas
from .memoria_calculo import AM, AZ, F, cel
from .vial import progresiva
from .vial_dxf import _perfil_cuneta, _v


def _cuneta_geo(c):
    """(ancho superior, área de concreto por m, longitud de la cara de apoyo, longitud de la superficie mojada, área hidráulica llena)."""
    s = c["seccion"]; pts = _perfil_cuneta(s)
    e = 0.075 if c["revestimiento"] == "concreto" else 0.0
    sup = LineString(pts)
    if e > 0:
        off = LineString(pts).offset_curve(-e)
        conc = Polygon(list(pts) + list(reversed(off.coords))).area; apoyo = off.length
    else:
        conc, apoyo = 0.0, sup.length
    agua = Polygon(list(pts) + [pts[0]]).area if s["tipo"] != "rectangular" else s["b"] * s["d"]
    return pts[-1][0] + e, conc, apoyo, sup.length, agua


def _alc_geo(a):
    s = a["seccion"]; circ = s["tipo"] == "circular"; D = s.get("D") or s["h"]; nc = s.get("celdas", 1)
    t = _v("vial_cab_espesor"); sc = _v("vial_cab_sobre_clave"); cim = _v("vial_cab_cimiento"); cama = _v("vial_cama"); rel = _v("vial_relleno_min")
    sep = max(0.40, min(1.00, D / 2)) if circ else 0.30
    luz = nc * (D if circ else s["b"]) + (nc - 1) * sep
    zanja = luz + 2 * 0.30
    ancho_cab = luz + 2 * 0.60; hcab = D + sc + cim
    hueco = nc * (math.pi * D * D / 4 if circ else s["b"] * D)
    ala = 1.2 * ((hcab) + (0.30 + cim)) / 2                     # trapecio de cada ala en elevación (1.20 m de largo)
    Lp = max(1.5, 1.5 * D)
    return {"circ": circ, "D": D, "nc": nc, "t": t, "cama": cama, "rel": rel, "zanja": zanja, "ancho_cab": ancho_cab, "hcab": hcab,
            "hueco": hueco, "ala": ala, "Lp": Lp, "L": a["longitud"] or 8.0}


def libro(ruta, R):
    via = R["via"]; pr = R["proyecto"]
    wb = Workbook(); ws = wb.active; ws.title = "METRADO"
    ws.cell(1, 1, "PLANILLA DE METRADOS – DRENAJE VIAL").font = F(size=11, bold=True)
    ws.cell(2, 1, "PROYECTO: " + pr.get("nombre", "") + ((" - " + pr["codigo"]) if pr.get("codigo") else "")).font = F(size=8)
    if pr.get("ubicacion"): ws.cell(3, 1, "UBICACIÓN: " + pr["ubicacion"]).font = F(size=8)
    cab = ["ÍTEM", "DESCRIPCIÓN", "UND.", "N.° VECES", "LARGO", "ANCHO", "ALTO", "PARCIAL", "TOTAL"]
    for j, (t, w) in enumerate(zip(cab, (9, 62, 7, 9, 10, 10, 10, 11, 12)), 1):
        cel(ws, 5, j, t, bold=True, fill=AZ, al="center"); ws.column_dimensions[chr(64 + j)].width = w
    st = {"r": 6, "res": []}

    def titulo(item, texto):
        r = st["r"]; cel(ws, r, 1, item, bold=True); cel(ws, r, 2, texto, bold=True)
        for j in range(3, 10): cel(ws, r, j, None)
        st["r"] += 1

    def partida(item, texto, und, filas, k=1):
        """filas: [(descripción, veces, largo, ancho, alto)] con None en las medidas que no aplican."""
        r0 = st["r"]; cel(ws, r0, 1, item, bold=True); cel(ws, r0, 2, texto, bold=True); cel(ws, r0, 3, und, al="center")
        for j in range(4, 9): cel(ws, r0, j, None)
        r = r0 + 1
        for d, n, L, A, H in filas:
            cel(ws, r, 2, "   " + d); cel(ws, r, 3, None)
            for j, v in zip((4, 5, 6, 7), (n, L, A, H)): cel(ws, r, j, None if v is None else round(v, 4), "0.00", fill=AM if v is not None else None)
            usados = [c for c, v in zip("DEFG", (n, L, A, H)) if v is not None]
            cel(ws, r, 8, "=" + "*".join("%s%d" % (c, r) for c in usados), "0.00"); cel(ws, r, 9, None)
            r += 1
        cel(ws, r0, 9, "=ROUND(SUM(H%d:H%d),2)" % (r0 + 1, r - 1), "0.00", bold=True)
        st["res"].append((item, texto, und, "I%d" % r0)); st["r"] = r

    junta = 3.0
    if R["cunetas"]:
        titulo("01", "CUNETAS REVESTIDAS DE CONCRETO")
        g = {c["id"]: _cuneta_geo(c) for c in R["cunetas"]}
        partida("01.01", "TRAZO Y REPLANTEO", "m2", [(c["id"], 1, c["longitud"], g[c["id"]][0], None) for c in R["cunetas"]])
        partida("01.02", "EXCAVACIÓN PARA ESTRUCTURAS EN MATERIAL SECO C/MAQUINARIA", "m3",
                [(c["id"] + " (sección hidráulica + revestimiento)", 1, c["longitud"], g[c["id"]][4] + g[c["id"]][1], None) for c in R["cunetas"]])
        partida("01.03", "REFINE, NIVELACIÓN Y COMPACTACIÓN DE TERRENO NATURAL", "m2", [(c["id"], 1, c["longitud"], g[c["id"]][2], None) for c in R["cunetas"]])
        con = [c for c in R["cunetas"] if c["revestimiento"] == "concreto"]
        if con:
            partida("01.04", "CONCRETO SIMPLE f'c = 175 kg/cm2", "m3", [(c["id"] + " (e = 0.075 m)", 1, c["longitud"], g[c["id"]][1], None) for c in con])
            rect = [c for c in con if c["seccion"]["tipo"] == "rectangular"]
            if rect: partida("01.05", "ENCOFRADO Y DESENCOFRADO", "m2", [(c["id"] + " (caras interiores)", 2, c["longitud"], None, c["seccion"]["d"]) for c in rect])
            partida("01.06", "JUNTAS ASFÁLTICAS E = 1\"", "m", [(c["id"] + " (cada %.1f m)" % junta, max(0, math.ceil(c["longitud"] / junta) - 1), g[c["id"]][3], None, None) for c in con])
    if R["alcantarillas"]:
        titulo("02", "ALCANTARILLAS")
        g = {a["id"]: _alc_geo(a) for a in R["alcantarillas"]}
        A = R["alcantarillas"]
        partida("02.01", "TRAZO Y REPLANTEO", "m2", [(a["id"], 1, g[a["id"]]["L"] + 2 * g[a["id"]]["t"] + 2 * g[a["id"]]["Lp"], g[a["id"]]["ancho_cab"], None) for a in A])
        partida("02.02", "EXCAVACIÓN PARA ESTRUCTURAS EN MATERIAL SECO C/MAQUINARIA", "m3",
                [(a["id"] + " (zanja hasta la clave)", 1, g[a["id"]]["L"], g[a["id"]]["zanja"], g[a["id"]]["D"] + g[a["id"]]["cama"]) for a in A] +
                [(a["id"] + " (cimiento de cabezales)", 2, g[a["id"]]["ancho_cab"], g[a["id"]]["t"], _v("vial_cab_cimiento")) for a in A])
        partida("02.03", "CAMA DE GRAVA ARENOSA E = %.2f m" % _v("vial_cama"), "m2", [(a["id"], 1, g[a["id"]]["L"], g[a["id"]]["zanja"], None) for a in A])
        partida("02.04", "RELLENO CON MATERIAL DE PRÉSTAMO SELECCIONADO", "m3",
                [(a["id"] + " (zanja menos conducto)", 1, g[a["id"]]["L"], None, g[a["id"]]["zanja"] * g[a["id"]]["D"] - g[a["id"]]["hueco"]) for a in A] +
                [(a["id"] + " (sobre la clave)", 1, g[a["id"]]["L"], g[a["id"]]["zanja"], g[a["id"]]["rel"]) for a in A])
        partida("02.05", "CONCRETO f'c = 210 kg/cm2 (CABEZALES Y ALAS)", "m3",
                [(a["id"] + " (cabezales menos huecos)", 2, None, g[a["id"]]["t"], g[a["id"]]["ancho_cab"] * g[a["id"]]["hcab"] - g[a["id"]]["hueco"]) for a in A] +
                [(a["id"] + " (alas)", 4, None, g[a["id"]]["t"], g[a["id"]]["ala"]) for a in A])
        partida("02.06", "ENCOFRADO Y DESENCOFRADO", "m2",
                [(a["id"] + " (cabezales, dos caras)", 4, None, None, g[a["id"]]["ancho_cab"] * g[a["id"]]["hcab"] - g[a["id"]]["hueco"]) for a in A] +
                [(a["id"] + " (alas, dos caras)", 8, None, None, g[a["id"]]["ala"]) for a in A])
        partida("02.07", "MAMPOSTERÍA CON PIEDRA EMBOQUILLADA", "m3",
                [(a["id"] + " (entrada y salida, e = 0.30 m)", 2, g[a["id"]]["Lp"], g[a["id"]]["ancho_cab"], 0.30) for a in A] +
                [(a["id"] + " (uñas)", 2, g[a["id"]]["ancho_cab"], 0.30, 0.50) for a in A])
        tipos = {}
        for a in A:
            if g[a["id"]]["circ"]: tipos.setdefault(a["seccion"]["pulg"], []).append(a)
        for k, (p, lista) in enumerate(sorted(tipos.items())):
            partida("02.%02d" % (8 + k), "ALCANTARILLA METÁLICA CIRCULAR TMC Ø = %.0f\"" % p, "m",
                    [(a["id"] + " " + progresiva(a["prog"]), g[a["id"]]["nc"], g[a["id"]]["L"], None, None) for a in lista])
        mca = [a for a in A if not g[a["id"]]["circ"]]
        if mca:
            partida("02.%02d" % (8 + len(tipos)), "ALCANTARILLA MARCO DE CONCRETO ARMADO (según diseño estructural)", "m",
                    [(a["id"] + " " + a["conducto"], g[a["id"]]["nc"], g[a["id"]]["L"], None, None) for a in mca])
    if R["badenes"]:
        titulo("03", "BADENES")
        anchos = {b["id"]: float(via.get("ancho_plataforma", 7.0)) for b in R["badenes"]}
        desarrollo = lambda b: b["seccion"]["b"] + 2 * b["profundidad"] * math.hypot(1, b["seccion"]["z1"])
        partida("03.01", "TRAZO Y REPLANTEO", "m2", [(b["id"], 1, desarrollo(b), anchos[b["id"]] + 2 * 1.5, None) for b in R["badenes"]])
        partida("03.02", "CONCRETO f'c = 210 kg/cm2 (LOSA E = 0.20 m)", "m3", [(b["id"], 1, desarrollo(b), anchos[b["id"]], 0.20) for b in R["badenes"]] +
                [(b["id"] + " (uñas de entrada y salida)", 2, desarrollo(b), 0.30, 0.80) for b in R["badenes"]])
        partida("03.03", "MAMPOSTERÍA CON PIEDRA EMBOQUILLADA (PROTECCIÓN AGUAS ABAJO)", "m3", [(b["id"], 1, desarrollo(b), 1.5, 0.30) for b in R["badenes"]])
    # resumen
    wr = wb.create_sheet("RESUMEN", 0)
    wr.cell(1, 1, "RESUMEN DE METRADOS – DRENAJE VIAL").font = F(size=11, bold=True)
    for j, (t, w) in enumerate(zip(["ÍTEM", "DESCRIPCIÓN", "UND.", "METRADO"], (9, 70, 8, 13)), 1):
        cel(wr, 3, j, t, bold=True, fill=AZ, al="center"); wr.column_dimensions[chr(64 + j)].width = w
    for i, (item, texto, und, ref) in enumerate(st["res"]):
        cel(wr, 4 + i, 1, item); cel(wr, 4 + i, 2, texto); cel(wr, 4 + i, 3, und, al="center"); cel(wr, 4 + i, 4, "=METRADO!%s" % ref, "0.00", bold=True)
    n = 6 + len(st["res"])
    for k, t in enumerate(["Cantidades de la geometría del cálculo; cabezal, cama, relleno, protección y juntas con las medidas típicas de la biblioteca "
                           "(criterio del proyectista, editables).",
                           "No incluye acero de refuerzo de cabezales, marcos ni badenes (según el diseño estructural), ni la eliminación de material excedente."]):
        wr.cell(n + k, 1, "• " + t).font = F(size=8)
    wb.save(ruta)
    return ruta
