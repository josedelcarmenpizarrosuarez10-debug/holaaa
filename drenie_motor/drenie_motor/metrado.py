"""Metrado general de drenaje pluvial (METRADO_DRENAJE_PLUVIAL.xlsx), con formulas, en el formato del metrado del CAR Mujeres
(CUI 2717013) y de Barranquita (C.I. 468901): nombres y estructura de partidas Delphin del Hogar de Refugio (CUI 2675514);
cantidades siempre del diseno de este proyecto.

Hojas: RESUMEN, PLANILLA GENERAL DE METRADOS, CONCRETO EN CUNETAS, ENCOFRADO Y DESENCOFRADO, METRADO ACERO, METRADO JUNTA DE
DILATACION, METRADO DE CURADO, METRADO DE REJILLAS, PLANILLA TAPAS INSPECCION, METRADO LLORADEROS, METRADO MONTANTES,
METRADO CANALETAS, METRADO  EN PISO (SUMIDERO), PARAMETROS y Listas. Cada eje lleva su ancho libre b (una medida por sector).
Criterios (conocimiento/06_metrado_acu_especificaciones.md, seccion 5):
- H del NCT a la cota de fondo, perfil cada 0.50 m; area de muro = integral de la altura libre del muro (bajo la losa tapa en tapados).
- Concreto: 2 muros x e x area + losa de fondo (b + 2e) x ef x L + losa tapa (b + 2e) x et x L tapado - tapas de inspeccion.
- Encofrado: caras interior y exterior de los 2 muros (exterior con la losa de fondo) + fondo y cantos de la losa tapa.
- Acero 3/8" solo con h >= 0.40 m o en tramo tapado: U @0.25 con ganchos 0.10 y longitudinales @0.20; losa tapa @0.20 y 3 barras.
- Curado: caras interiores de muros, fondo, coronas de muro (tramos con rejilla) y losa tapa.
- Excavacion: (b + 2e + 2 sobreanchos) x (H + losa de fondo + solado) x L; relleno a los lados; eliminacion con esponjamiento.
- Juntas de dilatacion cada 3.00 m en toda la seccion en U (2 (h + ef) + b + 2e); junta cuneta - piso a lo largo del borde con piso.
"""
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from shapely.geometry import LineString
from shapely.ops import unary_union
import shapely

_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
LETRA = "Arial Narrow"
CAB, SEC, PAR_, AZC, VER = "305496", "D9EAF7", "FFF2CC", "DDEBF7", "F0F7E8"
SOBREANCHO, ESPONJ, JUNTA_CADA, REC, KG38, SEP_T, SEP_L = 0.15, 1.25, 3.00, 0.04, 0.56, 0.25, 0.20
DADO = (0.40, 0.40, 0.35)
FC = {"concreto": 1.30 * 0.15 * 0.20, "encofrado": 2 * 1.30 * 0.15 + 1.30 * 0.20, "acero": (1.20 * 4 + 0.58 * 9) * 0.56}
TIPOS_PISO = ("vereda", "patio", "pavimento", "estacionamiento", "juegos", "techo")


def F(**k): return Font(name=LETRA, size=k.pop("size", 11), **k)


def cel(ws, r, c, v, fmt=None, bold=False, bg=None, color=None, al="center", wrap=False, borde=True, size=11):
    x = ws.cell(r, c, v); x.font = F(bold=bold, size=size, color=color)
    x.alignment = Alignment(horizontal=al, vertical="center", wrap_text=wrap)
    if fmt: x.number_format = fmt
    if bg: x.fill = PatternFill("solid", fgColor=bg)
    if borde: x.border = BOR
    return x


def cab_hoja(ws, titulo, ncol, col0=1, fila=1, lab=1):
    cel(ws, fila, col0, titulo, bold=True, size=14, borde=False)
    ws.merge_cells(start_row=fila, start_column=col0, end_row=fila, end_column=col0 + ncol - 1)
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        r = fila + 1 + i
        cel(ws, r, col0, k, bold=True, al="left"); cel(ws, r, col0 + lab, ref, al="left", wrap=True)
        if lab > 1: ws.merge_cells(start_row=r, start_column=col0, end_row=r, end_column=col0 + lab - 1)
        ws.merge_cells(start_row=r, start_column=col0 + lab, end_row=r, end_column=col0 + ncol - 1)
    ws.row_dimensions[fila + 2].height = 33


def nota(ws, r, c0, c1, texto, alto=30):
    cel(ws, r, c0, texto, al="left", wrap=True, borde=False)
    ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1); ws.row_dimensions[r].height = alto


def encabezado(ws, r, textos, anchos=None, c0=1, bg=CAB, color="FFFFFF", alto=32):
    for j, t in enumerate(textos):
        cel(ws, r, c0 + j, t, bold=True, bg=bg, color=color, wrap=True)
        if anchos: ws.column_dimensions[CL(c0 + j)].width = anchos[j]
    ws.row_dimensions[r].height = alto


# ---------------------------------------------------------------------------------------------- cantidades por eje
def _tap(r, s): return any(a - 1e-6 <= s <= b + 1e-6 for a, b in r["tapados"])


def cantidades(C, P, anchos, tapas, piso):
    EM, EF, ET, SOL = P["e_muro"], P["e_fondo"], P["e_tapa"], P["e_solado"]
    out = []
    for r in C["ramales"]:
        t0 = r["tramos"][0]; b = anchos.get((r["codigo"], t0["de"], t0["a"]), P["b_int"]); be = b + 2 * EM; bexc = be + 2 * SOBREANCHO
        Pf = r["perfil"]; L = r["L"]
        L_tap = sum(bb - a for a, bb in r["tapados"]); L_rej = L - L_tap
        int_h = int_hw = int_ho = int_hac = L_ac = hmax = 0.0
        for (s0, n0, f0), (s1, n1, f1) in zip(Pf[:-1], Pf[1:]):
            ds = s1 - s0
            if ds <= 0: continue
            h = ((n0 - f0) + (n1 - f1)) / 2; tap = _tap(r, (s0 + s1) / 2); hmax = max(hmax, n0 - f0, n1 - f1)
            hw = h - (ET if tap else 0.0)
            int_h += h * ds; int_hw += hw * ds; int_ho += (h + EF) * ds
            if tap or h >= P["h_acero"] - 1e-6: L_ac += ds; int_hac += h * ds
        h_prom = int_h / L if L else 0.0; h_ac = int_hac / L_ac if L_ac else 0.0
        n_tapas = sum(1 for t in tapas if t["eje"] == r["codigo"])
        exc = bexc * (int_ho + SOL * L); rell = (bexc - be) * (int_ho + SOL * L)
        jp = 0.0
        if piso is not None and not piso.is_empty and L > 0:
            ln = LineString([tuple(C["nudos"][n]) for n in r["nudos"]]); k = 0.25
            while k < ln.length:
                p = ln.interpolate(k); q = ln.interpolate(min(ln.length, k + 0.01)); o = ln.interpolate(max(0.0, k - 0.01))
                dx, dy = q.x - o.x, q.y - o.y; d = math.hypot(dx, dy) or 1.0; nx, ny = -dy / d, dx / d
                for sg in (1, -1):
                    if shapely.contains_xy(piso, p.x + sg * nx * (be / 2 + 0.15), p.y + sg * ny * (be / 2 + 0.15)): jp += 0.5
                k += 0.5
        out.append({"eje": r["codigo"], "b": round(b, 3), "be": round(be, 3), "bexc": round(bexc, 3), "L": round(L, 2), "L_tapado": round(L_tap, 2),
                    "L_rejilla": round(L_rej, 2), "h_prom": round(h_prom, 3), "h_max": round(hmax, 2), "area_muro": round(int_hw, 2),
                    "hw": round(int_hw / L, 3) if L else 0.0, "L_acero": round(L_ac, 2), "h_acero": round(h_ac, 3), "tapas": n_tapas,
                    "excavacion": round(exc, 3), "relleno": round(rell, 3), "junta_piso": round(min(jp, 2 * L), 2),
                    "tipo": "TAPADA" if L_rej <= 0 else ("MIXTA" if L_tap > 0 else "ABIERTA")})
    return out


# ---------------------------------------------------------------------------------------------- libro
PAR = [("Espesor de muros (m)", "e_muro", "em", "Secciones transversales"), ("Espesor de losa de fondo (m)", "e_fondo", "ef", "Secciones transversales"),
       ("Espesor de losa tapa en tramos tapados (m)", "e_tapa", "et", "Secciones transversales"),
       ("Espesor de solado (m)", "e_solado", "sol", "Solado bajo la losa de fondo"),
       ("Sobreancho de excavacion a cada lado (m)", SOBREANCHO, "sa", "Espacio de trabajo y encofrado"),
       ("Factor de esponjamiento para eliminacion", ESPONJ, "esp", "Eliminacion = (excavacion - relleno) x factor"),
       ("Junta de dilatacion en la seccion en U cada (m)", JUNTA_CADA, "jc", "Juntas en toda la seccion en U (EG-2013, seccion 510)"),
       ("Recubrimiento del acero (m)", REC, "rec", "Secciones transversales"), ("Peso de la barra de 3/8\" (kg/m)", KG38, "kg", "Hoja Listas"),
       ("Separacion de las U de 3/8\" (m)", SEP_T, "st", "Secciones transversales"), ("Separacion de barras longitudinales de 3/8\" (m)", SEP_L, "sl", "Secciones transversales"),
       ("Profundidad minima con acero (m): h >= 0.40 o tramo tapado", "h_acero", "hac", "Con h menor la cuneta abierta es de concreto simple"),
       ("Tapa de inspeccion: lado (m)", 0.60, "tl", "Planta general: tapa 0.60 x 0.60"), ("Tapa de inspeccion: espesor (m)", 0.10, "te", "Planta general"),
       ("Tapa de inspeccion: separacion de la parrilla de 3/8\" (m)", 0.15, "tsep", "Parrilla simple en ambos sentidos"),
       ("Lloradero: espaciamiento (m)", "llor_sep", "lls", "Solo en el muro que da a area verde"),
       ("Canaleta: separacion de soportes (m)", 0.60, "css", "Soportes = REDONDEAR.MAS(L / sep) + 1"),
       ("Dado de montante: largo x ancho x alto (m3)", DADO[0] * DADO[1] * DADO[2], "dado", "0.40 x 0.40 x 0.35"),
       ("Falsa columna: concreto por und (m3)", FC["concreto"], "fcc", "0.20 x 0.15 x 1.30"),
       ("Falsa columna: encofrado por und (m2)", FC["encofrado"], "fce", "2 x 1.30 x 0.15 + 1.30 x 0.20"),
       ("Falsa columna: acero por und (kg)", FC["acero"], "fca", "4 barras de 1.20 + 9 estribos de 0.58, 3/8\"")]


def libro(ruta, C, red, M, P, anchos, V, LLt, TT, SM, proyecto, cfg, tapas):
    piso = unary_union([c["geom"] for c in M.caras if c["tipo"] in TIPOS_PISO])
    q = cantidades(C, P, anchos, tapas, piso)
    Pk = {k: "PARAMETROS!$B$%d" % (5 + i) for i, (_, _, k, _) in enumerate(PAR)}
    pref = (cfg.get("metrado") or {}).get("prefijo", "01")
    wb = Workbook()
    wr = wb.active; wr.title = "RESUMEN"
    nombres = ["PLANILLA GENERAL DE METRADOS", "CONCRETO EN CUNETAS", "ENCOFRADO Y DESENCOFRADO", "METRADO ACERO", "METRADO JUNTA DE DILATACION",
               "METRADO DE CURADO", "METRADO DE REJILLAS", "PLANILLA TAPAS INSPECCION", "METRADO LLORADEROS", "METRADO MONTANTES",
               "METRADO CANALETAS", "METRADO  EN PISO (SUMIDERO)", "PARAMETROS", "Listas"]
    W = {n: wb.create_sheet(n) for n in nombres}
    wp, wc, we, wa, wj, wu, wg, wt, wll, wm, wca, wsu, wq, wl = (W[n] for n in nombres)

    # PARAMETROS y Listas
    cel(wq, 1, 1, "PARÁMETROS DEL METRADO", bold=True, size=14, borde=False); wq.merge_cells("A1:D1")
    cel(wq, 2, 1, "Valores editables: las hojas de detalle los usan en sus formulas. El ancho libre b de cada eje esta en cada hoja.", al="left", borde=False)
    encabezado(wq, 4, ["Descripción", "Valor", "Unidad", "Criterio"], [62, 12, 9, 58])
    for i, (t, v, k, crit) in enumerate(PAR):
        val = P[v] if isinstance(v, str) else v
        cel(wq, 5 + i, 1, t, al="left"); cel(wq, 5 + i, 2, val, "0.0000", bold=True, bg=PAR_); cel(wq, 5 + i, 3, "kg/m" if k == "kg" else ("-" if k == "esp" else "m"))
        cel(wq, 5 + i, 4, crit, al="left")
    cel(wl, 1, 1, "LISTAS Y FACTORES", bold=True, size=14, borde=False); wl.merge_cells("A1:F1")
    encabezado(wl, 3, ["DIÁMETRO", "FACTOR kg/m", "", "TIPO DE CUNETA"], [12, 13, 3, 18])
    for i, (d, f) in enumerate((("6 mm", 0.222), ("8 mm", 0.395), ("3/8\"", 0.56), ("12 mm", 0.888), ("1/2\"", 0.994), ("5/8\"", 1.552), ("3/4\"", 2.235), ("1\"", 3.973))):
        cel(wl, 4 + i, 1, d); cel(wl, 4 + i, 2, f, "0.000")
    for i, t in enumerate(("ABIERTA", "TAPADA", "MIXTA")): cel(wl, 4 + i, 4, t)

    # CONCRETO EN CUNETAS
    cab_hoja(wc, "CONCRETO F´C=175 KG/CM2 PARA CUNETAS - DRENAJE PLUVIAL", 14, col0=2, fila=2)
    cel(wc, 7, 2, "OBRAS DE CONCRETO ARMADO", bold=True, color="FFFFFF", bg="0070C0"); wc.merge_cells("B7:O7")
    for col, w in zip("ABCDEFGHIJKLMNO", [3, 14, 17, 12, 11, 11, 13, 4, 4, 14, 17, 12, 11, 11, 13]): wc.column_dimensions[col].width = w
    ref_conc = {}
    for i, x in enumerate(q):
        c0 = 2 if i % 2 == 0 else 10; r0 = 9 + (i // 2) * 16; cc = [CL(c0 + k) for k in range(6)]
        for k, t in ((0, "LONGITUDES Y AREAS"), (3, "m3"), (4, "Sub.TOTAL"), (5, "TOTAL (m3)")): cel(wc, r0, c0 + k, t, bold=True, color="FFFFFF", bg="0070C0")
        wc.merge_cells("%s%d:%s%d" % (cc[0], r0, cc[2], r0))
        be = "=%.2f+2*%s" % (x["b"], Pk["em"])
        filas = [("MUROS", "AREA", "=%.2f*%.3f" % (x["L"], x["hw"]), "0.00"), (None, "ESPESOR", "=%s" % Pk["em"], "0.00"), (None, "CANT.", 2, "0"),
                 ("LOSA", "ESPESOR", "=%s" % Pk["ef"], "0.00"), (None, "LONGITUD", x["L"], "0.00"), (None, "ANCHO (b + 2e)", be, "0.00"),
                 ("TAPA", "LONGITUD", x["L_tapado"], "0.00"), (None, "ESPESOR", "=%s" % Pk["et"], "0.00"), (None, "ANCHO (b + 2e)", be, "0.00"),
                 ("DESC. TAPAS", "CANT.", -x["tapas"], "0"), (None, "LADO", "=%s" % Pk["tl"], "0.00"), (None, "ESPESOR", "=%s" % Pk["te"], "0.00")]
        for k, (grupo, nom, v, fmt) in enumerate(filas):
            r = r0 + 1 + k
            cel(wc, r, c0 + 1, grupo, bold=True); cel(wc, r, c0 + 2, nom); cel(wc, r, c0 + 3, v, fmt)
            if k % 3 == 0:
                a, b_, c_ = "%s%d" % (cc[3], r), "%s%d" % (cc[3], r + 1), "%s%d" % (cc[3], r + 2)
                cel(wc, r, c0 + 4, ("=%s*%s*%s" % (a, b_, c_)) if k < 9 else ("=%s*%s*%s*%s" % (a, b_, b_, c_)), "0.000", bold=True)
            else: cel(wc, r, c0 + 4, None)
            cel(wc, r, c0 + 5, None)
        cel(wc, r0 + 1, c0, "CUNETA %s\nb = %.2f" % (x["eje"], x["b"]), bold=True, wrap=True)
        wc.merge_cells("%s%d:%s%d" % (cc[0], r0 + 1, cc[0], r0 + 12))
        cel(wc, r0 + 1, c0 + 5, "=SUM(%s%d:%s%d)" % (cc[4], r0 + 1, cc[4], r0 + 12), "0.000", bold=True, bg=PAR_)
        wc.merge_cells("%s%d:%s%d" % (cc[5], r0 + 1, cc[5], r0 + 12))
        for k in range(4): wc.merge_cells("%s%d:%s%d" % (cc[1], r0 + 1 + 3 * k, cc[1], r0 + 3 + 3 * k))
        ref_conc[x["eje"]] = "%s%d" % (cc[5], r0 + 1)

    # ENCOFRADO Y DESENCOFRADO
    cab_hoja(we, "ENCOFRADO Y DESENCOFRADO CUNETAS - DRENAJE PLUVIAL", 15, col0=2, fila=2)
    cel(we, 8, 2, "OBRAS DE CONCRETO ARMADO - ENCOFRADO Y DESENCOFRADO", bold=True, color="FFFFFF", bg="0070C0"); we.merge_cells("B8:P8")
    for col, w in zip("ABCDEFGHIJKLMNOP", [3, 12, 12, 17, 16, 11, 12, 13, 4, 12, 12, 17, 16, 11, 12, 13]): we.column_dimensions[col].width = w
    ref_enc = {}
    for i, x in enumerate(q):
        c0 = 2 if i % 2 == 0 else 10; r0 = 11 + (i // 2) * 11; cc = [CL(c0 + k) for k in range(7)]
        for k, t in ((0, "LONGITUDES Y AREAS"), (5, "Sub.TOTAL"), (6, "TOTAL (m2)")): cel(we, r0, c0 + k, t, bold=True, color="FFFFFF", bg="0070C0")
        we.merge_cells("%s%d:%s%d" % (cc[0], r0, cc[4], r0))
        filas = [("MURO EXTERNO", "AREA", "=%.2f*(%.3f+%s)" % (x["L"], x["hw"], Pk["ef"]), "=%s{r}*%s{r1}"), (None, "CANT.", 2, None),
                 ("MURO INTERNO", "AREA", "=%.2f*%.3f" % (x["L"], x["hw"]), "=%s{r}*%s{r1}"), (None, "CANTIDAD", 2, None),
                 ("LOSA TAPA", "LONGITUD", x["L_tapado"], "=%s{r}*%s{r1}"), (None, "FONDO + CANTOS", "=%.2f+2*%s" % (x["b"], Pk["et"]), None)]
        for k, (grupo, nom, v, f) in enumerate(filas):
            r = r0 + 1 + k
            cel(we, r, c0 + 2, grupo, bold=True); cel(we, r, c0 + 3, nom); cel(we, r, c0 + 4, v, "0.00")
            cel(we, r, c0 + 5, (f % (cc[4], cc[4])).format(r=r, r1=r + 1) if f else None, "0.00", bold=True); cel(we, r, c0 + 6, None)
            if k % 2 == 0: we.merge_cells("%s%d:%s%d" % (cc[2], r, cc[2], r + 1))
        cel(we, r0 + 1, c0, x["eje"], bold=True); cel(we, r0 + 1, c0 + 1, x["tipo"], bold=True)
        we.merge_cells("%s%d:%s%d" % (cc[0], r0 + 1, cc[0], r0 + 6)); we.merge_cells("%s%d:%s%d" % (cc[1], r0 + 1, cc[1], r0 + 6))
        cel(we, r0 + 1, c0 + 6, "=SUM(%s%d:%s%d)" % (cc[5], r0 + 1, cc[5], r0 + 6), "0.00", bold=True, bg=PAR_)
        we.merge_cells("%s%d:%s%d" % (cc[6], r0 + 1, cc[6], r0 + 6))
        ref_enc[x["eje"]] = "%s%d" % (cc[6], r0 + 1)

    # METRADO ACERO (columna X = ancho exterior be del eje)
    cab_hoja(wa, "METRADO DE ACERO EN CUNETAS - DRENAJE PLUVIAL", 24, lab=2)
    cel(wa, 8, 1, "METRADO DETALLADO POR EJE (ACERO SOLO CON h >= 0.40 m O EN TRAMO TAPADO)", bold=True, al="left", borde=False)
    hdr = ["N°", "EJE", "TIPO DE CUNETA", "LONGITUD EJE (m)", "LONG. CON ACERO (m)", "h PROM. CON ACERO (m)", "Ø CUERPO", "FACTOR (kg/m)",
           "SEP. LONG. (m)", "N° BARRAS LONG.", "LONG. ACERO LONG. (m)", "KG LONG.", "SEP. TRANSV. U (m)", "DESARROLLO U (m)", "CANT. U",
           "LONG. TRANSV. (m)", "KG U", "LONG. TAPADA (m)", "SEP. TRANSV. (m)", "DESARROLLO (m)", "BARRAS LONG.", "KG LOSA TAPA", "TOTAL ACERO (kg)",
           "ANCHO EXTERIOR b + 2e (m)"]
    encabezado(wa, 10, hdr, [5, 9, 11, 10, 10, 10, 8, 9, 8, 9, 11, 10, 9, 10, 9, 11, 10, 10, 9, 10, 9, 10, 12, 11], bg="5B9BD5", alto=45)
    fila_acero = {}
    for i, x in enumerate(q):
        r = 11 + i; fila_acero[x["eje"]] = r
        vals = [(i + 1, "0"), (x["eje"], None), (x["tipo"], None), (x["L"], "0.00"), (x["L_acero"], "0.00"), (x["h_acero"], "0.000"),
                ("3/8\"", None), ('=IFERROR(VLOOKUP(G{r},Listas!$A$4:$B$11,2,FALSE),0)', "0.000"), ("={sl}", "0.00"),
                ("=IF(E{r}>0,2*(CEILING(ROUND(F{r}/I{r},6),1)+1)+(CEILING(ROUND((X{r}-2*{rec})/I{r},6),1)+1),0)", "0"),
                ("=E{r}*J{r}", "0.00"), ("=K{r}*H{r}", "0.00"),
                ("={st}", "0.00"), ("=IF(E{r}>0,2*(F{r}+{ef}-2*{rec})+(X{r}-2*{rec})+0.2,0)", "0.000"), ("=IF(E{r}>0,E{r}/M{r},0)", "0.00"),
                ("=N{r}*O{r}", "0.00"), ("=P{r}*H{r}", "0.00"),
                (x["L_tapado"], "0.00"), ("={sl}", "0.00"), ("=IF(R{r}>0,X{r}-2*{rec}+0.2,0)", "0.000"), ("=IF(R{r}>0,3,0)", "0"),
                ("=IF(AND(E{r}>0,R{r}>0),((R{r}/S{r})*T{r}+U{r}*R{r})*H{r},0)", "0.00"), ("=L{r}+Q{r}+V{r}", "0.00"),
                ("=%.2f+2*{em}" % x["b"], "0.00")]
        for j, (v, fmt) in enumerate(vals, 1):
            if isinstance(v, str) and v.startswith("="): v = v.format(r=r, **Pk)
            cel(wa, r, j, v, fmt, bold=(j == 23), bg=PAR_ if j == 23 else None)
    rt = 11 + len(q)
    cel(wa, rt, 1, "TOTAL", bold=True); wa.merge_cells(start_row=rt, start_column=1, end_row=rt, end_column=22)
    cel(wa, rt, 23, "=SUM(W11:W%d)" % (rt - 1), "0.00", bold=True, bg=PAR_)

    # JUNTAS
    cab_hoja(wj, "PLANILLA DE METRADOS - JUNTAS DE CUNETAS", 10)
    cel(wj, 7, 1, "Espaciamiento entre juntas (m)", al="left"); cel(wj, 7, 2, "=%s" % Pk["jc"], "0.00", bg=PAR_)
    encabezado(wj, 9, ["ITEM", "EJE / DESCRIPCIÓN", "UND.", "LONGITUD EJE (m)", "h PROM. (m)", "ESPACIAMIENTO (m)", "N° DE JUNTAS",
                       "DESARROLLO DE LA U (m)", "b + 2e (m)", "PARCIAL (m)"], [14, 40, 7, 12, 10, 14, 10, 14, 10, 12], alto=45)
    rj0 = 10
    for i, x in enumerate(q):
        r = rj0 + i
        for j, (v, fmt) in enumerate([(i + 1, "0"), (x["eje"], None), ("m", None), (x["L"], "0.00"), (x["h_prom"], "0.000"), ("=$B$7", "0.00"),
                                       ("=INT(ROUND(D%d/F%d,6))" % (r, r), "0"), ("=2*(E%d+%s)+I%d" % (r, Pk["ef"], r), "0.000"),
                                       ("=%.2f+2*%s" % (x["b"], Pk["em"]), "0.00"), ("=G%d*H%d" % (r, r), "0.00")], 1):
            cel(wj, r, j, v, fmt)
    rj1 = rj0 + len(q)
    cel(wj, rj1, 1, "TOTAL JUNTA DE DILATACION", bold=True); wj.merge_cells(start_row=rj1, start_column=1, end_row=rj1, end_column=9)
    cel(wj, rj1, 10, "=SUM(J%d:J%d)" % (rj0, rj1 - 1), "0.00", bold=True, bg=PAR_)
    r = rj1 + 2
    encabezado(wj, r, ["ITEM", "EJE (junta cuneta - piso adyacente)", "UND.", "LONGITUD EJE (m)", "BORDE JUNTO A PISO (m)", "", "", "", "", "PARCIAL (m)"], alto=32)
    ok_p = [x for x in q if x["junta_piso"] > 0]; rp0 = r + 1
    for i, x in enumerate(ok_p):
        rr = rp0 + i
        for j, (v, fmt) in enumerate([(i + 1, "0"), (x["eje"], None), ("m", None), (x["L"], "0.00"), (x["junta_piso"], "0.00"),
                                       (None, None), (None, None), (None, None), (None, None), ("=E%d" % rr, "0.00")], 1):
            cel(wj, rr, j, v, fmt)
    rp1 = rp0 + len(ok_p)
    cel(wj, rp1, 1, "TOTAL JUNTA CUNETA - PISO", bold=True); wj.merge_cells(start_row=rp1, start_column=1, end_row=rp1, end_column=9)
    cel(wj, rp1, 10, "=SUM(J%d:J%d)" % (rp0, max(rp0, rp1 - 1)), "0.00", bold=True, bg=PAR_)

    # CURADO
    cab_hoja(wu, "METRADO DE CURADO DE CONCRETO EN CUNETAS", 14, lab=2)
    encabezado(wu, 7, ["N°", "EJE", "TIPO DE CUNETA", "LONGITUD (m)", "LONG. CON REJILLA (m)", "LONG. TAPADA (m)", "ANCHO LIBRE b (m)", "ALTURA LIBRE PROM. (m)",
                       "ÁREA 2 MUROS (m²)", "ÁREA FONDO (m²)", "CORONAS DE MURO (m²)", "ÁREA TAPA (m²)", "TOTAL CURADO (m²)", "UND."],
               [6, 10, 13, 12, 12, 12, 12, 12, 13, 12, 13, 12, 13, 7], alto=45)
    ru0 = 8
    for i, x in enumerate(q):
        r = ru0 + i
        vals = [(i + 1, "0"), (x["eje"], None), (x["tipo"], None), (x["L"], "0.00"), (x["L_rejilla"], "0.00"), (x["L_tapado"], "0.00"),
                (x["b"], "0.00"), (x["hw"], "0.000"), ("=2*D{r}*H{r}", "0.00"), ("=D{r}*G{r}", "0.00"),
                ("=2*{em}*E{r}", "0.00"), ("=F{r}*(G{r}+2*{em})", "0.00"), ("=SUM(I{r}:L{r})", "0.00"), ("m²", None)]
        for j, (v, fmt) in enumerate(vals, 1):
            if isinstance(v, str) and v.startswith("="): v = v.format(r=r, **Pk)
            cel(wu, r, j, v, fmt, bold=(j == 13), bg=VER if j == 13 else None)
    ru1 = ru0 + len(q)
    cel(wu, ru1, 1, "TOTAL CURADO DE CONCRETO EN CUNETA", bold=True); wu.merge_cells(start_row=ru1, start_column=1, end_row=ru1, end_column=12)
    cel(wu, ru1, 13, "=SUM(M%d:M%d)" % (ru0, ru1 - 1), "0.00", bold=True, bg=PAR_)

    # REJILLAS
    cab_hoja(wg, "METRADO DE REJILLAS PARA CUNETAS", 9, lab=2)
    encabezado(wg, 7, ["N°", "EJE", "TIPO DE CUNETA", "ANCHO LIBRE b (m)", "LONGITUD DEL EJE (m)", "DESCUENTO TRAMOS TAPADOS (m)", "LONGITUD CON REJILLA (m)",
                       "N° DE TAPAS DE INSPECCIÓN", "UND."], [6, 10, 14, 12, 14, 16, 15, 14, 7], alto=45)
    rg0 = 8
    for i, x in enumerate(q):
        r = rg0 + i
        for j, (v, fmt) in enumerate([(i + 1, "0"), (x["eje"], None), (x["tipo"], None), (x["b"], "0.00"), (x["L"], "0.00"), (x["L_tapado"], "0.00"),
                                       ("=IF(E{r}-F{r}<0,0,E{r}-F{r})".format(r=r), "0.00"), (x["tapas"], "0"), ("m", None)], 1):
            cel(wg, r, j, v, fmt, bold=(j == 7), bg=VER if j == 7 else None)
    rg1 = rg0 + len(q)
    cel(wg, rg1, 1, "TOTAL REJILLA METALICA", bold=True); wg.merge_cells(start_row=rg1, start_column=1, end_row=rg1, end_column=6)
    cel(wg, rg1, 7, "=SUM(G%d:G%d)" % (rg0, rg1 - 1), "0.00", bold=True, bg=PAR_); cel(wg, rg1, 8, "=SUM(H%d:H%d)" % (rg0, rg1 - 1), "0", bold=True)

    # TAPAS
    cab_hoja(wt, "METRADO DE TAPAS DE INSPECCIÓN EN CUNETAS TAPADAS", 11)
    encabezado(wt, 7, ["ITEM", "DESCRIPCIÓN", "Nº DE VECES", "LARGO", "ANCHO", "ALTURA", "FACT.", "AREA", "PARCIAL", "TOTAL", "UND"], [14, 52, 11, 10, 10, 10, 8, 10, 11, 11, 7])
    cel(wt, 8, 2, "TAPA DE INSPECCION 0.60m x 0.60m", bold=True, al="left", bg=PAR_); cel(wt, 8, 11, "und", bold=True, bg=PAR_)
    for i, t in enumerate(tapas):
        r = 9 + i
        cel(wt, r, 1, i + 1, "0"); cel(wt, r, 2, "%s - progresiva %.2f m (tramo tapado)" % (t["eje"], t["progresiva"]), al="left"); cel(wt, r, 3, 1, "0")
        cel(wt, r, 9, "=C%d" % r, "0"); cel(wt, r, 11, "und")
    cel(wt, 8, 10, "=SUM(I9:I%d)" % max(9, 8 + len(tapas)), "0", bold=True, bg=PAR_)

    # LLORADEROS
    cab_hoja(wll, "METRADO DE LLORADEROS (CUNETA COLINDANTE CON AREA VERDE)", 8)
    encabezado(wll, 7, ["N°", "EJE", "TRAMO", "LADO DEL MURO", "PROGRESIVA", "LONGITUD JUNTO A AREA VERDE (m)", "ESPACIAMIENTO (m)", "N° LLORADEROS"],
               [6, 10, 16, 14, 20, 18, 14, 14], alto=45)
    for i, t in enumerate(LLt):
        r = 8 + i
        for j, (v, fmt) in enumerate([(i + 1, "0"), (t["eje"], None), ("%s - %s" % (t["de"], t["a"]), None), (t["lado"].upper(), None),
                                       ("%.2f - %.2f" % (t["prog_desde"], t["prog_hasta"]), None), (t["L_m"], "0.00"), ("=%s" % Pk["lls"], "0.00"),
                                       ("=ROUNDUP(F%d/G%d,0)" % (r, r), "0")], 1):
            cel(wll, r, j, v, fmt, bold=(j == 8), bg=VER if j == 8 else None)
    rl1 = 8 + len(LLt)
    cel(wll, rl1, 1, "TOTAL DE LLORADEROS", bold=True); wll.merge_cells(start_row=rl1, start_column=1, end_row=rl1, end_column=7)
    cel(wll, rl1, 8, "=SUM(H8:H%d)" % max(8, rl1 - 1), "0", bold=True, bg=PAR_)
    nota(wll, rl1 + 2, 1, 8, "Lloradero de tuberia PVC 3\" a media altura del muro, con caja filtrante de 0.30 x 0.30 de grava envuelta en geotextil; solo en el "
         "muro que da a area verde (el que da a vereda o piso no lleva).")

    # MONTANTES (DrenIE: ubicacion, falsa columna, dado y salida enterrada; la altura del montante y sus accesorios los mete el proyectista)
    cab_hoja(wm, "MONTANTES PVC Ø4\": FALSA COLUMNA, DADO Y SALIDA ENTERRADA DEL DADO A LA CUNETA", 11)
    encabezado(wm, 7, ["MONTANTE", "TECHO", "LADO", "N° MONTANTES", "FALSA COLUMNA", "DADO", "CODO 90° EN DADO", "TUBERIA ENTERRADA (m)",
                       "PENDIENTE (%)", "DESCARGA A", "LLEGADA"], [11, 9, 9, 11, 11, 9, 11, 13, 11, 14, 16], alto=40)
    mts = (TT or {}).get("montantes", [])
    for i, m in enumerate(mts):
        r = 8 + i
        for j, (v, fmt) in enumerate([(m["montante"], None), (m["techo"], None), (m["lado"], None), (1, "0"), ("=D%d" % r, "0"), ("=D%d" % r, "0"),
                                       ("=D%d*%d" % (r, 1 if m.get("tubo_enterrado_m") else 0), "0"), (m.get("tubo_enterrado_m") or 0, "0.00"),
                                       (m.get("S_tubo_pct"), "0.00"), (m["cuneta"], None), ((m.get("llega") or "NO LLEGA").lower(), None)], 1):
            cel(wm, r, j, v, fmt)
    rm1 = 8 + len(mts)
    cel(wm, rm1, 1, "TOTAL", bold=True); wm.merge_cells(start_row=rm1, start_column=1, end_row=rm1, end_column=3)
    for j, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G"), (8, "H")):
        cel(wm, rm1, j, "=SUM(%s8:%s%d)" % (col, col, max(8, rm1 - 1)), "0.00" if col == "H" else "0", bold=True, bg=PAR_)
    nota(wm, rm1 + 2, 1, 11, "En el dado, un codo de 90° de 4\" y de ahi la tuberia PVC 4\" enterrada hasta la cuneta, llegando sobre el nivel del agua "
         "de la cuneta + 0.05 m (se mide en la hoja METRADO  EN PISO). La altura del montante, el tubo que sube a la canaleta, los codos del desvio, "
         "colgadores y abrazaderas los metra el proyectista con los cortes de arquitectura.")

    # CANALETAS
    cab_hoja(wca, "METRADO DE CANALETAS DE PLANCHA GALVANIZADA", 10)
    encabezado(wca, 7, ["TECHO", "LADO", "N° CANALETAS", "LONGITUD (m)", "TOTAL (m)", "SOPORTES", "TAPAS LATERALES", "BOQUILLAS", "AREA (m2)", "Q (l/s)"],
               [9, 9, 11, 11, 11, 11, 11, 11, 11, 10], alto=40)
    lds = (TT or {}).get("lados", [])
    for i, l in enumerate(lds):
        r = 8 + i
        for j, (v, fmt) in enumerate([(l["techo"], None), (l["lado"], None), (1, "0"), (l["L_canaleta_m"], "0.00"), ("=C%d*D%d" % (r, r), "0.00"),
                                       ("=C%d*IF(D%d>0,ROUNDUP(D%d/%s,0)+1,0)" % (r, r, r, Pk["css"]), "0"), ("=C%d*2" % r, "0"), ("=C%d*%d" % (r, l["boquillas"]), "0"),
                                       (l["A_m2"], "0.00"), (l["Q_ls"], "0.00")], 1):
            cel(wca, r, j, v, fmt)
    rc1 = 8 + len(lds)
    cel(wca, rc1, 1, "TOTAL", bold=True); wca.merge_cells(start_row=rc1, start_column=1, end_row=rc1, end_column=4)
    for j, col in ((5, "E"), (6, "F"), (7, "G"), (8, "H")):
        cel(wca, rc1, j, "=SUM(%s8:%s%d)" % (col, col, max(8, rc1 - 1)), "0.00" if col == "E" else "0", bold=True, bg=PAR_)

    # SUMIDEROS
    cab_hoja(wsu, "METRADO EN PISO (SUMIDEROS Y TUBERIA PVC Ø4\")", 8)
    encabezado(wsu, 7, ["SUMIDERO", "AREA (m2)", "CUNETA", "TUBERIA PVC-CP Ø4\" (m)", "SUMIDERO (und)", "CODO 90° (und)", "TEE (und)", "PENDIENTE (%)"],
               [11, 11, 11, 16, 13, 13, 10, 12], alto=40)
    sms = (SM or {}).get("sumideros", [])
    for i, s in enumerate(sms):
        r = 8 + i
        for j, (v, fmt) in enumerate([(s["sumidero"], None), (s["A_m2"], "0.0"), (s["cuneta"], None), (s["L_tubo_m"], "0.00"), (1, "0"),
                                       (s["codos90"], "0"), (0, "0"), (s.get("S_tubo_pct", 1.0), "0.00")], 1):
            cel(wsu, r, j, v, fmt)
    ment = [m for m in (TT or {}).get("montantes", []) if m.get("tubo_enterrado_m")]       # salidas de montante enterradas
    for i, m in enumerate(ment):
        r = 8 + len(sms) + i
        for j, (v, fmt) in enumerate([(m["montante"] + " (dado)", None), (None, "0.0"), (m["cuneta"], None), (m["tubo_enterrado_m"], "0.00"), (0, "0"),
                                       (1, "0"), (0, "0"), (m.get("S_tubo_pct", 1.0), "0.00")], 1):
            cel(wsu, r, j, v, fmt)
    rs1 = 8 + len(sms) + len(ment)
    cel(wsu, rs1, 1, "TOTAL", bold=True); wsu.merge_cells(start_row=rs1, start_column=1, end_row=rs1, end_column=3)
    for j, col in ((4, "D"), (5, "E"), (6, "F"), (7, "G")):
        cel(wsu, rs1, j, "=SUM(%s8:%s%d)" % (col, col, max(8, rs1 - 1)), "0.00" if col == "D" else "0", bold=True, bg=PAR_)

    # PLANILLA GENERAL
    cel(wp, 2, 1, "PLANILLA GENERAL DE METRADOS - DRENAJE PLUVIAL", bold=True, size=14, borde=False); wp.merge_cells("A2:L2")
    for i, (k, ref) in enumerate((("ENTIDAD", "=+RESUMEN!B4"), ("PROYECTO", "=+RESUMEN!B5"), ("UBICACIÓN", "=+RESUMEN!B6"), ("FECHA", "=+RESUMEN!B7"))):
        cel(wp, 3 + i, 1, k, bold=True, al="left"); cel(wp, 3 + i, 2, ref, al="left", wrap=True); wp.merge_cells(start_row=3 + i, start_column=2, end_row=3 + i, end_column=12)
    wp.row_dimensions[4].height = 33
    encabezado(wp, 8, ["ITEM", "DESCRIPCION", "Nº DE VECES", "LARGO", "ANCHO", "ALTURA", "FACT.", "AREA", "VOLUMEN", "PARCIAL", "TOTAL", "UND"],
               [16, 60, 10, 10, 10, 10, 8, 10, 10, 12, 13, 7])
    row = [9]; resumen = []

    def titulo(item, desc):
        r = row[0]; cel(wp, r, 1, item, bold=True, bg=SEC, al="left"); cel(wp, r, 2, desc, bold=True, bg=SEC, al="left")
        for j in range(3, 13): cel(wp, r, j, None, bg=SEC)
        resumen.append((item, desc, None, None)); row[0] += 1

    def partida(item, desc, und, filas):
        r0 = row[0]; cel(wp, r0, 1, item, bold=True, bg=PAR_, al="left"); cel(wp, r0, 2, desc, bold=True, bg=PAR_, al="left", wrap=True)
        for j in range(3, 11): cel(wp, r0, j, None, bg=PAR_)
        cel(wp, r0, 12, und, bold=True, bg=PAR_)
        row[0] += 1; ri = row[0]
        for f in filas:
            r = row[0]; d, C_, D_, E_, F_, G_, form = f
            cel(wp, r, 1, None); cel(wp, r, 2, d, al="left")
            for j, v in ((3, C_), (4, D_), (5, E_), (6, F_), (7, G_)):
                if isinstance(v, str) and v.startswith("="): v = v.format(r=r, **Pk)
                cel(wp, r, j, v, "0.00")
            cel(wp, r, 8, None); cel(wp, r, 9, None)
            cel(wp, r, 10, form.format(r=r, **Pk), "0.00"); cel(wp, r, 11, None); cel(wp, r, 12, None)
            row[0] += 1
        cel(wp, r0, 11, "=SUM(J%d:J%d)" % (ri, max(ri, row[0] - 1)), "0.00", bold=True, bg=PAR_)
        resumen.append((item, desc, und, "=+'PLANILLA GENERAL DE METRADOS'!K%d" % r0)); return r0

    p = lambda s: pref + s
    bx = lambda x: "=%.2f+2*{em}+2*{sa}" % x["b"]
    bt = lambda x: "=%.2f+2*{em}" % x["b"]
    nfc = len(mts)
    titulo(pref, "SISTEMA DE DRENAJE PLUVIAL")
    titulo(p(".01"), "TRABAJOS PRELIMINARES")
    partida(p(".01.01"), "LIMPIEZA MANUAL DE TERRENO", "m²", [(x["eje"], 1, x["L"], bx(x), None, None, "=C{r}*D{r}*E{r}") for x in q])
    partida(p(".01.02"), "TRAZO, NIVELES Y REPLANTEO PRELIMINAR", "m²", [(x["eje"], 1, x["L"], bx(x), None, None, "=C{r}*D{r}*E{r}") for x in q])
    titulo(p(".02"), "MOVIMIENTO DE TIERRAS")
    r_exc = partida(p(".02.01"), "EXCAVACIÓN MANUAL DE ZANJAS P/CUNETAS", "m³",
                    [(x["eje"] + " (h prom. %.3f m)" % x["h_prom"], 1, x["L"], bx(x), x["h_prom"], None, "=C{r}*D{r}*E{r}*(F{r}+{ef}+{sol})") for x in q])
    r_exd = partida(p(".02.02"), "EXCAVACIÓN MANUAL DE ZANJAS P/DADOS DE CONCRETO", "m³",
                    [("Dados de montante (%d und)" % nfc, nfc, None, None, None, None, "=C{r}*{dado}")] if nfc else [])
    partida(p(".02.03"), "REFINE, NIVELACIÓN Y COMPACTACIÓN EN TERRENO NORMAL", "m²", [(x["eje"], 1, x["L"], bt(x), None, None, "=C{r}*D{r}*E{r}") for x in q])
    r_rel = partida(p(".02.04"), "RELLENO COMPACTADO C/MATERIAL PROPIO SELECCIONADO", "m³",
                    [(x["eje"] + " (a los lados de la cuneta)", 1, x["L"], "=2*{sa}", x["h_prom"], None, "=C{r}*D{r}*E{r}*(F{r}+{ef}+{sol})") for x in q])
    partida(p(".02.05"), "ELIMINACIÓN DE MATERIAL EXCEDENTE D<1Km", "m³",
            [("(Excavacion + dados - relleno) x esponjamiento", None, None, None, None, "={esp}", "=(K%d+K%d-K%d)*G{r}" % (r_exc, r_exd, r_rel))])
    dsc = V.get("descargas", [])
    partida(p(".02.06"), "PERFILADO MANUAL DE TERRENO EN DESCARGAS", "m²",
            [("Descarga %s (C.F. %.3f)" % (d["D"], d["cf"]), 1, d.get("sale_a_m") or 0.5, "=%.2f" % (max(x["be"] for x in q) if q else 0.6), None, None, "=C{r}*D{r}*E{r}")
             for d in dsc])
    titulo(p(".03"), "OBRAS DE CONCRETO SIMPLE")
    partida(p(".03.01"), "CONCRETO f'c=100 kg/cm2 PARA SOLADO E=4''", "m²", [(x["eje"], 1, x["L"], bt(x), None, None, "=C{r}*D{r}*E{r}") for x in q])
    partida(p(".03.02"), "CONCRETO f'c=140 Kg/cm2 PARA DADOS DE CONCRETO", "m³",
            [("Dados de montante (%d und)" % nfc, nfc, None, None, None, None, "=C{r}*{dado}")] if nfc else [])
    titulo(p(".04"), "OBRAS DE CONCRETO ARMADO")
    titulo(p(".04.02"), "CUNETAS")
    partida(p(".04.02.01"), "CONCRETO f'c=175 kg/cm2 PARA CUNETAS", "m³",
            [(x["eje"] + " (hoja CONCRETO EN CUNETAS)", None, None, None, None, None, "=+'CONCRETO EN CUNETAS'!%s" % ref_conc[x["eje"]]) for x in q])
    partida(p(".04.02.02"), "ENCOFRADO Y DESENCOFRADO DE CUNETAS", "m²",
            [(x["eje"] + " (hoja ENCOFRADO Y DESENCOFRADO)", None, None, None, None, None, "=+'ENCOFRADO Y DESENCOFRADO'!%s" % ref_enc[x["eje"]]) for x in q])
    partida(p(".04.02.03"), "ACERO f'y=4200 kg/cm2 PARA CUNETA", "kg",
            [(x["eje"] + " (hoja METRADO ACERO)", None, None, None, None, None, "=+'METRADO ACERO'!W%d" % fila_acero[x["eje"]]) for x in q])
    partida(p(".04.02.04"), "CURADO DE CONCRETO EN CUNETA", "m²",
            [(x["eje"] + " (hoja METRADO DE CURADO)", None, None, None, None, None, "=+'METRADO DE CURADO'!M%d" % (ru0 + i)) for i, x in enumerate(q)])
    titulo(p(".04.03"), "VARIOS")
    partida(p(".04.03.01"), "REJILLA METALICA P/CUNETA C/PL 1''x 3/16'' Y ANGULO 1'' x 1'' x 3/16'', S/DETALLE", "m",
            [(x["eje"] + " (hoja METRADO DE REJILLAS)", None, None, None, None, None, "=+'METRADO DE REJILLAS'!G%d" % (rg0 + i)) for i, x in enumerate(q) if x["L_rejilla"] > 0])
    partida(p(".04.03.02"), "TAPA DE INSPECCION 0.60m x 0.60m", "und",
            [("Tapas de inspeccion (hoja PLANILLA TAPAS INSPECCION)", None, None, None, None, None, "=+'PLANILLA TAPAS INSPECCION'!J8")])
    partida(p(".04.03.03"), "JUNTA DE DILATACION E=1\" CON SELLADOR ASFALTO", "m",
            [(x["eje"] + " (hoja METRADO JUNTA DE DILATACION)", None, None, None, None, None, "=+'METRADO JUNTA DE DILATACION'!J%d" % (rj0 + i)) for i, x in enumerate(q)])
    partida(p(".04.03.04"), "LLORADERO DE TUBERIA PVC Ø3\" C/FILTRO DE GRAVA Y GEOTEXTIL @1.50 m (CUNETA COLINDANTE CON AREA VERDE)", "und",
            [("Lloraderos (hoja METRADO LLORADEROS)", None, None, None, None, None, "=+'METRADO LLORADEROS'!H%d" % rl1)])
    partida(p(".04.03.05"), "JUNTA E=1\" CON SELLADOR ASFALTICO ENTRE CUNETA Y PISO ADYACENTE", "m",
            [(x["eje"] + " (hoja METRADO JUNTA DE DILATACION)", None, None, None, None, None, "=+'METRADO JUNTA DE DILATACION'!J%d" % (rp0 + i)) for i, x in enumerate(ok_p)])
    if mts or sms or ment:
        titulo(p(".05"), "RED DE COLECCIÓN")
    if mts:
        titulo(p(".05.01"), "RED DE RECOLECCION TECHO")
        r_fc = partida(p(".05.01.02"), "FALSA COLUMNA P/MONTANTE DE 0.20m x 0.15m x 1.30m", "und",
                       [("Una por montante (hoja METRADO MONTANTES)", None, None, None, None, None, "=+'METRADO MONTANTES'!E%d" % rm1)])
        partida(p(".05.01.02.01"), "CONCRETO F´C=175 KG/CM2 EN FALSA COLUMNA", "m³", [("Falsas columnas x 0.20 x 0.15 x 1.30", "=K%d" % r_fc, None, None, None, None, "=C{r}*{fcc}")])
        partida(p(".05.01.02.02"), "ENCOFRADO Y DESENCOFRADO EN FALSA COLUMNA - MONTANTE PLUVIAL", "m²", [("Falsas columnas x 0.65 m2", "=K%d" % r_fc, None, None, None, None, "=C{r}*{fce}")])
        partida(p(".05.01.02.03"), "ACERO FY=4200 KG/CM2 EN FALSA COLUMNA (MONTANTE PLUVIAL)", "kg", [("Falsas columnas x 5.6112 kg", "=K%d" % r_fc, None, None, None, None, "=C{r}*{fca}")])
        partida(p(".05.01.03"), "CANALETA DE PLANCHA DE ACERO GALVANIZADO e=0.9mm, SECCIÓN 0.25m x 0.20m", "m",
                [("Canaletas (hoja METRADO CANALETAS)", None, None, None, None, None, "=+'METRADO CANALETAS'!E%d" % rc1)])
        partida(p(".05.01.04"), "SOPORTE DE PLATINA F°G° 1\" x 1/8\" P/CANALETA", "und", [("Soportes (hoja METRADO CANALETAS)", None, None, None, None, None, "=+'METRADO CANALETAS'!F%d" % rc1)])
        partida(p(".05.01.05"), "TAPA LATERAL DE PLANCHA GALVANIZADA P/CANALETA", "und", [("Tapas laterales (hoja METRADO CANALETAS)", None, None, None, None, None, "=+'METRADO CANALETAS'!G%d" % rc1)])
        partida(p(".05.01.06"), "BOQUILLA DE PLANCHA GALVANIZADA P/CONEXIÓN DE CANALETA A MONTANTE Ø4\"", "und",
                [("Boquillas (hoja METRADO CANALETAS)", None, None, None, None, None, "=+'METRADO CANALETAS'!H%d" % rc1)])
    if sms or ment:
        titulo(p(".05.02"), "RED DE RECOLECCION PISO")
        partida(p(".05.02.01"), "TUBERIA DESAGÜE PVC-CP Ø4\"", "m", [("Sumideros y salidas de montante enterradas a la cuneta (hoja METRADO  EN PISO (SUMIDERO))", None, None, None, None, None, "=+'METRADO  EN PISO (SUMIDERO)'!D%d" % rs1)])
    if sms:
        partida(p(".05.02.02"), "SUMIDERO METALICO C/PLATINA 1\" X 3/16\"", "und", [("Sumideros", None, None, None, None, None, "=+'METRADO  EN PISO (SUMIDERO)'!E%d" % rs1)])
    if sms or ment:
        partida(p(".05.02.03"), "CODO DESAGÜE PVC Ø4\" x 90°", "und", [("Codos de salida de sumideros y de los dados de montante", None, None, None, None, None,
                                                                         "=+'METRADO  EN PISO (SUMIDERO)'!F%d" % rs1)])
    if sms:
        partida(p(".05.02.04"), "TEE SANITARIA DESAGÜE PVC Ø4\"", "und", [("Tees", None, None, None, None, None, "=+'METRADO  EN PISO (SUMIDERO)'!G%d" % rs1)])

    # RESUMEN
    cel(wr, 2, 1, "RESUMEN METRADOS DE DRENAJE PLUVIAL", bold=True, size=14, borde=False); wr.merge_cells("A2:D2")
    nombre = proyecto.get("nombre", "") + ((" - " + proyecto["codigo"]) if proyecto.get("codigo") and proyecto["codigo"] not in proyecto.get("nombre", "") else "")
    for i, (k, v) in enumerate((("ENTIDAD", proyecto.get("entidad", "")), ("PROYECTO", nombre), ("UBICACIÓN", proyecto.get("ubicacion", "")), ("FECHA", proyecto.get("fecha", "")))):
        cel(wr, 4 + i, 1, k, bold=True); cel(wr, 4 + i, 2, v, al="left", wrap=True); wr.merge_cells(start_row=4 + i, start_column=2, end_row=4 + i, end_column=4)
    wr.row_dimensions[5].height = 48
    encabezado(wr, 8, ["Item", "Descripción", "Unid.", "Total."], [18, 78, 9, 15], alto=20)
    for i, (it, d, u, f) in enumerate(resumen):
        r = 9 + i; t = f is None
        cel(wr, r, 1, it, "@", bold=t, bg=SEC if t else None, al="left"); cel(wr, r, 2, d, bold=t, bg=SEC if t else None, al="left", wrap=True)
        cel(wr, r, 3, u, bg=SEC if t else None); cel(wr, r, 4, f, "#,##0.00", bg=SEC if t else None)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    wb.properties.creator = ""; wb.properties.lastModifiedBy = ""
    wb.save(ruta)
    return ruta, q, [{"codigo": it, "desc": d, "und": u or ""} for it, d, u, _ in resumen]


def recalcular(ruta, destino):
    """Copia recalculada del libro con LibreOffice (destino = carpeta); devuelve su ruta o None si no hay LibreOffice."""
    import glob, os, shutil, subprocess
    exe = shutil.which("soffice") or shutil.which("libreoffice") or next(iter(
        glob.glob(r"C:\Program Files*\LibreOffice\program\soffice.exe")), None)
    if not exe: return None
    os.makedirs(destino, exist_ok=True)
    try:
        subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", destino, ruta], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    except Exception:
        return None
    r = os.path.join(destino, os.path.basename(ruta))
    return r if os.path.exists(r) else None


def totales(recalculado):
    """Total de cada partida (hoja RESUMEN) del libro recalculado; {} si no hay copia recalculada."""
    import warnings
    from openpyxl import load_workbook
    if not recalculado: return {}
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        ws = load_workbook(recalculado, data_only=True)["RESUMEN"]
    return {str(r[0]): r[3] for r in ws.iter_rows(min_row=9, values_only=True) if r[0] and isinstance(r[3], (int, float))}
