"""Memoria de calculo hidrologico e hidraulico con formulas (MEMORIA_CALCULO_DRENAJE_PLUVIAL.xlsx).

Mismo libro de la I.E. N.° 0590 Barranquita (C.I. 468901, herramientas/barranquita/memoria_calculo_b.py), para cualquier
proyecto y con las hojas de los componentes de DrenIE:
RESUMEN, DATOS (parametros editables), SERIE, FRECUENCIAS (Normal, Log Normal, Log Pearson III, Gumbel), KS (Kolmogorov-Smirnov),
IDF (Dick y Peschke y modelo de Bell), SUPERFICIES (C ponderado), CAUDALES (metodo racional y Manning por tramo, con el ancho
de cada sector), CANALETAS Y MONTANTES (Manning y Wyly-Eaton), SUMIDEROS (tubo 4" y cota de llegada) y DESCARGAS.
Si el proyecto no trae la serie de una estacion (solo I_mmh), las hojas de lluvia se omiten y DATOS lleva la intensidad dada.
Los valores coinciden con los del diseno (hidrologia.py y verificacion.py); el tirante normal va como valor (Manning resuelto por
iteracion) y el Excel recalcula el caudal de Manning con el para comprobarlo.
"""
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

from . import normas


def _F(clave):
    """Fuente de un dato tomada de la biblioteca de normas (cita, tipo y estado)."""
    d = normas.biblioteca()["_datos"].get(clave) or normas.cuadro(clave)
    if d["clase"] == "criterio": return normas.cita(clave)
    return "%s (%s%s)" % (normas.cita(clave), normas.CLASES[d["clase"]].lower(), "" if d["estado"] == "verificado" else ", por verificar")

_b = Side(style="thin", color="808080"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
F = lambda **k: Font(name="Arial", size=k.pop("size", 9), **k)
AZ, AM, VE = "DDEBF7", "FFF2CC", "E2EFDA"
MESES = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SET", "OCT", "NOV", "DIC"]
NOM_SUP = {"techo": "Techos (coberturas)", "vereda": "Veredas", "patio": "Patios", "pavimento": "Pavimentos y areas libres de concreto",
           "estacionamiento": "Patio de maniobra y estacionamiento", "juegos": "Area de juegos", "cuneta": "Cunetas", "verde": "Areas verdes y terreno"}
DN = {"normal": "Normal", "lognormal": "Log Normal", "lp3": "Log Pearson III", "gumbel": "Gumbel"}


def cel(ws, r, c, v, fmt=None, bold=False, fill=None, borde=True, wrap=False, al=None):
    x = ws.cell(r, c, v); x.font = F(bold=bold)
    if fmt: x.number_format = fmt
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    if borde: x.border = BOR
    if wrap or al: x.alignment = Alignment(wrap_text=wrap, horizontal=al, vertical="center")
    return x


def tabla_cab(ws, r, textos, anchos=None):
    for j, t in enumerate(textos, 1):
        cel(ws, r, j, t, bold=True, fill=AZ, wrap=True, al="center")
        if anchos and j <= len(anchos): ws.column_dimensions[CL(j)].width = anchos[j - 1]
    ws.row_dimensions[r].height = 30


def yn(Q, S, n, b):
    lo, hi = 1e-6, 5.0
    for _ in range(100):
        y = (lo + hi) / 2; A = b * y; R = A / (b + 2 * y)
        if A * R ** (2 / 3) * math.sqrt(S) / n > Q: hi = y
        else: lo = y
    return (lo + hi) / 2


def libro(ruta, proyecto, P, Rh, LL, V, secciones, TT, SM, ES=None, TV=None, PF=None):
    hv = V["hidrologia"]
    cab_txt = ["PROYECTO: " + proyecto.get("nombre", "") + ((" - " + proyecto["codigo"]) if proyecto.get("codigo") else ""),
               "  -  ".join(x for x in (("ENTIDAD: " + proyecto["entidad"]) if proyecto.get("entidad") else "",
                                        ("UBICACION: " + proyecto["ubicacion"]) if proyecto.get("ubicacion") else "",
                                        ("FECHA: " + proyecto["fecha"]) if proyecto.get("fecha") else "") if x)]

    def cabecera(ws, titulo):
        ws.cell(1, 1, titulo).font = F(size=11, bold=True)
        for k, t in enumerate(cab_txt): ws.cell(2 + k, 1, t).font = F(size=8)

    wb = Workbook(); wr = wb.active; wr.title = "RESUMEN"
    wd = wb.create_sheet("DATOS")
    # ------------------------------------------------------------------ DATOS (parametros con nombre)
    cabecera(wd, "DATOS GENERALES Y PARAMETROS DE DISEÑO")
    tabla_cab(wd, 4, ["PARAMETRO", "VALOR", "", "FUENTE / CRITERIO"], [72, 13, 2, 82])
    D = {}

    def dato(clave, texto, valor, fmt, fuente):
        r = 5 + len(D); D[clave] = "DATOS!$B$%d" % r
        cel(wd, r, 1, texto); cel(wd, r, 2, valor, fmt, bold=True, fill=AM); cel(wd, r, 4, fuente)
    if Rh:
        e = Rh["estacion"] or {}
        dato("factor", "Factor aplicado a la serie de maximas", Rh["factor"], "0.0000",
             "1.0 = serie medida en la estacion, sin transposicion" if abs(Rh["factor"] - 1) < 1e-9 else "Transposicion por precipitaciones medias anuales (MTC)")
        dato("T", "Periodo de retorno de diseño T (años)", Rh["T"], "0", _F("T"))
        dato("tc", "Tiempo de concentracion de cunetas tc (min)", Rh["tc_min"], "0.0", _F("tc_min"))
    dato("n_cap", "Coeficiente de Manning para capacidad (concreto con rejilla)", P["n_cap"], "0.000", _F("n_cap"))
    dato("n_vel", "Coeficiente de Manning para velocidad (concreto liso)", P["n_vel"], "0.000", _F("n_vel"))
    dato("BL", "Borde libre minimo (m)", P["BL"], "0.00", _F("BL"))
    dato("Vmax", "Velocidad maxima en concreto (m/s)", P["Vmax"], "0.00", _F("Vmax"))
    dato("FS", "Factor de seguridad de los caudales FS", hv["FS"], "0.00", "Q = C I A FS; " + _F("FS"))
    for k, c in P["C"].items():
        dato("C_" + k, "Coeficiente de escorrentia - " + NOM_SUP.get(k, k).lower(), c, "0.00", _F("C_tabla1a") + "; adoptado: " + _F("C_adoptados"))
    if Rh:
        dato("P24", "Precipitacion maxima en 24 h de diseño (mm)", 0, "0.00", "Hoja FRECUENCIAS (%s)" % DN[Rh["distribucion"]])
        dato("Ptc", "Precipitacion en tc - Dick y Peschke P = P24 (tc/1440)^0.25 (mm)", "=%s*(%s/1440)^0.25" % (D["P24"], D["tc"]), "0.00", "Dick y Peschke: " + _F("dick_peschke"))
        dato("I", "Intensidad de diseño I = P x 60 / tc (mm/h)", "=%s*60/%s" % (D["Ptc"], D["tc"]), "0.00", "Para el metodo racional")
    else:
        dato("I", "Intensidad de diseño (mm/h)", hv["I"], "0.00", LL.get("fuente") or "Dato del proyecto")
    dato("n_can", "Canaleta: coeficiente de Manning (plancha galvanizada)", 0.013, "0.000", "Memorias de los CAR")
    dato("S_can", "Canaleta: pendiente (m/m)", 0.01, "0.000", _F("S_canaleta"))
    dato("b_can", "Canaleta: ancho (m)", 0.25, "0.00", "Partida Delphin 0.25 x 0.20")
    dato("h_can", "Canaleta: alto (m)", 0.20, "0.00", "Partida Delphin 0.25 x 0.20")
    dato("D4", "Montante: diametro (pulgadas)", 4, "0", "Siempre 4\" (regla del usuario)")
    dato("rs", "Montante: relacion de llenado (Wyly-Eaton)", 0.25, "0.000", "1/4; hasta 7/24 cuando el lado lleva 4 montantes")
    dato("rs_max", "Montante: relacion de llenado maxima", 7 / 24, "0.000", "7/24 (Wyly-Eaton)")
    dato("n_tubo", "Tuberia PVC 4\": coeficiente de Manning", 0.010, "0.000", "PVC")
    dato("D_tubo", "Tuberia PVC 4\": diametro interior (m)", 0.1016, "0.0000", "PVC 4\"")
    # ------------------------------------------------------------------ lluvia (si hay serie)
    if Rh:
        A = Rh["analisis"]; EL = Rh["distribucion"]
        ws_ = wb.create_sheet("SERIE"); wf = wb.create_sheet("FRECUENCIAS"); wk = wb.create_sheet("KS"); wi = wb.create_sheet("IDF")
        e = Rh["estacion"] or {}
        cabecera(ws_, "PRECIPITACION MAXIMA EN 24 HORAS - ESTACION %s%s" % (e.get("nombre", "").upper(), (" (%s, CODIGO %s)" % (e.get("entidad", ""), e["codigo"])) if e.get("codigo") else ""))
        cel(ws_, 5, 1, Rh["fuente_serie"], borde=False)
        tabla_cab(ws_, 7, ["AÑO"] + MESES + ["P MAX ANUAL (mm)", "P MAX DE DISEÑO (mm)", "LOG10 P"], [7] + [6] * 12 + [12, 13, 9])
        sm = [(a, m) for a, m in sorted(Rh["serie_mensual"].items()) if any(v is not None for v in m)]
        anios = {a for a, _ in Rh["maximos"]}; sm = [(a, m) for a, m in sm if a in anios]
        r0 = 8
        for i, (a, m) in enumerate(sm):
            r = r0 + i; cel(ws_, r, 1, a)
            for j, v in enumerate(m, 2): cel(ws_, r, j, v if v is not None else "S/D", "0.0")
            cel(ws_, r, 14, "=MAX(B%d:M%d)" % (r, r), "0.00", bold=True)
            cel(ws_, r, 15, "=N%d*%s" % (r, D["factor"]), "0.00", bold=True, fill=VE)
            cel(ws_, r, 16, "=LOG10(O%d)" % r, "0.0000")
        r1 = r0 + len(sm) - 1
        rngO, rngP = "SERIE!$O$%d:$O$%d" % (r0, r1), "SERIE!$P$%d:$P$%d" % (r0, r1)
        # FRECUENCIAS
        cabecera(wf, "ANALISIS DE FRECUENCIAS DE LA SERIE DE MAXIMAS ANUALES (mm)")
        est = [("Numero de datos n", "=COUNT(%s)" % rngO, "0"), ("Media aritmetica", "=AVERAGE(%s)" % rngO, "0.00"),
               ("Desviacion estandar", "=STDEV(%s)" % rngO, "0.00"), ("Coeficiente de asimetria", "=SKEW(%s)" % rngO, "0.000"),
               ("Media de log P", "=AVERAGE(%s)" % rngP, "0.0000"), ("Desviacion estandar de log P", "=STDEV(%s)" % rngP, "0.0000"),
               ("Coeficiente de asimetria de log P", "=SKEW(%s)" % rngP, "0.0000"),
               ("Gumbel: yn (Aparicio, n datos)", A["gumbel_yn"], "0.0000"), ("Gumbel: sn (Aparicio, n datos)", A["gumbel_sn"], "0.0000")]
        for i, (t, f_, fmt) in enumerate(est):
            cel(wf, 5 + i, 1, t); cel(wf, 5 + i, 2, f_, fmt, bold=True)
        tabla_cab(wf, 16, ["T (años)", "z", "NORMAL", "LOG NORMAL", "KT LP III", "LOG PEARSON III", "y GUMBEL", "GUMBEL", "DISEÑO (%s)" % DN[EL]],
                  [38, 9, 11, 12, 10, 13, 11, 11, 16])
        filaT = {}
        for i, T in enumerate(sorted(A["tabla"])):
            r = 17 + i; filaT[T] = r
            cel(wf, r, 1, T, "0", bold=True)
            w = "SQRT(LN(A%d^2))" % r
            cel(wf, r, 2, "=%s-(2.515517+0.802853*%s+0.010328*%s^2)/(1+1.432788*%s+0.189269*%s^2+0.001308*%s^3)" % ((w,) * 6), "0.000")
            cel(wf, r, 3, "=$B$6+B%d*$B$7" % r, "0.00")
            cel(wf, r, 4, "=10^($B$9+B%d*$B$10)" % r, "0.00")
            cel(wf, r, 5, "=B{r}+(B{r}^2-1)*($B$11/6)+1/3*(B{r}^3-6*B{r})*($B$11/6)^2-(B{r}^2-1)*($B$11/6)^3+B{r}*($B$11/6)^4+1/3*($B$11/6)^5".format(r=r), "0.000")
            cel(wf, r, 6, "=10^($B$9+E%d*$B$10)" % r, "0.00")
            cel(wf, r, 7, "=-LN(-LN(1-1/A%d))" % r, "0.000")
            cel(wf, r, 8, "=$B$6-$B$12*$B$7/$B$13+G%d*$B$7/$B$13" % r, "0.00")
            cel(wf, r, 9, "=%s%d" % ({"normal": "C", "lognormal": "D", "lp3": "F", "gumbel": "H"}[EL], r), "0.00", bold=True, fill=VE)
        wd.cell(int(D["P24"].split("$")[-1]), 2).value = "=FRECUENCIAS!I%d" % filaT[Rh["T"]]
        cel(wf, 17 + len(filaT) + 1, 1, "Se adopta la distribucion %s (menor estadistico de Kolmogorov-Smirnov, hoja KS). z por la aproximacion de Abramowitz y "
                                        "Stegun; KT de Log Pearson III por la expresion de Kite (k = Cs/6); Gumbel con yn y sn de Aparicio." % DN[EL], borde=False)
        # KS
        cabecera(wk, "PRUEBA DE BONDAD DE AJUSTE DE KOLMOGOROV-SMIRNOV (alfa = 0.05)")
        tabla_cab(wk, 8, ["m", "P ordenada de mayor a menor (mm)", "P(X>=x) observada m/(n+1)", "P teorica Normal", "P teorica Log Normal",
                          "P teorica Log Pearson III", "P teorica Gumbel", "|D| Normal", "|D| Log Normal", "|D| LP III", "|D| Gumbel"],
                  [6, 14, 14, 11, 12, 13, 11, 10, 11, 10, 10])
        n = A["n"]; L = A["ln"]
        par = [("n", "=COUNT(%s)" % rngO, "0"), ("media ln P", L["media"], "0.0000"), ("desv. ln P", L["desv"], "0.0000"),
               ("Cs ln P", L["cs"], "0.0000"), ("beta = (2/Cs)^2", "=(2/N12)^2", "0.0000"), ("alfa = s/raiz(beta)", "=N11/SQRT(N13)", "0.00000"),
               ("delta = media - alfa beta", "=N10-N14*N13", "0.0000"), ("Gumbel: media yn", A["gumbel_yn"], "0.0000"),
               ("Gumbel: desviacion sn", A["gumbel_sn"], "0.0000"), ("D critico (alfa 0.05)", "=SQRT(-LN(0.05/2)/(2*N9))-1/(6*N9)", "0.0000")]
        for i, (t, v, fmt) in enumerate(par):
            cel(wk, 9 + i, 13, t, bold=True); cel(wk, 9 + i, 14, v, fmt, fill=AM)
        wk.column_dimensions["M"].width = 32; wk.column_dimensions["N"].width = 12
        for i in range(n):
            r = 9 + i
            cel(wk, r, 1, i + 1, "0"); cel(wk, r, 2, "=LARGE(%s,A%d)" % (rngO, r), "0.00"); cel(wk, r, 3, "=A%d/($N$9+1)" % r, "0.0000")
            cel(wk, r, 4, "=1-NORMSDIST((B%d-FRECUENCIAS!$B$6)/FRECUENCIAS!$B$7)" % r, "0.0000")
            cel(wk, r, 5, "=1-NORMSDIST((LOG10(B%d)-FRECUENCIAS!$B$9)/FRECUENCIAS!$B$10)" % r, "0.0000")
            cel(wk, r, 6, "=CHIDIST(2*(LN(B%d)-$N$15)/$N$14,2*$N$13)" % r, "0.0000")
            cel(wk, r, 7, "=1-EXP(-EXP(-((B%d-FRECUENCIAS!$B$6)/FRECUENCIAS!$B$7*$N$17+$N$16)))" % r, "0.0000")
            for j, col in enumerate("DEFG"): cel(wk, r, 8 + j, "=ABS(%s%d-C%d)" % (col, r, r), "0.0000")
        rm = 9 + n
        cel(wk, rm, 7, "D max", bold=True)
        for j in range(4): cel(wk, rm, 8 + j, "=MAX(%s9:%s%d)" % (CL(8 + j), CL(8 + j), rm - 1), "0.0000", bold=True, fill=VE)
        cel(wk, rm + 1, 7, "D critico", bold=True)
        for j in range(4): cel(wk, rm + 1, 8 + j, "=$N$18", "0.0000")
        cel(wk, rm + 2, 7, "Ajuste", bold=True)
        for j in range(4): cel(wk, rm + 2, 8 + j, '=IF(%s%d<%s%d,"ACEPTA","RECHAZA")' % (CL(8 + j), rm, CL(8 + j), rm + 1))
        # IDF
        cabecera(wi, "LLUVIAS E INTENSIDADES MAXIMAS: DICK Y PESCHKE (DISEÑO) Y MODELO DE BELL (CURVAS IDF)")
        Ts = sorted(filaT); DUR = sorted(next(iter(Rh["idf"].values())).keys())
        cel(wi, 5, 1, "Lluvia (mm) para la duracion D: P = P24 (D / 1440)^0.25 (Dick y Peschke)", borde=False, bold=True)
        tabla_cab(wi, 6, ["D (min)"] + ["T = %d" % T for T in Ts], [10] + [10] * len(Ts))
        for i, Dm in enumerate(DUR):
            r = 7 + i; cel(wi, r, 1, Dm, "0", bold=True)
            for j, T in enumerate(Ts, 2): cel(wi, r, j, "=FRECUENCIAS!$I$%d*($A%d/1440)^0.25" % (filaT[T], r), "0.00")
        r2 = 7 + len(DUR) + 1
        cel(wi, r2, 1, "Intensidad I (mm/h) = P x 60 / D", borde=False, bold=True)
        tabla_cab(wi, r2 + 1, ["D (min)"] + ["T = %d" % T for T in Ts])
        for i, Dm in enumerate(DUR):
            r = r2 + 2 + i; cel(wi, r, 1, Dm, "0", bold=True)
            for j, T in enumerate(Ts):
                cel(wi, r, 2 + j, "=%s%d*60/$A%d" % (CL(2 + j), 7 + i, r), "0.00",
                    fill=VE if (abs(Dm - Rh["tc_min"]) < 1e-6 and T == Rh["T"]) else None)
        r3 = r2 + 2 + len(DUR) + 1
        cel(wi, r3, 1, "Modelo de Bell: P(T,t) = (0.21 ln T + 0.52)(0.54 t^0.25 - 0.50) P60(10);  P60(10) = 0.4602 P24(10)^0.876 =", borde=False, bold=True)
        cel(wi, r3, 9, "=0.4602*FRECUENCIAS!$I$%d^0.876" % filaT[10] if 10 in filaT else Rh["P60_10"], "0.00", bold=True, fill=AM)
        DB = sorted(next(iter(Rh["bell_i"].values())).keys())
        tabla_cab(wi, r3 + 1, ["T (años)"] + ["I t = %d min (mm/h)" % t for t in DB])
        for i, T in enumerate(sorted(Rh["bell_i"])):
            r = r3 + 2 + i; cel(wi, r, 1, T, "0", bold=True)
            for j, t in enumerate(DB):
                cel(wi, r, 2 + j, "=60*(0.21*LN($A%d)+0.52)*(0.54*%d^0.25-0.5)*$I$%d/%d" % (r, t, r3, t), "0.0")
        r4 = r3 + 2 + len(Rh["bell_i"]) + 1
        cel(wi, r4, 1, "Regresion multiple de las intensidades de Bell: I = K T^m / t^n", borde=False, bold=True)
        for k, (t, v) in enumerate((("K", Rh["K"]), ("m", Rh["m"]), ("n", Rh["n"]))):
            cel(wi, r4 + 1 + k, 1, t, bold=True); cel(wi, r4 + 1 + k, 2, round(v, 4), "0.0000", fill=AM)
    # ------------------------------------------------------------------ SUPERFICIES
    wsu = wb.create_sheet("SUPERFICIES")
    cabecera(wsu, "COEFICIENTE DE ESCORRENTIA PONDERADO (AREA DENTRO DEL CERCO)")
    tabla_cab(wsu, 6, ["TIPO DE SUPERFICIE", "AREA (m2)", "C", "C x A", "% DEL AREA"], [52, 12, 8, 12, 12])
    sup = [(k, v) for k, v in hv["superficies"].items() if v["A"] > 0]
    rt = 7 + len(sup)
    for i, (k, v) in enumerate(sup):
        r = 7 + i
        cel(wsu, r, 1, NOM_SUP.get(k, k)); cel(wsu, r, 2, v["A"], "#,##0"); cel(wsu, r, 3, "=" + D.get("C_" + k, str(v["C"])), "0.00")
        cel(wsu, r, 4, "=B%d*C%d" % (r, r), "#,##0.0"); cel(wsu, r, 5, "=B%d/$B$%d" % (r, rt), "0.0%")
    cel(wsu, rt, 1, "TOTAL / C PONDERADO", bold=True); cel(wsu, rt, 2, "=SUM(B7:B%d)" % (rt - 1), "#,##0", bold=True)
    cel(wsu, rt, 3, "=D%d/B%d" % (rt, rt), "0.000", bold=True, fill=VE); cel(wsu, rt, 4, "=SUM(D7:D%d)" % (rt - 1), "#,##0.0", bold=True)
    cel(wsu, rt, 5, "=SUM(E7:E%d)" % (rt - 1), "0.0%", bold=True)
    cel(wsu, rt + 2, 1, "Celdas de 1 m2 del area dentro del cerco, clasificadas por tipo de superficie y asignadas a la cuneta que las recibe.", borde=False)
    # ------------------------------------------------------------------ CAUDALES
    wc = wb.create_sheet("CAUDALES")
    cabecera(wc, "CAUDALES DE DISEÑO (METODO RACIONAL) Y VERIFICACION HIDRAULICA (MANNING) POR TRAMO")
    hdr = ["EJE", "SECTOR", "DE", "A", "AREA ACUM. (m2)", "C POND.", "I (mm/h)", "Q (m3/s)", "Q (l/s)", "S (%)", "b LIBRE (m)",
           "Q EN EL PUNTO CRITICO (l/s)", "h LIBRE EN ESE PUNTO (m)", "y (m) n cap.", "A = b y (m2)", "P = b + 2y (m)", "R (m)", "Q MANNING (l/s)",
           "BORDE LIBRE (m)", "CAPACIDAD", "y (m) n vel.", "V (m/s)", "VELOCIDAD"]
    tabla_cab(wc, 6, hdr, [8, 7, 7, 7, 10, 7, 9, 9, 8, 7, 8, 10, 9, 9, 9, 8, 8, 10, 9, 10, 9, 8, 10])
    Hh = V["hidraulica"]
    for i, f in enumerate(Hh):
        r = 7 + i; Qm = f["Q_ls"] / 1000.0; S_ = f["S_pct"] / 100.0; b = f["b"]
        Qc = f.get("Q_critico_ls", f["Q_ls"]) / 1000.0
        y = yn(Qc, S_, P["n_cap"], b) if Qc > 0 else 0.0; yv = yn(Qm, S_, P["n_vel"], b) if Qm > 0 else 0.0
        vals = [(f["eje"], None), (f["sector"], None), (f["de"], None), (f["a"], None), (f["A_m2"], "#,##0"), (f["C"], "0.000"),
                ("=" + D["I"], "0.00"), ("=F{r}*G{r}*E{r}*{fs}/3600000".format(r=r, fs=D["FS"]), "0.0000"), ("=H{r}*1000".format(r=r), "0.0"),
                (f["S_pct"], "0.00"), (b, "0.00"), (round(1000 * Qc, 2), "0.0"), (f["h"], "0.00"), (round(y, 4), "0.000"),
                ("=K{r}*N{r}".format(r=r), "0.0000"), ("=K{r}+2*N{r}".format(r=r), "0.000"), ("=IF(P{r}>0,O{r}/P{r},0)".format(r=r), "0.0000"),
                ("=1000*O{r}*Q{r}^(2/3)*SQRT(J{r}/100)/{n}".format(r=r, n=D["n_cap"]), "0.0"), ("=M{r}-N{r}".format(r=r), "0.000"),
                ('=IF(S{r}>={bl}-0.001,"CUMPLE","NO CUMPLE")'.format(r=r, bl=D["BL"]), None), (round(yv, 4), "0.000"),
                ("=IF(U{r}>0,H{r}/(K{r}*U{r}),0)".format(r=r), "0.00"), ('=IF(V{r}<={v},"CUMPLE","NO CUMPLE")'.format(r=r, v=D["Vmax"]), None)]
        for j, (v, fmt) in enumerate(vals, 1): cel(wc, r, j, v, fmt)
    rr = 7 + len(Hh)
    cel(wc, rr + 1, 1, "Punto critico: el de menor borde libre a lo largo del tramo (perfil cada 0.50 m); el caudal crece del que entra al tramo al de su "
                       "final en proporcion al largo y la altura libre es la del perfil (bajo la losa tapa donde esta tapado). y: tirante normal por "
                       "Manning (la columna Q MANNING lo comprueba). b: ancho libre del sector.", borde=False)
    # ------------------------------------------------------------------ CANALETAS Y MONTANTES
    if TT and TT["lados"]:
        wt = wb.create_sheet("CANALETAS Y MONTANTES")
        cabecera(wt, "CANALETAS (MANNING) Y MONTANTES PVC 4\" (WYLY-EATON) POR LADO DE TECHO")
        hdr = ["TECHO", "LADO", "AREA (m2)", "C", "I (mm/h)", "Q (l/s)", "MONTANTES 4\"", "Q / MONTANTE (l/s)", "LLENADO rs", "Q MAX 4\" (l/s)",
               "MONTANTE", "TIRANTE CANALETA (m)", "Q CAPACIDAD CANALETA (l/s)", "CANALETA"]
        tabla_cab(wt, 6, hdr, [8, 9, 10, 7, 9, 9, 11, 12, 10, 11, 11, 12, 13, 11])
        for i, l in enumerate(TT["lados"]):
            r = 7 + i
            rs = "IF(H{r}<=1.754*{rs}^(5/3)*{d}^(8/3),{rs},{rsm})".format(r=r, rs=D["rs"], rsm=D["rs_max"], d=D["D4"])
            hu = "({h}-{bl})".format(h=D["h_can"], bl=0.05)
            vals = [(l["techo"], None), (l["lado"], None), (l["A_m2"], "#,##0.00"), ("=" + D["C_techo"] if "C_techo" in D else 0.88, "0.00"),
                    ("=" + D["I"], "0.00"), ("=D{r}*E{r}*C{r}*{fs}/3600".format(r=r, fs=D["FS"]), "0.00"), (l["montantes"], "0"),
                    ("=F{r}/G{r}".format(r=r), "0.00"), ("=" + rs, "0.000"),
                    ("=1.754*I{r}^(5/3)*{d}^(8/3)".format(r=r, d=D["D4"]), "0.00"), ('=IF(H{r}<=J{r}+0.001,"CUMPLE","NO CUMPLE")'.format(r=r), None),
                    (l["tirante_m"], "0.000"),
                    ("=1000*{b}*{hu}*({b}*{hu}/({b}+2*{hu}))^(2/3)*SQRT({S})/{n}".format(b=D["b_can"], hu=hu, S=D["S_can"], n=D["n_can"]), "0.0"),
                    ('=IF(H{r}<=M{r},"CUMPLE","NO CUMPLE")'.format(r=r), None)]
            for j, (v, fmt) in enumerate(vals, 1): cel(wt, r, j, v, fmt)
        rt2 = 7 + len(TT["lados"])
        cel(wt, rt2 + 1, 1, "Wyly-Eaton: Q (l/s) = 1.754 rs^(5/3) D^(8/3), D en pulgadas; montantes siempre de 4\": a lo mas 4 por lado, con llenado 1/4 o "
                            "hasta 7/24. Canaleta: Manning con borde libre de 0.05 m; cada montante recibe Q / n del lado.", borde=False)
    # ------------------------------------------------------------------ SUMIDEROS
    if SM and SM["sumideros"]:
        wsm = wb.create_sheet("SUMIDEROS")
        cabecera(wsm, "SUMIDEROS DE PISO: CAUDAL, TUBERIA PVC 4\" Y COTA DE LLEGADA A LA CUNETA")
        hdr = ["SUMIDERO", "AREA (m2)", "C", "Q (l/s)", "Q TUBO 4\" (l/s)", "CAPACIDAD", "CUNETA", "L TUBO (m)", "S (%)", "NIVEL SALIDA",
               "LLEGADA", "COTA DE FONDO CUNETA", "TIRANTE CUNETA (m)", "NIVEL MINIMO (+0.05)", "LLEGA"]
        tabla_cab(wsm, 6, hdr, [10, 10, 7, 9, 10, 11, 9, 9, 7, 11, 10, 12, 10, 12, 16])
        for i, s in enumerate(SM["sumideros"]):
            r = 7 + i
            ck = "C_estacionamiento" if "C_estacionamiento" in D else next(k for k in D if k.startswith("C_"))
            vals = [(s["sumidero"], None), (s["A_m2"], "#,##0.0"), ("=" + D[ck], "0.00"), ("=B{r}*C{r}*{I}*{fs}/3600".format(r=r, I=D["I"], fs=D["FS"]), "0.00"),
                    ("=1000*PI()*{d}^2/4*({d}/4)^(2/3)*SQRT(I{r}/100)/{n}".format(d=D["D_tubo"], r=r, n=D["n_tubo"]), "0.00"),
                    ('=IF(D{r}<=E{r},"CUMPLE","NO CUMPLE")'.format(r=r), None), (s["cuneta"], None), (s["L_tubo_m"], "0.00"),
                    (s["S_tubo_pct"], "0.00"), (s["z_salida"], "0.000"), ("=J{r}-I{r}/100*H{r}".format(r=r), "0.000"), (s["cf_cuneta"], "0.000"),
                    (s["tirante_cuneta"], "0.000"), ("=L{r}+M{r}+0.05".format(r=r), "0.000"),
                    ('=IF(K{r}>=N{r}-0.0005,"SOBRE EL AGUA","BAJO EL AGUA")'.format(r=r), None)]
            for j, (v, fmt) in enumerate(vals, 1): cel(wsm, r, j, v, fmt)
        cel(wsm, 7 + len(SM["sumideros"]) + 1, 1, "Solo en areas libres sin cobertura. Tubo PVC 4\" al 1 % (minimo 0.5 %) hasta la cuneta mas cercana; su fondo debe llegar "
                                                  "sobre el nivel del agua de la cuneta (cota de fondo + tirante) + 0.05 m, para que no quede ahogado (si no, el programa "
                                                  "profundiza ese tramo de cuneta).", borde=False)
    # ------------------------------------------------------------------ SALIDAS DE MONTANTES
    ment = [m for m in (TT or {}).get("montantes", []) if m.get("tubo_enterrado_m")]
    if ment:
        wsa = wb.create_sheet("SALIDAS DE MONTANTES")
        cabecera(wsa, "SALIDA DE CADA MONTANTE: CODO 90 EN EL DADO Y TUBERIA PVC 4\" ENTERRADA HASTA LA CUNETA")
        tabla_cab(wsa, 6, ["MONTANTE", "TECHO", "CUNETA", "L TUBO (m)", "S (%)", "NIVEL SALIDA", "LLEGADA", "COTA DE FONDO CUNETA", "TIRANTE CUNETA (m)",
                           "NIVEL MINIMO (+0.05)", "LLEGA"], [10, 8, 9, 9, 7, 11, 10, 12, 10, 12, 16])
        for i, m in enumerate(ment):
            r = 7 + i
            vals = [(m["montante"], None), (m["techo"], None), (m["cuneta"], None), (m["tubo_enterrado_m"], "0.00"), (m["S_tubo_pct"], "0.00"),
                    (m["z_salida"], "0.000"), ("=F{r}-E{r}/100*D{r}".format(r=r), "0.000"), (m["cf_cuneta"], "0.000"), (m["tirante_cuneta"], "0.000"),
                    ("=H{r}+I{r}+0.05".format(r=r), "0.000"), ('=IF(G{r}>=J{r}-0.0005,"SOBRE EL AGUA","BAJO EL AGUA")'.format(r=r), None)]
            for j, (v, fmt) in enumerate(vals, 1): cel(wsa, r, j, v, fmt)
        cel(wsa, 7 + len(ment) + 1, 1, "En el dado de cada montante, un codo de 90 de 4\" y la tuberia PVC 4\" enterrada hasta la cara interior de la cuneta (hasta "
                                       "5 m), al 1 % (minimo 0.5 %); su fondo llega sobre el nivel del agua de la cuneta + 0.05 m. La cuneta puede ir a lo largo del "
                                       "alero o transversal: basta que el tubo de cada montante llegue.", borde=False)
    # ------------------------------------------------------------------ DESCARGAS
    wdc = wb.create_sheet("DESCARGAS")
    cabecera(wdc, "DESCARGAS POR EL CERCO PERIMETRICO")
    tabla_cab(wdc, 6, ["DESCARGA", "COTA DE FONDO (m)", "TERRENO EXTERIOR (m)", "SALE A (m del cerco)", "PERFILADO (m)", "ESTADO"], [11, 13, 13, 13, 12, 24])
    for i, d in enumerate(V["descargas"]):
        r = 7 + i
        for j, (v, fmt) in enumerate([(d["D"], None), (d["cf"], "0.000"), (d["terreno"], "0.00"), (d.get("sale_a_m"), "0.00"),
                                      (d.get("perfilado_max") or 0, "0.00"), (d["estado"], None)], 1):
            cel(wdc, r, j, v, fmt)
    # ------------------------------------------------------------------ RESUMEN
    cabecera(wr, "MEMORIA DE CALCULO - DRENAJE PLUVIAL (RED DE CUNETAS)")
    res = []
    if Rh:
        e = Rh["estacion"] or {}
        res += [("Estacion de referencia", "%s%s, %d - %d (%d años)" % (e.get("nombre", ""), (" (%s, codigo %s)" % (e.get("entidad", ""), e["codigo"])) if e.get("codigo") else "",
                                                                        Rh["maximos"][0][0], Rh["maximos"][-1][0], Rh["analisis"]["n"]), None),
                ("Distribucion adoptada", "%s (menor D de Kolmogorov-Smirnov)" % DN[Rh["distribucion"]], None),
                ("Periodo de retorno (años)", "=" + D["T"], "0"), ("Precipitacion maxima en 24 h de diseño (mm)", "=" + D["P24"], "0.00"),
                ("Tiempo de concentracion (min)", "=" + D["tc"], "0.0")]
    res += [("Intensidad de diseño (mm/h)", "=" + D["I"], "0.00"), ("Coeficiente de escorrentia ponderado", "=SUPERFICIES!C%d" % rt, "0.000"),
            ("Area de aporte dentro del cerco (m2)", "=SUPERFICIES!B%d" % rt, "#,##0"),
            ("Caudal maximo en un tramo (l/s)", "=MAX(CAUDALES!I7:I%d)" % (rr - 1), "0.0"),
            ("Menor borde libre (m)", "=MIN(CAUDALES!S7:S%d)" % (rr - 1), "0.000"),
            ("Velocidad maxima (m/s)", "=MAX(CAUDALES!V7:V%d)" % (rr - 1), "0.00"),
            ("Tramos que no cumplen capacidad", '=COUNTIF(CAUDALES!T7:T%d,"NO CUMPLE")' % (rr - 1), "0"),
            ("Tramos que no cumplen velocidad", '=COUNTIF(CAUDALES!W7:W%d,"NO CUMPLE")' % (rr - 1), "0"),
            ("Secciones adoptadas (ancho libre por sector)", ", ".join("%s: %.2f m" % (k, v) for k, v in sorted({s["sector"]: s["b_adoptado"] for s in secciones}.items())), None)]
    if TT and TT["lados"]:
        res.append(("Montantes PVC 4\" que no cumplen", '=COUNTIF(\'CANALETAS Y MONTANTES\'!K7:K%d,"NO CUMPLE")' % (6 + len(TT["lados"])), "0"))
    tabla_cab(wr, 6, ["CONCEPTO", "VALOR", ""], [52, 70, 2])
    for i, (t, v, fmt) in enumerate(res):
        cel(wr, 7 + i, 1, t, bold=True); cel(wr, 7 + i, 2, v, fmt, fill=VE if fmt else None)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0
    wb.properties.creator = ""; wb.properties.lastModifiedBy = ""
    est = [e for e in (ES or {}).get("estructuras", []) if e.get("codigo")]
    if est:
        # caida vertical con poza (Rand / USBR) con formulas: la profundidad de poza y la caida son datos del diseno
        we = wb.create_sheet("ESTRUCTURAS"); cabecera(we, "ESTRUCTURAS EN LOS CAMBIOS DE NIVEL (CAIDA VERTICAL CON POZA)")
        tabla_cab(we, 4, ["CODIGO", "EJE", "PROG. (m)", "DESNIVEL (m)", "Q (l/s)", "b (m)", "CAIDA DE FONDO (m)", "PROF. POZA (m)", "hd (m)", "q (m2/s)",
                          "D", "Ld (m)", "yp (m)", "y1 (m)", "y2 (m)", "Lj (m)", "L POZA (m)", "TIPO"],
                  [10, 9, 9, 10, 9, 8, 11, 10, 9, 10, 11, 9, 9, 9, 9, 9, 10, 44])
        for i, e in enumerate(est):
            r = 5 + i
            vals = [e["codigo"], e["eje"], e["progresiva"], e["desnivel"], e["Q_ls"], e["b"], max(e["caida_fondo"], 0.0), e["prof_poza"],
                    "=G{r}+H{r}", "=E{r}/1000/F{r}", "=J{r}^2/(9.81*I{r}^3)", "=4.3*I{r}*K{r}^0.27", "=1*I{r}*K{r}^0.22",
                    "=0.54*I{r}*K{r}^0.425", "=1.66*I{r}*K{r}^0.27", "=6.9*(O{r}-N{r})", "=MAX(0.6,CEILING(L{r}+P{r},0.05))", e["descripcion"]]
            fm = [None, None, "0.00", "0.00", "0.00", "0.00", "0.000", "0.00", "0.000", "0.0000", "0.000000", "0.000", "0.000", "0.000", "0.000", "0.000", "0.00", None]
            for j, (v, f) in enumerate(zip(vals, fm)):
                cel(we, r, j + 1, v.format(r=r) if isinstance(v, str) and v.startswith("=") else v, f, wrap=(j == 17))
        r = 6 + len(est)
        cel(we, r, 1, "Fuente: " + normas.cita_larga("caida_vertical") + ". Rangos: " + normas.dato("rangos_desnivel")["valor"] + " (criterio del proyectista).", borde=False)
    if TV:
        wt = wb.create_sheet("TAPA VEHICULAR"); cabecera(wt, "TAPA VEHICULAR DE CUNETAS EN ESTACIONAMIENTOS (RUEDA HL-93)")
        tabla_cab(wt, 4, ["DATO", "VALOR", "", "FORMULA / FUENTE"], [52, 13, 2, 70])
        filas = [("Rueda del camion HL-93 P (kN)", 72.5, "0.0", normas.cita("carga_vehicular")), ("Incremento por carga dinamica IM", 1.33, "0.00", "33 %"),
                 ("Factor de carga (Resistencia I)", 1.75, "0.00", "AASHTO LRFD / Manual de Puentes MTC"), ("Pu (kN)", "=B5*B6*B7", "0.0", "P x IM x 1.75"),
                 ("Luz de la losa L = b + e muro (m)", TV["luz"], "0.000", "ancho libre maximo mas espesor de muro"),
                 ("Espesor de losa h (m)", TV["h"], "0.00", "f'c = 280 kg/cm2"), ("Ancho de reparto E = 0.51 + 2h (m)", "=0.51+2*B10", "0.000", "huella de la llanta 0.51 m, reparto a 45 grados"),
                 ("Mu = Pu L / 4 / E (kN.m/m)", "=B8*B9/4/B11", "0.00", "rueda al centro de la luz"),
                 ("Acero adoptado", TV["acero"], None, "As = %.2f cm2/m (>= 0.0018 h, E.060)" % TV["As_cm2_m"]),
                 ("phi Vc punzonamiento (kN)", TV["phiVc_kN"], "0.0", "0.75 x 0.33 raiz(f'c) b0 d >= Pu")]
        for i, (t, v, f, s) in enumerate(filas):
            cel(wt, 5 + i, 1, t); cel(wt, 5 + i, 2, v, f, bold=True, fill=AM); cel(wt, 5 + i, 4, s)
    if PF and PF.get("ejes"):
        # perfil de flujo por el paso estandar: el tirante es el resultado iterado (valor); el Excel recalcula la energia y comprueba
        # el balance H(i) = H(i+1) + hf entre estaciones (residuo ~ 0 salvo en controles, caidas y resaltos, que se indican)
        wp = wb.create_sheet("PERFIL DE FLUJO"); cabecera(wp, "PERFIL DE FLUJO POR EL METODO DEL PASO ESTANDAR (FLUJO GRADUALMENTE VARIADO)")
        cel(wp, 4, 1, "n de Manning", bold=True); cel(wp, 4, 2, PF["n"], "0.000", fill=AM); cel(wp, 4, 3, "g (m/s2)", bold=True); cel(wp, 4, 4, 9.81, "0.00", fill=AM)
        cel(wp, 4, 6, "Fuente: " + _F("paso_estandar"), borde=False)
        tabla_cab(wp, 6, ["EJE", "PROG. s (m)", "COTA DE FONDO z (m)", "Q (l/s)", "b (m)", "TIRANTE y (m)", "A (m2)", "R (m)", "V (m/s)", "V2/2g (m)",
                          "ENERGIA H (m)", "Sf", "dx (m)", "hf (m)", "BALANCE H(i)-H(i+1)-hf (m)", "yc (m)", "Fr", "REGIMEN", "ALTURA LIBRE h (m)",
                          "BORDE LIBRE (m)", "OBSERVACION"],
                  [9, 9, 11, 8, 7, 9, 8, 8, 8, 9, 10, 10, 7, 8, 13, 8, 7, 12, 10, 10, 30])
        r = 7
        for e in PF["ejes"]:
            pts = e["estaciones"]; sal = {round(x["s"], 2) for x in e["saltos"]}
            for k, p in enumerate(pts):
                ult = k == len(pts) - 1
                sig = None if ult else pts[k + 1]
                obs = ""
                if ult: obs = "control: " + e["control"]
                elif p["z"] - sig["z"] > 0.05 and (p["z"] - sig["z"]) / max(sig["s"] - p["s"], 1e-6) > 0.10: obs = "caida de fondo (tirante critico en el borde)"
                elif round(sig["s"], 2) in sal: obs = "resalto hidraulico"
                elif abs(p["y"] - p["yc"]) < 0.002: obs = "tirante critico (control)"
                vals = [e["eje"], round(p["s"], 3), round(p["z"], 4), round(1000 * p["Q"], 3), p["b"], round(p["y"], 4),
                        "=E{r}*F{r}", "=G{r}/(E{r}+2*F{r})", "=D{r}/1000/G{r}", "=I{r}^2/(2*$D$4)", "=C{r}+F{r}+J{r}", "=($B$4*I{r}/H{r}^(2/3))^2",
                        None if ult else "=B{n}-B{r}", None if ult else "=M{r}*(L{r}+L{n})/2", None if ult else "=K{r}-K{n}-N{r}",
                        "=(D{r}^2/1000000/($D$4*E{r}^2))^(1/3)", "=I{r}/SQRT($D$4*F{r})", p["regimen"], round(p["h"], 3), "=S{r}-F{r}", obs]
                fm = [None, "0.00", "0.000", "0.00", "0.00", "0.000", "0.0000", "0.0000", "0.00", "0.0000", "0.000", "0.000000", "0.00", "0.0000",
                      "0.0000", "0.000", "0.00", None, "0.000", "0.000", None]
                for j, (v, f) in enumerate(zip(vals, fm)):
                    cel(wp, r, j + 1, v.format(r=r, n=r + 1) if isinstance(v, str) and v.startswith("=") else v, f,
                        fill=(AM if j == 5 else VE if j in (14, 19) else None))
                r += 1
        wp.freeze_panes = "C7"
        r += 1
        cel(wp, r, 1, "Tirante: resultado del paso estandar (iterado). Subcritico: de aguas abajo hacia aguas arriba desde el control; supercritico: de aguas "
                      "arriba hacia aguas abajo desde el tirante critico; en cada estacion el de mayor fuerza especifica. El balance es ~0 entre "
                      "estaciones del mismo tramo; en controles, caidas y resaltos no aplica.", borde=False)
        wr2 = wb.create_sheet("PERFIL DE FLUJO RESUMEN"); cabecera(wr2, "PERFIL DE FLUJO - RESUMEN POR EJE")
        tabla_cab(wr2, 4, ["EJE", "LONGITUD (m)", "CONTROL AGUAS ABAJO", "TIRANTE MAX. (m)", "BORDE LIBRE MIN. (m)", "PROG. (m)", "V MAX. (m/s)",
                           "SUPERCRITICO (m)", "RESALTOS", "ESTADO"], [9, 11, 52, 11, 12, 9, 10, 12, 10, 18])
        for i, e in enumerate(PF["ejes"]):
            for j, (v, f) in enumerate(zip([e["eje"], e["L"], e["control"], e["y_max"], e["BL_min"], e["s_BL_min"], e["V_max"], e["supercritico_m"],
                                            len(e["saltos"]), e["estado"]], [None, "0.00", None, "0.000", "0.000", "0.00", "0.00", "0.0", "0", None])):
                cel(wr2, 5 + i, j + 1, v, f, wrap=(j == 2), fill=(VE if e["estado"] == "CUMPLE" else AM) if j == 9 else None)
    normas.hoja_excel(wb)
    wb.save(ruta)
    return ruta
