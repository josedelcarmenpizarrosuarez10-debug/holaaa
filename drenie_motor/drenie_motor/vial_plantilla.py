"""Plantilla Excel de datos del drenaje vial (DATOS_DRENAJE_VIAL.xlsx): la llena el proyectista y el motor la lee.

Hojas: VIA (datos generales), CUNETAS (un tramo por fila, con hasta tres aportes), ALCANTARILLAS (paso de quebradas y alivios
propuestos) y BADENES. La lluvia viene de la pestaña LLUVIA de DrenIE (o P24 por T en la hoja VIA).
"""
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from .memoria_calculo import AM, F, cel, tabla_cab
from .vial import COBERTURAS, SUELOS

VIA = [("nombre", "Nombre del proyecto", ""), ("ubicacion", "Ubicación", ""), ("codigo", "Código (CUI)", ""),
       ("velocidad_kmh", "Velocidad de diseño (km/h)", 40), ("imda", "IMDA (veh/día)", 350),
       ("precipitacion_anual_mm", "Precipitación media anual (mm)", 2000), ("selva_alta", "Selva alta con cauces en V (SI / NO)", "NO"),
       ("talud_corte_z", "Talud de corte, z (H por 1 V)", 0.5), ("tc_min", "Tiempo de concentración mínimo (min)", 10),
       ("pendiente_alivio", "Pendiente de alcantarillas de alivio (m/m)", 0.02), ("longitud_alivio_m", "Longitud de alcantarillas de alivio (m)", 8),
       ("P24_T25", "P24 para T = 25 años (mm), solo si no hay serie en LLUVIA", ""), ("P24_T50", "P24 para T = 50 años (mm)", ""),
       ("P24_T100", "P24 para T = 100 años (mm)", "")]
CUN = [("id", "Código"), ("lado", "Lado"), ("prog_ini", "Prog. inicio (m)"), ("prog_fin", "Prog. fin (m)"), ("pendiente", "Pendiente (m/m)"),
       ("revestimiento", "Revestimiento"), ("tipo", "Sección"), ("d", "Profundidad d (m)"), ("a", "Ancho a o fondo b (m)"),
       ("z2", "Talud exterior z2"), ("ap1", "Aporte 1"), ("w1", "Ancho 1 (m)"), ("c1", "C 1"), ("ap2", "Aporte 2"), ("w2", "Ancho 2 (m)"),
       ("c2", "C 2"), ("ap3", "Aporte 3"), ("w3", "Ancho 3 (m)"), ("c3", "C 3")]
ALC = [("id", "Código"), ("prog", "Progresiva (m)"), ("tipo", "Tipo"), ("area_km2", "Área (km²)"), ("L_m", "L cauce (m)"), ("H_m", "Desnivel H (m)"),
       ("cobertura", "Cobertura"), ("suelo", "Suelo"), ("C", "C (opcional)"), ("CN", "CN (opcional)"), ("cunetas", "Recibe cunetas (códigos)"),
       ("tipo_conducto", "Conducto"), ("diametro_pulg", "Ø fijo (pulg)"), ("pendiente", "Pendiente (m/m)"), ("longitud_m", "Longitud (m)"),
       ("terreno_salida", "Terreno a la salida")]
BAD = [("id", "Código"), ("prog", "Progresiva (m)"), ("area_km2", "Área (km²)"), ("L_m", "L cauce (m)"), ("H_m", "Desnivel H (m)"),
       ("cobertura", "Cobertura"), ("suelo", "Suelo"), ("cunetas", "Recibe cunetas"), ("b", "Ancho de fondo b (m)"), ("z", "Talud z"),
       ("profundidad", "Profundidad (m)"), ("pendiente_transversal", "Pendiente transversal (m/m)")]


def _lista(ws, rango, valores):
    v = DataValidation(type="list", formula1='"%s"' % ",".join(valores), allow_blank=True); ws.add_data_validation(v); v.add(rango)


def crear(ruta, d=None):
    """Plantilla vacía, o llena con los datos d (mismo formato JSON del motor)."""
    d = d or {}; pr = d.get("proyecto", {}); via = d.get("via", {})
    wb = Workbook(); wv = wb.active; wv.title = "VIA"
    wv.cell(1, 1, "DATOS DEL DRENAJE VIAL").font = F(size=11, bold=True)
    tabla_cab(wv, 3, ["Dato", "Valor", "Clave"], [52, 22, 22])
    p24 = (d.get("lluvia") or {}).get("P24_por_T", {})
    for i, (k, t, v0) in enumerate(VIA):
        v = pr.get(k, via.get(k, v0))
        if k == "selva_alta": v = "SI" if via.get("selva_alta") else "NO"
        if k.startswith("P24_T"): v = p24.get(k[5:], "")
        cel(wv, 4 + i, 1, t); cel(wv, 4 + i, 2, v, fill=AM); cel(wv, 4 + i, 3, k)
    hojas = (("CUNETAS", CUN, d.get("cunetas", [])), ("ALCANTARILLAS", ALC, d.get("alcantarillas", [])), ("BADENES", BAD, d.get("badenes", [])))
    for nombre, cols, filas in hojas:
        ws = wb.create_sheet(nombre)
        tabla_cab(ws, 1, [t for _, t in cols], [14] * len(cols))
        for j, (k, _) in enumerate(cols, 1): ws.cell(2, j, k).font = F(size=7, color="808080")
        for i, x in enumerate(filas):
            s = x.get("seccion", {}); ap = x.get("aportes", [])
            for j, (k, _) in enumerate(cols, 1):
                if k in ("tipo", "d", "a", "z2") and nombre == "CUNETAS": v = s.get(k, s.get("b") if k == "a" else None)
                elif k[:-1] in ("ap", "w", "c") and k[-1].isdigit() and nombre == "CUNETAS":
                    m = int(k[-1]) - 1; a = ap[m] if m < len(ap) else {}
                    v = a.get({"ap": "descripcion", "w": "ancho", "c": "C"}[k[:-1]])
                elif k in ("cobertura", "suelo") and x.get("coberturas"): v = x["coberturas"][0].get(k)
                elif k == "cunetas": v = ", ".join(x.get("cunetas", []))
                else: v = x.get(k)
                cel(ws, 3 + i, j, v)
        n = 2 + max(len(filas), 60)
        col = {k: chr(64 + j) for j, (k, _) in enumerate(cols, 1)}
        if nombre == "CUNETAS":
            _lista(ws, "%s3:%s%d" % (col["lado"], col["lado"], n), ["izquierdo", "derecho"])
            _lista(ws, "%s3:%s%d" % (col["revestimiento"], col["revestimiento"], n), ["concreto", "tierra", "roca"])
            _lista(ws, "%s3:%s%d" % (col["tipo"], col["tipo"], n), ["triangular", "trapezoidal", "rectangular"])
        else:
            _lista(ws, "%s3:%s%d" % (col["cobertura"], col["cobertura"], n), list(COBERTURAS))
            _lista(ws, "%s3:%s%d" % (col["suelo"], col["suelo"], n), list(SUELOS))
            if nombre == "ALCANTARILLAS":
                _lista(ws, "%s3:%s%d" % (col["tipo"], col["tipo"], n), ["cruce", "alivio"])
                _lista(ws, "%s3:%s%d" % (col["tipo_conducto"], col["tipo_conducto"], n), ["TMC", "MCA"])
        ws.freeze_panes = "B3"
    wb.save(ruta)
    return ruta


def _v(x):
    if x is None or (isinstance(x, str) and not x.strip()): return None
    if isinstance(x, str):
        t = x.strip().replace(",", ".") if x.strip().replace(",", ".").replace(".", "", 1).lstrip("-").isdigit() else x.strip()
        try: return float(t)
        except ValueError: return t
    return x


def leer(ruta):
    """Datos del motor desde la plantilla llena."""
    wb = load_workbook(ruta, data_only=True)
    wv = wb["VIA"]; cfg = {}
    for r in range(4, wv.max_row + 1):
        k = wv.cell(r, 3).value
        if k: cfg[k] = _v(wv.cell(r, 2).value)
    pr = {k: cfg.pop(k) for k in ("nombre", "ubicacion", "codigo") if cfg.get(k) is not None}
    p24 = {k[5:]: cfg.pop(k) for k in list(cfg) if k.startswith("P24_T") and cfg[k] is not None}
    for k in [k for k in cfg if k.startswith("P24_T")]: cfg.pop(k)
    cfg["selva_alta"] = str(cfg.get("selva_alta") or "").upper().startswith("S")
    via = {k: v for k, v in cfg.items() if v is not None}

    def filas(nombre):
        if nombre not in wb.sheetnames: return []
        ws = wb[nombre]; claves = [ws.cell(2, j).value for j in range(1, ws.max_column + 1)]
        out = []
        for r in range(3, ws.max_row + 1):
            x = {k: _v(ws.cell(r, j + 1).value) for j, k in enumerate(claves) if k}
            if x.get("id") is None: continue
            out.append({k: v for k, v in x.items() if v is not None})
        return out
    cun = []
    for x in filas("CUNETAS"):
        sec = {"tipo": x.pop("tipo", "triangular")}
        for k in ("d", "a", "z2"):
            if k in x: sec[k] = x.pop(k)
        if sec["tipo"] != "triangular" and "a" in sec: sec["b"] = sec.pop("a")
        x["seccion"] = sec if "d" in sec else {}
        ap = []
        for m in (1, 2, 3):
            w, c, t = x.pop("w%d" % m, None), x.pop("c%d" % m, None), x.pop("ap%d" % m, None)
            if w is not None and c is not None: ap.append({"descripcion": t or "Aporte %d" % m, "ancho": w, "C": c})
        if ap: x["aportes"] = ap
        x["id"] = str(x["id"]); cun.append(x)
    obras = {}
    for nombre in ("ALCANTARILLAS", "BADENES"):
        obras[nombre] = []
        for x in filas(nombre):
            x["id"] = str(x["id"])
            if x.get("cunetas"): x["cunetas"] = [s.strip() for s in str(x["cunetas"]).split(",") if s.strip()]
            if x.get("cobertura"): x["coberturas"] = [{"fraccion": 1.0, "cobertura": x.pop("cobertura"), "suelo": x.pop("suelo", "semipermeable")}]
            obras[nombre].append(x)
    d = {"proyecto": pr, "via": via, "cunetas": cun, "alcantarillas": obras["ALCANTARILLAS"], "badenes": obras["BADENES"]}
    if p24: d["lluvia"] = {"P24_por_T": p24}
    return d
