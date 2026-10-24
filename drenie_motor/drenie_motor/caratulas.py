"""Caratulas de los documentos y caratula con indice de planos (Word), como en los expedientes de la I.E. N.° 0590 Barranquita
(C.I. 468901) y de los CAR (CUI 2705619, 2675514, 2717013).
- CARATULAS_DRENAJE_PLUVIAL.docx: una caratula por documento (memorias, memoria de calculo, especificaciones, metrado y planos).
- CARATULA_E_INDICE_DE_PLANOS.docx: caratula de los planos e indice por grupos (codigo, contenido, escala y formato), con los
  ejes de cada lamina de perfiles y las secciones de cada lamina de secciones leidos del DXF.
Sin nombre ni CIP del proyectista.
"""
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from . import word

DOCUMENTOS = [("MEMORIA DESCRIPTIVA", "Drenaje pluvial - red de cunetas"), ("ESTUDIO HIDROLÓGICO Y DE DRENAJE PLUVIAL", "Memoria hidrológica - red de cunetas"),
              ("MEMORIA DE CÁLCULO", "Hidrología e hidráulica de cunetas, canaletas, montantes y sumideros"),
              ("ESPECIFICACIONES TÉCNICAS", "Drenaje pluvial - red de cunetas"), ("METRADO", "Drenaje pluvial - red de cunetas"),
              ("PLANOS", "Drenaje pluvial - red de cunetas")]
GRUPOS = (("PLANTA GENERAL", "PLANTA GENERAL"), ("CUADROS", "CUADROS DE EJES, DESCARGAS Y COMPONENTES"), ("PERFILES", "PERFILES LONGITUDINALES"),
          ("SECCIONES", "SECCIONES TRANSVERSALES"), ("DESCARGA", "DETALLES DE DESCARGAS POR EL CERCO"), ("DETALLES", "DETALLES CONSTRUCTIVOS"))
TILDES = ((r"\bSECCION\b", "SECCIÓN"), (r"\bSECCIONES\b", "SECCIONES"), (r"\bISOMETRICO\b", "ISOMÉTRICO"), (r"\bINSPECCION\b", "INSPECCIÓN"),
          (r"\bELEVACION\b", "ELEVACIÓN"), (r"\bMETALICA\b", "METÁLICA"), (r"\bDILATACION\b", "DILATACIÓN"))


def tildes(t):
    for a, b in TILDES: t = re.sub(a, b, t)
    return t


def _sin_encabezado(D):
    s = D.d.sections[0]
    for p in s.header.paragraphs + s.footer.paragraphs:
        for r in p.runs: r.text = ""
        pPr = p._p.pPr
        if pPr is not None:
            for b in pPr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr"): pPr.remove(b)


def caratulas(ruta, proyecto):
    D = word.Doc(proyecto, "")
    _sin_encabezado(D)
    for i, (t, sub) in enumerate(DOCUMENTOS):
        D.caratula(t, sub)
    # la ultima caratula no deja una pagina en blanco
    ult = D.d.paragraphs[-1]; ult._p.getparent().remove(ult._p)
    return D.guardar(ruta)


def detalle_laminas(doc, marcos):
    """Ejes de las laminas de perfiles y secciones de las laminas de secciones, leidos de los textos del DXF."""
    txt = []
    for e in doc.modelspace().query("TEXT"):
        p = e.dxf.insert; txt.append((p.x, p.y, e.dxf.text))
    out = {}
    for m in marcos:
        dentro = [t for x, y, t in txt if m["x0"] <= x <= m["x1"] and m["y0"] <= y <= m["y1"]]
        ejes = sorted({re.match(r"PERFIL LONGITUDINAL (\S+)", t).group(1) for t in dentro
                       if t.startswith("PERFIL LONGITUDINAL ") and " DE LA DESCARGA " not in t})
        secs = sorted({int(t.split()[1]) for t in dentro if re.fullmatch(r"SECCION \d+", t)})
        if ejes: out[m["codigo"]] = "Ejes: " + ", ".join(ejes)
        elif secs: out[m["codigo"]] = "Secciones %02d a %02d" % (secs[0], secs[-1]) if len(secs) > 1 else "Sección %02d" % secs[0]
    return out


def indice_planos(ruta, proyecto, laminas, detalle):
    D = word.Doc(proyecto, "ÍNDICE DE PLANOS")
    D.caratula("PLANOS", "Drenaje pluvial - red de cunetas")
    p = D.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ÍNDICE DE PLANOS"); r.bold = True; r.font.size = Pt(14); r.font.color.rgb = word.AZUL
    filas, grupo, ng = [], None, 0
    for l in laminas:
        tit = l["titulo"].upper()
        g = next((n for k, n in GRUPOS if (k in tit if k == "DESCARGA" else tit.startswith(k) or k in tit)), "OTROS")
        if g != grupo:
            grupo = g; ng += 1; filas.append(["%d. %s" % (ng, g), "", "", ""])
        filas.append([l["codigo"], tildes(l["titulo"]) + (("\n" + detalle[l["codigo"]]) if l["codigo"] in detalle else ""), l.get("escala", ""), "A1"])
    t = D.tabla(None, ["LÁMINA", "CONTENIDO", "ESCALA", "FORMATO"], filas, [2.0, 10.2, 1.8, 2.0], tam=8.5, numerada=False)
    for i, f in enumerate(filas, 1):
        if f[1] == "":
            cells = t.rows[i].cells
            m = cells[0].merge(cells[3]); m.paragraphs[0].runs[0].bold = True; m.paragraphs[0].runs[0].font.color.rgb = word.AZUL
            for par in m.paragraphs[1:]: par._p.getparent().remove(par._p)          # sin las lineas vacias de las celdas unidas
            for c in t.rows[i].cells: word._sombra(c, "EEF3FA")
    D.p("Total: %d láminas. Medidas en metros; escala real (1 unidad del dibujo = 1 m)." % len(laminas), tam=9)
    return D.guardar(ruta)
