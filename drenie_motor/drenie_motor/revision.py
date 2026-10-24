"""Revision general de los entregables (REVISION_GENERAL.xlsx): las validaciones de conocimiento/01, seccion 8.2, que el programa
corre solo al terminar cada diseno, mas las reglas de diseno del usuario.

Cada fila: grupo, comprobacion, estado (CUMPLE / NO CUMPLE / AVISO) y detalle. NO CUMPLE = regla del usuario incumplida
(bloqueante); AVISO = dato a confirmar por el proyectista, sin cambiar nada.
Grupos:
- Diseno: secciones constructivas (0.30 a 0.40 libre, de 5 en 5 cm, una medida por sector), muros, losas y solado de 0.10,
  pendiente, profundidad, borde libre, velocidad, descargas por gravedad, comprobaciones del ciclo de verificacion.
- Componentes: montantes siempre de 4" y a lo mas 4 por lado, canaletas, sumideros solo en areas libres sin cobertura y con
  tuberia al 1 % (o 0.5 %) llegando sobre la cota de fondo, lloraderos @1.50 m solo junto a area verde.
- Metrado: sin errores de formula; cantidades iguales a las del diseno (rejilla, tapas, lloraderos, montantes, canaletas, sumideros).
- Documentos: idioma y terminos del usuario, sin nombres de personas ni CIP, sin rastros de herramientas en textos ni metadatos,
  sin citas a otros proyectos, laminas citadas existentes, cifras de las memorias y especificaciones iguales a las del metrado.
- Planos: escala real (cotas con texto igual a su medida), textos sin encimarse, dentro del marco, sin textos verticales,
  sin amarillo ni blanco puro.
"""
import os
import re
import warnings
import zipfile

from shapely.geometry import Point, box

CUMPLE, FALLA, AVISO = "CUMPLE", "NO CUMPLE", "AVISO"
OTROS = {"2705619": "CAR Varones", "2675514": "Hogar de Refugio", "2717013": "CAR Mujeres", "468901": "I.E. N.° 0590 Barranquita",
         "2673852": "Vía SM-107"}
NOMBRES_OTROS = ["CAR Varones", "CAR VARONES", "Mujeres Violentadas", "MUJERES VIOLENTADAS", "CAR Mujeres", "CAR MUJERES", "Hogar de Refugio",
                 "HOGAR DE REFUGIO", "Los Mangos", "LOS MANGOS", "Morales", "MORALES", "Barranquita", "BARRANQUITA", "Shamboyacu"]
LAMINAS_OTRAS = re.compile(r"\b(DPLV?|DPVL|DVPL|DP|DA|DD|DJ|DL)-\d{2}[A-Z]?\b")
PERSONAS = ["Katiuska", "Solange", "Wilma", "C.I.P.", " CIP ", "CIP N", "Ing. ", "Ingeniero "]
LISTA_NEGRA = [r"\binvert\b", r"\bcheck\b", r"\bflow\b", r"\bTODO\b", r"\bXXX\b", r"\[[^\]\n]{0,40}\]", r"\b(?:completar|insertar)\s+(?:aquí|aqui|dato|datos|valor|nombre|fecha)\b", r"\bClaude\b",
               r"\bAnthropic\b", r"\bChatGPT\b", r"\bOpenAI\b", r"inteligencia artificial", r"\bIA\b", r"\bpython", r"\bopenpyxl\b",
               r"\bezdxf\b", r"\bnudo\b", r"\bnudos\b", r"\bml\b"]
HERRAMIENTAS = ("python", "openpyxl", "ezdxf", "docx", "claude", "anthropic")


class Informe:
    def __init__(self): self.filas = []

    def add(self, grupo, comp, ok, detalle="", aviso=False):
        self.filas.append([grupo, comp, CUMPLE if ok else (AVISO if aviso else FALLA), detalle])

    def resumen(self):
        n = lambda e: sum(1 for f in self.filas if f[2] == e)
        return {"comprobaciones": len(self.filas), "cumple": n(CUMPLE), "no_cumple": n(FALLA), "avisos": n(AVISO),
                "fallas": [f[1] + ": " + f[3] for f in self.filas if f[2] == FALLA]}


def _es(v, a, tol): return v is not None and a is not None and abs(float(v) - float(a)) <= tol


# ============================================================================================== diseno y componentes
def diseno(I, P, V, C, sec, hq=None):
    H = V["hidraulica"]; g = "Diseño"
    bs = sorted({f["b"] for f in H})
    I.add(g, "Ancho libre entre 0.30 y %.2f m (total hasta %.2f m)" % (P["b_max"], P["b_max"] + 2 * P["e_muro"]),
          all(0.30 - 1e-9 <= b <= P["b_max"] + 1e-9 for b in bs), "anchos usados: " + ", ".join("%.2f" % b for b in bs))
    I.add(g, "Medidas constructivas de 5 en 5 cm", all(abs(round(b / 0.05) * 0.05 - b) < 1e-6 for b in bs) and
          all(abs(round(s["b_adoptado"] / 0.05) * 0.05 - s["b_adoptado"]) < 1e-6 for s in sec))
    por = {}
    for f in H: por.setdefault(f["sector"], set()).add(f["b"])
    I.add(g, "Una sola medida por sector de descarga", all(len(v) == 1 for v in por.values()),
          "; ".join("%s: %s" % (k, "/".join("%.2f" % b for b in sorted(v))) for k, v in sorted(por.items())))
    I.add(g, "Muros, losa de fondo, losa tapa y solado de 0.10 m",
          all(abs(P[k] - 0.10) < 1e-9 for k in ("e_muro", "e_fondo", "e_tapa", "e_solado")))
    smin = min(f["S_pct"] for f in H)
    I.add(g, "Pendiente de la cota de fondo >= %.1f %%" % (100 * P["S0"]), smin >= 100 * P["S0"] - 1e-6, "mínima %.2f %%" % smin)
    ini = []
    for r in C["ramales"]:
        t0 = r["tramos"][0]; tap = any(a <= 0.01 for a, b in r.get("tapados", []))
        h_ini = P["h0"] + (P["e_tapa"] if tap else 0)
        req = (hq or {}).get("%s|%s|%s" % (r["codigo"], t0["de"], t0["a"]))           # tramo profundizado (capacidad o tubo de sumidero)
        if req: h_ini = max(h_ini, req + (P["e_tapa"] if tap else 0))
        ini.append(abs(t0["h_de"] - h_ini) < 0.006)
    I.add(g, "Profundidad inicial %.2f m (más la losa si empieza tapada) y altura libre >= %.2f m" % (P["h0"], P["hmin"]),
          all(ini) and min(min(t["h_de"], t["h_a"]) for r in C["ramales"] for t in r["tramos"]) >= P["hmin"] - 1e-6,
          "%d de %d ejes empiezan con la profundidad mínima" % (sum(ini), len(ini)))
    bl = min(f["borde_libre"] for f in H); vm = max(f["V"] for f in H)
    I.add(g, "Borde libre >= %.2f m con el caudal de diseño" % P["BL"], bl >= P["BL"] - 1e-6, "mínimo %.3f m" % bl)
    I.add(g, "Velocidad <= %.1f m/s" % P["Vmax"], vm <= P["Vmax"] + 1e-6, "máxima %.2f m/s" % vm)
    for d in V["descargas"]:
        I.add(g, "Descarga %s por gravedad" % d["D"], d["estado"] == "SALE POR GRAVEDAD",
              "cota de fondo %.3f, terreno exterior %.2f" % (d["cf"], d["terreno"]))
    for c in V["comprobaciones"]:
        I.add(g, "Comprobación %d: %s" % (c["n"], c["comprobacion"]), c["estado"].startswith("CUMPLE"), c.get("detalle", ""))


def componentes(I, P, TT, SM, LLt, mapa_tipos):
    g = "Componentes"
    if TT["montantes"]:
        I.add(g, "Montantes siempre de 4\"", all(m["D_pulg"] == 4 for m in TT["montantes"]), "%d montantes" % len(TT["montantes"]))
        mas = [l for l in TT["lados"] if l["montantes"] > 4]
        I.add(g, "A lo más 4 montantes de 4\" por lado de techo", not mas,
              "; ".join("%s %s: %d (%.0f m²)" % (l["techo"], l["lado"], l["montantes"], l["A_m2"]) for l in mas))
        I.add(g, "Al menos 2 montantes por lado (esquinas)", all(l["montantes"] >= 2 for l in TT["lados"]))
        I.add(g, "Canaletas con borde libre", all(l["canaleta"] == "CUMPLE" for l in TT["lados"]))
        av = [a for a in TT["avisos"] if "no tiene cuneta" in a]
        ent = [m for m in TT["montantes"] if m.get("tubo_enterrado_m")]
        I.add(g, "Montantes con cuneta cercana (salida directa hasta 3 m o tubería enterrada hasta 25 m)", not av,
              "%d montantes sin cuneta a menos de 25 m: revisar si falta cuneta" % len(av) if av else "%d con tubería enterrada" % len(ent), aviso=True)
        for m in ent:
            I.add(g, "Salida enterrada de %s: tubo 4\" al 1 %% (o 0.5 %%) sobre la cota de fondo" % m["montante"],
                  m["S_tubo_pct"] in (1.0, 0.5) and m["z_llegada"] >= m["cf_cuneta"] - 1e-6,
                  "%.2f m a %s, S = %.2f %%, llega a %.3f, cota de fondo %.3f" % (m["tubo_enterrado_m"], m["cuneta"], m["S_tubo_pct"], m["z_llegada"], m["cf_cuneta"]))
        ale = [a for a in TT["avisos"] if "Altura de alero" in a]
        I.add(g, "Altura de alero por edificación", not ale, ale[0] if ale else "", aviso=True)
    for s in SM["sumideros"]:
        tipo = mapa_tipos(s["x"], s["y"])
        I.add(g, "Sumidero %s en área libre sin cobertura" % s["sumidero"], tipo not in ("verde", "vereda", "techo", "cuneta"),
              "superficie: %s" % (tipo or "sin identificar"))
        I.add(g, "Sumidero %s: tubo 4\" al 1 %% (o 0.5 %%) llegando sobre la cota de fondo" % s["sumidero"],
              s["S_tubo_pct"] in (1.0, 0.5) and s["z_llegada"] >= s["cf_cuneta"],
              "S = %.2f %%, llega a %.3f, cota de fondo %.3f" % (s["S_tubo_pct"], s["z_llegada"], s["cf_cuneta"]))
    if LLt:
        I.add(g, "Lloraderos @%.2f m" % P["llor_sep"], all(abs(x["sep_m"] - P["llor_sep"]) < 1e-6 for x in LLt),
              "%d lloraderos en %d tramos de muro junto a área verde" % (sum(x["n"] for x in LLt), len(LLt)))


# ============================================================================================== metrado
def metrado(I, recalculado, tot, partidas, R_tot):
    g = "Metrado"
    if not recalculado:
        I.add(g, "Fórmulas del metrado recalculadas", False, "no hay LibreOffice para recalcular: abrir el Excel y revisar", aviso=True); return
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(recalculado, data_only=True)
    err = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.rstrip("!?0/").upper() in ("#REF", "#VALUE", "#DIV", "#NAME", "#N/A", "#NUM", "#NULL"):
                    err.append("%s!%s" % (ws.title, c.coordinate))
    I.add(g, "Sin errores de fórmula", not err, ", ".join(err[:10]))
    sin = [p["codigo"] for p in partidas if p["und"] and tot.get(p["codigo"]) is None]
    I.add(g, "Todas las partidas con total", not sin, ", ".join(sin))

    def busca(*claves, und=None):
        for p in partidas:
            d = p["desc"].upper()
            if p["und"] and all(k in d for k in claves) and (und is None or p["und"] == und): return tot.get(p["codigo"]), p["codigo"]
        return None, None
    for nombre, claves, und, valor, tol in R_tot:
        v, cod = busca(*claves, und=und)
        if cod is None and not valor: continue
        I.add(g, "%s: metrado = diseño" % nombre, _es(v, valor, max(tol, 0.001 * abs(valor))), "%s metrado %s / diseño %s" % (cod or "-", "-" if v is None else round(v, 2), round(valor, 2)))


def memoria_calculo(I, recalculado, V, Rh, TT, SM):
    """La memoria de calculo con formulas, recalculada, debe dar lo mismo que el diseno."""
    g = "Memoria de cálculo"
    if not recalculado:
        I.add(g, "Fórmulas de la memoria de cálculo recalculadas", False, "no hay LibreOffice para recalcular", aviso=True); return
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(recalculado, data_only=True)
    err = ["%s!%s" % (ws.title, c.coordinate) for ws in wb.worksheets for row in ws.iter_rows() for c in row
           if isinstance(c.value, str) and c.value.startswith("#") and c.value[:4].upper() in ("#REF", "#VAL", "#DIV", "#NAM", "#N/A", "#NUM", "#NUL")]
    I.add(g, "Sin errores de fórmula", not err, ", ".join(err[:10]))
    res = {r[0]: r[1] for r in wb["RESUMEN"].iter_rows(min_row=7, values_only=True) if r[0]}
    H = V["hidraulica"]; hv = V["hidrologia"]
    comp = [("Intensidad de diseño (mm/h)", hv["I"], 0.01), ("Caudal maximo en un tramo (l/s)", max(f["Q_ls"] for f in H), 0.15),
            ("Menor borde libre (m)", min(f["borde_libre"] for f in H), 0.002), ("Velocidad maxima (m/s)", max(f["V"] for f in H), 0.02),
            ("Tramos que no cumplen capacidad", 0, 0), ("Tramos que no cumplen velocidad", 0, 0)]
    if Rh: comp += [("Precipitacion maxima en 24 h de diseño (mm)", Rh["P24"], 0.01)]
    if TT and TT["lados"]: comp.append(("Montantes PVC 4\" que no cumplen", 0, 0))
    for nom, val, tol in comp:
        v = res.get(nom)
        ok = isinstance(v, (int, float)) and abs(v - val) <= max(tol, 0.002 * abs(val)) + 1e-9      # redondeo de C y areas en la hoja
        I.add(g, "%s: Excel = diseño" % nom, ok, "Excel %s / diseño %s" % (round(v, 3) if isinstance(v, (int, float)) else v, round(val, 3)))
    if SM and SM["sumideros"] and "SUMIDEROS" in wb.sheetnames:
        mal = [r[0] for r in wb["SUMIDEROS"].iter_rows(min_row=7, values_only=True) if r[0] and str(r[0]).startswith("S-") and r[12] != "SOBRE EL FONDO"]
        I.add(g, "Sumideros: tubo sobre la cota de fondo (Excel)", not mal, ", ".join(mal))


# ============================================================================================== documentos
def _textos_docx(ruta):
    from docx import Document
    d = Document(ruta); t = [p.text for p in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            for c in row.cells: t.append(c.text)
    for s in d.sections:
        for p in s.header.paragraphs + s.footer.paragraphs: t.append(p.text)
    cp = d.core_properties
    meta = {"autor": cp.author, "modificado por": cp.last_modified_by, "comentarios": cp.comments, "titulo": cp.title,
            "palabras clave": cp.keywords, "categoria": cp.category}
    return "\n".join(t), meta


def _textos_xlsx(ruta):
    from openpyxl import load_workbook
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = load_workbook(ruta)
    t = []
    for ws in wb.worksheets:
        t.append(ws.title)
        for row in ws.iter_rows(values_only=True):
            t += [v for v in row if isinstance(v, str) and not v.startswith("=")]
    return "\n".join(t), {"autor": wb.properties.creator, "modificado por": wb.properties.lastModifiedBy}


def _textos_dxf(doc):
    t = []
    for e in doc.modelspace().query("TEXT MTEXT"):
        t.append(e.dxf.text if e.dxftype() == "TEXT" else e.plain_text())
    for b in doc.blocks:
        if b.name.startswith("*"): continue
        for e in b.query("TEXT MTEXT ATTDEF"):
            t.append(e.dxf.text if e.dxftype() != "MTEXT" else e.plain_text())
    return "\n".join(t), {"guardado por": doc.header.get("$LASTSAVEDBY", "")}


def documentos(I, archivos, proyecto, laminas, dxf_doc):
    g = "Documentos"
    propio = (proyecto.get("nombre", "") + " " + proyecto.get("nombre_corto", "") + " " + proyecto.get("codigo", "") + " " +
              proyecto.get("ubicacion", "") + " " + proyecto.get("entidad", ""))
    codigos = {x["codigo"] for x in laminas}
    pref = {re.match(r"[A-Z]+", c).group(0) for c in codigos if re.match(r"[A-Z]+", c)}
    for ruta in archivos:
        if not ruta or not os.path.exists(ruta): continue
        nom = os.path.basename(ruta)
        if ruta.endswith(".docx"): txt, meta = _textos_docx(ruta)
        elif ruta.endswith(".xlsx"): txt, meta = _textos_xlsx(ruta)
        elif ruta.endswith(".dxf"): txt, meta = _textos_dxf(dxf_doc if dxf_doc is not None and nom == "PLANOS_CUNETAS.dxf" else __import__("ezdxf").readfile(ruta))
        else: continue
        malas = set()
        for pat in LISTA_NEGRA:
            fl = 0 if pat in (r"\bIA\b", r"\bml\b", r"\bTODO\b", r"\bXXX\b") else re.I
            malas |= {m.group(0) for m in re.finditer(pat, txt, fl)}
        malas = sorted(malas)
        I.add(g, "%s: español y términos del usuario (sin 'invert', 'nudo', 'ml', marcas de borrador ni de herramientas)" % nom, not malas, ", ".join(malas))
        per = [p.strip() for p in PERSONAS if p in txt]
        I.add(g, "%s: sin nombres de personas ni CIP" % nom, not per, ", ".join(per))
        mt = [k for k, v in meta.items() if v and any(h in str(v).lower() for h in HERRAMIENTAS + ("python-docx",))] + \
             [k for k, v in meta.items() if k in ("autor", "modificado por", "guardado por") and v]
        if nom.endswith(".dxf"):
            with open(ruta, encoding="utf-8", errors="ignore") as fh: crudo = fh.read().upper()
            if "EZDXF" in crudo: mt.append("marca de la herramienta en el archivo"); meta["marca de la herramienta en el archivo"] = "%d veces" % crudo.count("EZDXF")
        I.add(g, "%s: metadatos vacíos" % nom, not mt, ", ".join(sorted(set("%s = %s" % (k, meta[k]) for k in mt))))
        otros = sorted({n for c, n in OTROS.items() if c in txt and c not in propio} |
                       {n for n in NOMBRES_OTROS if n in txt and n.lower() not in propio.lower()})
        lam = sorted({m.group(0) for m in LAMINAS_OTRAS.finditer(txt) if m.group(0) not in codigos and m.group(1) not in pref})
        I.add(g, "%s: sin citas a otros proyectos ni a sus láminas" % nom, not otros and not lam, ", ".join(otros + lam))
        if nom.endswith(".docx"):
            cit = sorted({m.group(0) for p in pref for m in re.finditer(r"\b%s-\d{2}\b" % p, txt)} - codigos)
            I.add(g, "%s: láminas citadas existen en el juego de planos" % nom, not cit, ", ".join(cit))


def cifras(I, docs, cif):
    """Cifras clave del diseno y del metrado que deben aparecer tal cual en las memorias y especificaciones."""
    g = "Coherencia"
    txt = {}
    for k, r in docs.items():
        if r and os.path.exists(r): txt[k] = _textos_docx(r)[0].replace(",", "")
    for nombre, valor, donde in cif:
        for k in donde:
            if k not in txt: continue
            I.add(g, "%s en %s" % (nombre, k), valor.replace(",", "") in txt[k], "se busca %s" % valor)


# ============================================================================================== planos
def planos(I, doc, marcos):
    from ezdxf import bbox
    g = "Planos"
    msp = doc.modelspace()
    # escala real: cotas con texto fijo igual a su medida
    malas = []; n = 0
    for e in msp.query("DIMENSION"):
        t = (e.dxf.get("text", "") or "").strip()
        if t in ("", "<>"): continue
        m = re.fullmatch(r"\s*(-?\d+(?:\.\d+)?)\s*(m|cm)?\s*", t)              # solo textos que son una medida (no "H/2", "1\"")
        if not m: continue
        try: med = e.get_measurement()
        except Exception: continue
        if not isinstance(med, (int, float)): continue
        n += 1; v = float(m.group(1)) / (100.0 if m.group(2) == "cm" else 1.0)
        if abs(v - med) > 0.006 and abs(v - med * 100) > 0.6: malas.append("%s (mide %.3f)" % (t, med))
    I.add(g, "Escala real: cotas con texto fijo iguales a su medida", not malas, "%d cotas con texto fijo revisadas; %s" % (n, ", ".join(malas[:8])))
    I.add(g, "Unidades del dibujo en metros (1 unidad = 1 m)", doc.header.get("$INSUNITS", 0) == 6, "INSUNITS = %s" % doc.header.get("$INSUNITS"))
    # textos
    cajas = [(m["codigo"], box(m["x0"], m["y0"], m["x1"], m["y1"]), m["esc"]) for m in marcos]
    tx = list(msp.query("TEXT MTEXT"))
    for ins in msp.query("INSERT"):
        try: tx += [v for v in ins.virtual_entities() if v.dxftype() in ("TEXT", "MTEXT")]
        except Exception: pass
    por = {c: [] for c, _, _ in cajas}; fuera, vert = [], []
    for e in tx:
        try: ex = bbox.extents([e], fast=False)
        except Exception: continue
        if not ex.has_data: continue
        b = box(ex.extmin.x, ex.extmin.y, ex.extmax.x, ex.extmax.y)
        txt = e.dxf.text if e.dxftype() == "TEXT" else e.plain_text()
        if not txt.strip(): continue
        rot = e.dxf.get("rotation", 0.0) % 180
        if 60 < rot < 120: vert.append(txt[:20])
        lam = next((c for c, cb, _ in cajas if cb.contains(b.centroid)), None)
        if lam is None: fuera.append(txt[:20]); continue
        cb = next(cb for c, cb, _ in cajas if c == lam)
        if not cb.buffer(1e-6).contains(b): fuera.append("%s (%s)" % (txt[:20], lam))
        por[lam].append((b, txt))
    I.add(g, "Textos dentro del marco de su lámina", not fuera, "; ".join(fuera[:8]) + (" ... (%d)" % len(fuera) if len(fuera) > 8 else ""))
    I.add(g, "Sin textos verticales (ni girados más de 60°)", not vert, "; ".join(vert[:8]) + (" ... (%d)" % len(vert) if len(vert) > 8 else ""))
    from shapely.strtree import STRtree
    for c, cb, esc in cajas:
        L = por[c]
        if len(L) < 2: continue
        geoms = [b for b, _ in L]; tree = STRtree(geoms); enc = []
        for i, (b, t) in enumerate(L):
            for j in tree.query(b):
                if j <= i: continue
                b2, t2 = L[j]; a = b.intersection(b2).area
                if a > 0.15 * min(b.area, b2.area): enc.append("'%s' / '%s'" % (t[:18], t2[:18]))
        I.add(g, "%s: textos sin encimarse" % c, not enc, "%d pares: %s" % (len(enc), "; ".join(enc[:6])) if enc else "%d textos" % len(L))
    # colores
    am = []
    for e in msp:
        col = e.dxf.get("color", 256); tc = e.dxf.get("true_color", None)
        if col == 2 or tc == 0xFFFFFF: am.append(e.dxftype())
    for ly in doc.layers:
        if ly.color == 2 or ly.dxf.get("true_color", None) == 0xFFFFFF: am.append("capa " + ly.dxf.name)
    I.add(g, "Sin amarillo ni blanco puro", not am, ", ".join(sorted(set(am))[:10]))


def entregables(I, salida, archivos, laminas=None):
    """Juego completo de la seccion 8.4 de conocimiento/01: lo que falta queda como aviso (pendiente)."""
    g = "Entregables"
    hay = lambda n: any(a and os.path.basename(a) == n for a in archivos) or os.path.exists(os.path.join(salida, n))
    for nom, archivo in (("Juego de planos DXF con presentaciones", "PLANOS_CUNETAS.dxf"), ("Metrado con fórmulas", "METRADO_DRENAJE_PLUVIAL.xlsx"),
                         ("Memoria hidrológica (Word)", "MEMORIA_HIDROLOGICA_DRENAJE_PLUVIAL.docx"),
                         ("Memoria descriptiva (Word)", "MEMORIA_DESCRIPTIVA_DRENAJE_PLUVIAL.docx"),
                         ("Especificaciones técnicas (Word)", "ESPECIFICACIONES_TECNICAS_DRENAJE_PLUVIAL.docx"),
                         ("Memoria de cálculo (Excel con fórmulas)", "MEMORIA_CALCULO_DRENAJE_PLUVIAL.xlsx"),
                         ("Carátula e índice de planos (Word)", "CARATULA_E_INDICE_DE_PLANOS.docx"),
                         ("Carátulas de los documentos (Word)", "CARATULAS_DRENAJE_PLUVIAL.docx")):
        I.add(g, nom, hay(archivo), archivo if hay(archivo) else "pendiente: " + archivo, aviso=True)
    tit = [x["titulo"].upper() for x in (laminas or [])]
    for nom, clave in (("Lámina de detalles (rejilla, tapa de inspección, juntas, lloradero, montante en falsa columna, sumidero)", "DETALLES"),
                       ("Láminas de descargas (planta, perfil e isométrico de cada descarga)", "DESCARGA")):
        ok = any((clave in t) if clave == "DESCARGA" else t.startswith(clave) for t in tit)
        I.add(g, nom, ok, "" if ok else "pendiente en el juego de planos", aviso=True)
    png = [n for n in os.listdir(salida) if n.startswith("LAMINA_") and n.endswith(".png")] if os.path.isdir(salida) else []
    falta = [x["codigo"] for x in (laminas or []) if "LAMINA_%s.png" % x["codigo"] not in png]
    I.add(g, "Vista previa PNG de cada lámina", bool(png) and not falta, "%d PNG" % len(png) + ("; faltan " + ", ".join(falta) if falta else ""), aviso=True)


def avisos_diseno(I, avisos):
    """Datos que el programa tomo por defecto o que el proyectista debe confirmar (no bloquean)."""
    for a in avisos: I.add("Datos a confirmar", a.split(":")[0][:90], False, a, aviso=True)


# ============================================================================================== libro
def escribir(ruta, I):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "REVISION"
    ws["A1"] = "REVISIÓN GENERAL DE LOS ENTREGABLES DE DRENAJE PLUVIAL"; ws["A1"].font = Font(bold=True, size=13)
    r = I.resumen()
    ws["A2"] = "%d comprobaciones: %d cumplen, %d no cumplen, %d avisos" % (r["comprobaciones"], r["cumple"], r["no_cumple"], r["avisos"])
    cab = ["Grupo", "Comprobación", "Estado", "Detalle"]
    for j, h in enumerate(cab, 1):
        c = ws.cell(4, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    color = {CUMPLE: "E2EFDA", FALLA: "F8CBAD", AVISO: "FFF2CC"}
    for i, f in enumerate(I.filas, 5):
        for j, v in enumerate(f, 1):
            c = ws.cell(i, j, v); c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(i, 3).fill = PatternFill("solid", fgColor=color[f[2]])
    for col, w in zip("ABCD", (14, 62, 12, 90)): ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    wb.properties.creator = ""; wb.properties.lastModifiedBy = ""
    wb.save(ruta); return ruta
