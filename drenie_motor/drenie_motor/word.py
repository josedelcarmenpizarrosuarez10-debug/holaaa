"""Word de los documentos de DrenIE (python-docx): memoria hidrologica, memoria descriptiva y especificaciones tecnicas.

Mismo formato de los expedientes de drenaje pluvial (I.E. N.° 0590 Barranquita, C.I. 468901; CAR Mujeres, CUI 2717013):
A4, margenes 3.0 / 2.5 / 2.5 / 2.2 cm, Arial 10.5, titulos azules numerados, tablas con encabezado sombreado y numeradas,
figuras con leyenda, encabezado con el proyecto y pie con el numero de pagina. Sin nombre ni CIP del proyectista.
"""
import os

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1F, 0x3A, 0x68)
GRIS = RGBColor(0x59, 0x59, 0x59)
ESPECIALIDAD = "INSTALACIONES SANITARIAS - DRENAJE PLUVIAL"


def _sombra(celda, color):
    tcPr = celda._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)


def _campo(run, instr):
    for t, v in (("begin", None), ("instr", instr), ("separate", None), ("end", None)):
        if t == "instr":
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = v
        else:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        run._r.append(e)


def _borde_inferior(par, color="1F3A68", tam=6):
    pPr = par._p.get_or_add_pPr(); b = OxmlElement("w:pBdr"); e = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(tam)), ("w:space", "2"), ("w:color", color)): e.set(qn(k), v)
    b.append(e); pPr.append(b)


def codigo_de(proyecto):
    """Codigo con el que se cita el proyecto: CUI N.° ... o C.I. N.° ... tal cual lo trae el proyecto."""
    c = (proyecto.get("codigo") or "").strip()
    if c: return c
    return ("CUI N.° %s" % proyecto["cui"]) if proyecto.get("cui") else ""


class Doc:
    def __init__(self, proyecto, corto):
        self.pr = proyecto; self.d = Document(); self.n = [0, 0, 0]; self.nt = 0; self.nf = 0; self.indice = []
        st = self.d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        pf = st.paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.15
        for k, tam in (("Heading 1", 13), ("Heading 2", 11.5), ("Heading 3", 10.5), ("Heading 4", 10)):
            h = self.d.styles[k]; h.font.name = "Arial"; h.font.size = Pt(tam); h.font.bold = True; h.font.italic = False
            h.font.color.rgb = AZUL if k in ("Heading 1", "Heading 2") else RGBColor(0, 0, 0)
            rf = h.element.rPr.rFonts; rf.set(qn("w:eastAsia"), "Arial")
            for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"): rf.attrib.pop(qn(a), None)
            h.paragraph_format.space_before = Pt(12); h.paragraph_format.space_after = Pt(6)
        s = self.d.sections[0]; s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = Cm(3.0); s.right_margin = Cm(2.5); s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.2)
        s.different_first_page_header_footer = True              # la caratula va sin encabezado ni pie
        hp = s.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        nom = (proyecto.get("nombre_corto") or "").upper(); cod = codigo_de(proyecto)
        r = hp.add_run(" - ".join(x for x in (nom, cod, corto) if x)); r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
        _borde_inferior(hp, "999999", 4)
        fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run("Página "); r.font.size = Pt(8); r2 = fp.add_run(); r2.font.size = Pt(8); _campo(r2, "PAGE")

    # ------------------------------------------------------------------ caratula e indice
    def caratula(self, titulo, subtitulo=None):
        d, pr = self.d, self.pr
        def linea(txt, tam, neg=False, color=None, antes=0, despues=6, cursiva=False):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(antes); p.paragraph_format.space_after = Pt(despues)
            r = p.add_run(txt); r.font.size = Pt(tam); r.bold = neg; r.italic = cursiva
            if color: r.font.color.rgb = color
            return p
        if pr.get("entidad"): linea(pr["entidad"].upper(), 14, True, AZUL, despues=0)
        if pr.get("gerencia"): linea(pr["gerencia"].upper(), 11, True, GRIS, despues=0)
        _borde_inferior(d.paragraphs[-1] if d.paragraphs else linea("", 6), "1F3A68", 12)
        linea("EXPEDIENTE TÉCNICO", 12, True, AZUL, antes=10)
        linea("“%s”" % pr.get("nombre", "").strip("“”\""), 11.5, True, despues=2)
        cod = codigo_de(pr)
        if cod: linea(cod, 11, True, AZUL, despues=10)
        foto = pr.get("foto")
        if foto and os.path.exists(foto):
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(foto, width=Cm(14.5))
            if pr.get("foto_leyenda"): linea(pr["foto_leyenda"], 8, cursiva=True, color=GRIS)
        else:
            for _ in range(6): d.add_paragraph()
        t = d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]; _sombra(c, "1F3A68"); c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(titulo); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = RGBColor(255, 255, 255)
        if subtitulo:
            p = c.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = p.add_run(subtitulo); r.bold = True; r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(255, 255, 255)
        linea("ESPECIALIDAD: " + (pr.get("especialidad_documentos") or ESPECIALIDAD), 10.5, True, antes=12, despues=2)
        if pr.get("ubicacion"): linea("UBICACIÓN: " + pr["ubicacion"], 10, despues=8)
        if pr.get("fecha"): linea(pr["fecha"].upper(), 12, True, AZUL)
        self.salto()

    def contenido(self, titulos):
        self.h_simple("CONTENIDO")
        for t in titulos:
            p = self.d.add_paragraph(t); p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.left_indent = Cm(0.6 * t.split(" ")[0].count("."))
        self.salto()

    def indice_campo(self, titulo="ÍNDICE"):
        """Indice automatico de Word (se actualiza al abrir: F9 o Referencias > Actualizar tabla)."""
        p = self.d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(titulo); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL
        _campo(self.d.add_paragraph().add_run(), 'TOC \\o "1-4" \\h \\z \\u')
        upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); self.d.settings.element.append(upd)
        self.salto()

    def salto(self):
        self.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ titulos y texto
    def h_simple(self, t):
        p = self.d.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(13); r.font.color.rgb = AZUL

    def h(self, nivel, t):
        self.n[nivel - 1] += 1
        for k in range(nivel, 3): self.n[k] = 0
        num = ".".join(str(x) for x in self.n[:nivel])
        txt = "%s%s %s" % (num, "." if nivel == 1 else "", t.upper() if nivel == 1 else t)
        self.d.add_heading(txt, level=nivel); self.indice.append(txt)

    def titulo(self, txt, nivel, salto=False):
        h = self.d.add_heading(txt, level=min(nivel, 4)); h.paragraph_format.keep_with_next = True
        if salto: h.paragraph_format.page_break_before = True
        return h

    def p(self, t, neg=None, centro=False, tam=None, cursiva=False):
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centro else WD_ALIGN_PARAGRAPH.JUSTIFY
        if neg:
            r = par.add_run(neg); r.bold = True
            if tam: r.font.size = Pt(tam)
        r = par.add_run(t); r.italic = cursiva
        if tam: r.font.size = Pt(tam)
        return par

    def rotulo(self, t):
        par = self.d.add_paragraph(); par.paragraph_format.space_before = Pt(6); par.paragraph_format.space_after = Pt(3)
        par.paragraph_format.keep_with_next = True
        r = par.add_run(t); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = AZUL

    def vinetas(self, items, estilo="List Bullet", izq=False):
        for it in items:
            par = self.d.add_paragraph(style=estilo); par.alignment = WD_ALIGN_PARAGRAPH.LEFT if izq else WD_ALIGN_PARAGRAPH.JUSTIFY
            par.paragraph_format.space_after = Pt(3)
            if isinstance(it, tuple):
                r = par.add_run(it[0]); r.bold = True; par.add_run(it[1])
            else:
                par.add_run(it)

    def formula(self, t):
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run(t); r.italic = True; r.font.size = Pt(11)

    # ------------------------------------------------------------------ tablas y figuras
    def tabla(self, titulo, cab, filas, anchos=None, tam=8.5, fuente=None, total=False, numerada=True):
        if titulo:
            if numerada: self.nt += 1
            par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.keep_with_next = True
            r = par.add_run(("Tabla N° %02d. %s" % (self.nt, titulo)) if numerada else titulo); r.bold = True; r.font.size = Pt(9)
        t = self.d.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, c in enumerate(cab):
            cc = t.rows[0].cells[j]; cc.text = ""; rr = cc.paragraphs[0].add_run(str(c)); rr.bold = True; rr.font.size = Pt(tam)
            cc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER; _sombra(cc, "D9E2F3")
        trPr = t.rows[0]._tr.get_or_add_trPr(); e = OxmlElement("w:tblHeader"); e.set(qn("w:val"), "true"); trPr.append(e)
        for i, f in enumerate(filas):
            row = t.add_row().cells
            for j, v in enumerate(f):
                row[j].text = ""; rr = row[j].paragraphs[0].add_run("" if v is None else str(v)); rr.font.size = Pt(tam)
                row[j].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if (j == 0 or (isinstance(v, str) and len(v) > 18)) else WD_ALIGN_PARAGRAPH.CENTER
                if total and i == len(filas) - 1:
                    rr.bold = True; _sombra(row[j], "F2F2F2")
        if anchos:
            t.autofit = False
            tblW = t._tbl.tblPr.find(qn("w:tblW"))
            if tblW is None:
                tblW = OxmlElement("w:tblW"); t._tbl.tblPr.append(tblW)
            tblW.set(qn("w:type"), "dxa"); tblW.set(qn("w:w"), str(int(sum(anchos) * 567)))
            for j, gc in enumerate(t._tbl.tblGrid.findall(qn("w:gridCol"))):
                if j < len(anchos): gc.set(qn("w:w"), str(int(anchos[j] * 567)))
            for row in t.rows:
                for j, w in enumerate(anchos): row.cells[j].width = Cm(w)
        if fuente:
            par = self.d.add_paragraph(); r = par.add_run("Fuente: " + fuente); r.italic = True; r.font.size = Pt(8)
        else:
            self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def recuadro(self, texto, titulo="Fuente"):
        """Recuadro de fuente: tabla de una celda sombreada con borde izquierdo azul, letra de 8 pt."""
        t = self.d.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        cc = t.rows[0].cells[0]; cc.text = ""; _sombra(cc, "EEF6FD")
        bordes = OxmlElement("w:tcBorders")
        for lado, val, sz, col in (("left", "single", "24", "1EA7F2"), ("top", "nil", "0", "auto"), ("bottom", "nil", "0", "auto"), ("right", "nil", "0", "auto")):
            e = OxmlElement("w:" + lado); e.set(qn("w:val"), val); e.set(qn("w:sz"), sz); e.set(qn("w:color"), col); bordes.append(e)
        cc._tc.get_or_add_tcPr().append(bordes)
        par = cc.paragraphs[0]; r = par.add_run(titulo + ". "); r.bold = True; r.font.size = Pt(8); r.font.color.rgb = AZUL
        r = par.add_run(texto); r.font.size = Pt(8)
        self.d.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def figura(self, ruta, titulo, ancho=15.5, fuente=None):
        if not ruta or not os.path.exists(ruta): return
        self.nf += 1
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.keep_with_next = True
        par.add_run().add_picture(ruta, width=Cm(ancho))
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run("Figura N° %02d. %s" % (self.nf, titulo)); r.bold = True; r.font.size = Pt(9)
        if fuente:
            r = par.add_run("\nFuente: " + fuente); r.italic = True; r.font.size = Pt(8)

    def cuadro_figuras(self, rutas, subtitulos, titulo, fuente=None, ancho=15.5, cols=2):
        """Varias imágenes en un cuadro de cols columnas (sin bordes), cada una con su subtítulo; una sola leyenda de figura."""
        rutas = [(r, t) for r, t in zip(rutas, subtitulos) if r and os.path.exists(r)]
        if not rutas: return
        self.nf += 1
        filas = (len(rutas) + cols - 1) // cols
        t = self.d.add_table(rows=filas, cols=cols); t.alignment = WD_TABLE_ALIGNMENT.CENTER
        w = ancho / cols - 0.3
        for k, (r, sub) in enumerate(rutas):
            cc = t.rows[k // cols].cells[k % cols]; cc.width = Cm(ancho / cols); cc.text = ""
            par = cc.paragraphs[0]; par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if sub:
                par.paragraph_format.keep_with_next = True
                rr = par.add_run(sub); rr.bold = True; rr.font.size = Pt(8)
                par = cc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.add_run().add_picture(r, width=Cm(w))
        par = self.d.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run("Figura N° %02d. %s" % (self.nf, titulo)); r.bold = True; r.font.size = Pt(9)
        if fuente:
            r = par.add_run("\nFuente: " + fuente); r.italic = True; r.font.size = Pt(8)

    def guardar(self, ruta):
        cp = self.d.core_properties
        for k in ("author", "last_modified_by", "comments", "title", "subject", "keywords", "category"): setattr(cp, k, "")
        self.d.save(ruta); return ruta


# ---------------------------------------------------------------------------------------------- especificaciones
PESOS = {3: [26, 50, 24], 5: [44, 8, 17, 15, 16]}


def bloques(D, lista):
    """Contenido de espec_textos: ["p", texto], ["l", [items]], ["t", cabecera, filas]."""
    for b in lista:
        if b[0] == "p": D.p(b[1])
        elif b[0] == "l": D.vinetas(b[1])
        elif b[0] == "t":
            n = len(b[1]); pesos = PESOS.get(n, [100.0 / n] * n)
            D.tabla(None, b[1], b[2], [15.5 * w / 100.0 for w in pesos], tam=8.5, numerada=False)


def especificaciones(J, ruta, proyecto):
    """Documento de especificaciones tecnicas: caratula, indice, capitulo I (generalidades) y capitulo II (una ficha por partida)."""
    D = Doc(proyecto, "ESPECIFICACIONES TÉCNICAS")
    D.caratula("ESPECIFICACIONES TÉCNICAS", "Drenaje pluvial - red de cunetas")
    ind = ["CAPÍTULO I. GENERALIDADES"] + ["1.%d  %s" % (i + 1, g[0]) for i, g in enumerate(J["generalidades"])]
    ind += ["CAPÍTULO II. ESPECIFICACIONES TÉCNICAS POR PARTIDA"] + ["%s  %s" % (it["codigo"], it["desc"]) for it in J["partidas"]]
    D.h_simple("ÍNDICE")
    for x in ind:
        par = D.d.add_paragraph(); par.paragraph_format.space_after = Pt(1)
        cod = x.split("  ")[0]; niv = cod.count(".") if cod[:1].isdigit() else 0
        par.paragraph_format.left_indent = Cm(0.5 * max(0, niv - 1)) if cod[:1].isdigit() else Cm(0)
        r = par.add_run(x); r.font.size = Pt(8.5 if niv > 1 else 9.5); r.bold = not cod[:1].isdigit() or niv <= 1
    D.salto()
    D.titulo("CAPÍTULO I. GENERALIDADES", 1)
    for i, (nom, cont) in enumerate(J["generalidades"]):
        D.titulo("1.%d  %s" % (i + 1, nom), 2); bloques(D, cont)
    D.titulo("CAPÍTULO II. ESPECIFICACIONES TÉCNICAS POR PARTIDA", 1, salto=True)
    D.p("Las partidas se presentan con la numeración, la descripción y la unidad del metrado de la especialidad. Los títulos agrupan "
        "partidas y no tienen medición ni pago propio.")
    for k, it in enumerate(J["partidas"]):
        D.titulo("%s  %s%s" % (it["codigo"], it["desc"], ("  (%s)" % it["und"]) if it["und"] else ""), min(2 + it["nivel"], 4),
                 salto=(it["nivel"] == 0 and k > 0))
        for nom, cont in it.get("secciones") or []:
            D.rotulo(nom); bloques(D, cont)
    return D.guardar(ruta)
