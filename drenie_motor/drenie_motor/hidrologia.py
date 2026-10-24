"""Lluvia de diseno: de los datos de una estacion SENAMHI a la intensidad de diseno de las cunetas.

Mismos metodos de la hoja ANALISIS ESTADISTICO de los CAR (usados en Varones, Refugio, Mujeres y Barranquita):
1. Serie de precipitacion maxima en 24 h: de los datos diarios de la estacion (maxima de cada mes y de cada anio) o dada por el
   usuario (maximas mensuales o anuales). Factor de transposicion opcional (1.0 = serie medida, sin transponer).
2. Distribuciones: Normal, Log Normal, Log Pearson III (factor de frecuencia de Kite con k = Cs/6 de los logaritmos) y Gumbel
   (yn y sn de Aparicio segun n). Variable normal para P = 1/T con la aproximacion de Abramowitz y Stegun.
3. Kolmogorov-Smirnov con la probabilidad de Weibull m/(n+1) sobre la serie de mayor a menor; se adopta la de menor D.
4. Lluvia de duracion D (Dick y Peschke): P_D = P24 (D/1440)^0.25; I = 60 P_D / D, con D = tc (minimo 10 min).
5. Curvas IDF: modelo de Bell para 5 a 60 min con P60 de 10 anios = 0.4602 P24(10)^0.876 y regresion I = K T^m / t^n.

Uso: python -m drenie_motor lluvia proyecto.json   -> salida/hidrologia.json, HIDROLOGIA.xlsx y CURVAS_IDF.png
"""
import csv
import datetime as dt
import math
import os
import re

import numpy as np

PERIODOS = [2, 5, 10, 20, 25, 50, 100, 200, 500]
DURACIONES = [5, 10, 15, 20, 30, 60, 90, 120, 240, 360, 720, 1440]
DUR_BELL = [5, 10, 15, 20, 30, 60]
GUMBEL_YN_SN = [(10, 0.4952, 0.9496), (15, 0.5128, 1.0206), (20, 0.5236, 1.0628), (25, 0.5309, 1.0914), (30, 0.5362, 1.1124),
                (90, 0.5586, 1.2007), (95, 0.5593, 1.2037)]
ALFA_KS = 0.05
NOMBRE_DIST = {"normal": "Normal", "lognormal": "Log Normal", "lp3": "Log Pearson III", "gumbel": "Gumbel"}
MIN_DIAS_MES = 20            # mes con datos suficientes para su maxima
MIN_DIAS_ANIO = 300          # anio con datos suficientes para la maxima anual


# ------------------------------------------------------------------------------------------- lectura de datos diarios
def _num(v):
    v = (v or "").strip().replace(",", ".")
    if v.upper() in ("T", "TR", "TRAZA"): return 0.0
    try: x = float(v)
    except ValueError: return None
    return None if x < 0 or x > 900 else x


def _filas(ruta):
    if ruta.lower().endswith((".xlsx", ".xls")):
        import openpyxl
        ws = openpyxl.load_workbook(ruta, data_only=True, read_only=True).active
        return [["" if c is None else (c.strftime("%Y-%m-%d") if hasattr(c, "strftime") else str(c)) for c in r]
                for r in ws.iter_rows(values_only=True)]
    txt = open(ruta, encoding="utf-8", errors="replace").read()
    if "-END HEADER-" in txt: txt = txt.split("-END HEADER-", 1)[1]
    lineas = [l for l in txt.splitlines() if l.strip()]
    muestra = "\n".join(lineas[:20])
    if muestra.count(",") >= muestra.count(";") and "," in muestra: sep = ","
    elif ";" in muestra: sep = ";"
    elif "\t" in muestra: sep = "\t"
    else: return [l.split() for l in lineas]
    return list(csv.reader(lineas, delimiter=sep))


def _norm(s):
    s = s.upper()
    for a, b in (("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"), ("Ñ", "N")): s = s.replace(a, b)
    return re.sub(r"[^A-Z0-9/]", "", s)


def leer_diario(ruta):
    """{fecha: mm} de un archivo diario. Formatos: SENAMHI (con encabezado AÑO/MES/DÍA o historico sin encabezado
    anio mes dia precipitacion ...; -99.9 = sin dato), NASA POWER (YEAR, MO, DY o DOY, PRECTOTCORR) y fecha + precipitacion."""
    filas = _filas(ruta); datos = {}
    i0, col = 0, None
    for i, f in enumerate(filas[:40]):
        h = [_norm(c) for c in f]
        if any(x in ("ANO", "YEAR", "FECHA", "DATE", "ANO/MES/DIA") or x.startswith("ANO") for x in h):
            col = h; i0 = i + 1; break

    def ip(nombres):
        for j, c in enumerate(col or []):
            if any(c.startswith(n) for n in nombres): return j
        return None
    if col:
        jp = ip(["PRECTOT", "PRECIP", "PP", "PREC", "LLUVIA", "RAIN", "CHIRPS"])
        jf = ip(["FECHA", "DATE", "ANO/MES/DIA"]); ja = ip(["ANO", "YEAR"]); jm = ip(["MES", "MO"]); jd = ip(["DIA", "DY"])
        jdoy = ip(["DOY"])
        if jp is None: jp = len(col) - 1
    else:
        jf, ja, jm, jd, jdoy, jp = None, 0, 1, 2, None, 3
    for f in filas[i0:]:
        try:
            if jf is not None and (ja is None or jf == ja):
                t = f[jf].strip().replace("/", "-")
                p = [int(x) for x in re.split(r"[-\s]", t)[:3]]
                d = dt.date(p[0], p[1], p[2]) if p[0] > 31 else dt.date(p[2], p[1], p[0])
            elif jdoy is not None and jm is None:
                d = dt.date(int(float(f[ja])), 1, 1) + dt.timedelta(days=int(float(f[jdoy])) - 1)
            else:
                d = dt.date(int(float(f[ja])), int(float(f[jm])), int(float(f[jd])))
        except (ValueError, IndexError):
            continue
        v = _num(f[jp]) if jp < len(f) else None
        if v is not None: datos[d] = v
    return datos


def serie_mensual(datos, desde=None, hasta=None):
    """Maxima de 24 h de cada mes {anio: [12 valores o None]}; un mes con menos de 20 dias con dato queda en blanco."""
    meses = {}
    for d, v in datos.items():
        if (desde and d.year < desde) or (hasta and d.year > hasta): continue
        meses.setdefault((d.year, d.month), []).append(v)
    out = {}
    for a in sorted({a for a, _ in meses}):
        ms = [meses.get((a, m)) for m in range(1, 13)]
        if sum(len(x) for x in ms if x) < MIN_DIAS_ANIO: continue          # anio incompleto: no entra en la serie
        out[a] = [round(max(x), 1) if x and len(x) >= MIN_DIAS_MES else None for x in ms]
    return out


# ------------------------------------------------------------------------------------------- estadistica
def _phi(z): return 0.5 * (1 + math.erf(z / math.sqrt(2)))


def estad(x):
    x = np.array(x, float); n = len(x); m = x.mean(); s = x.std(ddof=1)
    cs = n * ((x - m) ** 3).sum() / ((n - 1) * (n - 2) * s ** 3)
    return n, m, s, cs


def z_aprox(T):
    p = 1.0 / T; w = math.sqrt(math.log(1.0 / p ** 2))
    return w - (2.515517 + 0.802853 * w + 0.010328 * w * w) / (1 + 1.432788 * w + 0.189269 * w * w + 0.001308 * w ** 3)


def kt_lp3(z, k):
    return z + (z * z - 1) * k + (z ** 3 - 6 * z) * k * k / 3 - (z * z - 1) * k ** 3 + z * k ** 4 + k ** 5 / 3


def kt_gumbel(T): return -(math.sqrt(6) / math.pi) * (0.5772 + math.log(math.log(T / (T - 1))))


def _chi2_sf(x, gl):
    from scipy.special import gammaincc
    gl = int(gl)
    return float(gammaincc(gl / 2.0, x / 2.0)) if x > 0 else 1.0


def analisis(maximos, periodos=PERIODOS):
    """maximos: [(anio, P24 max)]. Tabla por periodo de retorno, Kolmogorov-Smirnov y distribucion de mejor ajuste."""
    x = [p for _, p in maximos]; n, m, s, cs = estad(x)
    lx = [math.log10(p) for p in x]; _, ml, sl, csl = estad(lx)
    a_g = math.sqrt(6) * s / math.pi; u_g = m - 0.5772 * a_g
    tab = {}
    for T in periodos:
        z = z_aprox(T); kp = kt_lp3(z, csl / 6)
        tab[T] = {"normal": m + z * s, "lognormal": 10 ** (ml + z * sl), "lp3": 10 ** (ml + kp * sl),
                  "gumbel": m + kt_gumbel(T) * s, "z": z, "kt_gumbel": kt_gumbel(T), "kt_lp3": kp}
    xd = sorted(x, reverse=True); E = [(i + 1) / (n + 1) for i in range(n)]
    lnx = [math.log(v) for v in x]; _, mL, sL, csL = estad(lnx)
    b1 = (2 / csL) ** 2 if csL > 0 else None; a1 = sL / math.sqrt(b1) if b1 else None; d1 = mL - a1 * b1 if b1 else None
    yn = float(np.interp(n, [t[0] for t in GUMBEL_YN_SN], [t[1] for t in GUMBEL_YN_SN]))
    sn = float(np.interp(n, [t[0] for t in GUMBEL_YN_SN], [t[2] for t in GUMBEL_YN_SN]))
    tabla_ks = []
    for i, v in enumerate(xd):
        fn = 1 - _phi((v - m) / s); fl = 1 - _phi((math.log10(v) - ml) / sl)
        fp = _chi2_sf(2 * (math.log(v) - d1) / a1, 2 * b1) if b1 else None
        fg = 1 - math.exp(-math.exp(-(((v - m) / s) * sn + yn)))
        tabla_ks.append({"m": i + 1, "P": v, "E": E[i], "normal": fn, "lognormal": fl, "lp3": fp, "gumbel": fg})
    ks = {d: max(abs(t[d] - t["E"]) for t in tabla_ks) if all(t[d] is not None for t in tabla_ks) else float("inf")
          for d in ("normal", "lognormal", "lp3", "gumbel")}
    crit = math.sqrt(-math.log(ALFA_KS / 2) / (2 * n)) - 1 / (6 * n)
    return {"n": n, "media": m, "desv": s, "cs": cs, "media_log": ml, "desv_log": sl, "cs_log": csl, "gumbel_alfa": a_g, "gumbel_u": u_g,
            "gumbel_yn": yn, "gumbel_sn": sn, "ln": {"media": mL, "desv": sL, "cs": csL, "beta": b1, "alfa": a1, "delta": d1},
            "tabla": tab, "ks": ks, "ks_critico": crit, "elegida": min(ks, key=ks.get), "tabla_ks": tabla_ks}


def p_dur(P24, D): return P24 * (D / 1440.0) ** 0.25


def bell(P60_10, T, t): return (0.21 * math.log(T) + 0.52) * (0.54 * t ** 0.25 - 0.50) * P60_10


# ------------------------------------------------------------------------------------------- calculo completo
def serie_de(lluvia, base):
    """Serie de maximas mensuales {anio: [12]} desde el proyecto: 'archivo' (datos diarios), 'serie' (maximas mensuales) o
    'maximos' (maxima anual)."""
    if lluvia.get("archivo"):
        ruta = lluvia["archivo"] if os.path.isabs(lluvia["archivo"]) else os.path.join(base, lluvia["archivo"])
        return serie_mensual(leer_diario(ruta), lluvia.get("desde"), lluvia.get("hasta")), "datos diarios de " + os.path.basename(ruta)
    if lluvia.get("serie"):
        return {int(a): m for a, m in lluvia["serie"].items()}, "maximas mensuales del proyecto"
    if lluvia.get("maximos"):
        return {int(a): [p] + [None] * 11 for a, p in lluvia["maximos"].items()}, "maximas anuales del proyecto"
    return None, None


def calcular(lluvia, base="."):
    """Resultado completo de la lluvia de diseno, o None si el proyecto solo trae I_mmh o P24."""
    sm, fuente = serie_de(lluvia, base)
    if not sm: return None
    factor = float(lluvia.get("factor", 1.0)); T = int(lluvia.get("T", 25)); tc = max(10.0, float(lluvia.get("tc_min", 10.0)))
    periodos = sorted(set(PERIODOS) | {T})
    maximos = [(a, round(max(v for v in m if v is not None) * factor, 2)) for a, m in sorted(sm.items()) if any(v is not None for v in m)]
    A = analisis(maximos, periodos)
    d = lluvia.get("distribucion") or A["elegida"]
    if d not in NOMBRE_DIST: raise ValueError("Distribucion desconocida: %s" % d)
    p24 = {t: A["tabla"][t][d] for t in periodos}
    idf = {t: {D: p_dur(p24[t], D) * 60.0 / D for D in DURACIONES} for t in periodos}
    P60_10 = 0.4602 * p24[10] ** 0.876
    bell_p = {t: {u: bell(P60_10, t, u) for u in DUR_BELL} for t in PERIODOS}
    bell_p[10][60] = P60_10
    bell_i = {t: {u: round(60.0 * bell_p[t][u] / u, 1) for u in DUR_BELL} for t in PERIODOS}
    X, Y = [], []
    for u in DUR_BELL:
        for t in sorted(PERIODOS, reverse=True):
            X.append([math.log10(t), math.log10(u), 1.0]); Y.append(math.log10(bell_i[t][u]))
    X, Y = np.array(X), np.array(Y)
    M = np.array([[X[:, 0].sum(), X[:, 1].sum(), len(Y)], [(X[:, 0] ** 2).sum(), (X[:, 0] * X[:, 1]).sum(), X[:, 0].sum()],
                  [(X[:, 0] * X[:, 1]).sum(), (X[:, 1] ** 2).sum(), X[:, 1].sum()]])
    V = np.array([Y.sum(), (X[:, 0] * Y).sum(), (X[:, 1] * Y).sum()])
    mm, mn, lk = np.linalg.solve(M, V)
    P24d = p24[T]; Pd = p_dur(P24d, tc); I = Pd * 60.0 / tc
    avisos = []
    if A["ks"][d] > A["ks_critico"]: avisos.append("La distribucion adoptada no pasa Kolmogorov-Smirnov (D = %.4f > %.4f)." % (A["ks"][d], A["ks_critico"]))
    if A["n"] < 15: avisos.append("La serie tiene solo %d anios; se recomiendan al menos 25." % A["n"])
    return {"estacion": lluvia.get("estacion", {}), "fuente_serie": fuente, "factor": factor, "serie_mensual": sm, "maximos": maximos,
            "analisis": A, "distribucion": d, "distribucion_nombre": NOMBRE_DIST[d], "T": T, "tc_min": tc, "p24": p24, "idf": idf,
            "P60_10": P60_10, "bell_i": bell_i, "K": round(10 ** lk, 2), "m": mm, "n": abs(mn),
            "P24": P24d, "P_tc": Pd, "I": I, "avisos": avisos}


def resumen(R):
    """Lo que necesita el diseno y lo que se muestra en el programa."""
    A = R["analisis"]
    return {"estacion": R["estacion"], "serie": R["fuente_serie"], "anios": A["n"], "desde": R["maximos"][0][0], "hasta": R["maximos"][-1][0],
            "factor_transposicion": R["factor"], "media_mm": round(A["media"], 2), "desv_mm": round(A["desv"], 2),
            "kolmogorov_smirnov": {NOMBRE_DIST[k]: round(v, 4) for k, v in A["ks"].items()}, "ks_critico": round(A["ks_critico"], 4),
            "distribucion": R["distribucion_nombre"], "T_anios": R["T"], "P24_mm": round(R["P24"], 2), "tc_min": R["tc_min"],
            "P_tc_mm": round(R["P_tc"], 2), "I_mmh": round(R["I"], 2),
            "P24_por_T": {str(t): round(v, 2) for t, v in R["p24"].items()},
            "IDF_bell": {"K": R["K"], "m": round(R["m"], 4), "n": round(R["n"], 4), "formula": "I = %.2f T^%.4f / t^%.4f" % (R["K"], R["m"], R["n"])},
            "avisos": R["avisos"]}


# ------------------------------------------------------------------------------------------- salidas
def png_idf(R, ruta):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(1, 2, figsize=(13, 5.2), dpi=130)
    D = DURACIONES
    for t in [2, 5, 10, 25, 50, 100]:
        if t in R["idf"]: ax[0].plot(D, [R["idf"][t][d] for d in D], marker="o", ms=3, lw=1.6 if t == R["T"] else 1, label="T = %d años" % t)
    ax[0].axvline(R["tc_min"], color="0.4", ls=":", lw=1); ax[0].plot([R["tc_min"]], [R["I"]], "o", color="crimson")
    ax[0].annotate("I = %.2f mm/h\n(T = %d, tc = %.0f min)" % (R["I"], R["T"], R["tc_min"]), (R["tc_min"], R["I"]), xytext=(25, R["I"] * 0.95), fontsize=9)
    ax[0].set_xscale("log"); ax[0].set_yscale("log"); ax[0].grid(True, which="both", alpha=.3)
    ax[0].set_xlabel("Duración (min)"); ax[0].set_ylabel("Intensidad (mm/h)"); ax[0].set_title("Curvas intensidad - duración - frecuencia (Dick y Peschke)", fontsize=10)
    ax[0].legend(fontsize=8)
    anios = [a for a, _ in R["maximos"]]; P = [p for _, p in R["maximos"]]
    ax[1].bar(anios, P, color="#4f7fb8"); ax[1].axhline(R["analisis"]["media"], color="0.3", ls="--", lw=1, label="media %.1f mm" % R["analisis"]["media"])
    ax[1].set_xlabel("Año"); ax[1].set_ylabel("P máx 24 h (mm)"); ax[1].legend(fontsize=8); ax[1].grid(axis="y", alpha=.3)
    est = R["estacion"].get("nombre", "")
    ax[1].set_title("Precipitación máxima en 24 h%s" % (" - estación " + est if est else ""), fontsize=10)
    fig.tight_layout(); fig.savefig(ruta); plt.close(fig)
    return ruta


def excel(R, ruta):
    """Libro de la lluvia de diseno: datos, analisis de frecuencias, Kolmogorov-Smirnov, IDF y resultado."""
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = openpyxl.Workbook()
    neg = Font(bold=True); cab = PatternFill("solid", fgColor="D9E2F3"); fino = Side(style="thin", color="999999")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)

    def tabla(ws, fila, col, encabezado, filas, formatos=None):
        for j, h in enumerate(encabezado):
            c = ws.cell(fila, col + j, h); c.font = neg; c.fill = cab; c.border = borde; c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i, f in enumerate(filas):
            for j, v in enumerate(f):
                c = ws.cell(fila + 1 + i, col + j, v); c.border = borde
                if formatos and formatos[j] and isinstance(v, (int, float)): c.number_format = formatos[j]
        return fila + len(filas) + 2
    est = R["estacion"]
    ws = wb.active; ws.title = "RESULTADO"
    ws["A1"] = "LLUVIA DE DISEÑO"; ws["A1"].font = Font(bold=True, size=13)
    filas = [("Estación", est.get("nombre", "")), ("Código", est.get("codigo", "")), ("Entidad", est.get("entidad", "SENAMHI")),
             ("Serie", R["fuente_serie"]), ("Años de la serie", R["analisis"]["n"]), ("Periodo", "%d - %d" % (R["maximos"][0][0], R["maximos"][-1][0])),
             ("Factor de transposición", R["factor"]), ("Distribución adoptada (menor D de Kolmogorov-Smirnov)", R["distribucion_nombre"]),
             ("Periodo de retorno T (años)", R["T"]), ("P máx 24 h para T (mm)", round(R["P24"], 2)), ("Tiempo de concentración tc (min)", R["tc_min"]),
             ("Lluvia de duración tc, Dick y Peschke (mm)", round(R["P_tc"], 2)), ("Intensidad de diseño I (mm/h)", round(R["I"], 2)),
             ("Curva IDF (modelo de Bell)", "I = %.2f T^%.4f / t^%.4f" % (R["K"], R["m"], R["n"]))]
    for i, (a, b) in enumerate(filas):
        ws.cell(3 + i, 1, a).border = borde; c = ws.cell(3 + i, 2, b); c.border = borde
    ws.cell(3 + 12, 2).font = Font(bold=True, color="C00000")
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 34
    # datos
    w = wb.create_sheet("DATOS")
    w["A1"] = "PRECIPITACIÓN MÁXIMA EN 24 H POR MES (mm)"; w["A1"].font = neg
    meses = ["ENE", "FEB", "MAR", "ABR", "MAY", "JUN", "JUL", "AGO", "SET", "OCT", "NOV", "DIC"]
    filas = [[a] + list(m) + [max(v for v in m if v is not None)] for a, m in sorted(R["serie_mensual"].items()) if any(v is not None for v in m)]
    tabla(w, 3, 1, ["AÑO"] + meses + ["MÁX ANUAL"], filas, [None] + ["0.0"] * 13)
    # analisis
    A = R["analisis"]; w = wb.create_sheet("ANALISIS")
    w["A1"] = "ANÁLISIS DE FRECUENCIAS"; w["A1"].font = neg
    f = tabla(w, 3, 1, ["Parámetro", "Valor"], [("n", A["n"]), ("Media (mm)", A["media"]), ("Desviación estándar (mm)", A["desv"]),
              ("Coeficiente de asimetría", A["cs"]), ("Media de log10 P", A["media_log"]), ("Desviación de log10 P", A["desv_log"]),
              ("Asimetría de log10 P", A["cs_log"]), ("Gumbel yn (Aparicio)", A["gumbel_yn"]), ("Gumbel sn (Aparicio)", A["gumbel_sn"])],
              [None, "0.0000"])
    per = sorted(A["tabla"])
    f = tabla(w, f, 1, ["T (años)", "z", "Normal", "Log Normal", "Log Pearson III", "Gumbel"],
              [[t, A["tabla"][t]["z"], A["tabla"][t]["normal"], A["tabla"][t]["lognormal"], A["tabla"][t]["lp3"], A["tabla"][t]["gumbel"]] for t in per],
              [None, "0.0000"] + ["0.00"] * 4)
    # Kolmogorov-Smirnov
    w = wb.create_sheet("KOLMOGOROV")
    w["A1"] = "PRUEBA DE KOLMOGOROV-SMIRNOV (probabilidad empírica de Weibull m/(n+1))"; w["A1"].font = neg
    f = tabla(w, 3, 1, ["m", "P (mm)", "P empírica", "Normal", "Log Normal", "Log Pearson III", "Gumbel"],
              [[t["m"], t["P"], t["E"], t["normal"], t["lognormal"], t["lp3"], t["gumbel"]] for t in A["tabla_ks"]],
              [None, "0.0", "0.0000", "0.0000", "0.0000", "0.0000", "0.0000"])
    tabla(w, f, 1, ["Distribución", "D máximo", "D crítico (α = 0.05)", "¿Ajusta?"],
          [[NOMBRE_DIST[k], v, A["ks_critico"], "Sí" if v <= A["ks_critico"] else "No"] for k, v in A["ks"].items()], [None, "0.0000", "0.0000", None])
    # IDF
    w = wb.create_sheet("IDF")
    w["A1"] = "INTENSIDADES (mm/h), Dick y Peschke con la distribución adoptada"; w["A1"].font = neg
    f = tabla(w, 3, 1, ["T (años)", "P24 (mm)"] + ["%d min" % d for d in DURACIONES],
              [[t, R["p24"][t]] + [R["idf"][t][d] for d in DURACIONES] for t in sorted(R["idf"])], [None, "0.00"] + ["0.00"] * len(DURACIONES))
    w.cell(f, 1, "Modelo de Bell (5 a 60 min), P60 de 10 años = 0.4602 P24(10)^0.876 = %.2f mm" % R["P60_10"]).font = neg
    tabla(w, f + 1, 1, ["T (años)"] + ["%d min" % d for d in DUR_BELL], [[t] + [R["bell_i"][t][d] for d in DUR_BELL] for t in PERIODOS], [None] + ["0.0"] * 6)
    wb.save(ruta)
    return ruta
