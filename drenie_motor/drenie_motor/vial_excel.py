"""Memoria de cálculo del drenaje vial con fórmulas (MEMORIA_CALCULO_DRENAJE_VIAL.xlsx).

Hojas: RESUMEN, DATOS (parámetros editables), LLUVIA (P24 por T e intensidad de Dick y Peschke), APORTES (anchos y C de cada
cuneta), CUNETAS (capacidad por Manning, longitud máxima y tramos hasta cada alivio), CUENCAS (C, tiempos de concentración de la
Tabla N.° 05 y caudal racional), ALCANTARILLAS (Manning en sección circular o marco con borde libre del 25 %), BADENES y NORMAS.
Los valores coinciden con vial.py; los tirantes normales van como valor (Manning resuelto por iteración) y el Excel recalcula con
ellos el caudal de Manning para comprobarlo.
"""
from openpyxl import Workbook
from openpyxl.utils import get_column_letter as CL

from . import normas
from .memoria_calculo import AM, AZ, VE, F, _F, cel, tabla_cab
from .vial import progresiva


def libro(ruta, R):
    pr = R["proyecto"]; via = R["via"]
    cab_txt = ["PROYECTO: " + pr.get("nombre", "") + ((" - " + pr["codigo"]) if pr.get("codigo") else ""),
               "  -  ".join(x for x in (("ENTIDAD: " + pr["entidad"]) if pr.get("entidad") else "",
                                        ("UBICACIÓN: " + pr["ubicacion"]) if pr.get("ubicacion") else "",
                                        ("FECHA: " + pr["fecha"]) if pr.get("fecha") else "") if x)]

    def cabecera(ws, titulo, sub=None):
        ws.cell(1, 1, titulo).font = F(size=11, bold=True)
        for k, t in enumerate(cab_txt): ws.cell(2 + k, 1, t).font = F(size=8)
        if sub: ws.cell(4, 1, sub).font = F(size=8, italic=True)
        ws.sheet_view.showGridLines = False; ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.page_setup.fitToHeight = 0

    wb = Workbook(); wr = wb.active; wr.title = "RESUMEN"
    # ------------------------------------------------------------------ DATOS
    wd = wb.create_sheet("DATOS"); cabecera(wd, "DATOS GENERALES Y PARÁMETROS DEL DRENAJE VIAL")
    tabla_cab(wd, 5, ["PARÁMETRO", "VALOR", "", "FUENTE / CRITERIO"], [70, 13, 2, 86])
    D = {}

    def dato(clave, texto, valor, fmt, fuente, formula=None):
        r = 6 + len(D); D[clave] = "DATOS!$B$%d" % r
        cel(wd, r, 1, texto); cel(wd, r, 2, formula or valor, fmt, bold=True, fill=VE if formula else AM); cel(wd, r, 4, fuente)
    dato("vel", "Velocidad de diseño de la vía (km/h)", float(via.get("velocidad_kmh", 40)), "0", "Dato del proyecto")
    dato("imda", "Índice medio diario anual IMDA (veh/día)", float(via.get("imda", 0)), "0", "Dato del proyecto")
    dato("pp", "Precipitación media anual (mm)", float(via.get("precipitacion_anual_mm", 2000)), "0",
         "Región " + R["region"].lower() + " (" + _F("vial_cuneta_tabla34") + ")")
    Pp = R["periodo_plataforma"]
    dato("R_pl", "Riesgo admisible, drenaje de la plataforma", Pp["R"], "0.00", _F("vial_riesgo_tabla02"))
    dato("n_pl", "Vida útil, drenaje de la plataforma (años)", Pp["n"], "0", _F("vial_riesgo_tabla02"))
    dato("Tc_pl", "Período de retorno calculado T = 1/(1−(1−R)^(1/n)) (años)", None, "0.0", "Ecuación 1, numeral 3.6",
         "=1/(1-(1-%s)^(1/%s))" % (D["R_pl"], D["n_pl"]))
    dato("T_pl", "Período de retorno adoptado para cunetas (entero superior)", None, "0", "Se adopta el entero superior",
         "=ROUNDUP(%s-0.000001,0)" % D["Tc_pl"])
    dato("tc", "Tiempo de concentración de la plataforma y mínimo de cuencas (min)", R["tc_min_norma"], "0.0", _F("tc_min"))
    dato("Lnorma", "Longitud máxima de cuneta hasta su alivio (m)", 200.0 if float(via.get("precipitacion_anual_mm", 2000)) >= 1600 else 250.0,
         "0", _F("vial_cuneta_longitud"))
    dato("bl", "Borde libre mínimo de alcantarillas (fracción de la altura)", normas.dato("vial_alc_borde")["valor"] / 100.0, "0.00", _F("vial_alc_borde"))
    dato("vmin", "Velocidad mínima en alcantarillas (m/s)", normas.dato("vial_alc_vmin")["valor"], "0.00", _F("vial_alc_vmin"))
    dato("g", "Aceleración de la gravedad (m/s²)", 9.81, "0.00", "")
    wd.freeze_panes = "A6"

    # ------------------------------------------------------------------ LLUVIA
    wl = wb.create_sheet("LLUVIA"); cabecera(wl, "LLUVIA DE DISEÑO", "Fuente de la serie: " + R["lluvia"]["fuente"])
    tabla_cab(wl, 5, ["T (años)", "P24 (mm)", "Uso", "", "Duración (min)", "I (mm/h) por Dick y Peschke"], [11, 12, 60, 2, 14, 26])
    LL = {}
    usos = {}
    usos.setdefault(Pp["T"], []).append("cunetas (plataforma)")
    for a in R["alcantarillas"]: usos.setdefault(a["periodo"]["T"], []).append(a["id"])
    for b in R["badenes"]: usos.setdefault(b["periodo"]["T"], []).append(b["id"])
    for i, (t, p) in enumerate(sorted(((int(t), p) for t, p in R["lluvia"]["P24_por_T"].items()))):
        r = 6 + i; LL[t] = r
        cel(wl, r, 1, t, "0"); cel(wl, r, 2, p, "0.00", fill=AM); cel(wl, r, 3, ", ".join(usos.get(t, [])), wrap=True)
    cel(wl, 6, 5, None, "0.0", fill=VE, bold=True).value = "=%s" % D["tc"]
    cel(wl, 6, 6, "=60*VLOOKUP(%s,$A$6:$B$%d,2,FALSE)*(E6/1440)^0.25/E6" % (D["T_pl"], 5 + len(LL)), "0.00", bold=True)
    cel(wl, 7 + len(LL), 1, "Dick y Peschke: P_D = P24 (D/1440)^0.25; I = 60 P_D / D. " + _F("dick_peschke"), borde=False)
    I_pl = "LLUVIA!$F$6"

    # ------------------------------------------------------------------ APORTES
    wa = wb.create_sheet("APORTES"); cabecera(wa, "ÁREAS DE APORTE A LAS CUNETAS (por metro de cuneta)")
    tabla_cab(wa, 5, ["Cuneta", "Aporte", "Ancho (m)", "C", "C × ancho (m)", "Fuente de C"], [10, 40, 11, 8, 13, 50])
    r = 6
    for c in R["cunetas"]:
        for a in c["aportes"]:
            cel(wa, r, 1, c["id"]); cel(wa, r, 2, a.get("descripcion", "")); cel(wa, r, 3, float(a["ancho"]), "0.00", fill=AM)
            cel(wa, r, 4, float(a["C"]), "0.00", fill=AM); cel(wa, r, 5, "=C%d*D%d" % (r, r), "0.000"); cel(wa, r, 6, a.get("fuente_C", "Dato del proyecto"))
            r += 1
    ra = r - 1

    # ------------------------------------------------------------------ CUNETAS
    wc = wb.create_sheet("CUNETAS")
    cabecera(wc, "CUNETAS: CAPACIDAD, LONGITUD MÁXIMA Y DESCARGAS",
             "Q = A R^(2/3) S^(1/2) / n (ecuación 117); capacidad = cuneta llena, limitada por la velocidad admisible (Tabla N.° 33).")
    cols = ["Cuneta", "Lado", "Prog. inicio", "Prog. fin", "L (m)", "S (m/m)", "Sección", "d (m)", "a o b (m)", "z1 (H:1)", "z2 (H:1)",
            "n", "A llena (m²)", "P llena (m)", "R (m)", "V llena (m/s)", "V adm. (m/s)", "Q capacidad (m³/s)", "Σ C·ancho (m)",
            "q (m³/s/m)", "L capacidad (m)", "L máx. (m)", "N.° de tramos", "L tramo (m)", "Estado"]
    tabla_cab(wc, 6, cols, [9, 9, 10, 10, 8, 8, 11, 7, 8, 8, 8, 7, 9, 9, 8, 9, 9, 11, 10, 11, 11, 9, 8, 9, 11])
    CU = {}
    for i, c in enumerate(R["cunetas"]):
        r = 7 + i; CU[c["id"]] = r; s = c["seccion"]
        cel(wc, r, 1, c["id"]); cel(wc, r, 2, c["lado"]); cel(wc, r, 3, c["prog_ini"], "0.00", fill=AM); cel(wc, r, 4, c["prog_fin"], "0.00", fill=AM)
        cel(wc, r, 5, "=D%d-C%d" % (r, r), "0.00"); cel(wc, r, 6, c["pendiente"], "0.0000", fill=AM); cel(wc, r, 7, s["tipo"])
        cel(wc, r, 8, s["d"], "0.00", fill=AM)
        if s["tipo"] == "triangular":
            cel(wc, r, 9, float(s["a"]), "0.00", fill=AM); cel(wc, r, 10, "=I%d/H%d" % (r, r), "0.00"); cel(wc, r, 11, s["z2"], "0.00", fill=AM)
            A = "H{0}^2*(J{0}+K{0})/2"; P = "H{0}*(SQRT(1+J{0}^2)+SQRT(1+K{0}^2))"
        elif s["tipo"] == "trapezoidal":
            cel(wc, r, 9, s["b"], "0.00", fill=AM); cel(wc, r, 10, s["z1"], "0.00", fill=AM); cel(wc, r, 11, s["z2"], "0.00", fill=AM)
            A = "H{0}*(I{0}+H{0}*(J{0}+K{0})/2)"; P = "I{0}+H{0}*(SQRT(1+J{0}^2)+SQRT(1+K{0}^2))"
        else:
            cel(wc, r, 9, s["b"], "0.00", fill=AM); cel(wc, r, 10, 0, "0.00"); cel(wc, r, 11, 0, "0.00")
            A = "H{0}*I{0}"; P = "I{0}+2*H{0}"
        cel(wc, r, 12, c["n"], "0.000", fill=AM)
        cel(wc, r, 13, "=" + A.format(r), "0.0000"); cel(wc, r, 14, "=" + P.format(r), "0.000"); cel(wc, r, 15, "=M%d/N%d" % (r, r), "0.0000")
        cel(wc, r, 16, "=O%d^(2/3)*SQRT(F%d)/L%d" % (r, r, r), "0.00"); cel(wc, r, 17, c["Vmax"], "0.00", fill=AM)
        # capacidad: llena si V llena ≤ V adm.; si no, el valor del tirante con V = V adm. calculado por el programa
        if c["V_lleno"] <= c["Vmax"] + 1e-9: cel(wc, r, 18, "=M%d*P%d" % (r, r), "0.0000", bold=True)
        else: cel(wc, r, 18, c["Q_capacidad"], "0.0000", bold=True)
        cel(wc, r, 19, "=SUMIF(APORTES!$A$6:$A$%d,A%d,APORTES!$E$6:$E$%d)" % (ra, r, ra), "0.000")
        cel(wc, r, 20, "=S%d*%s/3600000" % (r, I_pl), "0.000000")
        cel(wc, r, 21, "=(R%d-%s)/T%d" % (r, c["Q_inicial"], r), "0.0")
        cel(wc, r, 22, "=MIN(U%d,%s)" % (r, D["Lnorma"]), "0.0")
        cel(wc, r, 23, "=MAX(1,ROUNDUP(E%d/V%d-0.000001,0))" % (r, r), "0"); cel(wc, r, 24, "=E%d/W%d" % (r, r), "0.00")
        cel(wc, r, 25, '=IF(R%d>=T%d*X%d+%s,"CUMPLE","NO CUMPLE")' % (r, r, r, c["Q_inicial"]), bold=True)
    r0 = 9 + len(R["cunetas"])
    cel(wc, r0, 1, "TRAMOS DE CUNETA HASTA CADA DESCARGA", bold=True, borde=False)
    tabla_cab(wc, r0 + 1, ["Cuneta", "Desde", "Hasta", "Descarga en", "L (m)", "Q (m³/s)", "Tirante y (m)", "A (m²)", "P (m)",
                           "Q Manning con y (m³/s)", "V (m/s)", "V adm. (m/s)", "Verificación"])
    r = r0 + 2
    for c in R["cunetas"]:
        k = CU[c["id"]]; s = c["seccion"]
        for t in c["tramos"]:
            cel(wc, r, 1, c["id"]); cel(wc, r, 2, progresiva(t["desde"])); cel(wc, r, 3, progresiva(t["hasta"]))
            cel(wc, r, 4, progresiva(t["descarga_en"])); cel(wc, r, 5, t["longitud"], "0.00")
            cel(wc, r, 6, "=$T$%d*E%d%s" % (k, r, ("+%s" % c["Q_inicial"]) if c["Q_inicial"] and t is c["tramos"][0] else ""), "0.0000")
            cel(wc, r, 7, t["tirante"], "0.000", fill=AM)
            if s["tipo"] == "triangular": A, P = "G{0}^2*($J${1}+$K${1})/2", "G{0}*(SQRT(1+$J${1}^2)+SQRT(1+$K${1}^2))"
            elif s["tipo"] == "trapezoidal": A, P = "G{0}*($I${1}+G{0}*($J${1}+$K${1})/2)", "$I${1}+G{0}*(SQRT(1+$J${1}^2)+SQRT(1+$K${1}^2))"
            else: A, P = "G{0}*$I${1}", "$I${1}+2*G{0}"
            cel(wc, r, 8, "=" + A.format(r, k), "0.0000"); cel(wc, r, 9, "=" + P.format(r, k), "0.000")
            cel(wc, r, 10, "=H%d*(H%d/I%d)^(2/3)*SQRT($F$%d)/$L$%d" % (r, r, r, k, k), "0.0000")
            cel(wc, r, 11, "=J%d/H%d" % (r, r), "0.00"); cel(wc, r, 12, "=$Q$%d" % k, "0.00")
            cel(wc, r, 13, '=IF(AND(K%d<=L%d,G%d<=$H$%d),"CUMPLE","NO CUMPLE")' % (r, r, r, k), bold=True)
            r += 1
    wc.freeze_panes = "B7"

    # ------------------------------------------------------------------ CUENCAS
    wq = wb.create_sheet("CUENCAS")
    cabecera(wq, "CUENCAS: COEFICIENTE DE ESCORRENTÍA, TIEMPO DE CONCENTRACIÓN Y CAUDAL",
             "Tabla N.° 05 (tc en minutos) y método racional Q = 0.278 C I A (A < 10 km²) o racional modificado de Témez (10 a 770 km²).")
    tabla_cab(wq, 6, ["Obra", "Prog.", "A (km²)", "L cauce (m)", "H desnivel (m)", "S (m/m)", "C", "Fuente de C", "tc Kirpich (min)",
                      "tc California (min)", "tc Témez (min)", "tc adoptado (min)", "Método tc", "T (años)", "P24 (mm)", "I (mm/h)", "K",
                      "Q cuenca (m³/s)", "Método"],
              [8, 9, 8, 10, 10, 8, 7, 44, 10, 10, 10, 10, 10, 7, 9, 9, 7, 10, 30])
    QC = {}
    r = 7
    for x in R["alcantarillas"] + R["badenes"]:
        h = x.get("hidrologia")
        if not h: continue
        QC[x["id"]] = r
        cel(wq, r, 1, x["id"]); cel(wq, r, 2, progresiva(x["prog"])); cel(wq, r, 3, x["cuenca"]["area_km2"], "0.000", fill=AM)
        cel(wq, r, 4, x["cuenca"]["L_m"], "0.0", fill=AM); cel(wq, r, 5, x["cuenca"]["H_m"], "0.0", fill=AM); cel(wq, r, 6, "=E%d/D%d" % (r, r), "0.0000")
        cel(wq, r, 7, h["C"], "0.000", fill=AM); cel(wq, r, 8, h["fuente_C"], wrap=True)
        cel(wq, r, 9, "=0.01947*D%d^0.77*F%d^-0.385" % (r, r), "0.00")
        cel(wq, r, 10, "=0.0195*(D%d^3/E%d)^0.385" % (r, r), "0.00")
        cel(wq, r, 11, "=60*0.3*((D%d/1000)/F%d^0.25)^0.76" % (r, r), "0.00")
        met = h["tc_metodo"]
        f = {"promedio": "AVERAGE(I{0}:K{0})", "minimo": "MIN(I{0}:K{0})", "Kirpich": "I{0}", "California Culverts Practice": "J{0}",
             "Témez": "K{0}"}.get(met, None)
        if h["K"] == 1.0:
            cel(wq, r, 12, ("=MAX(%s,%s)" % (D["tc"], f.format(r))) if f else h["tc_min"], "0.00", bold=True)
        else:
            cel(wq, r, 12, "=K%d" % r, "0.00", bold=True)
        cel(wq, r, 13, met); T = x["periodo"]["T"]; cel(wq, r, 14, T, "0")
        cel(wq, r, 15, "=VLOOKUP(N%d,LLUVIA!$A$6:$B$%d,2,FALSE)" % (r, 5 + len(LL)), "0.00")
        if h["K"] == 1.0:
            cel(wq, r, 16, "=60*O%d*(L%d/1440)^0.25/L%d" % (r, r, r), "0.00"); cel(wq, r, 17, 1, "0.000")
            cel(wq, r, 18, "=0.278*G%d*P%d*C%d" % (r, r, r), "0.000", bold=True)
        else:                                      # Témez: tc en horas = L/60; kA, P corregida e intensidad
            tch = "(L{0}/60)".format(r)
            cel(wq, r, 16, "=((1-LOG10(C{0})/15)*O{0}/24)*11^((28^0.1-{1}^0.1)/(28^0.1-1))".format(r, tch), "0.00")
            cel(wq, r, 17, "=1+{0}^1.25/({0}^1.25+14)".format(tch), "0.000")
            cel(wq, r, 18, "=0.278*G%d*P%d*C%d*Q%d" % (r, r, r, r), "0.000", bold=True)
        cel(wq, r, 19, h["metodo"], wrap=True)
        r += 1
    wq.freeze_panes = "C7"

    # ------------------------------------------------------------------ ALCANTARILLAS
    wal = wb.create_sheet("ALCANTARILLAS")
    cabecera(wal, "ALCANTARILLAS: CAUDAL Y SECCIÓN (Manning, ecuación 47, con borde libre del 25 %)",
             "Sección circular: θ = 2·ACOS(1 − 2y/D); A = D²(θ − sen θ)/8; P = Dθ/2. Marco: A = b·y; P = b + 2y. d50 de salida por Laushey (ecuación 49).")
    tabla_cab(wal, 6, ["Obra", "Prog.", "Tipo", "Q cuenca (m³/s)", "Q cunetas (m³/s)", "Q diseño (m³/s)", "Conducto", "D o h (m)", "b (m)",
                       "Líneas", "S (m/m)", "n", "Tirante y (m)", "y / altura", "θ (rad)", "A (m²)", "P (m)", "Q Manning (m³/s)",
                       "V (m/s)", "d50 salida (m)", "Verificación", "Recibe"],
              [8, 9, 8, 10, 10, 10, 18, 8, 7, 7, 8, 7, 9, 8, 8, 8, 8, 10, 8, 9, 12, 40])
    r = 7
    for a in R["alcantarillas"]:
        s = a["seccion"]
        cel(wal, r, 1, a["id"]); cel(wal, r, 2, progresiva(a["prog"])); cel(wal, r, 3, a["tipo"])
        cel(wal, r, 4, ("=CUENCAS!R%d" % QC[a["id"]]) if a["id"] in QC else 0, "0.000")
        cel(wal, r, 5, a["Q_cunetas"], "0.0000", fill=AM); cel(wal, r, 6, "=D%d+E%d" % (r, r), "0.000", bold=True)
        cel(wal, r, 7, a["conducto"])
        circ = s["tipo"] == "circular"
        cel(wal, r, 8, s["D"] if circ else s["h"], "0.000", fill=AM); cel(wal, r, 9, None if circ else s["b"], "0.00", fill=None if circ else AM)
        cel(wal, r, 10, s.get("celdas", 1), "0", fill=AM); cel(wal, r, 11, a["pendiente"], "0.0000", fill=AM); cel(wal, r, 12, a["n"], "0.000", fill=AM)
        cel(wal, r, 13, a["tirante"], "0.000", fill=AM); cel(wal, r, 14, "=M%d/H%d" % (r, r), "0.000")
        if circ:
            cel(wal, r, 15, "=2*ACOS(1-2*M%d/H%d)" % (r, r), "0.0000")
            cel(wal, r, 16, "=J{0}*H{0}^2*(O{0}-SIN(O{0}))/8".format(r), "0.0000"); cel(wal, r, 17, "=J{0}*H{0}*O{0}/2".format(r), "0.000")
        else:
            cel(wal, r, 15, None); cel(wal, r, 16, "=J{0}*I{0}*M{0}".format(r), "0.0000"); cel(wal, r, 17, "=J{0}*(I{0}+2*M{0})".format(r), "0.000")
        cel(wal, r, 18, "=P{0}*(P{0}/Q{0})^(2/3)*SQRT(K{0})/L{0}".format(r), "0.000")
        cel(wal, r, 19, "=R%d/P%d" % (r, r), "0.00"); cel(wal, r, 20, "=S%d^2/(3.1*%s)" % (r, D["g"]), "0.000")
        cel(wal, r, 21, '=IF(AND(N{0}<=1-{1},R{0}>=F{0}*0.999,S{0}>={2},IF(ISBLANK(I{0}),TRUE,S{0}<=6)),"CUMPLE","REVISAR")'.format(r, D["bl"], D["vmin"]), bold=True)
        cel(wal, r, 22, "; ".join(a["desde_cunetas"]), wrap=True)
        r += 1
    wal.freeze_panes = "C7"

    # ------------------------------------------------------------------ BADENES
    if R["badenes"]:
        wb_ = wb.create_sheet("BADENES")
        cabecera(wb_, "BADENES: CANAL TRAPEZOIDAL EN RÉGIMEN UNIFORME (numeral 4.1.1.4.2)",
                 _F("vial_baden_bl") + "; " + _F("vial_baden_pt"))
        tabla_cab(wb_, 6, ["Obra", "Prog.", "Q cuenca (m³/s)", "Q cunetas (m³/s)", "Q diseño (m³/s)", "b (m)", "z (H:1)", "Profundidad (m)",
                           "S transversal", "n", "Tirante y (m)", "A (m²)", "P (m)", "Q Manning (m³/s)", "V (m/s)", "Borde libre (m)", "Verificación"],
                  [8, 9, 10, 10, 10, 7, 7, 10, 9, 7, 9, 8, 8, 10, 8, 9, 12])
        r = 7
        for b in R["badenes"]:
            s = b["seccion"]
            cel(wb_, r, 1, b["id"]); cel(wb_, r, 2, progresiva(b["prog"])); cel(wb_, r, 3, ("=CUENCAS!R%d" % QC[b["id"]]) if b["id"] in QC else 0, "0.000")
            cel(wb_, r, 4, b["Q_cunetas"], "0.0000", fill=AM); cel(wb_, r, 5, "=C%d+D%d" % (r, r), "0.000", bold=True)
            cel(wb_, r, 6, s["b"], "0.00", fill=AM); cel(wb_, r, 7, s["z1"], "0.0", fill=AM); cel(wb_, r, 8, b["profundidad"], "0.00", fill=AM)
            cel(wb_, r, 9, b["pendiente_transversal"], "0.000", fill=AM); cel(wb_, r, 10, b["n"], "0.000", fill=AM)
            cel(wb_, r, 11, b["tirante"], "0.000", fill=AM)
            cel(wb_, r, 12, "=K{0}*(F{0}+G{0}*K{0})".format(r), "0.000"); cel(wb_, r, 13, "=F{0}+2*K{0}*SQRT(1+G{0}^2)".format(r), "0.000")
            cel(wb_, r, 14, "=L{0}*(L{0}/M{0})^(2/3)*SQRT(I{0})/J{0}".format(r), "0.000"); cel(wb_, r, 15, "=N%d/L%d" % (r, r), "0.00")
            cel(wb_, r, 16, "=H%d-K%d" % (r, r), "0.000")
            cel(wb_, r, 17, '=IF(AND(P{0}>=0.3,N{0}>=E{0}*0.999),"CUMPLE","NO CUMPLE")'.format(r), bold=True)
            r += 1

    # ------------------------------------------------------------------ RESUMEN
    cabecera(wr, "MEMORIA DE CÁLCULO DEL DRENAJE VIAL – RESUMEN",
             "Manual de Hidrología, Hidráulica y Drenaje del MTC (R.D. N.° 20-2011-MTC/14). Celdas amarillas: datos; verdes: calculadas.")
    tabla_cab(wr, 6, ["Obra", "Progresiva", "Tipo", "T (años)", "Q (m³/s)", "Sección adoptada", "Estado"], [12, 14, 26, 9, 11, 34, 12])
    r = 7
    for c in R["cunetas"]:
        s = c["seccion"]
        dim = ("triangular d = %.2f m, a = %.2f m" % (s["d"], float(s["a"]))) if s["tipo"] == "triangular" else \
              ("%s b = %.2f m, d = %.2f m" % (s["tipo"], s["b"], s["d"]))
        fila = [c["id"], "%s a %s" % (progresiva(c["prog_ini"]), progresiva(c["prog_fin"])), "Cuneta lado " + c["lado"], Pp["T"],
                max(t["Q"] for t in c["tramos"]), dim + ", %d alivio(s)" % (len(c["tramos"]) - 1), "=CUNETAS!Y%d" % CU[c["id"]]]
        for j, v in enumerate(fila, 1): cel(wr, r, j, v, "0.000" if j == 5 else None, bold=(j == 7))
        r += 1
    for k, a in enumerate(R["alcantarillas"]):
        fila = [a["id"], progresiva(a["prog"]), "Alcantarilla de " + ("alivio" if a["tipo"] == "alivio" else "paso de quebrada"),
                a["periodo"]["T"], "=ALCANTARILLAS!F%d" % (7 + k), a["conducto"], "=ALCANTARILLAS!U%d" % (7 + k)]
        for j, v in enumerate(fila, 1): cel(wr, r, j, v, "0.000" if j == 5 else None, bold=(j == 7))
        r += 1
    for k, b in enumerate(R["badenes"]):
        fila = [b["id"], progresiva(b["prog"]), "Badén", b["periodo"]["T"], "=BADENES!E%d" % (7 + k),
                "trapezoidal b = %.2f m, z = %.0f, h = %.2f m" % (b["seccion"]["b"], b["seccion"]["z1"], b["profundidad"]), "=BADENES!Q%d" % (7 + k)]
        for j, v in enumerate(fila, 1): cel(wr, r, j, v, "0.000" if j == 5 else None, bold=(j == 7))
        r += 1
    if R["avisos"]:
        r += 1; cel(wr, r, 1, "OBSERVACIONES DEL CÁLCULO", bold=True, borde=False)
        for t in R["avisos"]:
            r += 1; cel(wr, r, 1, "• " + t, borde=False)
    normas.hoja_excel(wb)
    wb.save(ruta)
    return ruta
