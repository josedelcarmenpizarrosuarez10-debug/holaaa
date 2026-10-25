"""Observaciones y soluciones propuestas (ventana de soluciones del visor).

Lo que el programa no pudo resolver solo, o que el proyectista debe confirmar, no se esconde: cada caso sale como una observacion
con que paso, donde, por que y las soluciones posibles. Cada solucion es
- "automatica": el programa la aplica en el siguiente calculo (lleva un "cambio" que viaja en ediciones.json),
- "manual": el proyectista lo resuelve a mano (dibuja en el visor o corrige el plano), o
- "confirmar": el proyectista acepta lo que hizo el programa.
La eleccion del proyectista viaja en ediciones.json ("soluciones": [{"obs", "opcion", "tipo", "cambio"}]) y el motor la aplica
(cambio) o la registra (manual / confirmar) en el siguiente calculo.
Los id son estables entre calculos (no dependen del orden): "nivel", "techo:T-02", "montante:M-06", "tubo:S-01", ...
"""
import hashlib

MANUAL = {"id": "M", "texto": "Lo resuelvo a mano", "tipo": "manual", "cambio": {"manual": True},
          "efecto": "El programa deja este punto como está y lo anota como resuelto por el proyectista."}
DESPUES = "Lo analizo después (la observación sigue pendiente)."


def _obs(id_, gravedad, titulo, que, porque, donde=None, opciones=(), recomendada=None):
    ops = [dict(o) for o in opciones]
    if recomendada is None and ops: recomendada = ops[0]["id"]
    for o in ops: o["recomendada"] = o["id"] == recomendada
    return {"id": id_, "gravedad": gravedad, "titulo": titulo, "que_paso": que, "por_que": porque,
            "donde": {"x": round(donde[0], 3), "y": round(donde[1], 3)} if donde else None, "opciones": ops, "estado": "pendiente"}


def armar(cfg, nivel=None, plan_inf=(), TT=None, SM=None, V=None, avisos=(), revision=None, R=None, S=None, elegidas=None, da=None,
          sin_gravedad=None):
    """Lista de observaciones del calculo. elegidas: {obs: {"opcion", "tipo"}} que el proyectista ya eligio en el visor."""
    out = []
    npt = (cfg.get("niveles") or {}).get("npt")
    # ---------------------------------------------------------------- salida por gravedad
    if nivel and nivel[0] > 0:
        ops = []
        if len(nivel) > 3 and nivel[3]:
            ops.append({"id": "A", "tipo": "automatica", "texto": "Subir el nivel de piso %.2f m%s" % (nivel[0], (" (NPT %.2f)" % (npt + nivel[0])) if isinstance(npt, (int, float)) else ""),
                        "efecto": "Sube el NPT, el borde de cuneta, el área verde y las zonas %.2f m; verificado: con ese nivel toda la red sale por gravedad." % nivel[0],
                        "cambio": {"subir_nivel": nivel[0]}})
        ops.append({"id": "B", "tipo": "manual", "texto": "Indicar una descarga a un colector o canal existente",
                    "efecto": "Marque en el visor (comando DESCARGA) el punto del cerco donde sale el agua, o agregue el receptor al proyecto."})
        ops.append({"id": "C", "tipo": "automatica", "texto": "Descargar sin gravedad con cámara de bombeo",
                    "efecto": "Dos bombas sumergibles alternadas donde no hay gravedad (CE.040 21.2.2; IS.010 7.1 f; OS.080).",
                    "cambio": {"descarga_alternativa": {"tipo": "bombeo"}}})
        ops.append({"id": "D", "tipo": "automatica", "texto": "Descargar sin gravedad a pozos de absorción",
                    "efecto": "Necesita el test de percolación del lugar y la napa (IS.020); sin ellos se diseña con bombeo.",
                    "cambio": {"descarga_alternativa": {"tipo": "infiltracion"}}})
        ops.append(dict(MANUAL, efecto="La parte que sí sale queda diseñada; el resto lo resuelve el proyectista."))
        out.append(_obs("nivel", "decision", "El agua no sale por gravedad en toda la red",
                        "Con el nivel de piso dado: %s." % nivel[1],
                        "El terreno exterior queda más alto que la cota de fondo que alcanza la cuneta en el cerco; la cuneta nunca va enterrada.",
                        _punto_sin_salida(R, V), ops))
    # ---------------------------------------------------------------- techos donde no cabe una cuneta automatica
    for t in plan_inf or ():
        if t.get("decision") != "sin solucion automatica": continue
        p = _centro_techo(TT, t["techo"])
        out.append(_obs("techo:" + t["techo"], "decision", "%s: no cabe una cuneta automática alrededor del techo" % t["techo"],
                        "Ninguna combinación de cunetas a lo largo del alero o transversales deja a sus %d montantes a 5 m o menos." % t.get("montantes", 0),
                        "El cerco u otras edificaciones están demasiado cerca de sus lados.", p,
                        [{"id": "A", "tipo": "manual", "texto": "Dibujar la cuneta en el visor", "efecto": "Comando CUNETA: trace la cuneta que reciba esos montantes; el programa la une a la red."},
                         {"id": "B", "tipo": "manual", "texto": "Mover los montantes a otra esquina", "efecto": "En el visor, mueva los montantes (MT) hacia la esquina que tiene cuneta cerca."},
                         MANUAL]))
    # ---------------------------------------------------------------- montantes que no llegan (agrupados por techo)
    sin_cun = {}
    for m in (TT or {}).get("montantes", []):
        if m.get("tubo_enterrado_m"):
            if m.get("llega") and not m["llega"].startswith("SOBRE"):
                out.append(_obs("montante:" + m["montante"], "revisar", "%s: su tubo llega bajo el agua de la cuneta" % m["montante"],
                                "El tubo del dado llega a %.3f y el nivel mínimo (agua + 0.05) es %.3f en %s." % (m["z_llegada"], m["nivel_min"], m["cuneta"]),
                                "Para recibirlo, la cuneta tendría que pasar la profundidad máxima.", (m["x"], m["y"]),
                                [{"id": "A", "tipo": "manual", "texto": "Mover el montante a otra esquina", "efecto": "En el visor (MT), llévelo a la esquina con la cuneta más baja."},
                                 {"id": "B", "tipo": "manual", "texto": "Dibujar otra cuneta más cerca", "efecto": "Comando CUNETA en el visor."}, MANUAL]))
            continue
        sin_cun.setdefault(m["techo"], []).append(m)
    nivel_obs = next((o for o in out if o["id"] == "nivel"), None)
    for techo, ms in sorted(sin_cun.items()):
        if any(o["id"] == "techo:" + techo for o in out): continue
        ids = ", ".join(m["montante"] for m in ms)
        if nivel_obs is not None:
            # misma causa: esas cunetas no se construyeron porque el agua no tenia salida; se resuelven con la solucion de nivel
            nivel_obs.setdefault("relacionados", []).append("%s (%s)" % (techo, ids)); continue
        p = (sum(m["x"] for m in ms) / len(ms), sum(m["y"] for m in ms) / len(ms))
        out.append(_obs("montante:" + techo, "decision", "%s: %d montante%s sin cuneta a 5 m o menos" % (techo, len(ms), "s" if len(ms) > 1 else ""),
                        "%s: la cuneta más cercana está a más de 5 m del dado; el tubo terminaría muy profundo." % ids,
                        "No hay cuneta cerca de esas esquinas del techo.", p,
                        [{"id": "A", "tipo": "manual", "texto": "Dibujar una cuneta cerca de esas esquinas", "efecto": "Comando CUNETA en el visor (a lo largo del alero o transversal)."},
                         {"id": "B", "tipo": "manual", "texto": "Mover los montantes a otra esquina", "efecto": "Comando MT en el visor."}, MANUAL]))
    if nivel_obs is not None and nivel_obs.get("relacionados"):
        nivel_obs["que_paso"] += (" Por la misma causa quedan sin cuneta los montantes de %s; se resuelven al aplicar esta solución."
                                  % "; ".join(nivel_obs["relacionados"]))
    # ---------------------------------------------------------------- sumideros
    for s in (SM or {}).get("sumideros", []):
        if not s["llega"].startswith("SOBRE"):
            out.append(_obs("tubo:" + s["sumidero"], "revisar", "%s: su tubo llega bajo el agua de la cuneta" % s["sumidero"],
                            "Llega a %.3f y el nivel mínimo (agua + 0.05) es %.3f en %s." % (s["z_llegada"], s["nivel_min"], s["cuneta"]),
                            "La cuneta que lo recibe va muy llena en ese punto.", (s["x"], s["y"]),
                            [{"id": "A", "tipo": "manual", "texto": "Mover el sumidero hacia otra cuneta", "efecto": "Comando SU en el visor."}, MANUAL]))
    # ---------------------------------------------------------------- descargas indicadas que no salen
    if cfg.get("descargas") and not (nivel and nivel[0] > 0):
        for d in (V or {}).get("descargas", []):
            if d["estado"] != "SALE POR GRAVEDAD" and not d.get("alternativa"):
                q = (R.descargas.get(d["D"], {}) or {}).get("q") if R is not None else None
                out.append(_obs("descarga:" + d["D"], "revisar", "La descarga %s no sale por gravedad" % d["D"],
                                "Cota de fondo %.3f; el terreno exterior no baja de ella cerca del cerco." % d["cf"], "La descarga indicada está en un lado alto del cerco.", q,
                                [{"id": "A", "tipo": "automatica", "texto": "Dejar que el programa elija las descargas", "efecto": "Quita las descargas fijadas y busca las del lado bajo.",
                                  "cambio": {"quitar_descargas_fijas": True}},
                                 {"id": "B", "tipo": "manual", "texto": "Mover la descarga en el visor", "efecto": "Comando DESCARGA en el visor."}, MANUAL]))
    # ---------------------------------------------------------------- descargas sin gravedad (pozos de absorcion o camara de bombeo)
    if da:
        tipos = sorted({d["tipo"] for d in da.values()})
        hay_pozo = any(d["tipo"] == "pozo de absorción" for d in da.values())
        ops = [{"id": "A", "tipo": "confirmar", "texto": "Mantener %s" % " y ".join(tipos), "efecto": "Se deja así; el diseño completo está en la "
                "lámina, el metrado, la memoria y las especificaciones.", "cambio": {"confirmar": True}}]
        if hay_pozo:
            ops.append({"id": "B", "tipo": "automatica", "texto": "Cambiar a cámara de bombeo", "efecto": "Dos bombas sumergibles alternadas en cada descarga sin gravedad.",
                        "cambio": {"descarga_alternativa": {"tipo": "bombeo"}}})
        else:
            ops.append({"id": "B", "tipo": "automatica", "texto": "Cambiar a pozos de absorción (infiltración)", "efecto": "Necesita el test de percolación "
                        "del lugar (minutos para bajar 1 cm, IS.020) y la profundidad de la napa; sin ellos el programa vuelve a bombeo.",
                        "cambio": {"descarga_alternativa": {"tipo": "infiltracion"}}})
        if sin_gravedad and sin_gravedad.get("subir_npt_m"):
            ops.append({"id": "C", "tipo": "automatica", "texto": "Subir el nivel de piso %.2f m y descargar por gravedad" % sin_gravedad["subir_npt_m"],
                        "efecto": sin_gravedad.get("texto_nivel", ""), "cambio": {"subir_nivel": sin_gravedad["subir_npt_m"],
                                                                                 "descarga_alternativa": {"tipo": None}}})
        ops.append(MANUAL)
        d0 = sorted(da)[0]
        q = R.descargas[d0]["q"] if R is not None and d0 in getattr(R, "descargas", {}) else None
        out.append(_obs("sin-gravedad", "decision", "Descarga sin gravedad: %s" % ", ".join("%s a %s" % (k, v["tipo"]) for k, v in sorted(da.items())),
                        ("Con los niveles del proyecto el agua no sale por gravedad (%s)." % sin_gravedad["detalle"]) if sin_gravedad else
                        "El proyecto pide descargar sin gravedad en %s." % ", ".join(sorted(da)),
                        "El terreno exterior queda más alto que la cota de fondo a la que llegan las cunetas; la cuneta nunca va enterrada.",
                        q, ops, recomendada="C" if (sin_gravedad and sin_gravedad.get("verificado") and
                                                    0 < (sin_gravedad.get("subir_npt_m") or 9) <= 0.50) else "A"))
    # ---------------------------------------------------------------- remanso que la descarga ya no puede corregir (perfil de flujo)
    rem = next((x for x in (revision or {}).get("decisiones", []) if "Perfil de flujo" in x), None)
    if rem:
        import re
        Ds = re.findall(r"\b(D\d+) \(C\.F\.", rem)
        q = (R.descargas[Ds[0]]["q"] if R is not None and Ds and Ds[0] in getattr(R, "descargas", {}) else None)
        out.append(_obs("remanso", "decision", "El nivel del agua deja poco borde libre (remanso desde la descarga)", rem.split(" | ")[0],
                        "Con el caudal de diseño, el agua de los tramos finales sube hacia aguas arriba; para bajarla habría que bajar la "
                        "descarga, y el terreno exterior ya no lo permite.", q,
                        [{"id": "A", "tipo": "manual", "texto": "Adoptar un ancho libre mayor en los tramos finales",
                          "efecto": "Hoja DIMENSIONAMIENTO (ancho adoptado) o «secciones» del proyecto, en los ejes que llegan a la descarga."},
                         {"id": "B", "tipo": "manual", "texto": "Indicar otra descarga", "efecto": "Comando DESCARGA en el visor, en otro lado bajo del cerco."},
                         {"id": "C", "tipo": "confirmar", "texto": "Aceptar el borde libre calculado", "efecto": "Se deja así; queda anotado en la memoria.",
                          "cambio": {"confirmar": True}}, MANUAL]))
    # ---------------------------------------------------------------- lo que el programa supuso y conviene confirmar
    for a in avisos or ():
        if "techos sin flecha de caida" in a:
            out.append(_obs("caida-techos", "confirmar", "Caída de los techos supuesta", a,
                            "El plano no indica hacia dónde cae cada techo.", None,
                            [{"id": "A", "tipo": "confirmar", "texto": "Está bien: dos aguas con la cumbrera en el eje largo", "efecto": "Se deja asi.", "cambio": {"confirmar": True}},
                             {"id": "B", "tipo": "manual", "texto": "Dibujar las flechas de caída en el visor", "efecto": "Comando FLECHA en el visor, una por cada agua."}]))
        elif "sin identificar en el plano se tomaron como" in a:
            out.append(_obs("superficies-defecto", "confirmar", "Superficies sin identificar", a,
                            "El plano no tiene sombreado ni rótulo en esas áreas.", None,
                            [{"id": "A", "tipo": "confirmar", "texto": "Está bien el tipo por defecto", "efecto": "Se deja asi.", "cambio": {"confirmar": True}},
                             {"id": "B", "tipo": "manual", "texto": "Asignar el tipo en el visor", "efecto": "Pestaña Superficies: clic en cada área y elija su tipo."}]))
    # ---------------------------------------------------------------- toda otra falla de la revision general
    for f in (revision or {}).get("fallas", []):
        h = "revision:" + hashlib.md5(f.split(":")[0].encode("utf-8")).hexdigest()[:8]
        out.append(_obs(h, "revisar", "Revisión general: " + f.split(":")[0], f,
                        "La revisión automática de los entregables encontró esto y el programa no tiene una corrección automática para este caso.",
                        None, [MANUAL]))
    # ---------------------------------------------------------------- textos de otras partes del motor (sin tildes) con sus tildes
    from .espec_textos import acentos
    for o in out:
        o["que_paso"] = acentos(o["que_paso"])
    # ---------------------------------------------------------------- lo que ya eligio el proyectista
    for o in out:
        e = (elegidas or {}).get(o["id"])
        if e:
            o["estado"] = {"manual": "resuelto a mano por el proyectista", "confirmar": "confirmado por el proyectista",
                           "automatica": "solución aplicada; persiste: revisar"}.get(e.get("tipo"), "elegido")
            o["elegida"] = e.get("opcion")
    return out


def _punto_sin_salida(R, V):
    if R is not None and getattr(R, "sin_cubrir", None):
        a, b = R.sin_cubrir[0]; return ((a[0] + b[0]) / 2, (a[1] + b[1]) / 2)
    for d in (V or {}).get("descargas", []):
        if d["estado"] != "SALE POR GRAVEDAD" and not d.get("alternativa") and R is not None and d["D"] in R.descargas: return R.descargas[d["D"]]["q"]
    return None


def _centro_techo(TT, techo):
    ms = [m for m in (TT or {}).get("montantes", []) if m["techo"] == techo]
    if not ms: return None
    return (sum(m["x"] for m in ms) / len(ms), sum(m["y"] for m in ms) / len(ms))


def aplicar_en_cfg(cfg, soluciones):
    """Aplica al proyecto los cambios de las soluciones automaticas elegidas en el visor. Devuelve {obs: {"opcion", "tipo"}}."""
    elegidas = {}
    for s in soluciones or ():
        elegidas[s["obs"]] = {"opcion": s.get("opcion"), "tipo": s.get("tipo")}
        c = s.get("cambio") or {}
        if c.get("subir_nivel"):
            dz = float(c["subir_nivel"]); nv = cfg.setdefault("niveles", {})
            for k in ("npt", "borde_cuneta", "area_verde"):
                if isinstance(nv.get(k), (int, float)): nv[k] = round(nv[k] + dz, 3)
            for z in nv.get("zonas", []):
                if isinstance(z.get("nivel"), (int, float)): z["nivel"] = round(z["nivel"] + dz, 3)
        if c.get("quitar_descargas_fijas"): cfg.pop("descargas", None)
        if "descarga_alternativa" in c:
            da = dict(cfg.get("descarga_alternativa") or {}, **c["descarga_alternativa"])
            if da.get("tipo"): cfg["descarga_alternativa"] = da
            else:
                cfg.pop("descarga_alternativa", None); cfg["descarga_alternativa_auto"] = False
    return elegidas


def hoja_excel(ruta, obs):
    """Agrega la hoja SOLUCIONES al Excel de la revision general."""
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = load_workbook(ruta)
    if "SOLUCIONES" in wb.sheetnames: del wb["SOLUCIONES"]
    ws = wb.create_sheet("SOLUCIONES")
    ws["A1"] = "OBSERVACIONES Y SOLUCIONES PROPUESTAS"; ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Lo que el programa no resolvio solo o que conviene confirmar, con las soluciones posibles (elegirlas en la ventana de soluciones del visor)."
    cab = ["Observacion", "Tipo", "Que paso", "Por que", "Donde (E, N)", "Soluciones propuestas", "Estado"]
    for j, h in enumerate(cab, 1):
        c = ws.cell(4, j, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    tipo = {"decision": "Decision del proyectista", "revisar": "Revisar", "confirmar": "Confirmar"}
    for i, o in enumerate(obs, 5):
        sol = "\n".join("%s) %s%s - %s" % (p["id"], p["texto"], " (recomendada)" if p.get("recomendada") else "", p["efecto"]) for p in o["opciones"])
        fila = [o["titulo"], tipo.get(o["gravedad"], o["gravedad"]), o["que_paso"], o["por_que"],
                ("%.2f, %.2f" % (o["donde"]["x"], o["donde"]["y"])) if o["donde"] else "", sol, o["estado"]]
        for j, v in enumerate(fila, 1):
            c = ws.cell(i, j, v); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFG", (40, 18, 50, 40, 22, 70, 22)): ws.column_dimensions[col].width = w
    wb.properties.creator = ""; wb.properties.lastModifiedBy = ""
    wb.save(ruta)
    return ruta
