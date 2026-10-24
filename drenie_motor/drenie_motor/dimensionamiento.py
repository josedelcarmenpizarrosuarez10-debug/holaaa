"""Hoja de dimensionamiento (DIMENSIONAMIENTO.xlsx): lo que propone el programa y lo que adopta el proyectista.

- POR TRAMO: area y caudal de techos, pisos y areas verdes, pendiente, ancho que pide el tramo, ancho del sector, tirante, borde
  libre, velocidad, ancho total y excavacion.
- DIMENSIONAMIENTO: una fila por eje con su sector (descarga), el ancho propuesto (una sola medida por sector) y la columna
  B ADOPTADO (amarilla) que llena el proyectista. Con "dimensionamiento_adoptado": "DIMENSIONAMIENTO.xlsx" en el proyecto, el
  motor toma esos anchos y los verifica.
Medidas: libre b x h; ancho total = b + 2 muros; excavacion = h + losa de fondo + solado (0.30 x 0.30 libre -> 0.50 x 0.50).
"""
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HOJA = "DIMENSIONAMIENTO"


def escribir(ruta, filas, secciones, P):
    wb = openpyxl.Workbook()
    neg = Font(bold=True); cab = PatternFill("solid", fgColor="D9E2F3"); ama = PatternFill("solid", fgColor="FFF2CC")
    fino = Side(style="thin", color="999999"); borde = Border(left=fino, right=fino, top=fino, bottom=fino)

    def tabla(ws, fila, encabezado, datos, formatos, editable=None):
        for j, h in enumerate(encabezado):
            c = ws.cell(fila, 1 + j, h); c.font = neg; c.fill = cab; c.border = borde
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[fila].height = 42
        for i, f in enumerate(datos):
            for j, v in enumerate(f):
                c = ws.cell(fila + 1 + i, 1 + j, v); c.border = borde
                if formatos[j]: c.number_format = formatos[j]
                if editable is not None and j == editable: c.fill = ama; c.font = neg
        for j in range(len(encabezado)): ws.column_dimensions[openpyxl.utils.get_column_letter(1 + j)].width = 11
    ws = wb.active; ws.title = HOJA
    ws["A1"] = "DIMENSIONAMIENTO DE CUNETAS: PROPUESTO Y ADOPTADO"; ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Medidas libres interiores (m). Muros y losa de %.2f; ancho total = b + %.2f; excavación = h + %.2f (losa de fondo) + "
                "%.2f (solado). Una sola medida por sector (ejes que salen por la misma descarga). Llenar B ADOPTADO (amarillo) y "
                "volver a diseñar con \"dimensionamiento_adoptado\" en el proyecto." % (P["e_muro"], 2 * P["e_muro"], P["e_fondo"], P["e_solado"]))
    ws["A2"].alignment = Alignment(wrap_text=True); ws.merge_cells("A2:L2"); ws.row_dimensions[2].height = 44
    datos = [[s["sector"], s["eje"], s["L_m"], s["Q_max_ls"], s["b_pide"], s["b"], s["b_adoptado"], s["ancho_total"], s["h_min"], s["h_max"],
              s["excavacion_max"], s["ancho_por"]] for s in secciones]
    tabla(ws, 4, ["SECTOR (descarga)", "EJE", "L (m)", "Q máx (l/s)", "b que pide el eje (m)", "b PROPUESTO (m)", "B ADOPTADO (m)",
                  "Ancho total (m)", "h libre mín (m)", "h libre máx (m)", "Excavación máx (m)", "Criterio"],
          datos, [None, None, "0.00", "0.0", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", "0.00", None], editable=6)
    ws.column_dimensions["L"].width = 34
    w = wb.create_sheet("POR TRAMO")
    w["A1"] = "CAUDAL Y SECCIÓN POR TRAMO"; w["A1"].font = Font(bold=True, size=12)
    claves = [("sector", "Sector", None), ("eje", "Eje", None), ("de", "De", None), ("a", "A", None), ("L", "L (m)", "0.00"),
              ("A_techo_m2", "Área techos (m²)", "0"), ("A_pisos_m2", "Área pisos (m²)", "0"), ("A_verde_m2", "Área verde (m²)", "0"),
              ("Q_techo_ls", "Q techos (l/s)", "0.00"), ("Q_pisos_ls", "Q pisos (l/s)", "0.00"), ("Q_verde_ls", "Q área verde (l/s)", "0.00"),
              ("Q_ls", "Q total (l/s)", "0.00"), ("S_pct", "S (%)", "0.00"), ("b_pide", "b que pide (m)", "0.00"), ("b", "b (m)", "0.00"),
              ("h_de", "h libre inicio (m)", "0.000"), ("h_a", "h libre fin (m)", "0.000"), ("y", "Tirante (m)", "0.000"),
              ("borde_libre", "Borde libre (m)", "0.000"), ("V", "V (m/s)", "0.00"), ("ancho_total", "Ancho total (m)", "0.00"),
              ("excavacion_max", "Excavación máx (m)", "0.00"), ("cap", "Capacidad", None), ("vel", "Velocidad", None)]
    tabla(w, 3, [k[1] for k in claves], [[f[k[0]] for k in claves] for f in filas], [k[2] for k in claves])
    wb.properties.creator = ""; wb.properties.lastModifiedBy = ""
    wb.save(ruta)
    return ruta


def leer_adoptado(ruta):
    """{eje: b adoptado} de la hoja DIMENSIONAMIENTO (filas con B ADOPTADO lleno)."""
    ws = openpyxl.load_workbook(ruta, data_only=True)[HOJA]
    out = {}
    for fila in ws.iter_rows(min_row=5, values_only=True):
        if not fila or not fila[1]: continue
        try: b = float(fila[6])
        except (TypeError, ValueError): continue
        if b > 0: out[str(fila[1])] = b
    return out
