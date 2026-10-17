# -*- coding: utf-8 -*-
"""Motor del Plan de Seguridad y Salud en el Trabajo: lee DATOS_PROYECTO.xlsx, reconoce las actividades de la obra en las
partidas del presupuesto, arma el IPERC (R.M. N.° 050-2013-TR) y los programas del plan. Escribe entregables/_calc/psst.json.

Uso:  python3 nucleo.py <carpeta del proyecto>      (la carpeta contiene insumos/DATOS_PROYECTO.xlsx)"""
import collections, json, math, os, re, sys
from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
import biblioteca as B

NIVELES = [("Trivial", 4, 4), ("Tolerable", 5, 8), ("Moderado", 9, 16), ("Importante", 17, 24), ("Intolerable", 25, 36)]
INTERPRETACION = {
    "Trivial": "No se necesita adoptar ninguna acción.",
    "Tolerable": "No se necesita mejorar la acción preventiva; se consideran soluciones más rentables o mejoras sin carga económica importante y se "
                 "hacen comprobaciones periódicas para asegurar que se mantiene la eficacia de los controles.",
    "Moderado": "Se deben hacer esfuerzos para reducir el riesgo, con medidas implementadas en un periodo determinado. Cuando está asociado a "
                "consecuencias extremadamente dañinas, se precisa una acción posterior para establecer con más precisión la probabilidad de daño.",
    "Importante": "No debe comenzarse el trabajo hasta que se haya reducido el riesgo. Puede que se precisen recursos considerables. Cuando el riesgo "
                  "corresponde a un trabajo que se está realizando, debe remediarse el problema en un tiempo inferior al de los riesgos moderados.",
    "Intolerable": "No se debe comenzar ni continuar el trabajo hasta que se reduzca el riesgo. Si no es posible reducirlo, incluso con recursos "
                   "ilimitados, debe prohibirse el trabajo.",
}
COLOR = {"Trivial": "A9D08E", "Tolerable": "C6E0B4", "Moderado": "FFD966", "Importante": "F4B183", "Intolerable": "FF7C80"}
SIGNIFICATIVO = ("Importante", "Intolerable")
MESES_FREQ = {"al ingreso": 1, "mensual": 1, "trimestral": 3, "semestral": 6, "anual": 12}


def nivel(nr):
    for n, a, b in NIVELES:
        if a <= nr <= b: return n
    return "Intolerable"


def indice_a(n):
    return 1 if n <= 3 else (2 if n <= 12 else 3)


def leer(ruta):
    wb = load_workbook(ruta, data_only=True)
    pr = {}
    for f in wb["PROYECTO"].iter_rows(min_row=2, values_only=True):
        if f[0]: pr[str(f[0]).strip()] = "" if f[2] is None else str(f[2]).strip()
    partidas, titulos = [], {}
    for f in wb["PARTIDAS"].iter_rows(min_row=2, values_only=True):
        if not f[0] or not f[1]: continue
        it = str(f[0]).strip(); desc = re.sub(r"\s+", " ", str(f[1])).strip()
        if f[2]:
            padres = [titulos[k] for k in sorted(titulos, key=len) if it.startswith(k + ".")]
            partidas.append(dict(item=it, desc=desc, und=str(f[2]).strip(), metrado=f[3] or 0, titulos=padres))
        else:
            titulos[it] = desc
    personal = [(int(f[0]), int(f[1] or 0), f[2] or "") for f in wb["PERSONAL"].iter_rows(min_row=2, values_only=True) if f[0]]
    emerg = [list(f[:3]) for f in wb["EMERGENCIA"].iter_rows(min_row=2, values_only=True) if f[0]]
    return pr, partidas, titulos, personal, emerg


def construir(carpeta):
    ins = os.path.join(carpeta, "insumos", "DATOS_PROYECTO.xlsx")
    pr, partidas, titulos, personal, emerg = leer(ins)
    cond = {c.strip().lower() for c in pr.get("condiciones", "").split(",") if c.strip()}
    try:
        alt = max(float(x) for x in re.findall(r"\d[\d,]*", pr.get("altitud_msnm", "0").replace(",", "")))
    except ValueError:
        alt = 0
    if alt > 2500: cond.add("altura")
    pico = max([p[1] for p in personal] or [10])

    # --- reconocimiento de actividades en las partidas
    asig = collections.OrderedDict((a["cod"], []) for a in B.ACT)
    sin_act, sst = [], []
    for p in partidas:
        if B.SST_PARTIDA.search(p["desc"]): sst.append(p)
        if B.NO_FISICAS.search(p["desc"]): continue
        hit = False
        for a in B.ACT:
            if (a["rx"] and a["rx"].search(p["desc"])) or (a.get("rx_tit") and any(a["rx_tit"].search(t) for t in p["titulos"][1:])):
                asig[a["cod"]].append(p); hit = True
        if not hit: sin_act.append(p)
    maquinaria = any(asig[c] for c in ("MTMQ", "ASFA", "TRSP", "ROCA"))
    if maquinaria: cond.add("maquinaria")
    activas = []
    for a in B.ACT:
        if asig[a["cod"]] or (a["siempre"] and (a["condicion"] is None or a["condicion"] in cond)) or \
                (a["condicion"] == "maquinaria" and "maquinaria" in cond):
            activas.append(a)

    # --- IPERC
    filas, n = [], 0
    for a in activas:
        ps = asig[a["cod"]]
        comp = sorted({p["titulos"][1] if len(p["titulos"]) > 1 else (p["titulos"][0] if p["titulos"] else "") for p in ps} - {""})
        A = indice_a(min(a["cuadrilla"], pico))
        for z in a["peligros"]:
            n += 1
            B0, C0 = 2, 2; Br, Cr = 1, 1
            ip = A + B0 + C0 + z["D"]; nr = ip * z["S"]
            ipr = A + Br + Cr + z["Dr"]; nrr = ipr * z["Sr"]
            filas.append(dict(n=n, cod=a["cod"], proceso=a["proceso"], actividad=a["nombre"], componentes=comp,
                              items=[p["item"] for p in ps], n_partidas=len(ps), puestos=a["puestos"], rutina=z["rutina"],
                              peligro=z["peligro"], riesgo=z["riesgo"], consecuencia=z["consecuencia"], tipo=z["tipo"],
                              A=A, B=B0, C=C0, D=z["D"], IP=ip, S=z["S"], NR=nr, nivel=nivel(nr), sig="Sí" if nivel(nr) in SIGNIFICATIVO else "No",
                              elim=z["elim"], sust=z["sust"], ing=z["ing"], adm=z["adm"], epp=z["epp"],
                              Ar=A, Br=Br, Cr=Cr, Dr=z["Dr"], IPr=ipr, Sr=z["Sr"], NRr=nrr, nivel_r=nivel(nrr),
                              sig_r="Sí" if nivel(nrr) in SIGNIFICATIVO else "No", pets=a["pets"],
                              petar=", ".join(B.PETAR[k].split(" (")[0] for k in a["petar"])))

    # --- EPP por puesto
    epp_puesto = collections.OrderedDict()
    for a in activas:
        for pu in a["puestos"]:
            s = epp_puesto.setdefault(pu, [])
            for e in a["epp"]:
                if e not in s: s.append(e)
    epp_usados = [k for k in B.EPP if any(k in v for v in epp_puesto.values())]

    # --- PETS y PETAR
    pets, vistos = [], set()
    for a in activas:
        if a["pets"] in vistos: continue
        vistos.add(a["pets"]); pets.append(dict(cod="PETS-%02d" % (len(pets) + 1), nombre=a["pets"], actividad=a["nombre"],
                                                petar=[B.PETAR[k] for k in a["petar"]]))
    petar = collections.OrderedDict()
    for a in activas:
        for k in a["petar"]:
            petar.setdefault(k, dict(tipo=B.PETAR[k], actividades=[]))["actividades"].append(a["nombre"])
    # PETAR de explosivos solo si hay voladura explícita en el presupuesto
    # (el corte en roca puede hacerse con martillo hidráulico; queda condicionado al procedimiento constructivo)

    # --- capacitación
    meses = max(1, math.ceil(int(pr.get("plazo_dias", "30") or 30) / 30))
    cap = []
    for t, dirigido, h, fr in B.CAPACITACION_BASE:
        cap.append(dict(tema=t, dirigido=dirigido, horas=h, frecuencia=fr))
    for k in petar:
        t, d, h = B.CAPACITACION_PETAR[k]; cap.append(dict(tema=t, dirigido=d, horas=h, frecuencia="semestral"))
    if "selva" in cond:
        cap.append(dict(tema="Prevención de mordeduras de ofidios, picaduras y enfermedades metaxénicas", dirigido="Todo el personal", horas=1, frecuencia="semestral"))
    if "maquinaria" in cond:
        cap.append(dict(tema="Operación segura de maquinaria pesada, puntos ciegos y señales con el vigía", dirigido="Operadores, choferes y vigías", horas=2, frecuencia="semestral"))
    if "transito" in cond or asig["TRAN"]:
        cap.append(dict(tema="Control de tránsito en zona de trabajo (Manual de Dispositivos de Control del Tránsito del MTC)", dirigido="Controladores de tránsito y capataces", horas=2, frecuencia="semestral"))
    for c in cap:
        f = MESES_FREQ.get(c["frecuencia"], 6)
        c["meses"] = list(range(1, meses + 1)) if c["frecuencia"] == "al ingreso" else list(range(1, meses + 1, f))
        c["veces"] = len(c["meses"])

    # --- inspecciones
    insp = [("Pre-uso de maquinaria y vehículos (check-list)", "Operador / chofer", "Diaria"),
            ("Frentes de trabajo: orden, señalización y cumplimiento del ATS", "Supervisor de SST / capataz", "Diaria"),
            ("Herramientas manuales y eléctricas (cinta de color del mes)", "Capataz", "Mensual"),
            ("Equipos de protección personal", "Supervisor de SST", "Mensual"),
            ("Extintores, botiquines y camilla", "Supervisor de SST / brigadas", "Mensual"),
            ("Campamento, almacenes, comedor y servicios higiénicos", "Supervisor de SST / administración", "Mensual"),
            ("Instalaciones eléctricas provisionales y tableros", "Electricista", "Mensual"),
            ("Inspección planeada general de la obra con el comité", "Comité o supervisor de SST", "Mensual"),
            ("Almacén de combustibles y materiales peligrosos", "Almacenero / supervisor de SST", "Semanal")]
    if "altura" in petar: insp.append(("Andamios, escaleras, líneas de vida y arneses", "Supervisor de SST", "Diaria (tarjeta) / semanal"))
    if "excavacion" in petar: insp.append(("Excavaciones, taludes y entibados", "Persona competente designada", "Diaria y después de lluvias"))
    if "izaje" in petar: insp.append(("Grúas, aparejos y accesorios de izaje", "Rigger / operador", "Antes de cada izaje"))
    if "cauces" in petar: insp.append(("Equipos de rescate acuático y nivel de los cauces", "Supervisor de SST", "Diaria en temporada de lluvias"))
    if asig["TRAN"] or "transito" in cond: insp.append(("Dispositivos de control de tránsito y desvíos", "Supervisor de SST / responsable de tránsito", "Diaria"))

    # --- monitoreo de agentes ocupacionales (según los peligros de salud del IPERC)
    so = " ".join(f["peligro"] + " " + f["riesgo"] for f in filas if f["tipo"] == "SO").lower()
    moni = []
    for clave, agente, metodo, frec in [("ruido", "Ruido", "Dosimetría de ruido por puesto (TLV 85 dB(A) para 8 h)", "Semestral"),
                                        ("polvo", "Polvo inhalable y respirable", "Muestreo gravimétrico personal", "Semestral"),
                                        ("sílice", "Sílice cristalina respirable", "Muestreo personal y análisis de laboratorio", "Semestral"),
                                        ("vibración", "Vibración mano-brazo y cuerpo entero", "Medición con acelerómetro (ISO 5349 / ISO 2631)", "Anual"),
                                        ("calor", "Estrés térmico", "Índice WBGT en frentes a campo abierto", "En época de mayor temperatura"),
                                        ("vapores", "Vapores orgánicos y humos de asfalto", "Muestreo personal con tubos adsorbentes", "Semestral"),
                                        ("humos metálicos", "Humos metálicos de soldadura", "Muestreo personal", "Semestral"),
                                        ("monóxido", "Monóxido de carbono", "Detector de gases de lectura directa", "Antes de cada tarea con motobombas"),
                                        ("radiación", "Radiación ultravioleta solar", "Índice UV del SENAMHI y control de exposición", "Diaria"),
                                        ("postura", "Factores ergonómicos (manipulación de cargas y posturas)", "Evaluación ergonómica (R.M. N.° 375-2008-TR)", "Anual")]:
        if clave in so or (clave == "postura" and "sobreesfuerzo" in so) or (clave == "sílice" and "sílice" in so):
            moni.append(dict(agente=agente, metodo=metodo, frecuencia=frec))

    # --- personal y horas hombre
    hht = sum(p[1] for p in personal) * 8 * 26
    org = "comite" if pico >= 20 else "supervisor"
    sig_i = [f for f in filas if f["sig"] == "Sí"]; sig_r = [f for f in filas if f["sig_r"] == "Sí"]
    out = dict(proyecto=pr, condiciones=sorted(cond), personal=personal, pico=pico, hht=hht, organizacion=org, meses=meses,
               emergencia=emerg, actividades=[dict(cod=a["cod"], proceso=a["proceso"], nombre=a["nombre"], puestos=a["puestos"],
                                                    partidas=[p["item"] + " " + p["desc"] for p in asig[a["cod"]]], petar=a["petar"], pets=a["pets"])
                                               for a in activas],
               iperc=filas, epp_puesto=epp_puesto, epp_catalogo={k: B.EPP[k] for k in epp_usados}, epp_corto={k: B.EPP_CORTO[k] for k in epp_usados}, pets=pets, petar=petar,
               capacitacion=cap, inspecciones=[dict(tema=a, responsable=b, frecuencia=c) for a, b, c in insp], monitoreo=moni,
               partidas_sst=[dict(item=p["item"], desc=p["desc"], und=p["und"], metrado=p["metrado"]) for p in sst],
               sin_actividad=[p["item"] + " " + p["desc"] for p in sin_act],
               resumen=dict(n=len(filas), inh=collections.Counter(f["nivel"] for f in filas), res=collections.Counter(f["nivel_r"] for f in filas),
                            sig_i=len(sig_i), sig_r=len(sig_r), n_act=len(activas), n_part=len(partidas)),
               niveles=NIVELES, interpretacion=INTERPRETACION, color=COLOR)
    calc = os.path.join(carpeta, "entregables", "_calc"); os.makedirs(calc, exist_ok=True)
    json.dump(out, open(os.path.join(calc, "psst.json"), "w"), ensure_ascii=False, indent=1)
    return out


if __name__ == "__main__":
    o = construir(sys.argv[1])
    r = o["resumen"]
    print("actividades", r["n_act"], "peligros", r["n"], "inherente", dict(r["inh"]), "residual", dict(r["res"]),
          "significativos", r["sig_i"], "->", r["sig_r"])
    print("partidas sin actividad:", len(o["sin_actividad"]))
    for s in o["sin_actividad"]: print("   ", s)
    print("PETAR:", list(o["petar"]))
