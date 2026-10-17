# -*- coding: utf-8 -*-
"""Word del PLAN DE SEGURIDAD Y SALUD EN EL TRABAJO (PSST) de una obra, con el IPERC, los programas y los anexos.
Estructura según la Norma G.050 (contenido mínimo del plan) y el D.S. N.° 011-2019-TR; IPERC según la R.M. N.° 050-2013-TR.

Uso:  python3 psst_docx.py <carpeta del proyecto>      (antes: nucleo.py, iperc_xlsx.py y figuras.py)"""
import collections, json, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

AQUI = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, AQUI)
from docx_util import (fuente, par, vin, tit, tabla, ctab, fig, seccion, pie, encabezado, recuadro, banda, CONT, NUMT, NUMF, AZUL)
import biblioteca as B

C = WD_ALIGN_PARAGRAPH.CENTER; L_ = WD_ALIGN_PARAGRAPH.LEFT
N = "No aplica"


def colnivel(D):
    return lambda i, j, v: D["color"].get(v)


def main(carpeta):
    D = json.load(open(os.path.join(carpeta, "entregables", "_calc", "psst.json")))
    P = D["proyecto"]; R = D["iperc"]; RES = D["resumen"]; FIG = os.path.join(carpeta, "entregables", "_calc", "fig")
    cod = P.get("codigo", "OBRA"); cui = P.get("cui", "")
    sal = os.path.join(carpeta, "entregables", "PLAN_SEGURIDAD_SALUD_TRABAJO_%s.docx" % cod)
    xls = "MATRIZ_IPERC_PSST_%s.xlsx" % cod
    cond = set(D["condiciones"]); petar = D["petar"]; acts = {a["cod"] for a in D["actividades"]}
    vial = P.get("tipo_obra", "").lower().startswith("vial") or "TRAN" in acts
    org = D["organizacion"]; pico = D["pico"]
    cn = colnivel(D)
    corto = P.get("nombre", "")[:90] + ("…" if len(P.get("nombre", "")) > 90 else "")

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(10)
    s0 = doc.sections[0]; s0.page_width, s0.page_height = Cm(21.0), Cm(29.7)
    s0.left_margin = s0.right_margin = Cm(2.5); s0.top_margin = Cm(2.5); s0.bottom_margin = Cm(2.0)
    s0.different_first_page_header_footer = True              # portada sin encabezado ni pie

    # ================================================================ portada
    par(doc, P.get("encabezado_1", ""), 13, True, C, antes=6, despues=2, color=AZUL)
    par(doc, P.get("encabezado_2", ""), 11, True, C, despues=36, color=AZUL)
    banda(doc, "PLAN DE SEGURIDAD Y SALUD EN EL TRABAJO", tam=20)
    par(doc, "Identificación de peligros, evaluación de riesgos y medidas de control (IPERC), programas y plan de respuesta ante emergencias",
        10.5, False, C, antes=8, despues=26, cur=True)
    par(doc, "OBRA:", 10, True, C, despues=4)
    par(doc, "“" + P.get("nombre", "") + "”", 12, True, C, despues=8)
    par(doc, "CUI N.° %s" % cui, 12, True, C, despues=30)
    tabla(doc, ["DATO", "DESCRIPCIÓN"], [["Ubicación", P.get("ubicacion", "")], ["Plazo de ejecución", "%s días calendario" % P.get("plazo_dias", "")],
                                        ["Personal máximo previsto", "%d trabajadores" % pico], ["Modalidad", P.get("modalidad", "")],
                                        ["Ejecutor", P.get("ejecutor") or "________________________________"]], [4.5, 11.5], tam=9)
    par(doc, P.get("fecha", "").upper(), 11, True, C, antes=34)
    doc.add_page_break()

    par(doc, "CONTENIDO", 12, True, C, despues=8, color=AZUL)
    marca = doc.add_paragraph()
    doc.add_page_break()

    # ================================================================ 1. generalidades
    tit(doc, "1", "Generalidades")
    tit(doc, "1.1", "Objetivo", 2)
    par(doc, "Establecer las acciones para prevenir accidentes de trabajo, incidentes peligrosos y enfermedades ocupacionales durante la ejecución de la obra "
             "“%s”, CUI N.° %s, mediante la identificación de los peligros de cada actividad, la evaluación de sus riesgos y la aplicación de medidas de "
             "control, de modo que todas las tareas se realicen con un nivel de riesgo aceptable para los trabajadores, el personal de la entidad y de la "
             "supervisión, los visitantes y la población del entorno." % (P.get("nombre", ""), cui))
    par(doc, "Objetivos específicos:", neg=True, despues=3)
    for t in ["Identificar los peligros y evaluar los riesgos de todas las actividades de la obra (IPERC de línea base) y mantener la matriz actualizada.",
              "Aplicar las medidas de control en el orden de la jerarquía: eliminación, sustitución, controles de ingeniería, controles administrativos y equipo de protección personal.",
              "Definir la organización, las responsabilidades y los recursos para la gestión de la seguridad y salud en el trabajo en la obra.",
              "Capacitar a todo el personal en los peligros de su puesto y en los procedimientos de trabajo seguro.",
              "Contar con un plan de respuesta ante emergencias probado mediante simulacros.",
              "Cumplir la normativa nacional de seguridad y salud en el trabajo y medir el desempeño con indicadores."]: vin(doc, t)
    tit(doc, "1.2", "Alcance", 2)
    par(doc, "El plan es de cumplimiento obligatorio para todo el personal que participe en la obra: trabajadores del ejecutor, subcontratistas, proveedores "
             "que ingresen a la obra, personal de la supervisión y de la entidad, y visitantes. Comprende todas las actividades desde la movilización hasta "
             "la desmovilización y el cierre de las áreas auxiliares (campamento, patio de máquinas, canteras y depósitos de material excedente).")
    tit(doc, "1.3", "Datos de la obra", 2)
    ctab(doc, "Datos generales de la obra")
    tabla(doc, ["CONCEPTO", "DESCRIPCIÓN"],
          [["Obra", P.get("nombre", "")], ["CUI", cui], ["Entidad", P.get("entidad", "")], ["Unidad ejecutora", P.get("unidad_ejecutora", "")],
           ["Ubicación", P.get("ubicacion", "")], ["Altitud", "%s msnm" % P.get("altitud_msnm", "") if P.get("altitud_msnm") else "—"],
           ["Clima", P.get("clima", "") or "—"], ["Plazo de ejecución", "%s días calendario (%d meses)" % (P.get("plazo_dias", ""), D["meses"])],
           ["Modalidad", P.get("modalidad", "")], ["Jornada de trabajo", P.get("jornada", "")],
           ["Personal máximo previsto", "%d trabajadores (%s horas hombre en toda la obra)" % (pico, "{:,}".format(D["hht"]))]], [4.5, 11.5], tam=8.5)
    tit(doc, "1.4", "Descripción de la obra", 2)
    par(doc, P.get("descripcion", ""))
    # componentes del presupuesto (títulos de segundo nivel) con su número de partidas
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(carpeta, "insumos", "DATOS_PROYECTO.xlsx"), data_only=True)["PARTIDAS"]
    tit2, cuenta = collections.OrderedDict(), collections.Counter()
    for f in ws.iter_rows(min_row=2, values_only=True):
        if not f[0]: continue
        it = str(f[0]); niv = it.count(".")
        if not f[2] and niv == 1: tit2[it] = f[1]
        elif f[2]:
            k = ".".join(it.split(".")[:2]); cuenta[k] += 1
    ctab(doc, "Componentes del presupuesto de obra")
    tabla(doc, ["ÍTEM", "COMPONENTE", "PARTIDAS"], [[k, v, cuenta[k]] for k, v in tit2.items()] + [["", "TOTAL", sum(cuenta.values())]],
          [1.6, 11.0, 2.4], tam=8.5, alinear=[C, L_, C])
    tit(doc, "1.5", "Personal previsto", 2)
    par(doc, "El número de trabajadores varía con el avance de la obra. El histograma de personal define la cantidad de EPP, la frecuencia de las inducciones, "
             "el tamaño de las brigadas de emergencia y el órgano de SST que corresponde.")
    fig(doc, os.path.join(FIG, "f1_personal.png"), 15.5, "Histograma de personal previsto", "programa de ejecución de la obra (hoja PERSONAL de los datos del proyecto).")

    # ================================================================ 2. marco normativo
    tit(doc, "2", "Marco normativo")
    normas = [("Ley N.° 29783", "Ley de Seguridad y Salud en el Trabajo, y su modificatoria la Ley N.° 30222."),
              ("D.S. N.° 005-2012-TR", "Reglamento de la Ley N.° 29783 y modificatorias: comité o supervisor de SST, registros obligatorios, notificación de accidentes."),
              ("D.S. N.° 011-2019-TR", "Reglamento de Seguridad y Salud en el Trabajo para el Sector Construcción."),
              ("Norma G.050", "Seguridad durante la Construcción, del Reglamento Nacional de Edificaciones: contenido del plan y requisitos por actividad."),
              ("R.M. N.° 050-2013-TR", "Formatos referenciales de los registros obligatorios y guía para la identificación de peligros, evaluación de riesgos y medidas de control."),
              ("R.M. N.° 375-2008-TR", "Norma básica de ergonomía y procedimiento de evaluación de riesgo disergonómico."),
              ("R.M. N.° 312-2011-MINSA", "Protocolos de exámenes médico ocupacionales y guías de diagnóstico, y modificatorias."),
              ("Ley N.° 26790 y D.S. N.° 003-98-SA", "Seguro Complementario de Trabajo de Riesgo (SCTR) para actividades de construcción."),
              ("D.S. N.° 015-2005-SA", "Valores límite permisibles para agentes químicos en el ambiente de trabajo."),
              ("Ley N.° 30102", "Medidas preventivas contra los efectos nocivos para la salud por la exposición prolongada a la radiación solar."),
              ("Ley N.° 27942", "Ley de prevención y sanción del hostigamiento sexual y su reglamento."),
              ("NTP 399.010-1", "Señales de seguridad: colores, símbolos, formas y dimensiones."),
              ("NTP 350.043-1", "Extintores portátiles: selección, distribución, inspección, mantenimiento, recarga y prueba hidrostática."),
              ("Ley N.° 29664", "Sistema Nacional de Gestión del Riesgo de Desastres: respuesta ante emergencias.")]
    if vial: normas.append(("Manual de Dispositivos de Control del Tránsito (MTC)", "Manual de Dispositivos de Control del Tránsito Automotor para Calles y Carreteras: señalización de zonas de trabajo y desvíos."))
    if "electrico" in petar: normas.append(("R.M. N.° 111-2013-MEM/DM", "Reglamento de Seguridad y Salud en el Trabajo con Electricidad; distancias de seguridad del Código Nacional de Electricidad."))
    if "explosivos" in petar: normas.append(("Ley N.° 30299", "Ley de armas de fuego, municiones, explosivos, productos pirotécnicos y materiales relacionados de uso civil (SUCAMEC)."))
    ctab(doc, "Normas aplicables")
    tabla(doc, ["NORMA", "CONTENIDO APLICADO EN EL PLAN"], [list(n) for n in normas], [4.4, 11.6], tam=8)
    par(doc, "Referencias técnicas: ISO 45001:2018, Sistemas de gestión de la seguridad y salud en el trabajo; normas ANSI y ASTM citadas para los equipos de protección.",
        8.5, cur=True)

    # ================================================================ 3. política y objetivos
    tit(doc, "3", "Política, objetivos y metas de seguridad y salud en el trabajo")
    tit(doc, "3.1", "Política de seguridad y salud en el trabajo", 2)
    par(doc, "El ejecutor publica en el campamento y en los frentes de trabajo la siguiente política, firmada por su representante legal y por el residente de obra:")
    recuadro(doc, ["En la ejecución de la obra “%s” nos comprometemos a:" % corto,
                   "1. Proteger la seguridad y la salud de todos los trabajadores, previniendo lesiones, enfermedades e incidentes relacionados con el trabajo.",
                   "2. Cumplir la normativa nacional de seguridad y salud en el trabajo y los compromisos asumidos con la entidad.",
                   "3. Identificar los peligros, evaluar los riesgos y aplicar los controles en el orden de la jerarquía antes de iniciar cada tarea.",
                   "4. Garantizar la consulta y participación de los trabajadores y de sus representantes en todos los elementos del sistema de gestión.",
                   "5. Capacitar y entrenar al personal, y proporcionar sin costo el equipo de protección personal adecuado.",
                   "6. Mejorar de forma continua el desempeño en seguridad y salud en el trabajo e integrarlo a la gestión de la obra."], tam=9.5)
    tit(doc, "3.2", "Objetivos, metas e indicadores", 2)
    ctab(doc, "Objetivos, metas e indicadores del plan")
    tabla(doc, ["OBJETIVO", "META", "INDICADOR", "FRECUENCIA"],
          [["Evitar accidentes mortales", "0", "N.° de accidentes mortales", "Mensual"],
           ["Reducir los accidentes con tiempo perdido", "Índice de frecuencia ≤ 3.0", "Accidentes incapacitantes x 10⁶ / horas hombre trabajadas", "Mensual"],
           ["Reducir la gravedad de los accidentes", "Índice de gravedad ≤ 60", "Días perdidos x 10⁶ / horas hombre trabajadas", "Mensual"],
           ["Capacitar a todo el personal", "100 % del programa", "Capacitaciones ejecutadas / programadas", "Mensual"],
           ["Inspeccionar la obra", "≥ 95 % del programa", "Inspecciones ejecutadas / programadas", "Mensual"],
           ["Controlar los trabajos de alto riesgo", "100 %", "Tareas de alto riesgo con PETAR aprobado / total", "Semanal"],
           ["Asegurar la aptitud y cobertura del personal", "100 %", "Trabajadores con examen médico de ingreso y SCTR / total", "Mensual"],
           ["Atender las condiciones inseguras reportadas", "≥ 90 % en el plazo fijado", "Acciones correctivas cerradas / reportadas", "Mensual"],
           ["Estar preparados ante emergencias", "1 simulacro por semestre como mínimo", "Simulacros ejecutados / programados", "Semestral"]],
          [4.6, 3.0, 5.6, 2.8], tam=8)

    # ================================================================ 4. organización
    tit(doc, "4", "Organización y responsabilidades")
    tit(doc, "4.1", "Organigrama de seguridad y salud en el trabajo", 2)
    fig(doc, os.path.join(FIG, "f5_organigrama.png"), 15.5, "Organigrama de seguridad y salud en el trabajo de la obra", "elaboración propia según el D.S. N.° 011-2019-TR.")
    tit(doc, "4.2", "Funciones y responsabilidades", 2)
    ctab(doc, "Funciones y responsabilidades en seguridad y salud en el trabajo")
    tabla(doc, ["CARGO", "RESPONSABILIDADES"],
          [["Representante del ejecutor", "Aprueba la política y el plan; asigna los recursos; responde ante la entidad y la autoridad por el cumplimiento del plan."],
           ["Residente de obra", "Dirige la aplicación del plan; no autoriza el inicio de tareas sin IPERC, PETS y PETAR cuando corresponden; preside o integra el órgano de SST; "
                                 "participa en la investigación de accidentes."],
           ["Responsable de SST de la obra", "Elabora y actualiza el IPERC y el plan; capacita; inspecciona; verifica los PETAR; lleva los registros obligatorios y las "
                                             "estadísticas; dirige las brigadas y los simulacros; propone medidas correctivas y puede detener trabajos con riesgo inminente."],
           ["Supervisores de SST de frente", "Verifican en campo el ATS, el uso del EPP y las condiciones de trabajo; inspeccionan diariamente; reportan incidentes."],
           ["Ingenieros de frente, maestro de obra y capataces", "Planifican cada tarea con su ATS; dan la charla diaria de 5 minutos; firman los PETAR de su frente; corrigen de inmediato "
                                                                 "los actos y condiciones inseguras."],
           ["Médico ocupacional o enfermería", "Exámenes médicos ocupacionales; vigilancia de la salud; atención de primeros auxilios; control del botiquín y del tópico."],
           ["Comité o supervisor de SST", "Aprueba el programa anual de SST; inspecciona; investiga accidentes; vigila el cumplimiento del plan (Ley N.° 29783)."],
           ["Subcontratistas y proveedores", "Cumplen este plan; presentan su personal con SCTR, examen médico y EPP; asisten a la inducción; reportan incidentes."],
           ["Trabajadores", "Cumplen los procedimientos; usan correctamente el EPP; participan en las capacitaciones; reportan peligros e incidentes; pueden negarse a trabajar "
                            "ante un peligro inminente sin represalias."]], [4.6, 11.4], tam=8)
    tit(doc, "4.3", "Órgano de seguridad y salud en el trabajo", 2)
    if org == "comite":
        par(doc, "Con %d trabajadores en el periodo de mayor actividad, la obra debe contar con un comité de seguridad y salud en el trabajo paritario "
                 "(Ley N.° 29783: empleadores con 20 o más trabajadores), integrado en igual número por representantes del empleador y de los trabajadores, "
                 "estos últimos elegidos por votación. El comité se instala al inicio de la obra, sesiona una vez al mes y de forma extraordinaria ante un "
                 "accidente grave; sus acuerdos se registran en el libro de actas." % pico)
    else:
        par(doc, "Con %d trabajadores como máximo, la obra cuenta con un supervisor de seguridad y salud en el trabajo elegido por los trabajadores "
                 "(Ley N.° 29783: empleadores con menos de 20 trabajadores)." % pico)
    par(doc, "Independientemente del órgano de SST, el ejecutor designa un responsable de seguridad y salud en el trabajo de la obra, con la formación y "
             "experiencia que exigen el D.S. N.° 011-2019-TR y las bases, con dedicación en campo durante toda la ejecución, y los supervisores de "
             "SST de frente que requiera el número de frentes simultáneos.")

    # ================================================================ 5. IPERC
    tit(doc, "5", "Identificación de peligros, evaluación de riesgos y medidas de control (IPERC)")
    tit(doc, "5.1", "Metodología", 2)
    par(doc, "Se aplica el método de la R.M. N.° 050-2013-TR. Para cada actividad de la obra se identifican los peligros (fuentes, situaciones o actos con "
             "potencial de daño) y los riesgos que generan, y se evalúa la probabilidad y la severidad del daño:")
    vin(doc, "Índice de probabilidad = A (personas expuestas) + B (procedimientos existentes) + C (capacitación) + D (exposición al riesgo), con valores de 4 a 12.")
    vin(doc, "Índice de severidad de 1 a 3 (lesión sin incapacidad, con incapacidad temporal, con incapacidad permanente o muerte).")
    vin(doc, "Nivel de riesgo = probabilidad x severidad, de 4 a 36, en cinco niveles. Es significativo el riesgo de nivel importante o intolerable.")
    ctab(doc, "Índices de probabilidad y de severidad")
    tabla(doc, ["ÍNDICE", "A. PERSONAS EXPUESTAS", "B. PROCEDIMIENTOS EXISTENTES", "C. CAPACITACIÓN", "D. EXPOSICIÓN AL RIESGO", "SEVERIDAD"],
          [["1", "1 a 3", "Existen, son satisfactorios y suficientes", "Personal entrenado; conoce el peligro y lo previene",
            "Al menos una vez al año (S); esporádicamente (SO)", "Lesión sin incapacidad (S); disconfort o incomodidad (SO)"],
           ["2", "4 a 12", "Existen parcialmente y no son satisfactorios o suficientes", "Personal parcialmente entrenado; conoce el peligro pero no toma acciones de control",
            "Al menos una vez al mes (S); eventualmente (SO)", "Lesión con incapacidad temporal (S); daño a la salud reversible (SO)"],
           ["3", "Más de 12", "No existen", "Personal no entrenado; no conoce el peligro ni toma acciones de control",
            "Al menos una vez al día (S); permanentemente (SO)", "Lesión con incapacidad permanente (S); daño a la salud irreversible (SO)"]],
          [1.2, 2.0, 3.2, 3.4, 3.0, 3.2], tam=7.5, alinear=[C] + [L_] * 5)
    par(doc, "S: seguridad; SO: salud ocupacional.", 8, cur=True)
    ctab(doc, "Niveles de riesgo y su interpretación")
    tabla(doc, ["NIVEL", "P x S", "INTERPRETACIÓN"],
          [[n, "%d" % a if a == b else "%d a %d" % (a, b), D["interpretacion"][n]] for n, a, b in D["niveles"]], [2.6, 1.8, 11.6], tam=8,
          colores=lambda i, j, v: D["color"].get(v) if j == 0 else None, alinear=[C, C, L_])
    par(doc, "La evaluación inicial considera los procedimientos y la capacitación como parciales (B = 2, C = 2). La evaluación residual supone implementados "
             "los PETS, el ATS, la capacitación y la supervisión (B = 1, C = 1), y además la reducción de la exposición o de la severidad que logran los "
             "controles de ingeniería. El índice A se obtiene del tamaño de la cuadrilla de cada actividad.")
    tit(doc, "5.2", "Actividades evaluadas", 2)
    par(doc, "Las actividades se obtuvieron de las %d partidas del presupuesto de obra, agrupadas por el tipo de trabajo que se realiza, más las actividades "
             "transversales a toda la obra (trabajo a campo abierto, traslado de personal y mantenimiento de equipos). La relación partida por partida está en "
             "la hoja PARTIDAS del libro %s." % (RES["n_part"], xls))
    ctab(doc, "Actividades evaluadas en el IPERC")
    tabla(doc, ["COD.", "PROCESO", "ACTIVIDAD", "PARTIDAS", "PUESTOS EXPUESTOS"],
          [[a["cod"], a["proceso"], a["nombre"], len(a["partidas"]) or "Transversal", ", ".join(a["puestos"])] for a in D["actividades"]],
          [1.4, 3.0, 5.6, 1.8, 4.2], tam=7.5, alinear=[C, L_, L_, C, L_])
    tit(doc, "5.3", "Resultados", 2)
    ci, cr = RES["inh"], RES["res"]
    par(doc, "Se evaluaron %d peligros en %d actividades. Antes de aplicar los controles, %d riesgos son significativos (%d intolerables y %d importantes). "
             "Con los controles propuestos ningún riesgo queda en nivel importante o intolerable: %d quedan en nivel moderado y %d en tolerable o trivial. "
             "Los riesgos moderados residuales requieren supervisión permanente y verificación de que los controles se mantienen." %
        (RES["n"], RES["n_act"], RES["sig_i"], ci.get("Intolerable", 0), ci.get("Importante", 0), cr.get("Moderado", 0),
         cr.get("Tolerable", 0) + cr.get("Trivial", 0)) if RES["sig_r"] == 0 else
        "Se evaluaron %d peligros en %d actividades. Antes de los controles hay %d riesgos significativos; después de los controles quedan %d, que no "
        "pueden iniciarse sin reforzar las medidas." % (RES["n"], RES["n_act"], RES["sig_i"], RES["sig_r"]))
    fig(doc, os.path.join(FIG, "f3_niveles.png"), 15.0, "Número de peligros por nivel de riesgo, antes y después de los controles", "matriz IPERC (%s)." % xls)
    fig(doc, os.path.join(FIG, "f2_calor.png"), 15.5, "Mapa de calor: probabilidad x severidad antes y después de los controles", "matriz IPERC (%s)." % xls)
    procs = list(dict.fromkeys(f["proceso"] for f in R)); niv = [n for n, a, b in D["niveles"]]
    ctab(doc, "Resumen de niveles de riesgo por proceso")
    filas = []
    for p in procs:
        fi = [f for f in R if f["proceso"] == p]
        filas.append([p, len(fi)] + [sum(1 for f in fi if f["nivel"] == n) or "" for n in niv[2:]] + [sum(1 for f in fi if f["nivel_r"] == n) or "" for n in niv[1:3]])
    filas.append(["TOTAL", len(R)] + [ci.get(n, 0) for n in niv[2:]] + [cr.get(n, 0) for n in niv[1:3]])
    tabla(doc, ["PROCESO", "PELIGROS", "INICIAL\nMODERADO", "INICIAL\nIMPORTANTE", "INICIAL\nINTOLERABLE", "RESIDUAL\nTOLERABLE", "RESIDUAL\nMODERADO"], filas,
          [4.6, 1.6, 2.0, 2.0, 2.0, 2.0, 2.0], tam=7.5, alinear=[L_] + [C] * 6)
    fig(doc, os.path.join(FIG, "f4_procesos.png"), 14.5, "Peligros evaluados por proceso y tipo", "matriz IPERC (%s)." % xls)
    tit(doc, "5.4", "Riesgos intolerables antes de los controles", 2)
    par(doc, "Son los peligros que exigen controles antes de iniciar cualquier trabajo. Todos tienen procedimiento escrito y, cuando corresponde, permiso de "
             "trabajo de alto riesgo:")
    top = sorted([f for f in R if f["nivel"] == "Intolerable"], key=lambda f: -f["NR"])
    ctab(doc, "Riesgos intolerables y sus controles principales")
    tabla(doc, ["N.°", "ACTIVIDAD", "PELIGRO / RIESGO", "P x S", "CONTROLES PRINCIPALES", "RESIDUAL"],
          [[f["n"], f["actividad"], "%s: %s" % (f["peligro"], f["riesgo"].lower()), f["NR"],
            "; ".join(x for x in (f["ing"], f["adm"]) if x != N), "%s (%d)" % (f["nivel_r"], f["NRr"])] for f in top],
          [0.9, 3.2, 4.0, 1.1, 5.0, 1.8], tam=7, colores=lambda i, j, v: D["color"].get(str(v).split(" (")[0]) if j == 5 else None,
          alinear=[C, L_, L_, C, L_, C])
    tit(doc, "5.5", "Jerarquía de controles aplicada", 2)
    par(doc, "Los controles se eligieron en el orden que fija el artículo 21 de la Ley N.° 29783. La eliminación del peligro rara vez es posible en "
             "obra; por eso el plan se apoya en la sustitución (equipos y productos menos peligrosos), en los controles de ingeniería (barandas, "
             "entibados, guardas, delimitación, dispositivos de tránsito) y en los controles administrativos (procedimientos, permisos, capacitación, "
             "supervisión). El EPP es la última barrera y nunca la única.")
    fig(doc, os.path.join(FIG, "f6_jerarquia.png"), 14.5, "Jerarquía de controles y número de peligros en que se aplica cada nivel", "matriz IPERC (%s)." % xls)
    tit(doc, "5.6", "Actualización del IPERC", 2)
    par(doc, "El IPERC de línea base (Anexo 1 y libro %s) se actualiza: al iniciar una actividad no prevista o cambiar un método constructivo, equipo o "
             "material; después de todo accidente o incidente peligroso; cuando lo disponga el comité de SST; y al menos una vez al año. En el frente, cada "
             "cuadrilla elabora el análisis de trabajo seguro (ATS) antes de empezar la tarea, con los peligros específicos del día." % xls)

    # ================================================================ 6. mapa de riesgos y señalización
    tit(doc, "6", "Mapa de riesgos y señalización")
    par(doc, "Al instalar el campamento y cada frente, el responsable de SST elabora el mapa de riesgos con la ubicación de los peligros significativos, las "
             "zonas de trabajo de alto riesgo, las vías de circulación de equipos y peatones, los equipos de emergencia, las zonas seguras y las rutas de "
             "evacuación. El mapa se publica en el campamento y se actualiza con el avance de la obra.")
    ctab(doc, "Señalización de seguridad (NTP 399.010-1)")
    tabla(doc, ["TIPO", "COLOR Y FORMA", "EJEMPLOS EN LA OBRA"],
          [["Prohibición", "Círculo rojo con banda diagonal, fondo blanco", "Prohibido el ingreso a personal no autorizado; prohibido fumar; prohibido pasar bajo carga suspendida"],
           ["Obligación", "Círculo azul, pictograma blanco", "Uso obligatorio de casco, lentes, protector auditivo, arnés, chaleco reflectivo"],
           ["Advertencia", "Triángulo amarillo con borde negro", "Excavación profunda; maquinaria pesada en movimiento; riesgo eléctrico; caída a distinto nivel; material inflamable"],
           ["Emergencia y evacuación", "Rectángulo verde, pictograma blanco", "Zona segura; salida; botiquín; punto de reunión; ducha y lavaojos"],
           ["Equipos contra incendio", "Rectángulo rojo, pictograma blanco", "Extintor; alarma"]] +
          ([["Control de tránsito", "Según el Manual de Dispositivos del MTC", "Señales preventivas de obra, conos, barreras, tranqueras, paletas PARE/SIGA, luces intermitentes"]] if vial else []),
          [3.0, 4.6, 8.4], tam=8)

    # ================================================================ 7. procedimientos
    tit(doc, "7", "Procedimientos de trabajo")
    tit(doc, "7.1", "Procedimientos escritos de trabajo seguro (PETS)", 2)
    par(doc, "Cada actividad del IPERC cuenta con su PETS, que describe paso a paso cómo hacer la tarea, el personal, los equipos, los materiales, el EPP y "
             "las restricciones. Los PETS se elaboran con participación de los trabajadores, los aprueba el residente y se difunden antes de iniciar la actividad.")
    ctab(doc, "Procedimientos escritos de trabajo seguro")
    tabla(doc, ["CÓDIGO", "PROCEDIMIENTO", "PETAR ASOCIADO"], [[p["cod"], p["nombre"], "; ".join(x.split(" (")[0] for x in p["petar"]) or "—"] for p in D["pets"]],
          [1.8, 9.2, 5.0], tam=7.5, alinear=[C, L_, L_])
    tit(doc, "7.2", "Permisos de trabajo de alto riesgo (PETAR)", 2)
    par(doc, "Las siguientes tareas solo se inician con el PETAR firmado por el supervisor del frente, el responsable de SST y el residente, después de "
             "verificar en campo las condiciones y los controles. El permiso es válido para la tarea y la jornada indicadas y se archiva en obra:")
    ctab(doc, "Trabajos de alto riesgo identificados")
    tabla(doc, ["TRABAJO DE ALTO RIESGO", "ACTIVIDADES DE LA OBRA"], [[v["tipo"], "; ".join(v["actividades"])] for v in petar.values()], [5.6, 10.4], tam=8)
    if "explosivos" in petar:
        par(doc, "El uso de explosivos en el corte de roca queda condicionado al procedimiento constructivo aprobado; si se usa martillo hidráulico u otro "
                 "método mecánico, no se requiere el permiso de voladura.", 8.5, cur=True)
    tit(doc, "7.3", "Análisis de trabajo seguro (ATS) y charla diaria", 2)
    par(doc, "Antes de iniciar cada tarea la cuadrilla llena el ATS: pasos de la tarea, peligros, riesgos y controles, firmado por todos sus integrantes y "
             "por el capataz. Cada jornada empieza con una charla de 5 minutos sobre un peligro de las tareas del día. El formato de ATS está en el Anexo 3.")

    # ================================================================ 8. EPP y protección colectiva
    tit(doc, "8", "Equipos de protección personal y colectiva")
    tit(doc, "8.1", "Equipos de protección personal", 2)
    par(doc, "El ejecutor entrega el EPP sin costo para el trabajador, certificado según las normas indicadas, en tallas adecuadas, y lo repone cuando se "
             "deteriora. La entrega se registra con la firma del trabajador. La matriz de EPP por puesto está en el Anexo 2 y en la hoja EPP del libro %s." % xls)
    ctab(doc, "Equipos de protección personal de la obra")
    tabla(doc, ["N.°", "EQUIPO DE PROTECCIÓN PERSONAL", "PUESTOS QUE LO USAN"],
          [[i + 1, v, len([1 for l in D["epp_puesto"].values() if k in l])] for i, (k, v) in enumerate(D["epp_catalogo"].items())],
          [1.0, 11.0, 4.0], tam=8, alinear=[C, L_, C])
    tit(doc, "8.2", "Equipos de protección colectiva", 2)
    epc = ["Señalización de seguridad, cintas y mallas de delimitación de áreas de trabajo.", "Extintores PQS en campamento, almacenes, equipos y frentes; extintor de CO₂ en tableros eléctricos.",
           "Botiquines por frente, camilla rígida con inmovilizadores y lavaojos portátil.", "Tableros eléctricos con interruptor diferencial y puesta a tierra."]
    if "altura" in petar: epc.append("Barandas perimetrales, líneas de vida, puntos de anclaje y redes de seguridad.")
    if "excavacion" in petar: epc.append("Entibados, escaleras de acceso y barandas al borde de las excavaciones.")
    if "cauces" in petar: epc.append("Aros salvavidas, cuerdas de rescate y atagüías en los trabajos en cauces.")
    if vial: epc.append("Dispositivos de control de tránsito: señales, conos, barreras, tranqueras, paletas y luces.")
    if "maquinaria" in cond: epc.append("Alarmas de retroceso, luces estroboscópicas y estructuras ROPS/FOPS en maquinaria.")
    if "selva" in cond: epc.append("Mallas mosquiteras y control de vectores en campamento.")
    for t in epc: vin(doc, t)

    # ================================================================ 9. capacitación
    tit(doc, "9", "Capacitación y sensibilización")
    par(doc, "Todo trabajador recibe la inducción antes de ingresar a la obra y cuatro capacitaciones al año como mínimo (Ley N.° 29783), dentro de la jornada "
             "de trabajo. Las capacitaciones específicas de los trabajos de alto riesgo son requisito para firmar el PETAR. La asistencia se registra en el "
             "registro de inducción, capacitación, entrenamiento y simulacros de emergencia.")
    ctab(doc, "Programa de capacitación")
    tabla(doc, ["N.°", "TEMA", "DIRIGIDO A", "HORAS", "FRECUENCIA"],
          [[i + 1, c["tema"], c["dirigido"], c["horas"], c["frecuencia"]] for i, c in enumerate(D["capacitacion"])], [0.8, 7.6, 4.2, 1.2, 2.2], tam=7.5,
          alinear=[C, L_, L_, C, C])
    fig(doc, os.path.join(FIG, "f8_capacitacion.png"), 16.0, "Cronograma de capacitaciones", "programa de capacitación (hoja CAPACITACION del libro %s)." % xls)

    # ================================================================ 10. inspecciones
    tit(doc, "10", "Inspecciones, observaciones y auditorías")
    par(doc, "Las inspecciones verifican que los controles del IPERC se cumplen. Cada hallazgo se registra con su responsable y plazo de corrección, y se "
             "verifica su cierre. Las condiciones de riesgo inminente se corrigen de inmediato y, si no es posible, se detiene el trabajo.")
    ctab(doc, "Programa de inspecciones")
    tabla(doc, ["N.°", "INSPECCIÓN", "RESPONSABLE", "FRECUENCIA"], [[i + 1, x["tema"], x["responsable"], x["frecuencia"]] for i, x in enumerate(D["inspecciones"])],
          [0.8, 7.6, 4.6, 3.0], tam=7.5, alinear=[C, L_, L_, C])
    par(doc, "Se promueve el reporte de actos y condiciones inseguras por todos los trabajadores con tarjetas de observación. El plan se audita internamente "
             "una vez al año y, si la entidad lo dispone, por auditores externos registrados.")

    # ================================================================ 11. salud ocupacional
    tit(doc, "11", "Salud ocupacional y bienestar")
    tit(doc, "11.1", "Exámenes médicos ocupacionales y seguros", 2)
    for t in ["Examen médico ocupacional de ingreso, periódico y de retiro según la R.M. N.° 312-2011-MINSA, con los protocolos del puesto (audiometría para expuestos "
              "a ruido, espirometría para expuestos a polvo, evaluación para trabajos en altura, entre otros). Nadie ingresa a la obra sin el certificado de aptitud.",
              "Seguro Complementario de Trabajo de Riesgo (SCTR) de salud y pensión para todo el personal, incluidos subcontratistas, vigente desde el primer día.",
              "Vigilancia médica de los trabajadores con exposición a agentes del IPERC y registro de exámenes médicos ocupacionales."]: vin(doc, t)
    tit(doc, "11.2", "Monitoreo de agentes ocupacionales", 2)
    ctab(doc, "Monitoreo de agentes ocupacionales")
    tabla(doc, ["N.°", "AGENTE", "MÉTODO", "FRECUENCIA"], [[i + 1, m["agente"], m["metodo"], m["frecuencia"]] for i, m in enumerate(D["monitoreo"])],
          [0.8, 4.8, 7.0, 3.4], tam=7.5, alinear=[C, L_, L_, C])
    tit(doc, "11.3", "Bienestar y primeros auxilios", 2)
    bien = ["Agua potable o tratada en todos los frentes y en el campamento.", "Servicios higiénicos y lavamanos en proporción al número de trabajadores; servicios "
            "portátiles en frentes alejados, con limpieza diaria.", "Comedor techado con control sanitario de alimentos.",
            "Tópico en el campamento y botiquín en cada frente y vehículo; unidad de transporte disponible para emergencias durante toda la jornada.",
            "Protección contra la radiación solar (Ley N.° 30102): bloqueador, sombra, hidratación y pausas."]
    if "selva" in cond: bien.append("Prevención de enfermedades metaxénicas y coordinación con el establecimiento de salud sobre la disponibilidad de suero antiofídico.")
    if "altura" in cond: bien.append("Aclimatación del personal que llega de zonas bajas y evaluación médica para trabajos en altura geográfica.")
    for t in bien: vin(doc, t)

    # ================================================================ 12. subcontratistas
    tit(doc, "12", "Subcontratistas y proveedores")
    for t in ["Antes de ingresar presentan: relación del personal, SCTR, certificados de aptitud médica, EPP, PETS de sus trabajos y certificados de operatividad de sus equipos.",
              "Asisten a la inducción y cumplen este plan, el IPERC y los PETAR de la obra; designan un responsable de SST cuando su personal lo requiera.",
              "Reportan de inmediato todo accidente o incidente; su desempeño en SST se evalúa mensualmente y su incumplimiento puede motivar el retiro de la obra.",
              "Los proveedores que solo entregan materiales cumplen las reglas de circulación, velocidad, estacionamiento y descarga en obra."]: vin(doc, t)

    # ================================================================ 13. emergencias
    tit(doc, "13", "Plan de respuesta ante emergencias")
    tit(doc, "13.1", "Escenarios de emergencia", 2)
    esc = [("Accidente con lesión grave", "Caídas, golpes, atrapamientos, cortes", "Primeros auxilios, inmovilización y evacuación al establecimiento de salud")]
    esc.append(("Accidente de tránsito o volcadura", "Vehículos y maquinaria en vías y taludes", "Señalización del área, rescate, primeros auxilios, aviso a la PNP"))
    esc.append(("Incendio", "Combustibles, asfalto, instalaciones eléctricas, soldadura", "Uso de extintores, corte de energía, evacuación, aviso a bomberos"))
    esc.append(("Sismo", "Evento natural", "Alejarse de taludes y estructuras, ir a zonas seguras, conteo del personal, inspección posterior"))
    if "selva" in cond or "cauces" in cond:
        esc.append(("Lluvias intensas, crecidas e inundaciones", "Temporada de lluvias, cauces y quebradas", "Alerta temprana, suspensión de trabajos, evacuación a zonas altas"))
    if acts & {"MTMQ", "ROCA", "EXCA"}:
        esc.append(("Deslizamiento, caída de rocas o derrumbe de zanja", "Taludes inestables, excavaciones, lluvias", "Evacuación del área, rescate con equipo y personal entrenado, no ingresar sin evaluar"))
    if "selva" in cond:
        esc.append(("Mordedura de ofidio o picadura", "Trabajo en vegetación", "Inmovilizar la extremidad, no hacer cortes ni torniquetes, traslado inmediato para suero"))
    if "maquinaria" in cond:
        esc.append(("Derrame de combustibles o asfalto", "Abastecimiento y mantenimiento de equipos, imprimación", "Contención con kit antiderrame, limpieza y disposición como residuo peligroso"))
    if "electrico" in petar:
        esc.append(("Contacto eléctrico", "Líneas aéreas, tableros y equipos", "Cortar la energía antes de tocar a la víctima, reanimación y evacuación"))
    if "cauces" in petar:
        esc.append(("Caída al agua", "Trabajos en cauces y puentes", "Lanzar aro salvavidas o cuerda, no ingresar al agua sin equipo, rescate desde la orilla"))
    if "altura" in petar:
        esc.append(("Trabajador suspendido en el arnés", "Caída detenida en altura", "Rescate en menos de 15 minutos para evitar el trauma por suspensión"))
    if "confinado" in petar:
        esc.append(("Emergencia en espacio confinado", "Atmósfera peligrosa", "Rescate con trípode y equipo de respiración; nunca ingresar sin protección"))
    ctab(doc, "Escenarios de emergencia y respuesta")
    tabla(doc, ["ESCENARIO", "ORIGEN", "RESPUESTA"], [list(e) for e in esc], [4.0, 4.6, 7.4], tam=8)
    tit(doc, "13.2", "Brigadas y equipos de emergencia", 2)
    par(doc, "Se forman tres brigadas con trabajadores voluntarios capacitados, con al menos un brigadista de cada tipo en cada frente y turno: primeros "
             "auxilios; lucha contra incendios; y evacuación y rescate. Cada frente cuenta con botiquín, camilla rígida, extintor, radio o teléfono con "
             "cobertura y un vehículo disponible para la evacuación.")
    fig(doc, os.path.join(FIG, "f7_emergencia.png"), 15.5, "Secuencia de respuesta ante una emergencia", "plan de respuesta ante emergencias de la obra.")
    tit(doc, "13.3", "Directorio de emergencia", 2)
    ctab(doc, "Directorio de emergencia")
    tabla(doc, ["INSTITUCIÓN", "UBICACIÓN", "TELÉFONO"], [[e[0], e[1] or "", e[2] or "__________"] for e in D["emergencia"]], [8.0, 5.0, 3.0], tam=8,
          alinear=[L_, L_, C])
    par(doc, "El directorio se confirma y se publica en el campamento y en cada vehículo antes del inicio de la obra. Se realiza al menos un simulacro por "
             "semestre, alternando los escenarios.", 8.5, cur=True)

    # ================================================================ 14. investigación
    tit(doc, "14", "Notificación e investigación de accidentes e incidentes")
    ctab(doc, "Plazos de notificación")
    tabla(doc, ["EVENTO", "QUIÉN NOTIFICA", "PLAZO"],
          [["Accidente de trabajo mortal", "Empleador, al Ministerio de Trabajo y Promoción del Empleo", "Dentro de las 24 horas"],
           ["Incidente peligroso", "Empleador, al Ministerio de Trabajo y Promoción del Empleo", "Dentro de las 24 horas"],
           ["Accidente de trabajo no mortal", "Centro médico asistencial que atiende al trabajador", "Hasta el último día hábil del mes siguiente"],
           ["Todo accidente o incidente en obra", "Trabajador o testigo, al supervisor y al responsable de SST", "De inmediato"]], [4.6, 7.4, 4.0], tam=8)
    par(doc, "Todo accidente e incidente peligroso se investiga con el método de árbol de causas o de análisis de causa raíz, con participación del comité "
             "de SST, para identificar las causas inmediatas y básicas y definir medidas correctivas, que se incorporan al IPERC. Los resultados se difunden a "
             "todo el personal.")

    # ================================================================ 15. registros e indicadores
    tit(doc, "15", "Registros, indicadores y mejora continua")
    ctab(doc, "Registros obligatorios del sistema de gestión (D.S. N.° 005-2012-TR, art. 33; formatos de la R.M. N.° 050-2013-TR)")
    tabla(doc, ["N.°", "REGISTRO"], [[i + 1, t] for i, t in enumerate([
        "Accidentes de trabajo, enfermedades ocupacionales, incidentes peligrosos y otros incidentes, con la investigación y las medidas correctivas.",
        "Exámenes médicos ocupacionales.", "Monitoreo de agentes físicos, químicos, biológicos, psicosociales y factores de riesgo disergonómicos.",
        "Inspecciones internas de seguridad y salud en el trabajo.", "Estadísticas de seguridad y salud.", "Equipos de seguridad o emergencia.",
        "Inducción, capacitación, entrenamiento y simulacros de emergencia.", "Auditorías."])], [1.0, 15.0], tam=8, alinear=[C, L_])
    par(doc, "Los indicadores se calculan cada mes en la hoja INDICADORES del libro %s:" % xls)
    for t in ["Índice de frecuencia (IF) = accidentes incapacitantes x 1,000,000 / horas hombre trabajadas.",
              "Índice de gravedad (IG) = días perdidos x 1,000,000 / horas hombre trabajadas.",
              "Índice de accidentabilidad (IA) = IF x IG / 1,000."]: vin(doc, t)
    par(doc, "El residente y el responsable de SST revisan el plan cada trimestre con los resultados de las inspecciones, los indicadores y las "
             "investigaciones, y lo actualizan cuando cambian las condiciones de la obra.")

    # ================================================================ 16. recursos
    tit(doc, "16", "Recursos del plan en el presupuesto de obra")
    par(doc, "El presupuesto de obra incluye las siguientes partidas para implementar el plan (metrados del expediente):")
    ctab(doc, "Partidas de seguridad y salud en el trabajo del presupuesto")
    tabla(doc, ["ÍTEM", "PARTIDA", "UND.", "METRADO"], [[p["item"], p["desc"], p["und"], "{:,.2f}".format(p["metrado"])] for p in D["partidas_sst"]],
          [2.0, 10.2, 1.6, 2.2], tam=8, alinear=[C, L_, C, C])
    par(doc, "Los demás controles de ingeniería del IPERC (entibados, barandas, guardas, mantenimiento de equipos) forman parte de los costos de cada "
             "partida y de los gastos generales del ejecutor.", 8.5, cur=True)
    par(doc, "\n\n_______________________________          _______________________________\nResponsable de SST de la obra                       Residente de obra\n"
             "CIP: __________                                           CIP: __________", 9, al=C, antes=18)

    # ================================================================ anexos (horizontal)
    seccion(doc, True)
    tit(doc, "17", "Anexos")
    tit(doc, "17.1", "Anexo 1: Matriz IPERC de línea base", 2)
    par(doc, "A: personas expuestas; B: procedimientos; C: capacitación; D: exposición; IP: índice de probabilidad; S: severidad; NR: nivel de riesgo. "
             "Ing.: controles de ingeniería; Adm.: controles administrativos. La matriz completa, con las columnas de eliminación y sustitución y con fórmulas, "
             "está en el libro %s." % xls, 8, cur=True)
    filas = []
    for f in R:
        ctr = "; ".join(("Sust.: " + f["sust"]) if k == "sust" else ("Ing.: " + f["ing"]) if k == "ing" else ("Adm.: " + f["adm"]) if k == "adm" else ("EPP: " + f["epp"])
                        for k in ("sust", "ing", "adm", "epp") if f[k] != N)
        filas.append([f["n"], f["actividad"], f["peligro"], "%s. %s" % (f["riesgo"], f["consecuencia"]), f["tipo"],
                      "%d-%d-%d-%d" % (f["A"], f["B"], f["C"], f["D"]), f["IP"], f["S"], f["NR"], f["nivel"], ctr, f["IPr"], f["Sr"], f["NRr"], f["nivel_r"]])
    tabla(doc, ["N.°", "ACTIVIDAD", "PELIGRO", "RIESGO Y CONSECUENCIA", "TIPO", "A-B-C-D", "IP", "S", "NR", "NIVEL", "MEDIDAS DE CONTROL", "IP", "S", "NR", "NIVEL RESIDUAL"],
          filas, [0.9, 2.6, 2.7, 3.0, 0.9, 1.3, 0.7, 0.6, 0.8, 1.6, 6.7, 0.7, 0.6, 0.8, 1.6], tam=6,
          colores=lambda i, j, v: D["color"].get(v) if j in (9, 14) else None, alinear=[C, L_, L_, L_, C, C, C, C, C, C, L_, C, C, C, C])
    tit(doc, "17.2", "Anexo 2: Matriz de equipos de protección personal por puesto", 2)
    keys = list(D["epp_catalogo"]); corto_epp = D["epp_corto"]
    tabla(doc, ["PUESTO"] + [corto_epp[k] for k in keys], [[pu] + ["X" if k in l else "" for k in keys] for pu, l in D["epp_puesto"].items()],
          [3.8] + [round(21.4 / len(keys), 2)] * len(keys), tam=6.5, alinear=[L_] + [C] * len(keys),
          colores=lambda i, j, v: "DDEBF7" if v == "X" else None)
    tit(doc, "17.3", "Anexo 3: Formato de análisis de trabajo seguro (ATS)", 2)
    tabla(doc, ["DATO", "", "DATO", ""], [["Frente / sector", "", "Fecha y hora", ""], ["Tarea", "", "Supervisor / capataz", ""],
                                          ["PETAR requerido", "Sí ( )  No ( )   N.° ______", "EPP verificado", "Sí ( )  No ( )"]], [4.0, 8.4, 4.0, 8.4], tam=8)
    tabla(doc, ["N.°", "PASOS DE LA TAREA", "PELIGROS", "RIESGOS", "MEDIDAS DE CONTROL"], [[i, "", "", "", ""] for i in range(1, 7)], [1.0, 6.0, 5.4, 5.4, 7.0], tam=9)
    tabla(doc, ["N.°", "NOMBRE DEL TRABAJADOR", "DNI", "FIRMA", "N.°", "NOMBRE DEL TRABAJADOR", "DNI", "FIRMA"],
          [[i, "", "", "", i + 5, "", "", ""] for i in range(1, 6)], [0.8, 5.4, 2.2, 3.4, 0.8, 5.4, 2.2, 4.6], tam=9)
    tit(doc, "17.4", "Anexo 4: Lista de verificación para el PETAR", 2)
    filas = []
    for k, v in petar.items():
        chk = {"altura": "Arnés y línea de anclaje inspeccionados; puntos de anclaje; barandas; andamio con tarjeta verde; personal apto y capacitado; área inferior delimitada",
               "excavacion": "Talud o entibado; material a 0.60 m del borde; escaleras; redes ubicadas; evaluación después de lluvias; barandas y señalización",
               "caliente": "Retiro de combustibles; extintor PQS; vigía contra incendio; equipo de soldadura en buen estado; biombos",
               "confinado": "Medición de O₂, LEL, CO y H₂S; ventilación; vigía; trípode y plan de rescate; comunicación",
               "izaje": "Plan de izaje; certificado de grúa y aparejos; terreno firme; zona de exclusión; rigger certificado; viento menor del límite",
               "electrico": "Corte, bloqueo y etiquetado; prueba de ausencia de tensión; puesta a tierra; distancias de seguridad; guantes dieléctricos",
               "explosivos": "Autorización SUCAMEC; disparador autorizado; plan de voladura; cierre de vía y evacuación; señales sonoras",
               "cauces": "Nivel del cauce y pronóstico; chalecos salvavidas; aro y cuerda; vigía en la orilla; atagüía"}[k]
        filas.append([v["tipo"], chk, "Sí ( )  No ( )"])
    tabla(doc, ["TRABAJO", "VERIFICACIONES ANTES DE AUTORIZAR", "CONFORME"], filas, [6.0, 15.4, 3.2], tam=8)
    tit(doc, "17.5", "Anexo 5: Documentos que acompañan el plan", 2)
    vin(doc, "%s: datos de la obra, criterios, IPERC con fórmulas, resumen, mapa de calor, matriz de EPP, PETS y PETAR, programas de capacitación, "
             "inspecciones y monitoreo, indicadores, relación de partidas por actividad y directorio de emergencia." % xls)

    # ================================================================ contenido, encabezados y pies
    from docx.shared import Cm as _Cm
    for num, (t, nv) in CONT.items():
        p = doc.add_paragraph(); p.paragraph_format.left_indent = _Cm(0 if nv == 1 else 0.8); p.paragraph_format.space_after = Pt(2)
        fuente(p.add_run("%s  %s" % (num, t.upper() if nv == 1 else t)), 10 if nv == 1 else 9.5, nv == 1)
        marca._p.addprevious(p._p)
    marca._p.getparent().remove(marca._p)
    for i, s in enumerate(doc.sections):
        pie(s, "Plan de Seguridad y Salud en el Trabajo - CUI N.° %s" % cui)
        encabezado(s, "PLAN DE SEGURIDAD Y SALUD EN EL TRABAJO  |  " + corto)
    doc.save(sal); print(sal, "tablas", NUMT[0], "figuras", NUMF[0])
    return sal


if __name__ == "__main__":
    main(sys.argv[1])
