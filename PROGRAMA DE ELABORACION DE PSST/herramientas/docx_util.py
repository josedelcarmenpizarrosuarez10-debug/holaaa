# -*- coding: utf-8 -*-
"""Utilidades comunes para los documentos Word del PSST (estilo, títulos numerados, tablas de ancho fijo, figuras, pies)."""
import os, sys, json, collections
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FT = "Arial"; AZUL = RGBColor(0x1F, 0x4E, 0x79)


# ------------------------------------------------------------------ utilidades
def fuente(r, tam=10, neg=False, cur=False, color=None):
    r.font.name = FT; r.font.size = Pt(tam); r.bold = neg; r.italic = cur
    if color: r.font.color.rgb = color
    rpr = r._r.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    if rf is None: rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"): rf.set(qn(k), FT)
    return r


def par(doc, texto="", tam=10, neg=False, al=WD_ALIGN_PARAGRAPH.JUSTIFY, antes=0, despues=6, cur=False, color=None, sangria=0):
    p = doc.add_paragraph(); p.paragraph_format.alignment = al; p.paragraph_format.space_before = Pt(antes); p.paragraph_format.space_after = Pt(despues)
    p.paragraph_format.line_spacing = 1.15
    if sangria: p.paragraph_format.left_indent = Cm(sangria)
    if texto: fuente(p.add_run(texto), tam, neg, cur, color)
    return p


def mixto(doc, partes, tam=10, al=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = par(doc, al=al)
    for t, n in partes: fuente(p.add_run(t), tam, n)
    return p


def vin(doc, texto, nivel=0):
    p = par(doc, despues=3, sangria=0.9 + 0.6 * nivel); p.paragraph_format.first_line_indent = Cm(-0.45)
    fuente(p.add_run(("•  " if nivel == 0 else "–  ") + texto), 10); return p


CONT = collections.OrderedDict()


def tit(doc, num, texto, nivel=1):
    p = doc.add_paragraph(style="Heading %d" % nivel); p.paragraph_format.space_before = Pt(12 if nivel == 1 else 8); p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    fuente(p.add_run(("%s. " % num if nivel == 1 else "%s " % num) + texto.upper() if nivel == 1 else "%s %s" % (num, texto)), 12 if nivel == 1 else 10.5, True, color=AZUL)
    CONT[num] = (texto, nivel)
    return p


def sombra(c, color):
    tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), color); tcPr.append(sh)


def tabla(doc, cab, filas, anchos, tam=8, colores=None, alinear=None):
    t = doc.add_table(rows=1, cols=len(cab)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(cab):
        c = t.rows[0].cells[j]; c.text = ""; fuente(c.paragraphs[0].add_run(h), tam, True, color=RGBColor(255, 255, 255)); sombra(c, "1F4E79")
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.rows[0]._tr; trPr = tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    for i, fila in enumerate(filas):
        cells = t.add_row().cells
        for j, v in enumerate(fila):
            cells[j].text = ""; pp = cells[j].paragraphs[0]; fuente(pp.add_run(str(v)), tam)
            pp.alignment = (alinear[j] if alinear else WD_ALIGN_PARAGRAPH.LEFT)
            if colores and colores(i, j, v): sombra(cells[j], colores(i, j, v))
    t.autofit = False
    tblPr = t._tbl.tblPr; lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    tw = OxmlElement("w:tblW"); tw.set(qn("w:w"), str(int(sum(anchos) * 567))); tw.set(qn("w:type"), "dxa")
    for e in tblPr.findall(qn("w:tblW")): tblPr.remove(e)
    tblPr.append(tw)
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), anchos): gc.set(qn("w:w"), str(int(w * 567)))
    for row in t.rows:
        for j, w in enumerate(anchos): row.cells[j].width = Cm(w)
    par(doc, despues=4)
    return t


NUMT = [0]; NUMF = [0]


def ctab(doc, texto):
    NUMT[0] += 1; p = par(doc, "Tabla N.° %02d: %s" % (NUMT[0], texto), 9, True, WD_ALIGN_PARAGRAPH.CENTER, antes=4, despues=3)
    p.paragraph_format.keep_with_next = True


def fig(doc, ruta, ancho, texto, fuente_txt):
    if not os.path.exists(ruta): return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(ruta, width=Cm(ancho))
    NUMF[0] += 1; par(doc, "Figura N.° %02d: %s" % (NUMF[0], texto), 9, True, WD_ALIGN_PARAGRAPH.CENTER, despues=1)
    par(doc, "Fuente: " + fuente_txt, 8, cur=True, al=WD_ALIGN_PARAGRAPH.CENTER, despues=8)


C = WD_ALIGN_PARAGRAPH.CENTER; L_ = WD_ALIGN_PARAGRAPH.LEFT


def seccion(doc, horizontal):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    w, h = Cm(21.0), Cm(29.7)
    if horizontal: s.orientation = WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = h, w
    else: s.orientation = WD_ORIENT.PORTRAIT; s.page_width, s.page_height = w, h
    s.left_margin = s.right_margin = Cm(2.0); s.top_margin = Cm(2.0); s.bottom_margin = Cm(1.8)
    return s


def pie(sec, texto):
    sec.footer.is_linked_to_previous = False
    p = sec.footer.paragraphs[0]; p.text = ""; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fuente(p.add_run(texto + "   |   Página "), 7.5, color=RGBColor(0x59, 0x59, 0x59))
    r = p.add_run(); fuente(r, 7.5)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag); r._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; r._r.append(e)



def encabezado(sec, texto):
    sec.header.is_linked_to_previous = False
    p = sec.header.paragraphs[0]; p.text = ""; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fuente(p.add_run(texto), 7.5, color=RGBColor(0x59, 0x59, 0x59))
    pPr = p._p.get_or_add_pPr(); bd = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F4E79")): b.set(qn(k), v)
    bd.append(b); pPr.append(bd)


def recuadro(doc, texto, fill="DDEBF7", borde="1F4E79", tam=9.5, neg=False, ancho=16.0):
    """Texto dentro de un recuadro sombreado (tabla de una celda)."""
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.width = Cm(ancho); sombra(c, fill)
    tcPr = c._tc.get_or_add_tcPr(); bd = OxmlElement("w:tcBorders")
    for lado in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + lado); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "8"); e.set(qn("w:color"), borde); bd.append(e)
    tcPr.append(bd)
    lineas = texto if isinstance(texto, list) else [texto]
    for i, ln in enumerate(lineas):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph()
        p.paragraph_format.space_after = Pt(3); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        fuente(p.add_run(ln), tam, neg)
    par(doc, despues=4)
    return t


def banda(doc, texto, fill="1F4E79", tam=20, ancho=16.0, alto=1.6):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.width = Cm(ancho); sombra(c, fill)
    tr = t.rows[0]._tr; trPr = tr.get_or_add_trPr(); h = OxmlElement("w:trHeight"); h.set(qn("w:val"), str(int(alto * 567))); trPr.append(h)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(8)
    fuente(p.add_run(texto), tam, True, color=RGBColor(255, 255, 255))
    return t
