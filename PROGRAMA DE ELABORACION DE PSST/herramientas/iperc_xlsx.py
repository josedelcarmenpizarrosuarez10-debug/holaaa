# -*- coding: utf-8 -*-
"""Libro Excel del PSST: matriz IPERC con fórmulas (R.M. N.° 050-2013-TR), mapa de calor, EPP por puesto, PETS y PETAR,
programas de capacitación, inspecciones y monitoreo, indicadores y contactos de emergencia.
Celda amarilla = dato que se puede cambiar; celda verde = fórmula.

Uso:  python3 iperc_xlsx.py <carpeta del proyecto>      (antes: nucleo.py)"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L

AZUL = "1F4E79"; AZUL2 = "DDEBF7"; AMA = "FFF2CC"; VER = "E2EFDA"; GRIS = "F2F2F2"
_b = Side(style="thin", color="8C8C8C"); BOR = Border(left=_b, right=_b, top=_b, bottom=_b)
F = lambda **k: Font(name="Arial", size=k.pop("size", 8), **k)


def cab(ws, fila, col, texto, ancho=None, alto=None, fill=AZUL, color="FFFFFF", span=1):
    c = ws.cell(fila, col, texto); c.font = F(bold=True, color=color, size=8); c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOR
    if span > 1:
        ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + span - 1)
        for j in range(col, col + span): ws.cell(fila, j).border = BOR
    if ancho: ws.column_dimensions[L(col)].width = ancho
    if alto: ws.row_dimensions[fila].height = alto
    return c


def celda(ws, f, c, v, fill=None, neg=False, al="left", fmt=None, size=8):
    x = ws.cell(f, c, v); x.font = F(bold=neg, size=size); x.border = BOR
    x.alignment = Alignment(horizontal=al, vertical="center", wrap_text=True)
    if fill: x.fill = PatternFill("solid", fgColor=fill)
    if fmt: x.number_format = fmt
    return x


def titulo(ws, texto, sub, ncol):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol); ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws.cell(1, 1, texto).font = F(bold=True, size=12, color=AZUL); ws.cell(2, 1, sub).font = F(size=8, italic=True)
    ws.row_dimensions[2].height = 24; ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")


def formato_nivel(ws, rango, primera):
    for n, col in (("Trivial", "A9D08E"), ("Tolerable", "C6E0B4"), ("Moderado", "FFD966"), ("Importante", "F4B183"), ("Intolerable", "FF7C80")):
        ws.conditional_formatting.add(rango, FormulaRule(formula=['%s="%s"' % (primera, n)], fill=PatternFill("solid", fgColor=col), font=Font(bold=True)))


NIVEL_F = 'IF({0}<=4,"Trivial",IF({0}<=8,"Tolerable",IF({0}<=16,"Moderado",IF({0}<=24,"Importante","Intolerable"))))'


def main(carpeta):
    D = json.load(open(os.path.join(carpeta, "entregables", "_calc", "psst.json")))
    P = D["proyecto"]; R = D["iperc"]
    sal = os.path.join(carpeta, "entregables", "MATRIZ_IPERC_PSST_%s.xlsx" % P.get("codigo", "OBRA"))
    wb = Workbook()
    sub = "%s - CUI N.° %s" % (P.get("nombre", ""), P.get("cui", ""))

    # ---------------------------------------------------------------- ÍNDICE
    ws = wb.active; ws.title = "INDICE"; titulo(ws, "PLAN DE SEGURIDAD Y SALUD EN EL TRABAJO - MATRIZ IPERC Y PROGRAMAS", sub, 3)
    hojas = [("DATOS", "Datos generales de la obra"), ("CRITERIOS", "Criterios de evaluación de la R.M. N.° 050-2013-TR"),
             ("IPERC", "Matriz de identificación de peligros, evaluación de riesgos y medidas de control (con fórmulas)"),
             ("RESUMEN", "Resumen de niveles de riesgo por proceso (fórmulas sobre el IPERC)"), ("MAPA DE CALOR", "Probabilidad x severidad antes y después de los controles"),
             ("EPP", "Matriz de equipos de protección personal por puesto"), ("PETS-PETAR", "Procedimientos escritos de trabajo seguro y permisos de trabajo de alto riesgo"),
             ("CAPACITACION", "Programa de capacitación y sensibilización"), ("INSPECCIONES", "Programa de inspecciones"),
             ("MONITOREO", "Monitoreo de agentes ocupacionales"), ("INDICADORES", "Objetivos, metas e indicadores de SST (fórmulas)"),
             ("PARTIDAS", "Relación de partidas del presupuesto y actividades evaluadas"), ("EMERGENCIA", "Directorio de emergencia")]
    cab(ws, 4, 1, "N.°", 6); cab(ws, 4, 2, "HOJA", 18); cab(ws, 4, 3, "CONTENIDO", 90)
    for i, (h, d) in enumerate(hojas, 5):
        celda(ws, i, 1, i - 4, al="center"); c = celda(ws, i, 2, h, neg=True); c.hyperlink = "#'%s'!A1" % h; c.font = F(bold=True, color="0563C1", underline="single")
        celda(ws, i, 3, d)
    celda(ws, 19, 2, "Leyenda:", neg=True); celda(ws, 19, 3, "Celda amarilla = dato que se puede cambiar; celda verde = fórmula", fill=AMA)

    # ---------------------------------------------------------------- DATOS
    ws = wb.create_sheet("DATOS"); titulo(ws, "DATOS GENERALES DE LA OBRA", sub, 2)
    filas = [("Proyecto", P.get("nombre")), ("CUI", P.get("cui")), ("Entidad", P.get("entidad")), ("Unidad ejecutora", P.get("unidad_ejecutora")),
             ("Ubicación", P.get("ubicacion")), ("Tipo de obra", P.get("tipo_obra")), ("Descripción", P.get("descripcion")),
             ("Plazo de ejecución (días calendario)", P.get("plazo_dias")), ("Modalidad", P.get("modalidad")), ("Altitud (msnm)", P.get("altitud_msnm")),
             ("Clima", P.get("clima")), ("Jornada de trabajo", P.get("jornada")), ("Personal máximo previsto", D["pico"]),
             ("Horas hombre previstas", D["hht"]), ("Órgano de SST", "Comité de SST (20 o más trabajadores)" if D["organizacion"] == "comite" else "Supervisor de SST"),
             ("Actividades evaluadas", D["resumen"]["n_act"]), ("Peligros evaluados", D["resumen"]["n"])]
    cab(ws, 4, 1, "CONCEPTO", 34); cab(ws, 4, 2, "VALOR", 110)
    for i, (a, b) in enumerate(filas, 5):
        celda(ws, i, 1, a, neg=True); celda(ws, i, 2, b, fmt="#,##0" if isinstance(b, (int, float)) else None)

    # ---------------------------------------------------------------- CRITERIOS
    ws = wb.create_sheet("CRITERIOS"); titulo(ws, "CRITERIOS DE EVALUACIÓN DE RIESGOS (R.M. N.° 050-2013-TR)", "Probabilidad = A + B + C + D; nivel de riesgo = probabilidad x severidad", 6)
    for j, (t, w) in enumerate([("ÍNDICE", 8), ("A. PERSONAS EXPUESTAS", 22), ("B. PROCEDIMIENTOS EXISTENTES", 34), ("C. CAPACITACIÓN", 34),
                                ("D. EXPOSICIÓN AL RIESGO", 30), ("SEVERIDAD", 34)], 1): cab(ws, 4, j, t, w, 30)
    crit = [(1, "1 a 3", "Existen, son satisfactorios y suficientes", "Personal entrenado; conoce el peligro y lo previene",
             "Al menos una vez al año (S) / esporádicamente (SO)", "Lesión sin incapacidad (S) / disconfort o incomodidad (SO)"),
            (2, "4 a 12", "Existen parcialmente y no son satisfactorios o suficientes", "Personal parcialmente entrenado; conoce el peligro pero no toma acciones de control",
             "Al menos una vez al mes (S) / eventualmente (SO)", "Lesión con incapacidad temporal (S) / daño a la salud reversible (SO)"),
            (3, "Más de 12", "No existen", "Personal no entrenado; no conoce el peligro ni toma acciones de control",
             "Al menos una vez al día (S) / permanentemente (SO)", "Lesión con incapacidad permanente (S) / daño a la salud irreversible (SO)")]
    for i, f in enumerate(crit, 5):
        for j, v in enumerate(f, 1): celda(ws, i, j, v, al="center" if j == 1 else "left")
        ws.row_dimensions[i].height = 32
    cab(ws, 9, 1, "NIVEL", span=2); cab(ws, 9, 3, "RANGO (P x S)"); cab(ws, 9, 4, "INTERPRETACIÓN", span=3)
    for i, (n, a, b) in enumerate(D["niveles"], 10):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2); ws.merge_cells(start_row=i, start_column=4, end_row=i, end_column=6)
        celda(ws, i, 1, n, fill=D["color"][n], neg=True, al="center"); celda(ws, i, 3, "%d" % a if a == b else "%d a %d" % (a, b), al="center")
        celda(ws, i, 4, D["interpretacion"][n]); ws.row_dimensions[i].height = 40
        for j in (2, 5, 6): ws.cell(i, j).border = BOR
    celda(ws, 16, 1, "Riesgo significativo: nivel Importante o Intolerable. Evaluación inicial con B = 2 y C = 2 (procedimientos y capacitación parciales); "
                     "evaluación residual con B = 1 y C = 1 (PETS, ATS, capacitación y supervisión implementados).", fill=GRIS)
    ws.merge_cells("A16:F16"); ws.row_dimensions[16].height = 30

    # ---------------------------------------------------------------- IPERC
    ws = wb.create_sheet("IPERC"); titulo(ws, "MATRIZ DE IDENTIFICACIÓN DE PELIGROS, EVALUACIÓN DE RIESGOS Y MEDIDAS DE CONTROL (IPERC LÍNEA BASE)", sub, 34)
    grupos = [(1, 9, "IDENTIFICACIÓN"), (10, 18, "EVALUACIÓN INICIAL"), (19, 23, "MEDIDAS DE CONTROL (JERARQUÍA)"), (24, 32, "EVALUACIÓN RESIDUAL"), (33, 34, "PROCEDIMIENTOS")]
    for a, b, t in grupos: cab(ws, 4, a, t, span=b - a + 1, fill="2E75B6")
    cols = [("N.°", 5), ("PROCESO", 14), ("ACTIVIDAD", 22), ("PUESTOS", 16), ("R / NR", 5), ("PELIGRO", 24), ("RIESGO", 24), ("CONSECUENCIA", 18), ("S / SO", 5),
            ("A\nPERSONAS", 6), ("B\nPROCED.", 6), ("C\nCAPACIT.", 6), ("D\nEXPOSIC.", 6), ("PROBAB.\nA+B+C+D", 7), ("SEVERI-\nDAD", 6), ("P x S", 6), ("NIVEL", 11), ("SIGNIFI-\nCATIVO", 7),
            ("ELIMINACIÓN", 14), ("SUSTITUCIÓN", 16), ("CONTROLES DE INGENIERÍA", 28), ("CONTROLES ADMINISTRATIVOS", 34), ("EPP", 24),
            ("A", 5), ("B", 5), ("C", 5), ("D", 5), ("PROBAB.", 7), ("SEVERI-\nDAD", 6), ("P x S", 6), ("NIVEL", 11), ("SIGNIFI-\nCATIVO", 7),
            ("PETS", 26), ("PETAR", 22)]
    for j, (t, w) in enumerate(cols, 1): cab(ws, 5, j, t, w, 36)
    r0 = 6
    for i, f in enumerate(R):
        r = r0 + i
        vals = [f["n"], f["proceso"], f["actividad"], ", ".join(f["puestos"]), f["rutina"], f["peligro"], f["riesgo"], f["consecuencia"], f["tipo"]]
        for j, v in enumerate(vals, 1): celda(ws, r, j, v, al="center" if j in (1, 5, 9) else "left")
        for j, k in zip((10, 11, 12, 13), "ABCD"): celda(ws, r, j, f[k], fill=AMA, al="center")
        celda(ws, r, 14, "=SUM(J%d:M%d)" % (r, r), fill=VER, al="center"); celda(ws, r, 15, f["S"], fill=AMA, al="center")
        celda(ws, r, 16, "=N%d*O%d" % (r, r), fill=VER, al="center"); celda(ws, r, 17, "=" + NIVEL_F.format("P%d" % r), al="center")
        celda(ws, r, 18, '=IF(OR(Q{0}="Importante",Q{0}="Intolerable"),"Sí","No")'.format(r), fill=VER, al="center")
        for j, k in zip(range(19, 24), ("elim", "sust", "ing", "adm", "epp")): celda(ws, r, j, f[k])
        for j, k in zip((24, 25, 26, 27), ("Ar", "Br", "Cr", "Dr")): celda(ws, r, j, f[k], fill=AMA, al="center")
        celda(ws, r, 28, "=SUM(X%d:AA%d)" % (r, r), fill=VER, al="center"); celda(ws, r, 29, f["Sr"], fill=AMA, al="center")
        celda(ws, r, 30, "=AB%d*AC%d" % (r, r), fill=VER, al="center"); celda(ws, r, 31, "=" + NIVEL_F.format("AD%d" % r), al="center")
        celda(ws, r, 32, '=IF(OR(AE{0}="Importante",AE{0}="Intolerable"),"Sí","No")'.format(r), fill=VER, al="center")
        celda(ws, r, 33, f["pets"]); celda(ws, r, 34, f["petar"] or "—", al="center")
        ws.row_dimensions[r].height = 58
    rf = r0 + len(R) - 1
    formato_nivel(ws, "Q%d:Q%d" % (r0, rf), "Q%d" % r0); formato_nivel(ws, "AE%d:AE%d" % (r0, rf), "AE%d" % r0)
    ws.freeze_panes = "G6"; ws.auto_filter.ref = "A5:AH%d" % rf
    ws.sheet_view.zoomScale = 85
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A3; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = "4:5"
    RNG = lambda col: "IPERC!$%s$%d:$%s$%d" % (col, r0, col, rf)

    # ---------------------------------------------------------------- RESUMEN
    ws = wb.create_sheet("RESUMEN"); titulo(ws, "RESUMEN DE NIVELES DE RIESGO", "Fórmulas sobre la hoja IPERC: se actualiza al cambiar los índices", 12)
    niv = [n for n, a, b in D["niveles"]]
    cab(ws, 4, 1, "PROCESO", 30); cab(ws, 4, 2, "PELIGROS", 9)
    for j, n in enumerate(niv): cab(ws, 4, 3 + j, "INICIAL\n" + n.upper(), 11, 30, fill=D["color"][n], color="000000")
    for j, n in enumerate(niv): cab(ws, 4, 8 + j, "RESIDUAL\n" + n.upper(), 11, fill=D["color"][n], color="000000")
    procesos = list(dict.fromkeys(f["proceso"] for f in R))
    for i, p in enumerate(procesos, 5):
        celda(ws, i, 1, p, neg=True); celda(ws, i, 2, '=COUNTIF(%s,A%d)' % (RNG("B"), i), fill=VER, al="center")
        for j, n in enumerate(niv):
            celda(ws, i, 3 + j, '=COUNTIFS(%s,$A%d,%s,"%s")' % (RNG("B"), i, RNG("Q"), n), fill=VER, al="center")
            celda(ws, i, 8 + j, '=COUNTIFS(%s,$A%d,%s,"%s")' % (RNG("B"), i, RNG("AE"), n), fill=VER, al="center")
    t = 5 + len(procesos)
    celda(ws, t, 1, "TOTAL", neg=True, fill=AZUL2)
    for j in range(2, 13): celda(ws, t, j, "=SUM(%s5:%s%d)" % (L(j), L(j), t - 1), neg=True, fill=AZUL2, al="center")
    celda(ws, t + 2, 1, "Riesgos significativos antes de los controles", neg=True); celda(ws, t + 2, 2, '=COUNTIF(%s,"Sí")' % RNG("R"), fill=VER, al="center")
    celda(ws, t + 3, 1, "Riesgos significativos después de los controles", neg=True); celda(ws, t + 3, 2, '=COUNTIF(%s,"Sí")' % RNG("AF"), fill=VER, al="center")

    # ---------------------------------------------------------------- MAPA DE CALOR
    ws = wb.create_sheet("MAPA DE CALOR"); titulo(ws, "MAPA DE CALOR: PROBABILIDAD x SEVERIDAD", "Número de peligros en cada celda (fórmulas sobre el IPERC)", 9)
    for k, (col_ip, col_s, tit, c0) in enumerate((("N", "O", "EVALUACIÓN INICIAL", 1), ("AB", "AC", "EVALUACIÓN RESIDUAL", 6))):
        cab(ws, 4, c0, tit, span=4)
        cab(ws, 5, c0, "PROBABILIDAD", 13); [cab(ws, 5, c0 + s, "SEVERIDAD %d" % s, 12) for s in (1, 2, 3)]
        for i, ip in enumerate(range(12, 3, -1), 6):
            celda(ws, i, c0, ip, neg=True, al="center")
            for s in (1, 2, 3):
                n = ip * s; nv = [x for x, a, b in D["niveles"] if a <= n <= b][0]
                celda(ws, i, c0 + s, '=COUNTIFS(%s,%d,%s,%d)' % (RNG(col_ip), ip, RNG(col_s), s), fill=D["color"][nv], al="center", neg=True)
    ws.column_dimensions["E"].width = 4

    # ---------------------------------------------------------------- EPP
    ws = wb.create_sheet("EPP"); titulo(ws, "MATRIZ DE EQUIPOS DE PROTECCIÓN PERSONAL POR PUESTO", "X = uso obligatorio; según las actividades del IPERC", 2 + len(D["epp_catalogo"]))
    cab(ws, 4, 1, "PUESTO", 28)
    keys = list(D["epp_catalogo"])
    for j, k in enumerate(keys, 2): cab(ws, 4, j, D["epp_catalogo"][k], 13, 110)
    for i, (pu, lst) in enumerate(D["epp_puesto"].items(), 5):
        celda(ws, i, 1, pu, neg=True)
        for j, k in enumerate(keys, 2): celda(ws, i, j, "X" if k in lst else "", al="center", neg=True, fill=AZUL2 if k in lst else None)
    ws.freeze_panes = "B5"

    # ---------------------------------------------------------------- PETS-PETAR
    ws = wb.create_sheet("PETS-PETAR"); titulo(ws, "PROCEDIMIENTOS ESCRITOS DE TRABAJO SEGURO Y PERMISOS DE TRABAJO DE ALTO RIESGO", sub, 4)
    for j, (t, w) in enumerate([("CÓDIGO", 10), ("PROCEDIMIENTO", 55), ("ACTIVIDAD", 50), ("REQUIERE PETAR", 50)], 1): cab(ws, 4, j, t, w)
    for i, p in enumerate(D["pets"], 5):
        celda(ws, i, 1, p["cod"], al="center"); celda(ws, i, 2, p["nombre"]); celda(ws, i, 3, p["actividad"]); celda(ws, i, 4, "; ".join(p["petar"]) or "No")
    i0 = 6 + len(D["pets"])
    for j, (t, w) in enumerate([("N.°", 10), ("PETAR", 55), ("ACTIVIDADES", 50)], 1): cab(ws, i0, j, t)
    for i, (k, v) in enumerate(D["petar"].items(), i0 + 1):
        celda(ws, i, 1, i - i0, al="center"); celda(ws, i, 2, v["tipo"]); celda(ws, i, 3, "; ".join(v["actividades"]))

    # ---------------------------------------------------------------- CAPACITACIÓN
    m = D["meses"]
    ws = wb.create_sheet("CAPACITACION"); titulo(ws, "PROGRAMA DE CAPACITACIÓN Y SENSIBILIZACIÓN EN SST", "X = mes programado; total de horas por fórmula", 6 + m)
    for j, (t, w) in enumerate([("N.°", 5), ("TEMA", 50), ("DIRIGIDO A", 26), ("HORAS", 7), ("FRECUENCIA", 11)], 1): cab(ws, 4, j, t, w, 30)
    for k in range(m): cab(ws, 4, 6 + k, "M%d" % (k + 1), 4.2)
    cab(ws, 4, 6 + m, "VECES", 7); cab(ws, 4, 7 + m, "HORAS TOTAL", 8)
    for i, c in enumerate(D["capacitacion"], 5):
        celda(ws, i, 1, i - 4, al="center"); celda(ws, i, 2, c["tema"]); celda(ws, i, 3, c["dirigido"]); celda(ws, i, 4, c["horas"], fill=AMA, al="center")
        celda(ws, i, 5, c["frecuencia"], al="center")
        for k in range(m): celda(ws, i, 6 + k, "X" if (k + 1) in c["meses"] else "", al="center", fill=AMA)
        celda(ws, i, 6 + m, '=COUNTIF(%s%d:%s%d,"X")' % (L(6), i, L(5 + m), i), fill=VER, al="center")
        celda(ws, i, 7 + m, "=D%d*%s%d" % (i, L(6 + m), i), fill=VER, al="center")
    ws.freeze_panes = "F5"

    # ---------------------------------------------------------------- INSPECCIONES
    ws = wb.create_sheet("INSPECCIONES"); titulo(ws, "PROGRAMA DE INSPECCIONES DE SST", sub, 4)
    for j, (t, w) in enumerate([("N.°", 5), ("INSPECCIÓN", 60), ("RESPONSABLE", 40), ("FRECUENCIA", 28)], 1): cab(ws, 4, j, t, w)
    for i, x in enumerate(D["inspecciones"], 5):
        celda(ws, i, 1, i - 4, al="center"); celda(ws, i, 2, x["tema"]); celda(ws, i, 3, x["responsable"]); celda(ws, i, 4, x["frecuencia"], al="center")

    # ---------------------------------------------------------------- MONITOREO
    ws = wb.create_sheet("MONITOREO"); titulo(ws, "MONITOREO DE AGENTES OCUPACIONALES", "Según los peligros de salud ocupacional del IPERC", 4)
    for j, (t, w) in enumerate([("N.°", 5), ("AGENTE", 40), ("MÉTODO", 60), ("FRECUENCIA", 30)], 1): cab(ws, 4, j, t, w)
    for i, x in enumerate(D["monitoreo"], 5):
        celda(ws, i, 1, i - 4, al="center"); celda(ws, i, 2, x["agente"]); celda(ws, i, 3, x["metodo"]); celda(ws, i, 4, x["frecuencia"], al="center")

    # ---------------------------------------------------------------- INDICADORES
    ws = wb.create_sheet("INDICADORES"); titulo(ws, "OBJETIVOS, METAS E INDICADORES DE SST", "Registrar cada mes los accidentes y días perdidos (celdas amarillas)", 10)
    for j, (t, w) in enumerate([("MES", 7), ("TRABAJADORES", 13), ("HORAS HOMBRE TRABAJADAS", 14), ("ACCIDENTES INCAPACITANTES", 14), ("DÍAS PERDIDOS", 11),
                                ("ÍNDICE DE FRECUENCIA", 12), ("ÍNDICE DE GRAVEDAD", 12), ("ÍNDICE DE ACCIDENTABILIDAD", 14), ("CAPACITACIONES EJECUTADAS (%)", 14),
                                ("INSPECCIONES EJECUTADAS (%)", 14)], 1): cab(ws, 4, j, t, w, 42)
    for i, (mes, tr, _e) in enumerate(D["personal"], 5):
        celda(ws, i, 1, mes, al="center"); celda(ws, i, 2, tr, fill=AMA, al="center"); celda(ws, i, 3, "=B%d*8*26" % i, fill=VER, al="center", fmt="#,##0")
        celda(ws, i, 4, 0, fill=AMA, al="center"); celda(ws, i, 5, 0, fill=AMA, al="center")
        celda(ws, i, 6, "=IF(C%d=0,0,D%d*1000000/C%d)" % (i, i, i), fill=VER, al="center", fmt="0.00")
        celda(ws, i, 7, "=IF(C%d=0,0,E%d*1000000/C%d)" % (i, i, i), fill=VER, al="center", fmt="0.00")
        celda(ws, i, 8, "=F%d*G%d/1000" % (i, i), fill=VER, al="center", fmt="0.00")
        celda(ws, i, 9, 1, fill=AMA, al="center", fmt="0%"); celda(ws, i, 10, 1, fill=AMA, al="center", fmt="0%")
    t = 5 + len(D["personal"])
    celda(ws, t, 1, "TOTAL", neg=True, fill=AZUL2)
    for j, frm in ((2, "=MAX(B5:B%d)" % (t - 1)), (3, "=SUM(C5:C%d)" % (t - 1)), (4, "=SUM(D5:D%d)" % (t - 1)), (5, "=SUM(E5:E%d)" % (t - 1)),
                   (6, "=IF(C%d=0,0,D%d*1000000/C%d)" % (t, t, t)), (7, "=IF(C%d=0,0,E%d*1000000/C%d)" % (t, t, t)), (8, "=F%d*G%d/1000" % (t, t)),
                   (9, "=AVERAGE(I5:I%d)" % (t - 1)), (10, "=AVERAGE(J5:J%d)" % (t - 1))):
        celda(ws, t, j, frm, neg=True, fill=AZUL2, al="center", fmt="0%" if j >= 9 else ("#,##0" if j <= 5 else "0.00"))
    celda(ws, t + 2, 1, "Índice de frecuencia = accidentes incapacitantes x 1,000,000 / horas hombre trabajadas. Índice de gravedad = días perdidos x 1,000,000 / horas hombre "
                        "trabajadas. Índice de accidentabilidad = IF x IG / 1,000. Horas hombre = trabajadores x 8 h x 26 días.", fill=GRIS)
    ws.merge_cells(start_row=t + 2, start_column=1, end_row=t + 2, end_column=10); ws.row_dimensions[t + 2].height = 30

    # ---------------------------------------------------------------- PARTIDAS
    ws = wb.create_sheet("PARTIDAS"); titulo(ws, "PARTIDAS DEL PRESUPUESTO Y ACTIVIDADES EVALUADAS EN EL IPERC", "Trazabilidad entre el presupuesto y el IPERC", 3)
    cab(ws, 4, 1, "CÓDIGO", 10); cab(ws, 4, 2, "ACTIVIDAD DEL IPERC", 50); cab(ws, 4, 3, "PARTIDAS DEL PRESUPUESTO", 110)
    i = 5
    for a in D["actividades"]:
        celda(ws, i, 1, a["cod"], al="center"); celda(ws, i, 2, a["nombre"], neg=True)
        celda(ws, i, 3, ("%d partidas: " % len(a["partidas"]) + "; ".join(a["partidas"])) if a["partidas"] else "Actividad transversal a toda la obra")
        ws.row_dimensions[i].height = min(400, 14 + 11 * (len("; ".join(a["partidas"])) // 120)); i += 1

    # ---------------------------------------------------------------- EMERGENCIA
    ws = wb.create_sheet("EMERGENCIA"); titulo(ws, "DIRECTORIO DE EMERGENCIA", "Completar o confirmar los teléfonos locales antes del inicio de la obra", 3)
    cab(ws, 4, 1, "INSTITUCIÓN", 60); cab(ws, 4, 2, "UBICACIÓN", 40); cab(ws, 4, 3, "TELÉFONO", 18)
    for i, e in enumerate(D["emergencia"], 5):
        celda(ws, i, 1, e[0]); celda(ws, i, 2, e[1]); celda(ws, i, 3, e[2] or "", fill=AMA, al="center")

    for w in wb.worksheets: w.sheet_view.showGridLines = False
    wb.save(sal); print(sal)
    return sal


if __name__ == "__main__":
    main(sys.argv[1])
